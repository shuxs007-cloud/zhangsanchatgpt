#!/usr/bin/env python3
"""Build and consume one trusted, hash-bound scan cache for a fresh OpenSMS export."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import pickle
import tempfile
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CACHE_FORMAT = "opensms-export-scan-cache-v1"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_pickle(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=f".{path.name}.", dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
        pickle.dump(payload, handle, protocol=pickle.HIGHEST_PROTOCOL)
        handle.flush()
        os.fsync(handle.fileno())
    try:
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def atomic_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    encoded = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
    with tempfile.NamedTemporaryFile(
        prefix=f".{path.name}.", dir=path.parent, delete=False, mode="w", encoding="utf-8"
    ) as handle:
        temporary = Path(handle.name)
        handle.write(encoded)
        handle.flush()
        os.fsync(handle.fileno())
    try:
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def build_scan_cache(source: Path, output: Path, receipt: Path | None = None) -> dict[str, Any]:
    started = time.perf_counter()
    source = source.expanduser().resolve()
    output = output.expanduser().resolve()
    if not source.is_file():
        raise RuntimeError(f"OpenSMS源表不存在: {source}")
    source_hash = sha256(source)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        row_iterator = worksheet.iter_rows(values_only=True)
        try:
            headers = tuple(next(row_iterator))
        except StopIteration as exc:
            raise RuntimeError("OpenSMS源表为空") from exc
        rows = [tuple(row) for row in row_iterator]
        sheet_title = worksheet.title
    finally:
        workbook.close()
    payload = {
        "format": CACHE_FORMAT,
        "sourcePath": str(source),
        "sourceSha256": source_hash,
        "sourceSize": source.stat().st_size,
        "sheetTitle": sheet_title,
        "headers": headers,
        "rows": rows,
    }
    atomic_pickle(output, payload)
    verified = load_scan_cache(output, expected_source=source)
    result = {
        "status": "passed",
        "format": CACHE_FORMAT,
        "source": str(source),
        "sourceSha256": source_hash,
        "cache": str(output),
        "cacheSha256": sha256(output),
        "rowCount": len(verified["rows"]),
        "columnCount": len(verified["headers"]),
        "sheetTitle": verified["sheetTitle"],
        "durationSeconds": round(time.perf_counter() - started, 3),
        "hashBound": True,
        "candidateWorkbookReopenRequired": False,
    }
    if receipt is not None:
        receipt = receipt.expanduser().resolve()
        result["receipt"] = str(receipt)
        atomic_json(receipt, result)
    return result


def load_scan_cache(path: Path | str, expected_source: Path | str | None = None) -> dict[str, Any]:
    cache_path = Path(path).expanduser().resolve()
    if not cache_path.is_file():
        raise RuntimeError(f"OpenSMS扫描缓存不存在: {cache_path}")
    with cache_path.open("rb") as handle:
        payload = pickle.load(handle)
    if not isinstance(payload, dict) or payload.get("format") != CACHE_FORMAT:
        raise RuntimeError("OpenSMS扫描缓存格式不受支持")
    if not isinstance(payload.get("headers"), tuple) or not isinstance(payload.get("rows"), list):
        raise RuntimeError("OpenSMS扫描缓存结构损坏")
    if expected_source is not None:
        source = Path(expected_source).expanduser().resolve()
        if not source.is_file():
            raise RuntimeError(f"用于核验扫描缓存的源表不存在: {source}")
        current_hash = sha256(source)
        if current_hash != payload.get("sourceSha256"):
            raise RuntimeError("OpenSMS扫描缓存与当前源表SHA-256不一致，禁止复用")
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description="为本轮OpenSMS新鲜源表生成单次扫描缓存")
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--receipt", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(build_scan_cache(args.source, args.output, args.receipt), ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
