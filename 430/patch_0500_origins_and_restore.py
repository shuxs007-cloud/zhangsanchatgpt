#!/usr/bin/env python3
import glob
import json
import os
import re
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

from openpyxl import load_workbook

from opensms_artifact_paths import SNAPSHOT_DIR

BASE = "/Users/zhangsan/Downloads/430"
SNAPSHOTS = str(SNAPSHOT_DIR)
EXPORT_PATH = os.path.join(BASE, "export4day.xlsx")
ZERO_PATH = os.path.join(BASE, "opensms_zero_all_20260721T050336+0800_candidates.json")
SUMMARY_PATH = os.path.join(BASE, "opensms_current_analysis_export4day_export4day_20260721T050304+0800.json")
TAG = "20260721T050304+0800"
RESTORE_PATH = os.path.join(BASE, f"opensms_restore_e_replied_{TAG}_candidates.json")

E_CARRIER_NAMES = {"E-T休息区", "E-C休息区"}
UNKNOWN = "归属未识别/E区历史"


def norm_phone(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"\D", "", s)


def norm_slot(v):
    return "" if v is None else str(v).strip()


def key_for(row):
    return f"{norm_phone(row.get('电话') or row.get('phone'))}|{norm_slot(row.get('卡槽'))}"


def good_origin(v):
    if not v:
        return False
    s = str(v).strip()
    return bool(
        s
        and "休息" not in s
        and "换号" not in s
        and "未识别" not in s
    )


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def build_origin_map(keys=None, workers=4):
    """Read only needed identities, merging newest first, never completion order.

    None preserves the historical full-map API. Empty keys do not even enumerate
    history. Bounded batches drain before returning and have no write side effects.
    """
    if type(workers) is not int or not 1 <= workers <= 4:
        raise ValueError("历史归属读取线程数必须为1至4")
    wanted = None if keys is None else set(keys)
    if wanted == set():
        return {}
    search_dirs = (SNAPSHOTS, BASE)
    paths = []
    for directory in search_dirs:
        paths += glob.glob(os.path.join(directory, "opensms_zero_all_*_candidates.json"))
        paths += glob.glob(os.path.join(directory, "opensms_zero_e_gte40_a1_d1_d3u8_a11_*_candidates.json"))
        paths += glob.glob(os.path.join(directory, "opensms_validated_*_candidates.json"))
        paths += glob.glob(os.path.join(directory, "opensms_e_origin_backup_*.json"))
        paths += glob.glob(os.path.join(directory, "opensms_zero_opr_gte50_chatr_to_e_*_candidates.json"))
        paths += glob.glob(os.path.join(directory, "opensms_zero_b11_b21_gte50_to_e_*_candidates.json"))
        paths += glob.glob(os.path.join(directory, "opensms_restore_e_replied_*_candidates.json"))
    paths = [p for p in paths if os.path.abspath(p) != os.path.abspath(ZERO_PATH)]
    paths.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    def read_origins(path):
        result = {}
        try:
            data = load_json(path)
        except Exception:
            return result
        basename = os.path.basename(path)
        basename_lower = basename.lower()
        if basename.startswith("opensms_e_origin_backup_"):
            source_type = "e_origin_backup_json"
        elif "restore_e_replied" in basename_lower:
            source_type = "restore_candidate_json"
        elif basename.startswith("opensms_validated_"):
            source_type = "validated_e_transfer_snapshot_json"
        elif "chatr_to_e" in basename_lower or "b11_b21_gte50_to_e" in basename_lower:
            source_type = "e_transfer_snapshot_json"
        elif "zero_e_gte40_a1_d1_d3u8_a11" in basename_lower:
            source_type = "e_transfer_snapshot_json"
        else:
            source_type = "history_candidate_json"
        for row in data.get("rows", []):
            if basename.startswith("opensms_validated_"):
                target = (
                    row.get("目标运营商")
                    or row.get("targetCarrierName")
                    or data.get("targetCarrierName")
                    or (data.get("target") or {}).get("name")
                )
                if not str(target or "").strip().upper().startswith("E-"):
                    continue
            k = key_for(row)
            if not k or k in result or (wanted is not None and k not in wanted):
                continue
            origin = row.get("精确原运营商") or row.get("原运营商") or row.get("目标运营商") or row.get("targetCarrierName") or row.get("运营商")
            if good_origin(origin):
                result[k] = {"origin": str(origin).strip(), "source": path, "sourceType": source_type}
        return result

    origin_map = {}
    actual_workers = min(workers, os.cpu_count() or 1)
    with ThreadPoolExecutor(max_workers=actual_workers, thread_name_prefix="opensms-origin") as pool:
        for offset in range(0, len(paths), actual_workers):
            # map preserves input order even when newer files finish last.
            for result in pool.map(read_origins, paths[offset:offset + actual_workers]):
                for key, info in result.items():
                    origin_map.setdefault(key, info)
            if wanted is not None and wanted <= origin_map.keys():
                break
    return origin_map


def patch_zero(origin_map):
    data = load_json(ZERO_PATH)
    counts = Counter()
    missing = 0
    for row in data.get("rows", []):
        current = row.get("运营商")
        info = origin_map.get(key_for(row))
        if good_origin(current):
            origin = str(current).strip()
            row["归属识别来源"] = EXPORT_PATH
            row["归属识别来源类型"] = "current_carrier"
        elif info:
            origin = info["origin"]
            row["归属识别来源"] = info["source"]
            row["归属识别来源类型"] = info["sourceType"]
        else:
            origin = UNKNOWN
            row["归属识别来源"] = ""
            row["归属识别来源类型"] = "unidentified_e_history"
            missing += 1
        row["原运营商"] = origin
        counts[origin] += 1
    data["byOrigin"] = dict(counts)
    data["unidentifiedOriginCount"] = missing
    save_json(ZERO_PATH, data)
    return data


def as_int(v):
    if v in (None, ""):
        return 0
    try:
        return int(float(v))
    except Exception:
        return 0


def find_latest(headers):
    dates = [h[:5] for h in headers if re.fullmatch(r"\d{2}-\d{2}发送窗口数", h)]
    return sorted(dates)[-1]


def build_restore(origin_map):
    wb = load_workbook(EXPORT_PATH, read_only=True, data_only=True)
    ws = wb.active
    headers = {str(v).strip(): i for i, v in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))) if v is not None}
    latest = find_latest(headers)
    rows = []
    counts = Counter()
    missing = 0
    e_total = 0
    for vals in ws.iter_rows(min_row=2, values_only=True):
        def v(name):
            idx = headers.get(name)
            return vals[idx] if idx is not None and idx < len(vals) else None

        current_e_carrier = str(v("运营商") or "").strip()
        if current_e_carrier not in E_CARRIER_NAMES:
            continue
        e_total += 1
        send = as_int(v(f"{latest}发送窗口数"))
        reply = as_int(v(f"{latest}回复窗口数"))
        if reply <= 0:
            continue
        row = {
            "电话": norm_phone(v("电话")),
            "phone": norm_phone(v("电话")),
            "运营商": current_e_carrier,
            "opr": str(v("OPR") or v("opr") or "").strip(),
            "卡槽": norm_slot(v("卡槽")),
            "状态": str(v("状态") or "").strip(),
            "创建时间": str(v("创建时间") or "").strip(),
            "发送窗口数": send,
            "回复窗口数": reply,
            "latestDate": latest,
            "规则": "e_current_replied_reply_gt0_restore_origin",
        }
        info = origin_map.get(key_for(row))
        if info:
            target = info["origin"]
            row["归属识别来源"] = info["source"]
            row["归属识别来源类型"] = info["sourceType"]
        else:
            target = UNKNOWN
            row["归属识别来源"] = ""
            row["归属识别来源类型"] = "unidentified_e_history"
            missing += 1
        row["原运营商"] = target
        row["目标运营商"] = target
        row["targetCarrierName"] = target
        counts[target] += 1
        rows.append(row)
    wb.close()
    payload = {
        "createdAt": datetime.now().astimezone().isoformat(),
        "export": EXPORT_PATH,
        "latestDate": latest,
        "eTotal": e_total,
        "count": len(rows),
        "recoverableCount": sum(v for k, v in counts.items() if k != UNKNOWN),
        "unidentifiedOriginCount": missing,
        "byTarget": dict(counts),
        "rows": rows,
    }
    save_json(RESTORE_PATH, payload)
    for target, count in counts.items():
        if target == UNKNOWN:
            continue
        safe = re.sub(r"[^A-Za-z0-9_.+-]+", "_", target).strip("_")
        split = dict(payload)
        split["count"] = count
        split["recoverableCount"] = count
        split["unidentifiedOriginCount"] = 0
        split["targetCarrierName"] = target
        split["byTarget"] = {target: count}
        split["rows"] = [row for row in rows if row.get("目标运营商") == target]
        save_json(os.path.join(BASE, f"opensms_restore_e_replied_{TAG}_{safe}_candidates.json"), split)
    return payload


def patch_summary(zero_data):
    summary = load_json(SUMMARY_PATH)
    current_counts = Counter(row.get("原运营商") or UNKNOWN for row in zero_data.get("rows", []))
    today_rows = {}
    today_counts = Counter()
    day_prefix = TAG[:8]
    for path in sorted(glob.glob(os.path.join(BASE, f"opensms_zero_all_{day_prefix}T*_candidates.json"))):
        try:
            data = zero_data if os.path.abspath(path) == os.path.abspath(ZERO_PATH) else load_json(path)
        except Exception:
            continue
        for row in data.get("rows", []):
            k = key_for(row)
            if not k or k in today_rows:
                continue
            origin = row.get("原运营商")
            if not good_origin(origin):
                origin = UNKNOWN
            today_rows[k] = row
            today_counts[origin] += 1
    origin_rows = []
    for origin in sorted(today_counts, key=lambda k: (-today_counts[k], k)):
        total = today_counts[origin]
        current = current_counts.get(origin, 0)
        origin_rows.append({"carrier": origin, "today": total, "recovered": total - current, "current": current})
    summary["currentZero"] = sum(current_counts.values())
    summary["todayCumulative"] = sum(today_counts.values())
    summary["todayRecovered"] = summary["todayCumulative"] - summary["currentZero"]
    summary["origOperatorZero"] = origin_rows
    save_json(SUMMARY_PATH, summary)
    return summary


def main():
    origin_map = build_origin_map()
    zero = patch_zero(origin_map)
    restore = build_restore(origin_map)
    summary = patch_summary(zero)
    print(json.dumps({
        "zeroPath": ZERO_PATH,
        "zeroCount": zero.get("count"),
        "zeroByOrigin": zero.get("byOrigin"),
        "zeroUnidentified": zero.get("unidentifiedOriginCount"),
        "restorePath": RESTORE_PATH,
        "restoreCount": restore.get("count"),
        "restoreRecoverable": restore.get("recoverableCount"),
        "restoreUnidentified": restore.get("unidentifiedOriginCount"),
        "restoreByTarget": restore.get("byTarget"),
        "summaryPath": SUMMARY_PATH,
        "todayCumulative": summary.get("todayCumulative"),
        "todayRecovered": summary.get("todayRecovered"),
        "currentZero": summary.get("currentZero"),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
