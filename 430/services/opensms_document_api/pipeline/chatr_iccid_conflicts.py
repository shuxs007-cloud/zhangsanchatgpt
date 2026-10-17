"""Display-only contract for system/Google ICCID conflicts (no date decisions)."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = "ICCID冲突_系统为准"
DECISION = "系统ICCID为准；不使用冲突Google日期"
COLUMNS = (
    ("row", "主表行号"), ("phone", "电话"), ("slot", "卡槽"),
    ("system_iccid", "系统ICCID"), ("google_iccid", "Google ICCID（归一化）"),
    ("system_activation", "系统激活日期"), ("system_cutoff", "系统截止日"),
    ("google_activation", "Google激活日期（不采用）"), ("google_cutoff", "Google截止日（不采用）"),
    ("final_activation", "最终激活日期"), ("final_cutoff", "最终截止日"),
    ("activation_source", "激活日期来源"), ("cutoff_source", "截止日来源"),
    ("decision", "处理决定"),
)


def write_conflict_sheet(book, records):
    # One coordinator writes into the existing single-save workbook flow.
    if SHEET in book.sheetnames:
        del book[SHEET]
    sheet = book.create_sheet(SHEET)
    sheet.append([label for _, label in COLUMNS])
    for record in records:
        sheet.append([record[key] for key, _ in COLUMNS])
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Preserve ICCID precision; never execute source text.
                cell.number_format = "@"
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
            else:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for index, (key, _) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 46 if key == "decision" else 28 if "iccid" in key or "source" in key else 23
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def review_conflict_sheet(book, independently_expected):
    """Only compare display against independently recomputed source records."""
    if SHEET not in book.sheetnames:
        return False
    sheet = book[SHEET]
    expected = [[label for _, label in COLUMNS],
                *[[record[key] for key, _ in COLUMNS] for record in independently_expected]]
    actual = list(sheet.iter_rows())
    if len(actual) != len(expected) or sheet.max_column != len(COLUMNS):
        return False
    for cells, values in zip(actual, expected):
        for cell, value in zip(cells, values):
            if cell.value != value and not (cell.value is None and value == ""):
                return False
            if isinstance(value, str) and value and cell.data_type != "s":
                return False
    return True
