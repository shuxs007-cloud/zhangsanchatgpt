"""Main-sheet display column only; never changes routing or device identities."""
from copy import copy
import re
from openpyxl.utils import column_index_from_string, get_column_letter


def machine_number(slot):
    value = str(slot or '').strip()
    match = re.fullmatch(r'([A-Za-z0-9]+)_\d+(?:\.\d+)?', value)
    return match[1] if match else None


def add_machine_number(sheet):
    headers = [c.value for c in sheet[1]]
    if headers.count('卡槽') != 1 or '机器号码' in headers:
        raise ValueError('机器号码/卡槽列位置不唯一')
    column = headers.index('卡槽')+1
    dimensions = [(column_index_from_string(k),copy(v)) for k,v in sheet.column_dimensions.items()]
    sheet.insert_cols(column)
    sheet.column_dimensions.clear()
    for old, dimension in dimensions:
        new = old+(old>=column)
        dimension.index=get_column_letter(new)
        dimension.min=new
        dimension.max=(dimension.max or old)+((dimension.max or old)>=column)
        sheet.column_dimensions[dimension.index]=dimension
    sheet.column_dimensions[get_column_letter(column)].width=16
    unknown=0
    for row in sheet.iter_rows(min_col=column,max_col=column+1):
        cell,slot=row
        cell._style=copy(slot._style)
        cell.number_format='@'
        cell.value='机器号码' if cell.row==1 else machine_number(slot.value)
        if cell.row>1 and cell.value is None:
            unknown+=1
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref=sheet.dimensions
    return {'column':'机器号码','position':'卡槽左侧','rows':sheet.max_row-1,'unknown':unknown}


def verify_machine_number(headers, rows):
    if headers.count('机器号码')!=1 or headers.count('卡槽')!=1:
        raise ValueError('缺失或重复机器号码/卡槽列')
    m,s=headers.index('机器号码'),headers.index('卡槽')
    if m+1!=s:
        raise ValueError('机器号码必须紧邻卡槽左侧')
    for i,row in enumerate(rows,2):
        # Independently derive the display prefix, preserving leading zeros/letters.
        text=str(row[s] or '').strip()
        parts=text.split('_')
        valid=len(parts)==2 and bool(re.fullmatch('[A-Za-z0-9]+',parts[0])) and bool(re.fullmatch(r'\d+(?:\.\d+)?',parts[1]))
        expected=parts[0] if valid else None
        if row[m]!=expected:
            raise ValueError(f'第{i}行机器号码与卡槽不符')
    return {'matches':True,'rows':len(rows)}
