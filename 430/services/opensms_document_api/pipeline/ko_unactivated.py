"""Frozen Task F activation_missing cards present in this run's system inventory."""
from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import re

from openpyxl.styles import Alignment, Font, PatternFill

SHEET = 'ko系统有-未激活'
LEGACY_SHEET = 'ko未激活表'
AFTER_SHEET = '谷歌有-系统无'
SNAPSHOT_SCOPE = 'all_task_f_activation_missing'
SCOPE = 'task_f_activation_missing_in_current_system'
REPORT_SCOPE = 'task_f_activation_missing_in_10day_export'
REPORT_SLOT_SOURCE = 'current_10day_export.卡槽'
REPORT_RULE = 'F冻结库activation_missing与本轮10日表Kooda ICCID交集；仅列10日表和F库均无有效激活日期的卡；卡槽来自本轮10日表，不表示全平台库存或实时激活状态'
HEADERS = ('ICCID', '查询状态', '激活日期', '截止日', '最近记录时间', '号码（记录时）',
           '卡槽', 'OPR（记录时）', '运营商（记录时）', '购买日期', '激活申请时间', 'SKU', '数据来源')
FIELDS = ('queriedAt', 'phone', 'slot', 'opr', 'carrier', 'purchaseDate', 'actRequestDate', 'sku', 'dataSource')
DISPLAY_FIELDS = tuple('currentSystemSlot' if field == 'slot' else field for field in FIELDS)
SLOT_SOURCE = 'current_system_inventory.slotIdentif'
RULE = '任务F冻结库activation_missing与本轮完整系统Kooda库存ICCID交集；一ICCID一行，不限10日主表或上架状态；卡槽取系统当前slotIdentif，同ICCID多卡槽去重换行列全，空白不回填历史；仅表示F未取得激活日期，非实时激活确认'
COLOR = 'FFC7CE'


def load_snapshot(path, digest, run_id, required=False):
    path = Path(path)
    if not path.is_absolute() or not run_id or not re.fullmatch(r'[a-f0-9]{64}', digest or ''):
        raise ValueError('ko系统有-未激活缺少本轮冻结库路径、SHA256或runId')
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != digest:
        raise ValueError('ko系统有-未激活冻结库SHA256不一致')
    payload = json.loads(raw)
    section = payload.get('koUnactivated')
    if section is None and not required:
        return None  # Legacy input, never accepted by a new unified-run contract.
    if (payload.get('schemaVersion') != 1 or payload.get('status') != 'passed'
            or payload.get('runId') != run_id or payload.get('readOnly') is not True
            or payload.get('snapshotConsistent') is not True or not payload.get('createdAt')):
        raise ValueError('ko系统有-未激活必须绑定本轮一致的只读冻结库')
    if not isinstance(section, dict) or section.get('schemaVersion') != 1 or section.get('scope') != SNAPSHOT_SCOPE:
        raise ValueError('ko系统有-未激活缺少全部未激活快照')
    rows = section.get('records')
    if (not isinstance(rows, list) or type(section.get('recordCount')) is not int
            or section['recordCount'] != len(rows) or (not payload.get('sourceExists') and rows)):
        raise ValueError('ko系统有-未激活快照行数或数据库存在状态不一致')
    seen = set()
    valid_dates = {row.get('queryIccid') for row in payload.get('records', []) if isinstance(row, dict)}
    for row in rows:
        key = row.get('queryIccid') if isinstance(row, dict) else None
        if (not isinstance(key, str) or not re.fullmatch(r'[0-9]{18,22}', key) or key in seen
                or key in valid_dates or row.get('status') != 'activation_missing'
                or row.get('activationDate') is not None or row.get('cutoffDate') is not None
                or any(not isinstance(row.get(field), str) for field in FIELDS)):
            raise ValueError('ko系统有-未激活含非法、重复、状态或日期冲突记录')
        if any(len(row[field]) > 32767 for field in FIELDS):
            raise ValueError('ko系统有-未激活字段超过Excel长度上限，禁止截断')
        seen.add(key)
    if [r['queryIccid'] for r in rows] != sorted(seen):
        raise ValueError('ko系统有-未激活快照顺序不稳定')
    return {'rows': rows, 'proof': {'enabled': True, 'status': 'passed', 'scope': SNAPSHOT_SCOPE, 'sheet': SHEET,
        'run_id': run_id, 'snapshot_path': str(path), 'snapshot_sha256': digest,
        'frozen_at': payload['createdAt'], 'source_exists': payload.get('sourceExists'),
        'source_database': payload.get('sourceDatabase'), 'row_count': len(rows), 'rule': RULE}}


def select_system_records(evidence, system):
    """Reuse already authenticated raw rows; never infer presence from phone/slot."""
    if evidence is None:
        return None
    proof = evidence['proof']
    if (not system or system.get('run_id') != proof['run_id']
            or not Path(system.get('path') or '').is_absolute()
            or not re.fullmatch(r'[a-f0-9]{64}', system.get('sha256') or '')
            or not isinstance(system.get('raw_rows'), list)
            or type(system.get('total_rows')) is not int
            or len(system['raw_rows']) != system['total_rows']):
        raise ValueError('ko系统有-未激活缺少本轮完整系统库存，禁止回退全库')
    keys, invalid, slots_by_iccid = set(), 0, {}
    report_source = system.get('source_kind') == 'verified_10day_export'
    known = set()
    requested = {row['queryIccid'] for row in evidence['rows']}
    for row in system['raw_rows']:
        value = row.get('cardStore')
        key = re.sub(r'[Ff]$', '', value.strip()) if isinstance(value, str) else ''
        if re.fullmatch(r'[0-9]{18,22}', key):
            keys.add(key)
            if report_source and row.get('chatrActivateDate'):
                known.add(key)
            if key in requested:
                slot = row.get('slotIdentif')
                if slot is not None and not isinstance(slot, str):
                    raise ValueError('ko系统有-未激活的系统卡槽必须为文本或空值')
                if slot and slot.strip():
                    slots_by_iccid.setdefault(key, set()).add(slot)
        else:
            invalid += 1
    selected, filled, multiple = [], 0, 0
    for row in evidence['rows']:
        key = row['queryIccid']
        if key not in keys or key in known:
            continue
        slots = sorted(slots_by_iccid.get(key, set()))
        display = '\n'.join(slots)
        if len(display) > 32767:
            raise ValueError('ko系统有-未激活的系统卡槽超过Excel长度上限，禁止截断')
        selected.append({**row, 'currentSystemSlot': display})
        filled += bool(slots)
        multiple += len(slots) > 1
    return {'rows': selected, 'proof': {**proof, 'scope': REPORT_SCOPE if report_source else SCOPE,
        'rule': REPORT_RULE if report_source else RULE,
        'system_path': system['path'], 'system_sha256': system['sha256'],
        'system_row_count': system['total_rows'], 'system_unique_iccid_count': len(keys),
        'system_invalid_iccid_rows': invalid, 'f_snapshot_row_count': len(evidence['rows']),
        'slot_source': REPORT_SLOT_SOURCE if report_source else SLOT_SOURCE, 'current_slot_filled_count': filled,
        'current_slot_empty_count': len(selected) - filled, 'multiple_current_slot_count': multiple,
        'excluded_not_in_system_count': len(requested - keys),
        'excluded_activation_present_count': len(requested & known) if report_source else 0,
        'row_count': len(selected)}}


def load_from_env():
    return load_snapshot(os.environ.get('OPENSMS_REPORT_DATE_CACHE_PATH', ''),
        os.environ.get('OPENSMS_REPORT_DATE_CACHE_SHA256', ''), os.environ.get('OPENSMS_RUN_ID', ''),
        os.environ.get('OPENSMS_KO_UNACTIVATED_REQUIRED') == '1')


def write_sheet(book, evidence):
    if evidence is None:
        return None
    if (evidence['proof'].get('scope'), evidence['proof'].get('slot_source')) not in (
            (SCOPE, SLOT_SOURCE), (REPORT_SCOPE, REPORT_SLOT_SOURCE)):
        raise ValueError('ko系统有-未激活必须先按本轮系统ICCID取交集')
    for name in (SHEET, LEGACY_SHEET):
        if name in book.sheetnames:
            del book[name]
    position = book.sheetnames.index(AFTER_SHEET) + 1 if AFTER_SHEET in book.sheetnames else None
    ws = book.create_sheet(SHEET, position)
    ws.append(HEADERS)
    for number, row in enumerate(evidence['rows'], 2):
        values = [row['queryIccid'], '未激活（未取得激活日期）', None, None] + [row[field] for field in DISPLAY_FIELDS]
        for column, value in enumerate(values, 1):
            cell = ws.cell(number, column, value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.number_format = '@'
            cell.fill = PatternFill('solid', fgColor=COLOR)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column, width in zip('ABCDEFGHIJKLM', (26, 28, 14, 14, 30, 22, 20, 20, 26, 20, 25, 26, 34)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    summary = book['处理汇总']
    owned_keys = {name + suffix for name in (SHEET, LEGACY_SHEET) for suffix in ('行数', '口径')}
    for row in reversed(list(summary.iter_rows())):
        if row[0].value in owned_keys:
            summary.delete_rows(row[0].row)
    book['处理汇总'].append((SHEET + '行数', len(evidence['rows'])))
    book['处理汇总'].append((SHEET + '口径', evidence['proof']['rule']))
    return evidence['proof']


def review(book, proof, path, digest, run_id, required=False, system=None):
    evidence = load_snapshot(path, digest, run_id, required)
    if evidence is None:
        if proof or SHEET in book.sheetnames or LEGACY_SHEET in book.sheetnames:
            raise ValueError('ko系统有-未激活没有冻结来源')
        return None
    evidence = select_system_records(evidence, system)
    expected, errors = evidence['proof'], []
    if LEGACY_SHEET in book.sheetnames:
        errors.append('legacy_sheet_present')
    if proof != expected:
        errors.append('proof_mismatch')
    if SHEET not in book.sheetnames:
        errors.append('missing_sheet')
    else:
        ws = book[SHEET]
        if AFTER_SHEET in book.sheetnames and book.sheetnames.index(SHEET) != book.sheetnames.index(AFTER_SHEET) + 1:
            errors.append('sheet_order_mismatch')
        if tuple(c.value for c in next(ws.iter_rows(max_row=1))) != HEADERS:
            errors.append('headers_mismatch')
        actual = list(ws.iter_rows(min_row=2))
        if len(actual) != len(evidence['rows']):
            errors.append('row_count_mismatch')
        for cells, row in zip(actual, evidence['rows']):
            wanted = [row['queryIccid'], '未激活（未取得激活日期）', None, None]
            wanted.extend(row[key] or None for key in DISPLAY_FIELDS)
            for cell, value in zip(cells, wanted):
                blank_text = value is None and cell.value == '' and cell.data_type in ('s', 'inlineStr')
                if (cell.value != value and not blank_text) or (value is not None and cell.data_type != 's'):
                    errors.append('values_or_types_mismatch')
                if cell.number_format != '@':
                    errors.append('format_mismatch')
                if not (cell.fill and cell.fill.patternType == 'solid' and str(cell.fill.fgColor.rgb).endswith(COLOR)):
                    errors.append('fill_mismatch')
        summary = {r[0]: r[1] for r in book['处理汇总'].iter_rows(values_only=True)}
        if any(LEGACY_SHEET + suffix in summary for suffix in ('行数', '口径')):
            errors.append('legacy_summary_present')
        if summary.get(SHEET + '行数') != len(evidence['rows']) or summary.get(SHEET + '口径') != expected['rule']:
            errors.append('summary_mismatch')
    return {**expected, 'matches': not errors, 'errors': dict(Counter(errors))}
