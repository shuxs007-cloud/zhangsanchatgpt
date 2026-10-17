"""Task A card-retention evidence, request-time column and complete history sheets."""
from __future__ import annotations

from collections import Counter
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re

import copy

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill

FIELDS = ('cardNo', 'expireBefore', 'expireAfter', 'createTime', 'chatId', 'id')
HEADERS = ('卡号', '原到期日', '新到期日', '申请时间', '申请会话ID', '记录ID', '来源页码', '页内序号')


def short_phone(value):
    digits = re.sub(r'[^0-9]', '', str(value or ''))
    return digits[-10:]


def validate_source(source, run_id):
    if source.get('status') != 'passed' or (source.get('task') != 'A' or source.get('scope') != 'card_retention_report') or source.get('runId') != run_id:
        raise ValueError('任务A卡保留来源未通过或运行编号不符')
    if source.get('backendWrites') != 0 or source.get('readOnly') is not True:
        raise ValueError('任务A卡保留只读契约不符')
    if source.get('silentBrowserMode') is not True or source.get('foregroundActivationAttempted') is not False:
        raise ValueError('静默契约不符')
    if source.get('clientClosed') is not True or source.get('sharedTargetOwnedByRunner') is not True:
        raise ValueError('本轮卡保留API客户端清理未确认')
    passes = source.get('passes', [])
    if len(passes) != 2 or source.get('independentSourceReadbackPassed') is not True:
        raise ValueError('缺少独立全量回读')
    normalized = []
    for item in passes:
        pages = item.get('pages', [])
        if not pages:
            raise ValueError('缺少分页证据')
        total = pages[0]['response'].get('total')
        if isinstance(total, str) and re.fullmatch(r'[0-9]+', total):
            total = int(total)
        if type(total) is not int or total < 0:
            raise ValueError('无效总记录数')
        if len(pages) != max(1, (total + 199) // 200):
            raise ValueError('页数不完整')
        rows, ids = [], set()
        for number, page in enumerate(pages, 1):
            response = page.get('response', {})
            if page.get('pageNum') != number or page.get('pageSize') != 200 or response.get('code') != 200:
                raise ValueError('分页参数或业务状态不符')
            if str(response.get('total')) != str(total):
                raise ValueError('分页总数漂移')
            chunk = response.get('rows')
            if not isinstance(chunk, list) or len(chunk) != min(200, max(0, total - (number - 1) * 200)):
                raise ValueError('分页行数不符')
            for index, row in enumerate(chunk, 1):
                if not isinstance(row, dict) or not re.fullmatch(r'[0-9]{10}', short_phone(row.get('cardNo'))):
                    raise ValueError('无效卡号，未丢弃记录或猜测补齐')
                identity = str(row.get('id') or '').strip()
                if identity and identity in ids:
                    raise ValueError('重复后台记录ID')
                if identity:
                    ids.add(identity)
                rows.append((row, number, index))
        if len(rows) != total:
            raise ValueError('全量记录数不符')
        normalized.append(rows)
    if normalized[0] != normalized[1]:
        raise ValueError('两遍全量回读内容不一致')
    rows = normalized[0]
    phones = list(dict.fromkeys(short_phone(row['cardNo']) for row, _, _ in rows))
    retention = source.get('retention', {})
    if retention.get('totalRecords') != len(rows) or retention.get('fetchedRecords') != len(rows):
        raise ValueError('采集总数声明不符')
    if retention.get('cardNos') != phones or retention.get('uniqueCardCount') != len(phones):
        raise ValueError('去重号码清单不符')
    return rows, phones


def text_value(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(',', ':'))
    return str(value)


def append_text_row(sheet, values):
    if any(isinstance(value, str) and len(value) > 32767 for value in values):
        raise ValueError('字段超过Excel单元格限制；拒绝静默截断')
    sheet.append(list(values))
    for column in range(1, len(values) + 1):
        cell = sheet.cell(sheet._current_row, column)
        if isinstance(cell.value, str):
            cell.data_type = 's'
            cell.number_format = '@'


def expected_sheets(rows, phones):
    details = [HEADERS] + [tuple(text_value(row.get(key)) for key in FIELDS) + (page, index) for row, page, index in rows]
    counts = Counter(short_phone(row['cardNo']) for row, _, _ in rows)
    # Counts are audit observations, not a recalculating financial model.
    unique = [('卡号（去重，末10位）', '记录条数')] + [(phone, counts[phone]) for phone in phones]
    raw_keys = list(dict.fromkeys(key for row, _, _ in rows for key in row)) or list(FIELDS)
    raw = [tuple(raw_keys)] + [tuple(text_value(row.get(key)) for key in raw_keys) for row, _, _ in rows]
    return {'卡保留全部记录': details, '卡保留去重号码': unique, '卡保留原始字段': raw}


COLUMN = '需要保留时间'
SCOPE = 'card_retention_report'
MATCH_RULE = 'normalized_phone_last10_latest_createTime'


def load_source(path, digest, run_id):
    path = Path(path)
    if not path.is_absolute() or not re.fullmatch(r'[a-f0-9]{64}', digest or ''):
        raise ValueError('卡保留源路径或SHA256缺失')
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != digest:
        raise ValueError('卡保留源SHA256不一致')
    source = json.loads(raw)
    rows, phones = validate_source(source, run_id)
    return {'path': str(path), 'sha256': digest, 'run_id': run_id,
            'rows': rows, 'phones': phones, 'phone_set': set(phones), 'source': source}


def load_from_env():
    path = os.environ.get('OPENSMS_CARD_RETENTION_REPORT_PATH', '')
    digest = os.environ.get('OPENSMS_CARD_RETENTION_REPORT_SHA256', '')
    if not path and not digest and os.environ.get('OPENSMS_CARD_RETENTION_REPORT_REQUIRED') != '1':
        return None  # Historical offline builders do not fabricate a current snapshot.
    return load_source(path, digest, os.environ.get('OPENSMS_RUN_ID', ''))


def request_time(value):
    if not isinstance(value, str) or not re.fullmatch(r'20\d{2}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?', value.strip()):
        return None
    try:
        return datetime.fromisoformat(value.strip())
    except ValueError:
        return None


def latest_requests(rows):
    latest, unknown = {}, set()
    for raw, _, _ in rows:
        phone = short_phone(raw['cardNo'])
        instant = request_time(raw.get('createTime'))
        if instant is None:
            unknown.add(phone)
        elif phone not in latest or instant > latest[phone][0]:
            latest[phone] = (instant, raw['createTime'])
    return {phone: value for phone, (_, value) in latest.items() if phone not in unknown}, unknown


def add_retention_report(book, sheet, evidence):
    if evidence is None:
        return None
    headers = [cell.value for cell in sheet[1]]
    if headers.count('激活后天数') != 1 or headers.count('电话') != 1 or COLUMN in headers:
        raise ValueError('卡保留列位置或号码表头不唯一')
    column = headers.index('激活后天数') + 2
    dimensions = [(column_index_from_string(key), copy.copy(value)) for key, value in sheet.column_dimensions.items()]
    sheet.insert_cols(column)
    sheet.column_dimensions.clear()
    for old, dimension in dimensions:
        new = old + (old >= column)
        dimension.index = get_column_letter(new)
        dimension.min = new
        dimension.max = (dimension.max or old) + ((dimension.max or old) >= column)
        sheet.column_dimensions[dimension.index] = dimension
    sheet.column_dimensions[get_column_letter(column)].width = 24
    phone_column = headers.index('电话') + 1
    if phone_column >= column:
        phone_column += 1
    latest, unknown = latest_requests(evidence['rows'])
    matched, filled, matched_phones = 0, 0, set()
    # One workbook writer. Copy the row fill so red/yellow classifications remain whole-row.
    for row in range(1, sheet.max_row + 1):
        cell, neighbour = sheet.cell(row, column), sheet.cell(row, column - 1)
        cell._style = copy.copy(neighbour._style)
        cell.number_format = '@'
        cell.alignment = Alignment(vertical='center', wrap_text=False)
        if row == 1:
            cell.value = COLUMN
        else:
            phone = short_phone(sheet.cell(row, phone_column).value)
            if phone in evidence['phone_set']:
                matched += 1
                matched_phones.add(phone)
            value = latest.get(phone)
            if value is not None:
                cell.value = value
                cell.data_type = 's'
                filled += 1
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = sheet.dimensions
    expected = expected_sheets(evidence['rows'], evidence['phones'])
    for title, values in expected.items():
        if title in book.sheetnames:
            raise ValueError('拒绝覆盖已有卡保留历史sheet')
        detail = book.create_sheet(title)
        for values_row in values:
            append_text_row(detail, values_row)
        detail.freeze_panes = 'A2'
        detail.auto_filter.ref = detail.dimensions
        for cell in detail[1]:
            cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
            cell.fill = PatternFill('solid', fgColor='24537A')
        for index in range(1, detail.max_column + 1):
            detail.column_dimensions[get_column_letter(index)].width = 24
    result = {'enabled': True, 'status': 'passed', 'scope': SCOPE, 'run_id': evidence['run_id'],
              'source_path': evidence['path'], 'source_sha256': evidence['sha256'],
              'column': COLUMN, 'match_rule': MATCH_RULE,
              'record_count': len(evidence['rows']), 'unique_phone_count': len(evidence['phones']),
              'matched_rows': matched, 'filled_rows': filled,
              'matched_phone_count': len(matched_phones),
              'unmatched_phone_count': len(evidence['phones']) - len(matched_phones),
              'unknown_request_time_phone_count': len(unknown),
              'page_size': 200, 'page_count': len(evidence['source']['passes'][0]['pages']),
              'read_concurrency': evidence['source']['readConcurrency']['maxActive'],
              'backend_writes': 0, 'sheets': list(expected)}
    if '处理汇总' in book.sheetnames:
        for label, value in [('卡保留全部记录数', result['record_count']), ('卡保留去重号码数', result['unique_phone_count']),
                             ('需要保留时间已填行数', filled), ('需要保留时间来源', '卡保留记录createTime（申请时间），同号取最新；不是新到期日'),
                             ('卡保留记录独立全量回读', '两遍完整分页一致；无记录留空；不触发下卡通知')]:
            book['处理汇总'].append((label, value))
    return result


def review_retention_report(book, proof, main_headers=None, main_cells=None):
    if proof is None:
        if COLUMN in [c.value for c in book['sheet1'][1]] or os.environ.get('OPENSMS_CARD_RETENTION_REPORT_REQUIRED') == '1':
            raise ValueError('需要保留时间缺少绑定来源证据')
        return None
    source = load_from_env()
    if source is None or any(proof.get(k) != source[v] for k, v in
                             [('run_id','run_id'), ('source_path','path'), ('source_sha256','sha256')]):
        raise ValueError('卡保留独立核验缺少同轮源绑定')
    sheet = book['sheet1']
    headers = main_headers if main_headers is not None else [c.value for c in sheet[1]]
    errors = []
    prior = '使用天数' if '使用天数' in headers else '激活后天数'
    if headers.count(COLUMN) != 1 or headers.index(COLUMN) != headers.index(prior) + 1:
        raise ValueError('需要保留时间列位置不符')
    col, phone_col = headers.index(COLUMN), headers.index('电话')
    # Independently regroup and sort the original records, not the generator's lookup.
    grouped = {}
    for raw, _, _ in source['rows']:
        grouped.setdefault(short_phone(raw['cardNo']), []).append(raw.get('createTime'))
    expected = {}
    unknown = 0
    for phone, values in grouped.items():
        parsed = [(request_time(value), value) for value in values]
        if any(item[0] is None for item in parsed):
            unknown += 1
            continue
        expected[phone] = sorted(parsed, key=lambda item: item[0], reverse=True)[0][1]
    matched, filled, matched_phones = 0, 0, set()
    cells_to_review = main_cells if main_cells is not None else (
        (cells[phone_col], cells[col], cells[col - 1]) for cells in sheet.iter_rows(min_row=2))
    for row_number, (phone_cell, cell, prior_cell) in enumerate(cells_to_review, 2):
        phone = short_phone(phone_cell.value)
        wanted = expected.get(phone)
        if phone in grouped:
            matched += 1
            matched_phones.add(phone)
        if wanted is not None:
            filled += 1
        if (cell.value or None) != wanted or (wanted is not None and (cell.data_type != 's' or cell.number_format != '@')):
            errors.append(f'主表第{row_number}行申请时间或文本格式不一致')
        actual_fill = copy.copy(cell.fill) if cell.fill is not None and cell.fill.fill_type else None
        prior_fill = copy.copy(prior_cell.fill) if prior_cell.fill is not None and prior_cell.fill.fill_type else None
        if actual_fill != prior_fill:
            errors.append(f'主表第{row_number}行申请时间列填色不一致')
    for title, rows in expected_sheets(source['rows'], source['phones']).items():
        if title not in book.sheetnames:
            errors.append(f'缺少{title}')
            continue
        actual = list(book[title].iter_rows())
        if len(actual) != len(rows):
            errors.append(f'{title}行数不符')
            continue
        for cells, wanted in zip(actual, rows):
            if len(cells) != len(wanted) or any((c.value if c.value is not None else '') != value
                    or c.data_type in ('f', 'e') for c, value in zip(cells, wanted)):
                errors.append(f'{title}内容或类型不符')
                break
    metrics = {'record_count': len(source['rows']), 'unique_phone_count': len(grouped),
               'matched_rows': matched, 'filled_rows': filled, 'matched_phone_count': len(matched_phones),
               'unmatched_phone_count': len(grouped) - len(matched_phones),
               'unknown_request_time_phone_count': unknown}
    if any(proof.get(k) != v for k, v in metrics.items()) or proof.get('match_rule') != MATCH_RULE:
        errors.append('卡保留声明计数或匹配规则不一致')
    return {**proof, **metrics, 'matches': not errors, 'errors': errors}
