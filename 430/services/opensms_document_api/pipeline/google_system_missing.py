"""Read-only Google slot difference; never used for dates or carrier routing."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import unicodedata
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from .google_room_sheets import bound_room_sheets, select_room_sheets
except ImportError:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from google_room_sheets import bound_room_sheets, select_room_sheets

SHEET = "谷歌有-系统无"
LEGACY_SHEET = "所有谷歌有-系统没有"
PORT_SHEET = "系统无-谷歌表有"
LEGACY_PORT_SHEET = "未使用端口"
PORT_HEADERS = ("设备", "机房", "端口", "归一化卡槽", "电话", "运营商", "状态", "谷歌表对比", "谷歌ICCID", "平台ICCID")
PRESENT = "谷歌表存在"
SHEETS = ("二號機房", "三號機房", "四號機房")
HEADERS = ("Google页签", "Google行号", "卡槽", "归一化卡槽", "备注", "ICCID", "上卡日期", "电信商")
SCOPE = "google_slots_absent_from_current_system_export"
FLAG = "OPENSMS_GOOGLE_SYSTEM_MISSING_ENABLED"
YELLOW = "FFFF00"


def port_iccid_index(records):
    """Keep every distinct literal ICCID for a physical slot; never pick one."""
    result = {}
    for record in records:
        key = slot(record.get("slot", record.get("raw_slot")))
        value = "" if record.get("iccid") is None else str(record["iccid"]).strip()
        if key and value:
            if "\n" in value or "\r" in value:
                raise ValueError("ICCID包含换行，不能安全合并同卡槽记录")
            result.setdefault(key, set()).add(value)
    return {key: "\n".join(sorted(values)) for key, values in result.items()}


def attach_port_iccids(rows, google_records, system_records):
    google = port_iccid_index(google_records)
    system = port_iccid_index(system_records)
    return [list(row[:8]) + [google.get(slot(row[3]), ""), system.get(slot(row[3]), "")] for row in rows]


def enabled():
    return os.environ.get(FLAG) == "1"


def slot(value):
    if value is None:
        return ""
    value = unicodedata.normalize("NFKC", str(value).strip()).replace(" ", "").replace("-", "_")
    parts = []
    for part in value.split("_"):
        match = re.fullmatch(r"(\d+)(?:\.\d+)?", part)
        parts.append(str(int(match[1])) if match else part.lower())
    return "_".join(parts)


def source_from_snapshot(snapshot_path, snapshot_sha256, run_id):
    path = Path(snapshot_path)
    if not run_id or not path.is_absolute() or not re.fullmatch(r"[a-f0-9]{64}", snapshot_sha256 or ""):
        raise ValueError("Google差集缺少本轮快照绝对路径、runId或SHA256")
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != snapshot_sha256:
        raise ValueError("Google差集快照SHA256不一致")
    snapshot = json.loads(raw)
    if (snapshot.get("status") != "passed" or snapshot.get("runId") != run_id
            or snapshot.get("source") != "google_chatr_dates"
            or snapshot.get("sourceHashVerified") is not True):
        raise ValueError("Google差集快照未绑定本轮已核验来源")
    source = Path(snapshot.get("sourcePath") or "")
    digest = snapshot.get("sourceSha256")
    if not source.is_absolute() or not re.fullmatch(r"[a-f0-9]{64}", digest or ""):
        raise ValueError("Google差集原始文件路径或SHA256缺失")
    raw = source.read_bytes()
    if hashlib.sha256(raw).hexdigest() != digest:
        raise ValueError("Google差集原始文件SHA256不一致")
    source_sheets = bound_room_sheets(raw, snapshot.get("sourceSheets"))
    return raw, {"enabled": True, "status": "passed", "scope": SCOPE,
                 "run_id": run_id, "sheet": SHEET, "source_path": str(source),
                 "source_sha256": digest, "snapshot_path": str(path),
                 "snapshot_sha256": snapshot_sha256, "source_sheets": list(source_sheets)}


def read_rows(raw):
    """Retain every source row, including duplicate slots and blank dates."""
    values = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    formulas = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    records, counts, blanks = [], {}, 0
    aliases = {"slot": ("卡槽",), "note": ("備註", "备注"), "iccid": ("iccid", "ICCID"),
               "open": ("上卡日期",), "carrier": ("電信商/價格",), "activation": ("激活日期",)}
    try:
        source_sheets = select_room_sheets(values)
        for name in source_sheets:
            ws, fws = values[name], formulas[name]
            headers = [str(c.value or "").strip() for c in next(ws.iter_rows())]
            columns = {}
            for key, names in aliases.items():
                matches = [i for i, header in enumerate(headers) if header in names]
                if len(matches) != 1:
                    raise ValueError(f"Google差集 {name} 字段缺失或重复: {names[0]}")
                columns[key] = matches[0]
            counts[name] = 0
            for number, (row, frow) in enumerate(zip(ws.iter_rows(min_row=2), fws.iter_rows(min_row=2)), 2):
                if not any(c.value is not None for c in row):
                    continue
                for index in columns.values():
                    if row[index].data_type == "e" or (frow[index].data_type == "f" and row[index].value is None):
                        raise ValueError(f"Google差集 {name}!{row[index].coordinate} 错误或公式无缓存")
                record = {key: row[index].value for key, index in columns.items()}
                record.update(sheet=name, row=number)
                if not slot(record["slot"]):
                    blanks += 1
                    continue
                records.append(record)
                counts[name] += 1
        return records, counts, blanks
    finally:
        values.close()
        formulas.close()


def load_from_env():
    if not enabled():
        return None
    if os.environ.get("OPENSMS_GOOGLE_ASSOCIATION_ENABLED") != "0":
        raise ValueError("独立Google差集不能启用旧完整关联")
    raw, proof = source_from_snapshot(os.environ.get("OPENSMS_CHATR_GOOGLE_DATES_PATH", ""),
        os.environ.get("OPENSMS_CHATR_GOOGLE_DATES_SHA256", ""), os.environ.get("OPENSMS_RUN_ID", ""))
    rows, counts, blanks = read_rows(raw)
    return {"rows": rows, "proof": proof, "counts": counts, "blanks": blanks}


def write_sheet(book, evidence, system_slots, parse_date):
    if evidence is None:
        return None
    selected = [r for r in evidence["rows"] if slot(r["slot"]) not in system_slots]
    for name in (LEGACY_SHEET, SHEET):
        if name in book.sheetnames:
            del book[name]
    ws = book.create_sheet(SHEET, book.sheetnames.index(PORT_SHEET) + 1)
    ws.append(HEADERS)
    highlighted = 0
    for row_number, record in enumerate(selected, 2):
        opened = parse_date(record["open"])
        values = [record["sheet"], record["row"], record["slot"], slot(record["slot"]),
                  record["note"], str(record["iccid"]) if record["iccid"] is not None else None,
                  opened or record["open"], record["carrier"]]
        for column, value in enumerate(values, 1):
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError("Google差集字段超过Excel长度上限，禁止截断")
            cell = ws.cell(row_number, column, value)
            if isinstance(value, str):
                cell.data_type = "s"  # Untrusted external strings never become formulas.
        ws.cell(row_number, 6).number_format = "@"
        if isinstance(values[6], (date, datetime)):
            ws.cell(row_number, 7).number_format = "yyyy-mm-dd"
        ws.cell(row_number, 5).alignment = Alignment(wrap_text=True, vertical="top")
        if parse_date(record["activation"]):
            highlighted += 1
            for cell in ws[row_number]:
                cell.fill = PatternFill(fill_type="solid", fgColor=YELLOW)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGH", (16, 14, 18, 18, 52, 24, 16, 22)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    proof = {**evidence["proof"], "source_rows": len(evidence["rows"]),
             "source_sheet_rows": evidence["counts"], "blank_slot_rows": evidence["blanks"],
             "system_unique_slots": len(system_slots), "missing_rows": len(selected),
             "missing_unique_slots": len({slot(r["slot"]) for r in selected}),
             "activation_highlighted_rows": highlighted}
    book["处理汇总"].append((SHEET + "行数", len(selected)))
    book["处理汇总"].append((SHEET + "口径", "按归一化卡槽与本轮系统全部状态记录比较；保留重复Google行，不判断实际上架"))
    return proof


def review_port_presence(book, system_slots, google_slots, google_records=(), system_records=()):
    """Validate the annotation against raw source identities, never inferred status."""
    errors = []
    if PORT_SHEET not in book.sheetnames:
        return {"matches": False, "marked_rows": 0, "errors": {"missing_port_sheet": 1}}
    ws = book[PORT_SHEET]
    if tuple(c.value for c in next(ws.iter_rows(max_row=1))) != PORT_HEADERS:
        errors.append("port_headers_mismatch")
    marked = 0
    mismatched = 0
    # Independent reconstruction from bound raw records (not generated columns).
    indices = []
    for records in (google_records, system_records):
        index = {}
        for record in records:
            key = slot(record.get("slot", record.get("raw_slot")))
            value = "" if record.get("iccid") is None else str(record["iccid"]).strip()
            if key and value:
                index.setdefault(key, set()).add(value)
        indices.append(index)
    for cells in ws.iter_rows(min_row=2):
        if not any(c.value not in (None, "") for c in cells):
            continue
        if len(cells) != len(PORT_HEADERS):
            errors.append("port_columns_mismatch")
            continue
        key = slot(cells[3].value)
        expected = PRESENT if key and key not in system_slots and key in google_slots else ""
        marked += bool(expected)
        actual = "" if cells[7].value is None else cells[7].value
        if actual != expected or cells[7].data_type in ("f", "e") or (expected and cells[7].data_type != "s"):
            errors.append("port_presence_mismatch")
        if expected and cells[6].value != "系统中无":
            errors.append("port_presence_status_mismatch")
        iccids = ["\n".join(sorted(index.get(key, set()))) for index in indices]
        different = iccids[0] != iccids[1]
        mismatched += different
        for cell, wanted in zip(cells[8:10], iccids):
            if (str(cell.value or "") != wanted or cell.data_type in ("f", "e")
                    or cell.number_format != "@"):
                errors.append("port_iccid_values_or_format_mismatch")
        for cell in cells:
            yellow = cell.fill.patternType == "solid" and str(cell.fill.fgColor.rgb).endswith(YELLOW)
            if yellow != different:
                errors.append("port_iccid_highlight_mismatch")
    if any(name in book.sheetnames for name in (LEGACY_SHEET, LEGACY_PORT_SHEET)):
        errors.append("legacy_sheet_name")
    if SHEET in book.sheetnames and book.sheetnames.index(SHEET) != book.sheetnames.index(PORT_SHEET) + 1:
        errors.append("difference_sheet_not_adjacent")
    return {"matches": not errors, "marked_rows": marked, "iccid_mismatch_rows": mismatched, "errors": dict(Counter(errors))}


def review(book, raw_headers, raw_rows, proof, snapshot_path, snapshot_hash, run_id, parse_date):
    """Reopen original Google bytes, then independently rebuild the difference."""
    raw, expected = source_from_snapshot(snapshot_path, snapshot_hash, run_id)
    records, counts, blanks = read_rows(raw)
    source_index = raw_headers.index("卡槽")
    seen = set()
    for row in raw_rows:
        key = slot(row[source_index])
        if key:
            seen.add(key)
    wanted, keys, yellow = [], set(), 0
    for row in records:
        key = slot(row["slot"])
        if key in seen:
            continue
        opened = parse_date(row["open"])
        activated = bool(parse_date(row["activation"]))
        wanted.append(([row["sheet"], row["row"], row["slot"], key, row["note"],
                        str(row["iccid"]) if row["iccid"] is not None else None,
                        opened or row["open"], row["carrier"]], activated))
        keys.add(key)
        yellow += activated
    expected.update(source_rows=len(records), source_sheet_rows=counts, blank_slot_rows=blanks,
                    system_unique_slots=len(seen), missing_rows=len(wanted),
                    missing_unique_slots=len(keys), activation_highlighted_rows=yellow)
    errors = []
    if proof != expected:
        errors.append("source_or_counts_mismatch")
    iccid_column = next((i for i, name in enumerate(raw_headers) if str(name).strip().lower() == "iccid"), None)
    system_records = [{"slot": row[source_index], "iccid": row[iccid_column] if iccid_column is not None else None} for row in raw_rows]
    presence = review_port_presence(book, seen, {slot(r["slot"]) for r in records}, records, system_records)
    errors.extend(presence["errors"])
    if SHEET not in book.sheetnames:
        errors.append("missing_sheet")
    else:
        ws = book[SHEET]
        if tuple(c.value for c in next(ws.iter_rows(max_row=1))) != HEADERS:
            errors.append("headers_mismatch")
        actual = list(ws.iter_rows(min_row=2))
        if len(actual) != len(wanted):
            errors.append("rows_mismatch")
        for cells, (values, activated) in zip(actual, wanted):
            for cell, value in zip(cells, values):
                actual_value = None if cell.value == "" else cell.value
                expected_value = None if value == "" else value
                if (actual_value != expected_value or isinstance(actual_value, bool) != isinstance(expected_value, bool)
                        or (isinstance(value, str) and value and cell.data_type != "s")
                        or (isinstance(value, (int, float)) and cell.data_type != "n")):
                    errors.append("values_or_types_mismatch")
                is_yellow = bool(cell.fill and cell.fill.patternType == "solid" and str(cell.fill.fgColor.rgb).endswith(YELLOW))
                if is_yellow != activated:
                    errors.append("highlight_mismatch")
            if cells[5].number_format != "@":
                errors.append("iccid_format_mismatch")
            if isinstance(values[6], (date, datetime)) and cells[6].number_format != "yyyy-mm-dd":
                errors.append("date_format_mismatch")
        summaries = {row[0]: row[1] for row in book["处理汇总"].iter_rows(values_only=True)}
        if summaries.get(SHEET + "行数") != len(wanted):
            errors.append("summary_mismatch")
    return {**expected, "port_presence_review": presence, "matches": not errors, "errors": dict(Counter(errors))}
