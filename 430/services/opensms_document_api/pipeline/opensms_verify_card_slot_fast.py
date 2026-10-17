#!/usr/bin/env python3
"""Fast read-only verification for an already generated OpenSMS card-slot workbook."""

from __future__ import annotations

import hashlib
from io import BytesIO
import json
import math
import os
import posixpath
import re
import sys
import unicodedata
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

try:
    from .chatr_iccid_conflicts import SHEET as ICCID_CONFLICT_SHEET, DECISION as ICCID_CONFLICT_DECISION, review_conflict_sheet
    from .verify_device_daily_statistics import review_device_daily_statistics
    from .card_retention_report import review_retention_report
    from .machine_number import verify_machine_number
    from .creation_usage_days import verify_usage_days
    from .ko_unactivated import review as review_ko_unactivated
    from .google_room_sheets import bound_room_sheets
    from .google_system_missing import enabled as google_difference_enabled, review as review_google_difference, SHEET as GOOGLE_DIFFERENCE_SHEET, PORT_HEADERS, review_port_presence
except ImportError:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from chatr_iccid_conflicts import SHEET as ICCID_CONFLICT_SHEET, DECISION as ICCID_CONFLICT_DECISION, review_conflict_sheet
    from verify_device_daily_statistics import review_device_daily_statistics
    from card_retention_report import review_retention_report
    from machine_number import verify_machine_number
    from creation_usage_days import verify_usage_days
    from ko_unactivated import review as review_ko_unactivated
    from google_room_sheets import bound_room_sheets
    from google_system_missing import enabled as google_difference_enabled, review as review_google_difference, SHEET as GOOGLE_DIFFERENCE_SHEET, PORT_HEADERS, review_port_presence
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.text import Text
from openpyxl.reader.strings import read_string_table
from openpyxl.utils.datetime import from_excel

ROOM_CONFIG_PATH = Path(__file__).resolve().parents[3] / 'scripts/opensms_room_settings.json'
ROOM_CONFIG_BYTES = ROOM_CONFIG_PATH.read_bytes()
ROOM_CONFIG = json.loads(ROOM_CONFIG_BYTES)
ROOM_CONFIG_SHA256 = hashlib.sha256(ROOM_CONFIG_BYTES).hexdigest()
ROOM_DEVICE_MAP = {}
for _room in ROOM_CONFIG['rooms']:
    _count = 0
    for _start, _end in _room['deviceRanges']:
        if type(_start) is not int or type(_end) is not int or _start < 0 or _end < _start:
            raise ValueError('独立核验机房设备范围无效')
        for _device in range(_start, _end + 1):
            if _device in ROOM_DEVICE_MAP:
                raise ValueError('独立核验机房设备范围重叠')
            ROOM_DEVICE_MAP[_device] = _room['name']
            _count += 1
    if _count != _room['deviceCount']:
        raise ValueError('独立核验机房配置数量不符')
if len(ROOM_DEVICE_MAP) != ROOM_CONFIG['totalDeviceCount']:
    raise ValueError('独立核验机房配置总数不符')


EXPECTED_HEADERS = [
    "电话",
    "运营商",
    "过期状态",
    "换号状态",
    "koo是否2天",
    "激活后天数",
    "opr",
    "卡槽",
    "连续低于x%",
    "连续0回复",
    "开卡日期",
    "激活日期",
    "截止日",
    "机房",
    "Google電信商/價格",
    "OPR匹配预警",
]
REQUIRED_ASSOC_SHEETS = ["二號機房", "三號機房", "四號機房"]
GOOGLE_ONLY_SHEETS = (
    "未匹配卡槽", "电信商_OPR预警", "谷歌有-系统无", "所有谷歌有-系统没有",
    "谷歌文档中有激活日期-系统无法使用", "谷歌文档中有激活日期-系统没有",
    "Google有日期_主表缺失", "设备缺失端口",
)
GOOGLE_ONLY_HEADERS = ("Google電信商/價格", "OPR匹配预警")
FORMULA_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}
NUMERIC_TEXT_PATTERN = re.compile(
    r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?:[eE][+-]?\d+)?$"
)
def text(value) -> str:
    return "" if value is None else str(value).strip()


def review_task_f_missing_dates_new_card_row(opr, activation_date, cutoff_date, change_status) -> dict:
    """Independently check stored values; invalid nonempty dates are not absent."""
    eligible = (
        re.search(r"kooda|koodo|telus1", text(opr), re.IGNORECASE) is not None
        and text(activation_date) == ""
        and text(cutoff_date) == ""
    )
    return {
        "eligible": eligible,
        "expected": "未换号" if eligible else None,
        "actual": "" if change_status is None else str(change_status),
        "matches": not eligible or change_status == "未换号",
    }


def number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value) if math.isfinite(float(value)) else None
    raw = text(value).replace(",", "")
    percent = raw.endswith("%")
    try:
        parsed = float(raw[:-1] if percent else raw)
    except ValueError:
        return None
    return parsed / 100 if percent else parsed


def review_numeric_text(workbook) -> dict:
    text_numeric_cells = 0
    examples = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                normalized = cell.value.strip().replace(",", "")
                if normalized.endswith("%"):
                    normalized = normalized[:-1].strip()
                if not NUMERIC_TEXT_PATTERN.fullmatch(normalized):
                    continue
                text_numeric_cells += 1
                if len(examples) < 30:
                    examples.append(
                        {
                            "sheet": sheet.title,
                            "cell": cell.coordinate,
                            "value": cell.value,
                            "status": "preserved_as_text",
                        }
                    )
    return {
        "matches": True,
        "conversion_enabled": False,
        "preserved_text_numeric_cells": text_numeric_cells,
        "examples": examples,
        "policy": "preserve_source_cell_types",
    }


def review_numeric_text_ooxml(path: Path) -> dict:
    """Independently inspect stored string cells without a second cell model.

    This is an output read, not the generator's cache. Formula cells are excluded
    just as data_only=False excludes their cached results. Numeric cells remain
    numeric; shared and rich strings use openpyxl's own text decoding semantics.
    Workbook order and cell order preserve the historical first-30 examples.
    """
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    relation_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    count = 0
    examples = []
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relations
                   if item.attrib.get("TargetMode") != "External"}
        strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            with archive.open("xl/sharedStrings.xml") as stream:
                strings = read_string_table(stream)
        for sheet in workbook.find(f"{namespace}sheets"):
            target = targets[sheet.attrib[f"{relation_namespace}id"]]
            member = posixpath.normpath(target.lstrip("/") if target.startswith("/") else f"xl/{target}")
            if not member.startswith("xl/"):
                raise RuntimeError("工作表OOXML路径越界")
            root = ET.fromstring(archive.read(member))
            for cell in root.iter(f"{namespace}c"):
                kind = cell.attrib.get("t", "n")
                if kind not in {"s", "str", "inlineStr"} or cell.find(f"{namespace}f") is not None:
                    continue
                if kind == "inlineStr":
                    inline = cell.find(f"{namespace}is")
                    if inline is None:
                        continue
                    direct = inline.find(f"{namespace}t")
                    value = Text.from_tree(inline).content if inline.find(f"{namespace}r") is not None else (direct.text if direct is not None else "")
                else:
                    node = cell.find(f"{namespace}v")
                    if node is None or node.text is None:
                        continue
                    value = strings[int(node.text)] if kind == "s" else node.text
                if not isinstance(value, str):
                    continue
                normalized = value.strip().replace(",", "")
                if normalized.endswith("%"):
                    normalized = normalized[:-1].strip()
                if not NUMERIC_TEXT_PATTERN.fullmatch(normalized):
                    continue
                count += 1
                if len(examples) < 30:
                    examples.append({"sheet": sheet.attrib["name"], "cell": cell.attrib["r"],
                                     "value": value, "status": "preserved_as_text"})
    return {"matches": True, "conversion_enabled": False,
            "preserved_text_numeric_cells": count, "examples": examples,
            "policy": "preserve_source_cell_types"}


def phone_key(value) -> str:
    digits = "".join(character for character in text(value) if character.isdigit())
    return digits[-10:]


def normalize_slot(value) -> str:
    raw = text(value).replace(" ", "").replace("-", "_")
    normalized = []
    for part in raw.split("_"):
        match = re.fullmatch(r"(\d+)(?:\.\d+)?", part)
        normalized.append(str(int(match.group(1))) if match else part.lower())
    return "_".join(normalized)


def parse_device_port(value):
    normalized = normalize_slot(value)
    match = re.fullmatch(r"(\d+)_(\d+)", normalized)
    if not match:
        return None
    device, port_text = match.groups()
    port = int(port_text)
    if not 1 <= port <= 64:
        return None
    return str(int(device)), port


def room_for_device(device: str) -> str:
    value = text(device)
    return ROOM_DEVICE_MAP.get(int(value), '未配置') if value.isascii() and value.isdigit() else '未配置'


def review_room_configuration(book, receipt, summary):
    """Reconstruct mapping from config, then reread every affected workbook view."""
    errors = []
    expected_summary = {'机房配置生效日期': ROOM_CONFIG['effectiveDate'],
                        '机房配置SHA256': ROOM_CONFIG_SHA256,
                        '配置设备总数': ROOM_CONFIG['totalDeviceCount']}
    for room in ROOM_CONFIG['rooms']:
        ranges = '、'.join(str(a) if a == b else f'{a}–{b}' for a, b in room['deviceRanges'])
        expected_summary[room['name'] + '设备配置'] = f"{ranges}（{room['deviceCount']}台）"
    if receipt.get('room_configuration') != {'config': ROOM_CONFIG, 'sha256': ROOM_CONFIG_SHA256}:
        raise ValueError('机房配置回执与当前配置不一致')
    if any(summary.get(key) != value for key, value in expected_summary.items()):
        errors.append('处理汇总未完整显示最新机房配置')
    main_counts = Counter()
    checked = {}
    for sheet in book.worksheets:
        header_row = 5 if sheet.title == 'A11连续低于3%' else 1
        rows = sheet.iter_rows(min_row=header_row, values_only=True)
        headers = list(next(rows, ()))
        if '机房' not in headers:
            continue
        room_index = headers.index('机房')
        machine_index = next((headers.index(name) for name in ('机器号', '设备', '设备号') if name in headers), None)
        slot_index = headers.index('卡槽') if '卡槽' in headers else None
        if slot_index is None and machine_index is None:
            raise ValueError(f'{sheet.title}机房列缺少设备依据')
        count = 0
        for row_number, row in enumerate(rows, header_row + 1):
            if machine_index is not None:
                device = text(row[machine_index])
            else:
                match = re.match(r'^(\d+)_', normalize_slot(row[slot_index]))
                device = match.group(1) if match else ''
            expected = room_for_device(device)
            if row[room_index] != expected:
                errors.append(f'{sheet.title}第{row_number}行机房与最新配置不一致')
            count += 1
            if sheet.title == book.worksheets[0].title:
                main_counts[expected] += 1
        checked[sheet.title] = count
    expected_counts = {**{r['name']: main_counts[r['name']] for r in ROOM_CONFIG['rooms']},
                       '未配置': main_counts['未配置']}
    if receipt.get('room_counts') != expected_counts:
        errors.append('机房行数回执与独立主表计数不一致')
    for name, count in expected_counts.items():
        key = '未配置机房行数' if name == '未配置' else name + '行数'
        if summary.get(key) != count:
            errors.append('机房行数汇总与独立主表计数不一致')
    return {'matches': not errors, 'errors': errors, 'config_sha256': ROOM_CONFIG_SHA256,
            'effective_date': ROOM_CONFIG['effectiveDate'],
            'configured_devices': len(ROOM_DEVICE_MAP), 'checked_sheets': checked,
            'room_counts': expected_counts}


def unique_join(values) -> str:
    return "、".join(dict.fromkeys(text(value) for value in values if text(value)))


def normalize_carrier(value) -> str:
    raw = text(value).lower()
    if not raw:
        return ""
    if "koodo" in raw or "kooda" in raw or "telus1" in raw:
        return "koodo"
    if "chatr" in raw:
        return "chatr"
    if "kini" in raw:
        return "kini"
    if "lucky" in raw:
        return "lucky mobile"
    if "rogers" in raw:
        return "rogers"
    return re.sub(r"[^a-z\u3400-\u9fff]+", "", raw)


def date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 <= value <= 60000:
            return from_excel(value).date()
        if float(value).is_integer() and re.fullmatch(r"(?:19|20)\d{6}", str(int(value))):
            try:
                return datetime.strptime(str(int(value)), "%Y%m%d").date()
            except ValueError:
                return None
    raw = text(value).strip("'\"").replace("年", "-").replace("月", "-").replace("日", "")
    raw = raw.replace("／", "-").replace("－", "-").replace("．", "-").replace("/", "-").replace(".", "-")
    if re.fullmatch(r"(?:19|20)\d{6}", raw):
        try:
            return datetime.strptime(raw, "%Y%m%d").date()
        except ValueError:
            return None
    match = re.search(r"((?:19|20)\d{2})-(\d{1,2})-(\d{1,2})", raw)
    if match:
        try:
            return date(*map(int, match.groups()))
        except ValueError:
            return None
    ambiguous = re.search(r"(?<!\d)(\d{1,2})-(\d{1,2})-((?:19|20)?\d{2})(?!\d)", raw)
    if ambiguous:
        first, second, year = map(int, ambiguous.groups())
        year = year + (2000 if year < 70 else 1900) if year < 100 else year
        month, day = (second, first) if first > 12 else (first, second)
        try:
            return date(year, month, day)
        except ValueError:
            return None
    for pattern in ("%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    return None


def strict_system_date(value, label="系统日期"):
    """Read an ISO calendar date without inferring it or converting timezones."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        raise RuntimeError(f"{label}必须是有效ISO日期，不能使用非空数值或布尔值")
    raw = value
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
            return date.fromisoformat(raw)
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:\d{2})?", raw):
            return datetime.fromisoformat(raw.replace("Z", "+00:00")).date()
    except ValueError as error:
        raise RuntimeError(f"{label}不是有效ISO日期: {raw}") from error
    raise RuntimeError(f"{label}不是有效ISO日期: {raw}")


def system_kooda_key(slot, phone):
    normalized_slot = normalize_slot(slot)
    digits = re.sub(r"\D", "", text(phone))
    if not normalized_slot or len(digits) < 10:
        return None
    return normalized_slot, digits[-10:]


def system_kooda_iccid_header_matches(headers) -> bool:
    return headers.count("ICCID") == 1 and headers[-1] == "ICCID"


def review_export_schema_evidence(raw_headers: list[str], raw_values: list, main_headers: list[str], main_values: list) -> dict:
    """Retain immutable raw XLSX evidence, not duplicate visible date columns."""
    raw_index = {name: index for index, name in enumerate(raw_headers)}
    main_index = {name: index for index, name in enumerate(main_headers)}
    mismatches = []
    evidence_columns = {}
    for name in ("激活日期", "截止日"):
        if main_headers.count(name) != 1 or f"系统导出{name}" in main_headers:
            mismatches.append({"field": name, "reason": "date_column_missing_duplicate_or_redundant"})
    if len(raw_values) != len(main_values):
        mismatches.append({"field": "rows", "reason": "source_final_row_count_differs"})
    compared_iccid = 0
    empty_string_blank_cells = 0
    for row_number, (raw, final) in enumerate(zip(raw_values, main_values), start=2):
        for name in ("电话", "卡槽"):
            if text(raw[raw_index[name]]) != text(final[main_index[name]]):
                mismatches.append({"row": row_number, "field": name, "reason": "source_row_identity_changed"})
        if "ICCID" not in raw_index:
            continue
        source = raw[raw_index["ICCID"]]
        actual = final[main_index["ICCID"]]
        kooda = re.search(r"kooda|koodo|telus1", text(raw[raw_index["opr"]]), re.I) is not None
        if not kooda:
            # XLSX shared empty strings reopen as blank after openpyxl saves.
            # Only blank-to-blank is equivalent; retain strict nonempty identity.
            blank_equivalent = source in (None, "") and actual in (None, "")
            if not blank_equivalent and (source != actual or type(source) is not type(actual)):
                mismatches.append({"row": row_number, "field": "ICCID", "reason": "non_kooda_source_value_or_type_changed"})
        elif text(source):
            compared_iccid += 1
            source_key = re.sub(r"[Ff]$", "", text(source))
            actual_key = re.sub(r"[Ff]$", "", text(actual))
            if not isinstance(source, str) or not re.fullmatch(r"\d{18,22}", source_key, re.ASCII) or source_key != actual_key:
                mismatches.append({"row": row_number, "field": "ICCID", "reason": "source_system_standard_key_conflict"})
    return {
        "source_column_count": len(raw_headers), "source_has_iccid": "ICCID" in raw_index,
        "source_date_evidence_columns": evidence_columns,
        "source_rows": len(raw_values), "final_rows": len(main_values),
        "iccid_compared_rows": compared_iccid,
        "empty_string_dates_reopened_as_blank": empty_string_blank_cells,
        "mismatch_count": len(mismatches), "mismatches": mismatches[:50], "matches": not mismatches,
    }


def load_system_kooda_snapshot(path: Path, expected_sha256: str, expected_run_id: str) -> dict:
    """Independently authenticate and index this run's raw system receipt once."""
    expected_hash = text(expected_sha256).lower()
    if not re.fullmatch(r"[a-f0-9]{64}", expected_hash) or not text(expected_run_id):
        raise RuntimeError("Kooda系统独立核验缺少本轮SHA256或运行编号")
    raw_bytes = path.read_bytes()
    actual_hash = hashlib.sha256(raw_bytes).hexdigest()
    if actual_hash != expected_hash:
        raise RuntimeError("Kooda系统快照SHA256不匹配，拒绝读取")
    payload = json.loads(raw_bytes)
    if payload.get("status") != "passed" or payload.get("runId") != expected_run_id:
        raise RuntimeError("Kooda系统快照未通过采集或不属于本轮")
    from report_kooda_source import is_report_source, validate_projection
    report_source = is_report_source(payload)
    if report_source:
        validate_projection(payload, expected_run_id)
    for flag in (("silentBrowserMode",) if report_source else ("paginationTotalStable", "paginationBackendIdsUnique", "silentBrowserMode")):
        if payload.get(flag) is not True:
            raise RuntimeError(f"Kooda系统快照缺少有效证明: {flag}")
    if payload.get("browserForegroundActivationAttempted") is not False:
        raise RuntimeError("Kooda系统快照未证明浏览器全程静默")
    system_rows = payload.get("systemRows")
    if (not isinstance(system_rows, list) or type(payload.get("totalRows")) is not int
            or payload["totalRows"] != len(system_rows)):
        raise RuntimeError("Kooda系统快照全量原始行数不一致")
    indexed, by_phone = {}, {}
    backend_ids = set()
    unkeyed_rows = 0
    for number, source in enumerate(system_rows, start=1):
        if not isinstance(source, dict) or normalize_carrier(source.get("opr")) != "koodo":
            raise RuntimeError(f"Kooda系统快照包含非Kooda原始行: {number}")
        backend_id = text(source.get("sourceRowIdentity") if report_source else source.get("id"))
        if not backend_id or backend_id in backend_ids:
            raise RuntimeError("Kooda系统快照后台ID为空或重复")
        backend_ids.add(backend_id)
        iccid = source.get("cardStore")
        if iccid is not None and not isinstance(iccid, str):
            raise RuntimeError(f"Kooda系统ICCID不是文本，拒绝可能的精度损失: {number}")
        activation = strict_system_date(source.get("chatrActivateDate"), f"系统第{number}行chatrActivateDate")
        cutoff = strict_system_date(source.get("chatrPlanExpire"), f"系统第{number}行chatrPlanExpire")
        record = {"raw": source, "iccid": iccid, "activation": activation, "cutoff": cutoff}
        by_phone.setdefault(phone_key(source.get('cardNo')), []).append(record)
        key = system_kooda_key(source.get("slotIdentif"), source.get("cardNo"))
        if key is None:
            unkeyed_rows += 1
            continue
        if key in indexed and not report_source:
            raise RuntimeError("Kooda系统快照卡槽与手机号复合键重复")
        indexed[key] = record
    return {
        "path": str(path.resolve()), "sha256": actual_hash, "run_id": expected_run_id,
        "total_rows": len(system_rows), "unkeyed_rows": unkeyed_rows, "index": indexed,
        "phone_index": by_phone, "raw_rows": system_rows,
        "source_kind": payload.get('sourceKind'), "source_scope": payload.get('sourceScope'),
        "projection": payload if report_source else None,
        "source_row_index": {row['sourceRow']: {
            'raw': row, 'iccid': row['cardStore'],
            'activation': strict_system_date(row['chatrActivateDate'], 'source activation'),
            'cutoff': strict_system_date(row['chatrPlanExpire'], 'source cutoff')}
            for row in system_rows} if report_source else {},
    }


def find_system_record(snapshot, row, source_row=None):
    if snapshot.get('projection'):
        record = snapshot['source_row_index'].get(source_row)
        raw = record.get('raw') if record else None
        if (raw is None or phone_key(raw['cardNo']) != phone_key(row.get('电话'))
                or system_kooda_key(raw['slotIdentif'], raw['cardNo']) != system_kooda_key(row.get('卡槽'), row.get('电话'))
                # Excel reopens a saved empty text cell as None. Only these
                # two exact empty representations are equivalent; preserve
                # strict equality for every nonempty ICCID (including spaces).
                or not (raw['cardStore'] == row.get('ICCID')
                        or (raw['cardStore'] in (None, '') and row.get('ICCID') in (None, '')))):
            return None
        return record
    key = system_kooda_key(row.get('卡槽'), row.get('电话'))
    hit = snapshot['index'].get(key) if key else None
    if hit is not None:
        return hit
    # Independent reconstruction of the historical exception from raw sources.
    rows = snapshot.get('phone_index', {}).get(phone_key(row.get('电话')), [])
    if 'ICCID' not in row or not key or len(rows) != 1:
        return None
    record, raw = rows[0], rows[0]['raw']
    # Rebuild the allowed status pair independently, not by importing the
    # generator's helper. No other live-to-offline transition is accepted.
    old_archive = row.get('状态') == '下架' and raw.get('status') == '换卡已下线'
    swapped_archive = (row.get('状态') == '上架' and raw.get('status') == '换卡已下线'
                       and raw.get('displayStatus') == '卡槽异常'
                       and raw.get('showInTable') is False and raw.get('userAble') is False)
    empty_pool_archive = (row.get('状态') in ('上架', '下架')
                          and raw.get('status') == '猫池空号已下架')
    if not (old_archive or empty_pool_archive or swapped_archive):
        return None
    if raw.get('showAble') is not False or raw.get('slotIdentif') not in (None, ''):
        return None
    if str(row.get('卡槽') or '').split('_')[0] != str(raw.get('machineName') or ''):
        return None
    stamps = []
    for value in (row.get('创建时间'), raw.get('createTime')):
        if isinstance(value, datetime) and value.tzinfo is None:
            stamps.append(value)
        elif isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}', value):
            try:
                stamps.append(datetime.strptime(value, '%Y-%m-%d %H:%M:%S'))
            except ValueError:
                return None
        else:
            return None
    keys = []
    for value in (row['ICCID'], raw.get('cardStore')):
        if value in (None, ''):
            keys.append('')
        elif isinstance(value, str) and re.fullmatch(r'[0-9]{18,22}[Ff]?', value.strip()):
            keys.append(re.sub(r'[Ff]$', '', value.strip()))
        else:
            return None
    if empty_pool_archive and not keys[0]:
        return None
    return record if stamps[0] == stamps[1] and keys[0] == keys[1] else None


def review_system_kooda_row(snapshot: dict, row: dict, resolved_dates: dict | None = None, source_row=None) -> dict | None:
    """Compare a stored main-sheet row to exactly one raw system identity."""
    if normalize_carrier(row.get("opr")) != "koodo":
        return None
    key = system_kooda_key(row.get("卡槽"), row.get("电话"))
    record = find_system_record(snapshot, row, source_row)
    errors = []
    if record is None:
        errors.append({"field": "identity", "reason": "missing_exact_slot_phone_match"})
    else:
        expected_iccid = record["iccid"]
        actual_iccid = row.get("ICCID")
        iccid_matches = (
            isinstance(actual_iccid, str) and actual_iccid == expected_iccid
            if text(expected_iccid) else not text(actual_iccid)
        )
        if not iccid_matches:
            errors.append({"field": "ICCID", "expected": expected_iccid, "actual": actual_iccid})
        for column, source_key in (("激活日期", "activation"), ("截止日", "cutoff")):
            expected = resolved_dates[column] if resolved_dates is not None else record[source_key]
            try:
                actual = strict_system_date(row.get(column), f"报表{column}")
                matches = actual == expected
            except RuntimeError:
                actual = text(row.get(column))
                matches = False
            if not matches:
                errors.append({"field": column, "expected": str(expected) if expected else "", "actual": str(actual) if actual is not None else ""})
        if text(row.get("开卡日期")):
            errors.append({"field": "开卡日期", "expected": "", "actual": str(row["开卡日期"])})
    return {"matched": record is not None, "matches": not errors, "errors": errors}


def load_frozen_date_snapshot(path: Path, expected_sha256: str, run_id: str) -> dict:
    """Independent verifier implementation of frozen JSON validation."""
    if not run_id or not re.fullmatch(r"[a-fA-F0-9]{64}", expected_sha256):
        raise RuntimeError("独立日期复核缺少本轮SHA256或runId")
    content = path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    if digest != expected_sha256.lower():
        raise RuntimeError("独立日期缓存SHA256不匹配")
    payload = json.loads(content)
    records = payload.get("records")
    if (payload.get("schemaVersion") != 1 or payload.get("status") != "passed"
            or payload.get("runId") != run_id or not isinstance(records, list)
            or type(payload.get("recordCount")) is not int or payload["recordCount"] != len(records)
            or type(payload.get("sourceExists")) is not bool or not payload.get("createdAt")):
        raise RuntimeError("独立日期缓存契约或runId不匹配")
    if not payload["sourceExists"] and records:
        raise RuntimeError("独立日期缓存数据库不存在却包含日期记录")
    indexed = {}
    for record in records:
        key = record.get("queryIccid") if isinstance(record, dict) else None
        if not isinstance(key, str) or not re.fullmatch(r"\d{18,22}", key, re.ASCII) or key in indexed:
            raise RuntimeError("独立日期缓存ICCID非法或重复")
        activation = strict_system_date(record.get("activationDate"), "冻结激活日期")
        cutoff = strict_system_date(record.get("cutoffDate"), "冻结截止日")
        if record.get('renewal') is not None and record['renewal'].get('version') == 2:
            # Reconstruct this human-authorized rule independently of the writer.
            proof = record['renewal']
            if (proof.get('authorization') != 'current-user-verified-renewal-source'
                or proof.get('rule') != 'verified_source_cutoff_else_positive_balance_prior_effective_plus_29'
                or proof.get('queryIccid') != key or proof.get('activationDate') != record.get('activationDate')
                or proof.get('cutoffDate') != record.get('cutoffDate')
                or proof.get('identitySourceKind') != 'verified_system_export'
                or not re.fullmatch(r'[a-f0-9]{64}', proof.get('sourceSha256', ''))
                or not re.fullmatch(r'[a-f0-9]{64}', proof.get('identitySourceSha256', ''))
                or not re.fullmatch(r'\d{10}', proof.get('phone', ''))
                or not re.fullmatch(r'\d{8}T\d{6}\+0800', proof.get('runId', ''))
                or not re.fullmatch(r'\d{8}T\d{6}\+0800', proof.get('identityRunId', ''))
                or type(proof.get('sourceRow')) is not int or proof['sourceRow'] < 2
                or type(proof.get('identitySourceRow')) is not int or proof['identitySourceRow'] < 2
                or not proof.get('slot') or type(proof.get('planType')) is not int
                or proof['planType'] not in (29, 30) or type(proof.get('extensionDays')) is not int):
                raise RuntimeError('独立人工续费身份来源校验失败')
            amount = Decimal(proof['balanceCad'])
            base = strict_system_date(proof.get('baseCutoffDate'), '续费前F库截止日')
            if (activation is None or cutoff is None or base is None
                or not amount.is_finite() or amount < 0 or base - activation != timedelta(days=29)):
                raise RuntimeError('独立人工续费原周期或余额校验失败')
            if proof.get('dateMode') == 'source_explicit':
                expected = strict_system_date(proof.get('sourceCutoff'), '人工续费源截止日')
                if proof.get('extensionDays') != 0 or proof.get('provisional') is not False:
                    raise RuntimeError('独立人工续费不能追加天数')
            elif proof.get('dateMode') == 'positive_balance_prior_plus_29':
                prior = strict_system_date(proof.get('priorEffectiveCutoff'), '推算前有效截止日')
                if (prior is None or amount <= 0 or proof.get('sourceCutoff') is not None
                    or proof.get('extensionDays') != 29 or proof.get('provisional') is not True):
                    raise RuntimeError('独立余额临时推算校验失败')
                expected = prior + timedelta(days=29)
            else:
                raise RuntimeError('独立人工续费日期模式未知')
            if expected is None or cutoff != expected:
                raise RuntimeError('独立人工续费日期重算不一致')
        elif record.get('renewal') is not None:
            # Independently recompute the authorized renewal, not the generator's result.
            proof = record['renewal']
            if (not isinstance(proof, dict) or proof.get('version') != 1
                or proof.get('authorization') != 'current-user-renewal-source'
                or proof.get('rule') != 'current_source_cutoff_plus_30_if_balance_gte_30_else_source_cutoff'
                or proof.get('queryIccid') != key or proof.get('activationDate') != record.get('activationDate')
                or proof.get('cutoffDate') != record.get('cutoffDate')
                or proof.get('status') not in {'充值成功', '不需要充值'}
                or not re.fullmatch(r'[a-f0-9]{64}', proof.get('sourceSha256', ''))
                or not re.fullmatch(r'\d{10}', proof.get('phone', ''))
                or not re.fullmatch(r'\d{8}T\d{6}\+0800', proof.get('runId', ''))
                or type(proof.get('sourceRow')) is not int or proof['sourceRow'] < 2
                or not str(proof.get('backendId', '')).isdigit() or not proof.get('slot')):
                raise RuntimeError('独立日期缓存续费来源凭证无效')
            amount = Decimal(proof['balanceCad'])
            source = strict_system_date(proof.get('sourceCutoff'), '续费源截止日')
            base = strict_system_date(proof.get('baseCutoffDate'), '续费前F库截止日')
            extension = 30 if amount >= 30 else 0
            if (activation is None or cutoff is None or source is None or base is None
                or not amount.is_finite() or amount < 0
                or type(proof.get('extensionDays')) is not int or proof['extensionDays'] != extension
                or base - activation != timedelta(days=29) or cutoff != source + timedelta(days=extension)):
                raise RuntimeError('独立日期缓存续费截止日计算不一致')
        elif activation is None or cutoff is None or cutoff - activation != timedelta(days=29):
            raise RuntimeError("独立日期缓存日期不满足激活日加29天")
        indexed[key] = {"激活日期": activation, "截止日": cutoff}
    return {"path": str(path.resolve()), "sha256": digest, "run_id": run_id,
            "created_at": payload["createdAt"], "source_database": payload.get("sourceDatabase"),
            "source_exists": payload["sourceExists"], "record_count": len(indexed), "index": indexed}


def load_google_date_source(path: Path) -> dict:
    """Read original Google inputs; do not trust generator-resolved date values."""
    indexed = {}
    iccid_records = []
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for title in REQUIRED_ASSOC_SHEETS:
            sheet = book[title]
            rows = sheet.iter_rows(values_only=True)
            headers = {text(value): index for index, value in enumerate(next(rows))}
            for values in rows:
                def cell(name):
                    return values[headers[name]] if name in headers and headers[name] < len(values) else None
                slot = normalize_slot(cell("卡槽"))
                if not slot:
                    continue
                iccid_records.append({"slot": slot, "iccid": cell("iccid") or cell("ICCID")})
                score = sum(weight for field, weight in (("激活日期", 8), ("預計過期日", 4),
                            ("上卡日期", 2), ("電信商/價格", 1), ("卡槽", 1))
                            if text(cell(field)) and text(cell(field)).upper() not in FORMULA_ERRORS)
                if slot not in indexed or score > indexed[slot]["score"]:
                    indexed[slot] = {"score": score, "激活日期": date_value(cell("激活日期")),
                                     "截止日": date_value(cell("預計過期日"))}
    finally:
        book.close()
    return {"path": str(path.resolve()), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "index": indexed, "iccid_records": iccid_records}


def load_chatr_google_snapshot(path: Path, expected_sha256: str, run_id: str) -> dict:
    """Independently reread the original workbook and authenticate its date snapshot."""
    if not run_id or not re.fullmatch(r"[a-fA-F0-9]{64}", expected_sha256):
        raise RuntimeError("Chatr独立复核缺少快照SHA256或runId")
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != expected_sha256.lower():
        raise RuntimeError("Chatr日期快照SHA256不匹配")
    payload = json.loads(raw)
    records = payload.get("records")
    if (payload.get("schemaVersion") != 1 or payload.get("status") != "passed"
            or payload.get("source") != "google_chatr_dates" or payload.get("runId") != run_id
            or not payload.get("createdAt")
            or not isinstance(records, list) or type(payload.get("recordCount")) is not int
            or payload["recordCount"] != len(records)):
        raise RuntimeError("Chatr日期快照契约或来源不匹配")
    source_path = Path(payload.get("sourcePath") or "").resolve()
    source_hash = payload.get("sourceSha256", "")
    if not re.fullmatch(r"[a-f0-9]{64}", source_hash) or not source_path.is_file():
        raise RuntimeError("Chatr原始Google文件SHA256不匹配")
    source_bytes = source_path.read_bytes()
    if hashlib.sha256(source_bytes).hexdigest() != source_hash:
        raise RuntimeError("Chatr原始Google文件SHA256不匹配")
    try:
        source_sheets = bound_room_sheets(source_bytes, payload.get("sourceSheets"))
    except ValueError as exc:
        raise RuntimeError(f"Chatr独立复核来源范围错误: {exc}") from exc
    indexed, original, source_rows = {}, {}, {}
    def iccid_key(value):
        if not isinstance(value, str):
            return None
        normalized = re.sub(r"[Ff]$", "", re.sub(r"\s+", "", value))
        return normalized if re.fullmatch(r"\d{18,22}", normalized, re.ASCII) else None
    for row in records:
        if not isinstance(row, dict):
            raise RuntimeError("Chatr日期快照记录无效")
        slot = row.get("normalizedSlot")
        if (not isinstance(slot, str) or not slot or normalize_slot(slot) != slot or slot in indexed
                or normalize_slot(row.get("rawSlot")) != slot or normalize_carrier(row.get("carrier")) != "chatr"
                or row.get("sourceSheet") not in source_sheets
                or type(row.get("sourceRow")) is not int or row["sourceRow"] < 2):
            raise RuntimeError("Chatr日期快照卡槽、运营商或来源行无效")
        dates = {}
        for source_field, field in (("activationDate", "激活日期"), ("cutoffDate", "截止日")):
            value = row.get(source_field)
            if value is None:
                dates[field] = None
            else:
                if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:T\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?)?", value):
                    raise RuntimeError("Chatr日期快照非空字段不是有效ISO日期")
                try:
                    dates[field] = datetime.fromisoformat(value)
                except ValueError as error:
                    raise RuntimeError("Chatr日期快照日期无效") from error
        query_iccid = row.get("queryIccid")
        if query_iccid is not None and (query_iccid != iccid_key(query_iccid)
                or not any(isinstance(item, dict) and query_iccid == iccid_key(item.get("iccid"))
                           for item in row.get("provenance", [row]))):
            raise RuntimeError("Chatr日期快照ICCID来源不匹配")
        indexed[slot] = {**dates, "query_iccid": query_iccid}
    book = load_workbook(BytesIO(source_bytes), read_only=True, data_only=True)
    formula_book = load_workbook(BytesIO(source_bytes), read_only=True, data_only=False)
    try:
        for title in source_sheets:
            if title not in book.sheetnames:
                raise RuntimeError("Chatr原始Google缺少机房页签")
            rows = book[title].iter_rows()
            headers = [text(cell.value) for cell in next(rows)]
            required = ("卡槽", "電信商/價格", "激活日期", "預計過期日")
            if any(headers.count(name) != 1 for name in required):
                raise RuntimeError("Chatr原始Google必要表头缺失或重复")
            columns = {name: number for number, name in enumerate(headers)}
            iccid_columns = [name for name in ("ICCID", "iccid") if name in columns]
            if len(iccid_columns) > 1:
                raise RuntimeError("Chatr原始Google身份表头重复")
            formula_rows = formula_book[title].iter_rows(min_row=2)
            for row_number, (value_cells, formula_values) in enumerate(zip(rows, formula_rows), start=2):
                values = tuple(cell.value for cell in value_cells)
                def cell(name):
                    return values[columns[name]] if name in columns and columns[name] < len(values) else None
                if normalize_carrier(cell("電信商/價格")) != "chatr":
                    continue
                slot = normalize_slot(cell("卡槽"))
                if not re.fullmatch(r"[0-9]+_[0-9]+", slot):
                    continue
                item = {"query_iccid": iccid_key(cell(iccid_columns[0])) if iccid_columns else None}
                for source_field, field in (("激活日期", "激活日期"), ("預計過期日", "截止日")):
                    value = cell(source_field)
                    formula_cell = formula_values[columns[source_field]]
                    if formula_cell.data_type == "f" and value is None and value_cells[columns[source_field]].data_type not in {"str", "inlineStr"}:
                        raise RuntimeError("Chatr原始Google日期公式没有可核验缓存值")
                    try:
                        parsed = independent_chatr_source_timestamp(value, book.epoch)
                    except (ValueError, OverflowError, TypeError):
                        raise RuntimeError(f"Chatr原始Google日期无效: {title}:{row_number}:{source_field}")
                    item[field] = parsed
                source_rows[(title, row_number)] = {**item, "slot": slot, "raw_slot": cell("卡槽"), "carrier": cell("電信商/價格")}
                if slot in original:
                    old = original[slot]
                    for field in ("激活日期", "截止日", "query_iccid"):
                        if old[field] is not None and item[field] is not None and old[field] != item[field]:
                            raise RuntimeError("Chatr原始Google同卡槽日期或ICCID冲突")
                        if old[field] is None:
                            old[field] = item[field]
                else:
                    original[slot] = item
    finally:
        book.close()
        formula_book.close()
    if indexed != original:
        raise RuntimeError("Chatr日期快照与独立读取原始Google日期不一致")
    claimed_rows = set()
    for row in records:
        provenance = row.get("provenance", [row])
        if not isinstance(provenance, list) or not provenance:
            raise RuntimeError("Chatr日期快照缺少来源行证明")
        if (not isinstance(provenance[0], dict) or any(row.get(name) != provenance[0].get(name)
                for name in ("sourceSheet", "sourceRow", "rawSlot", "carrier", "iccid", "phone"))):
            raise RuntimeError("Chatr日期快照首条来源行证明不一致")
        for item in provenance:
            key = (item.get("sourceSheet"), item.get("sourceRow")) if isinstance(item, dict) else None
            actual = source_rows.get(key)
            if actual is None or key in claimed_rows or actual["slot"] != row["normalizedSlot"]:
                raise RuntimeError("Chatr日期快照来源行缺失、重复或卡槽不符")
            claimed_rows.add(key)
            if (item.get("normalizedSlot") != actual["slot"] or item.get("rawSlot") != actual["raw_slot"]
                    or item.get("carrier") != actual["carrier"] or item.get("queryIccid") != actual["query_iccid"]):
                raise RuntimeError("Chatr日期快照来源行身份不符")
            for source_field, field in (("activationDate", "激活日期"), ("cutoffDate", "截止日")):
                if timestamp_value(item.get(source_field)) != actual[field]:
                    raise RuntimeError("Chatr日期快照来源行日期不符")
    if claimed_rows != set(source_rows):
        raise RuntimeError("Chatr日期快照未覆盖全部有效原始行")
    return {"enabled": True, "status": "passed", "run_id": run_id,
            "snapshot_path": str(path.resolve()), "snapshot_sha256": digest,
            "source_path": str(source_path), "source_sha256": source_hash,
            "scope": "chatr_blank_activation_and_cutoff_only",
            "source": "current_run_system_then_fresh_google", "record_count": len(indexed), "index": indexed}


def independent_chatr_source_timestamp(value, epoch):
    """Strict original-cell parser, independent of the snapshot producer."""
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if type(value) in (int, float):
        if not math.isfinite(value):
            raise ValueError("nonfinite date")
        if value == int(value) and re.fullmatch(r"(?:19|20)\d{6}", str(int(value))):
            return datetime.strptime(str(int(value)), "%Y%m%d")
        if 20000 <= value <= 60000:
            return from_excel(value, epoch)
        raise ValueError("numeric date outside supported range")
    if not isinstance(value, str):
        raise ValueError("date must be text, date or serial")
    raw = unicodedata.normalize("NFKC", value.strip()).strip("'\"")
    raw = raw.replace("年", "-").replace("月", "-").replace("日", "")
    if re.fullmatch(r"\d{8}", raw):
        return datetime.strptime(raw, "%Y%m%d")
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})([T ]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?)?", raw)
    if match:
        year, month, day, suffix = match.groups()
        return datetime.fromisoformat(f"{int(year):04d}-{int(month):02d}-{int(day):02d}{suffix or ''}")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?(?:Z|[+-]\d{2}:\d{2})", raw):
        return datetime.fromisoformat(raw).replace(tzinfo=None)
    raise ValueError("invalid or ambiguous source date")


def review_resolved_dates(raw_headers, raw_values, main_headers, main_rows, system, cache, google, google_enabled=True, chatr_google=None) -> dict:
    """Recompute field priorities from four frozen sources, independently."""
    raw_index = {name: index for index, name in enumerate(raw_headers)}
    final_index = {name: index for index, name in enumerate(main_headers)}
    counts = {name: {source: 0 for source in ("system_export", "system_inventory_legacy", "local_cache", "google", "missing")}
              for name in ("激活日期", "截止日")}
    if chatr_google:
        for values in counts.values():
            values["google_chatr"] = 0
    chatr_matched_rows = chatr_filled_rows = 0
    chatr_iccid_conflicts = []
    chatr_rows = chatr_missing_activation_rows = chatr_missing_cutoff_rows = 0
    expected_rows, mismatches, missing_scopes = [], [], []
    archived_ids = set()
    creation_fallback_rows = []
    blank_system_rows = local_filled_rows = 0
    for row_number, (raw, final) in enumerate(zip(raw_values, main_rows), start=2):
        kooda = normalize_carrier(raw[raw_index["opr"]]) == "koodo"
        key = system_kooda_key(raw[raw_index["卡槽"]], raw[raw_index["电话"]])
        inventory = (find_system_record(system, dict(zip(raw_headers, raw)), row_number) or {}) if kooda else {}
        if inventory and key not in system['index'] and not system.get('projection'):
            backend_id = str(inventory['raw']['id'])
            if backend_id in archived_ids:
                mismatches.append({'row': row_number, 'field': 'identity',
                                   'reason': 'duplicate_archived_backend_identity'})
            archived_ids.add(backend_id)
        iccid = re.sub(r"[Ff]$", "", text(inventory.get("iccid")))
        cached = cache["index"].get(iccid, {})
        associated = google["index"].get(normalize_slot(raw[raw_index["卡槽"]]), {}) if google_enabled else {}
        chatr = normalize_carrier(raw[raw_index["opr"]]) == "chatr"
        chatr_rows += int(chatr)
        chatr_card = chatr_google["index"].get(normalize_slot(raw[raw_index["卡槽"]]), {}) if chatr_google and chatr else {}
        conflict = None
        if chatr_card:
            chatr_matched_rows += 1
            exported_iccid = raw[raw_index["ICCID"]] if "ICCID" in raw_index else None
            exported_key = re.sub(r"[Ff]$", "", re.sub(r"\s+", "", exported_iccid)) if isinstance(exported_iccid, str) else ""
            if re.fullmatch(r"\d{18,22}", exported_key, re.ASCII) and chatr_card["query_iccid"] and exported_key != chatr_card["query_iccid"]:
                conflict = {
                    "row": row_number, "phone": text(raw[raw_index["电话"]]), "slot": text(raw[raw_index["卡槽"]]),
                    "system_iccid": exported_iccid, "google_iccid": chatr_card["query_iccid"],
                    "system_activation": (timestamp_value(raw[raw_index["激活日期"]]).isoformat()
                                          if "激活日期" in raw_index and timestamp_value(raw[raw_index["激活日期"]]) else ""),
                    "system_cutoff": (timestamp_value(raw[raw_index["截止日"]]).isoformat()
                                      if "截止日" in raw_index and timestamp_value(raw[raw_index["截止日"]]) else ""),
                    "google_activation": chatr_card["激活日期"].isoformat() if chatr_card["激活日期"] else "",
                    "google_cutoff": chatr_card["截止日"].isoformat() if chatr_card["截止日"] else "",
                    "decision": ICCID_CONFLICT_DECISION,
                }
                chatr_card = {}
                if final[final_index["ICCID"]] != exported_iccid:
                    mismatches.append({"row": row_number, "field": "ICCID", "reason": "system_iccid_changed"})
        expected, sources, timestamps = {}, {}, {}
        for name, inventory_name in (("激活日期", "activation"), ("截止日", "cutoff")):
            value = date_value(raw[raw_index[name]]) if name in raw_index else None
            timestamp = timestamp_value(raw[raw_index[name]]) if name in raw_index else None
            if chatr_google and chatr and name in raw_index and text(raw[raw_index[name]]) and timestamp is None:
                raise RuntimeError("Chatr系统导出日期非空无效，不能回退Google")
            source = "system_export" if value else "missing"
            if kooda and name not in raw_index and inventory.get(inventory_name):
                value, source = inventory[inventory_name], "system_inventory_legacy"
            if kooda and name == "激活日期" and value is None:
                blank_system_rows += 1
            if value is None:
                fallback = cached.get(name) if kooda else chatr_card.get(name) if chatr_card else associated.get(name)
                if fallback is not None:
                    value, source = date_value(fallback), "local_cache" if kooda else "google_chatr" if chatr_card else "google"
                    timestamp = timestamp_value(fallback)
            timestamps[name] = timestamp or timestamp_value(value)
            expected[name], sources[name] = value, source
        if (all(value is None for value in expected.values())
                and all(not text(raw[raw_index[name]]) for name in expected if name in raw_index)):
            creation = timestamp_value(raw[raw_index["创建时间"]]) if "创建时间" in raw_index else None
            if creation is not None:
                from datetime import timedelta
                for name, offset in (("激活日期", 0), ("截止日", 29)):
                    timestamps[name] = creation + timedelta(days=offset)
                    expected[name] = timestamps[name].date()
                    sources[name] = "creation_date_fallback"
                creation_fallback_rows.append({"row": row_number, "phone": text(raw[raw_index["电话"]]),
                    "activation": timestamps["激活日期"].isoformat(), "cutoff": timestamps["截止日"].isoformat()})
        for name, value in expected.items():
            source = sources[name]
            counts[name][source] = counts[name].get(source, 0) + 1
            actual_raw = final[final_index[name]]
            actual = date_value(actual_raw)
            if actual != value or (text(actual_raw) and actual is None):
                mismatches.append({"row": row_number, "field": name, "source": source,
                                   "expected": str(value) if value else "", "actual": str(actual_raw)})
            if chatr_google and chatr and timestamp_value(actual_raw) != timestamps[name]:
                mismatches.append({"row": row_number, "field": name, "source": source, "reason": "chatr_authoritative_timestamp_changed"})
        if conflict is not None:
            conflict.update(final_activation=timestamps["激活日期"].isoformat() if timestamps["激活日期"] else "",
                            final_cutoff=timestamps["截止日"].isoformat() if timestamps["截止日"] else "",
                            activation_source=sources["激活日期"], cutoff_source=sources["截止日"])
            chatr_iccid_conflicts.append(conflict)
        local_filled_rows += int("local_cache" in sources.values())
        if kooda and (expected["激活日期"] is None or expected["截止日"] is None):
            missing_scopes.append({"row": row_number, "phone": text(raw[raw_index["电话"]]),
                "slot": text(raw[raw_index["卡槽"]]), "iccid": inventory.get("iccid"),
                "queryIccid": iccid if re.fullmatch(r"\d{18,22}", iccid, re.ASCII) else "",
                "activation_missing": expected["激活日期"] is None, "cutoff_missing": expected["截止日"] is None})
        chatr_filled_rows += int("google_chatr" in sources.values())
        if chatr:
            chatr_missing_activation_rows += int(expected["激活日期"] is None)
            chatr_missing_cutoff_rows += int(expected["截止日"] is None)
        expected_rows.append({**expected, "_sources": sources, "_timestamps": timestamps})
    result = {"source": "system_export_then_frozen_local_or_chatr_google" if chatr_google else "system_export_then_frozen_local_or_google" if google_enabled else "system_export_then_frozen_local", "field_counts": counts,
            "cache": {key: value for key, value in cache.items() if key != "index"},
            "association": {key: value for key, value in google.items() if key not in ("index", "iccid_records")} if google_enabled else None,
            "kooda_system_blank_activation_rows": blank_system_rows,
            "creation_date_fallback_rows": creation_fallback_rows,
            "kooda_local_cache_filled_rows": local_filled_rows,
            "missing_date_scopes": missing_scopes,
            "missing_activation_rows": sum(row["activation_missing"] for row in missing_scopes),
            "missing_cutoff_rows": sum(row["cutoff_missing"] for row in missing_scopes),
            "missing_either_date_rows": len(missing_scopes),
            "new_query_results_apply_next_export": True,
            "system_inventory_fallback_only_when_export_column_absent": True,
            "mismatch_count": len(mismatches), "mismatches": mismatches[:50],
            "matches": not mismatches and len(raw_values) == len(main_rows), "expected_rows": expected_rows}
    if chatr_google:
        result["chatr_google_dates"] = {key: value for key, value in chatr_google.items() if key != "index"}
        result["chatr_google_dates"].update(field_counts={name: values["google_chatr"] for name, values in counts.items()},
                                          matched_rows=chatr_matched_rows, filled_rows=chatr_filled_rows,
                                          chatr_rows=chatr_rows, unmatched_rows=chatr_rows - chatr_matched_rows,
                                          missing_activation_rows=chatr_missing_activation_rows,
                                          missing_cutoff_rows=chatr_missing_cutoff_rows,
                                          iccid_conflict_count=len(chatr_iccid_conflicts), iccid_conflicts=chatr_iccid_conflicts,
                                          iccid_conflict_sheet=ICCID_CONFLICT_SHEET, iccid_conflict_policy="system_wins_skip_google_dates")
    return result


def expiry_status(value, reference: date) -> str:
    expiry = date_value(value)
    if expiry is None:
        return ""
    days_until_expiry = (expiry - reference).days
    if days_until_expiry <= 0:
        return "已过期"
    if days_until_expiry <= 2:
        return "快过期"
    if days_until_expiry <= 6:
        return "6天"
    return ""


def parse_mmdd(value: str, reference: date) -> date:
    month, day = map(int, value.split("-"))
    candidate = date(reference.year, month, day)
    if candidate > reference + timedelta(days=1):
        candidate = date(reference.year - 1, month, day)
    return candidate


def aggregate_rate(points, count):
    selected = points[-count:]
    if len(selected) < count or any(send is None or reply is None for send, reply in selected):
        return None
    denominator = sum(send for send, _ in selected)
    if denominator <= 0:
        return None
    return sum(reply for _, reply in selected) / denominator


def aggregate_sends(points, count):
    selected = points[-count:]
    if len(selected) < count or any(send is None for send, _ in selected):
        return None
    return sum(send for send, _ in selected)


def trend_class(points):
    rates = [None if send is None or send <= 0 or reply is None else reply / send for send, reply in points]
    for count in (4, 3, 2):
        selected = rates[-count:]
        if len(selected) < count or any(value is None for value in selected):
            continue
        if all(selected[index] < selected[index - 1] for index in range(1, len(selected))):
            return f"持续{count}天降低"
    return ""


def longest_run(values) -> int:
    longest = 0
    current = 0
    for value in values:
        current = current + 1 if value else 0
        longest = max(longest, current)
    return longest


def trailing_run(values) -> int:
    count = 0
    for value in reversed(values):
        if not value:
            break
        count += 1
    return count


def low_class(days: int) -> str:
    if days >= 4:
        return "连续4天低于3%"
    if days >= 2:
        return f"连续{days}天低于3%"
    return ""


def zero_class(days: int) -> str:
    if days >= 4:
        return "连续4天0回复"
    if days >= 2:
        return f"连续{days}天0回复"
    return ""


def sheet_rows(sheet) -> int:
    count = 0
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number == 1:
            continue
        if any(value not in (None, "") for value in row):
            count += 1
    return count


def read_header(sheet):
    return [text(cell.value) for cell in sheet[1]]


def unique_phones(sheet) -> set[str]:
    headers = read_header(sheet)
    phone_index = next((index for index, name in enumerate(headers) if name in {"电话", "phone", "手机号"}), None)
    if phone_index is None:
        return set()
    result = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        key = phone_key(row[phone_index] if phone_index < len(row) else None)
        if key:
            result.add(key)
    return result


def phone_list(sheet) -> list[str]:
    headers = read_header(sheet)
    phone_index = next((index for index, name in enumerate(headers) if name in {"电话", "phone", "手机号"}), None)
    if phone_index is None:
        return []
    return [
        phone_key(row[phone_index] if phone_index < len(row) else None)
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if phone_key(row[phone_index] if phone_index < len(row) else None)
    ]


def keyed_rows(sheet, key_headers: tuple[str, ...]) -> list[str]:
    headers = read_header(sheet)
    indices = [headers.index(header) for header in key_headers]
    result = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = []
        for header, index in zip(key_headers, indices):
            value = row[index] if index < len(row) else None
            values.append(normalize_slot(value) if "卡槽" in header else phone_key(value) if header == "电话" else text(value))
        result.append("|".join(values))
    return result


def cell_is_solid_yellow(cell) -> bool:
    color = str(cell.fill.fgColor.rgb or "").upper()
    return cell.fill.fill_type == "solid" and color.endswith("FFFF00")


def cell_is_solid_task_f_missing_red(cell) -> bool:
    color = str(cell.fill.fgColor.rgb or "").upper()
    return cell.fill.fill_type == "solid" and color.endswith("FFC7CE")


def timestamp_value(value):
    """Preserve times for the strict 72-hour number-change boundary."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 20000 <= value <= 60000:
        return from_excel(value)
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?)?", value.strip()):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            return None
    day = date_value(value)
    if day is None:
        return None
    clock = re.search(r"[T ](\d{1,2}):(\d{1,2})(?::(\d{1,2}))?", text(value))
    try:
        return datetime(day.year, day.month, day.day, *(int(item or 0) for item in clock.groups())) if clock else datetime.combine(day, datetime.min.time())
    except ValueError:
        return None


def retained_excel_value_matches(source, actual) -> bool:
    """Compare retained Excel values, not parser-specific Python number classes.

    OOXML ``t=n`` represents both 106 and 106.0; openpyxl's round trip can
    change float to int without changing the numeric cell type or value. Empty
    source strings likewise become blank inline-string cells. No coercion is
    allowed for non-empty text, booleans, dates, or unequal numeric values.
    In particular, this deliberately does not use float conversion or a
    tolerance, which could hide large-integer precision loss or changed counts.
    """
    if source in (None, "") and actual in (None, ""):
        return True
    if type(source) in (int, float) and type(actual) in (int, float):
        if isinstance(source, float) and not math.isfinite(source):
            return False
        if isinstance(actual, float) and not math.isfinite(actual):
            return False
        return source == actual
    return type(source) is type(actual) and source == actual


def review_no_google_output(book, raw_headers, raw_values, headers, rows, process, dates, summary, resolved_dates):
    """Check retained content directly against this run's immutable source data."""
    dependency = {"enabled": False, "status": "skipped", "reason": "user_disabled",
                  "unused_port_device_scope": "system_source_only"}
    source_index = {name: index for index, name in enumerate(raw_headers)}
    index = {name: column for column, name in enumerate(headers)}
    unexpected_sheets = [name for name in GOOGLE_ONLY_SHEETS if name in book.sheetnames
                         and not (google_difference_enabled() and name == GOOGLE_DIFFERENCE_SHEET)]
    unexpected_headers = [name for name in GOOGLE_ONLY_HEADERS if name in headers]
    unlisted_headers = read_header(book["已激活未上架"])
    unexpected_headers.extend(name for name in ("Google页签", "Google行号") if name in unlisted_headers)
    fields_mismatches, change_mismatches, expected_unlisted = [], [], []
    checked_sources = Counter()
    changed_rows = 0
    for row_number, (raw, actual, resolved) in enumerate(zip(raw_values, rows, resolved_dates), start=2):
        kooda = normalize_carrier(raw[source_index["opr"]]) == "koodo"
        slot = normalize_slot(raw[source_index["卡槽"]])
        device = slot.split("_", 1)[0] if re.match(r"^\d+_", slot) else ""
        room = room_for_device(device) if device else "未配置"
        expected_core = {"机房": room, "开卡日期": ""}
        for name, expected in expected_core.items():
            if text(actual[index[name]]) != expected:
                fields_mismatches.append({"row": row_number, "field": name, "expected": expected, "actual": text(actual[index[name]])})
        for name, column in source_index.items():
            if name in {"激活日期", "截止日", "系统导出激活日期", "系统导出截止日", "ICCID"} or name not in index:
                continue
            if not retained_excel_value_matches(raw[column], actual[index[name]]):
                fields_mismatches.append({"row": row_number, "field": name, "reason": "source_value_or_type_changed"})
        activation = timestamp_value(raw[source_index["激活日期"]]) if "激活日期" in source_index else None
        if activation is None:
            activation = resolved.get("_timestamps", {}).get("激活日期") or timestamp_value(resolved["激活日期"])
        cutoff = timestamp_value(raw[source_index["截止日"]]) if "截止日" in source_index else None
        cutoff = cutoff or resolved.get("_timestamps", {}).get("截止日") or timestamp_value(resolved["截止日"])
        for name, expected_date in (("激活日期", activation), ("截止日", cutoff)):
            if timestamp_value(actual[index[name]]) != expected_date:
                fields_mismatches.append({"row": row_number, "field": name, "reason": "authoritative_timestamp_changed"})
        created = timestamp_value(raw[source_index["创建时间"]])
        if kooda and resolved["激活日期"] is None and resolved["截止日"] is None:
            expected_change = "未换号"
        elif activation is None or created is None:
            expected_change = ""
        else:
            expected_change = "已换号" if abs((created - activation).total_seconds()) > 259200 else ""
            checked_sources["激活日期"] += 1
        changed_rows += expected_change == "已换号"
        if text(actual[index["换号状态"]]) != expected_change:
            change_mismatches.append({"row": row_number, "expected": expected_change, "actual": text(actual[index["换号状态"]])})
        if activation and text(raw[source_index["状态"]]) != "上架":
            source = resolved["_sources"]["激活日期"]
            expected_unlisted.append([
                slot.split("_", 1)[0] if "_" in slot else "未识别", room,
                raw[source_index["卡槽"]], slot, raw[source_index["电话"]], raw[source_index["运营商"]],
                text(raw[source_index["状态"]]) or "空白", "激活日期", activation, None,
                cutoff, source,
            ])
    actual_unlisted = [list(row) for row in book["已激活未上架"].iter_rows(min_row=2, values_only=True)]
    # Compare date cells by their retained timestamp and blank cells uniformly.
    def comparable(values):
        return ["" if value is None else value for value in values]
    detail_matches = len(actual_unlisted) == len(expected_unlisted) and all(
        comparable(actual) == comparable(expected) for actual, expected in zip(actual_unlisted, expected_unlisted))
    if not detail_matches:
        fields_mismatches.append({"field": "已激活未上架", "reason": "retained_details_differ_from_sources"})
    google_zero = all(
        evidence.get("field_counts", {}).get(name, {}).get("google") == 0
        for evidence in (dates, process.get("date_resolution", {})) for name in ("激活日期", "截止日"))
    scope_label = text(summary.get("系统无-谷歌表有设备覆盖范围"))
    receipt_matches = process.get("google_dependency") == dependency
    change_receipt_matches = (process.get("changed_rows") == changed_rows
        and process.get("change_checked_by_activation_rows") == checked_sources["激活日期"]
        and process.get("change_checked_by_open_rows") == 0 and process.get("changed_by_open_rows") == 0)
    return {
        "google_dependency_review": {**dependency, "unexpected_sheets": unexpected_sheets,
            "unexpected_headers": unexpected_headers, "google_date_source_rows_zero": google_zero,
            "receipt_matches": receipt_matches,
            "matches": receipt_matches and google_zero and dates.get("association") is None
                and not unexpected_sheets and not unexpected_headers and "仅本轮系统源主表" in scope_label and "1至64" in scope_label},
        "number_change_review": {"rule": "Kooda双日期缺失为未换号；创建时间与已有激活日期绝对差超过72小时为已换号；其它为空",
            "rows": len(rows), "mismatch_count": len(change_mismatches), "mismatches": change_mismatches[:50],
            "receipt_matches": change_receipt_matches, "matches": not change_mismatches and change_receipt_matches},
        "system_only_fields_review": {"mismatch_count": len(fields_mismatches), "mismatches": fields_mismatches[:50],
            "activated_unlisted_details_match": detail_matches, "matches": not fields_mismatches},
    }


def read_stage_json(name: str) -> dict:
    """Large production receipts travel by file, never through exec's ARG_MAX.

    Keep the small legacy environment input for offline callers. If a file was
    supplied, missing/invalid file data must fail rather than fall back to it.
    """
    path = os.environ.get(name + "_PATH")
    value = json.loads(Path(path).read_text(encoding="utf-8") if path else os.environ[name])
    if not isinstance(value, dict):
        raise RuntimeError(f"{name} must contain a JSON object")
    return value


def main() -> int:
    if len(sys.argv) != 3:
        raise RuntimeError("用法: opensms_verify_card_slot_fast.py <结果表.xlsx> <export4day.xlsx>")
    final_path = Path(sys.argv[1]).expanduser().resolve()
    raw_path = Path(sys.argv[2]).expanduser().resolve()
    process = read_stage_json("PROCESS_JSON")
    google_enabled = os.environ.get("OPENSMS_GOOGLE_ASSOCIATION_ENABLED", "1") != "0"
    chatr_path = os.environ.get("OPENSMS_CHATR_GOOGLE_DATES_PATH", "")
    chatr_hash = os.environ.get("OPENSMS_CHATR_GOOGLE_DATES_SHA256", "")
    if bool(chatr_path) != bool(chatr_hash) or (chatr_path and google_enabled):
        raise RuntimeError("Chatr专用快照路径与SHA256必须成对提供且完整Google关联关闭")
    enrichment = read_stage_json("ENRICH_JSON")
    reference = date.fromisoformat(os.environ.get("EXPORT4DAY_SERVER_DATE") or process["expiry_reference_date"])
    if not os.environ.get("OPENSMS_SYSTEM_KOODA_PATH"):
        raise RuntimeError("Kooda系统独立核验缺少本轮系统快照路径")
    if not os.environ.get("OPENSMS_REPORT_DATE_CACHE_PATH") or (google_enabled and not os.environ.get("OPENSMS_ASSOC_PATH")):
        raise RuntimeError("日期独立核验缺少本轮冻结缓存或Google关联原文件")
    with ThreadPoolExecutor(max_workers=4 if google_enabled or chatr_path else 3, thread_name_prefix="opensms-verify-input") as executor:
        system_future = executor.submit(load_system_kooda_snapshot,
            Path(os.environ["OPENSMS_SYSTEM_KOODA_PATH"]),
            os.environ.get("OPENSMS_SYSTEM_KOODA_SHA256", ""), os.environ.get("OPENSMS_RUN_ID", ""))
        cache_future = executor.submit(load_frozen_date_snapshot,
            Path(os.environ["OPENSMS_REPORT_DATE_CACHE_PATH"]),
            os.environ.get("OPENSMS_REPORT_DATE_CACHE_SHA256", ""), os.environ.get("OPENSMS_RUN_ID", ""))
        google_future = executor.submit(load_google_date_source, Path(os.environ["OPENSMS_ASSOC_PATH"])) if google_enabled else None
        chatr_future = executor.submit(load_chatr_google_snapshot, Path(chatr_path), chatr_hash, os.environ.get("OPENSMS_RUN_ID", "")) if chatr_path else None
        raw_future = executor.submit(load_workbook, raw_path, read_only=True, data_only=True)
        system_snapshot, report_cache, google_dates, raw_book = (
            system_future.result(), cache_future.result(), google_future.result() if google_future else None, raw_future.result())
        chatr_dates = chatr_future.result() if chatr_future else None
    if google_enabled and os.environ.get("OPENSMS_ASSOC_SHA256") != google_dates["sha256"]:
        raise RuntimeError("Google关联源SHA256未绑定或不匹配")

    for path in (final_path, raw_path):
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
            if bad_member:
                raise RuntimeError(f"XLSX损坏: {path} -> {bad_member}")

    final_book = load_workbook(final_path, read_only=True, data_only=True)
    raw_sheet = raw_book.active
    main_sheet = final_book["sheet1"]
    raw_headers = read_header(raw_sheet)
    main_headers = read_header(main_sheet)
    raw_index = {name: index for index, name in enumerate(raw_headers) if name}
    main_index = {name: index for index, name in enumerate(main_headers) if name}
    iccid_header_matches = system_kooda_iccid_header_matches(main_headers)
    if not iccid_header_matches:
        raise RuntimeError("结果表必须在末列包含唯一ICCID表头")

    field_names = {
        "发送窗口数": "send_windows",
        "回复窗口数": "reply_windows",
        "总发送量": "total_sends",
        "总回复量": "total_replies",
    }
    daily: dict[str, dict[str, object]] = {}
    for index, header in enumerate(raw_headers):
        match = re.fullmatch(r"(\d{2}-\d{2})(发送窗口数|回复窗口数|总发送量|总回复量)", header)
        if not match:
            continue
        mmdd, label = match.groups()
        definition = daily.setdefault(mmdd, {"date": parse_mmdd(mmdd, reference), "mmdd": mmdd})
        definition[field_names[label]] = index
    date_columns = sorted(
        (
            definition
            for definition in daily.values()
            if all(name in definition for name in field_names.values())
        ),
        key=lambda definition: definition["date"],
    )
    latest_four = date_columns[-4:]
    consecutive_dates = len(latest_four) == 4 and all(
        latest_four[index]["date"] - latest_four[index - 1]["date"] == timedelta(days=1)
        for index in range(1, len(latest_four))
    )

    raw_rows = {}
    raw_values = []
    duplicate_phones = set()
    for row in raw_sheet.iter_rows(min_row=2, values_only=True):
        raw_values.append(row)
        key = phone_key(row[raw_index["电话"]])
        if key:
            if key in raw_rows:
                duplicate_phones.add(key)
            raw_rows[key] = row

    if system_snapshot.get('projection'):
        from report_kooda_source import verify_projection_rows
        verify_projection_rows(system_snapshot['projection'], raw_headers, raw_values,
                               os.environ['OPENSMS_RUN_ID'], raw_path)

    rate_mismatches = []
    classification_mismatches = []
    actual_trends = Counter()
    expected_trends = Counter()
    low_counts = Counter()
    zero_counts = Counter()
    use_days_count = 0
    expiry_status_mismatches = []
    activation_age_mismatches = []
    task_f_missing_dates_new_card_eligible_rows = 0
    task_f_missing_dates_new_card_mismatches = []
    system_kooda_rows = 0
    system_kooda_matched_rows = 0
    system_kooda_mismatches = []
    expected_zero_total_runs: dict[str, int] = {}
    # Independently read every final main-sheet cell once, retaining values for
    # all subsequent checks and checking its full-row style in the same stream.
    # A read-only worksheet reparses its entire XML on every iter_rows call.
    main_rows = []
    retention_main_cells = [] if process.get("card_retention_report") else None
    usage_main_cells = []
    actual_task_f_full_red_rows = set()
    for row_number, cells in enumerate(
        main_sheet.iter_rows(min_row=2, max_col=main_sheet.max_column), start=2
    ):
        main_rows.append(tuple(cell.value for cell in cells))
        if retention_main_cells is not None:
            prior_key = '使用天数' if '使用天数' in main_index else '激活后天数'
            retention_main_cells.append(tuple(cells[main_index[key]] for key in ('电话', '需要保留时间', prior_key)))
        usage_main_cells.append(tuple(cells[main_index[key]] for key in ('使用天数', '激活后天数')))
        if cells and all(cell_is_solid_task_f_missing_red(cell) for cell in cells):
            actual_task_f_full_red_rows.add(row_number)
    expected_task_f_missing_red_rows = {
        row_number
        for row_number, row in enumerate(main_rows, start=2)
        if normalize_carrier(row[main_index["opr"]]) == "koodo"
        and date_value(row[main_index["激活日期"]]) is None
    }
    task_f_missing_red_review = {
        "expected_rows": len(expected_task_f_missing_red_rows),
        "actual_full_red_rows": len(actual_task_f_full_red_rows),
        "missing_full_red_rows": sorted(expected_task_f_missing_red_rows - actual_task_f_full_red_rows)[:50],
        "unexpected_full_red_rows": sorted(actual_task_f_full_red_rows - expected_task_f_missing_red_rows)[:50],
        "fill_rgb": "FFC7CE",
        "next_export_auto_clear_when_activation_exists": True,
        "matches": expected_task_f_missing_red_rows == actual_task_f_full_red_rows,
    }
    source_schema_review = review_export_schema_evidence(raw_headers, raw_values, main_headers, main_rows)
    date_resolution_review = review_resolved_dates(
        raw_headers, raw_values, main_headers, main_rows, system_snapshot, report_cache, google_dates, google_enabled, chatr_dates)
    resolved_date_rows = date_resolution_review.pop("expected_rows")
    date_resolution_review["source_export"] = {"path": str(raw_path), "sha256": hashlib.sha256(raw_path.read_bytes()).hexdigest()}
    generator_resolution = process.get("date_resolution", {})
    receipt_mismatches = []
    for key in ("source", "field_counts", "cache", "association", "source_export",
                "creation_date_fallback_rows",
                "kooda_system_blank_activation_rows", "kooda_local_cache_filled_rows",
                "missing_date_scopes", "missing_activation_rows", "missing_cutoff_rows", "missing_either_date_rows",
                "new_query_results_apply_next_export", "system_inventory_fallback_only_when_export_column_absent", "chatr_google_dates"):
        if generator_resolution.get(key) != date_resolution_review.get(key):
            receipt_mismatches.append(key)
    date_resolution_review["receipt_mismatch_fields"] = receipt_mismatches
    date_resolution_review["matches"] = date_resolution_review["matches"] and not receipt_mismatches
    chatr_review = None
    if chatr_dates:
        proof = date_resolution_review["chatr_google_dates"]
        chatr_review = {**proof, "receipt_matches": process.get("google_chatr_date_fallback") == proof,
                        "matches": date_resolution_review["matches"] and process.get("google_chatr_date_fallback") == proof}
        chatr_review["conflict_sheet_matches"] = review_conflict_sheet(final_book, proof["iccid_conflicts"])
        chatr_review["matches"] = chatr_review["matches"] and chatr_review["conflict_sheet_matches"]
    elif process.get("google_chatr_date_fallback") or generator_resolution.get("chatr_google_dates"):
        chatr_review = {"enabled": False, "matches": False, "reason": "snapshot_not_provided"}
    main_phone_counts = Counter(phone_key(row[main_index["电话"]]) for row in main_rows)
    duplicate_phones.update(key for key, count in main_phone_counts.items() if key and count > 1)
    for row_number, row in enumerate(main_rows, start=2):
        key = phone_key(row[main_index["电话"]])
        system_row_review = review_system_kooda_row(system_snapshot, {
            name: row[main_index[name]]
            for name in ("opr", "电话", "卡槽", "ICCID", "开卡日期", "激活日期", "截止日", "状态", "创建时间")
            if name in main_index
        }, resolved_dates=resolved_date_rows[row_number - 2], source_row=row_number)
        if system_row_review is not None:
            system_kooda_rows += 1
            system_kooda_matched_rows += int(system_row_review["matched"])
            if not system_row_review["matches"]:
                system_kooda_mismatches.append({"row": row_number, "phone": key, **system_row_review})
        new_card_review = review_task_f_missing_dates_new_card_row(
            row[main_index["opr"]], row[main_index["激活日期"]],
            row[main_index["截止日"]], row[main_index["换号状态"]],
        )
        if new_card_review["eligible"]:
            task_f_missing_dates_new_card_eligible_rows += 1
            if not new_card_review["matches"]:
                task_f_missing_dates_new_card_mismatches.append({
                    "row": row_number,
                    "phone": key,
                    "opr": text(row[main_index["opr"]]),
                    "expected": new_card_review["expected"],
                    "actual": new_card_review["actual"],
                })
        raw_row = raw_rows.get(key)
        if raw_row is None:
            rate_mismatches.append({"row": row_number, "reason": "raw_missing"})
            classification_mismatches.append({"row": row_number, "phone": key, "reason": "raw_missing"})
            continue
        points = [
            (
                number(raw_row[int(definition["send_windows"])]),
                number(raw_row[int(definition["reply_windows"])]),
            )
            for definition in latest_four
        ]
        expected_three_sends = aggregate_sends(points, 3) if consecutive_dates else None
        expected_three = aggregate_rate(points, 3) if consecutive_dates else None
        expected_four_sends = aggregate_sends(points, 4) if consecutive_dates else None
        expected_four = aggregate_rate(points, 4) if consecutive_dates else None
        expected_trend = trend_class(points) if consecutive_dates else ""
        if expected_trend:
            expected_trends[expected_trend] += 1
        actual_three_sends = number(row[main_index["3日发送窗口数"]])
        actual_three = number(row[main_index["3日回复率"]])
        actual_four_sends = number(row[main_index["4日发送窗口数"]])
        actual_four = number(row[main_index["4日回复率"]])
        actual_trend = text(row[main_index["回复率持续降低"]])

        def same_rate(expected, actual):
            if expected is None:
                return actual is None
            return actual is not None and abs(expected - actual) <= 1e-8

        if (expected_three_sends != actual_three_sends
                or not same_rate(expected_three, actual_three)
                or expected_four_sends != actual_four_sends
                or not same_rate(expected_four, actual_four)
                or expected_trend != actual_trend):
            rate_mismatches.append({
                "row": row_number,
                "phone": key,
                "expected3Sends": expected_three_sends,
                "actual3Sends": actual_three_sends,
                "expected3": expected_three,
                "actual3": actual_three,
                "expected4Sends": expected_four_sends,
                "actual4Sends": actual_four_sends,
                "expected4": expected_four,
                "actual4": actual_four,
                "expectedTrend": expected_trend,
                "actualTrend": actual_trend,
            })
        if actual_trend:
            actual_trends[actual_trend] += 1
        carrier_and_opr = f"{text(raw_row[raw_index['运营商']])} {text(raw_row[raw_index['opr']])}"
        use_family = bool(re.search(r"kini|koo|telus1", carrier_and_opr, re.IGNORECASE))
        expected_use_days = "2天" if use_family and sum(
            number(raw_row[int(definition["send_windows"])]) is not None
            and number(raw_row[int(definition["send_windows"])]) > 80
            for definition in date_columns
        ) >= 2 else ""
        low_points = []
        zero_points = []
        zero_total_points = []
        for definition in date_columns:
            total_sends = number(raw_row[int(definition["total_sends"])])
            reply_windows = number(raw_row[int(definition["reply_windows"])])
            send_windows = number(raw_row[int(definition["send_windows"])])
            total_replies = number(raw_row[int(definition["total_replies"])])
            low_points.append(
                total_sends is not None
                and total_sends > 50
                and reply_windows is not None
                and reply_windows / total_sends < 0.03
            )
            zero_points.append(send_windows is not None and send_windows > 40 and reply_windows == 0)
            zero_total_points.append(total_replies == 0)
        expected_low = (
            low_class(longest_run(low_points))
            if re.match(r"^A11(?:-|$)", text(raw_row[raw_index["运营商"]]), re.IGNORECASE)
            else ""
        )
        expected_zero = zero_class(longest_run(zero_points))
        expected_zero_total_runs[key] = trailing_run(zero_total_points)
        low_label = text(row[main_index["连续低于x%"]])
        zero_label = text(row[main_index["连续0回复"]])
        actual_use_days = text(row[main_index["koo是否2天"]])
        if (actual_use_days, low_label, zero_label) != (expected_use_days, expected_low, expected_zero):
            classification_mismatches.append({
                "row": row_number,
                "phone": key,
                "expected_use_days": expected_use_days,
                "actual_use_days": actual_use_days,
                "expected_low": expected_low,
                "actual_low": low_label,
                "expected_zero": expected_zero,
                "actual_zero": zero_label,
            })
        if low_label:
            low_counts[low_label] += 1
        if zero_label:
            zero_counts[zero_label] += 1
        if actual_use_days:
            use_days_count += 1
        activation = date_value(row[main_index["激活日期"]])
        expected_activation_age = (reference - activation).days if activation else None
        if expected_activation_age is not None and expected_activation_age < 0:
            expected_activation_age = None
        actual_activation_age = number(row[main_index["激活后天数"]])
        if actual_activation_age != expected_activation_age:
            activation_age_mismatches.append({
                "row": row_number,
                "phone": key,
                "activation_date": activation.isoformat() if activation else None,
                "expected": expected_activation_age,
                "actual": actual_activation_age,
            })
        expected_expiry_status = expiry_status(row[main_index["截止日"]], reference)
        actual_expiry_status = text(row[main_index["过期状态"]])
        if expected_expiry_status != actual_expiry_status:
            expiry_status_mismatches.append({
                "row": row_number,
                "phone": key,
                "expiry_date": str(date_value(row[main_index["截止日"]]) or ""),
                "expected": expected_expiry_status,
                "actual": actual_expiry_status,
            })

    required_sheets = [
        "处理汇总",
        "未匹配卡槽",
        "电信商_OPR预警",
        "谷歌有-系统无",
        "谷歌文档中有激活日期-系统无法使用",
        "系统无-谷歌表有",
        "A11连续低于3%",
        "已激活未上架",
        "未上架机器汇总",
    ]
    if not google_enabled:
        required_sheets = [name for name in required_sheets if name not in GOOGLE_ONLY_SHEETS]
    missing_sheets = [name for name in required_sheets if name not in final_book.sheetnames]
    if missing_sheets:
        raise RuntimeError("结果表缺少工作表: " + ", ".join(missing_sheets))

    google_missing_review = {"status": "skipped", "reason": "user_disabled"}
    if google_enabled:
        obsolete_google_comparison_sheets = [
            name
            for name in ("设备缺失端口", "Google有日期_主表缺失", "谷歌文档中有激活日期-系统没有")
            if name in final_book.sheetnames
        ]
        google_all_sheet = final_book["谷歌有-系统无"]
        google_activated_sheet = final_book["谷歌文档中有激活日期-系统无法使用"]
        expected_google_all_headers = [
            "Google页签", "Google行号", "卡槽", "归一化卡槽", "备注", "ICCID", "上卡日期", "电信商",
        ]
        expected_google_activated_headers = [
            "Google页签", "Google行号", "卡槽", "归一化卡槽", "备注", "ICCID", "手机号",
            "上卡日期", "电信商", "激活日期", "截止日",
        ]
        actual_google_all_headers = read_header(google_all_sheet)
        actual_google_activated_headers = read_header(google_activated_sheet)
        google_all_keys = keyed_rows(google_all_sheet, ("Google页签", "Google行号", "归一化卡槽"))
        google_activated_keys = keyed_rows(
            google_activated_sheet,
            ("Google页签", "Google行号", "归一化卡槽"),
        )
        google_all_key_counts = Counter(google_all_keys)
        google_activated_key_counts = Counter(google_activated_keys)
        activated_is_subset = all(
            count <= google_all_key_counts.get(key, 0)
            for key, count in google_activated_key_counts.items()
        )
        activated_header_index = {
            name: index for index, name in enumerate(actual_google_activated_headers) if name
        }
        # These small detail rows support dates, completeness and identifier styles.
        # Retaining cells also avoids a fresh XML parse for each .cell() style read.
        activated_cells = list(google_activated_sheet.iter_rows(min_row=2))
        activated_rows = [tuple(cell.value for cell in cells) for cells in activated_cells]
        activated_date_index = activated_header_index.get("激活日期")
        activated_dates_valid = activated_date_index is not None and all(
            date_value(row[activated_date_index] if activated_date_index < len(row) else None) is not None
            for row in activated_rows
            if any(value not in (None, "") for value in row)
        )
        required_detail_fields = ("ICCID", "手机号", "上卡日期", "电信商")
        activated_field_counts = {
            name: sum(
                text(row[activated_header_index[name]] if activated_header_index[name] < len(row) else None) != ""
                for row in activated_rows
                if any(value not in (None, "") for value in row)
            )
            for name in required_detail_fields
            if name in activated_header_index
        }
        expected_activated_field_counts = process.get("google_activated_missing_field_counts") or {}
        activated_field_counts_match = (
            set(activated_field_counts) == set(required_detail_fields)
            and all(
                activated_field_counts[name] == int(expected_activated_field_counts.get(name) or 0)
                for name in required_detail_fields
            )
        )
        activated_identifier_columns_are_text = (
            all(name in activated_header_index for name in ("ICCID", "手机号"))
            and all(
                cells[activated_header_index[name]].number_format == "@"
                for name in ("ICCID", "手机号")
                for cells in activated_cells
            )
        )
        google_all_header_index = {
            name: index for index, name in enumerate(actual_google_all_headers) if name
        }
        google_all_key_indices = [
            google_all_header_index[name]
            for name in ("Google页签", "Google行号", "归一化卡槽")
        ]
        activated_key_set = set(google_activated_keys)
        yellow_highlighted_rows = 0
        yellow_highlight_mismatches = []
        for row_number, row in enumerate(google_all_sheet.iter_rows(min_row=2), start=2):
            values = [
                normalize_slot(row[index].value) if "卡槽" in header else text(row[index].value)
                for header, index in zip(
                    ("Google页签", "Google行号", "归一化卡槽"),
                    google_all_key_indices,
                )
            ]
            key = "|".join(values)
            if key not in activated_key_set:
                continue
            non_yellow = [cell.coordinate for cell in row[: len(expected_google_all_headers)] if not cell_is_solid_yellow(cell)]
            if non_yellow:
                yellow_highlight_mismatches.append({"row": row_number, "key": key, "cells": non_yellow})
            else:
                yellow_highlighted_rows += 1
        google_missing_review = {
            "all_headers_match": actual_google_all_headers[:len(expected_google_all_headers)]
            == expected_google_all_headers,
            "activated_headers_match": actual_google_activated_headers[:len(expected_google_activated_headers)]
            == expected_google_activated_headers,
            "all_rows": len(google_all_keys),
            "activated_rows": len(google_activated_keys),
            "all_rows_match_receipt": len(google_all_keys)
            == process.get("google_all_missing_from_main_rows"),
            "activated_rows_match_receipt": len(google_activated_keys)
            == process.get("google_activated_missing_from_main_rows"),
            "activated_is_subset_of_all": activated_is_subset,
            "all_activated_dates_valid": activated_dates_valid,
            "required_detail_fields": list(required_detail_fields),
            "detail_field_counts": activated_field_counts,
            "detail_field_counts_match_receipt": activated_field_counts_match,
            "identifier_columns_preserved_as_text": activated_identifier_columns_are_text,
            "activated_rows_expected_yellow": len(google_activated_keys),
            "activated_rows_highlighted_yellow": yellow_highlighted_rows,
            "yellow_highlight_matches": (
                yellow_highlighted_rows == len(google_activated_keys)
                and not yellow_highlight_mismatches
            ),
            "yellow_highlight_mismatch_count": len(yellow_highlight_mismatches),
            "yellow_highlight_mismatch_sample": yellow_highlight_mismatches[:20],
            "obsolete_sheets_present": obsolete_google_comparison_sheets,
        }
        google_missing_review["matches"] = all(
            google_missing_review[name]
            for name in (
                "all_headers_match",
                "activated_headers_match",
                "all_rows_match_receipt",
                "activated_rows_match_receipt",
                "activated_is_subset_of_all",
                "all_activated_dates_valid",
                "detail_field_counts_match_receipt",
                "identifier_columns_preserved_as_text",
                "yellow_highlight_matches",
            )
        ) and not obsolete_google_comparison_sheets

    system_by_slot = {}
    unused_devices = set()
    unused_source_rows = main_rows if google_enabled else raw_values
    unused_source_index = main_index if google_enabled else raw_index
    for row in unused_source_rows:
        parsed = parse_device_port(row[unused_source_index["卡槽"]])
        if not parsed:
            continue
        device, port = parsed
        slot = f"{device}_{port}"
        unused_devices.add(device)
        system_by_slot.setdefault(slot, []).append(row)
    if google_enabled:
        google_slot_index = actual_google_all_headers.index("归一化卡槽")
        for row in google_all_sheet.iter_rows(min_row=2, values_only=True):
            parsed = parse_device_port(row[google_slot_index] if google_slot_index < len(row) else None)
            if parsed:
                unused_devices.add(parsed[0])
    expected_unused_rows = []
    for device in sorted(unused_devices, key=int):
        for port in range(1, 65):
            slot = f"{device}_{port}"
            hits = system_by_slot.get(slot) or []
            if not hits:
                expected_unused_rows.append((device, room_for_device(device), port, slot, "", "", "系统中无"))
                continue
            statuses = [text(row[unused_source_index["状态"]]) for row in hits]
            if statuses and all(status == "注册失败" for status in statuses):
                expected_unused_rows.append((
                    device,
                    room_for_device(device),
                    port,
                    slot,
                    unique_join(row[unused_source_index["电话"]] for row in hits),
                    unique_join(row[unused_source_index["运营商"]] for row in hits),
                    "注册失败",
                ))
    unused_sheet = final_book["系统无-谷歌表有"]
    unused_headers = read_header(unused_sheet)
    expected_unused_headers = list(PORT_HEADERS)
    actual_unused_rows = []
    for row in unused_sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in row):
            continue
        actual_unused_rows.append((
            text(row[0]),
            text(row[1]),
            int(number(row[2]) or 0),
            normalize_slot(row[3]),
            text(row[4]),
            text(row[5]),
            text(row[6]),
        ))
    unused_mismatches = [
        {"index": index + 2, "expected": expected, "actual": actual}
        for index, (expected, actual) in enumerate(zip(expected_unused_rows, actual_unused_rows))
        if expected != actual
    ]
    if len(expected_unused_rows) != len(actual_unused_rows):
        unused_mismatches.append({
            "reason": "row_count",
            "expected": len(expected_unused_rows),
            "actual": len(actual_unused_rows),
        })
    expected_missing_count = sum(row[6] == "系统中无" for row in expected_unused_rows)
    expected_failed_count = sum(row[6] == "注册失败" for row in expected_unused_rows)
    unused_port_review = {
        "headers_match": unused_headers == expected_unused_headers,
        "sheet_name_matches": process.get("unused_port_sheet_name") == "系统无-谷歌表有" and "未使用端口" not in final_book.sheetnames,
        "device_port_capacity": 64,
        "device_scope": "fresh_system_and_google_numeric_device_union" if google_enabled else "system_source_only",
        "receipt_device_scope_matches": process.get("unused_port_device_scope") == ("fresh_system_and_google_numeric_device_union" if google_enabled else "system_source_only"),
        "device_count": len(unused_devices),
        "expected_rows": len(expected_unused_rows),
        "actual_rows": len(actual_unused_rows),
        "system_missing_rows": expected_missing_count,
        "registration_failed_rows": expected_failed_count,
        "receipt_device_count_matches": int(process.get("unused_port_device_count") or 0) == len(unused_devices),
        "receipt_total_matches": int(process.get("unused_port_count") or 0) == len(expected_unused_rows),
        "receipt_system_missing_matches": int(process.get("system_missing_port_count") or 0) == expected_missing_count,
        "receipt_registration_failed_matches": int(process.get("registration_failed_port_count") or 0) == expected_failed_count,
        "mismatch_count": len(unused_mismatches),
        "mismatch_sample": unused_mismatches[:20],
    }
    unused_port_review["matches"] = (
        unused_port_review["headers_match"]
        and unused_port_review["sheet_name_matches"]
        and unused_port_review["receipt_device_scope_matches"]
        and unused_port_review["receipt_device_count_matches"]
        and unused_port_review["receipt_total_matches"]
        and unused_port_review["receipt_system_missing_matches"]
        and unused_port_review["receipt_registration_failed_matches"]
        and not unused_mismatches
    )

    cancelled_review_sheets = [
        "已过期_未进过期下架卡",
        "快过期_今天明天",
        "今日总回复量为0",
        "今日昨天总回复量均为0",
    ]
    unexpected_cancelled_sheets = [name for name in cancelled_review_sheets if name in final_book.sheetnames]

    carrier_warning_mismatches = []
    expected_carrier_warning_keys = []
    actual_carrier_warning_keys = []
    if google_enabled:
        for row_number, row in enumerate(main_rows, start=2):
            google_key = normalize_carrier(row[main_index["Google電信商/價格"]])
            opr_key = normalize_carrier(row[main_index["opr"]])
            expected_warning = bool(google_key and opr_key and google_key != opr_key)
            actual_warning = bool(text(row[main_index["OPR匹配预警"]]))
            key = f"{phone_key(row[main_index['电话']])}|{normalize_slot(row[main_index['卡槽']])}"
            if expected_warning:
                expected_carrier_warning_keys.append(key)
            if expected_warning != actual_warning:
                carrier_warning_mismatches.append({
                    "row": row_number,
                    "key": key,
                    "expected": expected_warning,
                    "actual": actual_warning,
                })
        actual_carrier_warning_keys = keyed_rows(final_book["电信商_OPR预警"], ("电话", "归一化卡槽"))

    expected_activated_keys = [
        f"{normalize_slot(row[main_index['卡槽']])}|{phone_key(row[main_index['电话']])}"
        for row in main_rows
        if (date_value(row[main_index["激活日期"]]) or date_value(row[main_index["开卡日期"]]))
        and text(row[main_index["状态"]]) != "上架"
    ]
    actual_activated_keys = keyed_rows(final_book["已激活未上架"], ("归一化卡槽", "电话"))
    expected_machine_counts = Counter(
        (key.split("|", 1)[0].split("_", 1)[0] or "未识别")
        for key in expected_activated_keys
    )
    machine_headers = read_header(final_book["未上架机器汇总"])
    machine_index = machine_headers.index("机器号")
    count_index = machine_headers.index("已激活未上架数")
    actual_machine_counts = {
        text(row[machine_index]): int(number(row[count_index]) or 0)
        for row in final_book["未上架机器汇总"].iter_rows(min_row=2, values_only=True)
    }
    expected_over_five_machines = sorted(
        machine for machine, count in expected_machine_counts.items() if count > 5
    )
    actual_over_five_machines = sorted(
        machine for machine, count in actual_machine_counts.items() if count > 5
    )

    summary_rows = list(final_book["处理汇总"].iter_rows(values_only=True))
    summary = {text(row[0]): row[1] if len(row) > 1 else None for row in summary_rows if row}
    actual_mapping = {
        "ICCID": text(summary.get("ICCID来源")),
        "开卡日期": text(summary.get("开卡日期来源")),
        "激活日期": text(summary.get("激活日期来源")),
        "截止日": text(summary.get("截止日来源")),
    }
    expected_mapping = {
        "ICCID": "Kooda/Telus1使用本轮系统号码管理cardStore原值（保留末尾F）；其它OPR保持原规则",
        "开卡日期": "Kooda/Telus1留空（不使用Google上卡日期或系统创建时间）；其它OPR使用Google上卡日期",
        "激活日期": "逐字段优先本轮系统导出激活日期；Kooda缺失用本轮冻结F库；Chatr及其它OPR缺失用Google激活日期",
        "截止日": "逐字段优先本轮系统导出截止日；Kooda缺失用本轮冻结F库；Chatr及其它OPR缺失用Google預計過期日；旧无日期列导出允许本轮系统库存",
    }
    system_kooda_source_review = {
        "source": "system",
        "source_kind": system_snapshot.get('source_kind'),
        "source_scope": system_snapshot.get('source_scope'),
        "snapshot_path": system_snapshot["path"],
        "snapshot_sha256": system_snapshot["sha256"],
        "run_id": system_snapshot["run_id"],
        "total_rows": system_snapshot["total_rows"],
        "unkeyed_rows": system_snapshot["unkeyed_rows"],
        "kooda_rows": system_kooda_rows,
        "matched_rows": system_kooda_matched_rows,
        "missing_activation_rows": sum(normalize_carrier(row[main_index["opr"]]) == "koodo" and expected["激活日期"] is None for row, expected in zip(main_rows, resolved_date_rows)),
        "missing_cutoff_rows": sum(normalize_carrier(row[main_index["opr"]]) == "koodo" and expected["截止日"] is None for row, expected in zip(main_rows, resolved_date_rows)),
        "missing_either_date_rows": sum(normalize_carrier(row[main_index["opr"]]) == "koodo" and (expected["激活日期"] is None or expected["截止日"] is None) for row, expected in zip(main_rows, resolved_date_rows)),
        "mismatch_count": len(system_kooda_mismatches),
        "mismatches": system_kooda_mismatches,
        "iccid_header_matches": iccid_header_matches,
        "export_source_schema_evidence": source_schema_review,
        "google_date_fallback_enabled": False,
        "task_f_database_date_fallback_enabled": True,
        "cutoff_rule": "system_export_direct_else_frozen_cache",
        "matches": (
            not system_kooda_mismatches and system_kooda_matched_rows == system_kooda_rows
            and iccid_header_matches and source_schema_review["matches"] and date_resolution_review["matches"]
        ),
    }
    # Keep the receipt key for existing readers, but never restore the retired
    # F/Google date requirement or derive a 29-day cutoff here.
    task_f_kooda_date_review = system_kooda_source_review
    actual_assoc_sheets = text(summary.get("关联表页签")).split("、")
    assoc_sheet_rows_text = text(summary.get("关联表页签行数"))
    formula_error_cells = []
    for sheet in final_book.worksheets:
        if sheet.title == main_sheet.title:
            # Use the independently read values, including the unnormalized
            # original header, so whitespace/case checks retain their semantics.
            rows = iter((tuple(cell.value for cell in main_sheet[1]), *main_rows))
        else:
            rows = sheet.iter_rows(values_only=True)
        for row_number, row in enumerate(rows, start=1):
            for column_number, value in enumerate(row, start=1):
                if isinstance(value, str) and value.strip().upper() in FORMULA_ERRORS:
                    formula_error_cells.append({
                        "sheet": sheet.title,
                        "row": row_number,
                        "column": column_number,
                        "value": value,
                    })
    numeric_text_review = review_numeric_text_ooxml(final_path)

    trend_counts_match = actual_trends == expected_trends
    trend_summary_matches = all(
        number(summary.get(f"{label}号码数")) == expected_trends.get(label, 0)
        for label in ("持续2天降低", "持续3天降低", "持续4天降低")
    )
    classification_mismatch_count = len(classification_mismatches)

    expected_headers = EXPECTED_HEADERS if google_enabled else [name for name in EXPECTED_HEADERS if name not in GOOGLE_ONLY_HEADERS]
    if process.get('machine_number') or '机器号码' in main_headers:
        verify_machine_number(main_headers, main_rows)
        expected_headers = list(expected_headers)
        expected_headers.insert(expected_headers.index('卡槽'), '机器号码')
    retention_review = review_retention_report(final_book, process.get("card_retention_report"), main_headers, retention_main_cells)
    if retention_review is not None:
        expected_headers = list(expected_headers)
        expected_headers.insert(expected_headers.index("激活后天数") + 1, "需要保留时间")
    expected_headers = list(expected_headers)
    expected_headers.insert(expected_headers.index('激活后天数') + 1, '使用天数')
    usage_review = verify_usage_days(raw_headers, raw_values, main_headers, main_rows,
                                    usage_main_cells, reference, process.get('creation_usage_days'))
    no_google_review = {"status": "enabled"} if google_enabled else review_no_google_output(
        final_book, raw_headers, raw_values, main_headers, main_rows, process, date_resolution_review, summary, resolved_date_rows)
    device_daily_review = review_device_daily_statistics(final_book, raw_headers, raw_values)
    difference_review = None
    # Read only these comparison sheets with formulas visible. Cached blanks must
    # not hide a formula substituted for the literal presence flag.
    comparison_book = load_workbook(final_path, read_only=True, data_only=False)
    try:
        if google_difference_enabled():
            difference_review = review_google_difference(comparison_book, raw_headers, raw_values,
                process.get("google_system_missing_sheet"), chatr_path, chatr_hash,
                os.environ.get("OPENSMS_RUN_ID", ""), timestamp_value)
            presence_review = difference_review["port_presence_review"]
        else:
            comparison_slots = {
                normalize_slot(row[google_slot_index])
                for row in google_all_sheet.iter_rows(min_row=2, values_only=True)
            } if google_enabled else set()
            presence_review = review_port_presence(comparison_book,
                {normalize_slot(row[raw_index["卡槽"]]) for row in raw_values}, comparison_slots,
                google_dates["iccid_records"] if google_enabled else (),
                [{"slot": row[raw_index["卡槽"]], "iccid": row[next((i for i, name in enumerate(raw_headers) if str(name).strip().lower() == "iccid"), 0)] if any(str(name).strip().lower() == "iccid" for name in raw_headers) else None} for row in raw_values])
        ko_review = review_ko_unactivated(comparison_book, process.get('ko_unactivated_sheet'),
            os.environ.get('OPENSMS_REPORT_DATE_CACHE_PATH', ''), os.environ.get('OPENSMS_REPORT_DATE_CACHE_SHA256', ''),
            os.environ.get('OPENSMS_RUN_ID', ''), os.environ.get('OPENSMS_KO_UNACTIVATED_REQUIRED') == '1',
            system=system_snapshot)
    finally:
        comparison_book.close()
    unused_port_review["google_presence_review"] = presence_review
    unused_port_review["receipt_google_presence_matches"] = process.get("system_missing_google_present_count") == presence_review["marked_rows"]
    unused_port_review["matches"] = (unused_port_review["matches"] and presence_review["matches"]
        and unused_port_review["receipt_google_presence_matches"])
    room_configuration_review = review_room_configuration(final_book, process, summary)
    final_book.close()
    raw_book.close()
    result = {
        "room_configuration_review": room_configuration_review,
        "device_daily_statistics_review": device_daily_review,
        "card_retention_report_review": retention_review,
        "google_system_missing_sheet_review": difference_review,
        "ko_unactivated_sheet_review": ko_review,
        "creation_usage_days_review": usage_review,
        "headers_match": main_headers[:len(expected_headers)] == expected_headers,
        "headers": main_headers[:len(expected_headers)],
        "slot_alias_review": {
            "matches": normalize_slot("204_62") == normalize_slot("204_62.01") == "204_62",
            "google_source_has_alias": process.get("google_source_has_204_62_alias"),
        },
        "google_field_mapping_review": {
            "expected": expected_mapping,
            "actual": actual_mapping,
            "matches": actual_mapping == expected_mapping,
        },
        "google_sheet_scope_review": {
            "expected": REQUIRED_ASSOC_SHEETS,
            "actual": actual_assoc_sheets,
            "row_counts": assoc_sheet_rows_text,
            "matches": actual_assoc_sheets == REQUIRED_ASSOC_SHEETS
            and all(re.search(rf"{re.escape(name)}:\d+", assoc_sheet_rows_text) for name in REQUIRED_ASSOC_SHEETS),
        },
        "carrier_opr_review": {
            "expected_rows": len(expected_carrier_warning_keys),
            "sheet_rows": len(actual_carrier_warning_keys),
            "detail_mismatch_count": len(carrier_warning_mismatches),
            "matches": not carrier_warning_mismatches
            and actual_carrier_warning_keys == expected_carrier_warning_keys,
        },
        "raw_rows": len(raw_rows),
        "final_rows": len(main_rows),
        "duplicate_phone_count": len(duplicate_phones),
        "mismatch_count": classification_mismatch_count,
        "mismatch_sample": classification_mismatches[:10],
        "window_metric_headers_match": all(
            name in main_index for name in (
                "3日发送窗口数", "3日回复率", "4日发送窗口数", "4日回复率", "回复率持续降低"
            )
        ) and main_index["3日发送窗口数"] == main_index["10日回复率"] + 1
        and main_index["3日回复率"] == main_index["10日回复率"] + 2
        and main_index["4日发送窗口数"] == main_index["10日回复率"] + 3
        and main_index["4日回复率"] == main_index["10日回复率"] + 4
        and main_index["回复率持续降低"] == main_index["10日回复率"] + 5,
        "window_metric_mismatch_count": len(rate_mismatches),
        "window_metric_mismatch_sample": rate_mismatches[:10],
        "window_rate_review": {
            "dates": [item["date"].isoformat() for item in latest_four],
            "counts_match": trend_counts_match,
            "summary_matches": trend_summary_matches,
            "expected_decreasing_counts": dict(expected_trends),
            "actual_decreasing_counts": dict(actual_trends),
        },
        "activation_age_rule": "后台服务器日期 - 对应权威激活日期；同日为0天；缺失或未来日期留空",
        "activation_age_mismatch_count": len(activation_age_mismatches),
        "activation_age_mismatch_sample": activation_age_mismatches[:10],
        "expiry_status_rule": "已过期=服务器日期>=截止日；快过期=截止日-服务器日期为1到2天（含2天）；6天=相差3到6天",
        "expiry_status_mismatch_count": len(expiry_status_mismatches),
        "expiry_status_mismatch_sample": expiry_status_mismatches[:10],
        "task_f_kooda_date_review": task_f_kooda_date_review,
        "system_kooda_source_review": system_kooda_source_review,
        "date_resolution_review": date_resolution_review,
        "source_schema_review": source_schema_review,
        "task_f_missing_activation_red_review": task_f_missing_red_review,
        "task_f_missing_dates_new_card_review": {
            "eligible_rows": task_f_missing_dates_new_card_eligible_rows,
            "mismatch_count": len(task_f_missing_dates_new_card_mismatches),
            "mismatches": task_f_missing_dates_new_card_mismatches,
            "matches": not task_f_missing_dates_new_card_mismatches,
        },
        "cancelled_review_sheets": {
            "expected_absent": cancelled_review_sheets,
            "unexpected_present": unexpected_cancelled_sheets,
            "matches": not unexpected_cancelled_sheets,
        },
        "activated_unlisted_review": {
            "expected_rows": len(expected_activated_keys),
            "actual_rows": len(actual_activated_keys),
            "details_match": actual_activated_keys == expected_activated_keys,
            "expected_machine_count": len(expected_machine_counts),
            "actual_machine_count": len(actual_machine_counts),
            "machine_summary_matches": dict(expected_machine_counts) == actual_machine_counts,
            "expected_over_five_machines": expected_over_five_machines,
            "actual_over_five_machines": actual_over_five_machines,
            "over_five_machines_match": expected_over_five_machines == actual_over_five_machines,
            "summary_matches": number(summary.get("已激活未上架总数")) == len(expected_activated_keys)
            and number(summary.get("超过5张未上架机器数")) == len(expected_over_five_machines),
        },
        "formula_error_matches": len(formula_error_cells),
        "formula_error_cells": formula_error_cells[:300],
        "numeric_text_review": numeric_text_review,
        "missing_sheets": missing_sheets,
        "google_system_missing_review": google_missing_review,
        "unused_port_review": unused_port_review,
        "preview_mode": "skipped_optimized_python_verify",
        "previews": [],
        "fast_verification": True,
        "independent_detail_verification": True,
    }
    if not google_enabled:
        for key in ("google_field_mapping_review", "google_sheet_scope_review", "carrier_opr_review", "google_system_missing_review"):
            result[key] = {"status": "skipped", "reason": "user_disabled"}
        result.update(no_google_review)
    if chatr_review is not None:
        result["google_chatr_date_fallback_review"] = chatr_review
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"错误: {error}", file=sys.stderr)
        raise SystemExit(1)
