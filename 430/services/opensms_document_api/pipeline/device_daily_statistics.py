"""Device/day aggregates from the already loaded export; no I/O in workers."""
from __future__ import annotations

from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from datetime import date
from decimal import Decimal, InvalidOperation
import math
import os
import re
import time
import unicodedata

from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

SHEET_NAME = "设备每日统计"
UNKNOWN_DEVICE = "未识别设备"
FIELDS = ("发送窗口数", "回复窗口数", "回复率", "总发送量", "总回复量")
COUNTS = ("发送窗口数", "回复窗口数", "总发送量", "总回复量")
MAX_WORKERS = 8


def device_key(slot):
    value = unicodedata.normalize("NFKC", str(slot or "")).strip().replace(" ", "").replace("-", "_")
    match = re.fullmatch(r"([A-Za-z0-9]+)_(\d+)(?:\.\d+)?", value)
    # This is a grouping key, NOT the 64-physical-port availability check.
    # Never silently drop rows with port aliases or ports outside 1..64.
    return (str(int(match[1])) if match[1].isdigit() else match[1]) if match else UNKNOWN_DEVICE


def device_sort_key(key):
    if key == UNKNOWN_DEVICE:
        return 2, ""
    return (0, int(key)) if key.isdigit() else (1, key)


def count_value(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value >= 0 else None
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value:
            return None
        try:
            numeric = Decimal(value)
            return int(numeric) if numeric.is_finite() and numeric >= 0 and numeric == numeric.to_integral_value() else None
        except InvalidOperation:
            return None
    try:
        parsed = float(value)
    except (ValueError, TypeError, OverflowError):
        return None
    if not math.isfinite(parsed) or parsed < 0 or not parsed.is_integer():
        return None
    return int(parsed)


def aggregate_device_daily(headers, rows, workers=None):
    """Keep source date/metric order, with independent date jobs and exact counts.

    Missing/invalid source counts make that device/date/metric unknown (blank),
    never an apparently complete partial sum. A zero send denominator is blank.
    Source phone records are counted as rows; no phone-only deduplication.
    """
    started = time.perf_counter()
    headers = tuple(str(value or "").strip() for value in headers)
    rows = tuple(rows)
    if headers.count("卡槽") != 1:
        raise ValueError("设备每日统计需要唯一的卡槽列")
    layout, columns = [], {}
    for index, header in enumerate(headers):
        match = re.fullmatch(r"(\d{2}-\d{2})(发送窗口数|回复窗口数|回复率|总发送量|总回复量)", header)
        if not match:
            continue
        day, field = match.groups()
        month, day_number = map(int, day.split("-"))
        date(2000, month, day_number)  # validate MM-DD, including leap day
        definition = columns.setdefault(day, {})
        if field in definition:
            raise ValueError(f"设备每日统计发现重复日期指标列: {header}")
        definition[field] = index
        layout.append((day, field, index))
    if not columns or any(set(definition) != set(FIELDS) for definition in columns.values()):
        raise ValueError("设备每日统计要求每个日期拥有主表完整的5个指标列")
    requested = int(workers if workers is not None else os.environ.get("OPENSMS_DEVICE_DAILY_WORKERS", "4"))
    if not 1 <= requested <= MAX_WORKERS:
        raise ValueError(f"设备每日统计线程数必须为1..{MAX_WORKERS}")
    actual = min(requested, len(columns)) if rows else 0
    slot_index = headers.index("卡槽")
    devices = tuple(device_key(row[slot_index]) for row in rows)
    sizes = Counter(devices)
    keys = sorted(sizes, key=device_sort_key)

    def one_day(item):
        day, definition = item
        totals = {key: [0] * 4 for key in keys}
        missing = {key: [0] * 4 for key in keys}
        indexes = tuple(definition[field] for field in COUNTS)
        for key, row in zip(devices, rows):
            for position, index in enumerate(indexes):
                value = count_value(row[index])
                if value is None:
                    missing[key][position] += 1
                else:
                    totals[key][position] += value
        result, missing_cells = {}, 0
        for key in keys:
            metrics = {field: None if missing[key][i] else totals[key][i] for i, field in enumerate(COUNTS)}
            sends, replies = metrics["发送窗口数"], metrics["回复窗口数"]
            metrics["回复率"] = replies / sends if sends and replies is not None else None
            result[key] = metrics
            missing_cells += sum(missing[key])
        return day, result, missing_cells

    if actual > 1:
        with ThreadPoolExecutor(max_workers=actual, thread_name_prefix="opensms-device-day") as pool:
            results = list(pool.map(one_day, columns.items()))
    else:
        results = [one_day(item) for item in columns.items()]
    by_date = {day: values for day, values, _ in results}
    output_rows = [
        [key, sizes[key], *(by_date[day][key][field] for day, field, _ in layout)]
        for key in keys
    ]
    receipt = {
        "sheet_name": SHEET_NAME, "source_rows": len(rows),
        "device_count": len(keys) - int(UNKNOWN_DEVICE in sizes),
        "output_rows": len(keys), "unrecognized_device_rows": sizes.get(UNKNOWN_DEVICE, 0),
        "dates": list(columns), "date_count": len(columns),
        "source_metric_headers": [headers[index] for _, _, index in layout],
        "requested_workers": requested, "actual_workers": actual,
        "missing_or_invalid_count_cells": sum(missing for _, _, missing in results),
        "blank_aggregate_cells": sum(value is None for row in output_rows for value in row[2:]),
        "rate_rule": "同设备同日回复窗口数合计/发送窗口数合计；分母0或输入缺失时留空",
        "missing_rule": "任一源计数缺失或非法时，该设备当天对应合计留空，不发布不完整合计",
        "row_rule": "保留全部主表记录；数字设备号归一化，字母前缀保留（r110与110分别统计）；无法识别的归入未识别设备；不按手机号去重",
        "execution": "immutable_source_rows_parallel_dates_single_workbook_writer",
        "aggregation_seconds": round(time.perf_counter() - started, 4),
    }
    return {"headers": ["设备号", "记录数", *receipt["source_metric_headers"]],
            "rows": output_rows, "receipt": receipt}


def write_device_daily_sheet(book, main_sheet, result):
    """Only the coordinator touches openpyxl; preserve the main sheet unchanged."""
    started = time.perf_counter()
    if SHEET_NAME in book.sheetnames:
        del book[SHEET_NAME]
    sheet = book.create_sheet(SHEET_NAME, book.index(main_sheet) + 1)
    source_columns = {str(cell.value): cell.column for cell in main_sheet[1]}
    sheet.append(result["headers"])
    for row in result["rows"]:
        sheet.append(row)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = main_sheet.sheet_view.zoomScale
    sheet.sheet_view.showGridLines = main_sheet.sheet_view.showGridLines
    sheet.row_dimensions[1].height = main_sheet.row_dimensions[1].height
    sheet.sheet_format = copy(main_sheet.sheet_format)
    body_source_row = 2 if main_sheet.max_row > 1 else 1
    # Copy styles once per column, not every source record. Row-level expiry/
    # change-status/red-missing-date colors are not meaningful for a device sum.
    for column, header in enumerate(result["headers"], start=1):
        source_column = source_columns[header] if column > 2 else source_columns["卡槽"]
        original_header = main_sheet.cell(1, source_column)
        sheet.cell(1, column)._style = copy(original_header._style)
        sheet.cell(1, column).alignment = copy(original_header.alignment)
        source_letter = get_column_letter(source_column)
        source_dimension = main_sheet.column_dimensions.get(source_letter)
        width = source_dimension.width if source_dimension else (main_sheet.sheet_format.defaultColWidth or 13)
        sheet.column_dimensions[get_column_letter(column)].width = width if column > 2 else (14 if column == 1 else 12)
        body_style = copy(main_sheet.cell(body_source_row, source_column)._style)
        for row_number in range(2, len(result["rows"]) + 2):
            cell = sheet.cell(row_number, column)
            cell._style = copy(body_style)
            cell.fill = PatternFill()
            cell.number_format = "@" if column == 1 else "0.00%" if header.endswith("回复率") else "#,##0"
    receipt = dict(result["receipt"], write_seconds=round(time.perf_counter() - started, 4))
    return receipt
