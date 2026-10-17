"""Hash-bound Kooda projection of ONE fresh 10-day export, never live inventory.

No backend IDs are synthesized. sourceRowIdentity is only a workbook identity.
Generators consume the scan cache; document reviewers recompute from the XLSX
they independently opened. Old paginated inventory receipts remain readable.
"""
from datetime import date, datetime
import hashlib
from pathlib import Path
import re

KIND = 'verified_10day_export'
SCOPE = 'current_10day_export_only'
ELIGIBILITY = 'raw_10day_activation_missing_and_local_activation_missing'


def is_report_source(payload):
    return payload.get('sourceKind') == KIND


def iccid_key(value):
    raw = re.sub(r'[Ff]$', '', value.strip()) if isinstance(value, str) else ''
    return raw if re.fullmatch(r'[0-9]{18,22}', raw) else ''


def source_date(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        raw = value.strip()
        match = re.fullmatch(r'(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:\d{2})?)?', raw)
        if match:
            return date(*map(int, match.groups())).isoformat()
    raise RuntimeError('10日表非空激活日期/截止日无效，禁止当空白补查')


def json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def project_rows(headers, rows, run_id):
    names = [str(value or '').strip() for value in headers]
    if len([name for name in names if name]) != len(set(name for name in names if name)):
        raise RuntimeError('10日表源表头重复，禁止歧义映射')
    for name in ('电话', '运营商', 'opr', '卡槽', 'ICCID'):
        if names.count(name) != 1:
            raise RuntimeError(f'10日表缺少唯一{name}列，不回退全库存查询')
    index = {name: pos for pos, name in enumerate(names) if name}
    projected = []
    for number, values in enumerate(rows, 2):
        raw = {name: json_value(values[pos]) if pos < len(values) else None for name, pos in index.items()}
        if not re.search(r'kooda|koodo|telus1', str(raw.get('opr') or ''), re.I):
            continue
        iccid = raw['ICCID']
        if iccid is not None and not isinstance(iccid, str):
            raise RuntimeError('10日表ICCID不是文本，禁止丢精度后查询')
        projected.append({'sourceRow': number, 'sourceRowIdentity': f'{run_id}:export-row:{number}',
            'cardNo': raw['电话'], 'slotIdentif': raw['卡槽'], 'cardStore': iccid,
            'opr': raw['opr'], 'carrierName': raw['运营商'],
            'chatrActivateDate': source_date(raw.get('激活日期')),
            'chatrPlanExpire': source_date(raw.get('截止日')), 'sourceData': raw})
    return projected


def projection_payload(headers, rows, run_id, source_path, source_hash):
    projected = project_rows(headers, rows, run_id)
    return {'schemaVersion': 1, 'status': 'passed', 'runId': run_id,
        'sourceKind': KIND, 'sourceScope': SCOPE, 'dateEligibilityMode': ELIGIBILITY,
        'sourceWorkbookPath': str(Path(source_path).resolve()), 'sourceWorkbookSha256': source_hash,
        'sourceRowCount': len(rows), 'sourceHeaders': [str(v or '').strip() for v in headers],
        'systemRows': projected, 'totalRows': len(projected),
        'networkAccessed': False, 'backendInventoryQueried': False,
        'silentBrowserMode': True, 'browserForegroundActivationAttempted': False}


def validate_projection(payload, run_id):
    rows = payload.get('systemRows')
    if (not is_report_source(payload) or payload.get('schemaVersion') != 1
            or payload.get('status') != 'passed' or payload.get('runId') != run_id
            or payload.get('sourceScope') != SCOPE or payload.get('dateEligibilityMode') != ELIGIBILITY
            or payload.get('networkAccessed') is not False or payload.get('backendInventoryQueried') is not False
            or not Path(payload.get('sourceWorkbookPath') or '').is_absolute()
            or not re.fullmatch(r'[a-f0-9]{64}', payload.get('sourceWorkbookSha256') or '')
            or not isinstance(rows, list) or type(payload.get('totalRows')) is not int
            or payload['totalRows'] != len(rows) or type(payload.get('sourceRowCount')) is not int
            or payload['sourceRowCount'] < 0 or not isinstance(payload.get('sourceHeaders'), list)
            or payload.get('silentBrowserMode') is not True
            or payload.get('browserForegroundActivationAttempted') is not False):
        raise RuntimeError('10日表Kooda来源范围、哈希或行数证明无效')
    seen = set()
    for row in rows:
        number = row.get('sourceRow') if isinstance(row, dict) else None
        if (type(number) is not int or not 2 <= number <= payload['sourceRowCount'] + 1 or number in seen
                or row.get('sourceRowIdentity') != f'{run_id}:export-row:{number}' or row.get('id') is not None
                or not isinstance(row.get('sourceData'), dict)):
            raise RuntimeError('10日表源行身份缺失、重复或冒充后台ID')
        seen.add(number)
        # Reconstruct all selected fields from the preserved source object.
        raw = row['sourceData']
        expected = project_rows(payload['sourceHeaders'],
            [tuple(raw.get(h) for h in payload['sourceHeaders'])], run_id)
        if len(expected) != 1:
            raise RuntimeError('10日表Kooda来源OPR越界')
        expected[0].update(sourceRow=number, sourceRowIdentity=row['sourceRowIdentity'])
        if row != expected[0]:
            raise RuntimeError('10日表源字段投影不一致')
    return payload


def verify_projection_rows(payload, headers, rows, run_id, source_path):
    validate_projection(payload, run_id)
    with Path(source_path).open('rb') as stream:
        actual_hash = hashlib.file_digest(stream, 'sha256').hexdigest()
    if (actual_hash != payload['sourceWorkbookSha256']
            or payload['sourceRowCount'] != len(rows)
            or payload['sourceHeaders'] != [str(v or '').strip() for v in headers]
            or payload['systemRows'] != project_rows(headers, rows, run_id)):
        raise RuntimeError('10日表Kooda投影与独立重读原始XLSX不一致')


def verify_projection_workbook(payload, path, run_id):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = book.active.iter_rows(values_only=True)
        headers = next(iterator, ())
        verify_projection_rows(payload, headers, list(iterator), run_id, path)
    finally:
        book.close()


def missing_activation_rows(payload):
    """ICCID-wide exclusion: one real source activation excludes all its slots."""
    validate_projection(payload, payload['runId'])
    known = {iccid_key(row['cardStore']) for row in payload['systemRows'] if row['chatrActivateDate']}
    return [row for row in payload['systemRows']
            if not row['chatrActivateDate'] and (not iccid_key(row['cardStore']) or iccid_key(row['cardStore']) not in known)]
