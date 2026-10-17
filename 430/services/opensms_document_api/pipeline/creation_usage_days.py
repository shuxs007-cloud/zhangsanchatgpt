"""Display-only calendar age from the verified server date and source creation time."""
from copy import copy
from datetime import date, datetime
import re

from openpyxl.utils import column_index_from_string, get_column_letter

COLUMN = '使用天数'
RULE = '本轮后台服务器日期 - 源表创建时间的自然日期；同日0，不加1；缺失/非法留空，未来日期保留负数'


def creation_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str) or not value.strip():
        return None
    text = value.strip().replace('/', '-')
    match = re.fullmatch(r'(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2})(?:\.(\d{1,6}))?)?)?', text)
    if not match:
        return None
    year, month, day, hour, minute, second, fraction = match.groups()
    try:
        return datetime(int(year), int(month), int(day), int(hour or 0), int(minute or 0),
                        int(second or 0), int((fraction or '').ljust(6, '0'))).date()
    except ValueError:
        return None


def add_usage_days(sheet, reference):
    if type(reference) is not date:
        raise ValueError('使用天数需要已核验的后台服务器自然日期')
    headers = [cell.value for cell in sheet[1]]
    if COLUMN in headers or headers.count('激活后天数') != 1 or headers.count('创建时间') != 1:
        raise ValueError('使用天数/激活后天数/创建时间列缺失或重复')
    column = headers.index('激活后天数') + 2
    created_column = headers.index('创建时间') + 1
    dimensions = [(column_index_from_string(k), copy(v)) for k, v in sheet.column_dimensions.items()]
    sheet.insert_cols(column)
    sheet.column_dimensions.clear()
    for old, dimension in dimensions:
        new = old + (old >= column)
        dimension.index = get_column_letter(new)
        dimension.min = new
        dimension.max = (dimension.max or old) + ((dimension.max or old) >= column)
        sheet.column_dimensions[dimension.index] = dimension
    if created_column >= column:
        created_column += 1
    sheet.column_dimensions[get_column_letter(column)].width = 14
    known = unknown = future = 0
    for row in range(1, sheet.max_row + 1):
        cell = sheet.cell(row, column)
        cell._style = copy(sheet.cell(row, column - 1)._style)
        cell.number_format = '0'
        if row == 1:
            cell.value = COLUMN
            continue
        created = creation_date(sheet.cell(row, created_column).value)
        cell.value = (reference - created).days if created else None
        known += created is not None
        unknown += created is None
        future += created is not None and created > reference
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = sheet.dimensions
    return {'column': COLUMN, 'rule': RULE, 'server_date': reference.isoformat(),
            'rows': sheet.max_row - 1, 'known_rows': known, 'unknown_rows': unknown, 'future_rows': future}


def verify_usage_days(raw_headers, raw_rows, headers, rows, cells, reference, proof):
    """Recompute from reopened original source; never trust generated values/counts."""
    if (type(reference) is not date or headers.count(COLUMN) != 1
            or headers.count('激活后天数') != 1 or raw_headers.count('创建时间') != 1
            or headers.index(COLUMN) != headers.index('激活后天数') + 1):
        raise ValueError('使用天数缺少唯一来源或未紧跟激活后天数')
    if not (len(raw_rows) == len(rows) == len(cells)):
        raise ValueError('使用天数源表与结果行数不一致')
    created_index, output_index = raw_headers.index('创建时间'), headers.index(COLUMN)
    known = unknown = future = 0
    errors = []
    for row_number, (raw, row, (cell, neighbour)) in enumerate(zip(raw_rows, rows, cells), 2):
        # Separate date parser and arithmetic from the generator.
        value = raw[created_index]
        created = None
        if isinstance(value, datetime):
            created = date(value.year, value.month, value.day)
        elif isinstance(value, date):
            created = value
        elif isinstance(value, str):
            normalized = value.strip().replace('/', '-').replace('T', ' ')
            formats = ('%Y-%m-%d', '%Y-%m-%d %H:%M', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M:%S.%f')
            if not re.fullmatch(r'\d{4}-\d{1,2}-\d{1,2}(?: \d{1,2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?)?', normalized):
                formats = ()
            for fmt in formats:
                try:
                    parsed = datetime.strptime(normalized, fmt)
                    created = parsed.date()
                    break
                except ValueError:
                    pass
        wanted = reference.toordinal() - created.toordinal() if created else None
        known += created is not None
        unknown += created is None
        future += created is not None and created > reference
        actual_fill = copy(cell.fill) if cell.fill is not None and cell.fill.fill_type else None
        expected_fill = copy(neighbour.fill) if neighbour.fill is not None and neighbour.fill.fill_type else None
        if (row[output_index] != wanted or cell.value != wanted
                or (wanted is not None and (type(cell.value) not in (int, float) or cell.data_type != 'n'))
                or cell.number_format != '0' or actual_fill != expected_fill):
            errors.append(row_number)
    expected = {'column': COLUMN, 'rule': RULE, 'server_date': reference.isoformat(),
                'rows': len(raw_rows), 'known_rows': known, 'unknown_rows': unknown, 'future_rows': future}
    return {**expected, 'matches': not errors and proof == expected,
            'mismatch_count': len(errors), 'mismatch_rows': errors[:20]}
