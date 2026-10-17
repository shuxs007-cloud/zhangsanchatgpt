import contextlib
from datetime import datetime
import gzip
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
import opensms_restore_iccid_excel_history as scanner

KEY = '8912230000000000001'
OTHER = '8912230000000000002'


class RestoreTests(unittest.TestCase):
    def scan(self, path, keys=None):
        stat = path.stat()
        item = scanner.base.fingerprint(dict(path=str(path),size=stat.st_size,mtimeNs=stat.st_mtime_ns))
        return scanner.scan_book(item, set(keys or [KEY]))

    def test_all_fields_dates_epoch_and_trailing_f(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'input.xlsx'
            w=Workbook();w.epoch=CALENDAR_MAC_1904;s=w.active
            s.append(['ICCID','卡槽','上卡日期','创建时间','备注','激活日期','其它原字段'])
            s.append([KEY+'F','r216_07',datetime(2026,4,1),datetime(2026,4,2,3,4,5),'原始备注',None,'原字段内容'])
            s['C2'].number_format='yyyy-mm-dd';s['D2'].number_format='yyyy-mm-dd hh:mm:ss';w.save(path)
            r=self.scan(path)['rows'][0]['record'];v=scanner.extracted(r)
            self.assertEqual(v['上卡日期'],['2026-04-01'])
            self.assertEqual(v['创建时间'],['2026-04-02 03:04:05'])
            self.assertEqual(v['激活日期'],[])
            self.assertEqual(v['卡槽'],['r216_07'])
            self.assertTrue(any(x['header']=='其它原字段' and x['value']=='原字段内容' for x in r['fields']))

    def test_date_not_inferred_and_formula_inert(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'input.xlsx';w=Workbook();s=w.active
            s.append(['ICCID','创建时间','激活日期','备注'])
            s.append([KEY,'2026-04-02 03:04:05','2026-04-01','=1+1']);w.save(path)
            r=self.scan(path)['rows'][0]['record'];v=scanner.extracted(r)
            self.assertEqual(v['上卡日期'],[])
            formula=next(x for x in r['fields'] if x['header']=='备注')
            self.assertTrue(formula['formulaPresent']);self.assertEqual(formula['formula'],'1+1')
            self.assertEqual(v['备注'],[])

    def test_multiple_ids_untrusted_and_numeric_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'input.xlsx';w=Workbook();s=w.active
            s.append(['系统ICCID','GoogleICCID','卡槽','创建时间'])
            s.append([KEY,OTHER,'201_1','2026-04-01'])
            s.append([int(KEY),None,'201_2','2026-04-02']);w.save(path)
            result=self.scan(path)
            self.assertFalse(result['rows'][0]['record']['trusted'])
            self.assertEqual(scanner.extracted(result['rows'][0]['record'])['卡槽'],[])

    def test_headerless_preserves_whole_row_but_no_assignment(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'input.xlsx';w=Workbook();s=w.active
            s.append([KEY,'201_1','2026-04-01']);w.save(path)
            r=self.scan(path)['rows'][0]['record']
            self.assertFalse(r['trusted']);self.assertEqual(len(r['fields']),3)

    def test_rich_inline_split_key(self):
        xml=('<worksheet xmlns="'+scanner.base.NS+'"><sheetData><row r="1">'
             '<c r="A1" t="inlineStr"><is><t>ICCID</t></is></c><c r="B1" t="inlineStr"><is><t>卡槽</t></is></c></row>'
             '<row r="2"><c r="A2" t="inlineStr"><is><r><t>'+KEY[:8]+'</t></r><r><t>'+KEY[8:]+'</t></r></is></c>'
             '<c r="B2" t="inlineStr"><is><t>201_1</t></is></c></row></sheetData></worksheet>').encode()
        r=scanner.scan_sheet(xml,[],'none',{KEY},['General'],scanner.CALENDAR_WINDOWS_1900)
        self.assertEqual(len(r['rows']),1)

    def test_full_cli_scopes_duplicates_and_independent_output(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);listed=root/'list.txt';listed.write_text('ICCID\n'+KEY+'\n'+OTHER+'\n')
            path=root/'input_20260401.xlsx';w=Workbook();s=w.active
            s.append(['ICCID','卡槽','上卡日期','创建时间','额外字段'])
            s.append([KEY,'201_01','2026-04-01','2026-04-01 11:00:00','内容']);w.save(path)
            (root/'copy_20260501.xlsx').write_bytes(path.read_bytes())
            (root/'old_20260301.xlsx').write_bytes(path.read_bytes())
            (root/'log_20260401.json').write_text(json.dumps({'iccid':OTHER,'slot':'999_1'}))
            (root/'task_f.sqlite3').write_bytes(b'invalid-db')
            out=root/'output'
            args=['script','--iccid-list',str(listed),'--root',str(root),'--start','2026-03-24','--end','2026-09-24','--workers','1','--output',str(out)]
            with patch.object(sys,'argv',args),contextlib.redirect_stdout(io.StringIO()):scanner.main()
            report=json.loads((out/'summary.json').read_text())
            self.assertEqual(report['targetIccids'],2);self.assertEqual(report['matchedIccids'],1)
            self.assertEqual(report['withBothDates'],1);self.assertEqual(report['sourceRows'],2)
            self.assertEqual(report['selectedFiles'],2);self.assertEqual(report['uniqueContents'],1)
            self.assertFalse(report['databaseRead']);self.assertFalse(report['jsonLogsRead'])
            self.assertTrue(report['verification']['independentlyVerified'])
            w=load_workbook(report['workbook'],read_only=True)
            self.assertEqual(w['查询总览']['A2'].value,KEY)
            self.assertEqual(w['查询总览']['A2'].data_type,'s');w.close()
            with gzip.open(out/'all_records.jsonl.gz','rt') as f:
                rows=[json.loads(x) for x in f]
            self.assertEqual(len(rows),1)


if __name__=='__main__':unittest.main()
