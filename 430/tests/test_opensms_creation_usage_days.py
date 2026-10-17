"""Offline calendar-age and full generation/readback regressions."""
from datetime import date, datetime
import json
import sys
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from services.opensms_document_api.pipeline.creation_usage_days import add_usage_days, verify_usage_days


class CreationUsageDaysTests(unittest.TestCase):
    def test_calendar_boundaries_unknown_values_styles_and_types(self):
        book = Workbook()
        sheet = book.active
        headers = ['电话', 'koo是否2天', '激活后天数', '需要保留时间', '创建时间', 'ICCID']
        sheet.append(headers)
        reference = date(2026, 9, 20)
        cases = [('2026-09-20 00:00:00', 0), ('2026-09-19 23:59:59', 1),
                 ('2026/9/1 08:20:01', 19), (datetime(2026, 8, 31, 23), 20),
                 (date(2025, 12, 31), 263), ('2026-09-21', -1),
                 (None, None), ('', None), ('2026-02-30', None),
                 ('2026-09-20 25:01:00', None), (True, None), (46000, None)]
        for i, (created, _) in enumerate(cases):
            sheet.append([str(i), '2天', 99, '2026-10-01', created, '8912230200025159114F'])
        sheet.column_dimensions['F'].width = 26
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='FFC7CE' if cell.row % 2 else 'FFFF00')
        source_rows = list(sheet.values)[1:]
        proof = add_usage_days(sheet, reference)
        self.assertEqual(proof['known_rows'], 6)
        self.assertEqual(proof['unknown_rows'], 6)
        self.assertEqual(proof['future_rows'], 1)
        self.assertEqual(sheet.column_dimensions['G'].width, 26)
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        actual_headers = list(next(sheet.values))
        rows = list(sheet.values)[1:]
        cells = [(sheet.cell(i, 4), sheet.cell(i, 3)) for i in range(2, sheet.max_row + 1)]
        self.assertEqual([r[3] for r in rows], [expected for _, expected in cases])
        self.assertEqual(actual_headers[2:5], ['激活后天数', '使用天数', '需要保留时间'])
        review = verify_usage_days(headers, source_rows, actual_headers, rows, cells, reference, proof)
        self.assertTrue(review['matches'], review)
        sheet.cell(2, 4).value = '0'
        review = verify_usage_days(headers, source_rows, actual_headers, list(sheet.values)[1:], cells, reference, proof)
        self.assertFalse(review['matches'])
        with self.assertRaises(ValueError):
            add_usage_days(sheet, reference)
        book.close()

    def test_leap_day_and_missing_server_date(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['激活后天数', '创建时间'])
        sheet.append([100, '2024-02-28 23:59:59'])
        with self.assertRaises(ValueError):
            add_usage_days(sheet, None)
        add_usage_days(sheet, date(2024, 3, 1))
        self.assertEqual(sheet.cell(2, 2).value, 2)

    def test_real_pipeline_recomputes_from_source_and_rejects_tampering(self):
        from tests.test_opensms_document_no_google import NoGoogleDocumentTests, PIPELINE
        fixture = NoGoogleDocumentTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        source_book = load_workbook(fixture.raw)
        source_sheet = source_book.active
        # Match the real export's newest-day-first statistics order.
        for row_number, values in enumerate(list(source_sheet.values), 1):
            ordered = list(values[:13]) + [value for start in (28, 23, 18, 13) for value in values[start:start + 5]]
            for column, value in enumerate(ordered, 1):
                source_sheet.cell(row_number, column).value = value
        source_book.save(fixture.raw)
        source_book.close()
        before = [fixture.digest(p) for p in (fixture.raw, fixture.system, fixture.cache)]
        env = {**fixture.env, 'OPENSMS_MAIN_PATH': str(fixture.raw),
               'OPENSMS_OUTPUT_DIR': str(fixture.output.parent), 'EXPORT4DAY_SERVER_DATE': '2026-09-06'}
        process = fixture.command([sys.executable, '-B', str(PIPELINE / 'process_card_slots.py')], env)
        command = [sys.executable, '-B', str(PIPELINE / 'opensms_verify_card_slot_fast.py'), str(fixture.output), str(fixture.raw)]
        env.update({'PROCESS_JSON': json.dumps(process), 'ENRICH_JSON': json.dumps(process['enrichment'])})
        verified = fixture.command(command, env)
        self.assertTrue(verified['headers_match'])
        self.assertTrue(verified['creation_usage_days_review']['matches'], (process['creation_usage_days'], verified['creation_usage_days_review']))
        self.assertEqual(verified['creation_usage_days_review']['unknown_rows'], 1)
        self.assertEqual(verified['window_metric_mismatch_count'], 0)
        self.assertEqual(verified['activation_age_mismatch_count'], 0)
        from tests.test_opensms_parallel_audit import AUDIT
        structure = AUDIT.check_xlsx('result', fixture.output, inspect_result=True,
                                     server_date='2026-09-06', google_enabled=False)
        self.assertEqual(structure.status, AUDIT.PASS, structure.message)
        book = load_workbook(fixture.output)
        sheet = book['sheet1']
        self.assertEqual(sheet['G1'].value, '使用天数')
        self.assertEqual(sheet['G3'].value, 0)
        self.assertEqual(sheet['G4'].value, 5)
        self.assertEqual(sheet['E4'].value, '2天')
        self.assertEqual(sheet.cell(1, sheet.max_column).value, 'ICCID')
        sheet['G4'] = 6
        book.save(fixture.output)
        book.close()
        verified = fixture.command(command, env)
        self.assertFalse(verified['creation_usage_days_review']['matches'])
        self.assertEqual(verified['creation_usage_days_review']['mismatch_rows'], [4])
        book = load_workbook(fixture.output)
        book['sheet1'].delete_cols(7)
        book.save(fixture.output)
        book.close()
        structure = AUDIT.check_xlsx('result', fixture.output, inspect_result=True,
                                     server_date='2026-09-06', google_enabled=False)
        self.assertEqual(structure.status, AUDIT.FAIL)
        self.assertEqual(before, [fixture.digest(p) for p in (fixture.raw, fixture.system, fixture.cache)])


if __name__ == '__main__':
    unittest.main()
