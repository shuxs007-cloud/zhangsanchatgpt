import unittest

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from services.opensms_document_api.pipeline.google_system_missing import (
    attach_port_iccids, PORT_HEADERS, PORT_SHEET, review_port_presence,
)

from services.opensms_document_api.pipeline.process_card_slots import (
    GOOGLE_ACTIVATED_MISSING_HEADERS,
    build_unused_port_rows,
    classify_google_system_missing_rows,
    google_activated_missing_output_row,
    highlight_activated_google_missing_rows,
)


def google_row(slot, open_date=None, activation_date=None, expiry_date=None):
    return {
        "raw_slot": slot,
        "open_date": open_date,
        "activation_date": activation_date,
        "expiry_date": expiry_date,
        "note": "测试备注",
        "iccid": "89000000000000000001",
        "phone": "14385550123",
        "carrier_price": "Koodo",
        "source_sheet": "四號機房",
        "source_row": 2,
    }


class GoogleSystemMissingSheetTests(unittest.TestCase):
    def test_iccids_preserve_all_cards_without_expanding_scope(self):
        rows = [["130", "五號機房", 1, "130_1", "", "", "注册失败", ""],
                ["130", "五號機房", 2, "130_2", "", "", "系统中无", "谷歌表存在"]]
        google = [{"slot": "130_01.01", "iccid": " 89101 "},
                  {"slot": "130_1", "iccid": "89102"}, {"slot": "130_2", "iccid": "89103"},
                  {"slot": "888_1", "iccid": "89104"}]
        system = [{"slot": "130_1", "iccid": "89102"}, {"slot": "130_1", "iccid": "89101"}]
        enriched = attach_port_iccids(rows, google, system)
        self.assertEqual(len(enriched), 2)
        self.assertEqual(enriched[0][8:], ["89101\n89102", "89101\n89102"])
        self.assertEqual(enriched[1][8:], ["89103", ""])
        book = Workbook()
        ws = book.active
        ws.title = PORT_SHEET
        ws.append(PORT_HEADERS)
        for row in enriched:
            ws.append(row)
            for cell in ws[ws.max_row][8:]:
                cell.data_type = "s"
                cell.number_format = "@"
            if row[8] != row[9]:
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        result = review_port_presence(book, {"130_1"}, {"130_1", "130_2"}, google, system)
        self.assertTrue(result["matches"], result)
        self.assertEqual(result["iccid_mismatch_rows"], 1)
        ws.cell(2, 10).value = "wrong"
        self.assertFalse(review_port_presence(book, {"130_1"}, {"130_1", "130_2"}, google, system)["matches"])

    def test_both_blank_iccids_are_not_a_conflict(self):
        self.assertEqual(attach_port_iccids([["1", "", 1, "1_1", "", "", "系统中无", ""]], [], [])[0][8:], ["", ""])

    def test_unused_ports_expand_each_device_to_64_and_keep_exact_status(self):
        system_rows = [
            {"slot": "130_1", "phone": "1001", "carrier": "A1", "status": "上架"},
            {"slot": "130_2", "phone": "1002", "carrier": "B1", "status": "注册失败"},
            {"slot": "130_3", "phone": "1003", "carrier": "B1", "status": "注册失败"},
            {"slot": "130_3", "phone": "1004", "carrier": "B12", "status": "上架"},
        ]
        rows, devices = build_unused_port_rows(system_rows, ["131_1"])
        by_slot = {row[3]: row for row in rows}
        self.assertEqual(["130", "131"], devices)
        self.assertNotIn("130_1", by_slot)
        self.assertEqual("注册失败", by_slot["130_2"][6])
        self.assertEqual("1002", by_slot["130_2"][4])
        self.assertNotIn("130_3", by_slot)
        self.assertEqual("系统中无", by_slot["131_1"][6])
        self.assertEqual(126, len(rows))

    def test_google_annotation_does_not_expand_device_scope_or_mark_failed_ports(self):
        system = [
            {"slot": "130_1", "status": "注册失败"},
            {"slot": "130_2", "status": "下架"},
            {"slot": "130_3", "status": ""},
        ]
        rows, devices = build_unused_port_rows(system, [], ["130_01.01", "130_2", "130_3", "１３０-０４.01", "888_1", "r130_5"])
        by_slot = {row[3]: row for row in rows}
        self.assertEqual(devices, ["130"])
        self.assertEqual(by_slot["130_1"][7], "")
        self.assertNotIn("130_2", by_slot)
        self.assertNotIn("130_3", by_slot)
        self.assertEqual(by_slot["130_4"][7], "谷歌表存在")
        self.assertEqual(by_slot["130_5"][7], "")
        self.assertEqual(sum(row[7] == "谷歌表存在" for row in rows), 1)

    def test_all_google_missing_does_not_require_dates(self):
        rows = [google_row("130_01")]
        all_missing, activated_missing = classify_google_system_missing_rows(rows, set())
        self.assertEqual([item[1] for item in all_missing], ["130_1"])
        self.assertEqual(activated_missing, [])

    def test_activation_subset_requires_only_valid_activation_date(self):
        rows = [google_row("130_02", activation_date="2026/09/02")]
        all_missing, activated_missing = classify_google_system_missing_rows(rows, set())
        self.assertEqual(len(all_missing), 1)
        self.assertEqual(len(activated_missing), 1)
        self.assertIsNone(activated_missing[0][2])
        self.assertEqual(activated_missing[0][3].date().isoformat(), "2026-09-02")

    def test_activation_subset_output_includes_requested_google_details(self):
        rows = [google_row(
            "130_02",
            open_date="2026/09/01",
            activation_date="2026/09/02",
            expiry_date="2026/09/30",
        )]
        _all_missing, activated_missing = classify_google_system_missing_rows(rows, set())
        output = google_activated_missing_output_row(activated_missing[0])
        values = dict(zip(GOOGLE_ACTIVATED_MISSING_HEADERS, output))
        self.assertEqual("89000000000000000001", values["ICCID"])
        self.assertEqual("14385550123", values["手机号"])
        self.assertEqual("2026-09-01", values["上卡日期"].date().isoformat())
        self.assertEqual("Koodo", values["电信商"])

    def test_open_date_without_activation_is_not_in_activation_subset(self):
        rows = [google_row("130_03", open_date="2026/09/01")]
        all_missing, activated_missing = classify_google_system_missing_rows(rows, set())
        self.assertEqual(len(all_missing), 1)
        self.assertEqual(activated_missing, [])

    def test_system_present_and_empty_slots_are_excluded_after_normalization(self):
        rows = [
            google_row("204_62.01", activation_date="2026/09/01"),
            google_row("", activation_date="2026/09/01"),
        ]
        all_missing, activated_missing = classify_google_system_missing_rows(rows, {"204_62"})
        self.assertEqual(all_missing, [])
        self.assertEqual(activated_missing, [])

    def test_activation_subset_rows_are_highlighted_yellow_across_full_row(self):
        rows = [
            google_row("130_04", activation_date="2026/09/02"),
            google_row("130_05"),
        ]
        all_missing, _activated_missing = classify_google_system_missing_rows(rows, set())
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([
            "Google页签", "Google行号", "卡槽", "归一化卡槽",
            "备注", "ICCID", "上卡日期", "电信商",
        ])
        for record, norm_slot, *_dates in all_missing:
            worksheet.append([
                record["source_sheet"], record["source_row"], record["raw_slot"], norm_slot,
                record["note"], record["iccid"], record["open_date"], record["carrier_price"],
            ])

        highlighted = highlight_activated_google_missing_rows(worksheet, all_missing)

        self.assertEqual(highlighted, 1)
        self.assertTrue(all(cell.fill.fill_type == "solid" for cell in worksheet[2]))
        self.assertTrue(all(str(cell.fill.fgColor.rgb).upper().endswith("FFFF00") for cell in worksheet[2]))
        self.assertTrue(all(cell.fill.fill_type is None for cell in worksheet[3]))


if __name__ == "__main__":
    unittest.main()
