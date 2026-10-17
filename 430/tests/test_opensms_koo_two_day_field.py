"""Synthetic, offline rename regression; never connects to a browser."""
import json
import subprocess
import sys
import unittest

from openpyxl import load_workbook


class KooTwoDayFieldTests(unittest.TestCase):
    def test_threshold_family_and_nonconsecutive_dates_unchanged(self):
        from datetime import date
        from services.opensms_document_api.pipeline.process_card_slots import analyze_raw_metrics
        headers = ["电话", "运营商", "opr", "卡槽"]
        for day in ("09-03", "09-04", "09-05", "09-06"):
            headers += [day + name for name in ("发送窗口数", "回复窗口数", "总发送量", "总回复量")]
        for family, windows, expected in (
            ("Kooda", [80, 81, 0, 0], ""),
            ("Koodo", [81, 0, 81, 0], "2天"),
            ("Telus1", [81, 81, 81, 81], "2天"),
            ("Kini", [81, 81, 0, 0], "2天"),
            ("Chatr", [81, 81, 81, 81], ""),
        ):
            with self.subTest(family=family, windows=windows):
                row = ["14160000001", "B1", family, "101_1"]
                for sends in windows:
                    row += [sends, 0, 1000, 0]
                result = analyze_raw_metrics([headers, row], date(2026, 9, 6))
                self.assertEqual(result["use_days_by_phone"].get("14160000001", ""), expected)

    def test_real_generation_independent_readback_and_old_header_rejection(self):
        from tests.test_opensms_document_no_google import NoGoogleDocumentTests, PIPELINE
        fixture = NoGoogleDocumentTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        before = [fixture.digest(p) for p in (fixture.raw, fixture.system, fixture.cache)]
        env = {**fixture.env, "OPENSMS_MAIN_PATH": str(fixture.raw),
               "OPENSMS_OUTPUT_DIR": str(fixture.output.parent),
               "OPENSMS_REPORT_ROOT": str(fixture.root), "EXPORT4DAY_SERVER_DATE": "2026-09-06"}
        process = fixture.command([sys.executable, "-B", str(PIPELINE / "process_card_slots.py")], env)
        verify_env = {**env, "PROCESS_JSON": json.dumps(process),
                      "ENRICH_JSON": json.dumps(process["enrichment"])}
        command = [sys.executable, "-B", str(PIPELINE / "opensms_verify_card_slot_fast.py"),
                   str(fixture.output), str(fixture.raw)]
        verified = fixture.command(command, verify_env)
        self.assertTrue(verified["headers_match"])
        self.assertEqual(verified["formula_error_matches"], 0)
        self.assertEqual(verified["window_metric_mismatch_count"], 0)
        book = load_workbook(fixture.output)
        sheet = book["sheet1"]
        headers = [c.value for c in sheet[1]]
        self.assertEqual(headers[3:6], ["换号状态", "koo是否2天", "激活后天数"])
        self.assertEqual(headers.count("koo是否2天"), 1)
        self.assertEqual(headers[6], "使用天数")
        self.assertEqual(headers.count("使用天数"), 1)
        self.assertEqual(headers[-1], "ICCID")
        self.assertGreaterEqual(sheet.column_dimensions["E"].width, 14)
        for row, record in zip(list(sheet.values)[1:], fixture.records):
            self.assertEqual(row[4], None if record[0] == "Chatr" else "2天")
        # A stale field must not quietly pass the new independent schema.
        sheet.cell(1, 5).value = "使用天数"
        book.save(fixture.output)
        book.close()
        rejected = subprocess.run(command, env=verify_env, capture_output=True, text=True, timeout=45)
        if rejected.returncode == 0:
            self.assertFalse(json.loads(rejected.stdout)["headers_match"])
        self.assertEqual(before, [fixture.digest(p) for p in (fixture.raw, fixture.system, fixture.cache)])

    def test_old_candidate_field_is_rejected_before_backend(self):
        from tests.test_opensms_set_carrier_candidates_guard import ReplyRateSendGateTests
        for source, target, route in (
            ("B1-测试", "B11-测试", "b1_b11_use2_lt240_to_b11"),
            ("B1V-测试", "B11V-测试", "b1v_b11v_use2_lt240_to_b11v"),
        ):
            with self.subTest(route=route):
                candidate = {"route": route, "vZonePriority": 2, "tenDaySendThreshold": 240,
                    "useDaysFieldEquals": "2天", "sourceCarrierNames": [source], "targetCarrierName": target,
                    "rows": [{"电话": "14160000001", "原运营商": source, "opr": "Koodo",
                              "使用天数": "2天", "10日发送窗口数": 239,
                              "发送窗口数": 49, "回复窗口数": 1}]}
                result = ReplyRateSendGateTests().run_candidate(candidate, target)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("use_days_not_exactly_2", result.stderr)


if __name__ == "__main__":
    unittest.main()
