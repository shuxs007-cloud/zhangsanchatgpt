from copy import deepcopy
import hashlib
import importlib.util
import json
import os
import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
MODULE = ROOT / 'services/opensms_document_api/pipeline/card_retention_report.py'
SPEC = importlib.util.spec_from_file_location('retention_report', MODULE)
R = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(R)
RUN_ID = '20260917T001000+0800'


def source(count=201, raw_rows=None, run_id=RUN_ID):
    rows = raw_rows if raw_rows is not None else [
        {'id': str(i), 'cardNo': str(2000000000+i), 'expireBefore': '2026-09-10',
         'expireAfter': '2026-09-23', 'chatId': '12345678901234567890',
         'createTime': '2026-09-16 10:00:00', 'remark': '=untrusted text'} for i in range(count)]
    if raw_rows is None and len(rows)>1:
        rows[-1]['cardNo'] = rows[0]['cardNo']
    count = len(rows)
    pages = [{'pageNum': n+1, 'pageSize': 200, 'response': {'code': 200, 'total': count,
        'rows': rows[n*200:(n+1)*200]}} for n in range(max(1,(count+199)//200))]
    phones = list(dict.fromkeys(R.short_phone(r['cardNo']) for r in rows))
    return {'status': 'passed', 'task': 'A', 'scope': 'card_retention_report', 'runId': run_id,
            'readOnly': True, 'backendWrites': 0, 'clientClosed': True, 'sharedTargetOwnedByRunner': True,
            'silentBrowserMode': True, 'foregroundActivationAttempted': False,
            'independentSourceReadbackPassed': True, 'passes': [{'pages': deepcopy(pages)}, {'pages': deepcopy(pages)}],
            'retention': {'totalRecords': count, 'fetchedRecords': count, 'cardNos': phones, 'uniqueCardCount': len(phones)},
            'readConcurrency': {'maxActive': 1}, 'startedAt': '2026-09-16T15:00:00Z', 'completedAt': '2026-09-16T15:00:05Z'}


class CardRetentionTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.directory = Path(tmp.name)

    def evidence(self, payload=None):
        path = self.directory / 'source.json'
        path.write_text(json.dumps(payload or source()), encoding='utf-8')
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        run = (payload or source())['runId']
        self.env = {'OPENSMS_CARD_RETENTION_REPORT_PATH': str(path),
                    'OPENSMS_CARD_RETENTION_REPORT_SHA256': digest,
                    'OPENSMS_CARD_RETENTION_REPORT_REQUIRED': '1', 'OPENSMS_RUN_ID': run}
        return R.load_source(path, digest, run)

    def book(self):
        book = Workbook(); ws = book.active; ws.title = 'sheet1'
        ws.append(['电话', '运营商', '激活后天数', 'opr', '卡槽', 'ICCID'])
        ws.append(['12000000000', 'B1', 2, 'Kooda', '130_1', '8912230200081836700F'])
        ws.append(['2999999999', 'A1', 3, 'Chatr', '131_1', ''])
        ws.column_dimensions['D'].width = 32
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[2]: cell.fill = PatternFill('solid', fgColor='FFC7CE')
        book.create_sheet('处理汇总')
        return book

    def reread(self, book, proof):
        file = self.directory / 'output.xlsx'
        book.save(file); book.close()
        reread = load_workbook(file, read_only=True)
        try:
            with patch.dict(os.environ, self.env):
                return R.review_retention_report(reread, proof)
        finally: reread.close()

    def test_complete_export_position_values_precision_and_readback(self):
        evidence = self.evidence(); book = self.book()
        proof = R.add_retention_report(book, book['sheet1'], evidence)
        self.assertEqual(book['sheet1']['D1'].value, '需要保留时间')
        self.assertEqual(book['sheet1']['D2'].value, '2026-09-16 10:00:00')
        self.assertIsNone(book['sheet1']['D3'].value)
        self.assertEqual(book['sheet1']['G2'].value, '8912230200081836700F')
        self.assertEqual(book['sheet1'].column_dimensions['E'].width, 32)
        self.assertEqual(book['卡保留原始字段']['G2'].data_type, 's')
        self.assertEqual(book['卡保留全部记录']['E2'].value, '12345678901234567890')
        self.assertEqual(proof['record_count'], 201)
        self.assertEqual(proof['unique_phone_count'], 200)
        self.assertEqual(proof['filled_rows'], 1)
        self.assertTrue(self.reread(book, proof)['matches'])

    def test_latest_request_not_expiry_with_duplicate_source_and_main_phones(self):
        rows = [{'cardNo': '2000000000', 'createTime': when, 'expireAfter': '2099-01-01', 'id': str(i)}
                for i, when in enumerate(['2026-09-12 12:00:00', '2026-09-17 01:00:00', '2026-09-14 10:00:00'])]
        evidence = self.evidence(source(raw_rows=rows)); book = self.book()
        book['sheet1']['A3'] = '2000000000'
        proof = R.add_retention_report(book, book['sheet1'], evidence)
        self.assertEqual(book['sheet1']['D2'].value, '2026-09-17 01:00:00')
        self.assertEqual(book['sheet1']['D3'].value, '2026-09-17 01:00:00')
        self.assertEqual(book['卡保留全部记录'].max_row, 4)
        self.assertEqual(proof['filled_rows'], 2)
        self.assertTrue(self.reread(book, proof)['matches'])

    def test_empty_or_bad_request_time_keeps_history_and_blanks_main(self):
        for value in ['', None, '2026-02-30 00:00:00', '=formula', '2026-09-17']:
            rows = [{'cardNo': '2000000000', 'createTime': value, 'id': '1'},
                    {'cardNo': '2000000000', 'createTime': '2026-09-12 12:00:00', 'id': '2'}]
            evidence = self.evidence(source(raw_rows=rows)); book = self.book()
            proof = R.add_retention_report(book, book['sheet1'], evidence)
            self.assertIsNone(book['sheet1']['D2'].value)
            self.assertEqual(proof['unknown_request_time_phone_count'], 1)
            self.assertTrue(self.reread(book, proof)['matches'])

    def test_empty_source_still_has_column_and_detail_headers(self):
        evidence = self.evidence(source(0)); book = self.book()
        proof = R.add_retention_report(book, book['sheet1'], evidence)
        self.assertEqual(proof['filled_rows'], 0)
        self.assertEqual(book['卡保留全部记录'].max_row, 1)
        self.assertTrue(self.reread(book, proof)['matches'])

    def test_invalid_proofs_and_old_g_evidence_rejected(self):
        for field, value in [('task','G'), ('scope','other'), ('status','failed'), ('runId','other'),
                             ('backendWrites',1), ('readOnly',False), ('passes',[]), ('clientClosed',False)]:
            data = source(); data[field] = value
            with self.assertRaises(ValueError): R.validate_source(data, RUN_ID)

    def test_pagination_content_and_hash_tampering_rejected(self):
        data = source(); data['passes'][1]['pages'][1]['response']['rows'][0]['cardNo'] = '2999999999'
        with self.assertRaises(ValueError): R.validate_source(data, RUN_ID)
        data = source(); data['passes'][0]['pages'].pop()
        with self.assertRaises(ValueError): R.validate_source(data, RUN_ID)
        evidence = self.evidence()
        with self.assertRaises(ValueError): R.load_source(evidence['path'], '0'*64, RUN_ID)

    def test_output_tampering_fails_independent_review(self):
        evidence = self.evidence(); book = self.book()
        proof = R.add_retention_report(book, book['sheet1'], evidence)
        book['sheet1']['D2'] = '2099-01-01'
        self.assertFalse(self.reread(book, proof)['matches'])

    def test_long_text_and_duplicate_column_rejected(self):
        book = self.book()
        with self.assertRaises(ValueError): R.append_text_row(book.active, ['x'*32768])
        evidence = self.evidence(); R.add_retention_report(book, book['sheet1'], evidence)
        with self.assertRaises(ValueError): R.add_retention_report(book, book['sheet1'], evidence)
        book.close()

    def test_main_runner_wiring_and_g_entry_removed(self):
        runner = (ROOT/'scripts/opensms_patrol_runner.py').read_text()
        for marker in ['jobs["card_retention_report"] = pool.submit', 'verify_card_retention_acquisition',
                       '"OPENSMS_CARD_RETENTION_REPORT_REQUIRED": "1"', '"cardRetentionReport": retention_proof']:
            self.assertIn(marker, runner)
        # The deleted card-retention G entry stays deleted. The newly authorized
        # weekend-rest G has a different scope and no standalone business entry.
        for name in ['run_opensms_task_g.sh', 'opensms_collect_task_g.mjs']:
            self.assertFalse((ROOT/'scripts'/name).exists())
        self.assertIn('previous_day_total_send_weekend_rest_v1', (ROOT/'scripts/opensms_task_g.py').read_text())

    def test_final_carrier_refresh_preserves_retention_column_and_all_details(self):
        from test_opensms_refresh_workbook_carriers_fast import FastCarrierRefreshTests, MODULE as refresh
        evidence = self.evidence(); book = self.book()
        proof = R.add_retention_report(book, book['sheet1'], evidence)
        source_path = self.directory / 'before-refresh.xlsx'
        book.save(source_path); book.close()
        args = FastCarrierRefreshTests().fresh_args(self.directory, source_path, 'retention',
            [{'phone': '12000000000', 'liveCarrierName': 'A11'}])
        result = refresh.run(args)
        self.assertEqual(result['changed'], 1)
        reread = load_workbook(args.output, read_only=True)
        try:
            self.assertEqual(reread['sheet1']['B2'].value, 'A11')
            with patch.dict(os.environ, self.env):
                self.assertTrue(R.review_retention_report(reread, proof)['matches'])
        finally:
            reread.close()

    def test_document_resume_rejects_missing_or_unverified_retention_proof(self):
        from test_opensms_runner_no_google import RunnerNoGoogleTests, DEPENDENCY, RUNNER
        fixture = RunnerNoGoogleTests(); fixture.setUp(); self.addCleanup(fixture.doCleanups)
        fixture.document_completion(DEPENDENCY)
        document = fixture.runner.finish_document()
        self.assertEqual(fixture.runner.finish_document()['cardRetentionReport'], document['cardRetentionReport'])
        path = Path(document['cardReceipt'])
        raw = json.loads(path.read_text()); receipt = raw.get('cardSlotReport') or raw
        receipt['verification']['card_retention_report_review']['matches'] = False
        path.write_text(json.dumps(raw))
        # Hold the normal stage/file identity checks fixed to exercise the semantic gate.
        fixture.runner.reusable_stage = unittest.mock.Mock(return_value=document)
        with self.assertRaisesRegex(RUNNER.PatrolError, '卡保留独立核验'):
            fixture.runner.finish_document()

    def test_real_document_builder_and_independent_verifier(self):
        from test_opensms_document_no_google import NoGoogleDocumentTests
        fixture = NoGoogleDocumentTests(); fixture.setUp(); self.addCleanup(fixture.doCleanups)
        rows = [{'cardNo': '14385550001', 'createTime': '2026-09-16 13:14:15', 'id': '1'},
                {'cardNo': '14385550001', 'createTime': '2026-09-15 11:22:33', 'id': '2'},
                {'cardNo': '2999999999', 'createTime': '2026-09-16 14:00:00', 'id': '3'}]
        self.evidence(source(raw_rows=rows, run_id='offline-no-google'))
        fixture.env.update(self.env)
        pipeline = ROOT / 'services/opensms_document_api/pipeline'
        fixture.env.update({'OPENSMS_MAIN_PATH': str(fixture.raw), 'OPENSMS_OUTPUT_DIR': str(fixture.output.parent),
                            'EXPORT4DAY_SERVER_DATE': '2026-09-06'})
        process = fixture.command([sys.executable, str(pipeline/'process_card_slots.py')])
        process_path = self.directory / 'process.json'
        process_path.write_text(json.dumps(process), encoding='utf-8')
        fixture.env['PROCESS_JSON_PATH'] = str(process_path)
        enrich_path = self.directory / 'enrichment.json'
        enrich_path.write_text(json.dumps(process['enrichment']), encoding='utf-8')
        fixture.env['ENRICH_JSON_PATH'] = str(enrich_path)
        verification = fixture.command([sys.executable, str(pipeline/'opensms_verify_card_slot_fast.py'),
                                        str(fixture.output), str(fixture.raw)])
        proof = process['card_retention_report']
        self.assertEqual(proof['filled_rows'], 1)
        self.assertEqual(proof['unmatched_phone_count'], 1)
        self.assertTrue(verification['card_retention_report_review']['matches'], verification['card_retention_report_review'])
        self.assertTrue(verification['headers_match'])
        self.assertIn('需要保留时间及卡保留全量明细', (pipeline/'build_card_slot_report.sh').read_text())
        book = load_workbook(fixture.output, read_only=True)
        self.assertEqual(book['sheet1']['G1'].value, '使用天数')
        self.assertEqual(book['sheet1']['H1'].value, '需要保留时间')
        self.assertEqual(book['sheet1']['H2'].value, '2026-09-16 13:14:15')
        self.assertEqual(book['sheet1']['I1'].value, 'opr')
        headers = [cell.value for cell in book['sheet1'][1]]
        room_column = headers.index('机房') + 1
        self.assertEqual(book['sheet1'].cell(2, room_column).value, '六號機房')
        self.assertEqual(process['room_counts']['六號機房'], 8)
        self.assertNotIn('四號機房', process['room_counts'])
        summary = dict(book['处理汇总'].iter_rows(min_row=2, max_col=2, values_only=True))
        self.assertEqual(summary['六號機房行数'], 8)
        self.assertEqual(summary['五號機房设备配置'], '68–71、151–173（27台）')
        self.assertEqual(summary['六號機房设备配置'], '130–149、174–192、194–195（41台）')
        self.assertEqual(summary['七號機房设备配置'], '94–100（7台）')
        self.assertEqual(summary['七號機房行数'], 0)
        self.assertEqual(summary['配置设备总数'], 137)
        self.assertTrue(verification['room_configuration_review']['matches'])
        self.assertNotIn('四號機房行数', summary)
        book.close()


if __name__ == '__main__': unittest.main()
