"""Offline report/library reconciliation; no network or production data."""
import hashlib
import importlib.util
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("reconcile_test", ROOT / "scripts/opensms_task_f_reconcile_document.py")
RECONCILE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(RECONCILE)
RUN_ID = "20260907T020857+0800"
KEY = "8912230200278104811"
MISSING = "8912230200278150145"


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="task-f-reconcile-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.receipts = self.root / "receipts"
        self.receipts.mkdir()
        self.document = self.root / f"最终状态_{RUN_ID}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet1"
        sheet.append(["电话", "卡槽", "opr", "激活日期", "截止日", "ICCID"])
        sheet.append(["12493568234", "221_15.01", "302220 Koodo", None, None, KEY + "F"])
        sheet.append(["12493568323", "221_6.01", "Telus1", None, None, MISSING + "f"])
        sheet.append(["12493568066", "235_19.01", "Kooda", None, None, KEY])
        sheet.append(["12493568067", "235_20.01", "Chatr", None, None, "8912230200000000000F"])
        sheet.append(["12493568068", "235_21.01", "Kooda", None, None, 12345])
        workbook.save(self.document)
        workbook.close()
        self.database = self.root / "library.sqlite3"
        with sqlite3.connect(self.database) as connection:
            connection.execute("CREATE TABLE kooda_activation (" + ",".join(name + " TEXT" for name in RECONCILE.DB_FIELDS) + ")")
            record = {"query_iccid": KEY, "source_iccid": KEY + "F", "activation_date": "2026-09-06",
                      "cutoff_date": "2026-10-05", "lookup_status": "found", "queried_at": "2026-09-06T18:02:02.620Z",
                      "purchase_date_raw": "2026-08-31", "act_request_date": "2026-09-06 12:00",
                      "activation_date_raw": "2026-09-06 12:03", "sim_number": KEY, "sku": "SKU1",
                      "data_source": "grabandgo_incremental_lookup", "source_data_json": '[{"source":"full"}]', "record_schema_version": "2"}
            connection.execute("INSERT INTO kooda_activation VALUES (" + ",".join("?" for _ in RECONCILE.DB_FIELDS) + ")",
                               [record.get(name) for name in RECONCILE.DB_FIELDS])
        self.state_path = self.receipts / f"opensms_patrol_runner_state_{RUN_ID}.json"
        state = {"runId": RUN_ID, "status": "success", "stages": {
            "refresh": {"finalWorkbook": str(self.document), "finalWorkbookSha256": hashlib.sha256(self.document.read_bytes()).hexdigest()},
            "taskF": {"database": str(self.database)}}}
        self.state_path.write_text(json.dumps(state), encoding="utf-8")
        self.output = self.root / "new-reconciliation"

    def test_trailing_f_exact_match_deduplicates_and_preserves_sources(self):
        original_doc, original_db = self.document.read_bytes(), self.database.read_bytes()
        result = RECONCILE.reconcile(self.state_path, self.output)
        self.assertEqual((result["documentKoodaRows"], result["documentMissingActivationRows"], result["queryIccidCount"]), (4, 4, 2))
        self.assertEqual((result["readyForUploadCount"], result["pendingCount"], result["excludedCount"], result["duplicateIccidRows"]), (1, 1, 1, 1))
        self.assertEqual(result["rows"][0]["databaseRecord"]["sourceData"], [{"source": "full"}])
        self.assertEqual(result["rows"][1]["status"], "pending_no_local_record")
        self.assertEqual(self.document.read_bytes(), original_doc)
        self.assertEqual(self.database.read_bytes(), original_db)
        # Artifact exports need not have a cached dimension; inspect actual cells.
        workbook = load_workbook(result["files"]["manualUploadWorkbook"], read_only=False)
        sheet = workbook["截止日"]
        self.assertEqual([cell.value for cell in sheet[1]], ["ICCID", "截止日", "激活日期"])
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual((sheet["A2"].value, sheet["A2"].data_type, sheet["A2"].number_format), (KEY, "s", "@"))
        self.assertEqual(sheet["B2"].value.strftime("%Y-%m-%d"), "2026-10-05")
        workbook.close()
        self.assertFalse(result["networkAccess"])
        self.assertFalse(result["cloudUploadAttempted"])

    def test_hash_mismatch_is_rejected_before_output(self):
        self.document.write_bytes(self.document.read_bytes() + b"changed")
        with self.assertRaisesRegex(ValueError, "SHA256"):
            RECONCILE.reconcile(self.state_path, self.output)
        self.assertFalse(self.output.exists())

    def test_existing_directory_is_not_overwritten(self):
        self.output.mkdir()
        marker = self.output / "keep.txt"
        marker.write_text("keep")
        with self.assertRaises(FileExistsError):
            RECONCILE.reconcile(self.state_path, self.output)
        self.assertEqual(marker.read_text(), "keep")

    def test_missing_database_is_not_created(self):
        self.database.unlink()
        with self.assertRaisesRegex(ValueError, "禁止创建"):
            RECONCILE.reconcile(self.state_path, self.output)
        self.assertFalse(self.database.exists())
        self.assertFalse(self.output.exists())

    def test_iccid_removes_only_one_suffix_and_rejects_numeric(self):
        self.assertEqual(RECONCILE.normalize_iccid(KEY + "F"), KEY)
        self.assertEqual(RECONCILE.normalize_iccid(KEY + "f"), KEY)
        for invalid in (KEY + "FF", 1234567890123456789, " " + KEY, "８" * 19):
            self.assertEqual(RECONCILE.normalize_iccid(invalid), "")

    def raw_source(self, rows=None):
        source = self.root / f"原始导出_{RUN_ID}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet1"
        sheet.append(["电话", "ICCID", "opr", "卡槽", "激活日期", "截止日"])
        for row in rows or [
            ["12493568234", KEY + "F", "Koodo", "221_15.01", None, None],
            ["12493568323", MISSING, "Telus1", "221_6.01", None, None],
            ["12493568067", "8912230200000000000F", "Chatr", "235_20.01", None, None],
            ["12493568068", "8912230200000000001F", "Kooda", "235_21.01", "2026-09-06", None],
        ]:
            sheet.append(row)
        workbook.save(source)
        workbook.close()
        state = json.loads(self.state_path.read_text())
        state["stages"]["acquire"] = {"sourceArchive": str(source), "sourceVerified": True,
                                     "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest()}
        self.state_path.write_text(json.dumps(state))
        return source

    def test_raw_source_export_excludes_other_database_rows_and_preserves_inputs(self):
        source = self.raw_source()
        with sqlite3.connect(self.database) as conn:
            conn.execute("INSERT INTO kooda_activation(query_iccid,activation_date,cutoff_date,lookup_status) "
                         "VALUES ('8912230200000000999','2026-09-06','2026-10-05','found')")
        before = (source.read_bytes(), self.database.read_bytes())
        result = RECONCILE.reconcile(self.state_path, self.output, source_kind="system-export")
        self.assertEqual(result["sourceKind"], "system-export")
        self.assertEqual((result["documentMissingActivationRows"], result["readyForUploadCount"], result["pendingCount"]), (2, 1, 1))
        self.assertEqual((source.read_bytes(), self.database.read_bytes()), before)
        verification = json.loads((self.output / "template_verification.json").read_text())
        self.assertTrue(verification["verificationPassed"])
        self.assertEqual(verification["recordCount"], 1)
        self.assertEqual(verification["scope"], "source_document_missing_activation_only")
        self.assertIn(MISSING, Path(result["files"]["report"]).read_text())

    def test_raw_source_hash_or_freshness_failure_stops_before_output(self):
        source = self.raw_source()
        source.write_bytes(source.read_bytes() + b"changed")
        with self.assertRaisesRegex(ValueError, "SHA256"):
            RECONCILE.reconcile(self.state_path, self.output, source_kind="system-export")
        self.assertFalse(self.output.exists())
        self.raw_source()
        state = json.loads(self.state_path.read_text())
        state["stages"]["acquire"]["sourceVerified"] = False
        self.state_path.write_text(json.dumps(state))
        with self.assertRaisesRegex(ValueError, "新鲜源校验"):
            RECONCILE.read_report(self.state_path, source_kind="system-export")

    def test_source_present_activation_excluded_even_when_library_has_date(self):
        self.raw_source([["12493568234", KEY, "Kooda", "1_1", "2026-09-06", None]])
        _, meta, rows = RECONCILE.read_report(self.state_path, source_kind="system-export")
        self.assertEqual(meta["documentMissingActivationRows"], 0)
        self.assertEqual(rows, [])

    def test_source_mixed_present_and_missing_same_iccid_rejected(self):
        self.raw_source([["12493568234", KEY, "Kooda", "1_1", None, None],
                         ["12493568235", KEY + "F", "Kooda", "1_2", "2026-09-06", None]])
        with self.assertRaisesRegex(ValueError, "同时存在"):
            RECONCILE.read_report(self.state_path, source_kind="system-export")

    def test_source_invalid_nonblank_activation_is_not_a_missing_date(self):
        self.raw_source([["12493568234", KEY, "Kooda", "1_1", "invalid", None]])
        with self.assertRaises(ValueError):
            RECONCILE.read_report(self.state_path, source_kind="system-export")

    def test_source_cutoff_conflict_does_not_overwrite_system(self):
        self.raw_source([["12493568234", KEY, "Kooda", "1_1", None, "2026-10-06"]])
        with self.assertRaisesRegex(ValueError, "截止日与F库不一致"):
            RECONCILE.reconcile(self.state_path, self.output, source_kind="system-export")
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
