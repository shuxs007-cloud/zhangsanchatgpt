from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from scripts.opensms_parallel_audit import (
    PASS,
    PathResolver,
    check_c2_routing_rule,
    check_reply_rate_send_gate,
    check_rest_zone_entry_enabled,
    check_v_zone_priority_rule,
)


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "opensms_generate_opr_zero_candidates.py"


def write_c2_statistics_receipt(root: Path, send_size: int) -> Path:
    path = root / "c2_statistics.json"
    path.write_text(json.dumps({
        "replyRates": [{
            "group": "C2",
            "status": "ok",
            "endpoint": "/api/admin/statisticsList",
            "dailyRow": {"days": "2026-08-25", "sendSize": send_size},
        }],
    }), encoding="utf-8")
    return path


class CarrierRoutingTests(unittest.TestCase):
    def test_b1_standard_and_v_group_lt240_boundaries_follow_priority(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "export4day.xlsx"
            output_dir = root / "snapshots"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "电话", "运营商", "opr", "卡槽", "状态", "10日发送窗口数",
                "08-21发送窗口数", "08-21回复窗口数",
                "08-20发送窗口数", "08-20回复窗口数",
            ])
            rows = [
                ["14160000001", "B1-精-Te-120", "302220 Koodo", "101_1", "上架", 239, 90, 4, 81, 3],
                ["14160000002", "B1-精-Te-120", "302220 Koodo", "101_2", "上架", 240, 49, 1, 20, 1],
                ["14160000003", "B11-精-Te-120", "302220 Koodo", "101_3", "上架", 239, 90, 4, 81, 3],
                ["14160000004", "B12-精-Te-120", "302220 Koodo", "101_4", "上架", 300, 50, 0, 10, 1],
                ["14160000005", "B12-精-Te-120", "302220 Koodo", "101_5", "上架", 300, 50, 2, 10, 1],
                ["14160000006", "B12-精-Te-120", "302220 Koodo", "101_6", "上架", 300, 100, 5, 10, 1],
                ["14160000007", "B12-精-Te-120", "302220 Koodo", "101_7", "上架", 300, 49, 1, 10, 1],
                ["14160000008", "B13-精-Te-50", "302220 Koodo", "101_8", "上架", 300, 100, 5, 10, 1],
                ["14160000009", "B13-精-Te-50", "302220 Koodo", "101_9", "上架", 300, 100, 4, 10, 1],
                ["14160000010", "B13-精-Te-50", "302220 Koodo", "101_10", "上架", 300, 0, 0, 10, 1],
                ["14160000011", "A12-精普-C120", "302720 chatr", "101_11", "上架", 10, 1, 1, 10, 1],
                ["14160000012", "A12-精普-C120", "302720 chatr", "101_12", "上架", 10, 100, 4, 10, 1],
                ["14160000013", "A12-精普-C120", "302720 chatr", "101_13", "上架", 10, 100, 3, 10, 1],
                ["14160000014", "A1-精普-C120", "302220 Koodo", "101_14", "上架", 0, 0, 0, 0, 0],
                ["14160000015", "A21-股票-C-100", "302220 Kooda", "101_15", "上架", 0, 0, 0, 0, 0],
                ["14160000016", "A3-打T精普-C120", "302220 kOoDo", "101_16", "上架", 0, 0, 0, 0, 0],
                ["14160000017", "A1-精普-C120", "302720 chatr", "101_17", "上架", 10, 50, 0, 10, 1],
                ["14160000018", "A11-精普-C50", "302720 Chatr", "101_18", "上架", 10, 50, 0, 10, 1],
                ["14160000019", "A12-精普-C120", "302720 CHATR", "101_19", "上架", 10, 50, 0, 10, 1],
                ["14160000020", "A3-打T精普-C120", "302720 chatr", "101_20", "上架", 10, 50, 0, 10, 1],
                ["14160000021", "A3-打T精普-C120", "302720 chatr", "101_21", "上架", 10, 49, 0, 10, 1],
                ["14160000022", "A1-精普-C120", "302730 Fido", "101_22", "上架", 10, 50, 0, 10, 1],
                ["14160000023", "A1-精普-C120", "302720 chatr", "101_23", "上架", 10, 100, 1, 10, 1],
                ["14160000024", "A12-精普-C120", "302720 chatr", "101_24", "上架", 10, 50, 1, 10, 1],
                ["14160000025", "C2-测试区1-C150", "302730 Fido", "101_25", "上架", 300, 100, 4, 20, 0],
                ["14160000026", "C2-测试区1-C150", "302730 Fido", "101_26", "上架", 300, 100, 5, 20, 1],
                ["14160000027", "C2-测试区1-C150", "302730 Fido", "101_27", "上架", 300, 0, 0, 20, 1],
                ["14160000030", "A2-股票-C-120", "302220 Koodo", "101_30", "上架", 0, 0, 0, 0, 0],
                ["14160000031", "A11-精普-C50", "302720 chatr", "101_31", "上架", 10, 100, 4, 10, 1],
                ["14160000032", "B1V-精-Te-120", "Telus1", "101_32", "上架", 239, 90, 4, 81, 3],
                ["14160000033", "B11V-精-Te-120", "Telus1", "101_33", "上架", 239, 90, 4, 81, 3],
                ["14160000034", "B1V-精-Te-120", "Telus1", "101_34", "上架", 240, 49, 1, 20, 1],
                ["14160000035", "B12V-精-Te-120", "Telus1", "101_35", "上架", 300, 50, 0, 10, 1],
                ["14160000036", "B12V-精-Te-120", "Telus1", "101_36", "上架", 300, 50, 2, 10, 1],
                ["14160000037", "A12-精普-C120", "Telus1", "101_37", "上架", 0, 0, 0, 0, 0],
                ["14160000038", "B1-精-Te-120", "Telus1", "101_38", "上架", 240, 90, 4, 81, 3],
                ["14160000039", "B1-精-Te-120", "Kooda", "101_39", "上架", 299, 90, 4, 81, 3],
                ["14160000040", "B1-精-Te-120", "Koodo", "101_40", "上架", 300, 90, 4, 81, 3],
            ]
            for row in rows:
                sheet.append(row)
            workbook.save(workbook_path)
            statistics_path = write_c2_statistics_receipt(root, 200)

            completed = subprocess.run(
                [
                    sys.executable, str(SCRIPT), str(workbook_path), "--out-dir", str(output_dir),
                    "--now-bj", "2026-08-25T20:00:00+08:00",
                    "--c2-statistics-receipt", str(statistics_path),
                ],
                cwd=ROOT,
                check=True,
                capture_output=True,
                text=True,
            )
            manifest = json.loads(completed.stdout)

            def load_route(name: str) -> dict:
                return json.loads(Path(manifest["outputs"][name]["path"]).read_text(encoding="utf-8"))

            zero = load_route("b1_b11_b12_today_gte50_zero_to_e")
            under_240 = load_route("b1_b11_use2_lt240_to_b11")
            at_least_240 = load_route("b1_b11_gte240_to_b12")
            v_under_240 = load_route("b1v_b11v_use2_lt240_to_b11v")
            v_at_least_240 = load_route("b1v_b11v_gte240_to_b12v")
            chatr_zero = load_route("chatr_to_e")
            c2_low_reply = load_route("c2_previous_day_reply_rate_lt5_to_a11")
            a_low = load_route("a_zone_lt4_to_a12")
            a_high = load_route("a_zone_gte4_to_a1")
            a_kooda = load_route("a_zone_kooda_to_b12")

            for route in (under_240, v_under_240):
                self.assertEqual("2天", route["rows"][0]["koo是否2天"])
                self.assertIn("koo是否2天依据", route["rows"][0])
                self.assertNotIn("使用天数", route["rows"][0])

            self.assertEqual([], zero["rows"])
            self.assertEqual(
                ["14160000001"],
                [row["phone"] for row in under_240["rows"]],
            )
            self.assertEqual(
                ["14160000002", "14160000038", "14160000039", "14160000040"],
                [row["phone"] for row in at_least_240["rows"]],
            )
            self.assertEqual(["14160000032"], [row["phone"] for row in v_under_240["rows"]])
            self.assertEqual(["14160000034"], [row["phone"] for row in v_at_least_240["rows"]])
            self.assertEqual("B11V-精-Te-120", v_under_240["targetCarrierName"])
            self.assertEqual("B12V-精-Te-120", v_at_least_240["targetCarrierName"])
            self.assertEqual(2, v_under_240["vZonePriority"])
            self.assertEqual("2天", v_under_240["useDaysFieldEquals"])
            self.assertEqual(3, v_at_least_240["vZonePriority"])
            self.assertEqual(240, under_240["tenDaySendThreshold"])
            self.assertEqual(240, at_least_240["tenDaySendThreshold"])
            self.assertIn("lt240", Path(manifest["outputs"]["b1_b11_use2_lt240_to_b11"]["path"]).name)
            self.assertIn("gte240", Path(manifest["outputs"]["b1_b11_gte240_to_b12"]["path"]).name)
            self.assertEqual("E-T休息区", zero["targetCarrierName"])
            self.assertEqual(
                [],
                [row["phone"] for row in chatr_zero["rows"]],
            )
            self.assertEqual("E-C休息区", chatr_zero["targetCarrierName"])
            self.assertEqual(
                sorted(["A1-精普-C120", "A11-精普-C50", "A12-精普-C120", "A3-打T精普-C120"]),
                chatr_zero["sourceCarrierNames"],
            )
            self.assertEqual(["14160000025"], [row["phone"] for row in c2_low_reply["rows"]])
            self.assertEqual("A11-精普-C50", c2_low_reply["targetCarrierName"])
            self.assertAlmostEqual(0.0, c2_low_reply["rows"][0]["前一天回复率"])
            self.assertEqual(0.05, c2_low_reply["c2ReplyRateThresholdExclusive"])
            self.assertNotIn("b12_today_gte50_lt5_to_b13", manifest["outputs"])
            self.assertNotIn("b13_today_gte5_to_b12", manifest["outputs"])
            self.assertTrue(all(not output["target"].startswith("B13") for output in manifest["outputs"].values()))
            self.assertNotIn("a12_today_gte4_to_a1", manifest["outputs"])
            self.assertEqual(["14160000022", "14160000023"], [row["phone"] for row in a_low["rows"]])
            self.assertEqual(["14160000012", "14160000031"], [row["phone"] for row in a_high["rows"]])
            self.assertEqual("08-21", a_high["rows"][0]["统计日期"])
            self.assertAlmostEqual(0.04, a_high["rows"][0]["当天回复率"])
            self.assertEqual("latest_day", a_high["replyRateWindow"])
            self.assertEqual(
                ["14160000014", "14160000016", "14160000037"],
                [row["phone"] for row in a_kooda["rows"]],
            )
            self.assertEqual("B12-精-Te-120", a_kooda["targetCarrierName"])
            self.assertEqual(
                sorted(["A1-精普-C120", "A11-精普-C50", "A12-精普-C120", "A3-打T精普-C120"]),
                a_kooda["sourceCarrierNames"],
            )
            self.assertEqual(240, manifest["b1B11Use2ToB11TenDaySendThreshold"])
            self.assertEqual(240, manifest["b1B11ToB12TenDaySendThreshold"])
            self.assertEqual(240, manifest["b1vB11vTenDaySendThreshold"])
            self.assertFalse(manifest["restZoneEntryEnabled"])
            self.assertEqual("disabled_by_current_user_20260922", manifest["restZoneEntryPolicy"])
            self.assertEqual(manifest["restZoneEntryPolicy"], zero["restZoneEntryPolicy"])
            self.assertEqual(manifest["restZoneEntryPolicy"], chatr_zero["restZoneEntryPolicy"])
            self.assertEqual(0, manifest["restZoneEntryCandidateCount"])
            self.assertEqual(2, manifest["restZoneEntryBQualifiedHeldCount"])
            self.assertEqual(4, manifest["restZoneEntryChatrQualifiedHeldCount"])
            self.assertEqual(
                ["b1_b11_b12_today_gte50_zero_to_e", "chatr_to_e"],
                manifest["restZoneEntryDisabledRoutes"],
            )
            self.assertEqual("E-T休息区", manifest["bZoneRestTarget"])
            self.assertEqual("E-C休息区", manifest["aZoneChatrRestTarget"])
            self.assertEqual(["A1", "A11", "A12", "A3"], manifest["chatrZeroReplySendGateGroups"])
            self.assertTrue(manifest["aZoneChatrZeroToECPrecedesReplySplit"])
            self.assertEqual(0.04, manifest["a12ReplyRateThresholdInclusive"])
            self.assertEqual(50, manifest["replyRateLatestSendThresholdInclusive"])
            self.assertEqual(
                ["A1", "A11", "A12", "B1", "B11", "B12", "B1V", "B11V", "B12V"],
                manifest["replyRateSendGateGroups"],
            )
            self.assertNotIn("b1_b11_gte240_to_b12", manifest["replyRateSendGateRoutes"])
            self.assertNotIn("b1_b11_use2_lt240_to_b11", manifest["replyRateSendGateRoutes"])
            self.assertNotIn("a_zone_kooda_to_b12", manifest["replyRateSendGateRoutes"])
            self.assertIn("chatr_to_e", manifest["replyRateSendGateRoutes"])
            self.assertEqual(0.05, manifest["c2ReplyRateThresholdExclusive"])
            self.assertEqual(4000, manifest["c2LatestDaySendAlertThresholdExclusive"])
            self.assertEqual("08-20", manifest["c2PreviousDay"])
            self.assertEqual(200, manifest["c2DailyStatisticsSendTotal"])
            self.assertTrue(manifest["c2DailyStatisticsSendAlertTriggered"])
            self.assertTrue(manifest["c2DailyStatisticsSendAlert"]["triggered"])
            self.assertEqual("evaluation_completed", manifest["c2DailyExecution"]["status"])
            self.assertTrue(manifest["c2DailyExecution"]["eligible"])
            self.assertEqual(21, manifest["c2DailyCutoffHourExclusiveBj"])
            self.assertEqual(1, manifest["c2DailyMaxExecutions"])
            c2_audit = check_c2_routing_rule(
                {"carrierRoutingManifest": manifest["manifest"]},
                PathResolver(root / "config.json", None),
            )
            self.assertEqual(PASS, c2_audit.status)
            reply_gate_audit = check_reply_rate_send_gate(
                {"carrierRoutingManifest": manifest["manifest"]},
                PathResolver(root / "config.json", None),
            )
            self.assertEqual(PASS, reply_gate_audit.status)
            rest_entry_audit = check_rest_zone_entry_enabled(
                {
                    "carrierRoutingManifest": manifest["manifest"],
                    "groups": [],
                },
                PathResolver(root / "config.json", None),
            )
            self.assertEqual(PASS, rest_entry_audit.status)
            candidate_path = Path(manifest["outputs"]["chatr_to_e"]["path"])
            original_candidate = candidate_path.read_text(encoding="utf-8")
            tampered = json.loads(original_candidate)
            tampered["rows"] = [{"phone": "14160000017", "原运营商": "A1-精普-C120",
                                 "发送窗口数": 50, "回复窗口数": 0, "opr": "Chatr"}]
            candidate_path.write_text(json.dumps(tampered, ensure_ascii=False), encoding="utf-8")
            rejected_audit = check_rest_zone_entry_enabled(
                {"carrierRoutingManifest": manifest["manifest"], "groups": []},
                PathResolver(root / "config.json", None),
            )
            self.assertNotEqual(PASS, rejected_audit.status)
            candidate_path.write_text(original_candidate, encoding="utf-8")
            v_zone_audit = check_v_zone_priority_rule(
                {"carrierRoutingManifest": manifest["manifest"]},
                PathResolver(root / "config.json", None),
            )
            self.assertEqual(PASS, v_zone_audit.status)
            self.assertEqual(["B13"], manifest["abolishedCarrierGroups"])
            self.assertEqual(["B13"], manifest["forbiddenTargetCarrierGroups"])
            self.assertEqual(4, manifest["excluded"]["b12_family_held_b13_abolished"])
            self.assertEqual(1, manifest["b12vHeldWhenPriority1NotMatched"])
            self.assertEqual(1, manifest["candidateWorkbookPasses"])
            self.assertTrue(manifest["vZoneCandidateGenerationMergedWithMainRouting"])
            self.assertTrue(manifest["vZoneWriteExecutionSerial"])
            self.assertEqual("2天", manifest["useDaysFieldEquals"])
            self.assertEqual(
                [1, 2, 3, "hold"],
                [rule["priority"] for rule in manifest["vZonePriorityRules"]],
            )
            self.assertEqual([
                "b1_b11_b12_today_gte50_zero_to_e",
                "b1v_b11v_use2_lt240_to_b11v",
                "b1v_b11v_gte240_to_b12v",
            ], manifest["vZoneRoutePriority"])
            self.assertEqual(3, manifest["excluded"]["existing_b13_held_pending_migration_target"])
            self.assertEqual(["A21", "A2"], manifest["protectedSourceCarrierGroups"])
            self.assertEqual(2, manifest["protectedSourceExcludedCount"])
            self.assertEqual({"A21": 1, "A2": 1}, manifest["protectedSourceExcludedCounts"])
            self.assertEqual(1, manifest["protectedA21ExcludedCount"])
            self.assertEqual(1, manifest["excluded"]["protected_a21_held_no_carrier_change"])
            self.assertEqual(1, manifest["excluded"]["protected_a2_held_no_carrier_change"])
            routed_phones = []
            for route in manifest["outputs"]:
                routed_phones.extend(row["phone"] for row in load_route(route)["rows"])
            self.assertEqual(len(routed_phones), len(set(routed_phones)))
            self.assertNotIn("14160000035", routed_phones)
            self.assertTrue({"14160000004", "14160000017", "14160000018",
                             "14160000019", "14160000020"}.isdisjoint(routed_phones))
            self.assertNotRegex(json.dumps(manifest, ensure_ascii=False), r"(?i)b3[123]")
            self.assertEqual([
                "a1_kini_latest_lt4_to_b12",
                "a_zone_kooda_to_b12",
                "b1_b11_b12_today_gte50_zero_to_e",
                "b1_b11_use2_lt240_to_b11",
                "b1_b11_gte240_to_b12",
                "b1v_b11v_use2_lt240_to_b11v",
                "b1v_b11v_gte240_to_b12v",
                "chatr_to_e",
                "c2_previous_day_reply_rate_lt5_to_a11",
                "c2_rule_correction_restore_to_c2",
                "a_zone_gte4_to_a1",
                "a_zone_lt4_to_a12",
            ], manifest["routePriority"])

    def test_c2_exact_boundaries_do_not_trigger(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "export4day.xlsx"
            output_dir = root / "snapshots"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "电话", "运营商", "opr", "卡槽", "状态", "10日发送窗口数",
                "08-21发送窗口数", "08-21回复窗口数",
                "08-20发送窗口数", "08-20回复窗口数",
            ])
            sheet.append([
                "14160000028", "C2-测试区1-C150", "302730 Fido", "101_28", "上架",
                4000, 4000, 200, 4000, 200,
            ])
            workbook.save(workbook_path)
            statistics_path = write_c2_statistics_receipt(root, 4000)

            completed = subprocess.run(
                [
                    sys.executable, str(SCRIPT), str(workbook_path), "--out-dir", str(output_dir),
                    "--now-bj", "2026-08-25T20:00:00+08:00",
                    "--c2-statistics-receipt", str(statistics_path),
                ],
                cwd=ROOT,
                check=True,
                capture_output=True,
                text=True,
            )
            manifest = json.loads(completed.stdout)
            candidate = json.loads(
                Path(manifest["outputs"]["c2_previous_day_reply_rate_lt5_to_a11"]["path"])
                .read_text(encoding="utf-8")
            )

            self.assertEqual([], candidate["rows"])
            self.assertEqual(4000, manifest["c2DailyStatisticsSendTotal"])
            self.assertFalse(manifest["c2DailyStatisticsSendAlertTriggered"])
            self.assertFalse(manifest["c2DailyStatisticsSendAlert"]["triggered"])
            self.assertIsNone(manifest["c2DailyStatisticsSendAlert"]["message"])
            c2_audit = check_c2_routing_rule(
                {"carrierRoutingManifest": manifest["manifest"]},
                PathResolver(root / "config.json", None),
            )
            self.assertEqual(PASS, c2_audit.status)

    def test_c2_runs_once_per_beijing_day_and_never_at_or_after_21(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "export4day.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "电话", "运营商", "opr", "卡槽", "状态", "10日发送窗口数",
                "08-21发送窗口数", "08-21回复窗口数",
                "08-20发送窗口数", "08-20回复窗口数",
            ])
            sheet.append([
                "14160000029", "C2-测试区1-C150", "302730 Fido", "101_29", "上架",
                100, 100, 4, 100, 4,
            ])
            workbook.save(workbook_path)
            statistics_path = write_c2_statistics_receipt(root, 100)

            def run(output_dir: Path, now_bj: str) -> dict:
                completed = subprocess.run(
                    [
                        sys.executable, str(SCRIPT), str(workbook_path),
                        "--out-dir", str(output_dir), "--now-bj", now_bj,
                        "--c2-statistics-receipt", str(statistics_path),
                    ],
                    cwd=ROOT,
                    check=True,
                    capture_output=True,
                    text=True,
                )
                return json.loads(completed.stdout)

            daily_dir = root / "daily"
            first = run(daily_dir, "2026-08-25T20:00:00+08:00")
            second = run(daily_dir, "2026-08-25T20:30:00+08:00")
            after_cutoff = run(root / "after_cutoff", "2026-08-25T21:00:00+08:00")

            def c2_rows(manifest: dict) -> list[dict]:
                route_path = Path(manifest["outputs"]["c2_previous_day_reply_rate_lt5_to_a11"]["path"])
                return json.loads(route_path.read_text(encoding="utf-8"))["rows"]

            self.assertEqual(1, len(c2_rows(first)))
            self.assertEqual("evaluation_completed", first["c2DailyExecution"]["status"])
            self.assertTrue(first["c2DailyStatisticsSendAlertTriggered"])
            self.assertEqual([], c2_rows(second))
            self.assertEqual("skipped_already_executed", second["c2DailyExecution"]["status"])
            self.assertFalse(second["c2DailyStatisticsSendAlertTriggered"])
            self.assertEqual([], c2_rows(after_cutoff))
            self.assertEqual("skipped_after_cutoff", after_cutoff["c2DailyExecution"]["status"])
            self.assertFalse(after_cutoff["c2DailyStatisticsSendAlertTriggered"])
            resolver = PathResolver(root / "config.json", None)
            self.assertEqual(
                PASS,
                check_c2_routing_rule({"carrierRoutingManifest": second["manifest"]}, resolver).status,
            )
            self.assertEqual(
                PASS,
                check_c2_routing_rule({"carrierRoutingManifest": after_cutoff["manifest"]}, resolver).status,
            )


if __name__ == "__main__":
    unittest.main()
