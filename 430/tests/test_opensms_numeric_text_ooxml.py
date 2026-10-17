from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest
from xml.etree import ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "services/opensms_document_api/pipeline/opensms_verify_card_slot_fast.py"
spec = importlib.util.spec_from_file_location("numeric_text_verifier", SCRIPT)
verifier = importlib.util.module_from_spec(spec)
spec.loader.exec_module(verifier)
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class NumericTextOOXMLTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)

    def book(self, values):
        workbook = Workbook()
        for row in values:
            workbook.active.append(row)
        path = self.directory / "source.xlsx"
        workbook.save(path)
        return path

    def rewrite(self, source, replacements):
        target = self.directory / "modified.xlsx"
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, "w") as output:
            for item in original.infolist():
                output.writestr(item, replacements.pop(item.filename, original.read(item.filename)))
            for name, payload in replacements.items():
                output.writestr(name, payload)
        return target

    def same(self, path):
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            expected = verifier.review_numeric_text(book)
        finally:
            book.close()
        actual = verifier.review_numeric_text_ooxml(path)
        self.assertEqual(actual, expected)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)
        return actual

    def test_plain_inline_strings_and_numeric_precision_are_preserved(self):
        path = self.book([["0000123", "8912230200251591144", "1,000", "  -2.5% ", "中文", ""],
                          [12, 3.5, True, False, None, "not-a-number"]])
        result = self.same(path)
        self.assertEqual(result["preserved_text_numeric_cells"], 4)
        self.assertEqual(result["examples"][1]["value"], "8912230200251591144")

    def test_first_thirty_examples_keep_workbook_cell_order(self):
        path = self.book([[str(index), str(index + 1)] for index in range(40)])
        result = self.same(path)
        self.assertEqual(result["preserved_text_numeric_cells"], 80)
        self.assertEqual(result["examples"][-1]["cell"], "B15")

    def test_multiple_and_hidden_worksheets(self):
        path = self.book([["12"]])
        book = load_workbook(path)
        sheet = book.create_sheet("隐藏表")
        sheet.sheet_state = "hidden"
        sheet.append(["99"])
        book.save(path)
        result = self.same(path)
        self.assertEqual([row["sheet"] for row in result["examples"]], ["Sheet", "隐藏表"])

    def test_formula_cached_text_is_excluded_and_error_cells_remain_errors(self):
        path = self.book([["=1+1", "#DIV/0!", "123"]])
        with zipfile.ZipFile(path) as archive:
            tree = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        first = next(tree.iter(f"{NS}c"))
        first.set("t", "str")
        first.find(f"{NS}v").text = "42"
        path = self.rewrite(path, {"xl/worksheets/sheet1.xml": ET.tostring(tree)})
        result = self.same(path)
        self.assertEqual(result["preserved_text_numeric_cells"], 1)
        book = load_workbook(path, read_only=True, data_only=True)
        self.assertEqual(book.active["B1"].value, "#DIV/0!")
        book.close()

    def test_shared_strings_rich_text_and_escape_use_openpyxl_semantics(self):
        path = self.book([["12", "34", "56"]])
        with zipfile.ZipFile(path) as archive:
            tree = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            types = ET.fromstring(archive.read("[Content_Types].xml"))
        for index, cell in enumerate(tree.iter(f"{NS}c")):
            for child in list(cell):
                cell.remove(child)
            cell.set("t", "s")
            ET.SubElement(cell, f"{NS}v").text = str(index)
        ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                      PartName="/xl/sharedStrings.xml",
                      ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
        strings = (f'<sst xmlns="{NS[1:-1]}"><si><t>00012</t></si>'
                   '<si><r><t>3</t></r><r><t>4</t></r></si>'
                   '<si><t>_x005F_56</t></si></sst>')
        path = self.rewrite(path, {"xl/worksheets/sheet1.xml": ET.tostring(tree),
                                  "[Content_Types].xml": ET.tostring(types), "xl/sharedStrings.xml": strings})
        # Escaped underscore stays nonnumeric under openpyxl's shared decoder.
        self.assertEqual(self.same(path)["preserved_text_numeric_cells"], 2)

    def test_inline_rich_text_and_phonetic_text(self):
        path = self.book([["12"]])
        with zipfile.ZipFile(path) as archive:
            tree = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        cell = next(tree.iter(f"{NS}c"))
        cell.remove(cell.find(f"{NS}is"))
        inline = ET.SubElement(cell, f"{NS}is")
        for value in ("00", "12"):
            run = ET.SubElement(inline, f"{NS}r")
            ET.SubElement(run, f"{NS}t").text = value
        phonetic = ET.SubElement(inline, f"{NS}rPh", sb="0", eb="2")
        ET.SubElement(phonetic, f"{NS}t").text = "ignored"
        path = self.rewrite(path, {"xl/worksheets/sheet1.xml": ET.tostring(tree)})
        self.assertEqual(self.same(path)["examples"][0]["value"], "0012")

    def test_relative_and_absolute_sheet_targets_match(self):
        path = self.book([["12"]])
        with zipfile.ZipFile(path) as archive:
            tree = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        for child in tree:
            if child.attrib.get("Target", "").endswith("sheet1.xml"):
                child.set("Target", "worksheets/sheet1.xml")
        path = self.rewrite(path, {"xl/_rels/workbook.xml.rels": ET.tostring(tree)})
        self.assertEqual(self.same(path)["preserved_text_numeric_cells"], 1)

    def test_scanner_does_not_reopen_second_workbook_model(self):
        import inspect
        source = inspect.getsource(verifier.main)
        self.assertNotIn("final_formula_book", source)
        self.assertIn("review_numeric_text_ooxml(final_path)", source)
        self.assertIn("formula_error_cells", source)


if __name__ == "__main__":
    unittest.main()
