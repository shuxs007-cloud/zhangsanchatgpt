"""Offline renamed-room regression through extraction, consumers and readback."""
from datetime import datetime
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "services/opensms_document_api/pipeline"
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(PIPELINE))
from google_room_sheets import CURRENT_ROOM_SHEETS, LEGACY_ROOM_SHEETS
from google_system_missing import PORT_HEADERS, PORT_SHEET, read_rows, source_from_snapshot, write_sheet, review
from opensms_chatr_google_dates import freeze
from opensms_parallel_audit import audit_chatr_snapshot_rows


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


with patch.dict(os.environ, {"OPENSMS_GOOGLE_ASSOCIATION_ENABLED": "0"}):
    GENERATOR = module("room_layout_generator", PIPELINE / "process_card_slots.py")
VERIFY = module("room_layout_verify", PIPELINE / "opensms_verify_card_slot_fast.py")


class GoogleRoomLayoutTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="google-room-layout-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source.xlsx"
        self.snapshot = self.root / "snapshot.json"
        self.run_id = "offline-layout"

    def make_source(self, sheets=CURRENT_ROOM_SHEETS, *, hidden=(), extra=()):
        book = Workbook()
        book.remove(book.active)
        for position, name in enumerate((*sheets, *extra), 1):
            ws = book.create_sheet(name)
            ws.append(["卡槽", "電信商/價格", "激活日期", "預計過期日", "ICCID", "上卡日期", "備註"])
            ws.append([f"{100 + position}_1", "Chatr", "2026-09-01", "2026-09-30", None, None, ""])
            if name in hidden:
                ws.sheet_state = "hidden"
                ws["C2"] = "not-a-date"  # Hidden records must never contaminate dates/difference.
        book.save(self.source)
        book.close()

    def freeze_source(self):
        digest = hashlib.sha256(self.source.read_bytes()).hexdigest()
        return freeze(self.source, digest, self.snapshot, self.run_id)

    def consumers(self):
        digest = hashlib.sha256(self.snapshot.read_bytes()).hexdigest()
        with patch.dict(os.environ, {"OPENSMS_CHATR_GOOGLE_DATES_PATH": str(self.snapshot),
                                    "OPENSMS_CHATR_GOOGLE_DATES_SHA256": digest,
                                    "OPENSMS_RUN_ID": self.run_id}):
            generated = GENERATOR.load_chatr_google_dates()
        verified = VERIFY.load_chatr_google_snapshot(self.snapshot, digest, self.run_id)
        raw, proof = source_from_snapshot(str(self.snapshot), digest, self.run_id)
        return generated, verified, raw, proof

    def test_current_and_legacy_layouts_pass_all_consumers_and_difference_review(self):
        for sheets in (CURRENT_ROOM_SHEETS, LEGACY_ROOM_SHEETS):
            with self.subTest(sheets=sheets):
                self.make_source(sheets)
                self.snapshot.unlink(missing_ok=True)
                receipt = self.freeze_source()
                self.assertEqual(receipt["sourceSheets"], list(sheets))
                self.assertEqual(audit_chatr_snapshot_rows(self.source, json.loads(self.snapshot.read_text())), [])
                generated, verified, raw, proof = self.consumers()
                self.assertEqual(generated["record_count"], 3)
                self.assertEqual(verified["record_count"], 3)
                self.assertEqual(generated["index"]["103_1"]["activation_dt"], datetime(2026, 9, 1))
                self.assertEqual(verified["index"]["103_1"]["激活日期"], datetime(2026, 9, 1))
                records, counts, blanks = read_rows(raw)
                self.assertEqual(list(counts), list(sheets))
                self.assertEqual(records[-1]["sheet"], sheets[-1])
                book = Workbook()
                book.active.title = "处理汇总"
                book.create_sheet(PORT_SHEET).append(PORT_HEADERS)
                self.addCleanup(book.close)
                parse = GENERATOR.as_datetime
                difference = write_sheet(book, {"rows": records, "counts": counts, "blanks": blanks,
                                                "proof": proof}, set(), parse)
                checked = review(book, ["卡槽"], [], difference, str(self.snapshot),
                                 hashlib.sha256(self.snapshot.read_bytes()).hexdigest(), self.run_id, parse)
                self.assertTrue(checked["matches"], checked)
                self.assertEqual(difference["source_sheets"], list(sheets))

    def test_hidden_legacy_and_named_backups_are_not_substitutes(self):
        extras = ("四號機房", "四號合併前備份_20260923", "五號合併前備份_20260923")
        self.make_source(extra=extras, hidden=extras)
        receipt = self.freeze_source()
        self.assertEqual(receipt["recordCount"], 3)
        self.assertEqual(audit_chatr_snapshot_rows(self.source, json.loads(self.snapshot.read_text())), [])
        generated, verified, raw, proof = self.consumers()
        self.assertEqual(generated["record_count"], verified["record_count"])
        self.assertEqual(len(read_rows(raw)[0]), 3)
        self.assertEqual(proof["source_sheets"], list(CURRENT_ROOM_SHEETS))

    def test_missing_hidden_or_ambiguous_live_room_tabs_fail_without_publication(self):
        cases = (
            {"sheets": CURRENT_ROOM_SHEETS[:2], "extra": ("四號合併前備份_20260923",)},
            {"hidden": ("五號機房",)},
            {"sheets": ("三號機房", "五號機房")},
            {"extra": ("四號機房",)},
        )
        for options in cases:
            with self.subTest(options=options):
                self.make_source(**options)
                with self.assertRaises(ValueError):
                    self.freeze_source()
                self.assertFalse(self.snapshot.exists())
                with self.assertRaises(ValueError):
                    read_rows(self.source.read_bytes())

    def test_rehashed_snapshot_cannot_claim_old_or_backup_scope(self):
        self.make_source()
        self.freeze_source()
        original = json.loads(self.snapshot.read_text())
        for claimed in (list(LEGACY_ROOM_SHEETS), ["五號合併前備份_20260923"], None):
            with self.subTest(claimed=claimed):
                self.snapshot.write_text(json.dumps({**original, "sourceSheets": claimed}))
                digest = hashlib.sha256(self.snapshot.read_bytes()).hexdigest()
                self.assertTrue(audit_chatr_snapshot_rows(self.source, json.loads(self.snapshot.read_text())))
                with patch.dict(os.environ, {"OPENSMS_CHATR_GOOGLE_DATES_PATH": str(self.snapshot),
                                            "OPENSMS_CHATR_GOOGLE_DATES_SHA256": digest,
                                            "OPENSMS_RUN_ID": self.run_id}):
                    with self.assertRaisesRegex(RuntimeError, "来源范围"):
                        GENERATOR.load_chatr_google_dates()
                with self.assertRaisesRegex(RuntimeError, "来源范围"):
                    VERIFY.load_chatr_google_snapshot(self.snapshot, digest, self.run_id)
                with self.assertRaisesRegex(ValueError, "sourceSheets"):
                    source_from_snapshot(str(self.snapshot), digest, self.run_id)


if __name__ == "__main__":
    unittest.main()
