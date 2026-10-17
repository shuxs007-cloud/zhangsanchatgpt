"""Isolated synthetic document tests: no browser, network, production DB or inputs."""
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "services/opensms_document_api/pipeline"
DEPENDENCY = {"enabled": False, "status": "skipped", "reason": "user_disabled",
              "unused_port_device_scope": "system_source_only"}
FORBIDDEN = {"未匹配卡槽", "电信商_OPR预警", "所有谷歌有-系统没有", "谷歌有-系统无", "谷歌文档中有激活日期-系统无法使用"}


class NoGoogleDocumentTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="opensms-no-google-offline-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.raw, self.system, self.cache = (self.root / name for name in ("source.xlsx", "system.json", "cache.json"))
        self.output = self.root / "output" / "result.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "sheet1"
        headers = ["电话", "运营商", "opr", "卡槽", "创建时间", "状态", "激活日期", "截止日",
                   "10日发送窗口数", "10日回复窗口数", "10日回复率", "10日总发送量", "10日总回复量"]
        for day in ("09-03", "09-04", "09-05", "09-06"):
            headers += [day + suffix for suffix in ("发送窗口数", "回复窗口数", "回复率", "总发送量", "总回复量")]
        sheet.append(headers)
        self.records = [
            ("Chatr", "130_1", "注册失败", "2020-01-01", None, "2026-09-07"),
            ("Chatr", "130_2", "注册失败", "2026-09-06", None, None),
            ("Kooda", "130_3", "上架", "2026-09-01", "2026-09-01", "2026-12-15"),
            ("Koodo", "130_4", "注册失败", "2026-09-02", None, "2026-09-10"),
            ("Telus1", "130_5", "注册失败", "2020-01-01", None, None),
            ("Chatr", "130_6", "上架", "2026-09-04 00:00:00", "2026-09-01", "2026-09-12"),
            ("Chatr", "130_7", "下架", "2026-09-04 00:00:01", "2026-09-01", "2026-09-12"),
            ("Chatr", "130_2", "上架", "2026-08-28 23:59:59", "2026-09-01", "2026-09-12"),
            ("Chatr", "999_65", "注册失败", None, None, "2026-09-10"),
        ]
        system_rows = []
        for position, (opr, slot, status, created, activation, cutoff) in enumerate(self.records, start=1):
            phone = f"14385550{position:03d}"
            sheet.append([phone, "B1" if opr != "Chatr" else "A1", opr, slot, created, status, activation, cutoff,
                          "400", 40, 0.1, 400, 40] + [100, 10, 0.1, 100, 10] * 4)
            if opr != "Chatr":
                system_rows.append({"id": str(position), "slotIdentif": slot, "cardNo": phone,
                    "opr": opr, "cardStore": f"89122302000818367{position:02d}F",
                    "chatrActivateDate": "2026-09-01" if position == 3 else None,
                    "chatrPlanExpire": "2026-12-15" if position == 3 else None})
        sheet.auto_filter.ref = sheet.dimensions
        book.save(self.raw)
        book.close()
        system = {"status": "passed", "runId": "offline-no-google", "systemRows": system_rows,
            "totalRows": len(system_rows), "paginationTotalStable": True, "paginationBackendIdsUnique": True,
            "silentBrowserMode": True, "browserForegroundActivationAttempted": False}
        cache = {"schemaVersion": 1, "status": "passed", "runId": "offline-no-google",
            "createdAt": "2026-09-06T00:00:00+08:00", "sourceDatabase": str(self.root / "never-read.sqlite"),
            "sourceExists": True, "recordCount": 2, "records": [
                {"queryIccid": "8912230200081836703", "activationDate": "2026-09-05", "cutoffDate": "2026-10-04"},
                {"queryIccid": "8912230200081836704", "activationDate": "2026-09-02", "cutoffDate": "2026-10-01"},
            ]}
        self.system.write_text(json.dumps(system))
        self.cache.write_text(json.dumps(cache))
        self.env = {key: value for key, value in os.environ.items() if not key.startswith("OPENSMS_")}
        self.env.update({"PYTHON_BIN": sys.executable,
            "NODE_BIN": "/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node",
            "OPENSMS_GOOGLE_ASSOCIATION_ENABLED": "0", "OPENSMS_RUN_ID": "offline-no-google",
            "OPENSMS_SYSTEM_KOODA_PATH": str(self.system), "OPENSMS_SYSTEM_KOODA_SHA256": self.digest(self.system),
            "OPENSMS_REPORT_DATE_CACHE_PATH": str(self.cache), "OPENSMS_REPORT_DATE_CACHE_SHA256": self.digest(self.cache),
            "OPENSMS_TASK_F_DB_PATH": str(self.root / "never-read.sqlite"),
            "OPENSMS_OUTPUT_PATH": str(self.output), "OPENSMS_RECEIPT_DIR": str(self.root / "receipts"),
            "OPENSMS_ASSOC_PATH": str(self.root / "does-not-exist-google.xlsx"), "OPENSMS_ASSOC_SHA256": "not-a-hash",
        })

    def digest(self, path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def command(self, arguments, env=None):
        completed = subprocess.run(arguments, env=env or self.env, capture_output=True, text=True, timeout=45)
        self.assertEqual(completed.returncode, 0, (completed.stderr + completed.stdout)[-9000:])
        return json.loads(completed.stdout)

    def build(self):
        receipt = self.command(["bash", str(PIPELINE / "build_card_slot_report.sh"), str(self.raw), "-", str(self.output.parent), "2026-09-06"])
        self.assertEqual(receipt["result"], "成功")
        self.assertEqual(receipt["failed_checks"], [])
        return receipt

    def test_large_stage_receipts_use_files_not_exec_environment(self):
        import shutil
        pipeline = self.root / 'large-payload-pipeline'
        pipeline.mkdir()
        shutil.copyfile(PIPELINE / 'build_card_slot_report.sh', pipeline / 'build_card_slot_report.sh')
        # Run both real stages and their real independent workbook verification;
        # only add large unused receipt fields to reproduce exec size pressure.
        for filename, extras in (
            ('process_card_slots.py', "value['payload_size_regression'] = '数' * 400000\nvalue['enrichment']['payload_size_regression'] = '据' * 400000\n"),
            ('opensms_verify_card_slot_fast.py', "value['payload_size_regression'] = '验' * 400000\n"),
        ):
            (pipeline / filename).write_text(
                'import contextlib, io, json, runpy, sys\n'
                f'sys.path.insert(0, {str(PIPELINE)!r})\n'
                'buffer = io.StringIO()\n'
                'with contextlib.redirect_stdout(buffer):\n'
                '    try:\n'
                f'        runpy.run_path({str(PIPELINE / filename)!r}, run_name="__main__")\n'
                '    except SystemExit as error:\n'
                '        if error.code not in (0, None): raise\n'
                'value = json.loads(buffer.getvalue())\n' + extras +
                'print(json.dumps(value, ensure_ascii=False))\n')
        result = self.command(['bash', str(pipeline / 'build_card_slot_report.sh'),
                               str(self.raw), '-', str(self.output.parent), '2026-09-06'])
        self.assertEqual(result['result'], '成功')
        self.assertEqual(result['failed_checks'], [])
        self.assertEqual(len(result['process']['payload_size_regression']), 400000)
        self.assertEqual(len(result['enrichment']['payload_size_regression']), 400000)
        self.assertEqual(len(result['verification']['payload_size_regression']), 400000)
        payloads = list((self.root / 'receipts').glob('builder-*/*.json'))
        self.assertEqual(len(payloads), 3)
        self.assertTrue(all(path.stat().st_size > 1000000 for path in payloads))

    def verify(self, receipt):
        return self.command([sys.executable, str(PIPELINE / "opensms_verify_card_slot_fast.py"), str(self.output), str(self.raw)],
            {**self.env, "PROCESS_JSON": json.dumps(receipt["process"]), "ENRICH_JSON": json.dumps(receipt["enrichment"]),
             "EXPORT4DAY_SERVER_DATE": "2026-09-06"})

    def test_full_no_google_sources_dates_status_ports_and_types(self):
        before = [self.digest(path) for path in (self.raw, self.system, self.cache)]
        receipt = self.build()
        process, verify = receipt["process"], receipt["verification"]
        self.assertEqual(process["google_dependency"], DEPENDENCY)
        self.assertTrue(verify["google_dependency_review"]["matches"])
        self.assertTrue(verify["number_change_review"]["matches"])
        self.assertEqual(process["date_resolution"]["source"], "system_export_then_frozen_local")
        self.assertIsNone(process["date_resolution"]["association"])
        self.assertTrue(all(row["google"] == 0 for row in process["date_resolution"]["field_counts"].values()))
        self.assertEqual(process["date_resolution"]["kooda_local_cache_filled_rows"], 1)
        self.assertEqual(process["unused_port_devices"], ["130"])
        self.assertEqual(process["unused_port_count"], 60)
        self.assertEqual(process["registration_failed_port_count"], 3)
        self.assertEqual(process["device_daily_statistics"]["device_count"], 2)
        self.assertEqual(process["device_daily_statistics"]["date_count"], 4)
        self.assertEqual(process["device_daily_statistics"]["source_rows"], len(self.records))
        self.assertTrue(verify["device_daily_statistics_review"]["matches"])
        self.assertTrue(receipt["checks"]["设备每日统计独立重算"])
        self.assertEqual(process["system_missing_port_count"], 57)
        self.assertEqual(process["task_f_kooda_date_source"]["missing_dates_new_card_rows"], 0)
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        self.assertFalse(FORBIDDEN & set(book.sheetnames))
        sheet = book["sheet1"]
        index = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(index["换号状态"], 4)
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        self.assertEqual(sheet.cell(1, sheet.max_column).value, "ICCID")
        self.assertEqual(sheet.column_dimensions[sheet.cell(1, sheet.max_column).column_letter].width, 26)
        self.assertEqual(sheet.column_dimensions[sheet.cell(1, index["3日发送窗口数"]).column_letter].width, 16)
        self.assertEqual(sheet.column_dimensions[sheet.cell(1, index["4日回复率"]).column_letter].width, 12)
        self.assertNotIn("Google電信商/價格", index)
        self.assertNotIn("OPR匹配预警", index)
        self.assertIsNone(sheet.cell(2, index["换号状态"]).value)
        self.assertIsNone(sheet.cell(2, index["激活日期"]).value)
        self.assertEqual(sheet.cell(2, index["截止日"]).value, datetime(2026, 9, 7))
        self.assertEqual(sheet.cell(4, index["激活日期"]).value, datetime(2026, 9, 1))
        self.assertEqual(sheet.cell(4, index["截止日"]).value, datetime(2026, 12, 15))
        self.assertEqual(sheet.cell(5, index["激活日期"]).value, datetime(2026, 9, 2))
        self.assertEqual(sheet.cell(5, index["截止日"]).value, datetime(2026, 9, 10))
        self.assertIsNone(sheet.cell(6, index["换号状态"]).value)
        self.assertEqual(sheet.cell(6, index["激活日期"]).value, datetime(2020, 1, 1))
        self.assertEqual(sheet.cell(6, index["截止日"]).value, datetime(2020, 1, 30))
        self.assertFalse(any(cell.fill.fgColor.rgb == "00FFC7CE" for cell in sheet[6]))
        self.assertFalse(any(cell.fill.fgColor.rgb == "00FFC7CE" for cell in sheet[2]))
        self.assertIsNone(sheet.cell(7, index["换号状态"]).value)  # exactly 72h
        self.assertEqual(sheet.cell(8, index["换号状态"]).value, "已换号")  # 72h + 1s
        self.assertEqual(sheet.cell(9, index["换号状态"]).value, "已换号")  # absolute difference
        self.assertEqual(sheet.cell(2, index["10日发送窗口数"]).value, "400")
        self.assertEqual(sheet.cell(2, index["10日发送窗口数"]).data_type, "s")
        self.assertTrue(all(sheet.cell(row, index["开卡日期"]).value is None for row in range(2, sheet.max_row + 1)))
        self.assertEqual([self.digest(path) for path in (self.raw, self.system, self.cache)], before)
        self.assertFalse(Path(self.env["OPENSMS_TASK_F_DB_PATH"]).exists())
        self.assertFalse(Path(self.env["OPENSMS_ASSOC_PATH"]).exists())

    def test_numeric_xml_spelling_and_empty_text_roundtrip_preserve_excel_values(self):
        # The real exporter writes numeric XML values with a .0 suffix. Excel
        # has one numeric cell type; openpyxl serializes 40.0 as 40 on save.
        # Also reproduce the exporter's explicit empty text cells, which become
        # an empty inline-string cell when the report is saved.
        namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with ZipFile(self.raw) as source:
            entries = [(info, source.read(info.filename)) for info in source.infolist()]
        replacement = self.root / "source-numeric-spelling.xlsx"
        with ZipFile(replacement, "w") as target:
            for info, data in entries:
                if info.filename == "xl/worksheets/sheet1.xml":
                    tree = ElementTree.fromstring(data)
                    for coordinate, value in (("J2", "40.0"), ("N2", "100.0")):
                        cell = tree.find(f'.//x:c[@r="{coordinate}"]', namespace)
                        cell.find("x:v", namespace).text = value
                    cell = tree.find('.//x:c[@r="B2"]', namespace)
                    cell.clear()
                    cell.attrib.update(r="B2", t="inlineStr")
                    inline = ElementTree.SubElement(cell, "{" + namespace["x"] + "}is")
                    ElementTree.SubElement(inline, "{" + namespace["x"] + "}t").text = ""
                    data = ElementTree.tostring(tree, encoding="utf-8")
                target.writestr(info, data)
        replacement.replace(self.raw)
        book = load_workbook(self.raw, read_only=True)
        self.assertIs(type(book.active["J2"].value), float)
        self.assertEqual(book.active["B2"].value, "")
        book.close()
        receipt = self.build()
        self.assertTrue(receipt["verification"]["system_only_fields_review"]["matches"])
        book = load_workbook(self.output, read_only=True)
        sheet = book["sheet1"]
        index = {cell.value: cell.column for cell in next(sheet.iter_rows(max_row=1))}
        self.assertEqual(sheet.cell(2, index["10日回复窗口数"]).value, 40)
        self.assertEqual(sheet.cell(2, index["10日回复窗口数"]).data_type, "n")
        self.assertIsNone(sheet.cell(2, index["运营商"]).value)
        book.close()

    def test_independent_verifier_rejects_numeric_value_and_excel_type_changes(self):
        receipt = self.build()
        original = self.output.read_bytes()
        for value in ("40", True, 41, 40.00000001):
            with self.subTest(value=value):
                self.output.write_bytes(original)
                book = load_workbook(self.output)
                sheet = book["sheet1"]
                index = {cell.value: cell.column for cell in sheet[1]}
                sheet.cell(2, index["10日回复窗口数"]).value = value
                book.save(self.output)
                book.close()
                result = self.verify(receipt)
                self.assertFalse(result["system_only_fields_review"]["matches"])

    def test_independent_verifier_rejects_google_and_retained_field_tampering(self):
        receipt = self.build()
        book = load_workbook(self.output)
        sheet = book["sheet1"]
        index = {cell.value: cell.column for cell in sheet[1]}
        book.create_sheet("所有谷歌有-系统没有")
        sheet.cell(2, index["换号状态"]).value = "已换号"
        sheet.cell(2, index["截止日"]).value = "2030-01-01"
        sheet.cell(2, index["机房"]).value = "错误机房"
        for cell in sheet[2]:
            cell.fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
        book["系统无-谷歌表有"].cell(2, 4).value = "999_1"
        book.save(self.output)
        book.close()
        receipt["process"]["date_resolution"]["field_counts"]["激活日期"]["google"] = 1
        result = self.verify(receipt)
        for name in ("google_dependency_review", "number_change_review", "date_resolution_review",
                     "system_only_fields_review", "task_f_missing_activation_red_review", "unused_port_review"):
            self.assertFalse(result[name]["matches"], name)

    def test_process_cli_defaults_to_no_google_and_ignores_missing_association(self):
        env = {key: value for key, value in self.env.items() if key != "OPENSMS_GOOGLE_ASSOCIATION_ENABLED"}
        payload = self.command([sys.executable, str(ROOT / "scripts/opensms_process_document.py"),
            "--export", str(self.raw), "--association", str(self.root / "missing.xlsx"),
            "--server-date", "2026-09-06", "--output-dir", str(self.output.parent)], env)
        self.assertEqual(payload["inputs"]["google_dependency"], DEPENDENCY)
        self.assertIsNone(payload["inputs"]["association"])
        self.assertEqual(payload["api_attempts"], [])
        self.assertEqual(payload["local_result"]["process"]["google_dependency"], DEPENDENCY)

    def test_generate_cli_omits_association_and_publishes_system_only_report(self):
        payload = self.command([sys.executable, str(ROOT / "scripts/opensms_generate_from_downloads.py"),
            str(self.raw), "--backend-date", "2026-09-06", "--output-dir", str(self.output.parent),
            "--artifact-root", str(self.root / "bundle")])
        receipt = json.loads(Path(payload["receiptPath"]).read_text())
        self.assertIsNone(receipt["inputs"]["association"])
        self.assertEqual(receipt["inputs"]["google_dependency"], DEPENDENCY)
        report = Path(payload["reportPath"]).read_text()
        self.assertIn("Google关联表：已禁用", report)
        self.assertNotIn("| Google匹配行数 |", report)
        self.assertIn("Chatr激活日期允许空白", report)

    def test_three_argument_public_shell_entry_omits_association(self):
        result = self.command(["bash", str(ROOT / "scripts/opensms_build_card_slot_report.sh"),
            str(self.raw), str(self.output.parent), "2026-09-06"])
        self.assertEqual(result["local_result"]["process"]["google_dependency"], DEPENDENCY)

    def test_empty_frozen_cache_preserves_missing_activation_and_authoritative_cutoff(self):
        cache = json.loads(self.cache.read_text())
        cache.update(records=[], recordCount=0, sourceExists=False)
        self.cache.write_text(json.dumps(cache))
        self.env["OPENSMS_REPORT_DATE_CACHE_SHA256"] = self.digest(self.cache)
        receipt = self.build()
        self.assertEqual(receipt["process"]["date_resolution"]["kooda_local_cache_filled_rows"], 0)
        self.assertEqual(receipt["process"]["date_resolution"]["missing_activation_rows"], 1)
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        sheet = book["sheet1"]
        index = {cell.value: cell.column for cell in sheet[1]}
        self.assertIsNone(sheet.cell(5, index["激活日期"]).value)
        self.assertIsNone(sheet.cell(5, index["换号状态"]).value)
        self.assertEqual(sheet.cell(5, index["截止日"]).value, datetime(2026, 9, 10))
        self.assertTrue(all(cell.fill.fgColor.rgb == "00FFC7CE" for cell in sheet[5]))

    def test_no_google_ignores_inherited_legacy_enrichment_override(self):
        self.env["OPENSMS_LEGACY_DOCUMENT_ENRICHMENT"] = "1"
        receipt = self.build()
        self.assertTrue(receipt["stage_timings_seconds"]["legacy_second_full_workbook_pass_skipped"])
        self.assertTrue(receipt["verification"]["google_dependency_review"]["matches"])

    def test_explicit_legacy_low_level_mode_remains_compatible(self):
        google = self.root / "legacy-google.xlsx"
        book = Workbook()
        book.remove(book.active)
        for title in ("二號機房", "三號機房", "四號機房"):
            sheet = book.create_sheet(title)
            sheet.append(["卡槽", "電信商/價格", "上卡日期", "激活日期", "預計過期日", "備註", "iccid", "手機號"])
            if title == "四號機房":
                sheet.append(["130_1", "Chatr", "2026-09-01", "2026-09-01", "2026-09-20", None, None, None])
        book.save(google)
        book.close()
        receipt = self.command(["bash", str(PIPELINE / "build_card_slot_report.sh"), str(self.raw), str(google),
            str(self.output.parent), "2026-09-06"], {**self.env, "OPENSMS_GOOGLE_ASSOCIATION_ENABLED": "1"})
        self.assertEqual(receipt["result"], "成功")
        self.assertTrue(receipt["process"]["google_dependency"]["enabled"])
        self.assertEqual(receipt["process"]["date_resolution"]["field_counts"]["激活日期"]["google"], 1)


if __name__ == "__main__":
    unittest.main()
