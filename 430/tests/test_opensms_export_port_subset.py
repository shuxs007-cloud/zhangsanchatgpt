import importlib.util
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, Protection
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('subset', ROOT / 'scripts/opensms_export_port_subset.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class PortSubsetTests(unittest.TestCase):
    def test_user_requested_twenty_prefixes(self):
        expected = ('60902', '55972', '54782', '65440', '54422', '56580', '55802', '64106', '58828', '57156', '64416', '56924', '65370', '60710', '53564', '65476', '62050', '50360', '55072', '53322')
        self.assertEqual(module.PREFIXES, expected)
        self.assertEqual(len(set(module.PREFIXES)), 20)

    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.source, self.output, self.receipt = [self.root / name for name in ('source.xlsx', 'output.xlsx', 'receipt.json')]

    def source_book(self, slots):
        book = Workbook()
        sheet = book.active
        sheet.title = 'sheet1'
        sheet.append(['电话', '卡槽', 'ICCID'])
        for slot in slots:
            sheet.append(['01234567890', slot, '8912230200081836799F'])
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill('solid', fgColor='FFC7CE')
                cell.number_format = '@'
        book.create_sheet('处理汇总')
        book.save(self.source)
        book.close()

    def run_export(self):
        return module.export_subset(self.source, self.output, self.receipt, 'offline-run', module.digest(self.source))

    def test_all_prefixes_order_styles_types_and_no_middle_match(self):
        self.source_book(['x56924_1', None, '156924_1'] + [' '+p+'_1.01 ' for p in module.PREFIXES] + ['569240_2'])
        before = module.digest(self.source)
        result = self.run_export()
        self.assertEqual(result['rowCount'], 21)
        self.assertEqual(result['counts']['56924'], 2)
        self.assertEqual(before, module.digest(self.source))
        self.assertEqual(module.verify(self.source, self.output), 21)
        book = load_workbook(self.output)
        self.assertEqual(book.sheetnames, ['sheet1'])
        self.assertEqual(book['sheet1']['A2'].value, '01234567890')
        self.assertEqual(book['sheet1']['C2'].value, '8912230200081836799F')
        book['sheet1']['A2'] = 'tampered'
        book.save(self.output)
        book.close()
        with self.assertRaises(ValueError):
            module.verify(self.source, self.output)

    def test_no_match_header_only_and_no_overwrite(self):
        self.source_book(['other'])
        self.assertEqual(self.run_export()['rowCount'], 0)
        with self.assertRaises(FileExistsError):
            self.run_export()

    def test_sparse_unstyled_blank_cells_keep_default_styles(self):
        book = Workbook()
        sheet = book.active
        sheet.title = 'sheet1'
        sheet.append(['电话', '卡槽', '空列', 'ICCID'])
        sheet.append(['01234567890', '56924_1', None, '8912230200081836799F'])
        book.save(self.source)
        book.close()
        self.assertEqual(self.run_export()['rowCount'], 1)
        self.assertEqual(module.verify(self.source, self.output), 1)

    def test_wrong_source_hash_rejected_before_output(self):
        self.source_book(['56924_1'])
        with self.assertRaises(ValueError):
            module.export_subset(self.source, self.output, self.receipt, 'offline', '0'*64)
        self.assertFalse(self.output.exists())

    def test_cached_styles_preserve_components_and_detect_style_only_tamper(self):
        self.source_book(['56924_1', '56924_2', '56924_3'])
        book = load_workbook(self.source)
        sheet = book['sheet1']
        for row in (2, 3):
            cell = sheet.cell(row, 1)
            cell.font = Font(name='Arial', bold=True, color='123456')
            cell.alignment = Alignment(horizontal='right', wrap_text=True)
            cell.border = Border(left=Side(style='thin', color='AA0000'))
            cell.protection = Protection(locked=False, hidden=True)
            cell.comment = Comment('keep comment', 'tester')
        sheet.row_dimensions[2].height = 24
        sheet.column_dimensions['A'].width = 28
        sheet.freeze_panes = 'B2'
        book.save(self.source)
        book.close()
        result = self.run_export()
        self.assertTrue(result['performance']['independentSourceReopen'])
        book = load_workbook(self.output)
        sheet = book['sheet1']
        self.assertEqual(sheet['A3'].comment.text, 'keep comment')
        self.assertEqual(sheet['A3'].font.color.rgb, '00123456')
        self.assertEqual(sheet.row_dimensions[2].height, 24)
        self.assertEqual(sheet.column_dimensions['A'].width, 28)
        self.assertEqual(sheet.freeze_panes, 'B2')
        sheet['A3'].font = Font(name='Arial', bold=False)
        book.save(self.output)
        book.close()
        with self.assertRaisesRegex(ValueError, '样式核验失败'):
            module.verify(self.source, self.output)

    def test_style_ids_are_not_compared_between_workbooks(self):
        self.source_book(['56924_1'])
        self.run_export()
        original, result = load_workbook(self.source), load_workbook(self.output)
        try:
            # Exercise all cells against the uncached reference contract.
            for expected, actual in zip(module.selected_rows(original['sheet1']), result['sheet1'].iter_rows()):
                left, right = {}, {}
                for a, b in zip(expected, actual):
                    self.assertEqual(module.signature(a), module.signature(a, left))
                    self.assertEqual(module.signature(b), module.signature(b, right))
                    self.assertEqual(module.signature(a, left), module.signature(b, right))
        finally:
            original.close()
            result.close()

    def test_runner_no_longer_generates_port_subset_in_any_mode(self):
        runner = (ROOT / 'scripts/opensms_patrol_runner.py').read_text()
        body = runner.split('    def finalize(', 1)[1].split('    def summary(', 1)[0]
        self.assertNotIn('opensms_export_port_subset.py', body)
        self.assertNotIn('"port_subset"', body)
        self.assertIn('self.state["portSubset"] = None', body)
        self.assertNotIn('if self.task_a_only', body)

    def test_reports_do_not_offer_removed_file_even_with_old_result(self):
        spec = importlib.util.spec_from_file_location('chat_summary', ROOT / 'scripts/opensms_chat_summary.py')
        chat = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(chat)
        for status in ('success', 'failed'):
            result = {'status': status, 'portSubset': {
                'status': 'passed', 'rowCount': 10, 'output': '/old/subset.xlsx'}}
            rendered = chat.render_summary(result)
            self.assertNotIn('指定端口', rendered)
            self.assertNotIn('/old/subset.xlsx', rendered)
        report = (ROOT / 'scripts/opensms_finalize_patrol_report.py').read_text()
        self.assertNotIn("config['portSubset']", report)


if __name__ == '__main__':
    unittest.main()
