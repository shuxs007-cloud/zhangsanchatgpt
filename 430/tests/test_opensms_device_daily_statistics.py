"""Synthetic/offline only; never starts a patrol, network call or business write."""
from copy import copy
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from services.opensms_document_api.pipeline.device_daily_statistics import (
    SHEET_NAME, UNKNOWN_DEVICE, FIELDS, aggregate_device_daily, count_value,
    device_key, write_device_daily_sheet,
)
from services.opensms_document_api.pipeline.verify_device_daily_statistics import review_device_daily_statistics
from opensms_device_daily_report import run
from opensms_finalize_patrol_report import device_daily_report_lines


class DeviceDailyStatisticsTests(unittest.TestCase):
    def fixture(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "sheet1"
        headers = ["电话", "运营商", "卡槽", *(day + field for day in ("01-01", "12-31") for field in FIELDS)]
        sheet.append(headers)
        # Same phone in two rows must not cause an arbitrary row to disappear.
        rows = [
            ["001", "B1", "110_01", 100, 10, .99, 200, 15, 10, 0, .7, 20, 1],
            ["001", "B12", "110_2.01", 900, 180, .99, 1800, 270, 30, 4, .7, 60, 8],
            ["002", "A1", "9_65", 0, 0, .8, 0, 0, 4, 2, .2, 5, 3],
            ["003", "A1", "１１１＿０１", "1,000", "50", .8, "2,000", "60", 0, 0, .9, 0, 0],
        ]
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="336699")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="000000"))
        sheet.row_dimensions[1].height = 32
        sheet.column_dimensions["D"].width = 22
        sheet["D2"].fill = PatternFill("solid", fgColor="FFC7CE")
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        return book, headers, rows

    def test_weighted_rates_totals_date_order_and_no_phone_deduplication(self):
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        result = aggregate_device_daily(headers, rows, 4)
        self.assertEqual(result["headers"], ["设备号", "记录数", *headers[3:]])
        self.assertEqual(result["receipt"]["dates"], ["01-01", "12-31"])
        self.assertEqual(result["receipt"]["actual_workers"], 2)
        self.assertEqual(result["receipt"]["device_count"], 3)
        lookup = {row[0]: row for row in result["rows"]}
        self.assertEqual(lookup["110"], ["110", 2, 1000, 190, .19, 2000, 285, 40, 4, .1, 80, 9])
        self.assertIsNone(lookup["9"][4])
        self.assertEqual([row[0] for row in result["rows"]], ["9", "110", "111"])

    def test_single_thread_and_multiple_threads_are_identical(self):
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        reference = aggregate_device_daily(headers, rows, 1)
        for workers in (2, 4, 8):
            actual = aggregate_device_daily(headers, rows, workers)
            self.assertEqual(reference["headers"], actual["headers"])
            self.assertEqual(reference["rows"], actual["rows"])

    def test_device_aliases_and_unknown_records_are_not_dropped(self):
        for slot in ("110_01", "110_1.01", "110-01", " 0110 _ 01 ", "１１０＿０１"):
            self.assertEqual(device_key(slot), "110")
        self.assertEqual(device_key("110_65"), "110")
        for slot in (None, "", "110", "110_", "_01", "110_01_02"):
            self.assertEqual(device_key(slot), UNKNOWN_DEVICE)
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        rows[0][2] = None
        result = aggregate_device_daily(headers, rows, 1)
        self.assertEqual(result["rows"][-1][0], UNKNOWN_DEVICE)
        self.assertEqual(sum(row[1] for row in result["rows"]), len(rows))
        self.assertEqual(result["receipt"]["unrecognized_device_rows"], 1)

    def test_alphanumeric_device_prefix_is_preserved_not_merged_or_lost(self):
        self.assertEqual(device_key("r110_34.01"), "r110")
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        rows[0][2] = "r110_01"
        result = aggregate_device_daily(headers, rows, 2)
        self.assertEqual([row[0] for row in result["rows"]], ["9", "110", "111", "r110"])
        self.assertEqual(result["receipt"]["unrecognized_device_rows"], 0)
        write_device_daily_sheet(book, book.active, result)
        self.assertTrue(review_device_daily_statistics(book, headers, rows)["matches"])

    def test_missing_counts_do_not_publish_partial_sum_or_rate(self):
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        rows[0][3] = None
        rows[1][7] = "#VALUE!"
        result = aggregate_device_daily(headers, rows, 2)
        device = next(row for row in result["rows"] if row[0] == "110")
        self.assertEqual(device[2:7], [None, 190, None, 2000, None])
        self.assertEqual(result["receipt"]["missing_or_invalid_count_cells"], 2)
        write_device_daily_sheet(book, book.active, result)
        self.assertTrue(review_device_daily_statistics(book, headers, rows)["matches"])

    def test_counts_are_exact_and_nonnegative_integers_not_percentages(self):
        for value in (None, "", " ", True, "NaN", float("inf"), "-1", -1, "4%", 1.5):
            self.assertIsNone(count_value(value), repr(value))
        for value in (1000, 1000.0, "1,000", "1000.0", "1e3"):
            self.assertEqual(count_value(value), 1000)
        self.assertEqual(count_value("9007199254740993"), 9007199254740993)
        self.assertEqual(count_value("0"), 0)

    def test_invalid_schema_and_worker_limits_fail_closed(self):
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        for bad in (headers[:-1], [*headers, headers[-1]], [*headers[:2], "设备", *headers[3:]], []):
            with self.assertRaises(ValueError):
                aggregate_device_daily(bad, rows, 1)
        for workers in (0, -1, 9):
            with self.assertRaises(ValueError):
                aggregate_device_daily(headers, rows, workers)
        invalid_date = [name.replace("01-01", "13-01") for name in headers]
        with self.assertRaises(ValueError):
            aggregate_device_daily(invalid_date, rows, 1)

    def test_empty_input_has_headers_and_no_invented_devices(self):
        book, headers, _ = self.fixture()
        self.addCleanup(book.close)
        result = aggregate_device_daily(headers, [], 4)
        self.assertEqual(result["rows"], [])
        self.assertEqual(result["receipt"]["actual_workers"], 0)
        write_device_daily_sheet(book, book.active, result)
        self.assertTrue(review_device_daily_statistics(book, headers, [])["matches"])

    def test_style_matches_main_without_copying_business_row_colors(self):
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        main = book.active
        before = [(cell.value, cell.data_type, copy(cell._style)) for row in main for cell in row]
        result = aggregate_device_daily(headers, rows, 2)
        write_device_daily_sheet(book, main, result)
        sheet = book[SHEET_NAME]
        self.assertEqual(book.sheetnames[1], SHEET_NAME)
        self.assertEqual(sheet["C1"]._style, main["D1"]._style)
        self.assertEqual(sheet.column_dimensions["C"].width, 22)
        self.assertEqual(sheet.row_dimensions[1].height, 32)
        self.assertEqual(sheet["E2"].number_format, "0.00%")
        self.assertEqual(sheet["A2"].number_format, "@")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertIsNone(sheet["C2"].fill.patternType)
        self.assertEqual(sheet.freeze_panes, "C2")
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        self.assertEqual(before, [(cell.value, cell.data_type, cell._style) for row in main for cell in row])
        write_device_daily_sheet(book, main, result)
        self.assertEqual(book.sheetnames.count(SHEET_NAME), 1)

    def test_independent_readback_catches_corruption_and_missing_sheet(self):
        book, headers, rows = self.fixture()
        self.addCleanup(book.close)
        self.assertFalse(review_device_daily_statistics(book, headers, rows)["matches"])
        write_device_daily_sheet(book, book.active, aggregate_device_daily(headers, rows, 2))
        output = BytesIO()
        book.save(output)
        loaded = load_workbook(BytesIO(output.getvalue()))
        self.addCleanup(loaded.close)
        self.assertTrue(review_device_daily_statistics(loaded, headers, rows)["matches"])
        for cell, value in (("C3", 999), ("E3", .5), ("A3", "wrong"), ("C1", "wrong")):
            old = loaded[SHEET_NAME][cell].value
            loaded[SHEET_NAME][cell] = value
            self.assertFalse(review_device_daily_statistics(loaded, headers, rows)["matches"])
            loaded[SHEET_NAME][cell] = old

    def test_cli_is_offline_exclusive_and_preserves_all_original_sheets(self):
        with tempfile.TemporaryDirectory(prefix="opensms-device-offline-") as directory:
            source, output = (Path(directory) / name for name in ("source.xlsx", "output.xlsx"))
            book, _, _ = self.fixture()
            book.create_sheet("原始备注").append(["保持不变", 42])
            book.save(source)
            book.close()
            before = source.read_bytes()
            result = run(source, output, 4)
            self.assertEqual(result["status"], "passed")
            self.assertIn("not_fresh_export", result["mode"])
            self.assertTrue(all(result["original_sheet_values_preserved"].values()))
            self.assertEqual(source.read_bytes(), before)
            with self.assertRaises(ValueError):
                run(source, output, 4)
            with self.assertRaises(ValueError):
                run(source, source, 4)

    def test_cli_does_not_publish_failed_verification(self):
        with tempfile.TemporaryDirectory(prefix="opensms-device-failed-") as directory:
            source, output = (Path(directory) / name for name in ("source.xlsx", "output.xlsx"))
            book, _, _ = self.fixture()
            book.save(source)
            book.close()
            with patch("opensms_device_daily_report.review_device_daily_statistics", return_value={"matches": False}):
                with self.assertRaises(RuntimeError):
                    run(source, output)
            self.assertFalse(output.exists())
            self.assertEqual(list(Path(directory).glob(".device-daily-*")), [])

    def test_report_states_counts_and_real_verification_not_historical_success(self):
        self.assertEqual(device_daily_report_lines({}), [])
        proof = {"device_count": 3, "date_count": 2, "source_rows": 4, "actual_workers": 2,
                 "unrecognized_device_rows": 0, "missing_or_invalid_count_cells": 0}
        text = "\n".join(device_daily_report_lines({"process": {"device_daily_statistics": proof}}))
        self.assertIn("3台设备", text)
        self.assertIn("独立重算：未通过", text)
        text = "\n".join(device_daily_report_lines({"process": {"device_daily_statistics": proof},
              "verification": {"device_daily_statistics_review": {"matches": True}}}))
        self.assertIn("独立重算：通过", text)

    def test_real_builder_and_independent_verifier_synthetic_integration(self):
        # Invoke the real Python stages directly, not Bash here-docs (which are
        # incompatible with the offline runner's strict macOS write sandbox).
        import json
        from tests.test_opensms_document_no_google import NoGoogleDocumentTests, PIPELINE
        fixture = NoGoogleDocumentTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        env = {**fixture.env, "OPENSMS_MAIN_PATH": str(fixture.raw),
               "OPENSMS_OUTPUT_DIR": str(fixture.output.parent),
               "OPENSMS_REPORT_ROOT": str(fixture.root), "EXPORT4DAY_SERVER_DATE": "2026-09-06"}
        process = fixture.command([sys.executable, "-B", str(PIPELINE / "process_card_slots.py")], env)
        verify = fixture.command([sys.executable, "-B", str(PIPELINE / "opensms_verify_card_slot_fast.py"),
                                  str(fixture.output), str(fixture.raw)],
                                 {**env, "PROCESS_JSON": json.dumps(process), "ENRICH_JSON": json.dumps(process["enrichment"])})
        self.assertTrue(verify["device_daily_statistics_review"]["matches"])
        self.assertEqual(process["device_daily_statistics"]["source_rows"], 9)
        self.assertEqual(process["device_daily_statistics"]["device_count"], 2)
        self.assertEqual(process["device_daily_statistics"]["actual_workers"], 4)
        self.assertTrue(verify["system_only_fields_review"]["matches"])
        self.assertTrue(verify["number_change_review"]["matches"])
        self.assertTrue(verify["headers_match"])
        self.assertEqual(verify["formula_error_matches"], 0)
        self.assertEqual(verify["window_metric_mismatch_count"], 0)

    def test_carrier_refresh_preserves_device_sheet_xml_byte_for_byte(self):
        from tests.test_opensms_refresh_workbook_carriers_fast import MODULE, FastCarrierRefreshTests
        with tempfile.TemporaryDirectory(prefix="opensms-device-refresh-") as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            book, headers, rows = self.fixture()
            # Synthetic unique backend IDs: this test is about non-main sheet
            # preservation, not the refresh writer's duplicate-phone guard.
            book.active["A3"] = "004"
            write_device_daily_sheet(book, book.active, aggregate_device_daily(headers, rows, 2))
            book.save(source)
            book.close()
            helper = FastCarrierRefreshTests()
            args = helper.fresh_args(root, source, "device", [{"phone": "001", "liveCarrierName": "目标运营商"}])
            result = MODULE.run(args)
            self.assertEqual(result["changed"], 1)
            with ZipFile(source) as before, ZipFile(args.output) as after:
                original_path = dict(MODULE.workbook_sheets(before))[SHEET_NAME]
                output_path = dict(MODULE.workbook_sheets(after))[SHEET_NAME]
                self.assertEqual(before.read(original_path), after.read(output_path))


if __name__ == "__main__":
    unittest.main()
