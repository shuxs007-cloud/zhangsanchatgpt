import importlib.util
import json
from datetime import date, datetime, timedelta
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[1]


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


DB = load_module("opensms_task_f_db", ROOT / "scripts" / "opensms_task_f_db.py")
DOCUMENT = load_module(
    "opensms_process_card_slots_task_f",
    ROOT / "services" / "opensms_document_api" / "pipeline" / "process_card_slots.py",
)


class TaskFDatabaseTests(unittest.TestCase):
    def test_normalization_and_inclusive_30_day_cutoff(self):
        self.assertEqual(DB.normalize_iccid(" 8912230200081836799F "), "8912230200081836799")
        self.assertEqual(DB.cutoff_date_for_activation(date(2026, 8, 3)), date(2026, 9, 1))

    def test_terminal_query_is_not_scheduled_again_and_document_resolves_it(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            database = root / "task_f.sqlite3"
            cards = root / "cards.json"
            pending = root / "pending.json"
            lookup = root / "lookup.json"
            imported = root / "imported.json"
            cards.write_text(json.dumps({
                "rows": [{
                    "cardStore": "8912230200081836799F",
                    "phone": "14385550123",
                    "slot": "130_01",
                    "opr": "Telus1",
                    "carrierName": "B1",
                }]
            }), encoding="utf-8")
            first = DB.prepare_pending(database, cards, pending)
            self.assertEqual(first["pendingCount"], 1)
            lookup.write_text(json.dumps({
                "results": [{
                    **first["rows"][0],
                    "status": "found",
                    "purchaseDate": "2026-08-01",
                    "purchaseDateRaw": "2026-08-01",
                    "actRequestDate": "2026-08-03 09:00",
                    "activationDate": "2026-08-03",
                    "activationDateRaw": "2026-08-03 09:01",
                    "simNumber": "8912230200081836799",
                    "simNumberRaw": "8912230200081836799",
                    "sku": "TEST-SKU",
                    "detailSimIccid": "8912230200081836799",
                    "sourceOccurrences": [{"sourceData": {"cardStore": "8912230200081836799F"}}],
                    "webFieldPresence": {
                        "purchaseDate": True,
                        "actRequestDate": True,
                        "activationDate": True,
                        "simNumber": True,
                        "sku": True,
                    },
                    "workerId": 1,
                }]
            }), encoding="utf-8")
            receipt = DB.import_lookup(database, lookup, imported)
            self.assertEqual(receipt["foundImportedCount"], 1)
            second = DB.prepare_pending(database, cards, root / "pending_again.json")
            self.assertEqual(second["pendingCount"], 0)
            self.assertEqual(second["alreadyQueriedCount"], 1)

            indexes = DOCUMENT.load_task_f_kooda_dates(database)
            record, method = DOCUMENT.resolve_task_f_kooda_date(
                indexes,
                iccid="8912230200081836799F",
                phone="14385550123",
                slot="130_1.01",
            )
            self.assertIsNotNone(record)
            self.assertEqual(record["activation_date"], date(2026, 8, 3))
            self.assertEqual(record["cutoff_date"], date(2026, 9, 1))
            self.assertEqual(method, "ICCID去末尾单个F后精确匹配")

    def test_status_boundaries_use_platform_date(self):
        cutoff = date(2026, 9, 1)
        self.assertEqual(DOCUMENT.classify_expiry_status(cutoff, date(2026, 8, 26)), "6天")
        self.assertEqual(DOCUMENT.classify_expiry_status(cutoff, date(2026, 8, 30)), "快过期")
        self.assertEqual(DOCUMENT.classify_expiry_status(cutoff, date(2026, 8, 31)), "快过期")
        self.assertEqual(DOCUMENT.classify_expiry_status(cutoff, date(2026, 9, 1)), "已过期")
        self.assertEqual(DOCUMENT.classify_expiry_status(cutoff, date(2026, 9, 2)), "已过期")

    def test_missing_activation_red_has_priority_and_fresh_found_row_is_not_red(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["电话", "运营商", "激活日期"])
        sheet.append(["14385550123", "B1", None])
        for cell in sheet[2]:
            cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        DOCUMENT.fill_task_f_missing_activation_row(sheet, 2, 3)
        self.assertTrue(all(str(cell.fill.fgColor.rgb).upper().endswith("FFC7CE") for cell in sheet[2]))

        sheet.append(["14385550124", "B1", date(2026, 9, 3)])
        self.assertTrue(all(cell.fill.fill_type is None for cell in sheet[3]))


class TaskFNewCardNumberChangeTests(unittest.TestCase):
    def classify(self, *, opr="Kooda", activation=None, expiry=None, opened=None, created=None):
        return DOCUMENT.classify_number_change(
            kooda_row=DOCUMENT.is_kooda_opr(opr), activation_dt=activation,
            expiry_dt=expiry, open_dt=opened, created=created,
        )

    def test_kooda_aliases_with_both_dates_missing_are_new_cards(self):
        for opr in ("Kooda", "kooda", "KOODA", "Koodo", "Telus1", "telus1", "TELUS1"):
            for created in (None, datetime(2026, 9, 6)):
                with self.subTest(opr=opr, created=created):
                    self.assertEqual(self.classify(opr=opr, opened=datetime(2025, 1, 1),
                                                   created=created), ("未换号", "缺失日期按新卡"))

    def test_missing_kooda_dates_override_google_open_and_creation_difference(self):
        created = datetime(2026, 9, 6, 12)
        for opened in (None, created, created - timedelta(days=90), created + timedelta(days=90)):
            with self.subTest(opened=opened):
                self.assertEqual(self.classify(opened=opened, created=created),
                                 ("未换号", "缺失日期按新卡"))

    def test_normal_kooda_uses_activation_not_old_google_open_date(self):
        activation = datetime(2026, 9, 3)
        self.assertEqual(self.classify(activation=activation, expiry=activation + timedelta(days=29),
                                       opened=datetime(2025, 1, 1), created=activation),
                         ("", "激活日期"))
        self.assertEqual(self.classify(activation=activation, expiry=activation + timedelta(days=29),
                                       opened=datetime(2026, 9, 6), created=datetime(2026, 9, 8)),
                         ("已换号", "激活日期"))

    def test_chatr_retains_google_open_date_fallback(self):
        opened = datetime(2026, 8, 1)
        self.assertFalse(DOCUMENT.is_kooda_opr("Chatr"))
        self.assertEqual(self.classify(opr="Chatr", opened=opened, created=datetime(2026, 9, 6)),
                         ("已换号", "上卡日期"))
        self.assertEqual(self.classify(opr="Chatr", opened=opened, created=opened),
                         ("", "上卡日期"))

    def test_exact_three_day_boundary_and_absolute_second_difference(self):
        reference = datetime(2026, 9, 3, 12, 30)
        for sign in (-1, 1):
            for seconds, expected in ((3 * 86400, ""), (3 * 86400 + 1, "已换号")):
                for opr in ("Kooda", "Chatr"):
                    with self.subTest(sign=sign, seconds=seconds, opr=opr):
                        kwargs = ({"activation": reference, "expiry": reference + timedelta(days=29)}
                                  if opr == "Kooda" else {"opened": reference})
                        self.assertEqual(self.classify(opr=opr,
                            created=reference + timedelta(seconds=sign * seconds), **kwargs),
                            (expected, "激活日期" if opr == "Kooda" else "上卡日期"))

    def test_missing_created_or_reference_retains_blank_old_rule(self):
        reference = datetime(2026, 9, 3)
        self.assertEqual(self.classify(activation=reference, expiry=reference + timedelta(days=29)), ("", ""))
        self.assertEqual(self.classify(opr="Chatr", opened=reference), ("", ""))
        self.assertEqual(self.classify(opr="Chatr", created=reference), ("", ""))
        self.assertEqual(self.classify(opr="Chatr"), ("", ""))

    def test_only_one_missing_date_does_not_trigger_new_card_override(self):
        reference = datetime(2026, 8, 1)
        created = reference + timedelta(days=10)
        self.assertEqual(self.classify(activation=reference, created=created), ("已换号", "激活日期"))
        self.assertEqual(self.classify(expiry=reference + timedelta(days=29), opened=reference,
                                       created=created), ("已换号", "上卡日期"))

    def test_dates_filled_on_next_export_recalculate_instead_of_retaining_new_card_flag(self):
        created = datetime(2026, 9, 6)
        self.assertEqual(self.classify(opened=datetime(2025, 1, 1), created=created),
                         ("未换号", "缺失日期按新卡"))
        activation = datetime(2026, 8, 30)
        self.assertEqual(self.classify(activation=activation, expiry=activation + timedelta(days=29),
                                       opened=datetime(2025, 1, 1), created=created),
                         ("已换号", "激活日期"))
        self.assertEqual(self.classify(activation=created, expiry=created + timedelta(days=29),
                                       opened=datetime(2025, 1, 1), created=created),
                         ("", "激活日期"))


if __name__ == "__main__":
    unittest.main()
