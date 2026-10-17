"""Offline cross-implementation checks for the user's 2026-10-16 device list."""
import json
import os
import sys
import tempfile
from pathlib import Path
import subprocess
import unittest

import opensms_analyze_current as analysis
from services.opensms_document_api.pipeline import process_card_slots as process
from services.opensms_document_api.pipeline import opensms_verify_card_slot_fast as verify

ROOT = Path(__file__).resolve().parents[1]
ROOMS = {
    '二號機房': set(range(101, 122)),
    '三號機房': set(range(201, 242)),
    '五號機房': {*range(68, 72), *range(151, 174)},
    '六號機房': {*range(130, 150), *range(174, 193), *range(194, 196)},
    '七號機房': set(range(94, 101)),
}
LABELS = dict(zip(ROOMS, ['2号机房', '3号机房', '5号机房', '6号机房', '7号机房']))


def expected(device):
    return next((room for room, devices in ROOMS.items() if device in devices), '未配置')


class RoomMappingTests(unittest.TestCase):
    def test_explicit_list_counts_and_no_overlap(self):
        self.assertEqual([len(devices) for devices in ROOMS.values()], [21, 41, 27, 41, 7])
        self.assertEqual(len(set.union(*ROOMS.values())), 137)
        self.assertNotIn('四號機房', ROOMS)

    def test_all_devices_boundaries_and_unknowns(self):
        for device in [*range(301), 999, 56924]:
            with self.subTest(device=device):
                room = expected(device)
                self.assertEqual(process.room_for_slot(f'{device}_1.01'), room)
                self.assertEqual(verify.room_for_device(str(device)), room)
                self.assertEqual(analysis.room_key({'卡槽': f'{device}_1.01'}), LABELS.get(room, '未识别'))

    def test_reassigned_and_new_devices(self):
        moved = {*range(130, 150), *range(166, 196)}
        for device in moved:
            with self.subTest(device=device):
                self.assertEqual(process.room_for_slot(f'{device}_01'), expected(device))
                self.assertEqual(verify.room_for_device(str(device)), expected(device))
                self.assertEqual(analysis.room_key({'卡槽': f'{device}_01'}), LABELS.get(expected(device), '未识别'))

    def test_room6_replacement_leaves_193_unconfigured_and_adds_195(self):
        from services.opensms_document_api.pipeline import room_mapping as mapping
        self.assertEqual(mapping.CONFIG['effectiveDate'], '2026-10-16')
        self.assertEqual(mapping.DEVICE_ROOMS, {device: room for room, devices in ROOMS.items() for device in devices})
        for device, room in [(192, '六號機房'), (193, '未配置'), (194, '六號機房'), (195, '六號機房'), (196, '未配置')]:
            with self.subTest(device=device):
                self.assertEqual(mapping.room_for_device(device), room)
                self.assertEqual(process.room_for_slot(f'{device}_1'), room)
                self.assertEqual(verify.room_for_device(str(device)), room)
                self.assertEqual(analysis.room_key({'卡槽': f'{device}_1'}), LABELS.get(room, '未识别'))
        self.assertEqual(dict(mapping.summary_items())['六號機房设备配置'], '130–149、174–192、194–195（41台）')

    def test_room7_boundaries_and_configuration_summary(self):
        from services.opensms_document_api.pipeline import room_mapping as mapping
        for device in range(93, 102):
            with self.subTest(device=device):
                room = expected(device)
                self.assertEqual(mapping.room_for_device(device), room)
                self.assertEqual(process.room_for_slot(f'{device}_1'), room)
                self.assertEqual(verify.room_for_device(str(device)), room)
                self.assertEqual(analysis.room_key({'卡槽': f'{device}_1'}), LABELS.get(room, '未识别'))
        self.assertEqual(process.room_for_slot('0094_01.01'), '七號機房')
        self.assertEqual(process.room_for_slot('r94_1'), '未配置')
        self.assertEqual(dict(mapping.summary_items())['七號機房设备配置'], '94–100（7台）')
        self.assertEqual(dict(mapping.summary_items())['配置设备总数'], 137)

    def test_prefixes_preserved_and_zero_padding_normalized(self):
        for slot in ['r130_1', 'R134_2', '', None, 'abc_1']:
            self.assertEqual(process.room_for_slot(slot), '未配置')
            self.assertEqual(analysis.room_key({'卡槽': slot}), '未识别')
        self.assertEqual(process.room_for_slot('0130_01.01'), '六號機房')
        self.assertEqual(process.room_for_slot(' 0134-01.01 '), '六號機房')
        self.assertEqual(process.room_for_slot('0150_01.01'), '未配置')

    def test_unused_ports_follow_same_list_without_extra_devices(self):
        records = [{'slot': f'{device}_1', 'status': '注册失败', 'phone': '2000000000', 'carrier': 'B1'}
                   for device in [71, 94, 100, 101, 130, 134, 154, 193, 194, 195, 201, 999]]
        rows, devices = process.build_unused_port_rows(records, [])
        self.assertEqual(len(rows), len(records) * 64)
        self.assertEqual(set(devices), {'71', '94', '100', '101', '130', '134', '154', '193', '194', '195', '201', '999'})
        for row in rows:
            self.assertEqual(row[1], expected(int(row[0])))
            self.assertEqual(len(row), 8)
            self.assertEqual(row[6], '注册失败' if row[2] == 1 else '系统中无')
            self.assertEqual(row[7], '')  # No Google input in this fixture.

    def test_javascript_fallback_matches_every_device(self):
        source = (ROOT / 'services/opensms_document_api/pipeline/add_a11_low_reply_sheet.mjs').read_text()
        # Evaluate only these pure functions, never the browser/document entrypoint.
        functions = source[source.index('function normalizeSlot(value)'):source.index('function startOfDay(')]
        slots = [f'{device}_1.01' for device in range(301)] + ['r130_1', '0134_01.01', '']
        config = (ROOT / 'scripts/opensms_room_settings.json').read_text()
        script = ('const ROOM_CONFIG = ' + config + ';\nfunction asText(value) { return value == null ? "" : String(value).trim(); }\n'
                  + functions + '\nprocess.stdout.write(JSON.stringify(' + json.dumps(slots)
                  + '.map(roomForSlot)));')
        node = '/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
        result = subprocess.run([node, '-e', script], capture_output=True, text=True, check=True, timeout=20)
        self.assertEqual(json.loads(result.stdout), [expected(device) for device in range(301)]
                         + ['未配置', '六號機房', '未配置'])

    def test_independent_review_and_tampering(self):
        from openpyxl import Workbook
        from collections import Counter
        from services.opensms_document_api.pipeline import room_mapping as mapping
        book = Workbook()
        book.active.title = 'sheet1'
        book.active.append(['卡槽', '机房'])
        for device in [68, 93, 94, 100, 101, 121, 130, 149, 150, 151, 173, 174, 192, 193, 194, 195, 196, 201, 241]:
            book.active.append([f'{device}_1', expected(device)])
        rows = Counter(row[1] for row in book.active.iter_rows(min_row=2, values_only=True))
        counts = {**{name: rows[name] for name in ROOMS}, '未配置': rows['未配置']}
        summary = dict(mapping.summary_items())
        summary.update({(name + '行数' if name != '未配置' else '未配置机房行数'): count for name, count in counts.items()})
        receipt = {'room_configuration': {'config': mapping.CONFIG, 'sha256': mapping.CONFIG_SHA256}, 'room_counts': counts}
        proof = verify.review_room_configuration(book, receipt, summary)
        self.assertTrue(proof['matches'])
        self.assertEqual(proof['configured_devices'], 137)
        book.active['B4'] = '五號機房'
        failed = verify.review_room_configuration(book, receipt, summary)
        self.assertFalse(failed['matches'])
        self.assertIn('机房与最新配置不一致', failed['errors'][0])
        book.close()

    def test_standalone_audit_without_repository_pythonpath(self):
        from openpyxl import Workbook
        from services.opensms_document_api.pipeline import room_mapping as mapping
        with tempfile.TemporaryDirectory() as folder:
            book = Workbook()
            book.active.title = 'sheet1'
            book.active.append(['卡槽', '机房'])
            book.active.append(['174_1', '六號機房'])
            counts = {**{name: int(name == '六號機房') for name in ROOMS}, '未配置': 0}
            summary = dict(mapping.summary_items())
            summary.update({(name + '行数' if name != '未配置' else '未配置机房行数'): count for name, count in counts.items()})
            sheet = book.create_sheet('处理汇总')
            for row in summary.items():
                sheet.append(row)
            output = Path(folder) / 'result.xlsx'
            book.save(output)
            process_receipt = {'room_configuration': {'config': mapping.CONFIG, 'sha256': mapping.CONFIG_SHA256},
                               'room_counts': counts, 'output_path': str(output)}
            proof = verify.review_room_configuration(book, process_receipt, summary)
            receipt = Path(folder) / 'receipt.json'
            receipt.write_text(json.dumps({'result': '成功',
                'checks': {'最新机房配置及全部机房列独立回读': True}, 'process': process_receipt,
                'verification': {'room_configuration_review': proof}}))
            script = ('import sys; from pathlib import Path; sys.path.insert(0, sys.argv[1]); '
                      'from opensms_parallel_audit import check_card_receipt; '
                      'r=check_card_receipt(Path(sys.argv[2]), False); '
                      'assert "最新机房配置" not in r.message, r.message')
            environment = dict(os.environ)
            environment.pop('PYTHONPATH', None)
            subprocess.run([sys.executable, '-B', '-c', script, str(ROOT / 'scripts'), str(receipt)],
                           cwd=folder, env=environment, capture_output=True, text=True, check=True, timeout=20)
            book.close()


if __name__ == '__main__':
    unittest.main()
