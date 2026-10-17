"""Offline Chatr-only date fallback and independent source/provenance checks."""
from datetime import datetime
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "services/opensms_document_api/pipeline"
sys.path.insert(0, str(ROOT / "scripts"))
from opensms_chatr_google_dates import freeze


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


FIXTURE = load_module("chatr_document_fixture", ROOT / "tests/test_opensms_document_no_google.py")
VERIFY = load_module("chatr_document_verify", PIPELINE / "opensms_verify_card_slot_fast.py")
with mock.patch.dict(os.environ, {"OPENSMS_GOOGLE_ASSOCIATION_ENABLED": "0"}):
    GENERATOR = load_module("chatr_document_generator", PIPELINE / "process_card_slots.py")


class ChatrDocumentFallbackTests(unittest.TestCase):
    def setUp(self):
        self.fixture = FIXTURE.NoGoogleDocumentTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.source = self.fixture.root / "fresh-google.xlsx"
        self.snapshot = self.fixture.root / "chatr-dates.json"
        book = Workbook()
        book.remove(book.active)
        for name in ("二號機房", "三號機房", "四號機房"):
            sheet = book.create_sheet(name)
            sheet.append(["卡槽", "電信商/價格", "激活日期", "預計過期日", "ICCID", "上卡日期"])
        sheet = book["二號機房"]
        sheet.append(["130_01.01", "Chatr 25", "2026-09-03T00:00:01", "2026-11-01", None, "2020-01-01"])
        sheet.append(["130_2", "Chatr", "2026-09-03T00:00:01", None])
        sheet.append(["130_02.01", "Chatr", None, "2026-09-09", "8912230200081836702F"])
        for slot in ("130_3", "130_4", "130_5", "888_1"):
            sheet.append([slot, "Chatr", "2030-01-01", "2030-12-31"])
        sheet.append(["130_6", "Kooda", "2035-01-01", "2035-12-31"])
        book.save(self.source)
        book.close()
        freeze(self.source, self.fixture.digest(self.source), self.snapshot, self.fixture.env["OPENSMS_RUN_ID"])
        self.fixture.env.update(OPENSMS_CHATR_GOOGLE_DATES_PATH=str(self.snapshot),
                                OPENSMS_CHATR_GOOGLE_DATES_SHA256=self.fixture.digest(self.snapshot))

    def test_system_first_individual_blank_fallback_kooda_untouched_and_no_old_tabs(self):
        receipt = self.fixture.build()
        process, verification = receipt["process"], receipt["verification"]
        proof = process["google_chatr_date_fallback"]
        self.assertEqual(proof["field_counts"], {"激活日期": 2, "截止日": 1})
        self.assertEqual(proof["matched_rows"], 3)
        self.assertEqual(proof["filled_rows"], 2)
        self.assertEqual(proof["chatr_rows"], 6)
        self.assertEqual(proof["unmatched_rows"], 3)
        self.assertEqual(proof["missing_activation_rows"], 1)
        self.assertEqual(proof["missing_cutoff_rows"], 0)
        self.assertEqual(proof, verification["date_resolution_review"]["chatr_google_dates"])
        self.assertTrue(verification["google_chatr_date_fallback_review"]["matches"])
        self.assertEqual(process["google_dependency"], FIXTURE.DEPENDENCY)
        self.assertIsNone(process["date_resolution"]["association"])
        self.assertEqual(process["unused_port_devices"], ["130"])
        self.assertEqual(process["date_resolution"]["field_counts"]["激活日期"]["local_cache"], 1)
        book = load_workbook(self.fixture.output)
        self.addCleanup(book.close)
        self.assertFalse(FIXTURE.FORBIDDEN & set(book.sheetnames))
        sheet = book["sheet1"]
        index = {cell.value: cell.column for cell in sheet[1]}
        self.assertNotIn("Google電信商/價格", index)
        self.assertNotIn("OPR匹配预警", index)
        self.assertTrue(all(sheet.cell(row, index["开卡日期"]).value is None for row in range(2, sheet.max_row + 1)))
        self.assertEqual(sheet.cell(2, index["激活日期"]).value, datetime(2026, 9, 3, 0, 0, 1))
        self.assertEqual(sheet.cell(2, index["截止日"]).value, datetime(2026, 9, 7))
        self.assertEqual(sheet.cell(3, index["截止日"]).value, datetime(2026, 9, 9))
        self.assertIsNone(sheet.cell(3, index["换号状态"]).value)  # Less than 72 hours.
        self.assertEqual(sheet.cell(4, index["截止日"]).value, datetime(2026, 12, 15))
        self.assertEqual(sheet.cell(5, index["激活日期"]).value, datetime(2026, 9, 2))
        self.assertEqual(sheet.cell(6, index["激活日期"]).value, datetime(2020, 1, 1))
        self.assertEqual(sheet.cell(6, index["截止日"]).value, datetime(2020, 1, 30))
        self.assertEqual(sheet.cell(9, index["激活日期"]).value, datetime(2026, 9, 1))

    def test_dispatch_emits_explicit_snapshot_and_independent_proof(self):
        result = self.fixture.command([sys.executable, str(ROOT / "scripts/opensms_process_document.py"),
            "--export", str(self.fixture.raw), "--no-google", "--server-date", "2026-09-06",
            "--output-dir", str(self.fixture.output.parent)])
        self.assertEqual(result["inputs"]["chatr_google_dates"]["sha256"], self.fixture.digest(self.snapshot))
        self.assertTrue(result["google_chatr_date_fallback"]["enabled"])
        self.assertEqual(result["api_attempts"], [])

    def test_nonempty_invalid_system_date_fails_instead_of_fallback(self):
        book = load_workbook(self.fixture.raw)
        sheet = book.active
        columns = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, columns["激活日期"]).value = "invalid date"
        book.save(self.fixture.raw)
        book.close()
        run = subprocess.run(["bash", str(PIPELINE / "build_card_slot_report.sh"), str(self.fixture.raw), "-",
                              str(self.fixture.output.parent), "2026-09-06"], env=self.fixture.env, capture_output=True, text=True)
        self.assertNotEqual(run.returncode, 0)
        self.assertIn("Chatr系统导出日期非空但无效", run.stderr)

    def test_independent_verifier_rejects_forged_dates_counts_and_tabs(self):
        receipt = self.fixture.build()
        receipt["process"]["google_chatr_date_fallback"]["field_counts"]["激活日期"] = 999
        result = self.fixture.verify(receipt)
        self.assertFalse(result["google_chatr_date_fallback_review"]["matches"])
        book = load_workbook(self.fixture.output)
        sheet = book["sheet1"]
        columns = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, columns["激活日期"]).value = datetime(2026, 9, 3)  # Same date but changed time.
        book.create_sheet("所有谷歌有-系统没有")
        book.save(self.fixture.output)
        book.close()
        result = self.fixture.verify(receipt)
        self.assertFalse(result["date_resolution_review"]["matches"])
        self.assertFalse(result["google_dependency_review"]["matches"])

    def test_snapshot_provenance_and_original_hash_are_independently_checked(self):
        payload = json.loads(self.snapshot.read_text())
        payload["records"][0]["activationDate"] = "2026-09-05"
        self.snapshot.write_text(json.dumps(payload))
        digest = self.fixture.digest(self.snapshot)
        with self.assertRaisesRegex(RuntimeError, "独立读取原始Google日期"):
            VERIFY.load_chatr_google_snapshot(self.snapshot, digest, self.fixture.env["OPENSMS_RUN_ID"])
        self.source.write_bytes(self.source.read_bytes() + b"hash-change")
        self.fixture.env["OPENSMS_CHATR_GOOGLE_DATES_SHA256"] = digest
        with mock.patch.dict(os.environ, self.fixture.env), self.assertRaisesRegex(RuntimeError, "SHA256"):
            GENERATOR.load_chatr_google_dates()
        with self.assertRaisesRegex(RuntimeError, "SHA256"):
            VERIFY.load_chatr_google_snapshot(self.snapshot, digest, self.fixture.env["OPENSMS_RUN_ID"])

    def test_independent_verifier_rejects_forged_chatr_row_and_missing_counts(self):
        receipt = self.fixture.build()
        for field in ("chatr_rows", "unmatched_rows", "missing_activation_rows", "missing_cutoff_rows"):
            with self.subTest(field=field):
                forged = json.loads(json.dumps(receipt))
                for proof in (forged["process"]["google_chatr_date_fallback"],
                              forged["process"]["date_resolution"]["chatr_google_dates"]):
                    proof[field] += 1
                result = self.fixture.verify(forged)
                self.assertFalse(result["date_resolution_review"]["matches"])
                self.assertFalse(result["google_chatr_date_fallback_review"]["matches"])
                self.assertIn("chatr_google_dates", result["date_resolution_review"]["receipt_mismatch_fields"])
    def test_partial_snapshot_configuration_is_rejected(self):
        env = {**self.fixture.env, "OPENSMS_CHATR_GOOGLE_DATES_SHA256": ""}
        with mock.patch.dict(os.environ, env), self.assertRaisesRegex(RuntimeError, "SHA256"):
            GENERATOR.load_chatr_google_dates()

    def set_conflicting_source_iccid(self):
        book = load_workbook(self.fixture.raw)
        sheet = book.active
        columns = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(1, sheet.max_column + 1).value = "ICCID"
        sheet.cell(3, sheet.max_column).value = "8912230200081836799F"
        book.save(self.fixture.raw)
        book.close()
        return columns

    def test_conflict_system_wins_is_reported_and_creation_fallback_is_not_google(self):
        self.set_conflicting_source_iccid()
        receipt = self.fixture.build()
        proof = receipt["process"]["google_chatr_date_fallback"]
        self.assertEqual(proof["iccid_conflict_count"], 1)
        conflict = proof["iccid_conflicts"][0]
        self.assertEqual(conflict["system_iccid"], "8912230200081836799F")
        self.assertEqual(conflict["google_iccid"], "8912230200081836702")
        self.assertEqual(conflict["activation_source"], "creation_date_fallback")
        self.assertEqual(conflict["final_activation"], "2026-09-06T00:00:00")
        self.assertEqual(conflict["final_cutoff"], "2026-10-05T00:00:00")
        self.assertTrue(receipt["verification"]["google_chatr_date_fallback_review"]["conflict_sheet_matches"])
        book = load_workbook(self.fixture.output)
        self.addCleanup(book.close)
        index = {cell.value: cell.column for cell in book["sheet1"][1]}
        self.assertEqual(book["sheet1"].cell(3, index["ICCID"]).value, conflict["system_iccid"])
        sheet = book[proof["iccid_conflict_sheet"]]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 4).data_type, "s")
        sheet.cell(2, 4).value = conflict["google_iccid"]
        book.save(self.fixture.output)
        result = self.fixture.verify(receipt)
        self.assertFalse(result["google_chatr_date_fallback_review"]["matches"])
        self.assertFalse(result["google_chatr_date_fallback_review"]["conflict_sheet_matches"])

    def test_conflict_preserves_existing_system_dates_and_missing_single_date(self):
        columns = self.set_conflicting_source_iccid()
        book = load_workbook(self.fixture.raw)
        sheet = book.active
        sheet.cell(3, columns["激活日期"]).value = "2026-09-04T03:04:05"
        book.save(self.fixture.raw)
        book.close()
        receipt = self.fixture.build()
        conflict = receipt["process"]["google_chatr_date_fallback"]["iccid_conflicts"][0]
        self.assertEqual(conflict["final_activation"], "2026-09-04T03:04:05")
        self.assertEqual(conflict["final_cutoff"], "")
        self.assertEqual(conflict["activation_source"], "system_export")
        self.assertEqual(conflict["cutoff_source"], "missing")
        forged = json.loads(json.dumps(receipt))
        forged["process"]["google_chatr_date_fallback"]["iccid_conflict_count"] = 0
        self.assertFalse(self.fixture.verify(forged)["google_chatr_date_fallback_review"]["matches"])

    def test_normalized_equal_iccid_allows_google_and_other_opr_cannot_fallback(self):
        columns = self.set_conflicting_source_iccid()
        book = load_workbook(self.fixture.raw)
        sheet = book.active
        sheet.cell(3, sheet.max_column).value = "8912230200081836702f"
        book.save(self.fixture.raw)
        book.close()
        receipt = self.fixture.build()
        self.assertEqual(receipt["process"]["google_chatr_date_fallback"]["iccid_conflict_count"], 0)
        self.assertEqual(receipt["process"]["google_chatr_date_fallback"]["field_counts"], {"激活日期": 2, "截止日": 1})
        self.fixture.output.unlink()
        book = load_workbook(self.fixture.raw)
        sheet = book.active
        sheet.cell(3, columns["opr"]).value = "Lucky"
        book.save(self.fixture.raw)
        book.close()
        receipt = self.fixture.build()
        self.assertEqual(receipt["process"]["google_chatr_date_fallback"]["field_counts"], {"激活日期": 1, "截止日": 0})
        book = load_workbook(self.fixture.output)
        self.addCleanup(book.close)
        index = {cell.value: cell.column for cell in book["sheet1"][1]}
        self.assertEqual(book["sheet1"].cell(3, index["激活日期"]).value, datetime(2026, 9, 6))
        self.assertEqual(book["sheet1"].cell(3, index["截止日"]).value, datetime(2026, 10, 5))


if __name__ == "__main__":
    unittest.main()
