"""Synthetic-only template authoring/readback; never reads the production F DB."""
from datetime import datetime
from copy import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import opensms_export_task_f_deadlines as exporter
import opensms_system_missing_f_dates as difference

KEY = '08912230200278137092'
RECORD = {'lookup_status': 'found', 'activation_date': '2026-09-01',
          'cutoff_date': '2026-09-30', 'queried_at': '2026-10-02T00:00:00Z'}

class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='task-f-template-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_three_columns_preserve_template_style_and_leading_zero(self):
        output = self.root / 'example.xlsx'
        rows = [[KEY, '2026-09-30', '2026-09-01'], ['0' * 22, None, None]]
        before = exporter.DEFAULT_TEMPLATE.read_bytes()
        result = exporter.write_template_workbook(exporter.DEFAULT_TEMPLATE, rows, output)
        exporter.verify_output(output, rows)
        self.assertEqual(exporter.DEFAULT_TEMPLATE.read_bytes(), before)
        self.assertEqual(result['templateSha256'], exporter.DEFAULT_TEMPLATE_SHA256)
        original = load_workbook(exporter.DEFAULT_TEMPLATE)
        actual = load_workbook(output)
        try:
            self.assertEqual(actual.sheetnames, ['截止日'])
            for col in ('A', 'B', 'C'):
                self.assertEqual(copy(actual['截止日'][col+'1'].font), copy(original['截止日'][col+'1'].font))
                self.assertEqual(copy(actual['截止日'][col+'1'].fill), copy(original['截止日'][col+'1'].fill))
                self.assertAlmostEqual(actual['截止日'].column_dimensions[col].width,
                                       original['截止日'].column_dimensions[col].width, places=4)
            self.assertEqual(actual['截止日']['A2'].value, KEY)
            self.assertIsInstance(actual['截止日']['B2'].value, datetime)
        finally:
            original.close(); actual.close()

    def test_empty_template_has_no_example_rows(self):
        output = self.root / 'empty.xlsx'
        result = exporter.write_template_workbook(exporter.DEFAULT_TEMPLATE, [], output)
        exporter.verify_output(output, [])
        self.assertEqual(result['recordCount'], 0)

    def test_invalid_or_duplicate_rows_stop_before_output(self):
        cases = [[[KEY+'F', '2026-09-30', '2026-09-01']],
                 [[KEY, '2026-10-01', '2026-09-01']],
                 [[KEY, None, '2026-09-01']],
                 [[KEY, None, None], [KEY, None, None]]]
        for rows in cases:
            output = self.root / 'invalid.xlsx'
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                exporter.write_template_workbook(exporter.DEFAULT_TEMPLATE, rows, output)
            self.assertFalse(output.exists())

    def test_pinned_template_rejects_byte_change(self):
        fake = self.root / 'template.xlsx'
        fake.write_bytes(exporter.DEFAULT_TEMPLATE.read_bytes() + b'changed')
        with patch.object(exporter, 'DEFAULT_TEMPLATE', fake), self.assertRaisesRegex(ValueError, 'SHA256'):
            exporter.check_template(fake)

    def test_daily_difference_keeps_all_source_rows_but_imports_unique_iccids(self):
        source = self.root / 'system.xlsx'
        wb = Workbook(); sheet = wb.active; sheet.title = 'sheet1'
        sheet.append(['电话', 'ICCID', '激活日期', '截止日', '卡槽'])
        sheet.append(['14370000000', KEY+'F', None, None, '001_01'])
        sheet.append(['14370000001', KEY, None, None, '002_01'])
        wb.save(source); wb.close()
        receipts = self.root / 'receipts'; receipts.mkdir()
        with patch.object(difference, 'read_database', return_value={KEY: RECORD}) as read:
            result = difference.export_difference(source, difference.digest(source), 'test-run',
                                                  self.root, receipts)
        read.assert_called_once()
        self.assertEqual((result['rowCount'], result['uniqueIccidCount'], result['exportedRowCount']), (2,1,1))
        self.assertEqual(len(json.loads(Path(result['sourceDetail']).read_text())['rows']), 2)
        exporter.verify_output(Path(result['output']), [[KEY, '2026-09-30', '2026-09-01']])
        with patch.object(difference, 'read_database', side_effect=AssertionError('No new DB read')):
            self.assertEqual(difference.export_difference(source, difference.digest(source), 'test-run',
                                                         self.root, receipts), result)

    def test_conflicting_existing_system_cutoff_not_overwritten(self):
        with self.assertRaisesRegex(ValueError, '系统已有截止日'):
            difference.select_rows(['ICCID', '激活日期', '截止日'],
                                   [[KEY, None, '2026-10-01']], {KEY: RECORD})

    def test_existing_output_and_metadata_override_do_not_start_writer(self):
        output = self.root / 'keep.xlsx'; output.write_bytes(b'keep')
        with patch.object(exporter.subprocess, 'run') as writer, self.assertRaises(FileExistsError):
            exporter.write_template_workbook(exporter.DEFAULT_TEMPLATE, [], output)
        writer.assert_not_called(); self.assertEqual(output.read_bytes(), b'keep')
        with patch.object(exporter, 'read_rows') as read, self.assertRaises(ValueError):
            exporter.export_deadlines(self.root/'fake.sqlite', exporter.DEFAULT_TEMPLATE,
                self.root/'new.xlsx', self.root/'receipt.json', extra_receipt={'publicationMode':'fake'})
        read.assert_not_called()

if __name__ == '__main__':
    unittest.main()
