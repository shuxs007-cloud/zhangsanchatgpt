"""Offline report-scoped Task F lookup and monotonic date import tests."""
from datetime import datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("task_f_report_backfill_db", ROOT / "scripts/opensms_task_f_db.py")
DB = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(DB)


def iccid(number):
    return str(8912230200000000000 + number)


class ReportBackfillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="task-f-report-backfill-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.database = self.root / "test.sqlite3"
        self.cards = self.root / "inventory.json"
        self.document = self.root / "report.xlsx"
        self.pending = self.root / "pending.json"
        self.run_id = "current-test-round"

    def save_inputs(self, rows, inventory_changes=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet1"
        sheet.append(["电话", "opr", "卡槽", "ICCID", "激活日期"])
        source_rows = []
        for i, (opr, raw_iccid, activation) in enumerate(rows, 1):
            phone, slot = str(14385550000 + i), f"130_{i}"
            sheet.append([phone, opr, slot, raw_iccid, activation])
            if DB.re.search(r"kooda|koodo|telus1", opr, DB.re.I):
                source_rows.append({"id": f"id-{i}", "opr": opr, "slotIdentif": slot + ".01",
                                    "cardNo": phone[1:], "cardStore": raw_iccid, "auditField": "retained"})
        workbook.save(self.document)
        inventory = {"status": "passed", "runId": self.run_id, "totalRows": len(source_rows),
                     "paginationTotalStable": True, "paginationBackendIdsUnique": True,
                     "silentBrowserMode": True, "browserForegroundActivationAttempted": False,
                     "systemRows": source_rows, **(inventory_changes or {})}
        DB.atomic_json(self.cards, inventory)
        return hashlib.sha256(self.document.read_bytes()).hexdigest()

    def prepare(self, digest=None, **kwargs):
        return DB.prepare_pending(self.database, self.cards, self.pending, document=self.document,
                                  document_sha256=digest or DB.file_sha256(self.document),
                                  run_id=self.run_id, **kwargs)

    def seed(self, number, status, activation=None, cutoff=None, schema=2, data_source="test"):
        with DB.open_database(self.database) as connection:
            connection.execute("""INSERT INTO kooda_activation(
                query_iccid,source_iccid,lookup_status,activation_date,cutoff_date,
                queried_at,record_schema_version,receipt_json,data_source)
                VALUES (?,?,?,?,?,?,?,?,?)""",
                (iccid(number), iccid(number) + "F", status, activation, cutoff,
                 "2026-09-01T00:00:00+08:00", schema, "{}", data_source))
            connection.commit()

    def fallback_inputs(self):
        self.save_inputs([("Kooda", iccid(1) + "F", "2026-09-01"),
                          ("Kooda", iccid(2) + "f", "2026-09-01"),
                          ("Chatr", iccid(3), "2026-09-01")])
        book = load_workbook(self.document)
        sheet = book["sheet1"]
        sheet.cell(1, 6, "截止日")
        for number in range(2, 5):
            sheet.cell(number, 6, "2026-09-30")
        book.save(self.document)
        book.close()
        proof = {"result": "成功", "checks": {"independent": True}, "failed_checks": [],
                 "verification": {"system_kooda_source_review": {"matches": True,
                    "run_id": self.run_id, "snapshot_sha256": DB.file_sha256(self.cards)},
                    "date_resolution_review": {"matches": True, "creation_date_fallback_rows": [
                        {"row": 2, "phone": "14385550001", "activation": "2026-09-01", "cutoff": "2026-09-30"},
                        {"row": 4, "phone": "14385550003", "activation": "2026-09-01", "cutoff": "2026-09-30"}]}}}
        path = self.root / "proof.json"
        DB.atomic_json(path, proof)
        return path, proof

    def test_verified_creation_estimate_is_queried_but_identical_real_dates_and_chatr_are_not(self):
        path, _ = self.fallback_inputs()
        result = self.prepare(document_proof=path, document_proof_sha256=DB.file_sha256(path))
        self.assertEqual([iccid(1)], [row["queryIccid"] for row in result["rows"]])
        self.assertEqual(0, result["documentMissingDateRows"])
        self.assertEqual(1, result["documentCreationFallbackRows"])
        self.assertEqual(1, result["documentLookupNeededRows"])
        self.assertEqual("creation_date_fallback", result["rows"][0]["sourceOccurrences"][0]["dateSource"])

    def test_creation_fallback_reuses_valid_new_local_record_without_lookup(self):
        path, _ = self.fallback_inputs()
        self.seed(1, "found", "2026-09-04", "2026-10-03")
        result = self.prepare(document_proof=path, document_proof_sha256=DB.file_sha256(path))
        self.assertEqual(0, result["pendingCount"])
        self.assertEqual(1, result["readyForUploadCount"])

    def test_bad_fallback_provenance_or_row_identity_fails_before_database_write(self):
        path, proof = self.fallback_inputs()
        with self.assertRaisesRegex(RuntimeError, "SHA256"):
            self.prepare(document_proof=path, document_proof_sha256="0" * 64)
        for field, bad_value in (("phone", "14385559999"), ("activation", "2026-08-31")):
            with self.subTest(field=field):
                changed = json.loads(json.dumps(proof))
                changed["verification"]["date_resolution_review"]["creation_date_fallback_rows"][0][field] = bad_value
                DB.atomic_json(path, changed)
                with self.assertRaises(RuntimeError):
                    self.prepare(document_proof=path, document_proof_sha256=DB.file_sha256(path))
        self.assertFalse(self.database.exists())

    def test_creation_proof_from_other_run_is_rejected(self):
        path, proof = self.fallback_inputs()
        proof["verification"]["system_kooda_source_review"]["run_id"] = "old-run"
        DB.atomic_json(path, proof)
        with self.assertRaisesRegex(RuntimeError, "本轮独立核验"):
            self.prepare(document_proof=path, document_proof_sha256=DB.file_sha256(path))
        self.assertFalse(self.database.exists())

    def lookup(self, number, status="found", activation="2026-09-07", report_mode=False, **changes):
        row = {"queryIccid": iccid(number), "sourceIccid": iccid(number) + "F",
               "status": status, "activationDate": activation, "activationDateRaw": activation,
               "purchaseDate": "2026-09-01", "purchaseDateRaw": "2026-09-01",
               "actRequestDate": "2026-09-07 08:00", "simNumber": iccid(number),
               "simNumberRaw": iccid(number), "detailSimIccid": iccid(number), "sku": "SKU-preserved",
               "webFieldPresence": {name: True for name in
                   ("purchaseDate", "actRequestDate", "activationDate", "simNumber", "sku")},
               "sourceOccurrences": [{"cardStore": iccid(number) + "F",
                                       "sourceData": {"cardStore": iccid(number) + "F", "auditField": "retained"}}],
               **changes}
        path = self.root / "lookup.json"
        DB.atomic_json(path, {"results": [row], "reportBackfillMode": report_mode})
        return DB.import_lookup(self.database, path, self.root / "import.json")

    def read_activation(self, number):
        with sqlite3.connect(self.database) as connection:
            connection.row_factory = sqlite3.Row
            row = connection.execute("SELECT * FROM kooda_activation WHERE query_iccid=?", (iccid(number),)).fetchone()
        return dict(row) if row else None

    def test_proven_archived_off_shelf_gap_is_excluded_without_querying_new_occupant(self):
        self.save_inputs([('Kooda', iccid(1) + 'F', None)])
        book = load_workbook(self.document)
        sheet = book['sheet1']
        sheet.cell(1, 6, '状态'); sheet.cell(2, 6, '下架')
        sheet.cell(1, 7, '创建时间'); sheet.cell(2, 7, '2026-09-01 00:00:00')
        book.save(self.document); book.close()
        inventory = json.loads(self.cards.read_text())
        inventory['systemRows'][0].update(slotIdentif=None, status='换卡已下线', showAble=False,
                                          machineName='130', createTime='2026-09-01 00:00:00')
        self.cards.write_text(json.dumps(inventory))
        result = self.prepare()
        self.assertEqual(result['pendingCount'], 0)
        self.assertEqual(result['documentExcludedCount'], 1)
        self.assertEqual(result['documentExcludedRows'][0]['reason'], 'archived_off_shelf_no_current_slot')

    def test_only_document_missing_activations_are_selected_and_local_found_reused(self):
        digest = self.save_inputs([
            ("Kooda", iccid(1) + "F", datetime(2026, 9, 7)),
            ("Telus1", iccid(2) + "f", None), ("Koodo", iccid(3) + "F", None),
            ("KOODA", iccid(4), None), ("Kooda", iccid(5) + "F", None),
            ("Kooda", None, None), ("Kooda", int(iccid(7)), None), ("Chatr", iccid(8), None)])
        self.seed(2, "found", "2026-09-01", "2026-09-30", schema=1)
        self.seed(3, "activation_missing", data_source="grabandgo_complete_catalog_attachment")
        self.seed(4, "not_found")
        before = self.document.read_bytes()
        result = self.prepare(digest)
        self.assertTrue(result["reportBackfillMode"])
        self.assertEqual(result["documentKoodaRows"], 7)
        self.assertEqual(result["documentActivationPresentRows"], 1)
        self.assertEqual(result["documentMissingActivationRows"], 6)
        self.assertEqual(result["documentExcludedCount"], 2)
        self.assertEqual(result["documentMissingActivationIccidCount"], 4)
        self.assertEqual(result["validUniqueIccidCount"], 4)
        self.assertEqual(result["inventoryValidUniqueIccidCount"], 5)
        self.assertEqual(result["readyForUploadCount"], 1)
        self.assertEqual(result["readyForUploadLegacySchemaCount"], 1)
        self.assertEqual(result["alreadyQueriedCount"], 1)
        self.assertEqual(result["refreshExistingMissingCount"], 2)
        self.assertEqual(result["pendingCount"], 3)
        self.assertEqual({row["queryIccid"] for row in result["rows"]}, {iccid(3), iccid(4), iccid(5)})
        self.assertEqual(result["reuseRecords"][0]["queryIccid"], iccid(2))
        self.assertTrue(all(row["sourceData"]["auditField"] == "retained" for row in result["rows"]))
        self.assertEqual(self.document.read_bytes(), before)
        self.assertEqual(result["sourceCardsSha256"], DB.file_sha256(self.cards))

    def test_empty_pool_off_shelf_during_export_is_excluded_from_lookup(self):
        self.save_inputs([('Kooda', iccid(1) + 'F', None)])
        book = load_workbook(self.document)
        sheet = book['sheet1']
        sheet.cell(1, 6, '状态'); sheet.cell(2, 6, '上架')
        sheet.cell(1, 7, '创建时间'); sheet.cell(2, 7, '2026-09-01 00:00:00')
        book.save(self.document); book.close()
        inventory = json.loads(self.cards.read_text())
        inventory['systemRows'][0].update(slotIdentif=None, status='猫池空号已下架', showAble=False,
                                          machineName='130', createTime='2026-09-01 00:00:00')
        self.cards.write_text(json.dumps(inventory))
        before = self.document.read_bytes()
        result = self.prepare()
        self.assertEqual(result['pendingCount'], 0)
        self.assertEqual(result['documentExcludedCount'], 1)
        self.assertEqual(result['documentExcludedRows'][0]['reason'], 'archived_off_shelf_no_current_slot')
        self.assertEqual(result['documentExcludedRows'][0]['backendId'], 'id-1')
        self.assertEqual(self.document.read_bytes(), before)

    def test_fully_dated_document_never_queues_inventory_extras(self):
        self.save_inputs([("Kooda", iccid(1) + "F", "2026-09-07")])
        inventory = DB.load_json(self.cards)
        inventory["systemRows"].append({"id": "extra", "opr": "Kooda", "slotIdentif": "130_2",
                                       "cardNo": "14385550111", "cardStore": iccid(2) + "F"})
        inventory["totalRows"] += 1
        DB.atomic_json(self.cards, inventory)
        result = self.prepare()
        self.assertEqual(result["pendingCount"], 0)
        self.assertEqual(result["validUniqueIccidCount"], 0)
        self.assertEqual(result["inventoryValidUniqueIccidCount"], 2)

    def test_cutoff_only_gap_is_eligible_when_cutoff_column_exists(self):
        self.save_inputs([("Kooda", iccid(1) + "F", "2026-09-07")])
        workbook = load_workbook(self.document)
        workbook["sheet1"].cell(1, 6, "截止日")
        workbook.save(self.document)
        workbook.close()
        result = self.prepare()
        self.assertEqual(result["documentMissingActivationRows"], 0)
        self.assertEqual(result["documentMissingActivationIccidCount"], 0)
        self.assertEqual(result["documentMissingCutoffRows"], 1)
        self.assertEqual(result["documentMissingDateRows"], 1)
        self.assertEqual(result["documentMissingDateIccidCount"], 1)
        self.assertEqual(result["pendingCount"], 1)
        self.assertEqual(result["reportEligibleRecords"][0]["missingFields"], ["截止日"])

    def test_import_delta_marks_new_dates_but_not_reuse_or_schema_only_upgrade(self):
        first = self.lookup(1)
        self.assertEqual(first["activationDateUpdatedCount"], 1)
        self.assertTrue(first["rows"][0]["activationDateUpdated"])
        repeated = self.lookup(1)
        self.assertEqual(repeated["activationDateUpdatedCount"], 0)
        self.assertFalse(repeated["rows"][0]["activationDateUpdated"])
        self.seed(2, "found", "2026-09-07", "2026-10-06", schema=1)
        upgraded = self.lookup(2)
        self.assertEqual(upgraded["schemaUpgradedCount"], 1)
        self.assertEqual(upgraded["activationDateUpdatedCount"], 0)
        self.seed(3, "activation_missing")
        repaired = self.lookup(3, report_mode=True)
        self.assertEqual(repaired["activationDateUpdatedCount"], 1)

    def test_nonempty_invalid_document_activation_fails_closed(self):
        for activation in ("无效日期", "2026-02-30", "=1/0", True):
            with self.subTest(activation=activation):
                self.save_inputs([("Kooda", iccid(1), activation)])
                with self.assertRaisesRegex(RuntimeError, "激活日期非空但无效"):
                    self.prepare()
        self.assertFalse(self.database.exists())

    def test_document_hash_run_identity_mismatch_and_partial_arguments_fail_before_db_creation(self):
        digest = self.save_inputs([("Kooda", iccid(1), None)])
        with self.assertRaisesRegex(RuntimeError, "SHA256"):
            self.prepare("0" * 64)
        inventory = DB.load_json(self.cards)
        inventory["runId"] = "old-run"
        DB.atomic_json(self.cards, inventory)
        with self.assertRaisesRegex(RuntimeError, "runId"):
            self.prepare(digest)
        with self.assertRaisesRegex(RuntimeError, "同时提供"):
            DB.prepare_pending(self.database, self.cards, self.pending, document=self.document)
        self.assertFalse(self.database.exists())

    def test_document_triple_identity_mismatch_and_duplicate_inventory_fail(self):
        for kind in ("iccid", "phone", "slot", "duplicate"):
            with self.subTest(kind=kind):
                self.save_inputs([("Kooda", iccid(1), None)])
                inventory = DB.load_json(self.cards)
                raw = inventory["systemRows"][0]
                if kind == "iccid": raw["cardStore"] = iccid(2)
                if kind == "phone": raw["cardNo"] = "14385559999"
                if kind == "slot": raw["slotIdentif"] = "131_1"
                if kind == "duplicate":
                    inventory["systemRows"].append(dict(raw, id="other-id"))
                    inventory["totalRows"] += 1
                DB.atomic_json(self.cards, inventory)
                with self.assertRaises(RuntimeError):
                    self.prepare()
        self.assertFalse(self.database.exists())

    def test_inventory_flags_and_backend_ids_must_be_independently_complete(self):
        changes = [{"paginationTotalStable": False}, {"paginationBackendIdsUnique": False},
                   {"silentBrowserMode": False}, {"browserForegroundActivationAttempted": True}]
        for change in changes:
            with self.subTest(change=change):
                self.save_inputs([("Kooda", iccid(1), None)], change)
                with self.assertRaises(RuntimeError):
                    self.prepare()
        for missing in (False, True):
            with self.subTest(missing_id=missing):
                self.save_inputs([("Kooda", iccid(1), None), ("Kooda", iccid(2), None)])
                inventory = DB.load_json(self.cards)
                inventory["systemRows"][1]["id"] = "" if missing else inventory["systemRows"][0]["id"]
                DB.atomic_json(self.cards, inventory)
                with self.assertRaisesRegex(RuntimeError, "后台ID"):
                    self.prepare()
        self.assertFalse(self.database.exists())

    def test_named_sheet1_is_required_and_not_replaced_by_first_sheet(self):
        self.save_inputs([("Kooda", iccid(1), None)])
        workbook = load_workbook(self.document)
        workbook.create_sheet("not-the-main", 0)
        workbook["not-the-main"].append(["this worksheet must be ignored"])
        workbook.save(self.document)
        workbook.close()
        self.assertEqual(self.prepare()["pendingCount"], 1)
        workbook = load_workbook(self.document)
        workbook["sheet1"].title = "renamed-main"
        workbook.save(self.document)
        workbook.close()
        with self.assertRaisesRegex(RuntimeError, "sheet1"):
            self.prepare()

    def test_invalid_existing_found_dates_are_requeried(self):
        self.save_inputs([("Kooda", iccid(1), None)])
        self.seed(1, "found", "2026-09-01", "2026-10-01")
        result = self.prepare()
        self.assertEqual(result["pendingCount"], 1)
        self.assertEqual(result["readyForUploadCount"], 0)
        self.assertEqual(result["refreshExistingMissingCount"], 1)

    def test_shared_iccid_collapses_to_one_query_and_retains_all_report_identities(self):
        self.save_inputs([("Kooda", iccid(1) + "F", None), ("Telus1", iccid(1) + "f", None)])
        result = self.prepare()
        self.assertEqual(result["documentMissingActivationRows"], 2)
        self.assertEqual(len(result["reportEligibleRecords"]), 2)
        self.assertEqual(result["documentDuplicateIccidRows"], 1)
        self.assertEqual(result["pendingCount"], 1)
        self.assertEqual(result["rows"][0]["sourceOccurrenceCount"], 2)
        self.assertEqual({row["documentRow"] for row in result["rows"][0]["sourceOccurrences"]}, {2, 3})

    def test_successful_backfill_is_ready_next_time_without_mutating_current_report(self):
        self.save_inputs([("Kooda", iccid(1) + "F", None)])
        self.seed(1, "activation_missing")
        before = self.document.read_bytes()
        first = self.prepare()
        self.assertEqual(first["pendingCount"], 1)
        self.lookup(1)
        second = self.prepare()
        self.assertEqual(second["pendingCount"], 0)
        self.assertEqual(second["readyForUploadCount"], 1)
        self.assertEqual(second["documentMissingActivationRows"], 1)
        self.assertEqual(self.document.read_bytes(), before)

    def test_current_schema_missing_and_not_found_records_upgrade_to_found(self):
        for number, prior_status in ((1, "activation_missing"), (2, "not_found")):
            with self.subTest(prior_status=prior_status):
                self.seed(number, prior_status)
                result = self.lookup(number)
                self.assertEqual(result["foundImportedCount"], 1)
                self.assertEqual(result["refreshedMissingActivationCount"], 1)
                self.assertEqual(result["schemaUpgradedCount"], 0)
                record = self.read_activation(number)
                self.assertEqual(record["lookup_status"], "found")
                self.assertEqual(record["activation_date"], "2026-09-07")
                self.assertEqual(record["cutoff_date"], "2026-10-06")
                self.assertEqual(record["sku"], "SKU-preserved")
                self.assertEqual(record["purchase_date"], "2026-09-01")
                self.assertEqual(json.loads(record["source_data_json"])[0]["sourceData"]["auditField"], "retained")

    def test_failed_or_blank_lookup_never_downgrades_valid_date_even_old_schema(self):
        self.seed(1, "found", "2026-09-01", "2026-09-30", schema=1)
        before = self.read_activation(1)
        for status in ("network_error", "not_found", "activation_missing"):
            with self.subTest(status=status):
                result = self.lookup(1, status=status, activation=None)
                self.assertEqual(result["foundImportedCount"], 0)
                self.assertEqual(self.read_activation(1), before)

    def test_report_blank_retry_refreshes_local_attempt_but_never_downgrades_found(self):
        for number, status in ((1, "activation_missing"), (2, "not_found")):
            with self.subTest(status=status):
                self.seed(number, status)
                legacy = self.lookup(number, status=status, activation=None, sku="legacy-kept")
                self.assertEqual(legacy["alreadyTerminalKeptCount"], 1)
                self.assertIsNone(self.read_activation(number)["sku"])
                result = self.lookup(number, status=status, activation=None, report_mode=True,
                                     sku="latest-blank-attempt", queriedAt="2026-09-07T08:09:10Z")
                self.assertEqual(result["refreshedBlankRecordCount"], 1)
                self.assertEqual(result["rows"][0]["result"], "refreshed_blank_lookup_record")
                record = self.read_activation(number)
                self.assertEqual(record["queried_at"], "2026-09-07T08:09:10Z")
                self.assertEqual(record["sku"], "latest-blank-attempt")
                self.assertIn("auditField", record["source_data_json"])
                self.lookup(number, report_mode=True)
                found = self.read_activation(number)
                kept = self.lookup(number, status=status, activation=None, report_mode=True,
                                   sku="must-not-overwrite")
                self.assertEqual(kept["refreshedBlankRecordCount"], 0)
                self.assertEqual(self.read_activation(number), found)

    def test_lookup_source_sim_and_detail_identity_guards_prevent_cross_card_update(self):
        for field in ("sourceIccid", "simNumber", "detailSimIccid"):
            with self.subTest(field=field):
                result = self.lookup(1, **{field: iccid(2)})
                self.assertEqual(result["terminalImportedCount"], 0)
                self.assertEqual(result["transientErrorCount"], 1)
                self.assertIsNone(self.read_activation(1))

    def test_reissued_old_blank_catalog_cannot_erase_new_lookup_date(self):
        catalog = self.root / "catalog.json"
        row = {"ordinal": 1, "fields": {"lblPurchaseDate": "2026-09-01",
               "lblActivationRequestDate": "", "lblActivateDate": "", "lblSimCardNumber": iccid(1),
               "lblSKU": "OLD-SKU", "lblCost": "$12"}}
        payload = {"status": "passed", "total": 1, "extractedRowCount": 1, "rowCountMatchesTotal": True,
                   "missingRequiredFields": [], "rows": [row]}
        DB.atomic_json(catalog, payload)
        DB.import_catalog(self.database, catalog, self.root / "catalog-first.json")
        self.lookup(1)
        before = self.read_activation(1)
        repeat = DB.import_catalog(self.database, catalog, self.root / "catalog-repeat.json")
        self.assertTrue(repeat["alreadyImported"])
        payload["finishedAt"] = "2026-09-01T00:00:00Z"  # New hash, same old blank data.
        DB.atomic_json(catalog, payload)
        changed = DB.import_catalog(self.database, catalog, self.root / "catalog-reissued.json")
        self.assertEqual(changed["preservedValidLookupCount"], 1)
        self.assertEqual(self.read_activation(1), before)


if __name__ == "__main__":
    unittest.main()
