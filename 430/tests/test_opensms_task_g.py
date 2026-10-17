import copy
from datetime import date
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))
import opensms_task_g as g
NODE = Path('/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node')


class TaskGTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        for name in ("source", "cache", "workbook"):
            (self.dir / name).write_text(name)
        self.readback = self.dir / "readback.json"
        self.dump(self.readback, {"passed": True, "failureCount": 0, "rows": []})
        self.config = {"runId": "20261005T080000+0800", "serverDate": "2026-10-05",
                       "source": str(self.dir / "source"), "scanCache": str(self.dir / "cache"),
                       "postBusinessWorkbook": str(self.dir / "workbook"), "postBusinessReadback": str(self.readback)}
        self.cache = {"headers": ["电话", "运营商", "卡槽", "ICCID", "opr", "10-04总发送量"],
                      "rows": [["17000000001", "B1-test", "101_01", "089123456789012345678F", "Koodo", 161]]}
        self.mock = patch.object(g, "load_scan_cache", return_value=self.cache)
        self.mock.start()
        self.addCleanup(self.mock.stop)

    def dump(self, path, obj):
        path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n")

    def prepare(self):
        return g.prepare(self.config)

    def receipt(self, options=None):
        spec = self.prepare()
        path = self.dir / "input.json"
        self.dump(path, spec)
        args = [str(NODE), str(ROOT / "tests/test_opensms_task_g.mjs"), "--fixture", str(path)]
        if options is not None:
            option_path = self.dir / "options.json"
            self.dump(option_path, options)
            args.append(str(option_path))
        result = subprocess.run(args,
                                capture_output=True, text=True, check=True)
        receipt = json.loads(result.stdout)
        out = self.dir / "receipt.json"
        self.dump(out, receipt)
        return path, out, receipt

    def test_post_business_overlay_is_source(self):
        self.dump(self.readback, {"passed": True, "failureCount": 0, "rows": [
            {"phone": "17000000001", "liveCarrierName": "B12-test", "passed": True}]})
        self.assertEqual(self.prepare()["rows"][0]["sourceCarrier"], "B12-test")

    def test_other_task_isolation_excluded(self):
        self.dump(self.readback, {"passed": True, "failureCount": 0, "rows": [
            {"phone": "7000000001", "liveCarrierName": "过期下架卡T", "passed": True}]})
        self.assertEqual(self.prepare()["rows"], [])

    def test_exact_previous_date_column(self):
        self.cache["headers"][-1] = "10-05总发送量"
        with self.assertRaisesRegex(ValueError, "10-04总发送量"):
            self.prepare()

    def test_unknown_sends_not_zero(self):
        self.cache["rows"][0][-1] = None
        result = self.prepare()
        self.assertIsNone(result["rows"][0]["previousDayTotalSend"])
        self.assertEqual(result["unknownMetricRowCount"], 1)

    def test_bad_identity_is_not_coerced(self):
        self.cache["rows"][0][3] = 8912345678901234567
        with self.assertRaisesRegex(ValueError, "ICCID"):
            self.prepare()

    def test_duplicate_phone_conflict_stops(self):
        self.cache["rows"].append(["7000000001", "B12-test", "101_02", "", "Koodo", 161])
        with self.assertRaisesRegex(ValueError, "冲突"):
            self.prepare()

    def test_unverified_previous_stage_stops(self):
        self.dump(self.readback, {"passed": False, "failureCount": 1})
        with self.assertRaisesRegex(ValueError, "未通过"):
            self.prepare()

    def test_independent_node_python_receipt(self):
        path, out, _ = self.receipt()
        proof = g.audit(path, out)
        self.assertTrue(proof["passed"])
        self.assertEqual(proof["enteredCount"], 1)

    def test_public_threshold_160_strict_boundaries(self):
        for total, expected in ((150, 0), (151, 0), (159, 0), (160, 0), (161, 1)):
            with self.subTest(total=total):
                self.cache["rows"][0][-1] = total
                path, out, receipt = self.receipt()
                self.assertEqual(receipt["threshold"], 160)
                self.assertEqual(g.audit(path, out)["enteredCount"], expected)

    def test_source_kini_exclusion_for_all_six_groups_audits(self):
        self.cache["rows"] = [[str(17000000001 + n), group + "-test", "101_" + str(n + 1), "",
                               "kiNI" if n % 2 else "Canada Kini LTE", 1000]
                              for n, group in enumerate(g.GROUPS)]
        path, out, receipt = self.receipt()
        self.assertEqual(receipt["sourceKiniExcludedCount"], 6)
        self.assertEqual(receipt["batches"], [])
        self.assertEqual(g.audit(path, out)["enteredCount"], 0)

    def test_missing_or_duplicate_opr_column_stops(self):
        self.cache["headers"][4] = "not-opr"
        with self.assertRaisesRegex(ValueError, "OPR"):
            self.prepare()
        self.cache["headers"][4] = "opr"
        self.cache["headers"].append("OPR")
        with self.assertRaisesRegex(ValueError, "OPR"):
            self.prepare()

    def test_uppercase_opr_and_ascii_token_agree_with_executor(self):
        self.cache["headers"][4] = "OPR"
        for value, excluded in (("加拿大Kini", 1), ("KINI-LTE", 1), ("Kinix", 0), ("NotKini", 0)):
            with self.subTest(value=value):
                self.cache["rows"][0][4] = value
                path, out, receipt = self.receipt()
                self.assertEqual(receipt["sourceKiniExcludedCount"], excluded)
                self.assertEqual(g.audit(path, out)["enteredCount"], 1 - excluded)

    def test_unknown_source_opr_safely_excluded_and_audited(self):
        for value in (None, "", "  ", 123, True):
            with self.subTest(value=value):
                self.cache["rows"][0][4] = value
                path, out, receipt = self.receipt()
                self.assertEqual(receipt["sourceUnknownOprCount"], 1)
                self.assertEqual(g.audit(path, out)["enteredCount"], 0)

    def test_live_kini_or_unknown_opr_exclusion_audits(self):
        for opr in ("Kini", None):
            with self.subTest(opr=opr):
                path, out, receipt = self.receipt({"live": [{"phone": "7000000001", "id": "id7000000001",
                    "slot": "101_1", "iccid": "089123456789012345678", "carrierIdentif": "1", "opr": opr}]})
                self.assertEqual(g.audit(path, out)["enteredCount"], 0)
                self.assertEqual(receipt["liveKiniExcludedCount" if opr else "liveUnknownOprCount"], 1)

    def test_opr_change_immediately_before_write_independently_audits(self):
        for opr in ("Kini", None):
            with self.subTest(opr=opr):
                path, out, receipt = self.receipt({"queryOprChanges": [
                    {"queryNumber": 2, "phone": "7000000001", "opr": opr}]})
                self.assertEqual(g.audit(path, out)["enteredCount"], 0)
                self.assertTrue(receipt["skips"][0]["reason"].endswith("before_enter"))

    def test_forged_live_opr_skip_proof_rejected(self):
        path, out, receipt = self.receipt({"live": [{"phone": "7000000001", "id": "id7000000001",
            "slot": "101_1", "iccid": "089123456789012345678", "carrierIdentif": "1", "opr": "Kini"}]})
        receipt["skips"][0]["observed"]["opr"] = "Koodo"
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "OPR跳过"):
            g.audit(path, out)

    def test_independent_audit_rejects_kini_write_before_proof(self):
        path, out, receipt = self.receipt()
        receipt["batches"][0]["before"][0]["opr"] = "Kini"
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "非Kini门槛"):
            g.audit(path, out)

    def test_source_opr_and_exclusion_count_are_independently_rebuilt(self):
        path, out, receipt = self.receipt()
        receipt["sourceKiniExcludedCount"] = 1
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "规则/平台日期"):
            g.audit(path, out)
        receipt["sourceKiniExcludedCount"] = 0
        self.dump(out, receipt)
        spec = g.load(path)
        spec["rows"][0]["opr"] = "Kini"
        self.dump(path, spec)
        receipt["inputSha256"] = g.sha(path)
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "独立重建候选"):
            g.audit(path, out)

    def test_duplicate_phone_opr_conflict_stops_even_if_kini(self):
        row = list(self.cache["rows"][0])
        row[4] = "Kini"
        self.cache["rows"].append(row)
        with self.assertRaisesRegex(ValueError, "冲突"):
            self.prepare()

    def test_fixed_exclusion_configuration_cannot_be_disabled(self):
        settings = g.load(g.SETTINGS)
        settings["excludedEntryOprTokens"] = []
        with patch.object(g, "load", side_effect=lambda path: settings if Path(path) == g.SETTINGS else json.loads(Path(path).read_text())):
            with self.assertRaisesRegex(ValueError, "固定范围"):
                self.prepare()

    def test_receipt_wrong_target_detected(self):
        path, out, receipt = self.receipt()
        receipt["batches"][0]["after"][0]["carrierIdentif"] = "tv"
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "回读不一致"):
            g.audit(path, out)

    def test_altered_input_detected(self):
        path, out, _ = self.receipt()
        (self.dir / "source").write_text("changed")
        with self.assertRaisesRegex(ValueError, "发生变化"):
            g.audit(path, out)

    def test_altered_count_detected(self):
        path, out, receipt = self.receipt()
        receipt["enteredCount"] += 1
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "计数不符"):
            g.audit(path, out)

    def test_ledger_missing_detected(self):
        path, out, receipt = self.receipt()
        receipt["ledgerAfter"]["records"] = {}
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "哈希不一致"):
            g.audit(path, out)

    def test_next_day_restoration_audits_without_source_phone(self):
        _, _, first = self.receipt()
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"] = []
        path, out, second = self.receipt({"ledger": first["ledgerAfter"], "live": first["batches"][0]["after"]})
        proof = g.audit(path, out)
        self.assertEqual(proof["restoredCount"], 1)
        self.assertEqual(second["refreshGroups"][-1]["target"], "B1-test")

    def test_missing_live_slot_detaches_old_record_and_independently_audits(self):
        _, _, first = self.receipt({"useQueryAdapter": True})
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"] = []
        live = first["batches"][0]["after"]
        del live[0]["slot"]
        live[0]["carrierIdentif"] = "other"
        path, out, second = self.receipt({"ledger": first["ledgerAfter"], "live": live, "useQueryAdapter": True})
        proof = g.audit(path, out)
        self.assertEqual(proof["restoredCount"], 0)
        self.assertEqual(second["detachedCount"], 1)
        self.assertEqual(second["batches"], [])
        self.assertEqual(second["skips"][0]["observed"]["slot"], "")
        self.assertEqual(next(iter(second["ledgerAfter"]["records"].values()))["slot"], "101_1")

    def test_missing_live_slot_for_fresh_candidate_is_not_written(self):
        path, out, receipt = self.receipt({"useQueryAdapter": True, "live": [{
            "phone": "7000000001", "id": "id7000000001", "iccid": "089123456789012345678",
            "carrierIdentif": "1", "opr": "Koodo"}]})
        self.assertEqual(g.audit(path, out)["enteredCount"], 0)
        self.assertEqual(receipt["sourceChangedCount"], 1)
        self.assertEqual(receipt["skips"][0]["observed"]["slot"], "")

    def test_slot_disappears_before_restore_and_independently_audits(self):
        _, _, first = self.receipt({"useQueryAdapter": True})
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"] = []
        path, out, second = self.receipt({"ledger": first["ledgerAfter"], "live": first["batches"][0]["after"],
            "useQueryAdapter": True, "querySlotChanges": [{"queryNumber": 2, "phone": "7000000001", "slot": None}]})
        self.assertEqual(g.audit(path, out)["restoredCount"], 0)
        self.assertEqual(second["detachedCount"], 1)
        self.assertEqual(second["skips"][0]["reason"], "changed_before_restore")

    def test_forged_detached_proof_with_unchanged_identity_is_rejected(self):
        _, _, first = self.receipt({"useQueryAdapter": True})
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"] = []
        live = copy.deepcopy(first["batches"][0]["after"])
        live[0]["slot"] = ""
        path, out, second = self.receipt({"ledger": first["ledgerAfter"], "live": live, "useQueryAdapter": True})
        second["skips"][0]["observed"]["slot"] = "101_1"
        self.dump(out, second)
        with self.assertRaisesRegex(ValueError, "脱离区跳过与回读矛盾"):
            g.audit(path, out)

    def test_old_kini_ledger_restores_despite_new_entry_exclusion(self):
        _, _, first = self.receipt()
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"][0][4] = "Kini"
        live = first["batches"][0]["after"]
        live[0]["opr"] = "Kini"
        path, out, second = self.receipt({"ledger": first["ledgerAfter"], "live": live})
        self.assertEqual(g.audit(path, out)["restoredCount"], 1)
        self.assertEqual(second["enteredCount"], 0)
        self.assertEqual(second["sourceKiniExcludedCount"], 1)

    def test_independent_audit_blocks_forged_same_day_reentry(self):
        _, _, first = self.receipt()
        old = copy.deepcopy(next(iter(first["ledgerAfter"]["records"].values())))
        old.update(status="restored", restoredDate="2026-10-06")
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        path, out, second = self.receipt()
        second["ledgerBefore"]["records"][old["key"]] = old
        second["ledgerAfter"]["records"][old["key"]] = old
        second["ledgerAfterSha256"] = g.hashlib.sha256(
            (json.dumps(second["ledgerAfter"], ensure_ascii=False, indent=2) + "\n").encode()).hexdigest()
        self.dump(out, second)
        with self.assertRaisesRegex(ValueError, "恢复当天不得再次"):
            g.audit(path, out)

    def test_restoration_date_must_be_durable_and_correct(self):
        _, _, first = self.receipt()
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"] = []
        path, out, second = self.receipt({"ledger": first["ledgerAfter"], "live": first["batches"][0]["after"]})
        next(iter(second["ledgerAfter"]["records"].values()))["restoredDate"] = "2026-10-05"
        second["ledgerAfterSha256"] = g.hashlib.sha256(
            (json.dumps(second["ledgerAfter"], ensure_ascii=False, indent=2) + "\n").encode()).hexdigest()
        self.dump(out, second)
        with self.assertRaisesRegex(ValueError, "持久日期无依据"):
            g.audit(path, out)

    def test_old_changed_identity_is_not_restored_and_new_verified_card_can_enter(self):
        _, _, first = self.receipt()
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"][0][3] = "089123456789012345679F"
        path, out, second = self.receipt({"ledger": first["ledgerAfter"], "live": [{
            "phone": "7000000001", "id": "new-backend-id", "slot": "101_1",
            "iccid": "089123456789012345679", "carrierIdentif": "1", "opr": "Koodo"}]})
        proof = g.audit(path, out)
        self.assertEqual(proof["restoredCount"], 0)
        self.assertEqual(proof["enteredCount"], 1)
        self.assertEqual(second["detachedCount"], 1)

    def test_old_intent_reconcile_then_restore_independently_audits(self):
        _, _, first = self.receipt()
        next(iter(first["ledgerAfter"]["records"].values()))["status"] = "entering"
        self.config["serverDate"] = "2026-10-06"
        self.cache["headers"][-1] = "10-05总发送量"
        self.cache["rows"] = []
        path, out, _ = self.receipt({"ledger": first["ledgerAfter"], "live": first["batches"][0]["after"]})
        self.assertEqual(g.audit(path, out)["restoredCount"], 1)

    def test_eligible_source_cannot_be_silently_omitted(self):
        path, out, receipt = self.receipt()
        receipt["batches"] = []
        receipt["enteredCount"] = 0
        self.dump(out, receipt)
        with self.assertRaises(ValueError):
            g.audit(path, out)

    def test_storage_summary_not_fabricated(self):
        path, out, receipt = self.receipt()
        receipt["byOriginalCarrier"][0]["count"] += 1
        self.dump(out, receipt)
        with self.assertRaisesRegex(ValueError, "存入汇报"):
            g.audit(path, out)

    def test_real_xlsx_cache_final_refresh_and_parallel_audit(self):
        from openpyxl import Workbook, load_workbook
        from opensms_export_scan_cache import build_scan_cache
        from test_opensms_refresh_workbook_carriers_fast import FastCarrierRefreshTests, MODULE as refresh
        import opensms_parallel_audit as auditor
        self.mock.stop()
        source = self.dir / "fresh.xlsx"
        book = Workbook()
        book.active.title = "sheet1"
        book.active.append(self.cache["headers"])
        book.active.append(self.cache["rows"][0])
        book.active.append(["17000000002", "B1V-test", "101_02", "089123456789012345679F", "Koodo", 161])
        book.active.append(["17000000003", "B1-test", "101_03", "", "Koodo", 160])
        book.save(source)
        book.close()
        cached = self.dir / "scan.pkl"
        build_scan_cache(source, cached)
        self.config.update(source=str(source), scanCache=str(cached), postBusinessWorkbook=str(source))
        path, out, receipt = self.receipt()
        proof = g.audit(path, out)
        self.assertEqual(proof["enteredCount"], 2)
        rows = [{"phone": row["phone"], "liveCarrierName": batch["records"][0]["restCarrier"]["name"]}
                for batch in receipt["batches"] for row in batch["after"]]
        args = FastCarrierRefreshTests().fresh_args(self.dir, source, "g-final", rows)
        refresh.run(args)
        reread = load_workbook(args.output, read_only=True)
        self.assertEqual([r[1] for r in list(reread.active.values)[1:]], ["周末休息区T", "周末休息区TV", "B1-test"])
        reread.close()
        proof_path = self.dir / "proof.json"
        self.dump(proof_path, proof)
        cfg = {"taskG": {"status": "passed", "independentReadbackPassed": True,
                          "input": str(path), "receipt": str(out), "independentAudit": str(proof_path)}}
        resolver = auditor.PathResolver(ROOT, self.dir)
        self.assertEqual(auditor.check_task_g(cfg, resolver, args.output).status, "passed")

    def test_new_rule_runs_last_and_keeps_old_v_guard(self):
        import opensms_patrol_runner as runner
        self.assertEqual(runner.STAGE_ORDER[-1], "taskG")
        self.assertFalse(runner.execution_contract(True)["taskGWeekendRestEnabledForRun"])
        text = (ROOT / "scripts/opensms_patrol_runner.py").read_text()
        start = text.index("    def execute(self)")
        body = text[start:text.index("\ndef static_check", start)]
        self.assertLess(body.index("self.import_current_f_dates(acquire)"), body.index("self.run_task_g(acquire, before_g)"))
        self.assertLess(body.index("self.run_task_g(acquire, before_g)"), body.index("refresh = self.refresh_workbook(document)"))

    def test_chat_has_seven_sections_and_storage_records(self):
        from opensms_chat_summary import render_summary
        summary = render_summary({"status": "success", "taskG": {"status": "passed", "threshold": 160,
            "enteredCount": 2, "restoredCount": 1, "restingCount": 2, "byOriginalCarrier": [
                {"originalCarrier": "B12V-test", "restCarrier": "周末休息区TV", "count": 2,
                 "entryDate": "2026-10-05", "restoreDate": "2026-10-06"}]}})
        self.assertEqual(sum(line.startswith("## ") for line in summary.splitlines()), 7)
        self.assertIn("存入记录：B12V-test", summary)
        self.assertIn("2026-10-06", summary)
        self.assertIn("严格大于 160", summary)
        self.assertIn("Kini不转入", summary)

    def test_unknown_report_counts_remain_unknown(self):
        rendered = "\n".join(g.report_lines({"status": "failed_no_replay", "enteredCount": None}))
        self.assertIn("转入周末休息区 未知", rendered)
        self.assertNotIn("None", rendered)
        self.assertIn("源表排除 未知、实时排除 未知", rendered)


if __name__ == "__main__":
    unittest.main()
