"""Offline generic-XLSX regression checks; temporary artifacts, no live data."""
import json, subprocess, sys, tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'scripts'))
import opensms_task_i_pipeline as P
from openpyxl import load_workbook

def run():
    key='8912230200278094459';phone='14389392694'
    positions=dict(historicalSlots=['101_1'],historicalLastSlot='101_1',historicalLastSlots=['101_1'],
        historicalLastSlotObservedAt='2026-10-11T10:00:00',historicalLastSlotSources=[],historicalLastSlotStatus='confirmed')
    with tempfile.TemporaryDirectory(prefix='task-i-report-test-') as tmp:
        root=Path(tmp)
        for case,queries,has_data in [('empty',[key],False),('mixed',[key,phone],True)]:
            folder=root/case;folder.mkdir();qs=[];daily=[]
            for k in queries:
                qs.append(dict(query=k,historicalPhones=[phone] if has_data else [],documentCount=int(has_data),
                    dates=['2026-10-10'] if has_data else [],partialDates=[],unknownDailyRecords=0,
                    dailyRecordCount=int(has_data),totals=[100,5,130,8] if has_data else [None]*4,
                    replyRate=.05 if has_data else None,**positions))
                if has_data:daily.append(dict(query=k,date='2026-10-10',phone=phone[1:],iccid=key,values=[100,5,130,8],partial=False,
                    path='/test/20261011T100000+0800.xlsx',sheet='sheet1',row=2,iccidCell='B2',observed='2026-10-11T10:00:00',duplicateSnapshots=1))
            d=dict(queries=qs,daily=daily,conflicts=[],unknown=[],phoneSlots={phone[1:]:positions} if has_data else {},
                totals=None if has_data else [0]*4,indexedAt='2026-10-11T10:00:00',dateStart='2026-10-10' if has_data else None,dateEnd='2026-10-10' if has_data else None,
                verification=dict(status='passed',sampleCount=0))
            if has_data:d['portfolioWarning']='查询列表存在ICCID与手机号交叉，组合总计停用'
            P.prepare_result(d,{k:'测试' for k in queries},dict(documentCount=1))
            P.dump(folder/'query_result.json',d)
            with (folder/'build.log').open('w') as log:
                subprocess.run([P.NODE,str(Path(P.__file__).with_name('opensms_task_i_report.mjs')),'--data',str(folder/'query_result.json'),'--output',str(folder/'test.xlsx')],check=True,stdout=log,stderr=subprocess.STDOUT)
            checked=P.verify_workbook(folder/'test.xlsx',d)
            if has_data:
                b=load_workbook(folder/'test.xlsx',read_only=True,data_only=True)
                try:
                    row=next(b['逐手机号统计'].iter_rows(min_row=7,max_row=7))
                    assert [c.value for c in row[4:9]]==[100,8,5,130,8]
                    assert b['批次汇总']['F9'].value=='未知'
                finally:b.close()
            print(json.dumps(dict(case=case,status=checked['status'],queryRows=checked['queryRows'],phoneRows=checked['phoneRows'])))

if __name__=='__main__':run()
