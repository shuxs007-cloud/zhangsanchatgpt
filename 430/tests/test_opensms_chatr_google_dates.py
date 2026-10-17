"""Offline Chatr snapshot regression using temporary synthetic Google workbooks."""
from datetime import datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, to_excel

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from opensms_chatr_google_dates import REQUIRED_SHEETS, freeze, normalize_slot, parse_date


class ChatrGoogleDatesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="chatr-google-dates-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / "fresh-google.xlsx"
        self.output = self.root / "chatr.json"

    def workbook(self, rows=(), headers=None, omit=(), extra_sheets=None):
        wb = Workbook()
        wb.remove(wb.active)
        for name in REQUIRED_SHEETS:
            if name in omit:
                continue
            ws = wb.create_sheet(name)
            ws.append(headers or ["卡槽", "電信商/價格", "激活日期", "預計過期日", "ICCID", "手機號", "上卡日期"])
            if name == REQUIRED_SHEETS[0]:
                for row in rows:
                    ws.append(row)
        for name, rows in (extra_sheets or {}).items():
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        wb.save(self.source)
        wb.close()
        return hashlib.sha256(self.source.read_bytes()).hexdigest()

    def run_freeze(self, rows=(), **kwargs):
        digest = self.workbook(rows, **kwargs)
        receipt = freeze(self.source, digest, self.output, "current-run")
        return receipt, json.loads(self.output.read_text())

    def test_direct_chatr_dates_preserve_time_null_and_provenance_only(self):
        receipt, payload = self.run_freeze([
            ["１１５-０１.01", "Chatr / 35", datetime(2026, 9, 1, 13, 14, 15), "2026/9/24", "8912230200251591144F", "14385550123", "2026-08-01"],
            ["115_2", "chatr", None, "2026-09-20", None, None, "2026-09-02"],
            ["115_3", "CHATR", "2026-09-03", None, None, None, "2026-09-01"],
            ["115_4", "Chatr", None, None, None, None, "2026-08-01"],
            ["115_5", "Kooda", "not-a-date", "not-a-date"],
            ["115_6", "Chatr Koodo", "not-a-date", "not-a-date"],
            ["115_7", "Telus1 Chatr", "not-a-date", "not-a-date"],
        ])
        self.assertEqual(receipt["recordCount"], 4)
        self.assertEqual(receipt["ignoredNonChatrRows"], 3)
        self.assertEqual(payload["sourceSha256"], hashlib.sha256(self.source.read_bytes()).hexdigest())
        self.assertEqual(payload["runId"], "current-run")
        row = payload["records"][0]
        self.assertEqual(row["normalizedSlot"], "115_1")
        self.assertEqual(row["activationDate"], "2026-09-01T13:14:15")
        self.assertEqual(row["cutoffDate"], "2026-09-24")
        self.assertEqual(row["queryIccid"], "8912230200251591144")
        self.assertEqual(row["iccid"], "8912230200251591144F")
        self.assertEqual(row["provenance"][0]["sourceRow"], 2)
        self.assertIsNone(payload["records"][1]["activationDate"])
        self.assertIsNone(payload["records"][2]["cutoffDate"])
        self.assertIsNone(payload["records"][3]["activationDate"])
        self.assertEqual(receipt["missingBothDateRows"], 1)
        self.assertTrue(receipt["verificationPassed"])
        self.assertEqual(receipt["sha256"], hashlib.sha256(self.output.read_bytes()).hexdigest())

    def test_compatible_duplicates_merge_dates_and_keep_every_source_row(self):
        receipt, payload = self.run_freeze([
            ["204_62.01", "Chatr", "2026-09-01", None],
            ["204_62", "Chatr", None, "2026-09-21"],
            ["204_062", "Chatr", datetime(2026, 9, 1), "2026-09-21"],
        ])
        self.assertEqual(receipt["recordCount"], 1)
        self.assertEqual(receipt["mergedDuplicateRows"], 2)
        self.assertEqual(len(payload["records"][0]["provenance"]), 3)
        self.assertEqual(payload["records"][0]["cutoffDate"], "2026-09-21")

    def test_conflicting_calendar_time_or_identity_never_publishes(self):
        for second in (
            ["115_1", "Chatr", "2026-09-02", "2026-09-20", "8912230200251591144"],
            ["115_1", "Chatr", "2026-09-01 12:00:00", "2026-09-20", "8912230200251591144"],
            ["115_1", "Chatr", "2026-09-01", "2026-09-21", "8912230200251591144"],
            ["115_1", "Chatr", "2026-09-01", "2026-09-20", "8912230200251591152"],
        ):
            with self.subTest(second=second):
                digest = self.workbook([["115_01", "Chatr", "2026-09-01", "2026-09-20", "8912230200251591144"], second])
                with self.assertRaisesRegex(ValueError, "冲突"):
                    freeze(self.source, digest, self.output, "run")
                self.assertFalse(self.output.exists())

    def test_invalid_ambiguous_error_and_uncached_formula_dates_fail(self):
        for invalid in ("09/10/2026", "2026-02-30", "activated 2026-09-01", "#N/A", "=TODAY()", True):
            with self.subTest(invalid=invalid):
                digest = self.workbook([["115_1", "Chatr", invalid, None]])
                with self.assertRaises(ValueError):
                    freeze(self.source, digest, self.output, "run")
                self.assertFalse(self.output.exists())

    def test_missing_required_sheet_column_or_duplicate_column_fail(self):
        for kwargs in ({"omit": (REQUIRED_SHEETS[-1],)},
                       {"headers": ["卡槽", "電信商/價格", "激活日期", "截止日"]},
                       {"headers": ["卡槽", "電信商/價格", "激活日期", "預計過期日", "激活日期"]}):
            with self.subTest(kwargs=kwargs):
                digest = self.workbook(**kwargs)
                with self.assertRaises(ValueError):
                    freeze(self.source, digest, self.output, "run")
                self.assertFalse(self.output.exists())

    def test_only_four_required_headers_and_empty_dates_are_valid(self):
        receipt, payload = self.run_freeze([["115_1", "Chatr", None, None]], headers=["卡槽", "電信商/價格", "激活日期", "預計過期日"])
        self.assertEqual(receipt["recordCount"], 1)
        self.assertIsNone(payload["records"][0]["iccid"])

    def test_hash_run_id_and_existing_output_are_enforced(self):
        digest = self.workbook([["115_1", "Chatr", "2026-09-01", None]])
        for bad_hash, run in (("0" * 64, "run"), ("", "run"), (digest, " ")):
            with self.subTest(hash=bad_hash, run=run), self.assertRaises(ValueError):
                freeze(self.source, bad_hash, self.output, run)
        before = self.source.read_bytes()
        freeze(self.source, digest, self.output, "run")
        saved = self.output.read_bytes()
        with self.assertRaises(FileExistsError):
            freeze(self.source, digest, self.output, "run")
        self.assertEqual(self.output.read_bytes(), saved)
        self.assertEqual(self.source.read_bytes(), before)

    def test_invalid_slots_excluded_numeric_iccid_not_reconstructed(self):
        receipt, payload = self.run_freeze([
            [None, "Chatr", "2026-09-01", None],
            ["not-a-slot", "Chatr", "2026-09-01", None],
            ["115_1", "Chatr", None, None, 8912230200251591144],
        ])
        self.assertEqual(receipt["excludedInvalidSlotRows"], 2)
        self.assertEqual(len(payload["excludedRows"]), 2)
        self.assertEqual(receipt["numericIccidRows"], 1)
        self.assertIsNone(payload["records"][0]["queryIccid"])

    def test_normalization_matches_document_pipeline(self):
        spec = importlib.util.spec_from_file_location("chatr_slot_existing", ROOT / "services/opensms_document_api/pipeline/process_card_slots.py")
        document = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(document)
        for raw in ("115_01", "115_1", "115_1.01", "２０４-０６２.01", " 115 _ 01 ", None, "A_01"):
            with self.subTest(raw=raw):
                self.assertEqual(normalize_slot(raw), document.normalize_slot(raw))

    def test_excel_epochs_and_direct_time_preservation(self):
        for epoch in (CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904):
            value = datetime(2026, 9, 7, 16, 42, 30)
            self.assertEqual(parse_date(to_excel(value, epoch), epoch), value.isoformat())
        for raw in ("2026年9月7日", "20260907", 20260907, "2026/9/7"):
            self.assertEqual(parse_date(raw, CALENDAR_WINDOWS_1900), "2026-09-07")
        self.assertEqual(parse_date("2026-09-07T16:42:30-05:00", CALENDAR_WINDOWS_1900), "2026-09-07T16:42:30")

    def test_cli_receipt_has_hash_and_failure_does_not_publish(self):
        digest = self.workbook([["115_1", "Chatr", "2026-09-01", None]])
        command = [sys.executable, str(ROOT / "scripts/opensms_chatr_google_dates.py"),
                   "--source", str(self.source), "--source-sha256", digest,
                   "--run-id", "cli-run", "--output", str(self.output)]
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        receipt = json.loads(result.stdout)
        self.assertEqual(receipt["runId"], "cli-run")
        self.assertEqual(receipt["output"], str(self.output.resolve()))
        again = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(again.returncode, 2)
        self.assertEqual(json.loads(again.stdout)["status"], "failed")


if __name__ == "__main__":
    unittest.main()
