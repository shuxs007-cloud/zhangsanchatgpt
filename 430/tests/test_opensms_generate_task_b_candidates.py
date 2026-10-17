from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "opensms_generate_task_b_candidates.py"


class TaskBCandidateTests(unittest.TestCase):
    def test_protected_carrier_rows_are_included_in_both_task_b_routes(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "result.xlsx"
            output_dir = root / "snapshots"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "sheet1"
            sheet.append(["电话", "运营商", "opr", "卡槽", "过期状态", "截止日"])
            sheet.append(["1001", "A21-股票-C-100", "Chatr", "101_1", "已过期", "2026-08-20"])
            sheet.append(["1002", "A21-股票-C-100", "Koodo", "101_2", "快过期", "2026-08-28"])
            sheet.append(["1009", "A2-股票-C-120", "Chatr", "101_9", "已过期", "2026-08-20"])
            sheet.append(["1010", "A2-股票-C-120", "Koodo", "101_10", "快过期", "2026-08-28"])
            sheet.append(["1003", "A1-测试", "Chatr", "101_3", "已过期", "2026-08-20"])
            sheet.append(["1004", "A1-测试", "Kini", "101_4", "快过期", "2026-08-28"])
            sheet.append(["1011", "B1V-测试", "Telus1", "101_11", "快过期", "2026-08-28"])
            sheet.append(["1012", "快过期续费C", "Chatr", "101_12", "已过期", "2026-08-20"])
            sheet.append(["1013", "快过期续费C", "Chatr", "101_13", "快过期", "2026-08-28"])
            sheet.append(["1014", "快过期续费T", "Koodo", "101_14", "已过期", "2026-08-20"])
            sheet.append(["1015", "B14-C快下保留卡", "Chatr", "101_15", "已过期", "2026-08-20"])
            sheet.append(["1016", "B14-C快下保留卡", "Chatr", "101_16", "快过期", "2026-08-28"])
            sheet.append(["1017", "B14-T快下保留卡", "Telus1", "101_17", "已过期", "2026-08-20"])
            sheet.append(["1018", "B14-T快下保留卡", "Koodo", "101_18", "快过期", "2026-08-28"])
            sheet.append(["1019", "B14-T快下保留卡", "Koodo", "101_19", "快过期", "2026-08-29"])
            workbook.save(workbook_path)

            sheet.append(["1020", "三方隔离", "Koodo", "101_20", "已过期", "2026-08-20"])
            sheet.append(["1021", "三方隔离", "Chatr", "101_21", "快过期", "2026-08-28"])
            workbook.save(workbook_path)

            env = os.environ.copy()
            env["OPENSMS_FORCE_TASK_B"] = "1"
            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(workbook_path),
                    "--server-date",
                    "2026-08-27",
                    "--tag",
                    "test",
                    "--output-dir",
                    str(output_dir),
                ],
                check=False,
                capture_output=True,
                text=True,
                env=env,
                cwd=ROOT,
            )
            self.assertEqual(0, result.returncode, result.stderr)
            manifest = json.loads(result.stdout)
            expired = json.loads(Path(manifest["expiredCandidatePaths"]["C"]).read_text())
            tomorrow = json.loads(Path(manifest["tomorrowCandidatePaths"]["T"]).read_text())
            self.assertEqual(
                ["1001", "1009", "1003", "1012", "1015"],
                [row["电话"] for row in expired["rows"]],
            )
            self.assertEqual(
                ["1002", "1010", "1004", "1011", "1018"],
                [row["电话"] for row in tomorrow["rows"]],
            )
            self.assertEqual(["A21", "A2"], manifest["protectedSourceCarrierGroups"])
            self.assertTrue(manifest["taskBProtectedSourceExceptionEnabled"])
            self.assertEqual(4, manifest["protectedSourceIncludedCount"])
            self.assertEqual({"A21": 2, "A2": 2}, manifest["protectedSourceIncludedCounts"])
            self.assertEqual(0, manifest["protectedSourceExcludedCount"])
            self.assertEqual({"C": 5, "T": 3}, manifest["expiredCandidateCounts"])
            self.assertEqual({"C": 3, "T": 5}, manifest["tomorrowCandidateCounts"])
            self.assertEqual([], manifest["neverProcessSourceCarrierNames"])
            self.assertEqual(0, manifest["neverProcessSourceExcludedCount"])
            self.assertEqual(0, manifest["isolationCarrierHeldCount"])
            self.assertEqual({"快过期续费C": 0, "快过期续费T": 0}, manifest["neverProcessSourceExcludedCounts"])
            self.assertEqual({"快过期续费C": [], "快过期续费T": []}, manifest["neverProcessSourceExcludedRows"])
            self.assertFalse(manifest["retainedSourceProtectionEnforced"])
            self.assertEqual([], manifest["retainedSourceCarrierNames"])
            self.assertEqual(0, manifest["retainedSourceExcludedCount"])
            self.assertEqual({"B14-C快下保留卡": 0, "B14-T快下保留卡": 0}, manifest["retainedSourceExcludedCounts"])
            self.assertEqual({"B14-C快下保留卡": [], "B14-T快下保留卡": []}, manifest["retainedSourceExcludedRows"])
            self.assertEqual(0, manifest["expiredRetainedSourceExcludedCount"])
            self.assertEqual(0, manifest["tomorrowRetainedSourceExcludedCount"])
            for candidate_path in manifest["candidatePaths"]:
                candidate = json.loads(Path(candidate_path).read_text())
                self.assertFalse(candidate["retainedSourceProtectionEnforced"])
                self.assertEqual(manifest["retainedSourceCarrierNames"], candidate["retainedSourceCarrierNames"])
                self.assertFalse(set(candidate["sourceCarrierNames"]) & set(manifest["retainedSourceCarrierNames"]))
            self.assertTrue(manifest["expiredCoveragePassed"])
            self.assertEqual(manifest["expiredNonProtectedSourceCount"], manifest["expiredClassifiedCandidateCount"])
            self.assertTrue(manifest["tomorrowCoveragePassed"])
            self.assertEqual(manifest["tomorrowNonProtectedSourceCount"], manifest["tomorrowClassifiedCandidateCount"])
            self.assertTrue(manifest["allDueCoveragePassed"])

    def test_b14_protection_uses_exact_names_without_changing_tomorrow_threshold(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "result.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "sheet1"
            sheet.append(["电话", "运营商", "opr", "卡槽", "过期状态", "截止日"])
            for index, carrier in enumerate(["B14", "B14-C快下保留卡-备份", "B14-T快下保留卡其他"]):
                sheet.append([str(2001 + index), carrier, "Chatr", "101_1", "已过期", "2026-08-20"])
            sheet.append(["2004", "B14-其它", "Koodo", "101_4", "快过期", "2026-08-28"])
            sheet.append(["2005", "B14-其它", "Koodo", "101_5", "快过期", "2026-08-29"])
            workbook.save(workbook_path)
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(workbook_path), "--server-date", "2026-08-27",
                 "--tag", "exact", "--output-dir", str(root / "snapshots")],
                check=False, capture_output=True, text=True,
                env={**os.environ, "OPENSMS_FORCE_TASK_B": "1"}, cwd=ROOT,
            )
            self.assertEqual(0, result.returncode, result.stderr)
            manifest = json.loads(result.stdout)
            expired = json.loads(Path(manifest["expiredCandidatePaths"]["C"]).read_text())
            tomorrow = json.loads(Path(manifest["tomorrowCandidatePaths"]["T"]).read_text())
            self.assertEqual(["2001", "2002", "2003"], [row["电话"] for row in expired["rows"]])
            self.assertEqual(["2004"], [row["电话"] for row in tomorrow["rows"]])
            self.assertEqual("过期下架卡C", expired["targetCarrierName"])
            self.assertEqual("快过期隔离区T", tomorrow["targetCarrierName"])
            self.assertEqual(0, manifest["retainedSourceExcludedCount"])
            self.assertTrue(manifest["allDueCoveragePassed"])

    def test_unclassified_expired_non_protected_card_is_reported_as_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "result.xlsx"
            output_dir = root / "snapshots"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "sheet1"
            sheet.append(["电话", "运营商", "opr", "卡槽", "过期状态", "截止日"])
            sheet.append(["1001", "A1-测试", "未知OPR", "101_1", "已过期", "2026-08-20"])
            workbook.save(workbook_path)
            env = os.environ.copy()
            env["OPENSMS_FORCE_TASK_B"] = "1"
            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(workbook_path),
                    "--server-date", "2026-08-27",
                    "--tag", "test",
                    "--output-dir", str(output_dir),
                ],
                check=False,
                capture_output=True,
                text=True,
                env=env,
                cwd=ROOT,
            )
            self.assertEqual(0, result.returncode, result.stderr)
            manifest = json.loads(result.stdout)
            self.assertEqual(1, manifest['dataAnomalySkippedCount'])
            self.assertEqual('opr_unclassified', manifest['dataAnomalySkippedRows'][0]['reason'])
            self.assertEqual(0, manifest['expiredCandidateCount'])
            self.assertTrue(manifest['expiredCoveragePassed'])

    def test_unclassified_tomorrow_non_protected_card_is_reported_as_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "result.xlsx"
            output_dir = root / "snapshots"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "sheet1"
            sheet.append(["电话", "运营商", "opr", "卡槽", "过期状态", "截止日"])
            sheet.append(["1001", "A1-测试", "未知OPR", "101_1", "快过期", "2026-08-28"])
            workbook.save(workbook_path)
            env = os.environ.copy()
            env["OPENSMS_FORCE_TASK_B"] = "1"
            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(workbook_path),
                    "--server-date", "2026-08-27",
                    "--tag", "test",
                    "--output-dir", str(output_dir),
                ],
                check=False,
                capture_output=True,
                text=True,
                env=env,
                cwd=ROOT,
            )
            self.assertEqual(0, result.returncode, result.stderr)
            manifest = json.loads(result.stdout)
            self.assertEqual(1, manifest['dataAnomalySkippedCount'])
            self.assertEqual(0, manifest['tomorrowCandidateCount'])
            self.assertTrue(manifest['tomorrowCoveragePassed'])


if __name__ == "__main__":
    unittest.main()
