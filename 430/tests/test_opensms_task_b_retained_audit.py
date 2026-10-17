from __future__ import annotations

from contextlib import redirect_stdout
from datetime import date
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


def load_module(name: str, filename: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


AUDIT = load_module("b14_audit_under_test", "opensms_parallel_audit.py")
GENERATOR = load_module("b14_generator_fixture", "opensms_generate_task_b_candidates.py")
RETAINED = ["B14-C快下保留卡", "B14-T快下保留卡"]


class TaskBRetainedAuditTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="opensms-b14-audit-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "fake-source.xlsx"
        self.config_path = self.root / "config.json"
        self.write_json(self.config_path, {})
        self.resolver = AUDIT.PathResolver(self.config_path, self.root)

    @staticmethod
    def write_json(path, data):
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    @staticmethod
    def read_json(path):
        return json.loads(Path(path).read_text(encoding="utf-8"))

    def fixture(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet1"
        sheet.append(["电话", "运营商", "opr", "卡槽", "过期状态", "截止日"])
        rows = [
            ["1001", RETAINED[0], "Chatr", "1-1", "已过期", "2026-09-07"],
            ["1002", RETAINED[1], "Kooda", "1-2", "已过期", "2026-09-06"],
            ["1003", RETAINED[0], "Chatr", "1-3", "快过期", "2026-09-08"],
            ["1004", RETAINED[1], "Telus1", "1-4", "快过期", "2026-09-08"],
            ["1005", RETAINED[0], "Chatr", "1-5", "快过期", "2026-09-09"],
            ["1006", "快过期续费C", "Chatr", "1-6", "已过期", "2026-09-01"],
            ["1007", "快过期续费T", "Kooda", "1-7", "快过期", "2026-09-08"],
            ["2001", "A21-股票-C120", "Chatr", "2-1", "已过期", "2026-09-07"],
            ["2002", "A2-股票-T120", "Kini", "2-2", "已过期", "2026-09-07"],
            ["2003", "A21-股票-C120", "Chatr", "2-3", "快过期", "2026-09-08"],
            ["2004", "A2-股票-T120", "Koodo", "2-4", "快过期", "2026-09-08"],
            ["2005", "B14-C快下保留卡-其他", "Chatr", "2-5", "已过期", "2026-09-07"],
            ["2006", "b14-t快下保留卡", "Kooda", "2-6", "快过期", "2026-09-08"],
            ["2007", "A1", "Chatr", "2-7", "快过期", "2026-09-09"],
        ]
        for row in rows:
            sheet.append(row)
        workbook.save(self.source)
        workbook.close()
        args = [str(SCRIPTS / "opensms_generate_task_b_candidates.py"), str(self.source),
                "--server-date", "2026-09-07", "--tag", "fake", "--output-dir", str(self.root)]
        # All source data and receipts are synthetic; no patrol runner or browser is used.
        # Historical receipts keep their original protection contract; current
        # all-source routing is exercised by test_opensms_task_b_all_carriers.
        with patch.object(GENERATOR, "TASK_B_SOURCE_SCOPE", "legacy_protected_sources_v1"), patch.dict(os.environ, {"OPENSMS_FORCE_TASK_B": "1"}), patch.object(sys, "argv", args), redirect_stdout(io.StringIO()):
            GENERATOR.main()
        self.manifest_path = self.root / "opensms_task_b_fake_manifest.json"
        self.manifest = self.read_json(self.manifest_path)
        self.config = {
            "taskBExecuted": True, "taskBManifest": str(self.manifest_path),
            "taskBInputWorkbook": str(self.source), "serverDate": "2026-09-07",
            "groups": [{"title": f"任务B组{index}", "candidate": value}
                       for index, value in enumerate(self.manifest["candidatePaths"])],
        }
        return self.config

    def check(self):
        return AUDIT.check_task_b(self.config, self.resolver)

    def save_manifest(self):
        self.write_json(self.manifest_path, self.manifest)

    def candidate(self, branch="expired", family="C"):
        path = Path(self.manifest[f"{branch}CandidatePaths"][family])
        return path, self.read_json(path)

    def update_source(self, edit, rebind_hash=True):
        workbook = load_workbook(self.source)
        edit(workbook["sheet1"])
        workbook.save(self.source)
        workbook.close()
        if rebind_hash:
            source_sha = hashlib.sha256(self.source.read_bytes()).hexdigest()
            self.manifest["sourceSha256"] = source_sha
            self.save_manifest()
            for path in self.manifest["candidatePaths"]:
                candidate = self.read_json(path)
                candidate["sourceSha256"] = source_sha
                self.write_json(Path(path), candidate)

    def test_independent_partitions_preserve_a21_a2_and_exact_name_boundaries(self):
        self.fixture()
        result = self.check()
        self.assertEqual(AUDIT.PASS, result.status, result.message)
        self.assertTrue(result.details["source_coverage_independently_verified"])
        self.assertEqual({"expired": 3, "tomorrow": 3}, result.details["source_non_protected_counts"])
        self.assertEqual({RETAINED[0]: 2, RETAINED[1]: 2}, result.details["source_retained_excluded_counts"])
        self.assertEqual({"expired": 2, "tomorrow": 2}, result.details["source_retained_excluded_branch_counts"])
        self.assertEqual(2, result.details["source_renewal_excluded_count"])
        self.assertEqual(["快过期续费C", "快过期续费T"], result.details["never_process_source_carriers"])

    def test_omitted_non_protected_row_fails_even_with_consistent_manifest_counts(self):
        self.fixture()
        path, candidate = self.candidate()
        candidate["rows"] = candidate["rows"][1:]
        candidate["candidates"] = candidate["rows"]
        candidate["count"] = candidate["candidateCount"] = len(candidate["rows"])
        candidate["sourceCarrierNames"] = sorted({row["运营商"] for row in candidate["rows"]})
        self.write_json(path, candidate)
        for key in ("expiredNonProtectedSourceCount", "expiredClassifiedCandidateCount", "expiredCandidateCount"):
            self.manifest[key] -= 1
        self.manifest["expiredCandidateCounts"]["C"] -= 1
        self.save_manifest()
        result = self.check()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("非保护", result.message)

    def test_relabeling_actionable_row_as_retained_exclusion_fails(self):
        self.fixture()
        self.manifest["retainedSourceExcludedRows"][RETAINED[0]].append(9)
        self.manifest["retainedSourceExcludedCounts"][RETAINED[0]] += 1
        self.manifest["retainedSourceExcludedCount"] += 1
        self.manifest["expiredRetainedSourceExcludedCount"] += 1
        self.save_manifest()
        result = self.check()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("保护排除源行", result.message)

    def test_each_retained_exclusion_field_is_verified(self):
        self.fixture()
        original = json.loads(json.dumps(self.manifest))
        for field in ("retainedSourceProtectionEnforced", "retainedSourceCarrierNames", "retainedSourceExcludedRows", "retainedSourceExcludedCounts",
                      "retainedSourceExcludedCount", "expiredRetainedSourceExcludedCount", "tomorrowRetainedSourceExcludedCount"):
            with self.subTest(field=field):
                self.manifest = json.loads(json.dumps(original))
                del self.manifest[field]
                self.save_manifest()
                self.assertEqual(AUDIT.FAIL, self.check().status)

    def test_renewal_exclusions_remain_separate_and_verified(self):
        self.fixture()
        self.manifest["neverProcessSourceCarrierNames"] += RETAINED
        self.save_manifest()
        self.assertEqual(AUDIT.FAIL, self.check().status)

    def test_forged_retained_candidate_fails_source_coverage(self):
        self.fixture()
        path, candidate = self.candidate()
        for collection in ("rows", "candidates"):
            candidate[collection][0]["运营商"] = RETAINED[0]
            candidate[collection][0]["源行号"] = 2
        candidate["sourceCarrierNames"] = sorted({row["运营商"] for row in candidate["rows"]})
        self.write_json(path, candidate)
        result = self.check()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("候选含保护来源", result.message)

    def test_protected_source_is_verified_even_when_candidate_hides_source_name(self):
        self.fixture()
        path, candidate = self.candidate()
        for collection in ("rows", "candidates"):
            candidate[collection][0]["电话"] = "1001"
            candidate[collection][0]["源行号"] = 2
        self.write_json(path, candidate)
        self.assertEqual(AUDIT.FAIL, self.check().status)

    def test_source_hash_drift_fails(self):
        self.fixture()
        self.update_source(lambda sheet: setattr(sheet["B9"], "value", RETAINED[0]), rebind_hash=False)
        result = self.check()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("SHA-256", result.message)

    def test_non_protected_missing_phone_and_ambiguous_opr_cannot_be_hidden(self):
        for column, value in (("A9", None), ("C9", "Chatr Kooda"), ("C9", "unknown")):
            with self.subTest(column=column, value=value):
                if not hasattr(self, "manifest"):
                    self.fixture()
                def edit(sheet):
                    sheet["A9"] = "2001"
                    sheet["C9"] = "Chatr"
                    sheet[column] = value
                self.update_source(edit)
                result = self.check()
                self.assertEqual(AUDIT.FAIL, result.status)
                self.assertIn("缺号或OPR无法唯一分区" if value == "Chatr Kooda" else "异常数据排除", result.message)

    def test_protected_missing_phone_and_opr_do_not_reduce_non_protected_coverage(self):
        self.fixture()
        def edit(sheet):
            sheet["A2"] = None
            sheet["C2"] = "unknown"
        self.update_source(edit)
        result = self.check()
        self.assertEqual(AUDIT.PASS, result.status, result.message)
        self.assertEqual({"expired": 3, "tomorrow": 3}, result.details["source_non_protected_counts"])

    def test_all_due_rows_protected_allows_four_empty_candidate_files(self):
        self.fixture()
        self.update_source(lambda sheet: sheet.delete_rows(9, 7))
        for path in self.manifest["candidatePaths"]:
            candidate = self.read_json(path)
            candidate.update({"rows": [], "candidates": [], "count": 0, "candidateCount": 0, "sourceCarrierNames": []})
            self.write_json(Path(path), candidate)
        for branch in ("expired", "tomorrow"):
            self.manifest[f"{branch}NonProtectedSourceCount"] = 0
            self.manifest[f"{branch}ClassifiedCandidateCount"] = 0
            self.manifest[f"{branch}CandidateCount"] = 0
            self.manifest[f"{branch}CandidateCounts"] = {"C": 0, "T": 0}
        self.save_manifest()
        result = self.check()
        self.assertEqual(AUDIT.PASS, result.status, result.message)
        self.assertEqual({"expired": 0, "tomorrow": 0}, result.details["actual_classified_candidate_counts"])

    def test_two_day_fast_expiry_is_not_added_to_current_tomorrow_route(self):
        self.fixture()
        self.update_source(lambda sheet: setattr(sheet["F12"], "value", "2026-09-09"))
        self.assertEqual(AUDIT.FAIL, self.check().status)

    def test_target_cannot_change_to_renewal_without_policy_change(self):
        self.fixture()
        path, candidate = self.candidate("tomorrow")
        candidate["targetCarrierName"] = "快过期续费C"
        self.write_json(path, candidate)
        result = self.check()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("路由或目标", result.message)

    def test_candidate_alias_cannot_hide_extra_protected_row(self):
        self.fixture()
        path, candidate = self.candidate()
        candidate["candidates"].append({"电话": "1001", "运营商": RETAINED[0]})
        self.write_json(path, candidate)
        result = self.check()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("两份候选集合", result.message)

    def test_duplicate_source_row_cannot_preserve_coverage_by_count(self):
        self.fixture()
        path, candidate = self.candidate()
        candidate["rows"][1] = dict(candidate["rows"][0])
        candidate["candidates"] = candidate["rows"]
        self.write_json(path, candidate)
        self.assertEqual(AUDIT.FAIL, self.check().status)

    def test_actual_groups_must_reference_exact_four_candidate_files(self):
        self.fixture()
        self.config["groups"][1]["candidate"] = self.config["groups"][0]["candidate"]
        result = self.check()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("实际审计组", result.message)

    def test_disabled_task_does_not_read_any_workbook(self):
        with patch.object(AUDIT, "load_workbook", side_effect=AssertionError("must not read")):
            result = AUDIT.check_task_b({"taskBExecuted": False}, self.resolver)
        self.assertEqual(AUDIT.PASS, result.status)

    def test_audit_parses_excel_date_without_importing_generator_logic(self):
        self.assertEqual(date(2026, 9, 8), AUDIT.task_b_source_date("2026年9月8日", date(1899, 12, 30)))
        self.assertIsNone(AUDIT.task_b_source_date("invalid", date(1899, 12, 30)))

    def protected_check(self, payload, title="任务B组", field="candidate"):
        path = self.root / f"fake-{field}.json"
        self.write_json(path, payload)
        return AUDIT.check_protected_carrier_sources([{field: str(path), "title": title}], self.resolver)

    def test_retained_source_candidates_block_for_both_carriers_even_without_force_gate(self):
        for carrier in RETAINED:
            for force in (True, False):
                with self.subTest(carrier=carrier, force=force):
                    result = self.protected_check({"task": "B", "forceGate": force,
                                                   "rows": [{"电话": "1001", "运营商": carrier}]})
                    self.assertEqual(AUDIT.FAIL, result.status)
                    self.assertEqual(1, result.details["task_b_retained_source_hit_count"])

    def test_live_retained_source_in_write_receipt_is_blocked_via_candidate_path(self):
        candidate = self.root / "task-b.json"
        self.write_json(candidate, {"task": "B", "forceGate": True})
        result = self.protected_check({"candidatePath": str(candidate), "before": {
            "eligibleRows": [{"carrierName": RETAINED[1]}]}}, title="", field="write")
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertEqual(1, result.details["task_b_retained_source_hit_count"])

    def test_retained_protection_does_not_expand_to_other_tasks_or_similar_names(self):
        for task in ("A", "C", "D", "E"):
            with self.subTest(task=task):
                result = self.protected_check({"task": task, "rows": [{"运营商": RETAINED[0]}]}, title=f"任务{task}")
                self.assertEqual(AUDIT.PASS, result.status)
        for carrier in ("B14-C快下保留卡-其他", "b14-t快下保留卡"):
            with self.subTest(carrier=carrier):
                self.assertEqual(AUDIT.PASS, self.protected_check({"task": "B", "rows": [{"运营商": carrier}]}).status)

    def test_b14_target_for_task_d_remains_allowed(self):
        result = self.protected_check({"task": "D", "targetCarrierName": RETAINED[0],
                                       "rows": [{"运营商": "A1"}]}, title="任务D")
        self.assertEqual(AUDIT.PASS, result.status)


if __name__ == "__main__":
    unittest.main()
