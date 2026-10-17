"""Offline no-Google audit, provenance, delivery and report regressions."""
from __future__ import annotations

import contextlib
from copy import deepcopy
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


def module(name):
    alias = "no_google_audit_test_" + name
    spec = importlib.util.spec_from_file_location(alias, SCRIPTS / f"{name}.py")
    result = importlib.util.module_from_spec(spec)
    with patch.object(sys, "path", [str(SCRIPTS), *sys.path]), patch.dict(sys.modules, {alias: result}):
        spec.loader.exec_module(result)
    return result


AUDIT = module("opensms_parallel_audit")
REPORT = module("opensms_finalize_patrol_report")


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return str(path)


class NoGoogleAuditReportTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="opensms-offline-audit-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.config_path = self.root / "config.json"
        self.source = self.root / "source.xlsx"
        self.source.write_bytes(b"immutable current source fixture")
        self.result = self.root / "result.xlsx"
        self.result.write_bytes(b"immutable current result fixture")
        self.cache = self.root / "frozen_dates.json"
        created = "2026-09-07T00:00:00+08:00"
        write_json(self.cache, {"schemaVersion": 1, "runId": "offline-current", "status": "passed",
                               "recordCount": 0, "records": [], "createdAt": created})
        resolution = {"source": "system_export_then_frozen_local", "field_counts": {
            "激活日期": {"system_export": 0, "local_cache": 0, "google": 0, "missing": 1}},
            "cache": {"path": str(self.cache), "sha256": AUDIT.sha256_file(self.cache),
                      "run_id": "offline-current", "created_at": created},
            "source_export": {"path": str(self.source), "sha256": AUDIT.sha256_file(self.source)},
            "association": None, "new_query_results_apply_next_export": True,
            "redundant_system_date_columns_removed": True}
        self.card = {"result": "成功", "checks": {"3日4日回复率明细": True},
            "process": {"main_rows": 1, "matched_rows": 0, "unmatched_slot_count": 0,
                "google_dependency": deepcopy(AUDIT.NO_GOOGLE_DEPENDENCY), "date_resolution": resolution,
                "activated_unlisted_rows": 0, "activated_unlisted_machine_count": 0,
                "over_five_unlisted_machine_count": 0, "slot_alias_204_62_matches": True},
            "enrichment": {"decreasing_rate_counts": {}}, "verification": {
                "formula_error_matches": 0, "google_dependency_review": {
                    **AUDIT.NO_GOOGLE_DEPENDENCY, "matches": True},
                "unused_port_review": {"matches": True},
                "system_kooda_source_review": {"matches": True, "source": "system", "run_id": "offline-current"},
                "task_f_missing_activation_red_review": {"matches": True},
                "task_f_missing_dates_new_card_review": {"matches": True},
                "date_resolution_review": {"matches": True, "field_counts": deepcopy(resolution["field_counts"]),
                                           "cache": deepcopy(resolution["cache"])}}}
        self.card_path = self.root / "card.json"
        self.dispatch = {"result": "成功", "processing_mode": "local", "inputs": {
            "association": None, "google_dependency": deepcopy(AUDIT.NO_GOOGLE_DEPENDENCY),
            "export4day": {"sha256": AUDIT.sha256_file(self.source)}, "server_date": "2026-09-07"},
            "cell_type_policy": {"numeric_text_conversion_enabled": False}}
        self.dispatch_path = self.root / "dispatch.json"
        self.config = {"runId": "offline-current", "googleAssociationEnabled": False,
            "googleDisabledReason": "user_disabled", "groups": [],
            "commandTimings": [{"command": "fresh_export", "durationSeconds": 1, "exitCode": 0}],
            "sourceArchive": str(self.source), "cardResult": str(self.result), "cardReceipt": str(self.card_path),
            "sourceSha256": AUDIT.sha256_file(self.source),
            "reportDateCachePath": str(self.cache), "reportDateCacheSha256": AUDIT.sha256_file(self.cache),
            "analysis": write_json(self.root / "analysis.json", {"overall": {"send": 100, "reply": 5, "rate_pct": 5},
                "latestDate": "2026-09-07", "rows": 1, "currentZero": 0, "newSincePrev": 0, "recoveredSincePrev": 0}),
            "balance": write_json(self.root / "balance.json", {}),
            "extraReceipts": [{"label": "后台导出回执", "path": write_json(self.root / "export.json", {})},
                              {"label": "文档调度回执", "path": str(self.dispatch_path)}]}
        self.inputs = {"source": self.source, "result": self.result, "association": None,
                       "google_receipt": None, "card_receipt": self.card_path,
                       "dispatch_receipt": self.dispatch_path, "delivery_dir": self.root,
                       "server_date": "2026-09-07"}
        self.resolver = AUDIT.PathResolver(self.config_path, self.root)
        self.save_receipts()

    def save_receipts(self):
        write_json(self.card_path, self.card)
        write_json(self.dispatch_path, self.dispatch)

    def review(self):
        self.save_receipts()
        return AUDIT.check_google_dependency(self.config, self.resolver, self.inputs)

    def test_skip_is_verified_without_claiming_google_succeeded(self):
        with patch.object(AUDIT, "load_workbook", side_effect=AssertionError("unexpected workbook read")):
            result = self.review()
        self.assertEqual(AUDIT.SKIP, result.status)
        self.assertTrue(result.details["skip_contract_verified"])
        self.assertFalse(result.details["google_artifacts_read"])
        self.assertEqual("system_source_only", result.details["unused_port_device_scope"])

    def test_missing_commands_or_inconsistent_mode_fail(self):
        for key, value in (("commandTimings", []), ("googleAssociationEnabled", True),
                           ("googleDisabledReason", "")):
            with self.subTest(key=key), patch.dict(self.config, {key: value}):
                self.assertEqual(AUDIT.FAIL, self.review().status)

    def test_google_command_and_reference_rejected_without_opening(self):
        for mutation in ({"association": "/forbidden/old.xlsx"}, {"googleSha256": "1" * 64},
                         {"extraReceipts": [{"label": "Google下载回执", "path": "/forbidden/old.json"}]},
                         {"commandTimings": [{"command": "google_download"}]}):
            with self.subTest(mutation=mutation), patch.dict(self.config, mutation):
                self.assertEqual(AUDIT.FAIL, self.review().status)

    def test_old_google_artifact_in_run_directory_rejected_without_reading(self):
        (self.root / "opensms_google_card_slots_old.xlsx").write_bytes(b"must not read")
        with patch.object(AUDIT, "sha256_file", side_effect=AssertionError("unexpected hash/read")):
            result = self.review()
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertEqual(1, len(result.details["unexpected_artifacts"]))

    def test_missing_or_wrong_device_scope_proof_rejected(self):
        self.card["verification"]["google_dependency_review"]["unused_port_device_scope"] = "system_and_google"
        self.assertEqual(AUDIT.FAIL, self.review().status)
        self.card["verification"]["google_dependency_review"] = {}
        self.assertEqual(AUDIT.FAIL, self.review().status)

    def test_old_run_receipt_or_non_boolean_disabled_proof_is_rejected(self):
        self.card["verification"]["system_kooda_source_review"]["run_id"] = "prior-round"
        self.assertEqual(AUDIT.FAIL, self.review().status)
        self.card["verification"]["system_kooda_source_review"]["run_id"] = "offline-current"
        self.card["verification"]["google_dependency_review"]["enabled"] = 0
        self.assertEqual(AUDIT.FAIL, self.review().status)

    def test_google_dates_or_dispatch_input_rejected_without_reading(self):
        self.card["process"]["date_resolution"]["field_counts"]["激活日期"]["google"] = 1
        self.assertEqual(AUDIT.FAIL, self.review().status)
        self.card["process"]["date_resolution"]["field_counts"]["激活日期"]["google"] = 0
        self.dispatch["inputs"]["association"] = {"path": "/forbidden/old.xlsx", "sha256": "1" * 64}
        self.assertEqual(AUDIT.FAIL, self.review().status)

    def test_discovery_never_searches_or_opens_google_sources(self):
        original = AUDIT.PathResolver.glob
        def guarded(resolver, pattern):
            self.assertNotIn("google", pattern.lower())
            return original(resolver, pattern)
        with patch.object(AUDIT.PathResolver, "glob", guarded):
            _, inputs = AUDIT.discover_inputs(self.config_path, self.config, self.root, [])
        self.assertIsNone(inputs["association"])
        self.assertIsNone(inputs["google_receipt"])

    def test_task_plan_keeps_source_and_business_checks_without_google_files(self):
        tasks = AUDIT.build_tasks(self.config, self.resolver, self.inputs, 1)
        names = {task.name for task in tasks}
        self.assertTrue({"Google依赖关闭", "本轮源文件", "服务器日期", "后台导出回执", "日期来源与冻结缓存",
                         "卡槽关联完整复核", "运营商最新状态回写文档", "余量与按日回复率"} <= names)
        self.assertFalse({"本轮Google关联表", "Google关联表回执"} & names)

    def test_card_and_dispatch_do_not_fail_or_warn_for_skipped_matching(self):
        self.assertEqual(AUDIT.PASS, AUDIT.check_card_receipt(self.card_path, False).status)
        result = AUDIT.check_dispatch_receipt(self.dispatch_path, self.config["sourceSha256"], None,
                                             AUDIT.sha256_file(self.result), google_enabled=False)
        self.assertEqual(AUDIT.PASS, result.status)
        self.assertEqual(AUDIT.FAIL, AUDIT.check_card_receipt(self.card_path).status)

    def test_frozen_dates_remain_bound_to_current_source_and_run(self):
        self.assertEqual(AUDIT.PASS, AUDIT.check_date_resolution(self.config, self.resolver, self.inputs).status)
        self.config["runId"] = "wrong-run"
        self.assertEqual(AUDIT.FAIL, AUDIT.check_date_resolution(self.config, self.resolver, self.inputs).status)
        self.config["runId"] = "offline-current"
        self.source.write_bytes(b"different source")
        self.assertEqual(AUDIT.FAIL, AUDIT.check_date_resolution(self.config, self.resolver, self.inputs).status)

    def test_delivery_rejects_google_manifest_entry_before_hashing(self):
        (self.root / "SHA256SUMS.txt").write_text("0" * 64 + "  receipts/old_google.xlsx\n")
        (self.root / "FILES.tsv").write_text("path\n")
        with patch.object(AUDIT, "sha256_file", side_effect=AssertionError("Google hash/read")):
            result = AUDIT.check_delivery(self.root, google_disabled=True)
        self.assertEqual(AUDIT.FAIL, result.status)
        self.assertIn("未访问", result.details["checksum_failures"]["receipts/old_google.xlsx"])

    def test_workbook_requires_system_sheets_and_rejects_old_google_sheets(self):
        book = Workbook()
        book.active.title = "sheet1"
        for name in AUDIT.REQUIRED_RESULT_SHEETS - AUDIT.GOOGLE_ONLY_RESULT_SHEETS - {"sheet1"}:
            book.create_sheet(name)
        book.save(self.result)
        current = AUDIT.inspect_result_workbook(self.result, None, False)
        self.assertEqual([], current["missing_sheets"])
        historical = AUDIT.inspect_result_workbook(self.result, None)
        self.assertEqual(AUDIT.GOOGLE_ONLY_RESULT_SHEETS, set(historical["missing_sheets"]))
        book.create_sheet("所有谷歌有-系统没有")
        book.active.append(["Google電信商/價格", "OPR匹配预警"])
        book.save(self.result)
        book.close()
        current = AUDIT.inspect_result_workbook(self.result, None, False)
        self.assertTrue(any("关联工作表" in error for error in current["order_errors"]))
        self.assertTrue(any("关联列" in error for error in current["order_errors"]))

    def test_full_report_needs_no_google_hash_and_discloses_scope_with_three_links(self):
        self.config.update({"taskAOnly": True, "taskSelection": "A", "taskFStageStatus": "skipped",
            "taskFSkippedReason": "task_a_only", "offlineRecoveryEnabled": False,
            "runTime": "offline", "taskId": "offline", "taskDuration": "1s", "exportWorkflow": "offline fixture"})
        write_json(self.config_path, self.config)
        output = self.root / "report.md"
        with patch.object(sys, "argv", ["report", str(self.config_path), str(output)]), \
             contextlib.redirect_stdout(io.StringIO()):
            REPORT.main()
        rendered = output.read_text()
        for text in ("Google依赖：已关闭（user_disabled）", "Google卡槽匹配未执行", "system_source_only",
                     "系统主表完全没有的设备", "不推断已激活", "整行浅红", "Chatr空激活日期允许保留"):
            self.assertIn(text, rendered)
        self.assertNotIn("Google关联表SHA-256", rendered)
        self.assertNotIn("空白字段回退到本轮Google", rendered)
        self.assertEqual(3, sum(line.startswith("[打开") for line in rendered.splitlines()))


class RealNoGoogleDocumentAuditTests(unittest.TestCase):
    def test_generated_and_independently_verified_fixture_passes_current_audits(self):
        fixture_path = Path(__file__).with_name("test_opensms_document_no_google.py")
        spec = importlib.util.spec_from_file_location("no_google_document_audit_fixture", fixture_path)
        fixture_module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(fixture_module)
        fixture = fixture_module.NoGoogleDocumentTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        # The live export lists the newest statistics day first. Adapt this
        # shared synthetic fixture to that contract before its first generation.
        source_book = load_workbook(fixture.raw)
        source_sheet = source_book.active
        for row_number, values in enumerate(source_sheet.iter_rows(values_only=True), start=1):
            ordered = list(values[:13]) + [value for start in (28, 23, 18, 13) for value in values[start:start + 5]]
            for column, value in enumerate(ordered, start=1):
                source_sheet.cell(row_number, column).value = value
        source_book.save(fixture.raw)
        source_book.close()
        payload = fixture.build()
        receipt = fixture.root / "card_audit_fixture.json"
        write_json(receipt, payload)
        config = {"runId": "offline-no-google", "googleAssociationEnabled": False,
                  "googleDisabledReason": "user_disabled", "reportDateCachePath": str(fixture.cache),
                  "reportDateCacheSha256": AUDIT.sha256_file(fixture.cache)}
        resolver = AUDIT.PathResolver(fixture.root / "config.json", fixture.root)
        inputs = {"card_receipt": receipt, "source": fixture.raw, "association": None}
        checks = [AUDIT.check_card_receipt(receipt, False),
                  AUDIT.check_date_resolution(config, resolver, inputs),
                  AUDIT.check_xlsx("result", fixture.output, inspect_result=True,
                                   server_date="2026-09-06", google_enabled=False)]
        self.assertTrue(all(check.status == AUDIT.PASS for check in checks),
                        [(check.name, check.message) for check in checks])


if __name__ == "__main__":
    unittest.main()
