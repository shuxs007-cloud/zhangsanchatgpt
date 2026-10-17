"""Frozen per-field date precedence and no-same-run F mutation; offline only."""
from datetime import datetime
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import unittest
from unittest import mock

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "services/opensms_document_api/pipeline"


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


SCHEMA = module("date_priority_fixture", ROOT / "tests/test_opensms_export_schema_compat.py")
DOC = module("date_priority_generator", PIPELINE / "process_card_slots.py")
VERIFY = module("date_priority_verify", PIPELINE / "opensms_verify_card_slot_fast.py")


class ReportDatePriorityTests(unittest.TestCase):
    def setUp(self):
        self.fixture = SCHEMA.ExportSchemaCompatibilityTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.fixture.source(modern=True)
        self.base = self.fixture.fixture

    def cache(self, records, **overrides):
        payload = json.loads(self.base.cache.read_text())
        payload.update(sourceExists=bool(records), recordCount=len(records), records=records)
        payload.update(overrides)
        self.base.cache.write_text(json.dumps(payload))
        self.base.env["OPENSMS_REPORT_DATE_CACHE_SHA256"] = hashlib.sha256(self.base.cache.read_bytes()).hexdigest()

    def source_dates(self, dates):
        book = load_workbook(self.base.raw)
        sheet = book.active
        columns = {cell.value: cell.column for cell in sheet[1]}
        for row, activation, cutoff in dates:
            sheet.cell(row, columns["激活日期"]).value = activation
            sheet.cell(row, columns["截止日"]).value = cutoff
        book.save(self.base.raw)
        book.close()

    def read_dates(self):
        book = load_workbook(self.base.output)
        try:
            sheet = book["sheet1"]
            columns = {cell.value: cell.column for cell in sheet[1]}
            self.assertNotIn("系统导出激活日期", columns)
            self.assertNotIn("系统导出截止日", columns)
            return [[sheet.cell(row, columns[name]).value for name in ("激活日期", "截止日", "换号状态")]
                    for row in range(2, 5)]
        finally:
            book.close()

    def test_per_field_export_then_cache_or_google_and_export_beats_inventory(self):
        self.cache([
            {"queryIccid": "8912230200081836799", "activationDate": "2026-09-06", "cutoffDate": "2026-10-05"},
            {"queryIccid": "8912230200081836798", "activationDate": "2026-09-05", "cutoffDate": "2026-10-04"},
        ])
        self.source_dates([(2, "2026-09-03", None), (3, None, None), (4, "2026-09-04", None)])
        receipt = self.base.build()
        dates = self.read_dates()
        self.assertEqual(dates[0][:2], [datetime(2026, 9, 3), datetime(2026, 10, 5)])
        self.assertEqual(dates[1][:2], [datetime(2026, 9, 5), datetime(2026, 10, 4)])
        self.assertNotEqual(dates[1][2], "未换号")
        self.assertEqual(dates[2][:2], [datetime(2026, 9, 4), datetime(2026, 9, 20)])
        counts = receipt["process"]["date_resolution"]["field_counts"]
        self.assertEqual(counts["激活日期"], {"system_export": 2, "system_inventory_legacy": 0, "local_cache": 1, "google": 0, "missing": 0})
        self.assertEqual(counts["截止日"], {"system_export": 0, "system_inventory_legacy": 0, "local_cache": 2, "google": 1, "missing": 0})
        self.assertTrue(receipt["verification"]["date_resolution_review"]["matches"])
        self.assertEqual(receipt["verification"]["system_kooda_source_review"]["missing_either_date_rows"], 0)
        self.assertEqual(receipt["verification"]["task_f_missing_activation_red_review"]["expected_rows"], 0)
        self.assertEqual(receipt["process"]["source_system_comparison"]["activation_date_conflict_rows"], 1)

    def test_chatr_system_dates_do_not_require_google_match(self):
        book = load_workbook(self.base.google)
        book["四號機房"].delete_rows(2)
        book.save(self.base.google)
        book.close()
        self.base.env["OPENSMS_ASSOC_SHA256"] = hashlib.sha256(self.base.google.read_bytes()).hexdigest()
        receipt = self.base.build()
        self.assertEqual(self.read_dates()[2][:2], [datetime(2030, 1, 1), datetime(2030, 2, 1)])
        self.assertEqual(receipt["process"]["google_matched_rows"], 0)
        self.assertEqual(receipt["process"]["matched_rows"], 3)

    def test_missing_field_scope_is_after_frozen_fallback_not_inventory(self):
        self.source_dates([(2, "2026-09-01", None)])
        receipt = self.base.build()
        review = receipt["verification"]["system_kooda_source_review"]
        self.assertEqual((review["missing_activation_rows"], review["missing_cutoff_rows"], review["missing_either_date_rows"]), (0, 1, 1))
        scopes = receipt["process"]["date_resolution"]["missing_date_scopes"]
        self.assertEqual(len(scopes), 1)
        self.assertFalse(scopes[0]["activation_missing"])
        self.assertTrue(scopes[0]["cutoff_missing"])
        self.assertEqual(receipt["process"]["date_resolution"]["field_counts"]["激活日期"]["creation_date_fallback"], 1)

    def test_generator_never_reads_active_database_and_empty_snapshot_is_legal(self):
        with mock.patch.dict(os.environ, self.base.env), mock.patch.object(DOC.sqlite3, "connect", side_effect=AssertionError("active DB read forbidden")):
            self.assertEqual(DOC.load_report_date_cache()["record_count"], 0)
            # Simulates later background DB changes: JSON stays authoritative.
            Path(self.base.env["OPENSMS_TASK_F_DB_PATH"]).write_text("new data cannot enter frozen snapshot")
            self.assertEqual(DOC.load_report_date_cache()["index"], {})

    def test_snapshot_hash_run_identity_pair_and_duplicates_are_independently_rejected(self):
        good = {"queryIccid": "8912230200081836799", "activationDate": "2026-09-06", "cutoffDate": "2026-10-05"}
        cases = [([good], {"runId": "different-run"}), ([good], {"status": "failed"}),
                 ([good], {"recordCount": 2}), ([good], {"sourceExists": False}),
                 ([good, good], {}), ([{**good, "queryIccid": good["queryIccid"] + "F"}], {}),
                 ([{**good, "activationDate": ""}], {}), ([{**good, "cutoffDate": "2026-10-06"}], {})]
        for records, overrides in cases:
            self.cache(records, **overrides)
            with self.subTest(records=records, overrides=overrides):
                with mock.patch.dict(os.environ, self.base.env), self.assertRaises(RuntimeError):
                    DOC.load_report_date_cache()
                with self.assertRaises(RuntimeError):
                    VERIFY.load_frozen_date_snapshot(self.base.cache, self.base.env["OPENSMS_REPORT_DATE_CACHE_SHA256"], self.base.env["OPENSMS_RUN_ID"])
        self.cache([good])
        self.base.cache.write_text(self.base.cache.read_text() + " ")
        with mock.patch.dict(os.environ, self.base.env), self.assertRaisesRegex(RuntimeError, "SHA256"):
            DOC.load_report_date_cache()
        with self.assertRaisesRegex(RuntimeError, "SHA256"):
            VERIFY.load_frozen_date_snapshot(self.base.cache, self.base.env["OPENSMS_REPORT_DATE_CACHE_SHA256"], self.base.env["OPENSMS_RUN_ID"])

    def test_independent_review_rejects_forged_field_counts(self):
        receipt = self.base.build()
        receipt["process"]["date_resolution"]["field_counts"]["激活日期"]["local_cache"] = 999
        run = subprocess.run([sys.executable, str(PIPELINE / "opensms_verify_card_slot_fast.py"),
            str(self.base.output), str(self.base.raw)], env={**self.base.env,
                "PROCESS_JSON": json.dumps(receipt["process"]), "ENRICH_JSON": json.dumps(receipt["enrichment"]),
                "EXPORT4DAY_SERVER_DATE": "2026-09-06"}, capture_output=True, text=True, timeout=30)
        self.assertEqual(run.returncode, 0, run.stderr)
        result = json.loads(run.stdout)["date_resolution_review"]
        self.assertFalse(result["matches"])
        self.assertIn("field_counts", result["receipt_mismatch_fields"])


if __name__ == "__main__":
    unittest.main()
