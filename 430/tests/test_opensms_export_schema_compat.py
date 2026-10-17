"""Regression for old 58-column and new 61-column export schemas, offline only."""

import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import unittest

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "services/opensms_document_api/pipeline"
SPEC = importlib.util.spec_from_file_location("export_schema_base_fixture", ROOT / "tests/test_opensms_system_kooda_integration.py")
BASE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BASE)
VERIFY_SPEC = importlib.util.spec_from_file_location("export_schema_verifier", PIPELINE / "opensms_verify_card_slot_fast.py")
VERIFY = importlib.util.module_from_spec(VERIFY_SPEC)
VERIFY_SPEC.loader.exec_module(VERIFY)


class ExportSchemaCompatibilityTests(unittest.TestCase):
    def test_non_kooda_empty_iccid_roundtrip_and_identity_guards(self):
        headers = ["电话", "opr", "卡槽", "激活日期", "截止日", "ICCID"]
        def review(source, actual):
            base = ["14385550123", "Chatr", "130_1", None, None]
            return VERIFY.review_export_schema_evidence(headers, [base + [source]], headers, [base + [actual]])["matches"]
        for source, actual in [("", None), (None, ""), (None, None), ("", ""), ("00123F", "00123F")]:
            self.assertTrue(review(source, actual))
        for source, actual in [(" ", None), ("00123F", None), (None, "00123F"), ("123", 123), ("00123F", "00123"), (True, 1)]:
            self.assertFalse(review(source, actual))

    def setUp(self):
        self.fixture = BASE.SystemKoodaIntegrationTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)

    def source(self, modern, iccid_override=None):
        if modern:
            headers = ["电话", "ICCID", "运营商", "opr", "卡槽", "状态", "激活日期", "截止日",
                       "创建时间", "送达时间", "接收时间", "10日发送窗口数", "10日回复窗口数",
                       "10日回复率", "10日总发送量", "10日总回复量"]
        else:
            headers = ["电话", "运营商", "opr", "卡槽", "创建时间", "送达时间", "接收时间", "状态",
                       "10日发送窗口数", "10日回复窗口数", "10日回复率", "10日总发送量", "10日总回复量"]
        days = [f"09-0{value}" for value in range(1, 7)] + ["08-29", "08-30", "08-31"]
        for day in days:
            headers.extend(day + suffix for suffix in ("发送窗口数", "回复窗口数", "回复率", "总发送量", "总回复量"))
        self.assertEqual(len(headers), 61 if modern else 58)
        book = Workbook()
        sheet = book.active
        sheet.title = "sheet1"
        sheet.append(headers)
        for offset, (opr, activation, cutoff, iccid) in enumerate([
            ("Kooda", "2026-09-01", "2026-12-15", iccid_override or "8912230200081836799F"),
            ("Telus1", "", "", "8912230200081836798f"),
            ("Chatr", "2030-01-01", "2030-02-01", "KEEP-SOURCE-CHATR-ICCID"),
        ]):
            record = {"电话": str(14385550123 + offset), "运营商": "A1" if opr == "Chatr" else "B1",
                      "opr": opr, "卡槽": f"130_{offset + 1}", "创建时间": "2026-09-06 00:00:00",
                      "状态": "注册失败" if offset == 1 else "上架", "ICCID": iccid,
                      "激活日期": activation, "截止日": cutoff,
                      "10日发送窗口数": 900, "10日回复窗口数": 90, "10日回复率": .1,
                      "10日总发送量": 900, "10日总回复量": 90}
            for day in days:
                record.update({day + key: value for key, value in zip(
                    ("发送窗口数", "回复窗口数", "回复率", "总发送量", "总回复量"), (100, 10, .1, 100, 10))})
            sheet.append([record.get(name) for name in headers])
        book.save(self.fixture.raw)
        book.close()

    def test_old_58_columns_keep_full_independent_verification(self):
        self.source(modern=False)
        receipt = self.fixture.build()
        self.assertEqual(receipt["process"]["source_layout"]["source_column_count"], 58)
        self.assertFalse(receipt["process"]["source_layout"]["existing_iccid"])
        self.assertTrue(receipt["verification"]["source_schema_review"]["matches"])
        self.assertTrue(all(value is True for value in receipt["checks"].values()))

    def test_shared_string_empty_dates_may_reopen_blank_but_nonempty_values_cannot_change(self):
        source_headers = ["电话", "opr", "卡槽", "激活日期", "截止日"]
        final_headers = ["电话", "opr", "卡槽", "激活日期", "截止日"]
        original = [("14385550123", "Kooda", "130_1", "", "")]
        actual = [("14385550123", "Kooda", "130_1", None, None)]
        review = VERIFY.review_export_schema_evidence(source_headers, original, final_headers, actual)
        self.assertTrue(review["matches"])
        self.assertFalse(VERIFY.review_export_schema_evidence(
            source_headers, original, ["电话", "opr", "卡槽", "系统导出激活日期", "截止日"], actual
        )["matches"])

    def test_new_61_columns_preserve_evidence_and_unique_final_iccid(self):
        self.source(modern=True)
        receipt = self.fixture.build()
        self.assertEqual(receipt["process"]["source_layout"]["source_column_count"], 61)
        self.assertTrue(receipt["process"]["source_layout"]["existing_iccid"])
        self.assertEqual(receipt["process"]["source_system_comparison"]["iccid_conflict_rows"], 0)
        self.assertEqual(receipt["process"]["source_system_comparison"]["iccid_compared_rows"], 2)
        self.assertTrue(receipt["verification"]["source_schema_review"]["matches"])
        book = load_workbook(self.fixture.output)
        try:
            sheet = book["sheet1"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers[:4], ["电话", "运营商", "过期状态", "换号状态"])
            self.assertEqual(headers[-1], "ICCID")
            for name in ("ICCID", "激活日期", "截止日"):
                self.assertEqual(headers.count(name), 1)
            for name in ("系统导出激活日期", "系统导出截止日"):
                self.assertNotIn(name, headers)
            index = {name: position + 1 for position, name in enumerate(headers)}
            self.assertEqual(sheet.cell(4, index["ICCID"]).value, "KEEP-SOURCE-CHATR-ICCID")
            self.assertEqual(sheet.cell(4, index["激活日期"]).value.date().isoformat(), "2030-01-01")
            self.assertEqual(sheet.cell(4, index["截止日"]).value.date().isoformat(), "2030-02-01")
        finally:
            book.close()

    def test_source_system_iccid_key_conflict_fails_before_workbook_save(self):
        self.source(modern=True, iccid_override="8912230200081836790F")
        completed = subprocess.run(
            ["bash", str(PIPELINE / "build_card_slot_report.sh"), str(self.fixture.raw),
             str(self.fixture.google), str(self.fixture.output.parent), "2026-09-06"],
            env=self.fixture.env, text=True, capture_output=True, timeout=45,
        )
        self.assertNotEqual(completed.returncode, 0)
        self.assertIn("ICCID标准键冲突", completed.stderr)
        self.assertFalse(self.fixture.output.exists())

    def test_independent_verifier_rejects_tampered_source_evidence(self):
        self.source(modern=True)
        receipt = self.fixture.build()
        book = load_workbook(self.fixture.output)
        sheet = book["sheet1"]
        index = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(4, index["截止日"]).value = "2030-02-02"
        book.save(self.fixture.output)
        book.close()
        completed = subprocess.run(
            [sys.executable, str(PIPELINE / "opensms_verify_card_slot_fast.py"),
             str(self.fixture.output), str(self.fixture.raw)],
            env={**self.fixture.env, "PROCESS_JSON": json.dumps(receipt["process"]),
                 "ENRICH_JSON": json.dumps(receipt["enrichment"]), "EXPORT4DAY_SERVER_DATE": "2026-09-06"},
            text=True, capture_output=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        verification = json.loads(completed.stdout)
        self.assertFalse(verification["date_resolution_review"]["matches"])
        self.assertFalse(verification["system_kooda_source_review"]["matches"])


if __name__ == "__main__":
    unittest.main()
