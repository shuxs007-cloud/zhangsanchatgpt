"""Synthetic F/system-intersection tests; no network or production DB writes."""
import contextlib
import io
import json
import os
from pathlib import Path
import sqlite3
import sys
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'services/opensms_document_api/pipeline'))
sys.path.insert(0, str(ROOT / 'tests'))
from opensms_report_date_cache import freeze
from ko_unactivated import SHEET, LEGACY_SHEET, AFTER_SHEET, HEADERS, SCOPE, SLOT_SOURCE, load_snapshot, load_from_env, select_system_records, write_sheet, review
from opensms_verify_card_slot_fast import load_system_kooda_snapshot
from opensms_chatr_google_dates import freeze as freeze_google
import opensms_parallel_audit as audit
import test_opensms_document_no_google as fixture_module


class KoUnactivatedTests(unittest.TestCase):
    def setUp(self):
        self.f = fixture_module.NoGoogleDocumentTests()
        self.f.setUp()
        self.addCleanup(self.f.doCleanups)
        self.db = self.f.root / 'f.sqlite3'
        with sqlite3.connect(self.db) as conn:
            conn.execute('CREATE TABLE kooda_activation (query_iccid TEXT, activation_date TEXT, cutoff_date TEXT, '
                         'lookup_status TEXT, queried_at TEXT, source_phone TEXT, source_slot TEXT, source_opr TEXT, '
                         'source_carrier_name TEXT, purchase_date TEXT, act_request_date TEXT, sku TEXT, data_source TEXT)')
            for key, status, activation, cutoff in (
                ('8912230200081836701F', 'activation_missing', None, None),
                ('00012345678901234567', 'activation_missing', None, None),
                ('8912230200081836703', 'activation_missing', None, None),
                ('8912230200081836799', 'found', '2026-09-01', '2026-09-30'),
                ('8912230200081836798', 'not_found', None, None),
                ('8912230200081836797', 'error', None, None),
            ):
                conn.execute('INSERT INTO kooda_activation VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (key, activation, cutoff, status, '2026-09-05T01:02:03Z', '01234567890', 'r130_01.01',
                     'Kooda', 'B1', '', '', '=1+1', 'grabandgo_incremental_lookup'))
        self.refreeze()
        self.f.env['OPENSMS_KO_UNACTIVATED_REQUIRED'] = '1'

    def refreeze(self):
        self.f.cache.unlink(missing_ok=True)
        result = freeze(self.db, self.f.cache, self.f.env['OPENSMS_RUN_ID'])
        self.f.env['OPENSMS_REPORT_DATE_CACHE_SHA256'] = result['sha256']
        return result

    def config(self, receipt):
        path = self.f.root / 'ko-receipt.json'
        path.write_text(json.dumps(receipt))
        return ({'runId': self.f.env['OPENSMS_RUN_ID'], 'reportDateCachePath': str(self.f.cache),
                 'reportDateCacheSha256': self.f.digest(self.f.cache), 'koUnactivatedSheetEnabled': True,
                 'koUnactivatedSystemPath': str(self.f.system), 'koUnactivatedSystemSha256': self.f.digest(self.f.system),
                 'koUnactivatedSheet': receipt['process']['ko_unactivated_sheet']},
                {'card_receipt': path, 'result': self.f.output})

    def enable_google_difference(self):
        google = self.f.root / 'ko-order-google.xlsx'
        snapshot = self.f.root / 'ko-order-google.json'
        book = Workbook()
        book.remove(book.active)
        for name in ('二號機房', '三號機房', '四號機房'):
            sheet = book.create_sheet(name)
            sheet.append(['卡槽', '電信商/價格', '激活日期', '預計過期日', 'ICCID', '上卡日期', '備註'])
        book.worksheets[0].append(['130_8', 'Kooda', None, None, None, None, '离线排序测试'])
        book.save(google)
        book.close()
        freeze_google(google, self.f.digest(google), snapshot, self.f.env['OPENSMS_RUN_ID'])
        self.f.env.update(OPENSMS_CHATR_GOOGLE_DATES_PATH=str(snapshot),
                          OPENSMS_CHATR_GOOGLE_DATES_SHA256=self.f.digest(snapshot),
                          OPENSMS_GOOGLE_SYSTEM_MISSING_ENABLED='1')

    def test_google_ko_order_and_tampering_checked_independently(self):
        self.enable_google_difference()
        receipt = self.f.build()
        book = load_workbook(self.f.output)
        start = book.sheetnames.index('系统无-谷歌表有')
        self.assertEqual(book.sheetnames[start:start + 3], ['系统无-谷歌表有', AFTER_SHEET, SHEET])
        self.assertTrue(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
        self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.PASS)
        receipt_path = self.config(receipt)[1]['card_receipt']
        self.assertEqual(audit.check_card_receipt(receipt_path, google_enabled=False).status, audit.PASS)
        self.assertEqual(audit.inspect_result_workbook(self.f.output, '2026-09-06', False, True)['order_errors'], [])
        book.move_sheet(SHEET, 1)
        book.save(self.f.output)
        book.close()
        self.assertFalse(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
        self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.FAIL)
        self.assertIn('ko系统有-未激活未紧邻谷歌有-系统无右侧',
                      audit.inspect_result_workbook(self.f.output, '2026-09-06', False, True)['order_errors'])

    def test_empty_ko_sheet_keeps_requested_position(self):
        self.enable_google_difference()
        with sqlite3.connect(self.db) as conn:
            conn.execute("UPDATE kooda_activation SET lookup_status='not_found' WHERE lookup_status='activation_missing'")
        self.refreeze()
        receipt = self.f.build()
        book = load_workbook(self.f.output)
        self.addCleanup(book.close)
        self.assertEqual(book[SHEET].max_row, 1)
        self.assertEqual(book.sheetnames.index(SHEET), book.sheetnames.index(AFTER_SHEET) + 1)
        self.assertTrue(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])

    def test_only_f_unactivated_iccids_in_system_and_no_unknown_statuses(self):
        receipt = self.f.build()
        proof = receipt['process']['ko_unactivated_sheet']
        self.assertEqual(proof['scope'], SCOPE)
        self.assertEqual(proof['row_count'], 1)
        self.assertEqual(proof['f_snapshot_row_count'], 3)
        self.assertEqual(proof['excluded_not_in_system_count'], 2)
        self.assertTrue(receipt['checks']['ko系统有-未激活独立核验'])
        self.assertTrue(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
        book = load_workbook(self.f.output)
        self.addCleanup(book.close)
        sheet = book[SHEET]
        self.assertEqual(tuple(cell.value for cell in sheet[1]), HEADERS)
        self.assertEqual([row[0] for row in sheet.iter_rows(min_row=2, values_only=True)],
                         ['8912230200081836703'])
        self.assertEqual(sheet['L2'].value, '=1+1')
        self.assertEqual(sheet['L2'].data_type, 's')
        self.assertEqual(sheet['F2'].value, '01234567890')
        self.assertEqual(SHEET, 'ko系统有-未激活')
        self.assertNotIn(LEGACY_SHEET, book.sheetnames)
        self.assertEqual(sheet['G1'].value, '卡槽')
        self.assertEqual(sheet['G2'].value, '130_3')
        self.assertEqual(proof['slot_source'], SLOT_SOURCE)
        self.assertEqual((proof['current_slot_filled_count'], proof['current_slot_empty_count']), (1, 0))
        self.assertIsNone(sheet['C2'].value)
        self.assertIsNone(sheet['D2'].value)
        self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.PASS)

    def test_current_slots_preserve_text_deduplicate_and_list_all(self):
        snapshot = self.system_snapshot()
        snapshot['raw_rows'] = [{'cardStore': '00012345678901234567f', 'slotIdentif': slot}
            for slot in ('r0130_01.01', '00130_09', 'r0130_01.01', None, '', '  ')]
        snapshot['total_rows'] = len(snapshot['raw_rows'])
        frozen = self.evidence()
        original = json.dumps(frozen, sort_keys=True)
        selected = select_system_records(frozen, snapshot)
        self.assertEqual(selected['rows'][0]['currentSystemSlot'], '00130_09\nr0130_01.01')
        self.assertEqual(selected['proof']['multiple_current_slot_count'], 1)
        self.assertEqual(selected['proof']['current_slot_filled_count'], 1)
        self.assertEqual(json.dumps(frozen, sort_keys=True), original)

    def test_empty_current_slot_never_uses_historical_slot_or_other_fields(self):
        snapshot = self.system_snapshot()
        snapshot['raw_rows'] = [{'cardStore': '8912230200081836703F', 'slotIdentif': None,
                                'slot': 'wrong_alias', 'cardSlot': 'wrong_alias'}]
        snapshot['total_rows'] = 1
        selected = select_system_records(self.evidence(), snapshot)
        self.assertEqual(selected['rows'][0]['slot'], 'r130_01.01')
        self.assertEqual(selected['rows'][0]['currentSystemSlot'], '')
        self.assertEqual(selected['proof']['current_slot_empty_count'], 1)
        self.assertEqual(selected['proof']['current_slot_filled_count'], 0)

    def test_invalid_or_overlong_current_slot_stops_instead_of_guessing(self):
        for slot in (130, True, {}, 'x' * 32768):
            with self.subTest(slot_type=type(slot).__name__):
                snapshot = self.system_snapshot()
                snapshot['raw_rows'] = [{'cardStore': '8912230200081836703F', 'slotIdentif': slot}]
                snapshot['total_rows'] = 1
                with self.assertRaises(ValueError):
                    select_system_records(self.evidence(), snapshot)

    def test_replacing_legacy_sheet_and_summary_is_idempotent(self):
        book = Workbook()
        book.active.title = '处理汇总'
        book.active.append(('原有统计', 77))
        book.active.append((LEGACY_SHEET + '行数', 99))
        book.active.append((LEGACY_SHEET + '口径', '旧卡槽'))
        book.create_sheet(AFTER_SHEET)
        book.create_sheet(LEGACY_SHEET)
        selected = select_system_records(self.evidence(), self.system_snapshot())
        for _ in range(2):
            write_sheet(book, selected)
            self.assertNotIn(LEGACY_SHEET, book.sheetnames)
            self.assertEqual(book.sheetnames.index(SHEET), book.sheetnames.index(AFTER_SHEET) + 1)
            self.assertEqual(book['处理汇总'].max_row, 3)
            self.assertEqual(book['处理汇总']['B1'].value, 77)
            self.assertEqual(book[SHEET]['G2'].value, '130_3')
        book.close()

    def test_explicit_empty_text_is_blank_but_nonempty_text_is_not(self):
        book = Workbook()
        book.active.title = '处理汇总'
        system = self.system_snapshot()
        proof = write_sheet(book, select_system_records(self.evidence(), system))
        for address in ('C2', 'D2', 'K2'):
            book[SHEET][address] = ''
        args = (book, proof, self.f.cache, self.f.digest(self.f.cache), self.f.env['OPENSMS_RUN_ID'], True)
        self.assertTrue(review(*args, system=system)['matches'])
        book[SHEET]['C2'] = ' '
        self.assertFalse(review(*args, system=system)['matches'])
        book.close()

    def test_current_slot_and_legacy_name_tampering_fail_both_reviews(self):
        receipt = self.f.build()
        original = self.f.output.read_bytes()
        for mutation in ('historical_slot', 'old_header', 'old_sheet', 'extra_old_sheet'):
            with self.subTest(mutation=mutation):
                self.f.output.write_bytes(original)
                book = load_workbook(self.f.output)
                if mutation == 'historical_slot': book[SHEET]['G2'] = 'r130_01.01'
                elif mutation == 'old_header': book[SHEET]['G1'] = '卡槽（记录时）'
                elif mutation == 'old_sheet': book[SHEET].title = LEGACY_SHEET
                else: book.create_sheet(LEGACY_SHEET)
                book.save(self.f.output)
                book.close()
                self.assertFalse(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
                self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.FAIL)

    def system_snapshot(self, changes=None):
        if changes is not None:
            payload = json.loads(self.f.system.read_text())
            changes(payload)
            self.f.system.write_text(json.dumps(payload))
            self.f.env['OPENSMS_SYSTEM_KOODA_SHA256'] = self.f.digest(self.f.system)
        return load_system_kooda_snapshot(self.f.system, self.f.digest(self.f.system), self.f.env['OPENSMS_RUN_ID'])

    def evidence(self):
        return load_snapshot(self.f.cache, self.f.digest(self.f.cache), self.f.env['OPENSMS_RUN_ID'], True)

    def test_inventory_only_and_duplicate_iccid_keep_one_leading_zero_text_row(self):
        def add_rows(payload):
            for i, status in enumerate(('下架', '注册失败', ''), 20):
                payload['systemRows'].append({'id': str(i), 'opr': 'Kooda', 'slotIdentif': None,
                    'cardNo': str(14385550000 + i), 'cardStore': '00012345678901234567f', 'status': status})
            payload['totalRows'] = len(payload['systemRows'])
        self.system_snapshot(add_rows)
        receipt = self.f.build()
        self.assertEqual(receipt['process']['ko_unactivated_sheet']['row_count'], 2)
        book = load_workbook(self.f.output)
        self.addCleanup(book.close)
        self.assertEqual(book[SHEET]['A2'].value, '00012345678901234567')
        self.assertEqual(book[SHEET]['A2'].data_type, 's')
        self.assertTrue(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
        self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.PASS)

    def test_exact_iccid_only_not_phone_slot_or_invalid_suffix(self):
        snapshot = self.system_snapshot()
        snapshot['raw_rows'] = [{'cardStore': value, 'cardNo': '01234567890', 'slotIdentif': 'r130_01.01'}
            for value in ('8912230200081836701FF', '891223020008183670', '08912230200081836701',
                          '8.912230200081837e18', None, 8912230200081836701)]
        snapshot['total_rows'] = len(snapshot['raw_rows'])
        selected = select_system_records(self.evidence(), snapshot)
        self.assertEqual(selected['rows'], [])
        self.assertEqual(selected['proof']['excluded_not_in_system_count'], 3)

    def test_found_not_found_error_and_unknown_not_included_even_in_system(self):
        snapshot = self.system_snapshot()
        snapshot['raw_rows'] = [{'cardStore': '89122302000818367' + suffix} for suffix in ('97', '98', '99', '96')]
        snapshot['total_rows'] = 4
        self.assertEqual(select_system_records(self.evidence(), snapshot)['rows'], [])

    def test_no_intersection_keeps_empty_sheet_and_no_fallback(self):
        def remove_match(payload):
            payload['systemRows'][0]['cardStore'] = '8912230200081836788F'
        self.system_snapshot(remove_match)
        receipt = self.f.build()
        proof = receipt['process']['ko_unactivated_sheet']
        self.assertEqual((proof['row_count'], proof['f_snapshot_row_count'], proof['excluded_not_in_system_count']), (0, 3, 3))
        book = load_workbook(self.f.output)
        self.addCleanup(book.close)
        self.assertEqual(book[SHEET].max_row, 1)
        self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.PASS)

    def test_system_binding_and_unfiltered_write_fail_closed(self):
        evidence = self.evidence()
        for system in (None, {}, {**self.system_snapshot(), 'run_id': 'old'},
                       {**self.system_snapshot(), 'sha256': ''}):
            with self.subTest(system=system):
                with self.assertRaises(ValueError):
                    select_system_records(evidence, system)
        with self.assertRaises(ValueError):
            write_sheet(Workbook(), evidence)
        receipt = self.f.build()
        config, inputs = self.config(receipt)
        for field, value in (('koUnactivatedSystemSha256', '0' * 64),
                             ('koUnactivatedSystemPath', str(self.f.root / 'missing.json'))):
            with self.subTest(field=field):
                self.assertEqual(audit.check_ko_unactivated({**config, field: value}, inputs).status, audit.FAIL)

    def test_changed_system_membership_is_detected_by_both_reviews(self):
        receipt = self.f.build()
        self.system_snapshot(lambda payload: payload['systemRows'][0].update(cardStore='8912230200081836788F'))
        self.assertFalse(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
        self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.FAIL)

    def test_f_updates_only_next_snapshot_not_published_workbook(self):
        receipt = self.f.build()
        frozen, output = self.f.cache.read_bytes(), self.f.output.read_bytes()
        with sqlite3.connect(self.db) as conn:
            conn.execute("UPDATE kooda_activation SET lookup_status='found',activation_date='2026-09-01',"
                         "cutoff_date='2026-09-30' WHERE lookup_status='activation_missing'")
        self.assertEqual(self.f.cache.read_bytes(), frozen)
        self.assertEqual(self.f.output.read_bytes(), output)
        self.assertTrue(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
        self.f.output.unlink()
        self.refreeze()
        new = self.f.build()
        self.assertEqual(new['process']['ko_unactivated_sheet']['row_count'], 0)
        book = load_workbook(self.f.output)
        self.addCleanup(book.close)
        self.assertEqual(book[SHEET].max_row, 1)

    def test_sheet_does_not_change_main_or_other_sheets(self):
        self.f.build()
        book = load_workbook(self.f.output, data_only=True)
        before = {name: list(book[name].values) for name in book.sheetnames if name not in (SHEET, '处理汇总')}
        book.close()
        payload = json.loads(self.f.cache.read_text())
        del payload['koUnactivated']
        self.f.cache.write_text(json.dumps(payload))
        self.f.env['OPENSMS_REPORT_DATE_CACHE_SHA256'] = self.f.digest(self.f.cache)
        self.f.env.pop('OPENSMS_KO_UNACTIVATED_REQUIRED')
        self.f.output.unlink()
        self.f.build()
        book = load_workbook(self.f.output, data_only=True)
        self.addCleanup(book.close)
        self.assertNotIn(SHEET, book.sheetnames)
        for name, rows in before.items():
            self.assertEqual(list(book[name].values), rows, name)

    def test_tampering_fails_initial_and_final_independent_review(self):
        receipt = self.f.build()
        original = self.f.output.read_bytes()
        for kind in ('missing', 'wrong_iccid', 'fill', 'formula', 'extra', 'date'):
            with self.subTest(kind=kind):
                self.f.output.write_bytes(original)
                book = load_workbook(self.f.output)
                if kind == 'missing': del book[SHEET]
                elif kind == 'wrong_iccid': book[SHEET]['A2'] = 12345678901234567
                elif kind == 'fill': book[SHEET]['B2'].fill = PatternFill()
                elif kind == 'formula': book[SHEET]['L2'] = '=1+1'
                elif kind == 'date': book[SHEET]['C2'] = '2026-09-01'
                else: book[SHEET].append(['extra'])
                book.save(self.f.output)
                book.close()
                self.assertFalse(self.f.verify(receipt)['ko_unactivated_sheet_review']['matches'])
                self.assertEqual(audit.check_ko_unactivated(*self.config(receipt)).status, audit.FAIL)

    def test_required_snapshot_hash_scope_run_and_status_gates(self):
        payload = json.loads(self.f.cache.read_text())
        for mutation in ('missing_section', 'wrong_run', 'bad_status', 'dates', 'duplicate'):
            with self.subTest(mutation=mutation):
                altered = json.loads(json.dumps(payload))
                if mutation == 'missing_section': del altered['koUnactivated']
                elif mutation == 'wrong_run': altered['runId'] = 'old'
                elif mutation == 'bad_status': altered['koUnactivated']['records'][0]['status'] = 'not_found'
                elif mutation == 'dates': altered['koUnactivated']['records'][0]['activationDate'] = '2026-09-01'
                else: altered['koUnactivated']['records'][1] = altered['koUnactivated']['records'][0]
                self.f.cache.write_text(json.dumps(altered))
                with self.assertRaises(ValueError):
                    load_snapshot(self.f.cache, self.f.digest(self.f.cache), self.f.env['OPENSMS_RUN_ID'], True)
        with self.assertRaisesRegex(ValueError, 'SHA256'):
            load_snapshot(self.f.cache, '0' * 64, self.f.env['OPENSMS_RUN_ID'], True)

    def test_carrier_refresh_preserves_sheet(self):
        import test_opensms_refresh_workbook_carriers_fast as refresh
        self.enable_google_difference()
        receipt = self.f.build()
        args = refresh.FastCarrierRefreshTests().fresh_args(self.f.root, self.f.output, 'ko', [
            {'phone': '14385550001', 'liveCarrierName': '离线测试', 'route': 'offline', 'expectedTarget': '离线测试'}])
        result = refresh.MODULE.run(args)
        self.assertEqual(result['changed'], 1)
        config, inputs = self.config(receipt)
        inputs['result'] = Path(result['outputPath'])
        self.assertEqual(audit.check_ko_unactivated(config, inputs).status, audit.PASS)
        book = load_workbook(inputs['result'])
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames.index(SHEET), book.sheetnames.index(AFTER_SHEET) + 1)

    def test_runner_gate_rejects_missing_old_or_failed_evidence(self):
        receipt = self.f.build()
        import test_opensms_acquire_reliability as runner_fixture
        fixture = runner_fixture.AcquisitionReliabilityTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        runner = fixture.runner
        runner.state['koUnactivatedSheetEnabled'] = True
        proof = receipt['process']['ko_unactivated_sheet']
        runner.run_id = proof['run_id']
        runner.state['reportDateCache'] = {'path': proof['snapshot_path'], 'sha256': proof['snapshot_sha256']}
        runner.state.setdefault('stages', {})['acquire'] = {
            'processingSystemKooda': proof['system_path'], 'systemKoodaSha256': proof['system_sha256']}
        self.assertEqual(runner.verify_ko_unactivated_document(receipt), proof)
        for changes in ({'scope': 'all_task_f_activation_missing'}, {'system_sha256': '0' * 64},
                        {'system_path': '/old/system.json'}, {'sheet': LEGACY_SHEET}, {'slot_source': 'task_f_history'}):
            forged = json.loads(json.dumps(receipt))
            forged['process']['ko_unactivated_sheet'].update(changes)
            forged['verification']['ko_unactivated_sheet_review'].update(changes)
            with self.subTest(changes=changes), self.assertRaises(runner_fixture.RUNNER.PatrolError):
                runner.verify_ko_unactivated_document(forged)
        with self.assertRaises(runner_fixture.RUNNER.PatrolError):
            runner.verify_ko_unactivated_document({})
        receipt['verification']['ko_unactivated_sheet_review']['matches'] = False
        with self.assertRaises(runner_fixture.RUNNER.PatrolError):
            runner.verify_ko_unactivated_document(receipt)

    def test_inconsistent_database_dates_fail_before_publication(self):
        with sqlite3.connect(self.db) as conn:
            conn.execute("UPDATE kooda_activation SET activation_date='2026-09-01' WHERE lookup_status='activation_missing'")
        path = self.f.root / 'invalid.json'
        with self.assertRaisesRegex(ValueError, '状态.*日期冲突'):
            freeze(self.db, path, 'bad')
        self.assertFalse(path.exists())

    def test_report_counts_without_expanding_phone_lists(self):
        import test_opensms_no_google_audit_report as report_fixture
        fixture = report_fixture.NoGoogleAuditReportTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        fixture.card['process']['ko_unactivated_sheet'] = {'row_count': 1, 'f_snapshot_row_count': 3,
            'excluded_not_in_system_count': 2, 'frozen_at': '2026-09-06T00:00:00Z',
            'current_slot_filled_count': 1, 'current_slot_empty_count': 0, 'multiple_current_slot_count': 0}
        fixture.save_receipts()
        fixture.config.update({'koUnactivatedSheetEnabled': True, 'taskAOnly': True, 'taskSelection': 'A',
            'taskFStageStatus': 'skipped', 'taskFSkippedReason': 'task_a_only', 'offlineRecoveryEnabled': False,
            'runTime': 'offline', 'taskId': 'offline', 'taskDuration': '1s', 'exportWorkflow': 'fixture'})
        report_fixture.write_json(fixture.config_path, fixture.config)
        output = fixture.root / 'ko-report.md'
        with mock.patch.object(sys, 'argv', ['report', str(fixture.config_path), str(output)]), contextlib.redirect_stdout(io.StringIO()):
            report_fixture.REPORT.main()
        content = output.read_text()
        for text in ('ko系统有-未激活', '任务F未激活且本轮系统库存存在的记录 1 条', 'F冻结未激活总数 3 条',
                     '排除系统不存在 2 条', '按规范ICCID取交集并去重', '排除未找到/查询失败', '非实时上架或激活确认',
                     '卡槽来自本轮系统库存：已填 1 条', '空白 0 条', '同ICCID多卡槽 0 条', '系统卡槽为空不回填历史'):
            self.assertIn(text, content)


if __name__ == '__main__':
    unittest.main()
