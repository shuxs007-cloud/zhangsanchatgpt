import unittest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from services.opensms_document_api.pipeline.machine_number import add_machine_number, verify_machine_number

class MachineNumberTests(unittest.TestCase):
    def test_values_position_style_and_idempotency(self):
        w=Workbook();s=w.active
        s.append(['电话','卡槽','ICCID'])
        for slot in ['110_01','r110_34.01','001_2',None,'bad']:
            s.append(['00123',slot,'00999F'])
        s['B2'].fill=PatternFill('solid',fgColor='FFC7CE')
        s.auto_filter.ref=s.dimensions
        result=add_machine_number(s)
        self.assertEqual(result['unknown'],2)
        rows=list(s.values)
        self.assertEqual(rows[0],('电话','机器号码','卡槽','ICCID'))
        self.assertEqual([r[1] for r in rows[1:]],['110','r110','001',None,None])
        self.assertEqual(s['B2'].fill.fgColor.rgb,s['C2'].fill.fgColor.rgb)
        self.assertEqual(s.auto_filter.ref,'A1:D6')
        verify_machine_number(list(rows[0]),rows[1:])
        with self.assertRaises(ValueError):add_machine_number(s)
        s['B2']='111'
        with self.assertRaises(ValueError):verify_machine_number(list(rows[0]),list(s.values)[1:])

if __name__=='__main__':unittest.main()
