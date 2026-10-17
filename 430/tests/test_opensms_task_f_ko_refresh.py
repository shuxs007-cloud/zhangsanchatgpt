"""Offline synthetic ko-sheet refresh; no browser or production data access."""
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import opensms_task_f_db as db
import opensms_task_f_ko_refresh as refresh
from opensms_report_date_cache import freeze
import ko_unactivated as ko


class KoRefreshTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.database = self.root / 'test.sqlite3'
        self.cache = self.root / 'cache.json'
        self.cards = self.root / 'cards.json'
        self.document = self.root / 'report.xlsx'
        self.receipt = self.root / 'receipt.json'
        self.run = 'offline-ko-refresh'
        self.keys = ['00012345678901234567', '8912230200081836703', '8912230200081836799']
        with db.open_database(self.database) as conn:
            for key in self.keys:
                conn.execute('INSERT INTO kooda_activation (query_iccid,source_iccid,lookup_status,queried_at,receipt_json) '
                             'VALUES (?,?,?,?,?)', (key, key + 'F', 'activation_missing', '2026-09-01T00:00:00Z', '{}'))
            conn.commit()
        self.frozen = freeze(self.database, self.cache, self.run)
        rows = [{'id': f'id-{n}', 'opr': 'Kooda', 'cardStore': key + 'F',
                 'slotIdentif': slot, 'cardNo': f'438555000{n}'}
                for n, (key, slot) in enumerate([(self.keys[0], 'r0130_1'), (self.keys[1], None),
                                                (self.keys[0], '131_2')])]
        self.inventory = {'runId': self.run, 'status': 'passed', 'systemRows': rows, 'totalRows': len(rows),
            'paginationTotalStable': True, 'paginationBackendIdsUnique': True, 'silentBrowserMode': True,
            'browserForegroundActivationAttempted': False}
        db.atomic_json(self.cards, self.inventory)
        self.system = {'run_id': self.run, 'path': str(self.cards), 'sha256': db.file_sha256(self.cards),
                       'raw_rows': rows, 'total_rows': len(rows)}
        book = Workbook()
        sheet = book.active
        sheet.title = 'sheet1'
        sheet.append(['电话', 'opr', '卡槽', 'ICCID', '激活日期', '截止日'])
        # Complete system dates do NOT prevent refreshing the unactivated sheet.
        sheet.append(['4385550000', 'Kooda', 'r0130_1', self.keys[0] + 'F', '2026-09-01', '2026-09-30'])
        book.create_sheet('处理汇总')
        evidence = ko.select_system_records(ko.load_snapshot(self.cache, self.frozen['sha256'], self.run, True), self.system)
        proof = ko.write_sheet(book, evidence)
        book.save(self.document)
        book.close()
        book = load_workbook(self.document, read_only=True)
        review = ko.review(book, proof, self.cache, self.frozen['sha256'], self.run, True, self.system)
        book.close()
        self.proof = {'result': '成功', 'checks': {'ko独立核验': True}, 'failed_checks': [],
            'output_path': str(self.document), 'process': {'ko_unactivated_sheet': proof},
            'verification': {'ko_unactivated_sheet_review': review,
                'system_kooda_source_review': {'matches': True, 'run_id': self.run,
                    'snapshot_sha256': self.system['sha256'], 'missing_either_date_rows': 0},
                'task_f_missing_activation_red_review': {'matches': True, 'expected_rows': 0}}}
        db.atomic_json(self.receipt, self.proof)

    def select(self, **changes):
        book = load_workbook(self.document, read_only=True)
        try:
            args = dict(proof=self.proof, inventory=self.inventory, inventory_hash=self.system['sha256'],
                        run_id=self.run, cache_path=self.cache, cache_hash=self.frozen['sha256'])
            return refresh.selected_rows(book, **{**args, **changes})
        finally:
            book.close()

    def prepare(self):
        return db.prepare_pending(self.database, self.cards, self.root / 'pending.json',
            document=self.document, document_sha256=db.file_sha256(self.document), run_id=self.run,
            document_proof=self.receipt, document_proof_sha256=db.file_sha256(self.receipt),
            ko_cache=self.cache, ko_cache_sha256=self.frozen['sha256'])

    def test_exact_sheet_intersection_including_outside_main_and_blank_slots(self):
        rows = self.select()
        self.assertEqual([r['queryIccid'] for r in rows], self.keys[:2])
        self.assertEqual(len(rows[0]['systemMatches']), 2)
        self.assertEqual(rows[1]['slot'], '')
        self.assertEqual(rows[0]['documentRow'], ko.SHEET + '!2')

    def test_sheet_tamper_rejected_before_query(self):
        book = load_workbook(self.document)
        book[ko.SHEET]['A2'] = self.keys[2]
        book.save(self.document)
        book.close()
        with self.assertRaisesRegex(RuntimeError, '工作表独立回读失败'):
            self.select()

    def test_wrong_hash_run_or_receipt_rejected(self):
        for values in ({'cache_hash': '0' * 64}, {'run_id': 'other-run'}, {'inventory_hash': '0' * 64}):
            with self.subTest(values=values), self.assertRaises((RuntimeError, ValueError)):
                self.select(**values)
        proof = deepcopy(self.proof)
        proof['verification']['ko_unactivated_sheet_review']['row_count'] += 1
        with self.assertRaises(RuntimeError):
            self.select(proof=proof)

    def test_pending_rechecks_existing_missing_and_preserves_current_report(self):
        before = db.file_sha256(self.document)
        pending = self.prepare()
        self.assertEqual(pending['pendingCount'], 2)
        self.assertEqual(pending['documentKoRefreshRows'], 2)
        self.assertEqual(pending['documentMissingDateRows'], 0)
        self.assertEqual(db.file_sha256(self.document), before)

    def test_new_valid_local_result_is_reused_not_requeried(self):
        with db.open_database(self.database) as conn:
            conn.execute("UPDATE kooda_activation SET lookup_status='found', activation_date='2026-09-02', "
                         "cutoff_date='2026-10-01' WHERE query_iccid=?", (self.keys[0],))
            conn.commit()
        pending = self.prepare()
        self.assertEqual(pending['pendingCount'], 1)
        self.assertEqual(pending['readyForUploadCount'], 1)
        self.assertEqual(pending['rows'][0]['queryIccid'], self.keys[1])

    def test_main_gap_and_ko_row_are_one_lookup(self):
        book = load_workbook(self.document)
        book['sheet1'].cell(2, 5).value = None
        book['sheet1'].cell(2, 6).value = None
        book.save(self.document)
        book.close()
        pending = self.prepare()
        self.assertEqual(pending['pendingCount'], 2)
        self.assertEqual(pending['rows'][0]['sourceOccurrenceCount'], 2)
        self.assertEqual(pending['documentDuplicateIccidRows'], 1)

    def test_snapshot_next_round_removes_confirmed_activation(self):
        with db.open_database(self.database) as conn:
            conn.execute("UPDATE kooda_activation SET lookup_status='found', activation_date='2026-09-02', "
                         "cutoff_date='2026-10-01' WHERE query_iccid=?", (self.keys[0],))
            conn.commit()
        new = self.root / 'next.json'
        receipt = freeze(self.database, new, self.run)
        evidence = ko.select_system_records(ko.load_snapshot(new, receipt['sha256'], self.run, True), self.system)
        self.assertEqual([r['queryIccid'] for r in evidence['rows']], [self.keys[1]])

    def test_detached_dispatch_when_only_ko_sheet_needs_refresh(self):
        import opensms_task_f_background as background
        runner = SimpleNamespace(task_a_only=False, run_id=self.run, delivery=self.root / 'main',
            state={'reportDateCache': {'path': str(self.cache), 'sha256': self.frozen['sha256']}},
            args=SimpleNamespace(cdp='http://127.0.0.1:1'), child_environment=lambda: {},
            verify_system_kooda_acquisition=lambda acquire: None)
        with patch.object(background, 'ROOT', self.root), patch.object(background, 'JOBS_ROOT', self.root / 'jobs'), \
             patch.object(background.subprocess, 'Popen', return_value=SimpleNamespace(pid=123)) as spawn:
            result = background.prepare_and_start(runner,
                {'cardResult': str(self.document), 'cardReceipt': str(self.receipt)},
                {'systemKoodaInventory': str(self.cards), 'systemKoodaSha256': self.system['sha256']})
            spawn.assert_called_once()
            self.assertTrue(result['koRefreshEnabled'])
            self.assertEqual(result['documentKoRefreshRows'], 2)
            self.assertEqual(result['documentMissingDateRows'], 0)
            self.assertEqual(result['totalLookupNeededRows'], 2)
            job = background.validate_job(Path(result['backgroundJob']), result['backgroundJobSha256'])
            # Query preparation independently uses private copies, not originals.
            bound = job['inputs']
            pending = db.prepare_pending(self.database, Path(bound['inventory']['path']), self.root / 'detached.json',
                document=Path(bound['document']['path']), document_sha256=bound['document']['sha256'], run_id=self.run,
                document_proof=Path(bound['documentReceipt']['path']),
                document_proof_sha256=bound['documentReceipt']['sha256'],
                ko_cache=Path(bound['reportDateCache']['path']), ko_cache_sha256=bound['reportDateCache']['sha256'])
            self.assertEqual(pending['pendingCount'], 2)

    def test_independent_pending_audit_and_tamper_detection(self):
        import opensms_parallel_audit as audit
        pending = self.prepare()
        config = {'runId': self.run, 'taskFReportBackfillMode': True, 'taskFDocumentDateSource': 'system',
            'taskFResultsApplyNextExport': False,
            'taskFResultsWriteBackToDocument': False, 'taskFCloudUploadAttempted': False,
            'taskFDocumentPath': str(self.document), 'taskFDocumentSha256': db.file_sha256(self.document),
            'taskFDocumentRunId': self.run, 'taskFDocumentReceipt': str(self.receipt),
            'taskFInventoryReceipt': str(self.cards), 'taskFReadyForUploadCount': 0,
            'taskFDocumentExcludedCount': 0, 'taskFDocumentMissingActivationIccidCount': 0}
        lookup = {'results': [{'queryIccid': row['queryIccid'], 'status': 'not_found'} for row in pending['rows']]}
        resolver = audit.PathResolver(self.root / 'config.json', None)
        with patch.object(audit, 'check_card_receipt', return_value=SimpleNamespace(status=audit.PASS)):
            errors, details = audit.review_task_f_report_backfill(config, pending, self.inventory, lookup, resolver)
            self.assertEqual(errors, [])
            self.assertEqual(details['document_counts']['documentKoRefreshRows'], 2)
            pending['rows'][0]['sourceOccurrences'][0]['systemMatches'] = []
            errors, _ = audit.review_task_f_report_backfill(config, pending, self.inventory, lookup, resolver)
            self.assertTrue(errors)


if __name__ == '__main__':
    unittest.main()
