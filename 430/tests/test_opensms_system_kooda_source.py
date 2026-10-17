"""Offline regression for system-authoritative Kooda report fields."""
import contextlib
from datetime import datetime
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("system_kooda_document_test",
    ROOT / "services/opensms_document_api/pipeline/process_card_slots.py")
DOCUMENT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(DOCUMENT)


def system_row(**changes):
    return {"id": "backend-1", "opr": "Telus1", "slotIdentif": "130_01.01",
            "cardNo": "14385550123", "cardStore": "8912230200081836799F",
            "chatrActivateDate": "2026-09-01", "chatrPlanExpire": "2026-09-20",
            "createTime": "2026-01-01 00:00:00", **changes}


class SystemSourceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="system-kooda-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "system.json"

    def save_source(self, rows=None, **changes):
        rows = [system_row()] if rows is None else rows
        payload = {"status": "passed", "runId": "test-current-round", "systemRows": rows,
                   "totalRows": len(rows), "paginationTotalStable": True,
                   "paginationBackendIdsUnique": True, "silentBrowserMode": True,
                   "browserForegroundActivationAttempted": False, **changes}
        self.source.write_text(json.dumps(payload), encoding="utf-8")
        cache = self.root / "frozen-dates.json"
        cache.write_text(json.dumps({"schemaVersion": 1, "status": "passed", "runId": "test-current-round",
            "createdAt": "2026-09-07T01:00:00+08:00", "sourceDatabase": None, "sourceExists": False,
            "recordCount": 0, "records": []}))
        return {"OPENSMS_SYSTEM_KOODA_PATH": str(self.source),
                "OPENSMS_SYSTEM_KOODA_SHA256": hashlib.sha256(self.source.read_bytes()).hexdigest(),
                "OPENSMS_REPORT_DATE_CACHE_PATH": str(cache),
                "OPENSMS_REPORT_DATE_CACHE_SHA256": hashlib.sha256(cache.read_bytes()).hexdigest(),
                "OPENSMS_RUN_ID": "test-current-round"}

    def test_hash_bound_snapshot_resolves_composite_identity_and_direct_dates(self):
        with mock.patch.dict(os.environ, self.save_source()):
            source = DOCUMENT.load_system_kooda_source()
        row = DOCUMENT.resolve_system_kooda_row(source, slot="130_1", phone="4385550123")
        self.assertEqual(row["iccid"], "8912230200081836799F")
        self.assertEqual(row["activation_dt"], datetime(2026, 9, 1))
        self.assertEqual(row["expiry_dt"], datetime(2026, 9, 20))  # Not +29 days.
        self.assertEqual(source["source"], "system")
        self.assertEqual(source["indexed_rows"], 1)

    def test_snapshot_gate_tampering_is_rejected(self):
        for changes in ({"status": "failed"}, {"runId": "old-round"}, {"totalRows": 2},
                        {"paginationTotalStable": False}, {"paginationBackendIdsUnique": False},
                        {"silentBrowserMode": False}, {"browserForegroundActivationAttempted": True}):
            with self.subTest(changes=changes), mock.patch.dict(os.environ, self.save_source(**changes)):
                with self.assertRaises(RuntimeError):
                    DOCUMENT.load_system_kooda_source()
        env = self.save_source()
        self.source.write_text(self.source.read_text() + " ")
        with mock.patch.dict(os.environ, env), self.assertRaisesRegex(RuntimeError, "SHA256"):
            DOCUMENT.load_system_kooda_source()

    def test_missing_environment_does_not_fall_back_to_local_database(self):
        for missing in ("OPENSMS_SYSTEM_KOODA_PATH", "OPENSMS_SYSTEM_KOODA_SHA256", "OPENSMS_RUN_ID"):
            env = self.save_source()
            env[missing] = ""
            with self.subTest(missing=missing), mock.patch.dict(os.environ, env), \
                 mock.patch.object(DOCUMENT.sqlite3, "connect", side_effect=AssertionError("no database")), \
                 self.assertRaises(RuntimeError):
                DOCUMENT.load_system_kooda_source()

    def test_duplicate_key_backend_id_wrong_opr_and_numeric_iccid_fail(self):
        cases = [[system_row(), system_row(id="id-2", slotIdentif="130_1", cardNo="4385550123")],
                 [system_row(), system_row(slotIdentif="130_2", cardNo="14385550124")],
                 [system_row(opr="Chatr")], [system_row(cardStore=8912230200081836799)]]
        for rows in cases:
            with self.subTest(rows=rows), mock.patch.dict(os.environ, self.save_source(rows)), \
                 self.assertRaises(RuntimeError):
                DOCUMENT.load_system_kooda_source()

    def test_unkeyed_inventory_is_counted_but_main_unmatched_identity_fails(self):
        rows = [system_row(), system_row(id="unregistered", cardNo=None, slotIdentif="130_2")]
        with mock.patch.dict(os.environ, self.save_source(rows)):
            source = DOCUMENT.load_system_kooda_source()
        self.assertEqual(source["unkeyed_rows"], 1)
        self.assertEqual(source["total_rows"], 2)
        for slot, phone in (("130_2", "14385550123"), ("130_1", "14385550124"), ("", "")):
            with self.subTest(slot=slot, phone=phone), self.assertRaises(RuntimeError):
                DOCUMENT.resolve_system_kooda_row(source, slot=slot, phone=phone)

    def test_empty_snapshot_and_blank_dates_are_allowed_without_inference(self):
        with mock.patch.dict(os.environ, self.save_source([])):
            self.assertEqual(DOCUMENT.load_system_kooda_source()["index"], {})
        for raw in (None, "", " "):
            with self.subTest(raw=raw):
                self.assertIsNone(DOCUMENT.system_kooda_date(raw, "chatrActivateDate"))
        for field in ("chatrActivateDate", "chatrPlanExpire"):
            with self.subTest(field=field), mock.patch.dict(os.environ,
                    self.save_source([system_row(**{field: "not-a-date"})])), self.assertRaises(RuntimeError):
                DOCUMENT.load_system_kooda_source()

    def test_iso_dates_validate_calendar_and_do_not_convert_timezone(self):
        for raw in ("2026-09-01", "2026-09-01 13:04:05", "2026-09-01T23:00:01-05:00", "2026-09-01T13:00Z"):
            with self.subTest(raw=raw):
                self.assertEqual(DOCUMENT.system_kooda_date(raw, "date"), datetime(2026, 9, 1))
        for raw in ("2026-02-30", "2026/09/01", "2026-9-1", "20260901", "2026-09-01T27:00:00", 20260901):
            with self.subTest(raw=raw), self.assertRaises(RuntimeError):
                DOCUMENT.system_kooda_date(raw, "date")

    def test_process_uses_system_dates_and_no_kooda_google_fields_or_sqlite(self):
        class GuardedGoogleCard(dict):
            def __getitem__(self, key):
                if key in {"iccid", "open_date", "activation_date", "expiry_date"}:
                    raise AssertionError("Kooda cannot read Google " + key)
                return super().__getitem__(key)

        raw = self.root / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        headers = ["电话", "运营商", "opr", "卡槽", "创建时间", "送达时间", "接收时间", "状态",
                   "10日发送窗口数", "10日回复窗口数", "10日回复率", "10日总发送量", "10日总回复量"]
        for day in ("09-03", "09-04", "09-05", "09-06"):
            headers.extend(day + suffix for suffix in ("发送窗口数", "回复窗口数", "回复率", "总发送量", "总回复量"))
        sheet.append(headers)
        cards = [("14385550123", "Kooda", "130_1"), ("14385550124", "Telus1", "130_2"),
                 ("14385550125", "Chatr", "130_3"), ("14385550126", "Koodo", "130_4")]
        for phone, opr, slot in cards:
            sheet.append([phone, "B1" if opr != "Chatr" else "A1", opr, slot,
                          "2026-09-06 12:00:00", None, None,
                          "注册失败" if slot in {"130_1", "130_4"} else "上架", 400, 40, .1, 400, 40] +
                         [100, 10, .1, 100, 10] * 4)
        workbook.save(raw)
        rows = [system_row(), system_row(id="backend-2", slotIdentif="130_2", cardNo="14385550124",
                                        cardStore="8912230200081836798f", chatrActivateDate=None, chatrPlanExpire=None),
                system_row(id="backend-3", slotIdentif="130_4", cardNo="14385550126",
                           cardStore="8912230200081836797F")]
        env = self.save_source(rows)
        output = self.root / "output.xlsx"
        env.update({"EXPORT4DAY_SERVER_DATE": "2026-09-06", "OPENSMS_OUTPUT_PATH": str(output)})
        google = {}
        for i, (phone, opr, slot) in enumerate(cards, 1):
            if slot == "130_4":
                continue  # A system-authoritative Kooda row needs no Google card.
            card = {"raw_slot": slot, "phone": phone, "carrier_price": opr,
                    "source_sheet": "四號機房", "source_row": i + 1, "note": "保留Google其它字段",
                    "iccid": "GOOGLE-MUST-NOT-OVERRIDE", "open_date": "2020-01-01",
                    "activation_date": "2020-01-02", "expiry_date": "2020-02-01"}
            google[slot] = GuardedGoogleCard(card) if opr != "Chatr" else card
        log = io.StringIO()
        with mock.patch.dict(os.environ, env), mock.patch.object(DOCUMENT, "MAIN_PATH", raw), \
             mock.patch.object(DOCUMENT, "OUTPUT_DIR", self.root), \
             mock.patch.object(DOCUMENT, "build_assoc_lookup", return_value=(google, ["四號機房"], list(google.values()), {"四號機房": 3})), \
             mock.patch.object(DOCUMENT, "load_task_f_kooda_dates", side_effect=AssertionError("no Task F DB load")), \
             mock.patch.object(DOCUMENT.sqlite3, "connect", side_effect=AssertionError("no SQLite")), \
             contextlib.redirect_stdout(log):
            DOCUMENT.process()
        summary = json.loads(log.getvalue())
        self.assertEqual(summary["system_kooda_source"]["matched_rows"], 3)
        self.assertEqual(summary["system_kooda_source"]["missing_rows"], 0)
        self.assertEqual(summary["main_rows"], 4)
        self.assertEqual(summary["matched_rows"], 4)
        self.assertEqual(summary["google_matched_rows"], 3)
        self.assertEqual(summary["unmatched_slots"], ["130_4"])
        self.assertEqual(summary["google_all_missing_from_main_rows"], 0)
        self.assertEqual(summary["task_f_kooda_date_source"]["source"], "system_export_then_frozen_local")
        self.assertFalse(summary["task_f_kooda_date_source"]["database_exists"])
        result = load_workbook(output)
        try:
            main = result.active
            index = {cell.value: cell.column for cell in main[1]}
            self.assertEqual(main.cell(1, main.max_column).value, "ICCID")
            self.assertEqual(sum(cell.value == "ICCID" for cell in main[1]), 1)
            for row in (2, 3):
                self.assertIsNone(main.cell(row, index["开卡日期"]).value)
            self.assertEqual(main.cell(2, index["ICCID"]).value, "8912230200081836799F")
            self.assertEqual(main.cell(2, index["ICCID"]).number_format, "@")
            self.assertEqual(main.cell(2, index["激活日期"]).value, datetime(2026, 9, 1))
            self.assertEqual(main.cell(2, index["截止日"]).value, datetime(2026, 9, 20))
            self.assertEqual(main.cell(3, index["ICCID"]).value, "8912230200081836798f")
            self.assertIsNone(main.cell(3, index["激活日期"]).value)
            self.assertIsNone(main.cell(3, index["截止日"]).value)
            self.assertEqual(main.cell(3, index["换号状态"]).value, "未换号")
            self.assertTrue(all(str(cell.fill.fgColor.rgb).upper().endswith("FFC7CE") for cell in main[3]))
            self.assertEqual(main.cell(4, index["开卡日期"]).value, datetime(2020, 1, 1))
            self.assertEqual(main.cell(4, index["激活日期"]).value, datetime(2020, 1, 2))
            self.assertEqual(main.cell(4, index["截止日"]).value, datetime(2020, 2, 1))
            self.assertIsNone(main.cell(4, index["ICCID"]).value)
            self.assertEqual(main.cell(5, index["ICCID"]).value, "8912230200081836797F")
            self.assertIsNone(main.cell(5, index["开卡日期"]).value)
            self.assertEqual(main.cell(5, index["激活日期"]).value, datetime(2026, 9, 1))
            self.assertEqual(main.cell(5, index["截止日"]).value, datetime(2026, 9, 20))
            activated_rows = list(result["已激活未上架"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(activated_rows), 2)
            for activated in activated_rows:
                self.assertEqual(activated[11], "system_inventory_legacy")
                self.assertIsNone(activated[12])
        finally:
            result.close()


if __name__ == "__main__":
    unittest.main()
