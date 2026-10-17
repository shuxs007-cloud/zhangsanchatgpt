"""Offline 10-day/raw activation scope: real XLSX, private DB, no business IO."""
from copy import deepcopy
from datetime import datetime
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'services/opensms_document_api/pipeline'))
import opensms_task_f_db as DB
import opensms_parallel_audit as AUDIT
import report_kooda_source as SOURCE
import opensms_verify_card_slot_fast as VERIFY
import process_card_slots as GENERATE
import ko_unactivated as KO
import opensms_process_document as DISPATCH


def iccid(n):
    return str(8912230200000000000 + n)


class ReportOnlyFTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.run_id = '20261005T050000+0800'
        self.raw, self.doc = self.root / 'raw.xlsx', self.root / 'final.xlsx'
        self.inventory_path, self.pending_path = self.root / 'source.json', self.root / 'pending.json'
        self.database = self.root / 'f.sqlite3'
        self.headers = ['电话', '运营商', 'opr', '卡槽', 'ICCID', '激活日期', '截止日']
        # 1 both missing; 2 source activation present/cutoff blank; 3 local date;
        # 4 final creation estimate; 5/6 same ICCID, one real source date;
        # 7 duplicate unknown ICCID; 8 non-Kooda; 9 archived blank slot; 10 invalid.
        self.rows = []
        for n, card, opr, activation, cutoff in (
            (1, 1, 'Kooda', None, None), (2, 2, 'Kooda', '2026-10-01', None),
            (3, 3, 'Telus1', None, None), (4, 4, 'Koodo', None, '2026-10-30'),
            (5, 5, 'Kooda', None, None), (6, 5, 'Kooda', '2026-10-02', None),
            (7, 1, 'Kooda', None, None), (8, 8, 'Chatr', None, None),
            (9, 9, 'Kooda', None, None), (10, None, 'Kooda', None, None)):
            self.rows.append([str(14385550000 + n), 'operator', opr,
                None if n == 9 else f'130_{n}', iccid(card) + 'F' if card else None, activation, cutoff])
        self.write_book(self.raw, self.rows)
        final = deepcopy(self.rows)
        final[3][5] = '2026-10-01'  # An estimate must not become a real activation.
        final[2][5:7] = ['2026-10-01', '2026-10-30']
        self.write_book(self.doc, final)
        self.inventory = SOURCE.projection_payload(self.headers, self.rows, self.run_id, self.raw, DB.file_sha256(self.raw))
        DB.atomic_json(self.inventory_path, self.inventory)
        # Old-schema/date-present records still skip lookup; blank dates retry.
        with DB.open_database(self.database) as connection:
            for n, activation in ((3, '2026-10-01'), (4, None), (99, None)):
                connection.execute('''INSERT INTO kooda_activation(query_iccid,source_iccid,lookup_status,
                    activation_date,cutoff_date,queried_at,record_schema_version,receipt_json)
                    VALUES(?,?,?,?,?,?,?,?)''', (iccid(n), iccid(n) + 'F', 'activation_missing',
                    activation, None, '2026-10-01T00:00:00+08:00', 1, '{}'))
            connection.commit()
        self.proof_path = self.root / 'proof.json'
        self.proof = {'result': '成功', 'checks': {'offline': True}, 'failed_checks': [],
            'output_path': str(self.doc), 'verification': {
                'system_kooda_source_review': {'matches': True, 'run_id': self.run_id,
                    'source_kind': SOURCE.KIND, 'source_scope': SOURCE.SCOPE,
                    'snapshot_sha256': DB.file_sha256(self.inventory_path)},
                'date_resolution_review': {'matches': True, 'creation_date_fallback_rows': []}}}
        DB.atomic_json(self.proof_path, self.proof)

    def write_book(self, path, rows):
        book = Workbook()
        book.active.title = 'sheet1'
        book.active.append(self.headers)
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()

    def prepare(self):
        return DB.prepare_pending(self.database, self.inventory_path, self.pending_path,
            document=self.doc, document_sha256=DB.file_sha256(self.doc), run_id=self.run_id,
            document_proof=self.proof_path, document_proof_sha256=DB.file_sha256(self.proof_path))

    def test_only_both_missing_10day_iccids_and_dedup_once(self):
        pending = self.prepare()
        self.assertEqual([iccid(1), iccid(4), iccid(9)], [r['queryIccid'] for r in pending['rows']])
        self.assertEqual(2, pending['rows'][0]['sourceOccurrenceCount'])
        self.assertEqual([iccid(3)], [r['queryIccid'] for r in pending['alreadyQueried']])
        self.assertEqual(SOURCE.ELIGIBILITY, pending['dateEligibilityMode'])
        self.assertEqual(0, pending['documentKoRefreshRows'])
        self.assertEqual(1, pending['documentExcludedCount'])
        self.assertTrue(all(r['sourceRowIdentity'].startswith(self.run_id) for r in pending['reportEligibleRecords']))

    def test_cutoff_only_missing_and_any_real_duplicate_activation_never_query(self):
        selected, _ = DB.document_backfill_rows(self.inventory, self.doc, DB.file_sha256(self.doc), self.run_id)
        self.assertFalse({iccid(2), iccid(5), iccid(8), iccid(99)} & {r['queryIccid'] for r in selected})
        self.assertIn(iccid(4), {r['queryIccid'] for r in selected})

    def test_raw_projection_independently_recomputed_and_tamper_rejected(self):
        SOURCE.verify_projection_workbook(self.inventory, self.raw, self.run_id)
        for change in ({'sourceRowCount': -1}, {'sourceScope': 'all-system'}, {'runId': 'old'},
                       {'backendInventoryQueried': True}, {'silentBrowserMode': False}):
            with self.subTest(change=change), self.assertRaises(RuntimeError):
                SOURCE.validate_projection({**self.inventory, **change}, self.run_id)
        altered = deepcopy(self.inventory)
        altered['systemRows'][0]['id'] = 'fake-backend-id'
        with self.assertRaises(RuntimeError):
            SOURCE.validate_projection(altered, self.run_id)
        altered = deepcopy(self.inventory)
        altered['systemRows'] = altered['systemRows'][1:]
        altered['totalRows'] -= 1
        with self.assertRaises(RuntimeError):
            SOURCE.verify_projection_workbook(altered, self.raw, self.run_id)

    def test_local_cli_uses_existing_hash_bound_scan_without_browser(self):
        import contextlib
        import io
        sys.path.insert(0, str(ROOT))
        from opensms_export_scan_cache import build_scan_cache
        import opensms_build_report_kooda_source as builder
        cache, output = self.root / 'scan.pkl', self.root / 'local-source.json'
        build_scan_cache(self.raw, cache)
        argv = ['local-source', '--scan-cache', str(cache), '--source', str(self.raw),
                '--run-id', self.run_id, '--output', str(output)]
        with patch.object(sys, 'argv', argv), contextlib.redirect_stdout(io.StringIO()):
            builder.main()
        self.assertEqual(self.inventory, DB.load_json(output))

    def test_numeric_iccid_missing_column_and_bad_dates_fail_without_fallback(self):
        for field, value in ((4, int(iccid(1))), (5, 'not-a-date')):
            rows = deepcopy(self.rows)
            rows[0][field] = value
            with self.assertRaises(RuntimeError):
                SOURCE.project_rows(self.headers, rows, self.run_id)
        with self.assertRaises(RuntimeError):
            SOURCE.project_rows(self.headers[:-3], self.rows, self.run_id)
        self.assertEqual('2026-10-01', SOURCE.source_date(datetime(2026, 10, 1)))

    def test_final_row_identity_tamper_rejected(self):
        book = load_workbook(self.doc)
        book.active.cell(2, 5, iccid(99) + 'F')
        book.save(self.doc)
        book.close()
        with self.assertRaisesRegex(RuntimeError, '源行身份'):
            self.prepare()

    def test_new_snapshot_independent_verifier_supports_blank_slots(self):
        snapshot = VERIFY.load_system_kooda_snapshot(self.inventory_path, DB.file_sha256(self.inventory_path), self.run_id)
        row = dict(zip(self.headers, self.rows[8]))
        self.assertEqual(iccid(9) + 'F', VERIFY.find_system_record(snapshot, row, 10)['iccid'])
        self.assertIsNone(VERIFY.find_system_record(snapshot, row, 9))

    def test_excel_empty_iccid_roundtrip_all_three_consumers(self):
        # Production failure: source shared-string '' reopens as None after
        # saving the report. Blank identity is excluded from F, not fatal.
        for source_value in (None, ''):
            with self.subTest(source_value=source_value):
                payload = deepcopy(self.inventory)
                source = next(r for r in payload['systemRows'] if r['sourceRow'] == 11)
                source['cardStore'] = source['sourceData']['ICCID'] = source_value
                DB.atomic_json(self.inventory_path, payload)
                snapshot = VERIFY.load_system_kooda_snapshot(
                    self.inventory_path, DB.file_sha256(self.inventory_path), self.run_id)
                row = dict(zip(self.headers, self.rows[9]))
                row['ICCID'] = ''
                final_rows = deepcopy(self.rows)
                final_rows[9][4] = ''
                self.write_book(self.doc, final_rows)
                book = load_workbook(self.doc, read_only=True)
                try:
                    row['ICCID'] = book['sheet1'].cell(11, 5).value
                finally:
                    book.close()
                self.assertIsNone(row['ICCID'])
                self.assertTrue(VERIFY.review_system_kooda_row(snapshot, row, source_row=11)['matches'])
                with patch.dict('os.environ', {
                        'OPENSMS_SYSTEM_KOODA_PATH': str(self.inventory_path),
                        'OPENSMS_SYSTEM_KOODA_SHA256': DB.file_sha256(self.inventory_path),
                        'OPENSMS_RUN_ID': self.run_id}):
                    generated = GENERATE.load_system_kooda_source()
                GENERATE.resolve_system_kooda_row(generated, slot=row['卡槽'],
                    phone=row['电话'], iccid=row['ICCID'], source_row=11)
                selected, metadata = DB.document_backfill_rows(
                    payload, self.doc, DB.file_sha256(self.doc), self.run_id)
                self.assertTrue(metadata['documentExcludedCount'] >= 1)
                self.assertTrue(all(r['queryIccid'] for r in selected))

    def test_blank_iccid_tolerance_does_not_accept_identity_tamper(self):
        snapshot = VERIFY.load_system_kooda_snapshot(
            self.inventory_path, DB.file_sha256(self.inventory_path), self.run_id)
        blank_row = dict(zip(self.headers, self.rows[9]))
        for bad in (' ', '0', 0, False, iccid(1), '=1'):
            with self.subTest(bad=bad):
                altered = {**blank_row, 'ICCID': bad}
                self.assertIsNone(VERIFY.find_system_record(snapshot, altered, 11))
                final = deepcopy(self.rows)
                final[9][4] = bad
                self.write_book(self.doc, final)
                with self.assertRaises(RuntimeError):
                    DB.document_backfill_rows(self.inventory, self.doc, DB.file_sha256(self.doc), self.run_id)
        valid_row = dict(zip(self.headers, self.rows[0]))
        for bad in (None, '', iccid(1), iccid(1) + 'f', iccid(1) + 'F '):
            self.assertIsNone(VERIFY.find_system_record(snapshot, {**valid_row, 'ICCID': bad}, 2))
        self.assertIsNone(VERIFY.find_system_record(snapshot, {**blank_row, '卡槽': '131_10'}, 11))
        self.assertIsNone(VERIFY.find_system_record(snapshot, {**blank_row, '电话': '1234567890'}, 11))

    def test_ko_sheet_cannot_reintroduce_outside_or_activated_iccids(self):
        evidence = {'rows': [{'queryIccid': iccid(n)} for n in (1, 2, 5, 99)], 'proof': {'run_id': self.run_id}}
        result = KO.select_system_records(evidence, {'run_id': self.run_id, 'path': str(self.inventory_path),
            'sha256': DB.file_sha256(self.inventory_path), 'total_rows': len(self.inventory['systemRows']),
            'raw_rows': self.inventory['systemRows'], 'source_kind': SOURCE.KIND})
        self.assertEqual([iccid(1)], [r['queryIccid'] for r in result['rows']])
        self.assertEqual(KO.REPORT_SCOPE, result['proof']['scope'])
        self.assertEqual(1, result['proof']['excluded_not_in_system_count'])
        self.assertEqual(2, result['proof']['excluded_activation_present_count'])

    def dispatch_fixture(self, source_kind=SOURCE.KIND):
        source = deepcopy(self.inventory)
        if source_kind != SOURCE.KIND:
            source = {'status': 'passed', 'runId': self.run_id, 'systemRows': []}
        DB.atomic_json(self.inventory_path, source)
        environment = {'OPENSMS_KO_UNACTIVATED_REQUIRED': '1', 'OPENSMS_RUN_ID': self.run_id,
            'OPENSMS_SYSTEM_KOODA_PATH': str(self.inventory_path),
            'OPENSMS_SYSTEM_KOODA_SHA256': DB.file_sha256(self.inventory_path),
            'OPENSMS_REPORT_DATE_CACHE_SHA256': 'a' * 64}
        scope = KO.REPORT_SCOPE if source_kind == SOURCE.KIND else KO.SCOPE
        slot_source = KO.REPORT_SLOT_SOURCE if source_kind == SOURCE.KIND else KO.SLOT_SOURCE
        proof = {'enabled': True, 'status': 'passed', 'scope': scope, 'slot_source': slot_source,
            'sheet': KO.SHEET, 'run_id': self.run_id, 'snapshot_sha256': 'a' * 64,
            'system_path': str(self.inventory_path), 'system_sha256': environment['OPENSMS_SYSTEM_KOODA_SHA256']}
        google = {'enabled': False, 'status': 'skipped', 'reason': 'user_disabled',
            'unused_port_device_scope': 'system_source_only'}
        result = {'result': '成功', 'process': {'google_dependency': google, 'ko_unactivated_sheet': proof},
            'verification': {'google_dependency_review': {**google, 'matches': True},
                'ko_unactivated_sheet_review': {**proof, 'matches': True}}}
        return environment, result

    def dispatch(self, environment, result):
        child = DISPATCH.subprocess.CompletedProcess([], 0, json.dumps(result), '')
        with patch.dict(DISPATCH.os.environ, environment, clear=True), \
                patch.object(DISPATCH.subprocess, 'run', return_value=child):
            return DISPATCH.run_local(Path('builder'), self.raw, None, self.root, '2026-10-04')

    def test_document_dispatch_accepts_both_bound_source_contracts(self):
        for kind in (SOURCE.KIND, 'legacy'):
            with self.subTest(kind=kind):
                environment, result = self.dispatch_fixture(kind)
                self.assertEqual(result, self.dispatch(environment, result))

    def test_document_dispatch_rejects_cross_scope_and_independent_review_tamper(self):
        for kind in (SOURCE.KIND, 'legacy'):
            for field, value in (('scope', 'all-system'), ('slot_source', 'historical.slot'),
                                 ('run_id', 'old'), ('system_sha256', '0' * 64),
                                 ('snapshot_sha256', '0' * 64)):
                with self.subTest(kind=kind, field=field):
                    environment, result = self.dispatch_fixture(kind)
                    result['process']['ko_unactivated_sheet'][field] = value
                    result['verification']['ko_unactivated_sheet_review'][field] = value
                    with self.assertRaises(RuntimeError):
                        self.dispatch(environment, result)
            environment, result = self.dispatch_fixture(kind)
            result['verification']['ko_unactivated_sheet_review']['matches'] = False
            with self.assertRaises(RuntimeError):
                self.dispatch(environment, result)

    def test_document_dispatch_rejects_projection_hash_and_metadata_tamper(self):
        environment, result = self.dispatch_fixture()
        self.inventory_path.write_bytes(b'changed')
        with self.assertRaisesRegex(RuntimeError, 'SHA256'):
            self.dispatch(environment, result)
        for field, value in (('sourceScope', 'all-system'), ('runId', 'old'),
                             ('backendInventoryQueried', True)):
            with self.subTest(field=field):
                environment, result = self.dispatch_fixture()
                source = DB.load_json(self.inventory_path)
                source[field] = value
                DB.atomic_json(self.inventory_path, source)
                environment['OPENSMS_SYSTEM_KOODA_SHA256'] = DB.file_sha256(self.inventory_path)
                result['process']['ko_unactivated_sheet']['system_sha256'] = environment['OPENSMS_SYSTEM_KOODA_SHA256']
                result['verification']['ko_unactivated_sheet_review']['system_sha256'] = environment['OPENSMS_SYSTEM_KOODA_SHA256']
                with self.assertRaises(RuntimeError):
                    self.dispatch(environment, result)

    def test_real_pending_matches_independent_audit_and_scope_tamper_fails(self):
        pending = self.prepare()
        config = {'runId': self.run_id, 'taskFReportBackfillMode': True,
            'taskFDocumentDateSource': 'system', 'taskFResultsApplyNextExport': False,
            'taskFResultsWriteBackToDocument': False, 'taskFCloudUploadAttempted': False,
            'taskFDocumentPath': str(self.doc), 'taskFDocumentSha256': DB.file_sha256(self.doc),
            'taskFDocumentRunId': self.run_id, 'taskFDocumentReceipt': str(self.proof_path),
            'taskFInventoryReceipt': str(self.inventory_path), 'taskFReadyForUploadCount': 1,
            'taskFDocumentExcludedCount': 1, 'taskFDocumentMissingActivationIccidCount': 4}
        lookup = {'results': [{'queryIccid': r['queryIccid'], 'status': 'not_found'} for r in pending['rows']]}
        resolver = AUDIT.PathResolver(self.root / 'config.json', None)
        with patch.object(AUDIT, 'check_card_receipt', return_value=SimpleNamespace(status='passed')):
            errors, _ = AUDIT.review_task_f_report_backfill(config, pending, self.inventory, lookup, resolver)
            self.assertEqual([], errors)
            altered = deepcopy(pending)
            altered['sourceScope'] = 'all-system'
            errors, _ = AUDIT.review_task_f_report_backfill(config, altered, self.inventory, lookup, resolver)
            self.assertTrue(errors)


class ReportOnlyBackgroundTests(unittest.TestCase):
    def setUp(self):
        import test_opensms_task_f_background as fixtures
        self.fixture = fixtures.BackgroundFTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.background = fixtures.background

    def configure(self, activation=None, cache_activation=None):
        f = self.fixture
        headers = ['电话', '运营商', 'opr', '卡槽', 'ICCID', '激活日期', '截止日']
        rows = [['14385550123', 'operator', 'Kooda', '130_1', iccid(1) + 'F', activation, None]]
        raw = f.root / 'raw.xlsx'
        for path in (raw, f.document):
            book = Workbook()
            book.active.title = 'sheet1'
            book.active.append(headers)
            book.active.append(rows[0])
            book.save(path)
            book.close()
        inventory = SOURCE.projection_payload(headers, rows, f.run_id, raw, DB.file_sha256(raw))
        DB.atomic_json(f.inventory, inventory)
        f.acquire['systemKoodaSha256'] = DB.file_sha256(f.inventory)
        proof = DB.load_json(f.proof)
        proof['verification']['system_kooda_source_review']['snapshot_sha256'] = DB.file_sha256(f.inventory)
        DB.atomic_json(f.proof, proof)
        if cache_activation:
            cache = DB.load_json(f.cache)
            cache.update(records=[{'queryIccid': iccid(1), 'activationDate': cache_activation}], recordCount=1)
            DB.atomic_json(f.cache, cache)
            f.runner.state['reportDateCache']['sha256'] = DB.file_sha256(f.cache)

    def test_known_activation_or_cutoff_only_gap_no_spawn(self):
        for source, local in (('2026-10-01', None), (None, '2026-10-01')):
            with self.subTest(source=source, local=local):
                self.configure(source, local)
                result, spawn = self.fixture.prepare()
                self.assertEqual('skipped', result['status'])
                spawn.assert_not_called()
                self.assertFalse((self.fixture.root / 'jobs').exists())

    def test_missing_dates_job_binds_raw_workbook_and_no_ko_extension(self):
        self.configure()
        result, spawn = self.fixture.prepare()
        self.assertEqual(1, spawn.call_count)
        job = self.background.validate_job(Path(result['backgroundJob']), result['backgroundJobSha256'])
        self.assertIn('rawExport', job['inputs'])
        self.assertFalse(job['koRefreshEnabled'])
        self.assertEqual(SOURCE.ELIGIBILITY, job['dateEligibilityMode'])
        raw = Path(job['inputs']['rawExport']['path'])
        raw.write_bytes(b'tampered')
        with self.assertRaises(RuntimeError):
            self.background.validate_job(Path(result['backgroundJob']), result['backgroundJobSha256'])


class ReportOnlyDocumentIntegrationTests(unittest.TestCase):
    def test_original_builder_and_independent_verifier_in_process(self):
        # Exercise the actual Python pipeline without Bash here-document temp
        # IO or loosening the existing OS network/write isolation.
        import contextlib
        import io
        import os
        import test_opensms_system_kooda_integration as fixtures
        f = fixtures.SystemKoodaIntegrationTests()
        f.setUp()
        self.addCleanup(f.doCleanups)
        book = load_workbook(f.raw)
        sheet = book.active
        col = sheet.max_column + 1
        for offset, header in enumerate(('ICCID', '激活日期', '截止日')):
            sheet.cell(1, col + offset, header)
        sheet.cell(2, col, '8912230200081836799F')
        sheet.cell(2, col + 1, '2026-09-01')
        sheet.cell(2, col + 2, '2026-12-15')
        sheet.cell(3, col, '8912230200081836798f')
        # An archived record needs no live backend slot/ID lookup.
        sheet.cell(3, 4).value = None
        values = list(sheet.iter_rows(values_only=True))
        book.save(f.raw)
        book.close()
        projection = SOURCE.projection_payload(values[0], values[1:], f.env['OPENSMS_RUN_ID'], f.raw, DB.file_sha256(f.raw))
        DB.atomic_json(f.system, projection)
        f.env.update(OPENSMS_SYSTEM_KOODA_SHA256=DB.file_sha256(f.system), OPENSMS_OUTPUT_PATH=str(f.output),
            OPENSMS_REPORT_ROOT=str(f.root), OPENSMS_MAIN_PATH=str(f.raw),
            OPENSMS_OUTPUT_DIR=str(f.output.parent), EXPORT4DAY_SERVER_DATE='2026-09-06')
        spec = importlib.util.spec_from_file_location('report_only_document_integration',
            ROOT / 'services/opensms_document_api/pipeline/process_card_slots.py')
        document = importlib.util.module_from_spec(spec)
        with patch.dict(os.environ, f.env), contextlib.redirect_stdout(io.StringIO()) as out:
            spec.loader.exec_module(document)
            document.process()
        process = json.loads(out.getvalue())
        f.env.update(PROCESS_JSON=json.dumps(process), ENRICH_JSON=json.dumps(process['enrichment']))
        with patch.dict(os.environ, f.env), patch.object(sys, 'argv', ['verify', str(f.output), str(f.raw)]), \
                contextlib.redirect_stdout(io.StringIO()) as out:
            VERIFY.main()
        review = json.loads(out.getvalue())
        for key in ('system_kooda_source_review', 'date_resolution_review', 'source_schema_review',
                    'task_f_missing_activation_red_review', 'task_f_missing_dates_new_card_review'):
            self.assertTrue(review[key]['matches'], (key, review[key]))
        self.assertEqual(SOURCE.KIND, review['system_kooda_source_review']['source_kind'])
        rows, _ = DB.document_backfill_rows(projection, f.output, DB.file_sha256(f.output), f.env['OPENSMS_RUN_ID'])
        self.assertEqual(['8912230200081836798'], [row['queryIccid'] for row in rows])


if __name__ == '__main__':
    unittest.main()
