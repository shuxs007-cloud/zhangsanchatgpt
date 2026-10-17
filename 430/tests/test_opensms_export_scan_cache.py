import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from opensms_analyze_current import read_export
from opensms_export_scan_cache import build_scan_cache, load_scan_cache


class OpenSmsExportScanCacheTests(unittest.TestCase):
    def make_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "电话", "运营商", "opr", "卡槽", "状态", "创建时间",
            "09-01发送窗口数", "09-01回复窗口数",
            "09-02发送窗口数", "09-02回复窗口数",
        ])
        sheet.append(["17000000001", "A1", "Chatr", "130_1", "上架", "2026-09-01", 80, 2, 90, 0])
        workbook.save(path)
        workbook.close()

    def test_cache_is_hash_bound_and_preserves_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            cache = root / "source.scan.pkl"
            receipt = root / "source.scan.json"
            self.make_workbook(source)

            result = build_scan_cache(source, cache, receipt)
            payload = load_scan_cache(cache, expected_source=source)

            self.assertEqual(result["status"], "passed")
            self.assertEqual(result["rowCount"], 1)
            self.assertEqual(payload["headers"][0], "电话")
            self.assertEqual(payload["rows"][0][0], "17000000001")

            workbook = load_workbook(source)
            workbook.active["B2"] = "A12"
            workbook.save(source)
            workbook.close()
            with self.assertRaisesRegex(RuntimeError, "SHA-256不一致"):
                load_scan_cache(cache, expected_source=source)

    def test_analysis_read_matches_direct_workbook_read(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            cache = root / "source.scan.pkl"
            receipt = root / "source.scan.json"
            self.make_workbook(source)
            build_scan_cache(source, cache, receipt)

            direct = read_export(str(source))
            cached = read_export(str(source), str(cache))

            self.assertEqual(direct["latestDate"], cached["latestDate"])
            self.assertEqual(direct["rows"], cached["rows"])
            self.assertEqual(direct["zero"], cached["zero"])
            self.assertTrue(cached["sharedScanCacheUsed"])
            self.assertEqual(cached["candidateProcessWorkbookReopens"], 0)


if __name__ == "__main__":
    unittest.main()
