"""Offline real pending-builder -> independent report-backfill audit contracts."""
from copy import deepcopy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]


def module(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / filename)
    value = importlib.util.module_from_spec(spec)
    sys.modules[name] = value
    spec.loader.exec_module(value)
    return value


DB = module("report_backfill_audit_db", "opensms_task_f_db.py")
AUDIT = module("report_backfill_independent_audit", "opensms_parallel_audit.py")


def iccid(number):
    return str(8912230200000000000 + number)


class ReportBackfillAuditTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="report-backfill-audit-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.document = self.root / "current.xlsx"
        self.inventory_path = self.root / "inventory.json"
        self.pending_path = self.root / "pending.json"
        self.database = self.root / "database.sqlite3"
        self.receipt_path = self.root / "card.json"
        self.run_id = "offline-current-round"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "sheet1"
        sheet.append(["电话", "卡槽", "opr", "激活日期", "ICCID"])
        source_rows = []
        for number, opr, activation, card in (
            (1, "Kooda", "2026-09-01", iccid(1) + "F"),
            (2, "Telus1", None, iccid(2) + "f"),
            (3, "Koodo", None, iccid(3) + "F"),
            (4, "KOODA", None, iccid(4) + "F"),
            (5, "Kooda", None, None),
            (6, "Chatr", None, iccid(6)),
        ):
            phone, slot = str(14385550000 + number), f"130_{number}"
            sheet.append([phone, slot, opr, activation, card])
            if opr != "Chatr":
                source_rows.append({"id": f"id-{number}", "opr": opr, "slotIdentif": slot + ".01",
                                    "cardNo": phone[1:], "cardStore": card, "fullRawField": "retained"})
        workbook.save(self.document)
        workbook.close()
        self.inventory = {"status": "passed", "runId": self.run_id, "totalRows": len(source_rows),
                          "systemRows": source_rows, "paginationTotalStable": True,
                          "paginationBackendIdsUnique": True, "silentBrowserMode": True,
                          "browserForegroundActivationAttempted": False}
        DB.atomic_json(self.inventory_path, self.inventory)
        with DB.open_database(self.database) as connection:
            for number, status, activation, cutoff, schema in (
                (2, "found", "2026-09-01", "2026-09-30", 1),
                (3, "activation_missing", None, None, 2), (4, "not_found", None, None, 2),
            ):
                connection.execute("""INSERT INTO kooda_activation(
                    query_iccid,source_iccid,lookup_status,activation_date,cutoff_date,
                    queried_at,record_schema_version,receipt_json,data_source)
                    VALUES (?,?,?,?,?,?,?,?,?)""", (iccid(number), iccid(number) + "F", status,
                    activation, cutoff, "2026-09-01T00:00:00+08:00", schema, "{}", "grabandgo_complete_catalog_attachment"))
            connection.commit()
        digest = AUDIT.sha256_file(self.document)
        self.pending = DB.prepare_pending(self.database, self.inventory_path, self.pending_path,
                                         document=self.document, document_sha256=digest, run_id=self.run_id)
        self.lookup = {"results": [{"queryIccid": row["queryIccid"], "status": "not_found",
                                    "sourceOccurrences": row["sourceOccurrences"]} for row in self.pending["rows"]]}
        self.receipt = {
            "result": "成功", "output_path": str(self.document), "checks": {"offline-fixture": True},
            "process": {"main_rows": 6, "matched_rows": 6},
            "verification": {
                "unused_port_review": {"matches": True},
                "system_kooda_source_review": {"source": "system", "matches": True, "run_id": self.run_id},
                "task_f_missing_activation_red_review": {"matches": True},
                "task_f_missing_dates_new_card_review": {"matches": True},
            },
        }
        DB.atomic_json(self.receipt_path, self.receipt)
        self.config = {
            "runId": self.run_id, "taskFReportBackfillMode": True,
            "taskFDocumentDateSource": "system", "taskFResultsApplyNextExport": False,
            "taskFResultsWriteBackToDocument": False, "taskFCloudUploadAttempted": False,
            "taskFDocumentPath": str(self.document), "taskFDocumentSha256": digest,
            "taskFDocumentRunId": self.run_id, "taskFDocumentReceipt": str(self.receipt_path),
            "taskFInventoryReceipt": str(self.inventory_path), "taskFReadyForUploadCount": 1,
            "taskFDocumentExcludedCount": 1, "taskFDocumentMissingActivationIccidCount": 3,
        }
        self.resolver = AUDIT.PathResolver(self.root / "config.json", None)

    def review(self):
        return AUDIT.review_task_f_report_backfill(self.config, self.pending, self.inventory, self.lookup, self.resolver)

    def assert_rejected(self, substring):
        errors, _ = self.review()
        self.assertTrue(any(substring in error for error in errors), errors)

    def test_provenance_bound_creation_estimate_passes_independent_audit_and_tamper_fails(self):
        book = load_workbook(self.document)
        sheet = book["sheet1"]
        sheet.cell(1, 6, "截止日")
        sheet.cell(2, 6, "2026-09-30")
        book.save(self.document)
        book.close()
        self.receipt["failed_checks"] = []
        self.receipt["verification"]["system_kooda_source_review"]["snapshot_sha256"] = DB.file_sha256(self.inventory_path)
        self.receipt["verification"]["date_resolution_review"] = {"matches": True, "creation_date_fallback_rows": [
            {"row": 2, "phone": "14385550001", "activation": "2026-09-01", "cutoff": "2026-09-30"}]}
        DB.atomic_json(self.receipt_path, self.receipt)
        digest = DB.file_sha256(self.document)
        self.pending = DB.prepare_pending(self.database, self.inventory_path, self.pending_path,
            document=self.document, document_sha256=digest, run_id=self.run_id,
            document_proof=self.receipt_path, document_proof_sha256=DB.file_sha256(self.receipt_path))
        self.config["taskFDocumentSha256"] = digest
        self.lookup = {"results": [{"queryIccid": row["queryIccid"], "status": "not_found"} for row in self.pending["rows"]]}
        errors, details = self.review()
        self.assertEqual([], errors)
        self.assertEqual(1, details["document_counts"]["documentCreationFallbackRows"])
        self.assertEqual(5, details["document_counts"]["documentLookupNeededRows"])
        # A physically filled report must not be labelled no-gap while the
        # independent provenance still marks a Kooda creation estimate.
        self.receipt["verification"]["system_kooda_source_review"]["missing_either_date_rows"] = 0
        DB.atomic_json(self.receipt_path, self.receipt)
        skip_config = {**self.config, "taskFSkippedReason": "no_missing_authoritative_dates",
            "taskFStageStatus": "skipped", "taskFExecuted": False, "taskFBackground": False,
            "taskFOnlineLookupTriggered": False, "taskFDocumentMissingDateRows": 0}
        skipped = AUDIT.check_task_f_no_missing_dates(skip_config, self.resolver)
        self.assertEqual(AUDIT.FAIL, skipped.status)
        self.assertIn("创建日期兜底行却跳过", skipped.message)
        self.pending["documentProofSha256"] = "0" * 64
        self.assert_rejected("SHA256")

    def test_creation_provenance_date_mismatch_is_rejected_independently(self):
        proof = deepcopy(self.receipt)
        proof["failed_checks"] = []
        proof["verification"]["system_kooda_source_review"]["snapshot_sha256"] = "a" * 64
        proof["verification"]["date_resolution_review"] = {"matches": True, "creation_date_fallback_rows": [
            {"row": 2, "phone": "14385550001", "activation": "2026-09-01", "cutoff": "2026-10-01"}]}
        errors, rows = AUDIT.task_f_creation_fallback_rows(proof, self.run_id, "a" * 64)
        self.assertTrue(errors)
        self.assertEqual({}, rows)

    def test_real_pending_catalog_blank_and_not_found_requery_pass_with_found_reuse(self):
        errors, details = self.review()
        self.assertEqual(errors, [])
        self.assertTrue(details["matches"])
        self.assertEqual({row["previousLookupStatus"] for row in self.pending["rows"]}, {"activation_missing", "not_found"})
        self.assertEqual(self.pending["readyForUploadLegacySchemaCount"], 1)
        self.assertEqual(details["document_counts"]["documentMissingActivationRows"], 4)

    def test_cutoff_only_gap_is_eligible_but_does_not_inflate_activation_gap_count(self):
        workbook = load_workbook(self.document)
        sheet = workbook["sheet1"]
        sheet.cell(1, 6, "截止日")
        for row_number in range(3, 8):
            sheet.cell(row_number, 6, "2026-09-30")
        workbook.save(self.document)
        workbook.close()
        digest = AUDIT.sha256_file(self.document)
        self.pending = DB.prepare_pending(self.database, self.inventory_path, self.pending_path,
                                         document=self.document, document_sha256=digest, run_id=self.run_id)
        self.config["taskFDocumentSha256"] = digest
        self.lookup = {"results": [{"queryIccid": row["queryIccid"], "status": "not_found"}
                                   for row in self.pending["rows"]]}
        errors, details = self.review()
        self.assertEqual([], errors)
        self.assertEqual(1, details["document_counts"]["documentMissingCutoffRows"])
        self.assertEqual(3, details["document_counts"]["documentMissingActivationIccidCount"])
        self.assertEqual(4, details["document_counts"]["documentMissingDateIccidCount"])
        cutoff_only = next(row for row in self.pending["rows"] if row["queryIccid"] == iccid(1))
        self.assertEqual(["截止日"], cutoff_only["sourceOccurrences"][0]["missingFields"])
        cutoff_only["sourceOccurrences"][0]["missingFields"] = ["激活日期"]
        self.assert_rejected("源出现位置或完整系统对象")

    def test_changed_document_bytes_rejected(self):
        with self.document.open("ab") as stream:
            stream.write(b"changed-after-pending")
        self.assert_rejected("SHA256")

    def test_cross_run_and_missing_independent_proof_rejected(self):
        self.pending["documentRunId"] = "prior-round"
        self.assert_rejected("运行编号")
        self.pending["documentRunId"] = self.run_id
        self.receipt_path.unlink()
        self.assert_rejected("独立回读")

    def test_receipt_for_other_document_or_round_rejected(self):
        self.receipt["verification"]["system_kooda_source_review"]["run_id"] = "prior"
        DB.atomic_json(self.receipt_path, self.receipt)
        self.assert_rejected("独立回读回执")

    def test_pending_cannot_add_already_activated_or_inventory_only_iccid(self):
        forged = deepcopy(self.pending["rows"][0])
        forged["queryIccid"] = iccid(1)
        self.pending["rows"].append(forged)
        self.assert_rejected("不属于主表空激活日期")

    def test_pending_cannot_omit_eligible_iccid_or_forge_counts(self):
        self.pending["rows"].pop()
        self.assert_rejected("互斥完整分区")
        self.pending["documentMissingActivationRows"] = 100
        self.assert_rejected("计数不一致")

    def test_complete_source_occurrences_and_backend_identity_are_required(self):
        self.pending["rows"][0]["sourceOccurrences"][0]["sourceData"]["fullRawField"] = "forged"
        self.assert_rejected("完整系统对象")

    def test_date_missing_record_cannot_be_relabelled_as_valid_reuse(self):
        self.pending["alreadyQueried"][0]["activationDate"] = None
        self.pending["alreadyQueried"][0]["cutoffDate"] = None
        self.assert_rejected("无有效激活日期")

    def test_excluded_identity_and_reason_are_audited(self):
        self.pending["documentExcludedRows"][0]["reason"] = ""
        self.assert_rejected("排除行未记录原因")
        self.pending["documentExcludedRows"][0]["phone"] = "19999999999"
        self.assert_rejected("报表排除行")

    def test_upload_and_document_writeback_flags_must_all_be_false(self):
        for name in ("taskFCloudUploadAttempted", "taskFResultsWriteBackToDocument", "taskFResultsApplyNextExport"):
            with self.subTest(name=name):
                self.config[name] = True
                self.assert_rejected("无主表回填及无云端上传")
                self.config[name] = False

    def test_lookup_cannot_include_unscheduled_iccid(self):
        self.lookup["results"][0]["queryIccid"] = iccid(1)
        self.assert_rejected("查询结果ICCID")

    def test_invalid_nonempty_activation_does_not_silently_expand_scope(self):
        workbook = load_workbook(self.document)
        workbook["sheet1"]["D2"] = "invalid-date"
        workbook.save(self.document)
        workbook.close()
        digest = AUDIT.sha256_file(self.document)
        self.pending["documentSha256"] = self.config["taskFDocumentSha256"] = digest
        self.assert_rejected("非空无效激活日期")

    def full_audit_inputs(self):
        fields = ["Purchase Date", "Act. Request Date", "Activation Date", "SIM Number", "SKU"]
        normalization = "strip_exactly_one_final_F_case_insensitive"
        cutoff_rule = "activation_date_plus_29_days_inclusive_30_day_period"
        catalog = {"status": "passed", "sourceRowCount": 3, "normalizedUniqueIccidCount": 3,
                   "foundCount": 1, "activationMissingCount": 2, "invalidCount": 0,
                   "trailingFNormalization": normalization, "activationField": "lblActivateDate", "cutoffRule": cutoff_rule}
        self.inventory.update({"pageSize": 500, "oprSearch": "koo", "fullSourceDataIncluded": True,
                               "rows": [{"sourceData": raw} for raw in self.inventory["systemRows"]],
                               "uniqueIccidCount": 4, "missingCardStoreCount": 1, "invalidCardStoreCount": 0})
        DB.atomic_json(self.inventory_path, self.inventory)
        self.pending["sourceCardsSha256"] = AUDIT.sha256_file(self.inventory_path)
        DB.atomic_json(self.pending_path, self.pending)
        self.lookup.update({"status": "passed", "loginStatus": "authenticated", "inputCount": 2,
                            "mode": "task-f-grabandgo-read-only-webforms-independent-http-sessions",
                            "perWorkerDelayMs": 1000, "silentBrowserMode": True,
                            "browserForegroundActivationAttempted": False,
                            "transportPriority": ["same-origin-http", "background-page-javascript"],
                            "recordSchemaVersion": 2, "recordedWebFields": fields, "fullSourceDataIncluded": True,
                            "requestedWorkers": 8, "effectiveWorkers": 2, "adaptiveConcurrency": True,
                            "sharedSessionForcesSerial": True, "independentAspNetSessions": True,
                            "uniqueAspNetSessionCount": 2, "sharedAspNetSessionConcurrentWorkers": False,
                            "strictIdentityGuard": "queryIccid_equals_listSimNumber_equals_detailSimIccid",
                            "identityMismatchStoredAsTerminal": False, "batchMode": True,
                            "databaseCheckpointAfterEveryBatch": True, "batchCount": 1, "completedBatchCount": 1})
        imported = {"status": "passed", "transientErrorCount": 0, "foundImportedCount": 0,
                    "cutoffRule": cutoff_rule, "recordSchemaVersion": 2, "recordedWebFields": fields,
                    "fullSourceDataIncluded": True, "batchMode": True, "databaseCheckpointAfterEveryBatch": True}
        library = {"status": "passed", "recordCount": 3, "uniqueNormalizedIccidCount": 3,
                   "statusCounts": {"found": 1, "activation_missing": 1, "not_found": 1},
                   "foundCount": 1, "activationMissingCount": 1,
                   "trailingFNormalization": normalization, "activationField": "lblActivateDate",
                   "cutoffRule": cutoff_rule, "includesCompleteSourceData": True}
        for key, value in (("taskFCatalogReceipt", catalog), ("taskFLoginReceipt", {"status": "authenticated"}),
                           ("taskFLookupReceipt", self.lookup), ("taskFImportReceipt", imported),
                           ("taskFIccidLibrary", library), ("taskFLatestIccidLibrary", library)):
            path = self.root / (key + ".json")
            DB.atomic_json(path, value)
            self.config[key] = str(path)
        self.config.update({"taskFExecuted": True, "taskFAutomaticCatalogSync": True,
                            "taskFPostDocumentSupplemental": True, "taskFDocumentGeneratedBeforeLookup": True,
                            "taskFDatabaseSnapshot": str(self.database), "taskFPendingReceipt": str(self.pending_path),
                            "taskFCatalogSourceRowCount": 3, "taskFCatalogFoundCount": 1,
                            "taskFCatalogActivationMissingCount": 2, "taskFInventoryCount": 4,
                            "taskFAlreadyQueriedCount": 1, "taskFPendingCount": 2,
                            "taskFFoundImportedCount": 0, "taskFIccidLibraryRecordCount": 3,
                            "taskFIccidLibraryFoundCount": 1, "taskFIccidLibraryActivationMissingCount": 1,
                            "taskFDatabaseCurrentSchemaCount": 2})

    def test_full_task_f_audit_accepts_new_report_mode_and_legacy_found_reuse(self):
        self.full_audit_inputs()
        result = AUDIT.check_task_f(self.config, self.resolver)
        self.assertNotEqual(result.status, AUDIT.FAIL, result.message)
        self.assertTrue(result.details["report_backfill_review"]["matches"])

    def test_full_task_f_audit_rejects_reuse_disagreeing_with_database_snapshot(self):
        self.full_audit_inputs()
        with DB.open_database(self.database) as connection:
            connection.execute("UPDATE kooda_activation SET activation_date='2026-09-02',cutoff_date='2026-10-01' WHERE query_iccid=?", (iccid(2),))
            connection.commit()
        result = AUDIT.check_task_f(self.config, self.resolver)
        self.assertEqual(result.status, AUDIT.FAIL)
        self.assertIn("复用待人工上传日期与本轮数据库快照不一致", result.message)

    def test_new_gate_is_called_only_for_report_mode_preserving_standalone_path(self):
        self.config.update({"taskFExecuted": True, "taskFCurrentUserAuthorized": True,
                            "taskFDatabaseSnapshot": str(self.database)})
        for key in ("taskFCatalogReceipt", "taskFLoginReceipt", "taskFLookupReceipt", "taskFImportReceipt", "taskFIccidLibrary", "taskFLatestIccidLibrary"):
            path = self.root / (key + ".json")
            DB.atomic_json(path, {})
            self.config[key] = str(path)
        self.config["taskFPendingReceipt"] = str(self.pending_path)
        with patch.object(AUDIT, "review_task_f_report_backfill", return_value=(["offline-new-gate-marker"], {})) as reviewed:
            result = AUDIT.check_task_f(self.config, self.resolver)
            reviewed.assert_called_once()
            self.assertIn("offline-new-gate-marker", result.message)
        self.config["taskFReportBackfillMode"] = False
        self.pending["reportBackfillMode"] = False
        DB.atomic_json(self.pending_path, self.pending)
        with patch.object(AUDIT, "review_task_f_report_backfill") as reviewed:
            AUDIT.check_task_f(self.config, self.resolver)
            reviewed.assert_not_called()


if __name__ == "__main__":
    unittest.main()
