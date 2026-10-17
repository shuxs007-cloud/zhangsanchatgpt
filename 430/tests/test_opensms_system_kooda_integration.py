"""Synthetic XLSX -> local builder -> complete independent verifier; no network."""
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "services/opensms_document_api/pipeline"
SPEC = importlib.util.spec_from_file_location("system_kooda_integration_audit", ROOT / "scripts/opensms_parallel_audit.py")
AUDIT = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = AUDIT
SPEC.loader.exec_module(AUDIT)


class SystemKoodaIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.raw = self.root / "source.xlsx"
        self.google = self.root / "google.xlsx"
        self.system = self.root / "system.json"
        self.cache = self.root / "report-date-cache.json"
        self.output = self.root / "built" / "document.xlsx"
        run_id = "20260907T010000+0800"

        book = Workbook()
        sheet = book.active
        sheet.title = "sheet1"
        headers = ["电话", "运营商", "opr", "卡槽", "创建时间", "送达时间", "接收时间", "状态",
                   "10日发送窗口数", "10日回复窗口数", "10日回复率", "10日总发送量", "10日总回复量"]
        for day in ("09-03", "09-04", "09-05", "09-06"):
            headers.extend(day + suffix for suffix in ("发送窗口数", "回复窗口数", "回复率", "总发送量", "总回复量"))
        sheet.append(headers)
        for phone, opr, slot, status, created in (
            ("14385550123", "Kooda", "130_1", "上架", "2026-09-01 00:00:00"),
            ("14385550124", "Telus1", "130_2", "注册失败", "2020-01-01 00:00:00"),
            ("14385550125", "Chatr", "130_3", "上架", "2026-09-06 00:00:00"),
        ):
            sheet.append([phone, "A1" if opr == "Chatr" else "B1", opr, slot, created, None, None,
                          status, 400, 40, 0.1, 400, 40] + [100, 10, 0.1, 100, 10] * 4)
        book.save(self.raw)
        book.close()

        book = Workbook()
        book.remove(book.active)
        for title in ("二號機房", "三號機房", "四號機房"):
            sheet = book.create_sheet(title)
            sheet.append(["卡槽", "電信商/價格", "上卡日期", "激活日期", "預計過期日", "備註", "iccid", "手機號"])
            if title == "四號機房":
                # Both Kooda rows deliberately have no Google card at all.
                sheet.append(["130_3", "Chatr", "2026-08-28", "2026-09-02", "2026-09-20",
                              "Chatr保留Google来源", "GOOGLE-CHATR-ICCID", "14385550125"])
        book.save(self.google)
        book.close()

        payload = {
            "status": "passed", "runId": run_id, "totalRows": 2,
            "paginationTotalStable": True, "paginationBackendIdsUnique": True,
            "silentBrowserMode": True, "browserForegroundActivationAttempted": False,
            "systemRows": [
                {"id": "system-1", "slotIdentif": "130_01.01", "cardNo": "4385550123", "opr": "Kooda",
                 "cardStore": "8912230200081836799F", "chatrActivateDate": "2026-09-01", "chatrPlanExpire": "2026-12-15"},
                {"id": "system-2", "slotIdentif": "130_2", "cardNo": "4385550124", "opr": "Telus1",
                 "cardStore": "8912230200081836798f", "chatrActivateDate": None, "chatrPlanExpire": None},
            ],
        }
        encoded = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.system.write_bytes(encoded)
        cache_bytes = json.dumps({"schemaVersion": 1, "status": "passed", "runId": run_id,
            "createdAt": "2026-09-07T01:00:00+08:00", "sourceDatabase": str(self.root / "never-read-task-f.sqlite"),
            "sourceExists": False, "recordCount": 0, "records": []}).encode()
        self.cache.write_bytes(cache_bytes)
        self.env = {
            **os.environ,
            "PYTHON_BIN": sys.executable,
            "NODE_BIN": "/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node",
            "OPENSMS_SYSTEM_KOODA_PATH": str(self.system),
            "OPENSMS_SYSTEM_KOODA_SHA256": hashlib.sha256(encoded).hexdigest(),
            "OPENSMS_RUN_ID": run_id,
            "OPENSMS_REPORT_DATE_CACHE_PATH": str(self.cache),
            "OPENSMS_REPORT_DATE_CACHE_SHA256": hashlib.sha256(cache_bytes).hexdigest(),
            "OPENSMS_ASSOC_PATH": str(self.google),
            "OPENSMS_ASSOC_SHA256": hashlib.sha256(self.google.read_bytes()).hexdigest(),
            "OPENSMS_OUTPUT_PATH": str(self.output),
            "OPENSMS_RECEIPT_DIR": str(self.root / "receipts"),
            "OPENSMS_TASK_F_DB_PATH": str(self.root / "never-read-task-f.sqlite"),
            "OPENSMS_LEGACY_DOCUMENT_ENRICHMENT": "0",
        }

    def build(self):
        completed = subprocess.run(
            ["bash", str(PIPELINE / "build_card_slot_report.sh"), str(self.raw), str(self.google),
             str(self.output.parent), "2026-09-06"],
            env=self.env, text=True, capture_output=True, timeout=45,
        )
        self.assertEqual(completed.returncode, 0, (completed.stderr + completed.stdout)[-8000:])
        receipt = json.loads(completed.stdout)
        self.assertEqual(receipt["result"], "成功")
        self.assertEqual(receipt["failed_checks"], [])
        return receipt

    def test_archived_off_shelf_row_builds_and_independent_verifier_rejects_wrong_proof(self):
        book = load_workbook(self.raw)
        sheet = book['sheet1']
        col = sheet.max_column + 1
        sheet.cell(1, col, 'ICCID')
        sheet.cell(2, col, '8912230200081836799F')
        sheet.cell(3, col, '8912230200081836798f')
        sheet.cell(2, 8, '下架')
        book.save(self.raw); book.close()
        payload = json.loads(self.system.read_text())
        payload['systemRows'][0].update(slotIdentif=None, showAble=False, status='换卡已下线',
            machineName='130', createTime='2026-09-01 00:00:00')
        self.system.write_text(json.dumps(payload))
        self.env['OPENSMS_SYSTEM_KOODA_SHA256'] = hashlib.sha256(self.system.read_bytes()).hexdigest()
        receipt = self.build()
        self.assertTrue(receipt['verification']['system_kooda_source_review']['matches'])
        self.assertEqual(receipt['verification']['system_kooda_source_review']['matched_rows'], 2)
        spec = importlib.util.spec_from_file_location('archive_independent_verify', PIPELINE / 'opensms_verify_card_slot_fast.py')
        verifier = importlib.util.module_from_spec(spec); spec.loader.exec_module(verifier)
        snapshot = verifier.load_system_kooda_snapshot(self.system, self.env['OPENSMS_SYSTEM_KOODA_SHA256'], self.env['OPENSMS_RUN_ID'])
        raw = {'电话': '14385550123', '卡槽': '130_1', '状态': '下架',
               '创建时间': '2026-09-01 00:00:00', 'ICCID': '8912230200081836799F'}
        self.assertIsNotNone(verifier.find_system_record(snapshot, raw))
        for key, wrong in [('状态', '上架'), ('创建时间', '2026-09-02 00:00:00'),
                           ('卡槽', '131_1'), ('ICCID', '8912230200081836798F')]:
            self.assertIsNone(verifier.find_system_record(snapshot, {**raw, key: wrong}))

    def test_complete_builder_and_verifier_use_system_without_kooda_google_match(self):
        source_hashes = [hashlib.sha256(path.read_bytes()).hexdigest() for path in (self.raw, self.google, self.system)]
        receipt = self.build()
        self.assertEqual(receipt["process"]["main_rows"], 3)
        self.assertEqual(receipt["process"]["matched_rows"], 3)
        self.assertEqual(receipt["process"]["google_matched_rows"], 1)
        self.assertEqual(receipt["verification"]["system_kooda_source_review"]["matched_rows"], 2)
        self.assertTrue(receipt["verification"]["system_kooda_source_review"]["matches"])
        self.assertEqual(receipt["verification"]["task_f_missing_activation_red_review"]["expected_rows"], 1)
        self.assertTrue(receipt["verification"]["task_f_missing_activation_red_review"]["matches"])
        self.assertEqual(receipt["verification"]["task_f_missing_dates_new_card_review"]["eligible_rows"], 1)
        self.assertTrue(receipt["verification"]["task_f_missing_dates_new_card_review"]["matches"])

        book = load_workbook(self.output)
        try:
            sheet = book["sheet1"]
            index = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(1, sheet.max_column).value, "ICCID")
            self.assertEqual(sheet.cell(2, index["ICCID"]).value, "8912230200081836799F")
            self.assertIsNone(sheet.cell(2, index["开卡日期"]).value)
            self.assertEqual(sheet.cell(2, index["截止日"]).value.date().isoformat(), "2026-12-15")
            self.assertEqual(sheet.cell(3, index["换号状态"]).value, "未换号")
            self.assertTrue(all(str(cell.fill.fgColor.rgb).upper().endswith("FFC7CE") for cell in sheet[3]))
            self.assertEqual(sheet.cell(4, index["开卡日期"]).value.date().isoformat(), "2026-08-28")
            self.assertEqual(sheet.cell(4, index["激活日期"]).value.date().isoformat(), "2026-09-02")
            self.assertEqual(sheet.cell(4, index["截止日"]).value.date().isoformat(), "2026-09-20")
        finally:
            book.close()
        self.assertFalse(Path(self.env["OPENSMS_TASK_F_DB_PATH"]).exists())
        self.assertEqual(source_hashes, [hashlib.sha256(path.read_bytes()).hexdigest() for path in (self.raw, self.google, self.system)])
        # Google-only missing slots are warnings, not rejection of system-owned Kooda rows.
        self.assertNotEqual(AUDIT.check_card_receipt(Path(receipt["receipt_path"])).status, AUDIT.FAIL)

    def test_complete_verifier_detects_tampered_system_date_and_new_card_status(self):
        receipt = self.build()
        book = load_workbook(self.output)
        sheet = book["sheet1"]
        index = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, index["截止日"]).value = "2026-09-30"  # Incorrect F-style +29 fallback.
        sheet.cell(3, index["换号状态"]).value = "已换号"
        book.save(self.output)
        book.close()
        completed = subprocess.run(
            [sys.executable, str(PIPELINE / "opensms_verify_card_slot_fast.py"), str(self.output), str(self.raw)],
            env={**self.env, "PROCESS_JSON": json.dumps(receipt["process"]),
                 "ENRICH_JSON": json.dumps(receipt["enrichment"]), "EXPORT4DAY_SERVER_DATE": "2026-09-06"},
            text=True, capture_output=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        verification = json.loads(completed.stdout)
        self.assertFalse(verification["system_kooda_source_review"]["matches"])
        self.assertFalse(verification["task_f_missing_dates_new_card_review"]["matches"])
        changed_receipt = {**receipt, "verification": verification}
        changed_path = self.root / "tampered-receipt.json"
        changed_path.write_text(json.dumps(changed_receipt), encoding="utf-8")
        self.assertEqual(AUDIT.check_card_receipt(changed_path).status, AUDIT.FAIL)

    def test_document_only_cli_publishes_full_verified_receipt_and_report(self):
        env = {key: value for key, value in self.env.items()
               if key not in ("OPENSMS_OUTPUT_PATH", "OPENSMS_RECEIPT_DIR")}
        completed = subprocess.run(
            [sys.executable, str(ROOT / "scripts/opensms_generate_from_downloads.py"),
             str(self.raw), str(self.google), "--backend-date", "2026-09-06",
             "--system-kooda", str(self.system),
             "--system-kooda-sha256", self.env["OPENSMS_SYSTEM_KOODA_SHA256"],
             "--run-id", self.env["OPENSMS_RUN_ID"],
             "--output-dir", str(self.root / "local-cli"),
             "--artifact-root", str(self.root / "local-receipts")],
            env=env, text=True, capture_output=True, timeout=45,
        )
        self.assertEqual(completed.returncode, 0, (completed.stderr + completed.stdout)[-8000:])
        result = json.loads(completed.stdout)
        self.assertEqual(result["result"], "成功")
        receipt = json.loads(Path(result["receiptPath"]).read_text())
        card = receipt["cardSlotReport"]
        self.assertTrue(card["checks"])
        self.assertTrue(all(value is True for value in card["checks"].values()))
        self.assertEqual(card["failed_checks"], [])
        self.assertNotIn("google_matched_rows", card["process"])
        self.assertEqual(card["process"]["google_dependency"], {
            "enabled": False, "status": "skipped", "reason": "user_disabled",
            "unused_port_device_scope": "system_source_only",
        })
        self.assertTrue(card["verification"]["google_dependency_review"]["matches"])
        self.assertIsNone(receipt["inputs"]["association"])
        self.assertEqual(card["verification"]["system_kooda_source_review"]["matched_rows"], 2)
        self.assertFalse(receipt["liveBackendActionsExecuted"])
        self.assertTrue(Path(result["outputPath"]).is_file())
        self.assertIn("系统cardStore", Path(result["reportPath"]).read_text())

    def test_document_only_cli_freezes_empty_database_when_no_snapshot_supplied(self):
        env = {key: value for key, value in self.env.items()
               if key not in ("OPENSMS_OUTPUT_PATH", "OPENSMS_RECEIPT_DIR",
                              "OPENSMS_REPORT_DATE_CACHE_PATH", "OPENSMS_REPORT_DATE_CACHE_SHA256")}
        run = subprocess.run([sys.executable, str(ROOT / "scripts/opensms_generate_from_downloads.py"),
            str(self.raw), str(self.google), "--backend-date", "2026-09-06",
            "--system-kooda", str(self.system), "--system-kooda-sha256", self.env["OPENSMS_SYSTEM_KOODA_SHA256"],
            "--run-id", self.env["OPENSMS_RUN_ID"], "--task-f-db", self.env["OPENSMS_TASK_F_DB_PATH"],
            "--output-dir", str(self.root / "local-cli-freeze"), "--artifact-root", str(self.root / "local-cli-receipts")],
            env=env, capture_output=True, text=True, timeout=45)
        self.assertEqual(run.returncode, 0, run.stderr[-4000:])
        result = json.loads(run.stdout)
        receipt = json.loads(Path(result["receiptPath"]).read_text())
        snapshot = receipt["inputs"]["reportDateCache"]
        self.assertFalse(snapshot["provided"])
        self.assertFalse(snapshot["sourceExists"])
        self.assertEqual(snapshot["recordCount"], 0)
        self.assertEqual(json.loads(Path(snapshot["path"]).read_text())["runId"], self.env["OPENSMS_RUN_ID"])
        self.assertTrue(receipt["cardSlotReport"]["verification"]["date_resolution_review"]["matches"])
        self.assertFalse(Path(self.env["OPENSMS_TASK_F_DB_PATH"]).exists())


if __name__ == "__main__":
    unittest.main()
