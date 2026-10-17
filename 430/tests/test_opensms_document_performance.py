import unittest
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection

from services.opensms_document_api.pipeline.process_card_slots import fill_row_with_registered_fill


class RegisteredFillTests(unittest.TestCase):
    def workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["电话", "发送", "截止日", "公式", "未创建单元格"])
        sheet.append(["001234567890", 49, "2026-09-05", "=B2+1"])
        sheet.append(["001234567891", 50, None, False])
        sheet["A2"].number_format = "@"
        sheet["A2"].font = Font(bold=True, color="003366")
        sheet["B2"].alignment = Alignment(wrap_text=True)
        sheet["D2"].protection = Protection(locked=False)
        sheet["C2"].fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
        # Loaded workbooks can share a StyleArray. Preserve descriptor semantics.
        sheet["A3"]._style = sheet["A2"]._style
        return workbook

    def test_registered_fill_matches_descriptor_and_preserves_other_style_fields(self):
        old, new = self.workbook(), self.workbook()
        fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        for row in (2, 3):
            for col in range(1, 6):
                old.active.cell(row, col).fill = fill
        fill_id = new._fills.add(fill)
        for row in (2, 3):
            fill_row_with_registered_fill(new.active, row, 5, fill_id)
        for old_row, new_row in zip(old.active.iter_rows(), new.active.iter_rows()):
            for before, after in zip(old_row, new_row):
                self.assertEqual(before.value, after.value)
                self.assertEqual(before.data_type, after.data_type)
                self.assertEqual(before._style, after._style)
        archives = []
        for workbook in (old, new):
            stream = BytesIO()
            workbook.save(stream)
            archives.append(ZipFile(stream))
        self.assertEqual(archives[0].read("xl/worksheets/sheet1.xml"), archives[1].read("xl/worksheets/sheet1.xml"))
        self.assertEqual(archives[0].read("xl/styles.xml"), archives[1].read("xl/styles.xml"))


if __name__ == "__main__":
    unittest.main()
