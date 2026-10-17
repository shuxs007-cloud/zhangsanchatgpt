"""Offline fixtures only: restored difference never changes main-sheet routing inputs."""
import importlib.util
import contextlib
import io
import json
import os
from pathlib import Path
import sys
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "services/opensms_document_api/pipeline"))
from google_system_missing import SHEET, HEADERS, PORT_SHEET, PORT_HEADERS, load_from_env, read_rows
from opensms_chatr_google_dates import freeze
import opensms_parallel_audit as audit


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


FIXTURE = module("restored_google_fixture", ROOT / "tests/test_opensms_document_no_google.py")


class RestoredGoogleDifferenceTests(unittest.TestCase):
    def setUp(self):
        self.f = FIXTURE.NoGoogleDocumentTests()
        self.f.setUp()
        self.addCleanup(self.f.doCleanups)
        self.google, self.snapshot = self.f.root / "google.xlsx", self.f.root / "snapshot.json"
        book = Workbook()
        book.remove(book.active)
        for name in ("二號機房", "三號機房", "四號機房"):
            sheet = book.create_sheet(name)
            sheet.append(["卡槽", "電信商/價格", "激活日期", "預計過期日", "ICCID", "上卡日期", "備註"])
        sheet = book.worksheets[0]
        for values in (
            ["130_01.01", "Kooda", "2026-09-01", None, "9999999999999999999", None, "注册失败也算系统存在"],
            ["130_7", "Kooda", None, None, None, None, "下架也算系统存在"],
            ["888_1", "Kooda", "2026-09-01", None, "8912230200081836701F", "2026/09/01", "=1+1"],
            ["888_01.01", "Kooda", None, None, "0001230200081836701", None, "重复保留"],
            ["r130_1", "Kooda", "20260901", None, None, None, "字母设备不能并入130"],
            [None, "Kooda", None, None, None, None, "没有卡槽无法比较"],
        ):
            sheet.append(values)
            sheet.cell(sheet.max_row, 7).data_type = "s"
        book.worksheets[2].append(["151_2", "Kooda", None, None, None, None, "历史机房页签继续读取"])
        book.save(self.google)
        book.close()
        freeze(self.google, self.f.digest(self.google), self.snapshot, self.f.env["OPENSMS_RUN_ID"])
        self.f.env.update(OPENSMS_CHATR_GOOGLE_DATES_PATH=str(self.snapshot),
                          OPENSMS_CHATR_GOOGLE_DATES_SHA256=self.f.digest(self.snapshot),
                          OPENSMS_GOOGLE_SYSTEM_MISSING_ENABLED="1")

    def config(self, receipt):
        card = self.f.root / "receipt.json"
        card.write_text(json.dumps(receipt))
        proof = receipt["process"]["google_system_missing_sheet"]
        config = {"runId": self.f.env["OPENSMS_RUN_ID"], "googleAssociationEnabled": False,
                  "googleChatrDateFallbackEnabled": True, "googleSystemMissingSheetEnabled": True,
                  "googleSystemMissingSheet": proof, "chatrGoogleSourceSha256": self.f.digest(self.google),
                  "chatrGoogleDatesSha256": self.f.digest(self.snapshot)}
        return config, {"card_receipt": card, "source": self.f.raw, "result": self.f.output}

    def test_exact_difference_duplicates_formatting_and_final_audit(self):
        receipt = self.f.build()
        proof = receipt["process"]["google_system_missing_sheet"]
        self.assertEqual((proof["missing_rows"], proof["missing_unique_slots"], proof["activation_highlighted_rows"]), (4, 3, 2))
        self.assertEqual(proof["source_rows"], 6)
        self.assertEqual(proof["blank_slot_rows"], 1)
        self.assertTrue(receipt["verification"]["google_system_missing_sheet_review"]["matches"])
        self.assertTrue(receipt["verification"]["google_dependency_review"]["matches"])
        book = load_workbook(self.f.output)
        self.addCleanup(book.close)
        sheet = book[SHEET]
        self.assertEqual(tuple(c.value for c in sheet[1]), HEADERS)
        self.assertEqual([sheet.cell(r, 4).value for r in range(2, 6)], ["888_1", "888_1", "r130_1", "151_2"])
        self.assertEqual(sheet["E2"].value, "=1+1")
        self.assertEqual(sheet["E2"].data_type, "s")
        self.assertEqual(sheet["F3"].value, "0001230200081836701")
        self.assertEqual(sheet["F3"].data_type, "s")
        self.assertFalse((FIXTURE.FORBIDDEN - {SHEET}) & set(book.sheetnames))
        self.assertEqual(receipt["process"]["unused_port_devices"], ["130"])
        self.assertEqual(book.sheetnames.index(SHEET), book.sheetnames.index(PORT_SHEET) + 1)
        self.assertEqual(tuple(c.value for c in book[PORT_SHEET][1]), PORT_HEADERS)
        self.assertNotIn("未使用端口", book.sheetnames)
        self.assertNotIn("所有谷歌有-系统没有", book.sheetnames)
        config, inputs = self.config(receipt)
        self.assertEqual(audit.check_google_difference(config, inputs).status, audit.PASS)
        structure = audit.inspect_result_workbook(self.f.output, "2026-09-06", False, True)
        self.assertFalse(structure["missing_sheets"])
        self.assertTrue(audit.inspect_result_workbook(self.f.output, "2026-09-06", False)["order_errors"])

    def test_toggle_does_not_change_main_dates_routing_or_unlisted_sheets(self):
        first = self.f.build()
        book = load_workbook(self.f.output, data_only=True)
        before = {name: list(book[name].values) for name in ("sheet1", "已激活未上架", "未上架机器汇总", "系统无-谷歌表有", "设备每日统计")}
        book.close()
        self.f.output.unlink()
        self.f.env["OPENSMS_GOOGLE_SYSTEM_MISSING_ENABLED"] = "0"
        second = self.f.build()
        book = load_workbook(self.f.output, data_only=True)
        self.addCleanup(book.close)
        self.assertNotIn(SHEET, book.sheetnames)
        for name, values in before.items():
            if name == PORT_SHEET:
                # Disabling Google changes only its new ICCID evidence column;
                # existing eight-column port coverage stays identical.
                self.assertEqual([r[:8] for r in book[name].values], [r[:8] for r in values], name)
                self.assertTrue(all(not r[8] for r in list(book[name].values)[1:]))
            else:
                self.assertEqual(list(book[name].values), values, name)
        self.assertEqual(first["process"]["date_resolution"], second["process"]["date_resolution"])

    def test_verifiers_reject_tampered_counts_values_types_highlights_and_missing_sheet(self):
        receipt = self.f.build()
        original = self.f.output.read_bytes()
        for kind in ("value", "type", "highlight", "deleted_sheet", "extra_row", "wrong_flag", "flag_formula", "removed_flag_column", "moved_sheet", "legacy_sheet"):
            with self.subTest(kind=kind):
                self.f.output.write_bytes(original)
                book = load_workbook(self.f.output)
                if kind == "value": book[SHEET]["D2"] = "130_1"
                elif kind == "type": book[SHEET]["F2"] = 8912230200081836701
                elif kind == "highlight": book[SHEET]["A2"].fill = PatternFill()
                elif kind == "deleted_sheet": del book[SHEET]
                elif kind == "wrong_flag": book[PORT_SHEET]["H2"] = "谷歌表存在"
                elif kind == "flag_formula": book[PORT_SHEET]["H2"] = '=""'
                elif kind == "removed_flag_column": book[PORT_SHEET].delete_cols(8)
                elif kind == "moved_sheet": book.move_sheet(SHEET, -1)
                elif kind == "legacy_sheet": book.create_sheet("未使用端口")
                else: book[SHEET].append(["fake"])
                book.save(self.f.output)
                book.close()
                self.assertFalse(self.f.verify(receipt)["google_system_missing_sheet_review"]["matches"])
                self.assertEqual(audit.check_google_difference(*self.config(receipt)).status, audit.FAIL)
        self.f.output.write_bytes(original)
        receipt["process"]["google_system_missing_sheet"]["missing_rows"] = 999
        self.assertFalse(self.f.verify(receipt)["google_system_missing_sheet_review"]["matches"])

    def test_presence_uses_fresh_google_slots_and_independent_audit(self):
        book = load_workbook(self.google)
        for values in (
            ["１３０-０８.01", "Chatr", None, None, None, None, "无日期也标记"],
            ["130_8", "Kooda", None, None, None, None, "同槽重复只标一个端口"],
            ["130_2", "Kooda", None, None, None, None, "失败和上架混合仍属系统有"],
            ["r130_9", "Kooda", None, None, None, None, "不能标到130_9"],
        ):
            book.worksheets[0].append(values)
        book.save(self.google)
        book.close()
        self.snapshot.unlink()
        freeze(self.google, self.f.digest(self.google), self.snapshot, self.f.env["OPENSMS_RUN_ID"])
        self.f.env["OPENSMS_CHATR_GOOGLE_DATES_SHA256"] = self.f.digest(self.snapshot)
        receipt = self.f.build()
        self.assertEqual(receipt["process"]["system_missing_google_present_count"], 1)
        self.assertTrue(receipt["verification"]["unused_port_review"]["matches"])
        book = load_workbook(self.f.output)
        by_slot = {row[3]: row for row in book[PORT_SHEET].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(by_slot["130_8"][7], "谷歌表存在")
        self.assertIsNone(by_slot["130_1"][7])
        self.assertIsNone(by_slot["130_9"][7])
        self.assertNotIn("130_2", by_slot)
        self.assertNotIn("130_7", by_slot)
        self.assertFalse(any(key.startswith("888_") for key in by_slot))
        self.assertEqual(audit.check_google_difference(*self.config(receipt)).status, audit.PASS)
        target = next(row for row in book[PORT_SHEET].iter_rows(min_row=2) if row[3].value == "130_8")
        target[7].value = None
        book.save(self.f.output)
        book.close()
        self.assertFalse(self.f.verify(receipt)["unused_port_review"]["matches"])
        self.assertEqual(audit.check_google_difference(*self.config(receipt)).status, audit.FAIL)

    def test_source_run_hash_and_missing_header_fail_closed(self):
        with mock.patch.dict(os.environ, {**self.f.env, "OPENSMS_RUN_ID": "other"}), self.assertRaisesRegex(ValueError, "本轮"):
            load_from_env()
        with mock.patch.dict(os.environ, {**self.f.env, "OPENSMS_CHATR_GOOGLE_DATES_SHA256": "0" * 64}), self.assertRaisesRegex(ValueError, "SHA256"):
            load_from_env()
        self.google.write_bytes(self.google.read_bytes() + b"changed")
        with mock.patch.dict(os.environ, self.f.env), self.assertRaisesRegex(ValueError, "SHA256"):
            load_from_env()

    def test_missing_header_sheet_and_uncached_formula_rejected(self):
        original = self.google.read_bytes()
        for kind in ("missing_header", "missing_sheet", "uncached_formula"):
            with self.subTest(kind=kind):
                self.google.write_bytes(original)
                book = load_workbook(self.google)
                if kind == "missing_header": book.worksheets[0]["G1"] = "unknown"
                elif kind == "missing_sheet": del book["四號機房"]
                else: book.worksheets[0]["C4"] = "=TODAY()"
                book.save(self.google)
                book.close()
                with self.assertRaises(ValueError):
                    read_rows(self.google.read_bytes())

    def test_actual_carrier_refresh_preserves_difference_and_final_review(self):
        receipt = self.f.build()
        refresh = module("difference_refresh_fixture", ROOT / "tests/test_opensms_refresh_workbook_carriers_fast.py")
        fixture = refresh.FastCarrierRefreshTests()
        book = load_workbook(self.f.output)
        main = book.worksheets[0]
        headers = [cell.value for cell in main[1]]
        phone = str(main.cell(2, headers.index("电话") + 1).value)
        before = list(book[SHEET].values)
        book.close()
        args = fixture.fresh_args(self.f.root, self.f.output, "restored-difference", [
            {"phone": phone, "liveCarrierName": "离线核验运营商", "route": "offline", "expectedTarget": "离线核验运营商"}])
        result = refresh.MODULE.run(args)
        self.assertEqual(result["changed"], 1)
        config, inputs = self.config(receipt)
        inputs["result"] = Path(result["outputPath"])
        self.assertEqual(audit.check_google_difference(config, inputs).status, audit.PASS)
        book = load_workbook(inputs["result"])
        self.addCleanup(book.close)
        self.assertEqual(list(book[SHEET].values), before)
        self.assertEqual(book.sheetnames.index(SHEET), book.sheetnames.index(PORT_SHEET) + 1)
        self.assertEqual(tuple(cell.value for cell in book[PORT_SHEET][1]), PORT_HEADERS)

    def test_empty_difference_keeps_header_and_zero_counts(self):
        book = load_workbook(self.google)
        for sheet in book:
            if sheet.max_row > 1:
                sheet.delete_rows(2, sheet.max_row - 1)
        book.worksheets[0].append(["130_1", "Kooda", "2026-09-01", None, None, None, None])
        book.save(self.google)
        book.close()
        self.snapshot.unlink()
        freeze(self.google, self.f.digest(self.google), self.snapshot, self.f.env["OPENSMS_RUN_ID"])
        self.f.env["OPENSMS_CHATR_GOOGLE_DATES_SHA256"] = self.f.digest(self.snapshot)
        receipt = self.f.build()
        self.assertEqual(receipt["process"]["google_system_missing_sheet"]["missing_rows"], 0)
        book = load_workbook(self.f.output)
        self.addCleanup(book.close)
        self.assertEqual(book[SHEET].max_row, 1)

    def test_report_discloses_difference_scope_and_counts(self):
        import test_opensms_no_google_audit_report as report_test
        fixture = report_test.NoGoogleAuditReportTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        proof = {"missing_rows": 4, "missing_unique_slots": 3, "activation_highlighted_rows": 2}
        fixture.card["process"]["google_system_missing_sheet"] = proof
        fixture.card["verification"]["google_system_missing_sheet_review"] = {**proof, "matches": True}
        fixture.save_receipts()
        fixture.config.update({"googleChatrDateFallbackEnabled": True, "googleSystemMissingSheetEnabled": True,
            "taskAOnly": True, "taskSelection": "A", "taskFStageStatus": "skipped", "taskFSkippedReason": "task_a_only",
            "offlineRecoveryEnabled": False, "runTime": "offline", "taskId": "offline", "taskDuration": "1s", "exportWorkflow": "fixture"})
        report_test.write_json(fixture.config_path, fixture.config)
        output = fixture.root / "restored-difference-report.md"
        with mock.patch.object(sys, "argv", ["report", str(fixture.config_path), str(output)]), contextlib.redirect_stdout(io.StringIO()):
            report_test.REPORT.main()
        rendered = output.read_text()
        for value in ("谷歌有-系统无：4 条Google原记录", "3 个唯一卡槽", "标黄 2 行", "独立回读通过",
                      "按归一化卡槽与本轮系统全部状态记录比较", "不参与日期补值或号码操作"):
            self.assertIn(value, rendered)


if __name__ == "__main__":
    unittest.main()
