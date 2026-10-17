import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import opensms_task_f_slot_history as scanner

KEY = '8912230000000000001'


class HistoryTests(unittest.TestCase):
    def test_iccid_strict(self):
        self.assertEqual(scanner.normalize_iccid(KEY + 'F'), KEY)
        self.assertIsNone(scanner.normalize_iccid(int(KEY)))
        self.assertIsNone(scanner.normalize_iccid(KEY + 'FF'))
        self.assertIsNone(scanner.normalize_iccid('8.91223E18'))

    def test_filename_date_wins(self):
        day, basis = scanner.file_date(Path('/archive/20260824/a.xlsx'),
                                      type('Stat', (), {'st_mtime': 1790200000})())
        self.assertEqual(day, '2026-08-24')

    def make_and_scan(self, values):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.xlsx'
            wb = Workbook()
            ws = wb.active
            for row in values:
                ws.append(row)
            wb.save(path)
            stat = path.stat()
            item = scanner.fingerprint({'path': str(path), 'size': stat.st_size,
                                        'mtimeNs': stat.st_mtime_ns})
            return scanner.scan_workbook(item, {KEY})

    def test_text_slot_and_f(self):
        result = self.make_and_scan([['ICCID', '卡槽'], [KEY+'F', 'r112_45.01']])
        self.assertNotIn('error', result)
        self.assertEqual(result['evidence'][0]['slot'], 'r112_45.01')
        self.assertEqual(result['evidence'][0]['iccidCell'], 'A2')

    def test_no_headers_not_recoverable(self):
        result = self.make_and_scan([[KEY, 'r112_45.01']])
        self.assertFalse(result['evidence'])
        self.assertEqual(len(result['unresolved']), 1)

    def test_formula_and_partial_slot(self):
        result = self.make_and_scan([['ICCID', '卡槽'], [KEY, '45.01'], ['="'+KEY+'"', '112_45.01']])
        self.assertFalse(result['evidence'])

    def test_hidden_sheets(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)/'hidden.xlsx'
            w = Workbook(); s = w.create_sheet('hidden'); s.sheet_state='hidden'
            s.append(['SIM Number', '卡槽标识']);s.append([KEY, '112_1.01']);w.save(p)
            st = p.stat()
            result = scanner.scan_workbook(scanner.fingerprint({'path':str(p),'size':st.st_size,'mtimeNs':st.st_mtime_ns}),{KEY})
            self.assertEqual(result['evidence'][0]['sheet'],'hidden')

    def test_conflict_not_r_prefix_merge(self):
        target = {KEY: {'inventorySlot': ''}}
        ev = [{'iccid':KEY,'slot':slot,'ambiguousColumns':False,'sourceKind':'系统',
               'date':'2026-09-23','file':'a.xlsx','sheet':'sheet1','row':2} for slot in ['112_1.01','r112_1.01']]
        self.assertEqual(scanner.summarize(target,ev)[0]['status'],'多卡槽记录待核对')

    def test_shared_strings_and_numeric_rejection(self):
        xml = (f'<worksheet xmlns="{scanner.NS}"><sheetData>'
               '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>'
               '<row r="2"><c r="A2" t="s"><v>2</v></c><c r="B2" t="s"><v>3</v></c></row>'
               f'<row r="3"><c r="A3"><v>{KEY}</v></c><c r="B3" t="s"><v>3</v></c></row>'
               '</sheetData></worksheet>').encode()
        result = scanner.scan_sheet(xml, ['ICCID', '卡槽', KEY+'F', '00112_01.01'], 'test1', {KEY})
        self.assertEqual([(r['row'],r['slot']) for r in result['evidence']],[(2,'00112_01.01')])
        self.assertEqual(result['rejected']['ICCID非可靠文本或公式'],1)

    def test_multi_identity_columns_marked_ambiguous(self):
        result = self.make_and_scan([['系统ICCID', 'GoogleICCID', '卡槽'],[KEY,KEY,'112_1.01']])
        self.assertEqual(len(result['evidence']),2)
        self.assertTrue(all(r['ambiguousColumns'] for r in result['evidence']))

    def test_output_readback_no_formula_injection(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)/'result.xlsx'
            rows=[{'iccid':KEY,'status':'未找到','slots':[],'latestFileDate':'',
                   'inventorySlotBeforeSearch':'','latestFile':'=1+1','latestSheet':''}]
            digest=scanner.write_excel(p,rows,[],[],[])
            self.assertEqual(len(digest),64)

    def test_download_response_not_excel(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'fake.xlsx';p.write_bytes(b'"not-an-excel-download"')
            st=p.stat();item=scanner.fingerprint({'path':str(p),'size':st.st_size,'mtimeNs':st.st_mtime_ns})
            result=scanner.scan_workbook(item,{KEY})
            self.assertIn('不是OOXML',result['error'])
            self.assertFalse(result['evidence'])


if __name__ == '__main__':
    unittest.main()
