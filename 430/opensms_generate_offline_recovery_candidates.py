#!/usr/bin/env python3
import argparse
import json
import os
import re
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from opensms_artifact_paths import SNAPSHOT_DIR, ensure_artifact_dirs


TZ = ZoneInfo("Asia/Hong_Kong")
PROTECTED_SOURCE_PATTERN = re.compile(r"^(A21|A2)(?:-|\s|$)", re.IGNORECASE)
OFFLINE_RECOVERY_ENABLED = False


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("export")
    parser.add_argument("--tag", required=True)
    args = parser.parse_args()

    ensure_artifact_dirs()

    export_path = os.path.abspath(args.export)
    if not OFFLINE_RECOVERY_ENABLED:
        output = {
            "createdAt": datetime.now(TZ).isoformat(),
            "source": export_path,
            "status": "disabled",
            "rule": "下架恢复功能已取消；不生成候选、不提交恢复、不重试、不回读",
            "protectedSourceCarrierGroups": ["A21", "A2"],
            "protectedSourceExcludedCounts": {group: 0 for group in ("A21", "A2")},
            "protectedSourceExcludedCount": 0,
            "count": 0,
            "rows": [],
        }
        output_path = os.path.join(
            str(SNAPSHOT_DIR),
            f"opensms_offline_recover_{args.tag}_candidates.json",
        )
        with open(output_path, "w", encoding="utf-8") as handle:
            json.dump(output, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        print(json.dumps({
            "path": output_path,
            "status": "disabled",
            "count": 0,
        }, ensure_ascii=False))
        return

    workbook = load_workbook(export_path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = list(next(iterator))
    indexes = {
        str(value or "").strip(): index
        for index, value in enumerate(headers)
        if value not in (None, "")
    }
    dates = [
        match.group(1)
        for header in headers
        if isinstance(header, str)
        for match in [re.fullmatch(r"(\d{2}-\d{2})发送窗口数", header.strip())]
        if match
    ]
    if not dates:
        raise RuntimeError("报表没有日期统计列")

    latest = max(dates, key=lambda value: tuple(map(int, value.split("-"))))
    send_column = indexes[f"{latest}发送窗口数"]
    reply_column = indexes[f"{latest}回复窗口数"]
    rows = []
    protected_excluded_counts = {group: 0 for group in ("A21", "A2")}
    for values in iterator:
        status = str(values[indexes["状态"]] or "").strip()
        send = number(values[send_column])
        if status != "下架" or send <= 0:
            continue
        carrier = str(values[indexes.get("运营商")] or "").strip() if "运营商" in indexes else ""
        protected_match = PROTECTED_SOURCE_PATTERN.match(carrier)
        if protected_match:
            protected_excluded_counts[protected_match.group(1).upper()] += 1
            continue
        phone = str(values[indexes["电话"]] or "").strip()
        if phone.endswith(".0"):
            phone = phone[:-2]
        rows.append({
            "电话": phone,
            "phone": phone,
            "卡槽": str(values[indexes.get("卡槽")] or "").strip() if "卡槽" in indexes else "",
            "状态": status,
            "运营商": carrier,
            "发送窗口数": send,
            "回复窗口数": number(values[reply_column]),
            "统计日期": latest,
        })
    workbook.close()

    output = {
        "createdAt": datetime.now(TZ).isoformat(),
        "source": export_path,
        "reportLatestDate": latest,
        "rule": "latest date has sending windows and current status is down",
        "protectedSourceCarrierGroups": list(protected_excluded_counts),
        "protectedSourceExcludedCounts": protected_excluded_counts,
        "protectedSourceExcludedCount": sum(protected_excluded_counts.values()),
        "count": len(rows),
        "rows": rows,
    }
    output_path = os.path.join(
        str(SNAPSHOT_DIR),
        f"opensms_offline_recover_{args.tag}_candidates.json",
    )
    with open(output_path, "w", encoding="utf-8") as handle:
        json.dump(output, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(json.dumps({"path": output_path, "latestDate": latest, "count": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
