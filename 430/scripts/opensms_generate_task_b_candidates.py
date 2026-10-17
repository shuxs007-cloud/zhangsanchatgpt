#!/usr/bin/env python3
import argparse
import hashlib
import json
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from opensms_temporary_carrier_hold import temporary_carrier_held, isolation_carrier_held
from opensms_data_anomalies import source_anomaly, anomaly_fields, assert_anomaly_identity_separation


EXPIRED_TARGETS = {"C": "过期下架卡C", "T": "过期下架卡T"}
TOMORROW_TARGETS = {"C": "快过期隔离区C", "T": "快过期隔离区T"}
CHATR_PATTERN = re.compile(r"chatr", re.IGNORECASE)
T_FAMILY_PATTERN = re.compile(r"kooda|koodo|kini|telus1|\blucky\b", re.IGNORECASE)
REQUIRED_HEADERS = {"电话", "运营商", "卡槽", "过期状态", "截止日"}
TASK_B_PROTECTED_SOURCE_EXCEPTION_GROUPS = ("A21", "A2")
TASK_B_PROTECTED_SOURCE_EXCEPTION_PATTERN = re.compile(
    r"^(A21|A2)(?:-|\s|$)", re.IGNORECASE
)
NEVER_PROCESS_SOURCE_CARRIERS = ("快过期续费C", "快过期续费T")
TASK_B_RETAINED_SOURCE_CARRIERS = ("B14-C快下保留卡", "B14-T快下保留卡")
TASK_B_SOURCE_SCOPE = "all_carriers_v1"


def parse_args():
    parser = argparse.ArgumentParser(description="从本轮卡槽关联结果表生成任务B候选")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--server-date", required=True)
    parser.add_argument("--tag", required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def parse_date(value, epoch):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else None
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        try:
            converted = from_excel(float(text), epoch)
            return converted.date() if isinstance(converted, datetime) else None
        except (TypeError, ValueError, OverflowError):
            pass
    match = re.fullmatch(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?", text)
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def deadline_status(deadline, server_date):
    """The resolved 截止日, never age or an inherited label, decides expiry."""
    if deadline is None:
        return ""
    remaining = (deadline - server_date).days
    return "已过期" if remaining <= 0 else "快过期" if remaining <= 2 else "6天" if remaining <= 6 else ""


def phone_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", "", str(value).strip()).removesuffix(".0")


def json_value(value, epoch):
    if value in (None, ""):
        return ""
    parsed = parse_date(value, epoch)
    if parsed and isinstance(value, (date, datetime)):
        return parsed.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def unique_phones(rows, label):
    seen = set()
    duplicates = []
    for row in rows:
        phone = row["电话"]
        if phone in seen:
            duplicates.append(phone)
        seen.add(phone)
    if duplicates:
        raise ValueError(f"{label}候选存在重复号码: {sorted(set(duplicates))[:20]}")


def classify_opr(value):
    text = str(value or "").strip()
    is_chatr = bool(CHATR_PATTERN.search(text))
    is_t_family = bool(T_FAMILY_PATTERN.search(text))
    if is_chatr and is_t_family:
        return "ambiguous"
    if is_chatr:
        return "C"
    if is_t_family:
        return "T"
    return "unclassified"


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def candidate_payload(*, source, source_sha, server_date, target, route, rule, rows):
    source_names = sorted({str(row.get("运营商", "")).strip() for row in rows if str(row.get("运营商", "")).strip()})
    return {
        "createdAt": datetime.now().astimezone().isoformat(),
        "sourcePath": str(source.resolve()),
        "sourceSha256": source_sha,
        "sourceSheet": "sheet1",
        "serverDate": server_date.isoformat(),
        "tomorrowDate": (server_date + timedelta(days=1)).isoformat(),
        "rule": rule,
        "candidateCount": len(rows),
        "count": len(rows),
        "task": "B",
        "forceGate": True,
        "unconfiguredCarrierPolicy": "skip_no_business_write_v1",
        "expiryDecisionBasis": "deadline_minus_server_date_v1",
        "sourceCarrierScope": TASK_B_SOURCE_SCOPE,
        "sourceCarrierExclusions": [] if TASK_B_SOURCE_SCOPE == "all_carriers_v1" else list(NEVER_PROCESS_SOURCE_CARRIERS + TASK_B_RETAINED_SOURCE_CARRIERS),
        "retainedSourceProtectionEnforced": TASK_B_SOURCE_SCOPE != "all_carriers_v1",
        "retainedSourceCarrierNames": [] if TASK_B_SOURCE_SCOPE == "all_carriers_v1" else list(TASK_B_RETAINED_SOURCE_CARRIERS),
        "route": route,
        "targetCarrierName": target,
        "sourceCarrierNames": source_names,
        "rows": rows,
        "candidates": rows,
    }


def main():
    hold_evaluated_at = datetime.now().astimezone()
    args = parse_args()
    if os.environ.get("OPENSMS_FORCE_TASK_B") != "1":
        raise RuntimeError("OPENSMS_FORCE_TASK_B 未严格等于 1，拒绝生成任务B候选")
    server_date = date.fromisoformat(args.server_date)
    if not args.workbook.is_file():
        raise FileNotFoundError(args.workbook)

    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    if "sheet1" not in workbook.sheetnames:
        raise ValueError(f"结果主表 sheet1 不存在: {workbook.sheetnames}")
    sheet = workbook["sheet1"]
    header_values = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    headers = [str(value).strip() if value is not None else "" for value in header_values]
    missing_headers = sorted(REQUIRED_HEADERS - set(headers))
    if missing_headers:
        raise ValueError(f"结果主表缺少字段: {missing_headers}")

    expired_rows = {"C": [], "T": []}
    tomorrow_rows = {"C": [], "T": []}
    held_opr_rows = {"ambiguous": [], "unclassified": []}
    skipped_missing_phone = []
    protected_source_included_rows = {
        group: [] for group in TASK_B_PROTECTED_SOURCE_EXCEPTION_GROUPS
    }
    never_process_source_rows = {carrier: [] for carrier in NEVER_PROCESS_SOURCE_CARRIERS}
    retained_source_rows = {carrier: [] for carrier in TASK_B_RETAINED_SOURCE_CARRIERS}
    expired_retained_source_rows = []
    tomorrow_retained_source_rows = []
    expired_non_protected_rows = []
    tomorrow_non_protected_rows = []
    temporary_held_rows = []
    isolation_held_rows = []
    all_sources = TASK_B_SOURCE_SCOPE == "all_carriers_v1"
    source_coverage = {}
    status_mismatches = []
    unconfigured_carrier_rows = []
    data_anomalies = []
    due_anomaly_counts = {"expired": 0, "tomorrow": 0}
    tomorrow_date = server_date + timedelta(days=1)
    keep_headers = ["电话", "运营商", "opr", "卡槽", "状态", "创建时间", "过期状态", "机房", "开卡日期", "激活日期", "截止日"]

    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = {headers[index]: cell.value for index, cell in enumerate(cells) if index < len(headers) and headers[index]}
        source_carrier = str(values.get("运营商") or "").strip()
        coverage = source_coverage.setdefault(source_carrier, {
            "sourceCarrierName": source_carrier, "sourceRowCount": 0,
            "expiredSourceCount": 0, "expiredCandidateCount": 0,
            "tomorrowSourceCount": 0, "tomorrowCandidateCount": 0,
            "expiredDataAnomalyCount": 0, "tomorrowDataAnomalyCount": 0,
        })
        coverage["sourceRowCount"] += 1
        if not source_carrier:
            unconfigured_carrier_rows.append(row_number)
            continue
        original_status = str(values.get("过期状态") or "").strip()
        raw_deadline = values.get("截止日")
        deadline = parse_date(raw_deadline, workbook.epoch)
        blank_deadline = raw_deadline is None or (isinstance(raw_deadline, str) and not raw_deadline.strip())
        if deadline is None and (not blank_deadline or original_status):
            data_anomalies.append(source_anomaly(row_number, values, "deadline_missing_or_invalid", phone_text(values.get("电话"))))
            continue
        expiry_status = deadline_status(deadline, server_date)
        if original_status != expiry_status:
            status_mismatches.append({"源行号": row_number, "电话": phone_text(values.get("电话")),
                                      "截止日": deadline.isoformat() if deadline else "",
                                      "原过期状态": original_status, "截止日计算状态": expiry_status,
                                      "距截止日天数": (deadline - server_date).days if deadline else None})
        is_expired = deadline is not None and deadline <= server_date
        is_tomorrow = deadline == tomorrow_date
        if not is_expired and not is_tomorrow:
            continue
        if not all_sources and isolation_carrier_held(source_carrier):
            isolation_held_rows.append(row_number)
            continue
        if not all_sources and temporary_carrier_held(source_carrier, hold_evaluated_at):
            temporary_held_rows.append(row_number)
            continue
        if not all_sources and source_carrier in NEVER_PROCESS_SOURCE_CARRIERS:
            never_process_source_rows[source_carrier].append(row_number)
            continue
        if not all_sources and source_carrier in TASK_B_RETAINED_SOURCE_CARRIERS:
            retained_source_rows[source_carrier].append(row_number)
            if is_expired:
                expired_retained_source_rows.append(row_number)
            if is_tomorrow:
                tomorrow_retained_source_rows.append(row_number)
            continue
        if is_expired:
            expired_non_protected_rows.append(row_number)
            coverage["expiredSourceCount"] += 1
        if is_tomorrow:
            tomorrow_non_protected_rows.append(row_number)
            coverage["tomorrowSourceCount"] += 1
        protected_match = TASK_B_PROTECTED_SOURCE_EXCEPTION_PATTERN.match(source_carrier)
        if protected_match:
            protected_source_included_rows[protected_match.group(1).upper()].append(row_number)
        phone = phone_text(values.get("电话"))
        family = classify_opr(values.get("opr"))
        if family == "ambiguous":
            raise ValueError(f"任务B第{row_number}行OPR身份冲突，不能作为缺失数据跳过")
        if not phone:
            skipped_missing_phone.append(row_number)
        if not phone or family == "unclassified":
            data_anomalies.append(source_anomaly(row_number, values,
                "phone_missing" if not phone else "opr_unclassified", phone))
            branch = "expired" if is_expired else "tomorrow"
            coverage[f"{branch}DataAnomalyCount"] += 1
            due_anomaly_counts[branch] += 1
            continue
        selected = {header: json_value(values.get(header), workbook.epoch) for header in keep_headers if header in values}
        selected["电话"] = phone
        selected["截止日"] = deadline.isoformat() if deadline else ""
        selected["文档原过期状态"] = original_status
        selected["过期状态"] = expiry_status
        selected["距截止日天数"] = (deadline - server_date).days
        selected["源行号"] = row_number
        if family not in {"C", "T"}:
            held_opr_rows[family].append({
                "电话": phone,
                "运营商": source_carrier,
                "opr": str(values.get("opr") or ""),
                "过期状态": expiry_status,
                "源行号": row_number,
            })
            continue
        selected["OPR分类"] = "Chatr" if family == "C" else "Kooda/T系"
        if is_expired:
            selected["targetCarrierName"] = EXPIRED_TARGETS[family]
            expired_rows[family].append(selected)
            coverage["expiredCandidateCount"] += 1
        if is_tomorrow:
            selected["targetCarrierName"] = TOMORROW_TARGETS[family]
            tomorrow_rows[family].append(selected)
            coverage["tomorrowCandidateCount"] += 1

    workbook.close()
    assert_anomaly_identity_separation(data_anomalies,
        expired_rows['C'] + expired_rows['T'] + tomorrow_rows['C'] + tomorrow_rows['T'])
    held_due_rows = [
        row
        for rows in held_opr_rows.values()
        for row in rows
        if row.get("过期状态") in {"已过期", "快过期"}
    ]
    if held_due_rows:
        raise ValueError(
            "任务B存在无法按Chatr或Kooda/T系分区的已过期或明日到期非保护号码，"
            f"拒绝静默遗漏或部分执行: {held_due_rows[:20]}"
        )
    unique_phones(expired_rows["C"] + expired_rows["T"], "已过期")
    unique_phones(tomorrow_rows["C"] + tomorrow_rows["T"], "明日到期")
    if all_sources:
        seen_keys = set()
        for row in expired_rows["C"] + expired_rows["T"] + tomorrow_rows["C"] + tomorrow_rows["T"]:
            key = re.sub(r"\D", "", row["电话"])[-10:]
            if not key or key in seen_keys:
                raise ValueError("任务B全运营商到期候选号码无效或标准化后重复，禁止跨分支漏号或重复执行")
            seen_keys.add(key)

    source_sha = sha256(args.workbook)
    expired_paths = {
        family: args.output_dir / f"opensms_task_b_expired_to_offline_{family.lower()}_{args.tag}_candidates.json"
        for family in ("C", "T")
    }
    tomorrow_paths = {
        family: args.output_dir / f"opensms_task_b_tomorrow_fast_to_isolation_{family.lower()}_{args.tag}_candidates.json"
        for family in ("C", "T")
    }
    manifest_path = args.output_dir / f"opensms_task_b_{args.tag}_manifest.json"

    for family in ("C", "T"):
        expired = candidate_payload(
            source=args.workbook,
            source_sha=source_sha,
            server_date=server_date,
            target=EXPIRED_TARGETS[family],
            route=f"task_b_expired_to_offline_{family.lower()}",
            rule=(
                "结果主表有效截止日<=服务器日期，重算为‘已过期’，且OPR按"
                "Chatr归C、Kooda/T系及Lucky归T"
            ),
            rows=expired_rows[family],
        )
        tomorrow = candidate_payload(
            source=args.workbook,
            source_sha=source_sha,
            server_date=server_date,
            target=TOMORROW_TARGETS[family],
            route=f"task_b_tomorrow_fast_to_isolation_{family.lower()}",
            rule=(
                "结果主表有效截止日严格等于"
                f"服务器日期次日{tomorrow_date.isoformat()}，OPR按Chatr归C、"
                "Kooda/T系及Lucky归T"
            ),
            rows=tomorrow_rows[family],
        )
        write_json(expired_paths[family], expired)
        write_json(tomorrow_paths[family], tomorrow)
    manifest = {
        "createdAt": datetime.now().astimezone().isoformat(),
        "task": "B",
        **anomaly_fields(data_anomalies, manifest_path),
        "expiredDataAnomalyCount": due_anomaly_counts["expired"],
        "tomorrowDataAnomalyCount": due_anomaly_counts["tomorrow"],
        "forceGate": True,
        "unconfiguredCarrierPolicy": "skip_no_business_write_v1",
        "unconfiguredCarrierSkippedCount": len(unconfigured_carrier_rows),
        "unconfiguredCarrierSkippedRows": unconfigured_carrier_rows,
        "sourceCarrierScope": TASK_B_SOURCE_SCOPE,
        "sourceCarrierExclusions": [] if all_sources else list(NEVER_PROCESS_SOURCE_CARRIERS + TASK_B_RETAINED_SOURCE_CARRIERS),
        "sourceCarrierCoverage": [source_coverage[name] for name in sorted(source_coverage)],
        "sourceCarrierCoveragePassed": all(
            row["expiredSourceCount"] == row["expiredCandidateCount"] + row["expiredDataAnomalyCount"]
            and row["tomorrowSourceCount"] == row["tomorrowCandidateCount"] + row["tomorrowDataAnomalyCount"]
            for row in source_coverage.values()
        ),
        "expiryDecisionBasis": "deadline_minus_server_date_v1",
        "expiryStatusRecomputedCount": len(status_mismatches),
        "expiryStatusRecomputedRows": status_mismatches,
        "serverDate": server_date.isoformat(),
        "tomorrowDate": tomorrow_date.isoformat(),
        "sourcePath": str(args.workbook.resolve()),
        "sourceSha256": source_sha,
        "oprTargetRouting": {
            "Chatr": "C",
            "Kooda/T系": "T",
            "Lucky": "T",
            "ambiguous": "hold_no_write",
            "unclassified": "hold_no_write",
        },
        "expiredCandidatePaths": {
            family: str(expired_paths[family].resolve()) for family in ("C", "T")
        },
        "expiredCandidateCounts": {
            family: len(expired_rows[family]) for family in ("C", "T")
        },
        "expiredCandidateCount": sum(len(expired_rows[family]) for family in ("C", "T")),
        "tomorrowCandidatePaths": {
            family: str(tomorrow_paths[family].resolve()) for family in ("C", "T")
        },
        "tomorrowCandidateCounts": {
            family: len(tomorrow_rows[family]) for family in ("C", "T")
        },
        "tomorrowCandidateCount": sum(len(tomorrow_rows[family]) for family in ("C", "T")),
        "candidatePaths": [
            str(expired_paths[family].resolve()) for family in ("C", "T")
        ] + [
            str(tomorrow_paths[family].resolve()) for family in ("C", "T")
        ],
        "heldOprCounts": {
            key: len(rows) for key, rows in held_opr_rows.items()
        },
        "heldOprRows": held_opr_rows,
        "temporaryB11HeldRows": temporary_held_rows,
        "isolationCarrierHeldRows": isolation_held_rows,
        "isolationCarrierHeldCount": len(isolation_held_rows),
        "temporaryCarrierHoldEvaluatedAt": hold_evaluated_at.isoformat(),
        "temporaryB11HeldCount": len(temporary_held_rows),
        "taskBProtectedSourceExceptionEnabled": True,
        "protectedSourceCarrierGroups": list(TASK_B_PROTECTED_SOURCE_EXCEPTION_GROUPS),
        "protectedSourceIncludedCount": sum(
            len(rows) for rows in protected_source_included_rows.values()
        ),
        "protectedSourceIncludedCounts": {
            group: len(protected_source_included_rows[group])
            for group in TASK_B_PROTECTED_SOURCE_EXCEPTION_GROUPS
        },
        "protectedSourceExcludedCount": 0,
        "protectedSourceExcludedCounts": {
            group: 0 for group in TASK_B_PROTECTED_SOURCE_EXCEPTION_GROUPS
        },
        "neverProcessSourceCarrierNames": [] if all_sources else list(NEVER_PROCESS_SOURCE_CARRIERS),
        "neverProcessSourceExcludedCount": sum(len(rows) for rows in never_process_source_rows.values()),
        "neverProcessSourceExcludedCounts": {
            carrier: len(never_process_source_rows[carrier])
            for carrier in NEVER_PROCESS_SOURCE_CARRIERS
        },
        "neverProcessSourceExcludedRows": never_process_source_rows,
        "retainedSourceProtectionEnforced": not all_sources,
        "retainedSourceCarrierNames": [] if all_sources else list(TASK_B_RETAINED_SOURCE_CARRIERS),
        "retainedSourceExcludedCount": sum(len(rows) for rows in retained_source_rows.values()),
        "retainedSourceExcludedCounts": {
            carrier: len(retained_source_rows[carrier])
            for carrier in TASK_B_RETAINED_SOURCE_CARRIERS
        },
        "retainedSourceExcludedRows": retained_source_rows,
        "expiredRetainedSourceExcludedCount": len(expired_retained_source_rows),
        "tomorrowRetainedSourceExcludedCount": len(tomorrow_retained_source_rows),
        "expiredNonProtectedSourceCount": len(expired_non_protected_rows),
        "expiredClassifiedCandidateCount": sum(len(expired_rows[family]) for family in ("C", "T")),
        "expiredCoveragePassed": len(expired_non_protected_rows) == sum(
            len(expired_rows[family]) for family in ("C", "T")
        ) + due_anomaly_counts["expired"],
        "tomorrowNonProtectedSourceCount": len(tomorrow_non_protected_rows),
        "tomorrowClassifiedCandidateCount": sum(len(tomorrow_rows[family]) for family in ("C", "T")),
        "tomorrowCoveragePassed": len(tomorrow_non_protected_rows) == sum(
            len(tomorrow_rows[family]) for family in ("C", "T")
        ) + due_anomaly_counts["tomorrow"],
        "allDueCoveragePassed": (
            len(expired_non_protected_rows) == sum(len(expired_rows[family]) for family in ("C", "T")) + due_anomaly_counts["expired"]
            and len(tomorrow_non_protected_rows) == sum(len(tomorrow_rows[family]) for family in ("C", "T")) + due_anomaly_counts["tomorrow"]
        ),
    }
    write_json(manifest_path, manifest)
    manifest["manifestPath"] = str(manifest_path.resolve())
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
