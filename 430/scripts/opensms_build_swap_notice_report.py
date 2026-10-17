#!/usr/bin/env python3
"""Join fresh read-only swap population to local transactional notice evidence."""
import argparse
import hashlib
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--population', type=Path, required=True)
    parser.add_argument('--history', type=Path, required=True)
    parser.add_argument('--out-dir', type=Path, required=True)
    args = parser.parse_args()
    population = json.loads(args.population.read_text())
    history = json.loads(args.history.read_text())
    if population.get('status') != 'passed' or not population.get('readOnly'):
        raise ValueError('Fresh population read did not pass')
    phones = history['phones']
    stamp = datetime.fromisoformat(population['completedAt'].replace('Z', '+00:00')).astimezone(ZoneInfo('Asia/Hong_Kong'))
    rows = []
    details = []
    for item in sorted(population['rows'], key=lambda r: (r['carrierName'], r['phone10'])):
        raw = item['raw']
        entry = phones.get(item['phone10'], {})
        confirmed = entry.get('lastConfirmed')
        uncertain = entry.get('lastUncertain')
        legacy = entry.get('lastLegacyRecovered')
        not_found = entry.get('lastNotFound')
        if confirmed:
            label = '有已确认通知回执'
        elif uncertain:
            label = '未确证：存在不确定提交'
        elif legacy:
            label = '未确证：仅历史恢复凭据'
        elif not_found:
            label = '未确证：曾提交但后台未匹配'
        else:
            label = '未查到已确认通知记录'
        display_phone = str(raw['cardNo'])
        if len(display_phone) == 10 and str(raw.get('callingCode')) == '1':
            display_phone = '1' + display_phone
        proof = confirmed or uncertain or legacy or not_found or {}
        detail = {'phone': display_phone, 'phone10': item['phone10'], 'carrier': item['carrierName'],
                  'classification': label, 'confirmed': bool(confirmed), 'history': entry}
        details.append(detail)
        rows.append([
            display_phone, item['carrierName'], str(raw.get('opr', '')), str(raw.get('slotIdentif', '')),
            str(raw.get('status', '')), str(raw.get('displayStatus', '')), label,
            str((confirmed or {}).get('expireDate', '')), str((confirmed or {}).get('task', '')),
            str((confirmed or {}).get('operationStartedAt', '')),
            str(proof.get('receiptPath') or proof.get('evidencePath') or proof.get('path') or proof.get('file') or ''),
            item['phone10'], json.dumps(entry, ensure_ascii=False, separators=(',', ':')) if entry else '',
        ])
    if len(rows) != population['total'] or len({row[0] for row in rows}) != len(rows):
        raise ValueError('Population identity/count mismatch')
    missing = [row for row, detail in zip(rows, details) if not detail['confirmed']]
    headers = ['电话', '当前运营商', 'OPR', '卡槽', '系统状态', '上架状态', '通知核对结果',
               '最近确证通知到期时间', '最近确证任务', '回执操作起始时间', '主要证据路径', '规范化号码', '完整历史核对证据']
    wb = Workbook()
    wb.remove(wb.active)
    for title, sheet_rows in [('未查到确证通知', missing), ('全部换号区核对', rows)]:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in sheet_rows:
            ws.append(row)
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 34
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='16324F')
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        widths = [18, 17, 26, 18, 19, 16, 34, 25, 16, 29, 70, 18, 85]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.column_dimensions['L'].hidden = True
        ws.column_dimensions['M'].hidden = True
        for row in ws.iter_rows(min_row=2):
            has_confirmation = row[6].value == '有已确认通知回执'
            color = 'E9F4EE' if has_confirmation else 'FFF2CC'
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.fill = PatternFill('solid', fgColor=color)
                cell.alignment = Alignment(vertical='center')
                if isinstance(cell.value, str):
                    cell.data_type = 's'
            row[0].number_format = '@'
            row[11].number_format = '@'
    ws = wb.create_sheet('核对说明')
    notes = [
        ['项目', '说明'],
        ['实时查询完成时间', stamp.strftime('%Y-%m-%d %H:%M:%S %z')],
        ['当前范围', 'BitBrowser 430；精确换号区C、换号区T；每页500，全部页面；不限制回复量'],
        ['确证口径', '存在可核对精确号码与输入/命中/未匹配数量的历史同步事务业务回执。'],
        ['未查到确证通知', '仅表示可访问的本地证据中没有找到确证记录，不等于后台证明从未通知。'],
        ['历史覆盖', json.dumps(history.get('coverage', {}), ensure_ascii=False)],
        ['历史扫描统计', json.dumps(history.get('summary', {}), ensure_ascii=False)],
        ['重要限制', '后台没有到期标记独立查询接口。人工操作、未保存或已丢失回执不一定在本次范围内。'],
        ['时间含义', '到期时间是最近已确认通知所请求的日期，不是系统卡截止日，也不是当前后台独立回读值。'],
        ['身份口径', '按手机号对照历史；保留记录是用户申请日志，不能充当下卡通知日志。'],
        ['较晚不确定操作', '存在历史确证也不保证后续没有人工/未确证操作；完整证据保存在隐藏M列与JSON。'],
        ['本次写入', '无运营商修改、无下卡通知、未执行任务B/C/D/E/F。'],
        ['实时源文件', str(args.population.resolve())],
        ['历史证据文件', str(args.history.resolve())],
    ]
    for row in notes:
        ws.append(row)
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 120
    for row in ws:
        for cell in row:
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = 's'
    ws.freeze_panes = 'A2'
    args.out_dir.mkdir(parents=True, exist_ok=True)
    tag = stamp.strftime('%Y%m%dT%H%M%S%z')
    excel = args.out_dir / f'换号区_未查到确证下卡通知_{tag}.xlsx'
    report = args.out_dir / f'换号区_下卡通知核对_{tag}.md'
    result_file = args.out_dir / f'换号区_下卡通知核对_{tag}.json'
    for target in (excel, report, result_file):
        if target.exists():
            raise FileExistsError(target)
    wb.save(excel)
    with ZipFile(excel) as archive:
        if archive.testzip() is not None:
            raise ValueError('Excel ZIP verification failed')
    check = load_workbook(excel, read_only=True, data_only=False)
    for title, expected in [('未查到确证通知', missing), ('全部换号区核对', rows)]:
        values = list(check[title].values)
        expected_values = [tuple(headers)] + [tuple(value if value != '' else None for value in row) for row in expected]
        if values != expected_values:
            raise ValueError(f'Independent Excel readback failed: {title}')
        if any(cell.data_type in ('e', 'f') for row in check[title] for cell in row):
            raise ValueError('Unexpected formula/error in static evidence table')
    check.close()
    by_carrier = []
    for group in population['groups']:
        subset = [d for d in details if d['carrier'] == group['carrierName']]
        by_carrier.append({'carrier': group['carrierName'], 'total': len(subset),
                           'confirmed': sum(d['confirmed'] for d in subset),
                           'withoutConfirmed': sum(not d['confirmed'] for d in subset)})
    result = {'status': 'passed', 'readOnly': True, 'currentPopulation': population['total'],
              'withoutConfirmed': len(missing), 'confirmed': len(rows) - len(missing),
              'byCarrier': by_carrier, 'classificationCounts': dict(Counter(d['classification'] for d in details)),
              'queriedAt': stamp.isoformat(), 'historySummary': history.get('summary'),
              'files': {'excel': str(excel.resolve()), 'report': str(report.resolve())},
              'verification': {'zip': 'passed', 'independentRowReadback': 'passed', 'formulaErrors': 0,
                               'populationSha256': hashlib.sha256(args.population.read_bytes()).hexdigest(),
                               'historySha256': hashlib.sha256(args.history.read_bytes()).hexdigest()},
              'rows': details}
    result_file.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    lines = ['# 换号区下卡通知只读核对', '', f'当前运营商查询完成：{stamp:%Y-%m-%d %H:%M:%S %z}。', '',
             '| 当前区 | 当前号码 | 有确证通知 | 未查到确证通知 |', '| --- | ---: | ---: | ---: |']
    lines += [f"| {g['carrier']} | {g['total']} | {g['confirmed']} | {g['withoutConfirmed']} |" for g in by_carrier]
    lines += ['', f"共 {len(missing)} 个当前换号区号码未查到已确认通知记录，完整清单见 Excel 第一个工作表。", '',
              '## 证据与限制', '',
              '- 本次只读取当前换号区全部分页和本地历史D/E回执；没有补发通知、修改运营商或运行其它任务。',
              '- “未查到”不等于证明从未通知：后台未提供到期标记查询接口，未保存的人工操作可能缺少凭据。',
              '- 不以卡保留记录、候选清单、弹窗、HTTP 200 或整个流水线状态代替逐批事务证据。',
              '- 本表按手机号追溯已有确认，不证明当前卡的生命周期内有通知；号码重新分配、后续人工变更需另查。',
              '- Excel 依表格技能保留号码文本格式、筛选和冻结表头；已独立逐行回读及ZIP核验。', '',
              '## 分类统计', '']
    lines += [f'- {key}：{value}' for key, value in result['classificationCounts'].items()]
    lines += ['', '## 历史扫描覆盖', '', '```json', json.dumps({'summary': history.get('summary'), 'coverage': history.get('coverage')}, ensure_ascii=False, indent=2), '```', '',
              '## 文件', '', f'- [号码清单 Excel]({excel.resolve()})', f'- [完整对账 JSON]({result_file.resolve()})',
              f'- [当前换号区原始只读数据]({args.population.resolve()})', f'- [历史凭据索引]({args.history.resolve()})', '']
    report.write_text('\n'.join(lines))
    print(json.dumps({**result, 'rows': None, 'sampleWithoutConfirmed': [d['phone'] for d in details if not d['confirmed']][:15]}, ensure_ascii=False))


if __name__ == '__main__':
    main()
