#!/usr/bin/env python3
"""Generate Task C carrier-change candidates from a fresh card-slot workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from opensms_temporary_carrier_hold import temporary_carrier_held, isolation_carrier_held
from opensms_task_c_daily_guard import guard_metadata, exclusion_reason
from opensms_data_anomalies import source_anomaly, anomaly_fields, assert_anomaly_identity_separation


REQUIRED_HEADERS = {
    "电话",
    "运营商",
    "opr",
    "卡槽",
    "4日发送窗口数",
    "10日发送窗口数",
    "4日回复率",
    "换号状态",
    "过期状态",
    "截止日",
}
KOODA_PATTERN = re.compile(r"kooda|koodo|telus1|kini", re.IGNORECASE)
CHATR_PATTERN = re.compile(r"chatr", re.IGNORECASE)
UNCHANGED_STATUSES = {"", "未换号"}
TASK_C_TARGET_NAMES = {"C": "换号区C", "T": "换号区T"}
PROTECTED_SOURCE_GROUPS = ("A21", "A2")
PROTECTED_SOURCE_PATTERN = re.compile(r"^(A21|A2)(?:-|\s|$)", re.IGNORECASE)
NEVER_PROCESS_SOURCE_CARRIERS = ("快过期续费C", "快过期续费T")
TASK_B_PRIORITY_TARGET_CARRIERS = (
    "过期下架卡C", "过期下架卡T", "快过期隔离区C", "快过期隔离区T",
)
REPLY_RATE_SEND_GATED_SOURCE_GROUPS = (
    "A1", "A11", "A12", "B1", "B11", "B12", "B1V", "B11V", "B12V",
)
REPLY_RATE_SEND_GATED_SOURCE_PATTERN = re.compile(
    r"^(A1|A11|A12|B1|B11|B12|B1V|B11V|B12V)(?:-|\s|$)", re.IGNORECASE,
)
REPLY_RATE_SEND_THRESHOLD = 50
TEN_DAY_SEND_THRESHOLD_EXCLUSIVE = 360


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从本轮卡槽关联结果表生成任务C候选")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--target-carrier-c", required=True, help="后台实时唯一解析出的换号区C完整运营商名称")
    parser.add_argument("--target-carrier-t", required=True, help="后台实时唯一解析出的换号区T完整运营商名称")
    parser.add_argument("--server-date", required=True, help="本轮后台服务器日期 YYYY-MM-DD")
    parser.add_argument("--tag", required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def parse_date(value: Any, epoch) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        try:
            converted = from_excel(float(text), epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            pass
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def phone_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", "", str(value).strip()).removesuffix(".0")


def parse_number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def parse_rate(value: Any) -> float | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    is_percent = text.endswith("%")
    number = parse_number(text[:-1] if is_percent else value)
    if number is None or number < 0:
        return None
    if is_percent or number > 1:
        return number / 100
    return number


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def unique_phones(rows: list[dict[str, Any]], label: str) -> None:
    seen: set[str] = set()
    duplicates: set[str] = set()
    for row in rows:
        phone = row["电话"]
        if phone in seen:
            duplicates.add(phone)
        seen.add(phone)
    if duplicates:
        raise ValueError(f"{label}候选存在重复号码: {sorted(duplicates)[:20]}")


def candidate_payload(
    *,
    source: Path,
    source_sha: str,
    target: str,
    route: str,
    rule: str,
    threshold: float,
    server_date: date,
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    source_names = sorted({row["运营商"] for row in rows if row.get("运营商")})
    return {
        "createdAt": datetime.now().astimezone().isoformat(),
        "sourcePath": str(source.resolve()),
        "sourceSha256": source_sha,
        "sourceSheet": "sheet1",
        "task": "C",
        "forceGate": True,
        "unconfiguredCarrierPolicy": "skip_no_business_write_v1",
        "taskBPriorityLockEnforced": True,
        "serverDate": server_date.isoformat(),
        "tomorrowDate": (server_date + timedelta(days=1)).isoformat(),
        "route": route,
        "rule": rule,
        "rateWindowDays": 4,
        "replyRateThresholdExclusive": threshold,
        "comparison": "strictly_less_than",
        "replyRateSendThresholdInclusive": REPLY_RATE_SEND_THRESHOLD,
        "replyRateSendMetric": "4日发送窗口数",
        "replyRateSendGateGroups": list(REPLY_RATE_SEND_GATED_SOURCE_GROUPS),
        "tenDaySendThresholdExclusive": TEN_DAY_SEND_THRESHOLD_EXCLUSIVE,
        "tenDaySendMetric": "10日发送窗口数",
        "recentDailyReplyRateGuard": {**guard_metadata(server_date), "thresholdExclusive": threshold},
        "targetCarrierName": target,
        "sourceCarrierNames": source_names,
        "candidateCount": len(rows),
        "count": len(rows),
        "rows": rows,
        "candidates": rows,
    }


def main() -> None:
    args = parse_args()
    if os.environ.get("OPENSMS_FORCE_TASK_C") != "1":
        raise RuntimeError("OPENSMS_FORCE_TASK_C 未严格等于 1，拒绝生成任务C候选")
    target_carrier_c = args.target_carrier_c.strip()
    target_carrier_t = args.target_carrier_t.strip()
    server_date = date.fromisoformat(args.server_date)
    tomorrow_date = server_date + timedelta(days=1)
    if target_carrier_c != TASK_C_TARGET_NAMES["C"]:
        raise ValueError(
            f"任务C Chatr目标运营商必须精确等于{TASK_C_TARGET_NAMES['C']}: "
            f"{target_carrier_c or '空'}"
        )
    if target_carrier_t != TASK_C_TARGET_NAMES["T"]:
        raise ValueError(
            f"任务C Kooda/T系目标运营商必须精确等于{TASK_C_TARGET_NAMES['T']}: "
            f"{target_carrier_t or '空'}"
        )
    if not args.workbook.is_file():
        raise FileNotFoundError(args.workbook)

    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    if "sheet1" not in workbook.sheetnames:
        raise ValueError(f"结果主表 sheet1 不存在: {workbook.sheetnames}")
    sheet = workbook["sheet1"]
    header_values = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    headers = [str(value).strip() if value is not None else "" for value in header_values]
    missing_headers = sorted(REQUIRED_HEADERS - set(headers))
    if missing_headers:
        raise ValueError(f"结果主表缺少字段: {missing_headers}")
    daily_guard = guard_metadata(server_date)
    daily_headers = daily_guard["headers"]
    if any(headers.count(header) != 1 for header in daily_headers):
        raise ValueError(f"任务C服务器当天/前一天回复率字段缺失或重复: {daily_headers}")
    daily_excluded = {"above_threshold": [], "missing_or_invalid": []}

    kooda_rows: list[dict[str, Any]] = []
    chatr_rows: list[dict[str, Any]] = []
    skipped_missing_phone: list[int] = []
    skipped_missing_carrier: list[int] = []
    skipped_invalid_rate: list[int] = []
    skipped_ten_day_send: dict[str, list[int]] = {"missing_or_invalid": [], "not_above_360": []}
    skipped_reply_rate_send_gate: dict[str, list[int]] = {
        group: [] for group in REPLY_RATE_SEND_GATED_SOURCE_GROUPS
    }
    skipped_unknown_change_status: list[int] = []
    data_anomalies = []
    skipped_protected_rows: dict[str, list[int]] = {
        group: [] for group in PROTECTED_SOURCE_GROUPS
    }
    skipped_never_process_rows: dict[str, list[int]] = {
        carrier: [] for carrier in NEVER_PROCESS_SOURCE_CARRIERS
    }
    skipped_task_b_priority_rows: dict[str, list[int]] = {
        "已过期": [],
        "明日到期": [],
        "任务B目标运营商": [],
    }
    selected_phones: set[str] = set()
    temporary_held_rows = []
    isolation_held_rows = []

    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = {
            headers[index]: cell.value
            for index, cell in enumerate(cells)
            if index < len(headers) and headers[index]
        }
        opr = str(values.get("opr") or "").strip()
        source_carrier = str(values.get("运营商") or "").strip()
        if not source_carrier:
            skipped_missing_carrier.append(row_number)
            continue
        if isolation_carrier_held(source_carrier):
            isolation_held_rows.append(row_number)
            continue
        if temporary_carrier_held(source_carrier):
            temporary_held_rows.append(row_number)
            continue
        if source_carrier in NEVER_PROCESS_SOURCE_CARRIERS:
            skipped_never_process_rows[source_carrier].append(row_number)
            continue
        expiry_status = str(values.get("过期状态") or "").strip()
        deadline = parse_date(values.get("截止日"), workbook.epoch)
        if expiry_status == "已过期":
            skipped_task_b_priority_rows["已过期"].append(row_number)
            continue
        if expiry_status == "快过期" and deadline == tomorrow_date:
            skipped_task_b_priority_rows["明日到期"].append(row_number)
            continue
        if source_carrier in TASK_B_PRIORITY_TARGET_CARRIERS:
            skipped_task_b_priority_rows["任务B目标运营商"].append(row_number)
            continue
        protected_match = PROTECTED_SOURCE_PATTERN.match(source_carrier)
        if protected_match:
            skipped_protected_rows[protected_match.group(1).upper()].append(row_number)
            continue
        is_kooda = bool(KOODA_PATTERN.search(opr))
        is_chatr = bool(CHATR_PATTERN.search(opr))
        if not is_kooda and not is_chatr:
            continue

        send_count = parse_number(values.get("4日发送窗口数"))
        rate = parse_rate(values.get("4日回复率"))
        if send_count is None or send_count <= 0 or rate is None:
            skipped_invalid_rate.append(row_number)
            continue
        reply_rate_gate_match = REPLY_RATE_SEND_GATED_SOURCE_PATTERN.match(source_carrier)
        if reply_rate_gate_match and send_count < REPLY_RATE_SEND_THRESHOLD:
            skipped_reply_rate_send_gate[reply_rate_gate_match.group(1).upper()].append(row_number)
            continue

        ten_day_send = parse_number(values.get("10日发送窗口数"))
        if ten_day_send is None or not math.isfinite(ten_day_send) or ten_day_send < 0:
            skipped_ten_day_send["missing_or_invalid"].append(row_number)
            continue
        if ten_day_send <= TEN_DAY_SEND_THRESHOLD_EXCLUSIVE:
            skipped_ten_day_send["not_above_360"].append(row_number)
            continue

        change_status = str(values.get("换号状态") or "").strip()
        kooda_match = is_kooda and rate < 0.05
        chatr_rate_match = is_chatr and rate < 0.04
        if not kooda_match and not chatr_rate_match:
            continue
        if chatr_rate_match and change_status not in UNCHANGED_STATUSES:
            if change_status != "已换号":
                skipped_unknown_change_status.append(row_number)
                data_anomalies.append(source_anomaly(row_number, values, "change_status_unknown", phone_text(values.get("电话"))))
            continue

        daily_reason = exclusion_reason(values, daily_headers, 0.05 if kooda_match else 0.04)
        if daily_reason:
            daily_excluded[daily_reason].append(row_number)
            continue

        phone = phone_text(values.get("电话"))
        if not phone:
            skipped_missing_phone.append(row_number)
            data_anomalies.append(source_anomaly(row_number, values, "phone_missing", phone))
            continue
        if phone in selected_phones:
            raise ValueError(f"任务C号码重复或跨组重叠: {phone}")

        selected = {
            "电话": phone,
            "运营商": source_carrier,
            "opr": opr,
            "卡槽": str(values.get("卡槽") or "").strip(),
            "4日发送窗口数": send_count,
            "10日发送窗口数": ten_day_send,
            "4日回复率": rate,
            "4日回复率百分比": rate * 100,
            "换号状态": change_status,
            "过期状态": expiry_status,
            "截止日": deadline.isoformat() if deadline else "",
            "源行号": row_number,
            **{header: values[header] for header in daily_headers},
        }
        selected_phones.add(phone)
        if kooda_match:
            kooda_rows.append(selected)
        else:
            chatr_rows.append(selected)

    workbook.close()
    assert_anomaly_identity_separation(data_anomalies, kooda_rows + chatr_rows)
    unique_phones(kooda_rows, "Kooda/T系 4日回复率<5%")
    unique_phones(chatr_rows, "Chatr未换号且4日回复率<4%")

    source_sha = sha256(args.workbook)
    kooda_path = args.output_dir / f"opensms_task_c_kooda_4day_lt5_to_swap_{args.tag}_candidates.json"
    chatr_path = args.output_dir / f"opensms_task_c_chatr_unchanged_4day_lt4_to_swap_{args.tag}_candidates.json"
    manifest_path = args.output_dir / f"opensms_task_c_{args.tag}_manifest.json"

    kooda = candidate_payload(
        source=args.workbook,
        source_sha=source_sha,
        target=target_carrier_t,
        route="task_c_kooda_4day_reply_rate_lt5_to_swap",
        rule="OPR属于Kooda/T系；10日发送窗口数严格>360；A1/A11/A12/B1/B11/B12/B1V/B11V/B12V来源4日发送窗口数>=50，其他来源>0；4日回复率严格低于5%；服务器当天和前一天回复率均有效且均不大于5%",
        threshold=0.05,
        server_date=server_date,
        rows=kooda_rows,
    )
    chatr = candidate_payload(
        source=args.workbook,
        source_sha=source_sha,
        target=target_carrier_c,
        route="task_c_chatr_unchanged_4day_reply_rate_lt4_to_swap",
        rule="OPR包含Chatr；10日发送窗口数严格>360；A1/A11/A12/B1/B11/B12/B1V/B11V/B12V来源4日发送窗口数>=50，其他来源>0；4日回复率严格低于4%，且换号状态为空或未换号；服务器当天和前一天回复率均有效且均不大于4%",
        threshold=0.04,
        server_date=server_date,
        rows=chatr_rows,
    )
    write_json(kooda_path, kooda)
    write_json(chatr_path, chatr)
    manifest = {
        **anomaly_fields(data_anomalies, manifest_path),
        "createdAt": datetime.now().astimezone().isoformat(),
        "task": "C",
        "forceGate": True,
        "taskBPriorityLockEnforced": True,
        "unconfiguredCarrierPolicy": "skip_no_business_write_v1",
        "unconfiguredCarrierSkippedCount": len(skipped_missing_carrier),
        "unconfiguredCarrierSkippedRows": skipped_missing_carrier,
        "serverDate": server_date.isoformat(),
        "tomorrowDate": tomorrow_date.isoformat(),
        "sourcePath": str(args.workbook.resolve()),
        "sourceSha256": source_sha,
        "sourceSheet": "sheet1",
        "targetCarrierNames": {"C": target_carrier_c, "T": target_carrier_t},
        "groupPriority": [
            "task_c_kooda_4day_reply_rate_lt5_to_swap",
            "task_c_chatr_unchanged_4day_reply_rate_lt4_to_swap",
        ],
        "koodaReplyRateThresholdExclusive": 0.05,
        "chatrReplyRateThresholdExclusive": 0.04,
        "replyRateSendThresholdInclusive": REPLY_RATE_SEND_THRESHOLD,
        "replyRateSendMetric": "4日发送窗口数",
        "replyRateSendGateGroups": list(REPLY_RATE_SEND_GATED_SOURCE_GROUPS),
        "tenDaySendThresholdExclusive": TEN_DAY_SEND_THRESHOLD_EXCLUSIVE,
        "tenDaySendMetric": "10日发送窗口数",
        "tenDaySendGateExcludedCount": sum(map(len, skipped_ten_day_send.values())),
        "tenDaySendGateExcludedCounts": {reason: len(rows) for reason, rows in skipped_ten_day_send.items()},
        "tenDaySendGateExcludedRows": skipped_ten_day_send,
        "recentDailyReplyRateGuard": {**daily_guard, "thresholds": {"T": 0.05, "C": 0.04}},
        "recentDailyReplyRateExcludedCount": sum(map(len, daily_excluded.values())),
        "recentDailyReplyRateExcludedCounts": {key: len(value) for key, value in daily_excluded.items()},
        "recentDailyReplyRateExcludedRows": daily_excluded,
        "replyRateSendGateExcludedCount": sum(
            len(rows) for rows in skipped_reply_rate_send_gate.values()
        ),
        "replyRateSendGateExcludedCounts": {
            group: len(skipped_reply_rate_send_gate[group])
            for group in REPLY_RATE_SEND_GATED_SOURCE_GROUPS
        },
        "koodaCandidatePath": str(kooda_path.resolve()),
        "koodaCandidateCount": len(kooda_rows),
        "chatrCandidatePath": str(chatr_path.resolve()),
        "chatrCandidateCount": len(chatr_rows),
        "candidatePhoneCount": len(selected_phones),
        "temporaryB11HeldRows": temporary_held_rows,
        "isolationCarrierHeldRows": isolation_held_rows,
        "isolationCarrierHeldCount": len(isolation_held_rows),
        "temporaryB11HeldCount": len(temporary_held_rows),
        "protectedSourceCarrierGroups": list(PROTECTED_SOURCE_GROUPS),
        "protectedSourceExcludedCount": sum(len(rows) for rows in skipped_protected_rows.values()),
        "protectedSourceExcludedCounts": {
            group: len(skipped_protected_rows[group]) for group in PROTECTED_SOURCE_GROUPS
        },
        "protectedA21ExcludedCount": len(skipped_protected_rows["A21"]),
        "neverProcessSourceCarrierNames": list(NEVER_PROCESS_SOURCE_CARRIERS),
        "neverProcessSourceExcludedCount": sum(len(rows) for rows in skipped_never_process_rows.values()),
        "neverProcessSourceExcludedCounts": {
            carrier: len(skipped_never_process_rows[carrier])
            for carrier in NEVER_PROCESS_SOURCE_CARRIERS
        },
        "neverProcessSourceExcludedRows": skipped_never_process_rows,
        "taskBPriorityTargetCarrierNames": list(TASK_B_PRIORITY_TARGET_CARRIERS),
        "taskBPriorityExcludedCount": sum(
            len(rows) for rows in skipped_task_b_priority_rows.values()
        ),
        "taskBPriorityExcludedCounts": {
            reason: len(rows) for reason, rows in skipped_task_b_priority_rows.items()
        },
        "taskBPriorityExcludedRows": skipped_task_b_priority_rows,
        "skippedInvalidRateRowCount": len(skipped_invalid_rate),
        "mutuallyExclusive": True,
    }
    write_json(manifest_path, manifest)
    manifest["manifestPath"] = str(manifest_path.resolve())
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
