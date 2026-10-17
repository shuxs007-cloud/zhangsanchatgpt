#!/usr/bin/env python3
"""Local-only, snapshot-consistent Task F database -> unchanged deadline template."""
from __future__ import annotations

import argparse
from collections import Counter
import ctypes
import errno
from datetime import date, datetime, timedelta
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import warnings
import zipfile

import openpyxl  # Independent readback only; authoring is done by artifact-tool.

SCRIPTS = Path(__file__).resolve().parent
ROOT = SCRIPTS.parent
DEPENDENCIES = Path('/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies')
NODE = DEPENDENCIES / 'node/bin/node'
NODE_MODULES = DEPENDENCIES / 'node/node_modules'
DEFAULT_DB = ROOT / 'opensms_artifacts/state/task_f_kooda_activation.sqlite3'
DEFAULT_TEMPLATE = SCRIPTS / 'templates/task_f_deadline_import_1790891063096.xlsx'
DEFAULT_TEMPLATE_SHA256 = '3ca5d981113a6da94d2fabaea2085da2a1d00045064af70faf2c2f6727ce86ea'
HEADERS = ['ICCID', '截止日', '激活日期']
SHEET = '截止日'


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def publish_no_replace(source: Path, destination: Path) -> str:
    """Never overwrite; use verified exclusive copy when exFAT lacks rename-excl."""
    if sys.platform == 'darwin':
        # macOS SDK sys/stdio.h: RENAME_EXCL=0x4. APFS supports atomic no-replace.
        rename = ctypes.CDLL(None, use_errno=True).renamex_np
        rename.argtypes = [ctypes.c_char_p, ctypes.c_char_p, ctypes.c_uint]
        rename.restype = ctypes.c_int
        if rename(os.fsencode(source), os.fsencode(destination), 0x4) != 0:
            error = ctypes.get_errno()
            if error not in {errno.ENOTSUP, errno.EOPNOTSUPP}:
                raise OSError(error, os.strerror(error), str(destination))
        else:
            return 'atomic-no-replace'
    else:
        try:
            os.link(source, destination)
            return 'atomic-no-replace'
        except OSError as exc:
            if exc.errno not in {errno.ENOTSUP, errno.EOPNOTSUPP}:
                raise
    # exFAT has neither hard links nor RENAME_EXCL. Exclusive creation prevents
    # races; the success receipt is published only after copied bytes verify.
    with source.open('rb') as reader, destination.open('xb') as writer:
        shutil.copyfileobj(reader, writer)
        writer.flush()
        os.fsync(writer.fileno())
    if sha256(source) != sha256(destination):
        raise ValueError('发布文件哈希不一致；未生成成功回执')
    return 'exclusive-copy-verified'


def normalize_iccid(value: object) -> str:
    if not isinstance(value, str):
        raise ValueError('ICCID必须为文本，禁止从数字恢复精度')
    result = re.sub(r'[Ff]$', '', value)
    if not re.fullmatch(r'\d{18,22}', result, flags=re.ASCII):
        raise ValueError('ICCID必须为18至22位数字或仅含一个末尾F')
    return result


def valid_date(value: object) -> date | None:
    if value is None or value == '':
        return None
    if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        raise ValueError('数据库日期必须为yyyy-MM-dd')
    result = date.fromisoformat(value)
    if result.isoformat() != value:
        raise ValueError('无效自然日')
    return result


def read_rows(database: Path) -> tuple[list[list], dict]:
    """One read-only transaction, including count and data; no catalog/network work."""
    with sqlite3.connect(database.resolve().as_uri() + '?mode=ro', uri=True, timeout=30) as conn:
        conn.execute('PRAGMA query_only=ON')
        conn.execute('BEGIN')
        if conn.execute('PRAGMA quick_check').fetchone()[0] != 'ok':
            raise ValueError('任务F数据库完整性检查失败')
        count = conn.execute('SELECT COUNT(*) FROM kooda_activation').fetchone()[0]
        source = conn.execute('SELECT query_iccid,cutoff_date,activation_date,lookup_status '
                              'FROM kooda_activation ORDER BY query_iccid').fetchall()
    if len(source) != count:
        raise ValueError('数据库总数不一致')
    seen, rows, statuses = set(), [], Counter()
    for iccid, cutoff, activation, status in source:
        key = normalize_iccid(iccid)
        if key in seen:
            raise ValueError('归一化后ICCID重复，停止导出')
        seen.add(key)
        a, c = valid_date(activation), valid_date(cutoff)
        if bool(a) != bool(c) or (a and c != a + timedelta(days=29)):
            raise ValueError('截止日必须等于激活日+29天，且缺日期必须成对留空')
        if status not in {'found', 'activation_missing', 'not_found'}:
            raise ValueError('未知任务F记录状态')
        if (status == 'found') != bool(a):
            raise ValueError('任务F状态与日期不一致')
        rows.append([key, c.isoformat() if c else None, a.isoformat() if a else None])
        statuses[status] += 1
    rows.sort(key=lambda row: row[0])
    return rows, dict(statuses)


def check_template(path: Path) -> None:
    if path.resolve() == DEFAULT_TEMPLATE.resolve() and sha256(path) != DEFAULT_TEMPLATE_SHA256:
        raise ValueError('固定任务F模板SHA256不一致，停止导出')
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        # The supplied one-row template declares A1:A1 despite holding A1:C1.
        # Normal read mode computes actual cell bounds instead of trusting that cache.
        wb = openpyxl.load_workbook(path, read_only=False)
    try:
        if wb.sheetnames != [SHEET]:
            raise ValueError('模板必须只包含截止日工作表')
        sheet = wb[SHEET]
        if sheet.max_column != 3 or sheet.max_row != 1 or list(next(sheet.values)) != HEADERS:
            raise ValueError('模板列名/顺序不符或已有数据，停止以防覆盖')
    finally:
        wb.close()


def verify_output(path: Path, rows: list[list]) -> None:
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError('导出XLSX压缩结构损坏')
    # The writer may omit dimension metadata; parse actual cells rather than
    # depending on a cached range for independent row and identity validation.
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        if wb.sheetnames != [SHEET]:
            raise ValueError('输出工作表结构变化')
        sheet = wb[SHEET]
        if sheet.max_column != 3 or sheet.max_row != len(rows) + 1:
            raise ValueError('输出行列总数不一致')
        if list(next(sheet.values)) != HEADERS:
            raise ValueError('输出列头或顺序变化')
        actual = sheet.iter_rows(min_row=2)
        for index, expected in enumerate(rows, 2):
            cells = next(actual)
            iccid = cells[0]
            if iccid.data_type != 's' or iccid.number_format != '@' or iccid.value != expected[0]:
                raise ValueError(f'第{index}行ICCID文本精度校验失败')
            for col in (1, 2):
                cell, value = cells[col], cells[col].value
                if cell.number_format != 'yyyy-mm-dd' or cell.data_type in {'f', 'e'}:
                    raise ValueError(f'第{index}行日期类型/格式不正确')
                if value is None:
                    decoded = None
                elif isinstance(value, (datetime, date)):
                    decoded = value.strftime('%Y-%m-%d')
                    if isinstance(value, datetime) and value.time() != datetime.min.time():
                        raise ValueError('日期含非零时间')
                else:
                    raise ValueError(f'第{index}行日期不是Excel日期')
                if decoded != expected[col]:
                    raise ValueError(f'第{index}行日期与数据库不一致')
    finally:
        wb.close()


def write_template_workbook(template: Path, rows: list[list], output: Path) -> dict:
    """Shared pure writer for F increments, library exports and A gap exports."""
    template, output = template.resolve(), output.resolve()
    if output == template or output.suffix.lower() != '.xlsx':
        raise ValueError('模板和输出不得重叠，输出必须为xlsx')
    if os.path.lexists(output):
        raise FileExistsError('输出已存在；禁止覆盖历史')
    check_template(template)
    template_hash = sha256(template)
    seen = set()
    for row in rows:
        if len(row) != 3 or normalize_iccid(row[0]) != row[0] or row[0] in seen:
            raise ValueError('模板数据必须是唯一、已规范化ICCID的三列集合')
        seen.add(row[0])
        cutoff, activation = valid_date(row[1]), valid_date(row[2])
        if bool(cutoff) != bool(activation) or (activation and cutoff != activation + timedelta(days=29)):
            raise ValueError('模板日期必须成对有效，截止日等于激活日期+29天')
    payload = json.dumps(rows, ensure_ascii=False, separators=(',', ':'))
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.task_f_deadline_', dir=output.parent) as tmp:
        work = Path(tmp)
        (work / 'node_modules').symlink_to(NODE_MODULES, target_is_directory=True)
        writer = work / 'writer.mjs'
        shutil.copyfile(SCRIPTS / 'opensms_task_f_deadline_writer.mjs', writer)
        source, staged = work / 'rows.json', work / 'result.xlsx'
        source.write_text(payload, encoding='utf-8')
        write_start = time.monotonic()
        completed = subprocess.run([str(NODE), str(writer), str(template), str(source), str(staged)],
                                   capture_output=True, text=True, timeout=120)
        if completed.returncode:
            raise RuntimeError('模板生成程序失败: ' + completed.stderr[-1500:])
        writer_result = json.loads(completed.stdout.strip().splitlines()[-1])
        if writer_result.get('status') != 'success' or writer_result.get('recordCount') != len(rows):
            raise ValueError('模板生成程序回执数量不符')
        write_seconds = time.monotonic() - write_start
        verify_start = time.monotonic()
        verify_output(staged, rows)
        if sha256(template) != template_hash:
            raise ValueError('导出过程中模板发生变化')
        result = {'template': str(template), 'templateSha256': template_hash,
                  'headers': HEADERS, 'sheet': SHEET, 'dateFormat': 'yyyy-MM-dd',
                  'recordCount': len(rows), 'verificationPassed': True,
                  'sourceRowsSha256': hashlib.sha256(payload.encode()).hexdigest(),
                  'sha256': sha256(staged),
                  'writeSeconds': round(write_seconds, 3),
                  'verifySeconds': round(time.monotonic() - verify_start, 3)}
        result['publicationMode'] = publish_no_replace(staged, output)
    return result


def export_deadlines(database: Path, template: Path, output: Path, receipt: Path, *,
                     selected_rows: list[list] | None = None, extra_receipt: dict | None = None) -> dict:
    started = time.monotonic()
    database, template, output, receipt = [p.resolve() for p in (database, template, output, receipt)]
    if output.suffix.lower() != '.xlsx' or receipt.suffix.lower() != '.json':
        raise ValueError('输出必须为xlsx，回执必须为json')
    if len({database, template, output, receipt}) != 4:
        raise ValueError('输入、输出与回执路径不得重叠')
    if output.exists() or receipt.exists():
        raise FileExistsError('输出或回执已存在；请使用新路径，禁止覆盖历史')
    reserved = {'status', 'executionMode', 'networkAccess', 'databaseReadOnly', 'database',
                'template', 'templateSha256', 'sourceRowsSha256', 'output', 'receipt',
                'recordCount', 'activationCount', 'missingDateCount', 'statusCounts',
                'headers', 'iccidNormalized', 'dateFormat', 'verificationPassed', 'sha256',
                'createdAt', 'elapsedSeconds', 'performance', 'publicationMode'}
    if extra_receipt and set(extra_receipt) & reserved:
        raise ValueError('增量元数据不得覆盖导出校验字段')
    rows, statuses = read_rows(database)
    if selected_rows is not None:
        current = {row[0]: row for row in rows}
        keys = [row[0] for row in selected_rows]
        if not selected_rows or len(keys) != len(set(keys)):
            raise ValueError('增量导出必须是非空、唯一 ICCID 集合')
        for expected in selected_rows:
            if len(expected) != 3 or not expected[1] or not expected[2] or current.get(expected[0]) != expected:
                raise ValueError('增量入库回执与数据库当前已提交日期不一致')
        rows = sorted(selected_rows, key=lambda row: row[0])
        statuses = {'found': len(rows)}
    scan_seconds = time.monotonic() - started
    receipt.parent.mkdir(parents=True, exist_ok=True)
    written = write_template_workbook(template, rows, output)
    result = {
            'status': 'success', 'executionMode': 'local-program-only', 'networkAccess': False,
            'databaseReadOnly': True, 'database': str(database), 'template': str(template),
            'templateSha256': written['templateSha256'], 'sourceRowsSha256': written['sourceRowsSha256'],
            'output': str(output), 'receipt': str(receipt), 'recordCount': len(rows),
            'activationCount': statuses.get('found', 0),
            'missingDateCount': len(rows) - statuses.get('found', 0), 'statusCounts': statuses,
            'headers': HEADERS, 'iccidNormalized': True, 'dateFormat': 'yyyy-MM-dd',
            'verificationPassed': True, 'sha256': written['sha256'],
            'createdAt': datetime.now().astimezone().isoformat(),
            'elapsedSeconds': round(time.monotonic() - started, 3),
            'performance': {'readSeconds': round(scan_seconds, 3), 'writeSeconds': written['writeSeconds'],
                            'verifySeconds': written['verifySeconds']},
        }
    if extra_receipt:
        if set(extra_receipt) & set(result):
            raise ValueError('增量元数据不得覆盖导出校验字段')
        result.update(extra_receipt)
    # The receipt is the commit marker; a crash before it is never success.
    result['publicationMode'] = written['publicationMode']
    result['elapsedSeconds'] = round(time.monotonic() - started, 3)
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=receipt.parent,
                                         prefix='.deadline_receipt_', delete=False) as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
        handle.flush()
        os.fsync(handle.fileno())
        receipt_tmp = Path(handle.name)
    try:
        publish_no_replace(receipt_tmp, receipt)
    finally:
        receipt_tmp.unlink(missing_ok=True)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description='只读任务F数据库，全量填入固定三列截止日模板；不联网')
    parser.add_argument('--db', type=Path, default=DEFAULT_DB)
    parser.add_argument('--template', type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--receipt', type=Path)
    parser.add_argument('--json', action='store_true')
    args = parser.parse_args()
    stamp = datetime.now().astimezone().strftime('%Y%m%dT%H%M%S_%f%z')
    output = args.output or ROOT / 'outputs/task_f_deadlines' / stamp / f'Kooda截止日导入_{stamp}.xlsx'
    receipt = args.receipt or output.with_suffix('.receipt.json')
    try:
        result = export_deadlines(args.db, args.template, output, receipt)
    except Exception as exc:
        print(json.dumps({'status': 'failed', 'error': str(exc), 'output': None,
                          'networkAccess': False, 'databaseReadOnly': True}, ensure_ascii=False))
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=None if args.json else 2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
