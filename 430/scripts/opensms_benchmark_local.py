#!/usr/bin/env python3
"""Offline, same-input port/style and concurrency benchmark; no live operations.

Requires an exact successful state. Historical files are read-only fixtures,
never fresh deliverables. Workers have OS-enforced network/write isolation.
Results are advice only: this command never changes production configuration.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import shutil
import statistics
import subprocess
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
SANDBOX = Path('/usr/bin/sandbox-exec')


def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def save(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def load_module(path):
    spec = importlib.util.spec_from_file_location('offline_target', path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def sandbox_profile(output):
    return ('(version 1)(allow default)(deny network*)(deny file-write*)'
            '(allow file-write-data (literal "/dev/null"))'
            f'(allow file-write* (subpath {json.dumps(str(output.resolve()))}))'
            '(deny process-exec (literal "/usr/bin/security")'
            ' (literal "/usr/bin/open") (literal "/usr/bin/osascript"))')


def worker(args):
    directory = args.output.resolve()
    if args.worker == 'compare':
        module = load_module(args.script)
        count = module.verify(args.source, args.comparison)
        result = {'status': 'passed', 'rowCount': count,
                  'baselineSource': str(args.source), 'optimizedOutput': str(args.comparison),
                  'baselineSha256': digest(args.source), 'optimizedSha256': digest(args.comparison),
                  'baselineVerifierUsed': True, 'valueTypeAndStylesEqual': True}
        save(directory / 'comparison.json', result)
        print(json.dumps(result))
    elif args.worker == 'port':
        module = load_module(args.script)
        result = module.export_subset(args.source, directory / 'subset.xlsx',
                                      directory / 'receipt.json', 'offline-benchmark', digest(args.source))
        print(json.dumps({'rowCount': result['rowCount'], 'counts': result['counts'],
                          'independentReadbackPassed': result['independentReadbackPassed']}))
    else:
        from openpyxl import load_workbook
        module = load_module(ROOT / 'services/opensms_document_api/pipeline/device_daily_statistics.py')
        book = load_workbook(args.source, read_only=True, data_only=False)
        try:
            iterator = book['sheet1'].iter_rows(values_only=True)
            headers, rows = tuple(next(iterator)), tuple(iterator)
        finally:
            book.close()
        # Parse once, alternate order to limit cache/thermal bias. Only the pure
        # aggregation is timed; workbook loading is intentionally excluded.
        records, expected = [], None
        for iteration in range(args.repeat):
            order = args.workers if iteration % 2 == 0 else list(reversed(args.workers))
            for count in order:
                started = time.perf_counter()
                result = module.aggregate_device_daily(headers, rows, workers=count)
                duration = time.perf_counter() - started
                normalized = {'headers': result['headers'], 'rows': result['rows'],
                              'receipt': {k: v for k, v in result['receipt'].items()
                                          if k not in {'aggregation_seconds', 'requested_workers', 'actual_workers'}}}
                semantic_hash = hashlib.sha256(json.dumps(normalized, sort_keys=True).encode()).hexdigest()
                if expected is not None and semantic_hash != expected:
                    raise RuntimeError('Device/day results differ across concurrency levels')
                expected = semantic_hash
                records.append({'iteration': iteration + 1, 'workers': count,
                                'seconds': duration, 'semanticSha256': semantic_hash})
        save(directory / 'device.json', records)
        print(json.dumps({'records': records, 'allResultsEqual': True}))


def validate_workers(text):
    values = list(dict.fromkeys(int(item) for item in text.split(',')))
    if not values or any(value < 1 or value > min(16, os.cpu_count() or 1) for value in values):
        raise ValueError('workers must be 1..min(CPU count,16)')
    return values


def recommend_workers(medians):
    """Choose the smallest pool within 5% of the fastest measured median."""
    if not medians:
        return None
    values = {int(count): float(seconds) for count, seconds in medians.items()}
    if any(count < 1 or not math.isfinite(seconds) or seconds <= 0
           for count, seconds in values.items()):
        raise ValueError('Invalid worker timing samples')
    fastest = min(values, key=lambda count: (values[count], count))
    chosen = min(count for count, seconds in values.items() if seconds <= values[fastest] * 1.05)
    return {'workers': chosen, 'fastestMeasuredWorkers': fastest,
            'medianSeconds': values[chosen], 'nearFastestTolerancePct': 5,
            'scope': 'this_fixture_and_machine_only', 'autoApplied': False}


def audit_semantics(audit):
    rows = []
    for item in audit['checks']:
        details = dict(item['details'])
        if item['name'] == '最终交付校验':
            details.pop('hash_workers', None)
        rows.append((item['name'], item['status'], item['message'], details))
    return rows


def validate_output_directory(output, root=ROOT):
    """Accept the configured external outputs link, but not link escapes."""
    output = output.resolve()
    allowed = (root / 'outputs').resolve(strict=True)
    if output == allowed or not output.is_relative_to(allowed):
        raise ValueError('Benchmark output must be a new private directory under 430/outputs')
    return output


def run(args):
    if sys.platform != 'darwin' or not SANDBOX.is_file():
        raise RuntimeError('OS isolation unavailable; refusing unprotected benchmark')
    state = json.loads(args.state.read_text())
    if state.get('status') != 'success':
        raise ValueError('An exact successful state is required')
    source = Path(state['summary']['finalWorkbook'])
    expected_sha = state['stages']['refresh']['finalWorkbookSha256']
    if digest(source) != expected_sha:
        raise ValueError('Final workbook SHA256 does not match this state')
    output = validate_output_directory(args.output)
    if args.state.resolve().parent in output.parents or source.resolve().parent in output.parents:
        raise ValueError('Benchmark output cannot be inside the evidence source directory')
    output.mkdir(parents=True, exist_ok=False)
    local_source = output / 'readonly_fixture.xlsx'
    shutil.copyfile(source, local_source)
    config = Path(state['stages']['finalize']['config'])
    source_hashes = {str(path): digest(path) for path in (source, args.state, local_source, config)}
    records = []
    environment = {key: value for key, value in os.environ.items() if not key.startswith('OPENSMS_')}
    environment['PYTHONDONTWRITEBYTECODE'] = '1'

    def execute(label, command):
        directory = output / label
        directory.mkdir()
        environment.update(TMPDIR=str(directory), PYTHONPYCACHEPREFIX=str(directory / 'pycache'))
        started = time.perf_counter()
        completed = subprocess.run([str(SANDBOX), '-p', sandbox_profile(directory), *map(str, command(directory))],
                                   cwd=ROOT, env=environment, capture_output=True, text=True, timeout=240)
        elapsed = time.perf_counter() - started
        (directory / 'stdout.json').write_text(completed.stdout)
        (directory / 'stderr.txt').write_text(completed.stderr)
        records.append({'label': label, 'seconds': elapsed, 'exitCode': completed.returncode})
        save(output / 'progress.json', records)
        print(json.dumps(records[-1]), flush=True)
        if completed.returncode:
            raise RuntimeError(f'Offline benchmark failed; inspect {directory}')
        return directory, completed.stdout

    scripts = {} if args.skip_port else {'current': ROOT / 'scripts/opensms_export_port_subset.py'}
    if args.baseline:
        if args.skip_port:
            raise ValueError('--baseline requires the legacy port benchmark')
        scripts['baseline'] = args.baseline.resolve()
    frozen_scripts = {}
    for label, path in scripts.items():
        frozen_scripts[label] = output / f'{label}.py'
        shutil.copyfile(path, frozen_scripts[label])
    port_results = []
    for iteration in range(args.repeat):
        order = list(frozen_scripts) if iteration % 2 == 0 else list(reversed(frozen_scripts))
        for label in order:
            directory, stdout = execute(f'port-{label}-{iteration + 1}', lambda d, script=frozen_scripts[label]: [
                sys.executable, __file__, '--worker', 'port', '--source', local_source,
                '--script', script, '--output', d])
            port_results.append(json.loads(stdout))
    if any(row != port_results[0] for row in port_results):
        raise RuntimeError('Port subset counts/readback differ')
    cross_verified = False
    if 'baseline' in frozen_scripts:
        # Use the old uncached verifier to compare BOTH generated workbooks;
        # this catches mistakes shared by the new builder and its verifier.
        execute('cross-verification', lambda d: [sys.executable, __file__, '--worker', 'compare',
                '--script', frozen_scripts['baseline'], '--source', output / 'port-baseline-1/subset.xlsx',
                '--comparison', output / 'port-current-1/subset.xlsx', '--output', d])
        cross_verified = True
    device_workers = [value for value in args.workers if value <= 8]
    if not device_workers:
        device_workers = [1]
    _, stdout = execute('device', lambda d: [sys.executable, __file__, '--worker', 'device',
                        '--source', local_source, '--output', d, '--repeat', max(3, args.repeat),
                        '--workers', ','.join(map(str, device_workers))])
    device = json.loads(stdout)
    audit_results = []
    if not args.skip_audit:
        for iteration in range(args.repeat):
            order = args.workers if iteration % 2 == 0 else list(reversed(args.workers))
            for count in order:
                directory, _ = execute(f'audit-{count}-{iteration + 1}', lambda d, count=count: [
                    sys.executable, ROOT / 'scripts/opensms_parallel_audit.py', '--config', config,
                    '--delivery-dir', state['delivery'], '--output', d / 'audit.json',
                    '--workers', count, '--json'])
                audit = json.loads((directory / 'audit.json').read_text())
                # Preserve full independent audit receipts; compare all business
                # outcomes. Timing and per-thread scheduler fields vary by design.
                semantic = audit_semantics(audit)
                audit_results.append(semantic)
        if any(item != audit_results[0] for item in audit_results):
            raise RuntimeError('Audit outcomes differ across concurrency levels')
    unchanged = all(digest(Path(path)) == value for path, value in source_hashes.items())
    if not unchanged:
        raise RuntimeError('Benchmark source changed; measurements rejected')
    port_median = {label: statistics.median(row['seconds'] for row in records if row['label'].startswith(f'port-{label}-'))
                   for label in frozen_scripts}
    audit_median = {str(count): statistics.median(row['seconds'] for row in records if row['label'].startswith(f'audit-{count}-'))
                    for count in args.workers} if audit_results else {}
    device_median = {str(count): statistics.median(row['seconds'] for row in device['records'] if row['workers'] == count)
                     for count in device_workers}
    result = {'scope': 'offline-fixture-only-not-a-fresh-export', 'sourceRunId': state['runId'],
              'networkDeniedByOS': True, 'writesRestrictedToPrivateOutput': True, 'inputsUnchanged': unchanged,
              'inputSha256': source_hashes, 'cpuCount': os.cpu_count(), 'repeat': args.repeat,
              'portBenchmarkExecuted': bool(port_results),
              'portMedianSeconds': port_median, 'portResultsEqual': bool(port_results),
              'baselineIndependentCrossVerification': cross_verified,
              'auditMedianSeconds': audit_median, 'auditOutcomesEqual': bool(audit_results),
              'deviceMedianSeconds': device_median, 'deviceResultsEqual': device['allResultsEqual'],
              'recommendations': {'deviceDaily': recommend_workers(device_median),
                                  'audit': recommend_workers(audit_median)},
              'configurationChanged': False, 'records': records}
    save(output / 'benchmark.json', result)
    print(json.dumps({k: v for k, v in result.items() if k not in {'records', 'inputSha256'}}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--state', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--baseline', type=Path)
    parser.add_argument('--repeat', type=int, default=2, choices=range(1, 6))
    parser.add_argument('--workers', type=validate_workers, default=validate_workers('1,2,4,8'))
    parser.add_argument('--skip-audit', action='store_true')
    parser.add_argument('--skip-port', action='store_true', help='Benchmark only active report calculations/audit, not the cancelled port workbook')
    parser.add_argument('--worker', choices=('port', 'device', 'compare'))
    parser.add_argument('--source', type=Path)
    parser.add_argument('--script', type=Path)
    parser.add_argument('--comparison', type=Path)
    args = parser.parse_args()
    if args.worker:
        worker(args)
    elif args.state:
        run(args)
    else:
        parser.error('--state is required')


if __name__ == '__main__':
    main()
