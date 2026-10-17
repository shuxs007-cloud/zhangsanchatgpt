#!/usr/bin/env python3
import argparse
import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def parse_args():
    parser = argparse.ArgumentParser(description="从本轮卡槽结果主表生成指定机房全量下架候选")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--room", default="2号机房")
    parser.add_argument("--tag", required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def normalize_room(value):
    text = str(value or "").strip().replace("號", "号").replace("機", "机")
    aliases = {
        "二号机房": "2号机房",
        "三号机房": "3号机房",
        "四号机房": "4号机房",
    }
    return aliases.get(text, text)


def phone_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", "", str(value).strip()).removesuffix(".0")


def json_value(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main():
    args = parse_args()
    if not args.workbook.is_file():
        raise FileNotFoundError(args.workbook)
    target_room = normalize_room(args.room)
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    if "sheet1" not in workbook.sheetnames:
        raise ValueError(f"结果主表sheet1不存在: {workbook.sheetnames}")
    sheet = workbook["sheet1"]
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    missing = sorted({"电话", "机房"} - set(headers))
    if missing:
        raise ValueError(f"结果主表缺少字段: {missing}")
    keep = ["电话", "运营商", "opr", "卡槽", "状态", "创建时间", "机房", "过期状态", "截止日"]
    rows = []
    missing_phone_rows = []
    seen = set()
    duplicates = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = {headers[index]: cell.value for index, cell in enumerate(cells) if index < len(headers) and headers[index]}
        if normalize_room(values.get("机房")) != target_room:
            continue
        phone = phone_text(values.get("电话"))
        if not phone:
            missing_phone_rows.append(row_number)
            continue
        if phone in seen:
            duplicates.append(phone)
        seen.add(phone)
        selected = {header: json_value(values.get(header)) for header in keep if header in values}
        selected["电话"] = phone
        selected["机房"] = target_room
        selected["源行号"] = row_number
        rows.append(selected)
    workbook.close()
    if missing_phone_rows:
        raise ValueError(f"二号机房命中行缺少号码: {missing_phone_rows[:20]}")
    if duplicates:
        raise ValueError(f"二号机房候选存在重复号码: {sorted(set(duplicates))[:20]}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    output = args.output_dir / f"opensms_room_2_downlist_{args.tag}_candidates.json"
    payload = {
        "createdAt": datetime.now().astimezone().isoformat(),
        "sourcePath": str(args.workbook.resolve()),
        "sourceSha256": sha256(args.workbook),
        "sourceSheet": "sheet1",
        "rule": "结果主表机房归一化后严格等于2号机房的全部号码设置为下架",
        "targetRoom": target_room,
        "candidateCount": len(rows),
        "count": len(rows),
        "rows": rows,
        "candidates": rows,
    }
    with output.open("x", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(json.dumps({"path": str(output.resolve()), **payload, "rows": None, "candidates": None}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
