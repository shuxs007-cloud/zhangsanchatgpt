#!/usr/bin/env python3
"""Task I history pipeline, with an explicitly authorized current-position step.

History stays local. Only the current-position opt-in invokes the canonical
read-only, serial number-management search/export. No business writes.
"""
from __future__ import annotations
import csv, fcntl, hashlib, json, math, os, re, subprocess, time, zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from concurrent.futures import ProcessPoolExecutor
import opensms_task_i as I

NODE='/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
MARKER='/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/plugins/openai-primary-runtime/plugins/spreadsheets/skills/spreadsheets/container_tools/mark_artifact_operation_started.mjs'
WORKDIR='/Users/zhangsan/Downloads/430'

def dump(path,data):
    temporary=path.with_suffix(path.suffix+'.tmp')
    temporary.write_text(json.dumps(data,ensure_ascii=False,indent=2))
    os.replace(temporary,path)

def normalize_inputs(identifiers,lists):
    groups=[('查询清单',identifiers)]+[(p.stem if p.stem!='合计' else '清单_合计',p.read_text(encoding='utf-8-sig').splitlines()) for p in lists]
    keys=[];batches={};input_count=0
    for label,lines in groups:
        for value in lines:
            value=value.strip()
            if not value:continue
            input_count+=1
            key=I.iccid_key(value)
            if key is None:
                phone=I.phone_key(value)
                if phone is None:raise ValueError('无效查询标识: '+value)
                key='1'+phone
            if key not in batches:keys.append(key);batches[key]=label
    if not keys:raise ValueError('请提供ICCID、手机号或每行一个标识的清单')
    return keys,batches,input_count

def validate_current_position_request(args,keys):
    online=bool(getattr(args,'current_positions',False))
    authorization=getattr(args,'authorization_source',None)
    if online:
        if authorization!='current-user':raise ValueError('最新位置实时查询需要本次用户明确授权：--authorization-source current-user')
        if getattr(args,'current_search_receipt',None):raise ValueError('实时查询与已有搜索回执不能同时使用')
        if len(keys)>500 or any(I.iccid_key(k) is None for k in keys):raise ValueError('最新位置实时查询仅支持1至500个精确ICCID；手机号历史查询仍保持离线')
    elif authorization is not None:raise ValueError('授权来源只能与 --current-positions 同时使用')
    return online

def export_current_positions(result,out):
    """Exactly one fresh export attempt, using the existing guarded exporter."""
    query_path=out/'current_position_queries.json'
    dump(query_path,dict(queries=result['queries']))
    directory=out/'current_search_export'
    print(json.dumps(dict(event='task_i_current_position_started',requestedIccids=len(result['queries']),mode='serial-read-only-search-export')),flush=True)
    with (out/'current_position_export.log').open('w') as log:
        try:
            subprocess.run([NODE,str(Path(__file__).with_name('opensms_task_i_current_slots.mjs')),
                'export',str(directory),str(query_path),'--current-user-authorized-task-i-search-export'],
                check=True,cwd=WORKDIR,stdout=log,stderr=subprocess.STDOUT)
        except subprocess.CalledProcessError as e:
            failure=directory/'failure.json'
            reason=json.loads(failure.read_text()).get('error') if failure.is_file() else None
            raise RuntimeError('最新位置搜索导出失败，未重试：'+str(reason or '见本次 current_position_export.log')) from e
    receipt_path=directory/'export_receipt.json'
    I.attach_current_positions(result,receipt_path)
    result['currentPositionSource']['exportReceiptPath']=str(receipt_path)
    result['currentPositionSource']['authorizationSource']='current-user'
    return json.loads(receipt_path.read_text())

def prepare_result(result,batch,index_receipt):
    for q in result['queries']:q['batch']=batch[q['query']]
    for row in result['daily']:row['batch']=batch[row['query']]
    result['batches']=[]
    for label in list(dict.fromkeys(batch.values()))+['合计']:
        qs=[q for q in result['queries'] if label=='合计' or q['batch']==label]
        rows=[r for r in result['daily'] if label=='合计' or r['batch']==label]
        identities=[(r['date'],r['iccid'],r['phone']) for r in rows]
        total=[sum(r['values'][i] for r in rows) for i in range(4)] if len(identities)==len(set(identities)) else None
        result['batches'].append(dict(label=label,count=len(qs),withDocuments=sum(q['documentCount']>0 for q in qs),
            withDailyStats=sum(q['dailyRecordCount']>0 for q in qs),totals=total,
            replyRate=total[1]/total[0] if total and total[0] else None))
    counts=Counter(r['path'] for r in result['daily'])
    source_ids={p:n+1 for n,p in enumerate(sorted(counts))}
    result['usedSources']=[dict(id=source_ids[p],path=p,filename=Path(p).name,selectedDailyRows=n) for p,n in sorted(counts.items())]
    for r in result['daily']:r['sourceId']=source_ids[r['path']]
    result['indexReceipt']=index_receipt
    return result

def choose_samples(rows,limit=12):
    if not rows:return []
    picks=[]
    for key in (lambda r:r['date'],*(lambda r,n=n:r['values'][n] for n in range(4))):
        for r in (min(rows,key=key),max(rows,key=key)):
            if r not in picks:picks.append(r)
    for r in rows[::max(1,len(rows)//limit)]:
        if r not in picks:picks.append(r)
        if len(picks)>=limit:break
    return picks[:limit]

def _verify_source_group(job):
    """Independent parser verifies exact identity and the chosen daily counts."""
    path,sheet,rows=job;p=Path(path);before=p.stat();digest=I.content_key(p)
    from openpyxl import load_workbook
    from openpyxl.utils.cell import column_index_from_string
    book=None;found=[];wanted=defaultdict(list)
    for r in rows:wanted[r['row']].append(r)
    if p.suffix.lower() in ('.csv','.tsv'):
        raw=p.read_bytes();encodings=('utf-16','utf-8-sig','gb18030') if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','gb18030')
        for enc in encodings:
            try:text=raw.decode(enc);break
            except UnicodeDecodeError:continue
        else:raise ValueError('独立回读无法识别文本编码')
        import io
        iterator=enumerate(csv.reader(io.StringIO(text),delimiter='\t' if p.suffix.lower()=='.tsv' else ','),1)
        column=lambda ref:int(ref)
    else:
        book=load_workbook(p,read_only=True,data_only=True)
        iterator=enumerate(book[sheet].iter_rows(values_only=True),1)
        column=lambda ref:column_index_from_string(re.sub(r'[0-9]','',ref))-1
    headers=None
    try:
        for rn,values in iterator:
            if headers is None and rn<=15 and any(I.header(v) in I.IC_HEADERS|I.PH_HEADERS for v in values):headers=values
            if rn in wanted:
                if headers is None:raise ValueError('独立回读未发现身份表头')
                for r in wanted[rn]:
                    if r['iccid']:
                        if not r['iccidCell'] or I.iccid_key(values[column(r['iccidCell'])])!=r['iccid']:raise ValueError('独立原表ICCID不一致')
                    else:
                        ph=[I.phone_key(values[i]) for i,v in enumerate(headers) if I.header(v) in I.PH_HEADERS]
                        if r['phone'] not in ph:raise ValueError('独立原表手机号不一致')
                    mapped=I.metric_columns({str(i):(str(h),'str',False,'') for i,h in enumerate(headers)},r['observed'])
                    extracted=[I.count_cell((str(values[int(mapped[r['date']][metric])]) if values[int(mapped[r['date']][metric])] is not None else '', 'n',False,'')) for metric in I.METRICS]
                    if extracted!=r['values']:raise ValueError('独立原表计数不一致')
                    found.append(dict(path=path,sheet=sheet,row=rn,query=r['query'],date=r['date'],values=extracted))
            if rn>=max(wanted):break
    finally:
        if book:book.close()
    after=p.stat()
    if (before.st_size,before.st_mtime_ns)!=(after.st_size,after.st_mtime_ns) or digest!=I.content_key(p):raise ValueError('独立回读期间源文档变化')
    if len(found)!=len(rows):raise ValueError('独立原表回读行数不足')
    return found

def verify_sources(result,workers):
    samples=choose_samples(result['daily']);groups=defaultdict(list)
    for r in samples:groups[(r['path'],r['sheet'])].append(r)
    jobs=[(p,s,rows) for (p,s),rows in groups.items()]
    if jobs:
        with ProcessPoolExecutor(max_workers=min(workers,len(jobs)),initializer=I.configure_worker_priority) as pool:
            checked=[r for part in pool.map(_verify_source_group,jobs) for r in part]
    else:checked=[]
    for q in result['queries']:
        rows=[r for r in result['daily'] if r['query']==q['query']]
        total=[sum(r['values'][i] for r in rows) for i in range(4)] if rows else [None]*4
        if total!=q['totals']:raise ValueError('逐查询计数勾稽不一致')
    result['verification']=dict(status='passed',independentReader='openpyxl / CSV read-only',sampleCount=len(checked),
        originalsUnchanged=True,allQueryTotalsReconciled=True,samples=checked)
    return result['verification']

def slot_display(q):
    history='、'.join(q['historicalSlots']) or '未知'
    last=q['historicalLastSlot'] or ('待核对：'+'、'.join(q['historicalLastSlots']) if q['historicalLastSlots'] else '未知')
    current='、'.join(q.get('currentSlots',[])) or ('搜索未命中' if q.get('currentSlotStatus')=='search-not-found' else '最新导出卡槽为空' if q.get('currentSlotStatus')=='slot-missing' else '未实时查询')
    return [history,last,current]

def verify_workbook(p,result):
    """Read the saved workbook independently; test every requested metric."""
    from openpyxl import load_workbook
    with zipfile.ZipFile(p) as z:
        if z.testzip() is not None:raise ValueError('Excel ZIP CRC校验失败')
    b=load_workbook(p,read_only=True,data_only=True)
    try:
        if b.sheetnames!=['逐卡统计','逐手机号统计','批次汇总','逐日统计','采用来源','历史变更与缺失']:raise ValueError('报表页结构异常')
        for s in b:
            for row in s.iter_rows():
                if any(c.data_type=='e' for c in row):raise ValueError('报表包含公式错误: '+s.title)
        for cells,q in zip(b['逐卡统计'].iter_rows(min_row=7,max_row=6+len(result['queries'])),result['queries']):
            if cells[0].value!=q['query'] or cells[0].data_type!='s' or cells[0].number_format!='@':raise ValueError('长号码文本精度核验失败')
            expected=[q['totals'][0],q['totals'][3],q['totals'][1],q['totals'][2],q['totals'][3]] if q['dailyRecordCount'] else ['未知']*5
            if [c.value for c in cells[6:11]]!=expected:raise ValueError('逐查询五项计数不一致')
            rate=cells[5].value
            if (rate!='未知' if q['replyRate'] is None else not isinstance(rate,(int,float)) or not math.isclose(rate,q['replyRate'],rel_tol=1e-12)):raise ValueError('逐查询回复率不一致')
            if [c.value for c in cells[15:18]]!=slot_display(q):raise ValueError('历史/最新位置混淆')
        groups=defaultdict(lambda:[0,0,0,0]);seen={}
        for r in result['daily']:
            ident=(r['date'],r['iccid'],r['phone'])
            if ident in seen:
                if seen[ident]!=r['values']:raise ValueError('同一流量交叉计数冲突')
                continue
            seen[ident]=r['values']
            for i,v in enumerate(r['values']):groups[r['phone']][i]+=v
        if groups:
            for cells,(phone,t) in zip(b['逐手机号统计'].iter_rows(min_row=7,max_row=6+len(groups)),sorted(groups.items())):
                if cells[0].value!='1'+phone or cells[0].data_type!='s':raise ValueError('手机号文本核验失败')
                if [c.value for c in cells[4:9]]!=[t[0],t[3],t[1],t[2],t[3]]:raise ValueError('逐手机号计数不一致')
                rate=cells[3].value
                if (rate!='未知' if not t[0] else not isinstance(rate,(int,float)) or not math.isclose(rate,t[1]/t[0],rel_tol=1e-12)):raise ValueError('逐手机号回复率不一致')
                if [c.value for c in cells[12:15]]!=slot_display(result['phoneSlots'][phone]):raise ValueError('手机号位置不一致')
        for cells,r in zip(b['逐日统计'].iter_rows(min_row=2,max_row=1+len(result['daily'])),result['daily']):
            if cells[0].value!=r['query'] or [c.value for c in cells[4:8]]!=r['values'] or cells[11].value!=r['row']:raise ValueError('逐日来源行核验失败')
        totalrow=7+len(result['batches']);t=result['totals']
        expected=[t[0],t[1],t[3],t[2],t[3]] if t and result['daily'] else ['未知']*5
        if [c.value for c in next(b['批次汇总'].iter_rows(min_row=totalrow,max_row=totalrow,min_col=6,max_col=10))]!=expected:raise ValueError('汇总总计不一致')
    finally:b.close()
    formulas=load_workbook(p,read_only=True,data_only=False)
    try:
        rows=formulas['逐卡统计'].iter_rows(min_row=7,max_row=6+len(result['queries']),min_col=7,max_col=7)
        for cells,q in zip(rows,result['queries']):
            if q['dailyRecordCount'] and cells[0].data_type!='f':raise ValueError('统计公式未保留')
    finally:formulas.close()
    return dict(status='passed',queryRows=len(result['queries']),phoneRows=len(groups),dailyRows=len(result['daily']),
        allSixMetricsReconciled=True,allPositionColumnsVerified=True,longIdentifiersStoredAsExactText=True,
        savedFormulaCachesVerified=True,sourceSampleReadback=result['verification']['sampleCount'],output=str(p),bytes=p.stat().st_size)

def run_pipeline(args):
    if not 1<=args.workers<=8:raise ValueError('并发数必须1至8')
    start=time.monotonic();timings={};keys,batch,input_count=normalize_inputs(args.identifiers,args.list)
    online=validate_current_position_request(args,keys)
    I.validate_storage(args.db)
    out=args.output_dir or args.db.parent/'reports'/datetime.now().strftime('%Y%m%dT%H%M%S_%f')
    # Task I report directories are excluded from the discovery index. Reusing
    # a directory is forbidden so old report evidence can never be overwritten.
    if 'task_i' not in out.resolve().parts or not str(out.resolve()).startswith('/Volumes/zs01/'):
        raise ValueError('输出目录必须位于zs01的独立task_i目录，防止本地回退与输出回流历史索引')
    out.mkdir(parents=True,exist_ok=False)
    with (args.db.parent/'pipeline.lock').open('a') as lock:
        try:fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError:raise RuntimeError('任务I全流程正在运行，不启动第二份')
        phase='incremental-index'
        try:
            priority=I.configure_worker_priority();dump(out/'input_scope.json',dict(inputCount=input_count,uniqueCount=len(keys),identifiers=keys,batches=batch))
            t=time.monotonic()
            receipt=I.index(SimpleNamespace(db=args.db,root=args.root,workers=args.workers,omit=str(out),scan_evidence=None,manifest=None,receipt=out/'index_receipt.json',full_audit=args.full_audit))
            timings['incrementalIndex']=round(time.monotonic()-t,3)
            phase='history-query';t=time.monotonic();result=I.parallel_query(args.db,keys,args.workers,args.locations);timings['queryAndDedup']=round(time.monotonic()-t,3)
            prepare_result(result,batch,receipt)
            phase='independent-source-readback';t=time.monotonic();verify_sources(result,args.workers);timings['independentSourceReadback']=round(time.monotonic()-t,3)
            current_export=None
            if online:
                phase='current-position-search-export';t=time.monotonic()
                current_export=export_current_positions(result,out)
                timings['currentPositionSearchExportAndReadback']=round(time.monotonic()-t,3)
            elif args.current_search_receipt:
                phase='current-position-receipt-readback';I.attach_current_positions(result,args.current_search_receipt)
            dump(out/'query_result.json',result)
            filename=out/f'任务I_{len(keys)}个标识历史统计.xlsx'
            print(json.dumps(dict(event='task_i_report_started',queryCount=len(keys),dailyRows=len(result['daily']),actualWorkers=result['execution']['actualWorkers'])),flush=True)
            phase='workbook-build';t=time.monotonic()
            # Mandatory artifact-authoring marker is called once immediately
            # before this pipeline's sole XLSX authoring execution.
            subprocess.run([NODE,MARKER,'--operation-kind','create','--expected-output-count','1','--output-format','xlsx'],check=True,cwd=WORKDIR)
            with (out/'report_build.log').open('w') as log:
                subprocess.run([NODE,str(Path(__file__).with_name('opensms_task_i_report.mjs')),'--data',str(out/'query_result.json'),'--output',str(filename)],check=True,cwd=WORKDIR,stdout=log,stderr=subprocess.STDOUT)
            timings['workbookBuildAndRender']=round(time.monotonic()-t,3)
            phase='saved-workbook-readback';t=time.monotonic();verified=verify_workbook(filename,result);timings['savedWorkbookReadback']=round(time.monotonic()-t,3);dump(out/'independent_workbook_check.json',verified)
            final=dict(status='passed',workflow='task-i-history-and-current-positions' if online else 'task-i-local-pipeline',output=str(filename),resultJson=str(out/'query_result.json'),
                inputCount=input_count,queryCount=len(keys),requestedWorkers=args.workers,execution=result['execution'],coordinatorPriority=priority,
                index=receipt,timingsSeconds=timings,totalSeconds=round(time.monotonic()-start,3),
                verification=verified,networkRequests=None if online else 0,businessWrites=0,taskFUntouched=True,
                onlineCurrentPositionQuery=online,currentPositionSource=result.get('currentPositionSource'),
                currentPositionStatus='fresh-search-export-attached' if online or args.current_search_receipt else 'not-queried-local-only')
            if current_export:final['numberManagementRequestCount']=current_export.get('numberManagementRequestCount')
            dump(out/'run_receipt.json',final)
            print(json.dumps(final,ensure_ascii=False),flush=True)
            return final
        except Exception as e:
            dump(out/'run_receipt.json',dict(status='failed',error=str(e),failedPhase=phase,timingsSeconds=timings,totalSeconds=round(time.monotonic()-start,3),
                networkRequests=None if online else 0,businessWrites=0,onlineCurrentPositionQuery=online,
                currentPositionStatus='failed' if phase.startswith('current-position') else 'not-completed',taskFUntouched=True))
            raise
