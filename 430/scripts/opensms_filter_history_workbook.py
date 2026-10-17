#!/usr/bin/env python3
"""Remove an explicit ICCID list from a generated historical workbook; never overwrite it."""
import argparse
from concurrent.futures import ThreadPoolExecutor
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

from lxml import etree
from openpyxl import load_workbook

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ROW = re.compile(rb'<row\b[^>]*(?:/>|>.*?</row>)', re.S)
TOKEN = re.compile(rb'(?<![0-9])[0-9]{18,22}[Ff]?(?![0-9])')
CELL_REF = re.compile(rb'(<c\b[^>]*\br=")([A-Z]+)[0-9]+(")')
DATA_SHEETS = {'查询总览','卡槽日期历史','卡槽记录关联','历史记录','全部字段','逐条来源','缺失与待核对','不可靠标识命中'}


def normalize(value):
    value = str(value or '').strip()
    if value.endswith(('F','f')):
        value = value[:-1]
    return value


def values(fragment):
    root = etree.fromstring(fragment, etree.XMLParser(resolve_entities=False, no_network=True))
    result = {}
    for c in root:
        col = re.sub(r'\d','',c.get('r',''))
        result[col] = ''.join(c.itertext())
    return result


def content_hash(fragment):
    fragment = re.sub(rb'(<row\b[^>]*\br=")[0-9]+(")', rb'\g<1>0\2', fragment)
    return CELL_REF.sub(rb'\g<1>\g<2>0\3', fragment)


def filter_sheet(name, xml, targets, remaining):
    if re.search(rb'<(?:\w+:)?f(?:\s|>)',xml) or b'mergeCell' in xml or b'tablePart' in xml:
        raise ValueError('需另行处理的公式/合并/表格结构: '+name)
    if b't="s"' in xml:
        raise ValueError('共享字符串必须另行重建，禁止保留删除数据: '+name)
    wanted = {x.encode() for x in targets}
    removed, kept, found = 0, 0, set()
    expected = hashlib.sha256()
    def process(match):
        nonlocal removed, kept
        row = match.group()
        original_number = int(re.search(rb'\br="(\d+)"',row).group(1))
        matches = {normalize(x.group().decode()) for x in TOKEN.finditer(row)} & targets
        if original_number > 1 and matches:
            v = values(row)
            if name not in DATA_SHEETS or normalize(v.get('A')) not in targets:
                raise ValueError('删除标识出现在其他记录或元数据中，需核对: '+name+':'+str(original_number))
            removed += 1
            found.update(matches)
            return b''
        if name == '说明' and original_number > 1:
            v = values(row)
            if v.get('A') == '目标':
                root = etree.fromstring(row)
                c = next(c for c in root if c.get('r','').startswith('B'))
                for child in list(c):
                    c.remove(child)
                c.set('t','inlineStr')
                text = etree.SubElement(etree.SubElement(c,'is'),'t')
                text.text = f'本副本已删除用户指定的{len(targets)}个ICCID及全部关联数据，保留{remaining}个；原始工作簿未改动。'
                row = etree.tostring(root)
        kept += 1
        row = re.sub(rb'(<row\b[^>]*\br=")[0-9]+(")',lambda x:x[1]+str(kept).encode()+x[2],row)
        row = CELL_REF.sub(lambda x:x[1]+x[2]+str(kept).encode()+x[3],row)
        expected.update(content_hash(row))
        return row
    result = ROW.sub(process,xml)
    result = re.sub(rb'(<autoFilter\b[^>]*\bref="[A-Z]+1:[A-Z]+)[0-9]+(")',lambda m:m[1]+str(kept).encode()+m[2],result)
    result = re.sub(rb'(<dimension\b[^>]*\bref="[A-Z]+1:[A-Z]+)[0-9]+(")',lambda m:m[1]+str(kept).encode()+m[2],result)
    if {normalize(x.group().decode()) for x in TOKEN.finditer(result)} & targets:
        raise ValueError('工作表仍有删除标识残留: '+name)
    return result, dict(removed=removed,remaining=kept-1,found=sorted(found),retainedContentHash=expected.hexdigest())


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source',type=Path,required=True)
    p.add_argument('--remove-list',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    p.add_argument('--receipt',type=Path,required=True)
    args=p.parse_args()
    if args.output.exists() or args.receipt.exists() or args.source.resolve()==args.output.resolve():
        raise ValueError('输出/回执必须为新文件，禁止覆盖原文件')
    requested=args.remove_list.read_text(encoding='utf-8-sig').splitlines()
    if any(not re.fullmatch(r'[0-9]{18,22}',x) for x in requested):
        raise ValueError('非法删除清单')
    targets=set(requested)
    source_hash=hashlib.sha256(args.source.read_bytes()).hexdigest()
    with zipfile.ZipFile(args.source) as source:
        if 'xl/sharedStrings.xml' in source.namelist():
            raise ValueError('共享字符串结构不在本安全入口范围内')
        wb=etree.fromstring(source.read('xl/workbook.xml'))
        rels=etree.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        members={r.get('Id'):r.get('Target').lstrip('/') if r.get('Target').startswith('/') else 'xl/'+r.get('Target') for r in rels}
        sheets=[(s.get('name'),members[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]) for s in wb.findall('.//{'+NS+'}sheet')]
        index_member=dict(sheets)['查询总览']
        original_keys={normalize(values(m.group()).get('A')) for m in list(ROW.finditer(source.read(index_member)))[1:]}
        missing=targets-original_keys
        if missing:
            raise ValueError('以下删除ICCID不在原总览，需核对: '+','.join(sorted(missing)))
        remaining=len(original_keys-targets)
        payloads={name:source.read(member) for name,member in sheets}
        with ThreadPoolExecutor(max_workers=4) as pool:
            results=list(pool.map(lambda x:filter_sheet(x[0],payloads[x[0]],targets,remaining),sheets))
        replacements={member:result[0] for (_,member),result in zip(sheets,results)}
        stats={name:result[1] for (name,_),result in zip(sheets,results)}
        for defined in wb.findall('.//{'+NS+'}definedName'):
            if defined.get('name')=='_xlnm._FilterDatabase':
                index=int(defined.get('localSheetId'))
                count=stats[sheets[index][0]]['remaining']+1
                defined.text=re.sub(r'(\$[A-Z]+\$)\d+$',lambda m:m[1]+str(count),defined.text)
        replacements['xl/workbook.xml']=etree.tostring(wb)
        with zipfile.ZipFile(args.output,'x',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as out:
            for item in source.infolist():
                out.writestr(item.filename,replacements.get(item.filename,source.read(item.filename)))
    print(json.dumps(dict(event='saved_verifying',requested=len(targets),remainingIccids=remaining,
                          removedRows=sum(x['removed'] for x in stats.values())),ensure_ascii=False),flush=True)
    with zipfile.ZipFile(args.output) as output:
        assert output.testzip() is None
        for member in output.namelist():
            data=output.read(member)
            if member.endswith(('.xml','.rels')):
                assert not ({normalize(x.group().decode()) for x in TOKEN.finditer(data)} & targets), member
        for name,member in sheets:
            data=output.read(member);h=hashlib.sha256();count=0
            for count,m in enumerate(ROW.finditer(data),1):
                assert int(re.search(rb'\br="(\d+)"',m.group())[1])==count
                h.update(content_hash(m.group()))
            assert h.hexdigest()==stats[name]['retainedContentHash'],name
            assert count==stats[name]['remaining']+1,name
    checked=load_workbook(args.output,read_only=True,data_only=False)
    actual_keys=set()
    try:
        for sheet in checked:
            count=0
            for index,row in enumerate(sheet.iter_rows(),1):
                count+=1
                for cell in row:
                    if cell.data_type in ('f','e'):
                        raise ValueError('输出存在公式或错误: '+sheet.title)
                    if isinstance(cell.value,str) and any(normalize(m.group()) in targets for m in re.finditer(r'(?<![0-9])[0-9]{18,22}[Ff]?(?![0-9])',cell.value)):
                        raise ValueError('独立回读发现删除标识: '+sheet.title)
                if index>1 and sheet.title in DATA_SHEETS:
                    assert row[0].data_type=='s'
                if index>1 and sheet.title=='查询总览':
                    actual_keys.add(normalize(row[0].value))
            assert count==stats[sheet.title]['remaining']+1
        assert actual_keys==original_keys-targets
    finally:
        checked.close()
    assert hashlib.sha256(args.source.read_bytes()).hexdigest()==source_hash
    receipt=dict(status='passed',source=str(args.source),sourceSha256=source_hash,output=str(args.output),
                 outputSha256=hashlib.sha256(args.output.read_bytes()).hexdigest(),requestedIccids=len(targets),
                 removedIccids=len(targets),remainingIccids=remaining,removedRows=sum(x['removed'] for x in stats.values()),
                 sourceUnchanged=True,allPartsNoRemovedIdentifiers=True,independentReadback=True,
                 preservedRemainingValuesAndStyles=True,sheets=stats)
    with args.receipt.open('x',encoding='utf-8') as f:
        json.dump(receipt,f,ensure_ascii=False,indent=2)
    print(json.dumps(receipt,ensure_ascii=False),flush=True)


if __name__=='__main__':main()
