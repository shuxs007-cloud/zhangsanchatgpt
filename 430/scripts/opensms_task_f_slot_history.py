#!/usr/bin/env python3
"""Offline, read-only Task F ICCID/slot evidence search in historical Excel files.

Does not update the production database, existing workbooks or any remote system.
All dates are file/archive dates, not proof of the card's current location.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor, as_completed
from datetime import date, datetime
import hashlib
import io
import json
import os
from pathlib import Path
import posixpath
import re
import sqlite3
import subprocess
import threading
import time
import zipfile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB = Path('/Users/zhangsan/Downloads/430/opensms_artifacts/state/task_f_kooda_activation.sqlite3')
DEFAULT_ROOTS = ['/Users/zhangsan/Downloads', '/Users/zhangsan/Documents',
                 '/Users/zhangsan/Desktop', '/Volumes/zs01']
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ROW = re.compile(rb'<(?:[\w.-]+:)?row\b[^>]*(?:/>|>.*?</(?:[\w.-]+:)?row\s*>)', re.S)
SLOT = re.compile(r'[A-Za-z0-9]+_[0-9]+(?:\.[0-9]+)?\Z')
DATE = re.compile(r'(?<!\d)(20\d{2})[-_]?([01]\d)[-_]?([0-3]\d)(?!\d)')
ICCID_HEADERS = {'iccid', '卡存储', '卡存儲', 'simnumber', 'sim卡号', 'sim卡號',
                 '系统iccid', '系統iccid', 'googleiccid', 'queryiccid', 'sourceiccid'}
SLOT_HEADERS = {'卡槽', '卡槽标识', '卡槽標識', '卡槽（记录时）', '卡槽(记录时)',
                'slot', 'slotidentif', '卡槽标识符', '归一化卡槽', '卡槽（記錄時）'}
EXCLUDED_PARTS = {'node_modules', '.git', '.venv', '__pycache__', 'tests', 'fixtures',
                  '.Trash', '.Trashes', 'slot_history', 'opensms_slot_history'}
SHEET_CACHE = {}
CACHE_LOCK = threading.Lock()


def normalize_iccid(value):
    if not isinstance(value, str):
        return None
    value = value.strip()
    if value[-1:] in ('F', 'f'):
        value = value[:-1]
    return value if re.fullmatch(r'[0-9]{18,22}', value) else None


def header(value):
    return re.sub(r'\s+', '', str(value or '')).lower()


def local_name(element):
    return etree.QName(element).localname


def file_date(path, stat):
    # Closest dated filename/directory wins over copy/restore mtime.
    for part in reversed(path.parts):
        for match in reversed(list(DATE.finditer(part))):
            try:
                return date(*map(int, match.groups())).isoformat(), '文件名或目录日期'
            except ValueError:
                pass
    return datetime.fromtimestamp(stat.st_mtime).date().isoformat(), '文件修改日期（非采集时间）'


def targets_from_db(path):
    with sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True) as conn:
        conn.execute('PRAGMA query_only=ON')
        conn.execute('BEGIN')
        rows = conn.execute("SELECT query_iccid, source_slot, queried_at FROM kooda_activation "
                            "WHERE lookup_status='activation_missing' ORDER BY query_iccid").fetchall()
        inventory = dict(conn.execute('SELECT query_iccid, slot FROM kooda_inventory'))
    if any(normalize_iccid(row[0]) != row[0] for row in rows):
        raise ValueError('生产库有非法 ICCID，停止')
    return {r[0]: {'activationSlot': r[1] or '', 'queriedAt': r[2],
                   'inventorySlot': inventory.get(r[0]) or ''} for r in rows}


def discover(roots, start, end, output):
    existing = [str(Path(p).resolve()) for p in roots if Path(p).is_dir()]
    missing = [str(p) for p in roots if not Path(p).is_dir()]
    command = ['rg', '--files', '--hidden', '--no-ignore', *existing,
               '-g', '*.xlsx', '-g', '*.xlsm', '-g', '*.xls', '-g', '*.xlsb',
               '-g', '*.xltx', '-g', '*.xltm', '-g', '!._*', '-g', '!~$*', '-g', '!.~*',
               '-g', '!node_modules/**', '-g', '!.git/**', '-g', '!.venv/**']
    proc = subprocess.run(command, capture_output=True, text=True, check=False)
    if proc.returncode not in (0, 1):
        raise RuntimeError('文件发现失败: ' + proc.stderr[:2000])
    selected, skipped = [], Counter()
    for name in sorted(set(proc.stdout.splitlines())):
        path = Path(name)
        if set(path.parts) & EXCLUDED_PARTS or output in path.parents:
            skipped['开发测试或本扫描输出'] += 1
            continue
        stat = path.stat()
        day, basis = file_date(path, stat)
        if not start <= day <= end:
            skipped['一个月范围外'] += 1
            continue
        selected.append({'path': str(path), 'date': day, 'dateBasis': basis,
                         'size': stat.st_size, 'mtimeNs': stat.st_mtime_ns})
    return selected, {'roots': existing, 'missingRoots': missing,
                      'discovered': len(set(proc.stdout.splitlines())), 'excluded': dict(skipped)}


def fingerprint(item):
    result = dict(item)
    try:
        digest = hashlib.sha256()
        with open(item['path'], 'rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
        stat = Path(item['path']).stat()
        if (stat.st_size, stat.st_mtime_ns) != (item['size'], item['mtimeNs']):
            raise ValueError('读取期间源文件发生变化')
        result['sha256'] = digest.hexdigest()
    except Exception as exc:
        result['error'] = str(exc)
    return result


def rich_text(element):
    return ''.join(t.text or '' for t in element.iter() if local_name(t) == 't'
                   and (t.getparent() is None or local_name(t.getparent()) != 'rPh'))


def cell_values(fragment, strings, namespaces):
    declarations = ' '.join(('xmlns' + (':' + key if key else '') + '="' + value + '"')
                            for key, value in namespaces.items())
    root = etree.fromstring(('<wrapper ' + declarations + '>').encode() + fragment + b'</wrapper>',
                           etree.XMLParser(resolve_entities=False, no_network=True))
    row = root[0]
    values = {}
    for cell in row:
        if local_name(cell) != 'c':
            continue
        ref = cell.get('r', '')
        col = re.sub(r'\d', '', ref)
        kind = cell.get('t', 'n')
        formula = any(local_name(c) == 'f' for c in cell)
        raw = next((c.text or '' for c in cell if local_name(c) == 'v'), '')
        if kind == 's':
            value = strings[int(raw)] if raw else ''
        elif kind == 'inlineStr':
            value = rich_text(cell)
        else:
            value = raw
        values[col] = (value, kind, formula, ref)
    return int(row.get('r', '0')), values


def scan_sheet(xml, strings, string_hash, targets):
    cache_key = (string_hash, hashlib.sha256(xml).hexdigest(), tuple(sorted(targets)))
    with CACHE_LOCK:
        if cache_key in SHEET_CACHE:
            return SHEET_CACHE[cache_key]
    string_ids = [str(i).encode() for i, value in enumerate(strings) if normalize_iccid(value) in targets]
    # Scan long digit tokens once, then O(1) set lookup; do not compare every cell
    # against hundreds of long ICCID alternatives. Never truncate a longer token.
    literal = re.compile(rb'[0-9]{18,}[Ff]?')
    target_bytes = {key.encode() for key in targets}
    reference = (re.compile(rb'<(?:[\w.-]+:)?v>\s*(?:' + b'|'.join(string_ids)
                           + rb')\s*</(?:[\w.-]+:)?v>') if string_ids else None)
    def contains(data):
        return (any(match.group().removesuffix(b'F').removesuffix(b'f') in target_bytes
                    for match in literal.finditer(data))
                or (reference is not None and reference.search(data)))
    # Decompression, hash and C-level byte search precede Python row decoding.
    if not contains(xml):
        result = {'evidence': [], 'unresolved': [], 'rejected': {}}
    else:
        namespaces = next(etree.iterparse(io.BytesIO(xml), events=('start',),
                                         resolve_entities=False, no_network=True))[1].nsmap
        iccid_cols, slot_cols, header_row = [], [], 0
        for match in list_first_rows(xml, 50):
            rownum, values = cell_values(match, strings, namespaces)
            ids = [col for col, v in values.items() if header(v[0]) in ICCID_HEADERS]
            slots = [col for col, v in values.items() if header(v[0]) in SLOT_HEADERS]
            if ids and slots:
                iccid_cols, slot_cols, header_row = ids, slots, rownum
                break
        evidence, unresolved, rejected = [], [], Counter()
        for match in ROW.finditer(xml):
            fragment = match.group()
            if not contains(fragment):
                continue
            rownum, values = cell_values(fragment, strings, namespaces)
            for col, (value, kind, formula, ref) in values.items():
                key = normalize_iccid(value)
                if key not in targets:
                    continue
                base = {'iccid': key, 'row': rownum, 'iccidCell': ref}
                if kind not in ('s', 'str', 'inlineStr') or formula:
                    rejected['ICCID非可靠文本或公式'] += 1
                    unresolved.append({**base, 'reason': 'ICCID非可靠文本或公式'})
                    continue
                if not header_row or rownum <= header_row or col not in iccid_cols:
                    unresolved.append({**base, 'reason': '命中ICCID但缺少同表明确ICCID及卡槽列'})
                    continue
                slots = []
                for slotcol in slot_cols:
                    v, t, f, cellref = values.get(slotcol, ('', 'n', False, ''))
                    if v and not f and t in ('s', 'str', 'inlineStr') and SLOT.fullmatch(v.strip()):
                        slots.append((v.strip(), cellref))
                if not slots:
                    unresolved.append({**base, 'reason': '卡槽空白或不是明确完整卡槽'})
                    continue
                for slot, cellref in slots:
                    evidence.append({**base, 'slot': slot, 'slotCell': cellref,
                                     'ambiguousColumns': len(iccid_cols) != 1 or len(slot_cols) != 1})
        result = {'evidence': evidence, 'unresolved': unresolved, 'rejected': dict(rejected)}
    with CACHE_LOCK:
        SHEET_CACHE[cache_key] = result
    return result


def list_first_rows(xml, count):
    for number, match in enumerate(ROW.finditer(xml)):
        if number >= count:
            break
        yield match.group()


def scan_workbook(item, targets):
    result = {'sha256': item['sha256'], 'evidence': [], 'unresolved': [], 'sheets': 0, 'rejected': Counter()}
    try:
        if not zipfile.is_zipfile(item['path']):
            raise ValueError('文件内容不是OOXML工作簿；可能为下载响应或不支持的旧二进制格式，不作为卡槽证据')
        with zipfile.ZipFile(item['path']) as book:
            # .xls files containing OOXML are accepted; binary XLS/XLSB remain explicit errors.
            workbook = etree.fromstring(book.read('xl/workbook.xml'))
            rels = etree.fromstring(book.read('xl/_rels/workbook.xml.rels'))
            targets_by_id = {r.get('Id'): posixpath.normpath('xl/' + r.get('Target', ''))
                             if not r.get('Target', '').startswith('/') else r.get('Target').lstrip('/')
                             for r in rels if r.get('TargetMode') != 'External'}
            raw_strings = book.read('xl/sharedStrings.xml') if 'xl/sharedStrings.xml' in book.namelist() else b''
            strings = [rich_text(si) for si in etree.fromstring(raw_strings)] if raw_strings else []
            string_hash = hashlib.sha256(raw_strings).hexdigest()
            sheets = workbook.findall('.//{' + NS + '}sheet')
            if not sheets:
                raise ValueError('工作簿没有可识别的标准工作表')
            for sheet in sheets:
                rid = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                member = targets_by_id.get(rid, '')
                if not member.endswith('.xml'):
                    raise ValueError('不支持的工作表格式: ' + member)
                xml = book.read(member)
                scanned = scan_sheet(xml, strings, string_hash, targets)
                result['sheets'] += 1
                for group in ('evidence', 'unresolved'):
                    result[group].extend({**row, 'sheet': sheet.get('name'), 'sheetMember': member}
                                         for row in scanned[group])
                result['rejected'].update(scanned['rejected'])
        stat = Path(item['path']).stat()
        if (stat.st_size, stat.st_mtime_ns) != (item['size'], item['mtimeNs']):
            raise ValueError('扫描期间源文件发生变化')
    except Exception as exc:
        result['error'] = type(exc).__name__ + ': ' + str(exc)
        result['evidence'] = []  # A partially decoded workbook is not accepted evidence.
    result['rejected'] = dict(result['rejected'])
    return result


def source_kind(path, sheet):
    text = (path + '/' + sheet).lower()
    if 'google' in text or '谷歌' in text:
        return 'Google或含Google关联来源（需系统核验）'
    if 'ko未激活' in sheet:
        return '本地库衍生表（不是新增来源）'
    if any(x in path.lower() for x in ('opensms', 'export4day', 'phone_config')):
        return '系统或历史处理表（非实时）'
    return '其它历史表（需核对来源）'


def summarize(targets, evidence):
    grouped = defaultdict(list)
    for row in evidence:
        grouped[row['iccid']].append(row)
    rows = []
    for key, original in sorted(targets.items()):
        found = grouped[key]
        slots = sorted({r['slot'] for r in found})
        qualified = [r for r in found if not r['ambiguousColumns']
                     and r['sourceKind'] != '本地库衍生表（不是新增来源）']
        qslots = sorted({r['slot'] for r in qualified})
        status = '未找到' if not slots else '多卡槽记录待核对' if len(slots) > 1 else '唯一历史卡槽'
        if slots and not qualified:
            status = '仅模糊列或衍生记录待核对'
        latest = max(found, key=lambda r: (r['date'], r['file'], r['sheet'], r['row'])) if found else {}
        rows.append({'iccid': key, 'status': status, 'slots': slots, 'qualifiedSlots': qslots,
                     'evidenceCount': len(found), 'latestFileDate': latest.get('date', ''),
                     'latestFile': latest.get('file', ''), 'latestSheet': latest.get('sheet', ''),
                     'inventorySlotBeforeSearch': original['inventorySlot']})
    return rows


def write_json(path, value):
    with path.open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)


def write_excel(path, summary, evidence, unresolved, files):
    book = Workbook(write_only=True)
    tables = [
        ('查询结果', ['ICCID', '查询结论', '历史卡槽集合', '最新来源日期（非当前位置）', '原库存卡槽', '最新来源文件', '工作表'],
         [[r['iccid'], r['status'], '；'.join(r['slots']), r['latestFileDate'], r['inventorySlotBeforeSearch'], r['latestFile'], r['latestSheet']] for r in summary]),
        ('精确匹配证据', ['ICCID', '历史卡槽', '来源日期', '日期依据', '来源类型', '来源文件', 'SHA256', '工作表', '行号', 'ICCID单元格', '卡槽单元格', '列映射需核对'],
         [[r[k] for k in ('iccid', 'slot', 'date', 'dateBasis', 'sourceKind', 'file', 'sha256', 'sheet', 'row', 'iccidCell', 'slotCell', 'ambiguousColumns')] for r in evidence]),
        ('命中但不能恢复', ['ICCID', '原因', '来源日期', '来源文件', '工作表', '行号', 'ICCID单元格'],
         [[r.get(k, '') for k in ('iccid', 'reason', 'date', 'file', 'sheet', 'row', 'iccidCell')] for r in unresolved]),
        ('扫描文件清单', ['来源文件', '日期', '日期依据', 'SHA256', '状态', '错误'],
         [[r.get(k, '') for k in ('path', 'date', 'dateBasis', 'sha256', 'scanStatus', 'error')] for r in files]),
    ]
    expected = {}
    for name, headers, records in tables:
        ws = book.create_sheet(name)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = 'A1:' + get_column_letter(len(headers)) + str(len(records)+1)
        for column, title in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(column)].width = 65 if '文件' in title else 30
        for rownum, row in enumerate([headers, *records], 1):
            cells = []
            for value in row:
                cell = WriteOnlyCell(ws, value=value)
                cell.font = Font(name='Arial', bold=rownum == 1,
                                 color='FFFFFF' if rownum == 1 else '000000')
                if rownum == 1:
                    cell.fill = PatternFill('solid', fgColor='284B63')
                if isinstance(value, str):
                    if len(value) > 32767:
                        raise ValueError('Excel字段过长，禁止截断')
                    cell.data_type = 's'
                    cell.number_format = '@'
                cells.append(cell)
            ws.append(cells)
        expected[name] = (headers, records)
    book.save(path)
    book.close()
    # Separate parser/open: exact values, ICCID text precision, zero formulas/errors.
    checked = load_workbook(path, read_only=True, data_only=False)
    try:
        for name, (headers, records) in expected.items():
            iterator = checked[name].iter_rows()
            if [c.value for c in next(iterator)] != headers:
                raise ValueError('输出独立回读表头错误')
            count = 0
            for actual, wanted in zip(iterator, records, strict=True):
                count += 1
                if [c.value for c in actual] != [v if v != '' else None for v in wanted]:
                    raise ValueError('输出独立回读值不一致')
                if any(c.data_type in ('f', 'e') for c in actual):
                    raise ValueError('输出包含公式或Excel错误')
                if name != '扫描文件清单' and actual[0].data_type != 's':
                    raise ValueError('ICCID不是文本')
            if count != len(records):
                raise ValueError('输出独立回读数量不一致')
    finally:
        checked.close()
    with zipfile.ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError('输出压缩包校验失败')
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--db', type=Path, default=DB)
    parser.add_argument('--root', action='append')
    parser.add_argument('--start', required=True, type=date.fromisoformat)
    parser.add_argument('--end', required=True, type=date.fromisoformat)
    parser.add_argument('--workers', type=int, default=min(8, os.cpu_count() or 4))
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if not 1 <= args.workers <= 8 or args.start > args.end:
        parser.error('workers须1至8，起止日期须有效')
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    started = time.monotonic()
    implementation_sha = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    targets = targets_from_db(args.db)
    write_json(output / 'target_snapshot.json', {'createdAt': datetime.now().astimezone().isoformat(),
               'database': str(args.db), 'readOnly': True, 'records': targets})
    files, discovery = discover(args.root or DEFAULT_ROOTS, args.start.isoformat(), args.end.isoformat(), output)
    print(json.dumps({'event': 'discovered', 'targetCards': len(targets), 'files': len(files),
                      'workers': args.workers}, ensure_ascii=False), flush=True)
    fingerprints = []
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        for item in pool.map(fingerprint, files):
            fingerprints.append(item)
    groups = defaultdict(list)
    for item in fingerprints:
        if 'error' not in item:
            groups[item['sha256']].append(item)
    print(json.dumps({'event': 'hash_complete', 'uniqueWorkbooks': len(groups),
                      'duplicateFiles': sum(len(v)-1 for v in groups.values())}), flush=True)
    evidence, unresolved, errors = [], [], []
    done, last_progress = 0, time.monotonic()
    with ProcessPoolExecutor(max_workers=args.workers) as pool:
        futures = {pool.submit(scan_workbook, rows[0], set(targets)): digest for digest, rows in groups.items()}
        for future in as_completed(futures):
            digest = futures[future]
            result = future.result()
            aliases = groups[digest]
            for item in aliases:
                item['scanStatus'] = 'failed' if result.get('error') else 'scanned'
                if result.get('error'):
                    item['error'] = result['error']
                    errors.append({'file': item['path'], 'error': result['error']})
                for group, destination in [('evidence', evidence), ('unresolved', unresolved)]:
                    for row in result[group]:
                        destination.append({**row, 'file': item['path'], 'date': item['date'],
                                            'dateBasis': item['dateBasis'], 'sha256': digest,
                                            'sourceKind': source_kind(item['path'], row['sheet'])})
            done += 1
            if time.monotonic() - last_progress >= 15 or done == len(groups):
                progress = {'event': 'scan_progress', 'completed': done, 'total': len(groups),
                            'matchedCards': len({r['iccid'] for r in evidence}),
                            'updatedAt': datetime.now().astimezone().isoformat()}
                temporary = output / 'progress.json.tmp'
                with temporary.open('w', encoding='utf-8') as stream:
                    json.dump(progress, stream, ensure_ascii=False)
                temporary.replace(output / 'progress.json')
                print(json.dumps(progress), flush=True)
                last_progress = time.monotonic()
    for item in fingerprints:
        if 'error' in item and 'scanStatus' not in item:
            item['scanStatus'] = 'failed'
            errors.append({'file': item['path'], 'error': item['error']})
    evidence.sort(key=lambda r: (r['iccid'], r['date'], r['file'], r['sheet'], r['row'], r['slot']))
    summary = summarize(targets, evidence)
    write_json(output / 'evidence.json', evidence)
    write_json(output / 'unresolved.json', unresolved)
    write_json(output / 'files.json', fingerprints)
    write_json(output / 'results.json', summary)
    excel = output / '任务F未激活卡_近一个月历史卡槽核对.xlsx'
    digest = write_excel(excel, summary, evidence, unresolved, fingerprints)
    counts = dict(Counter(r['status'] for r in summary))
    kinds_by_card = defaultdict(set)
    for row in evidence:
        kinds_by_card[row['iccid']].add(row['sourceKind'])
    source_coverage = {'googleOnlyCards': sum(all('Google' in kind for kind in kinds)
                                             for kinds in kinds_by_card.values()),
                       'hasSystemOrProcessedEvidenceCards': sum('系统或历史处理表（非实时）' in kinds
                                                               for kinds in kinds_by_card.values())}
    report = {'status': 'completed_with_errors' if errors else 'completed', 'readOnlyProduction': True,
              'databaseUpdated': False, 'backendAccessed': False, 'start': str(args.start), 'end': str(args.end),
              'targetCards': len(targets), 'selectedFiles': len(fingerprints), 'uniqueWorkbooks': len(groups),
              'duplicateFiles': len(fingerprints)-len(groups), 'counts': counts, 'evidenceRows': len(evidence),
              'sourceCoverage': source_coverage,
              'unresolvedRows': len(unresolved), 'scanErrors': errors, 'discovery': discovery,
              'workers': args.workers, 'parallelism': 'threaded_hashing_process_pool_xml_scanning_single_workbook_writer',
              'scriptSha256AtStart': implementation_sha,
              'durationSeconds': round(time.monotonic()-started, 2), 'workbook': str(excel),
              'workbookSha256': digest, 'independentOutputReadbackPassed': True}
    write_json(output / 'summary.json', report)
    md = ['# 任务F未激活卡历史Excel卡槽查询', '',
          f'范围：{args.start} 至 {args.end}（含首尾）。只读生产库及历史文件，无联网、无回填。', '',
          f'目标 {len(targets)} 张；文件 {len(fingerprints)} 份；按SHA256去重后 {len(groups)} 份。', '',
          '| 查询结论 | 卡数 |', '|---|---:|', *[f'| {k} | {v} |' for k, v in counts.items()], '',
          '卡槽均为历史证据，不能冒充当前位置。带r设备不合并；多卡槽、Google来源和模糊列须核对。',
          f'仅Google来源卡数：{source_coverage["googleOnlyCards"]}；有系统或历史处理表证据卡数：{source_coverage["hasSystemOrProcessedEvidenceCards"]}。处理表也不等于原始系统权威记录。',
          '18至22位ICCID仅接受文本，去末尾恰好一个F/f后精确匹配；数字或公式不作可靠恢复依据。',
          '有明确文件名/归档日期时优先使用，否则按文件修改日期筛选并标注。无法读取的文件不计为未命中。', '',
          f'文件级错误：{len(errors)}；详细清单见 files.json 及工作簿。', '',
          f'[核对Excel]({excel})', f'[完整汇总]({output / "summary.json"})',
          f'[逐条来源证据]({output / "evidence.json"})']
    with (output / '报告.md').open('x', encoding='utf-8') as stream:
        stream.write('\n'.join(md) + '\n')
    print(json.dumps(report, ensure_ascii=False), flush=True)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
