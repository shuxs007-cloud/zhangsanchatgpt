#!/usr/bin/env python3
"""Read-only analysis of an existing report; never runs any business task."""
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from zipfile import ZipFile

B14 = {'B14-C快下保留卡', 'B14-T快下保留卡'}

def encode(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value)

def main():
    mode, source, folder = sys.argv[1:4]
    source, folder = Path(source), Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    if mode == 'prepare':
        workbook = load_workbook(source, read_only=True, data_only=True)
        iterator = workbook['sheet1'].values
        headers = next(iterator)
        selected = []
        for n, values in enumerate(iterator, 2):
            row = dict(zip(headers, values))
            if row.get('过期状态') == '快过期' or row.get('运营商') in B14:
                phone = str(row['电话'])
                key = phone[1:] if len(phone) == 11 and phone.startswith('1') else phone
                if len(key) != 10 or not key.isdigit():
                    raise ValueError('Invalid phone')
                selected.append({'phone10': key, 'rowNumber': n, 'report': row})
        workbook.close()
        payload = {'scope': 'existing_report_fast_expiry_and_b14', 'sourceWorkbook': str(source), 'rows': selected}
        with (folder / 'input.json').open('x') as f:
            json.dump(payload, f, ensure_ascii=False, default=encode)
        print(json.dumps({'selected': len(selected), 'input': str(folder / 'input.json')}))
        return
    if mode != 'report':
        raise ValueError('Mode must be prepare or report')
    data = json.loads((folder / 'input.json').read_text())
    live = json.loads((folder / 'live.json').read_text())
    live_map = {r['phone10']: r for r in live['rows']}
    wb = Workbook()
    ws = wb.active
    ws.title = '快过期核对'
    reserve = wb.create_sheet('B14保留区核对')
    header = ['电话', '报表运营商', '当前运营商', '过期状态', '截止日', 'OPR', '卡槽', '核对结论', '报表原始行号']
    ws.append(header)
    reserve.append(header)
    stats, retained, details = Counter(), Counter(), []
    for item in data['rows']:
        r = item['report']
        current = live_map.get(item['phone10'])
        carrier = current['carrierName'] if current else '后台未匹配'
        if current is None:
            conclusion = '后台未匹配，不能确认当前区'
        elif carrier in B14:
            conclusion = '按本次要求应保护，任务B现无B14排除'
        elif carrier in {'快过期隔离区C', '快过期隔离区T'}:
            conclusion = '已在快过期隔离区'
        else:
            conclusion = '报表快过期，当前未在隔离区且非B14'
        row = [str(r['电话']), str(r['运营商']), carrier, str(r.get('过期状态') or ''),
               str(r.get('截止日') or '')[:10], str(r.get('opr') or ''), str(r.get('卡槽') or ''), conclusion, item['rowNumber']]
        if r.get('过期状态') == '快过期':
            ws.append(row)
            stats[carrier] += 1
            details.append(dict(zip(header, row)))
        if carrier in B14:
            reserve.append(row)
            retained[str(r.get('过期状态') or '无过期标记')] += 1
    for sheet in wb:
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        for col, width in zip('ABCDEFGHI', [18, 25, 25, 17, 18, 27, 20, 58, 18]):
            sheet.column_dimensions[col].width = width
        for cell in sheet[1]:
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='16324F')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
            row[0].number_format = '@'
    excel = folder / '快过期与B14保护_只读核查.xlsx'
    if excel.exists():
        raise FileExistsError(excel)
    expected = {s.title: list(s.values) for s in wb}
    wb.save(excel)
    with ZipFile(excel) as z:
        assert z.testzip() is None
    verify = load_workbook(excel, read_only=True, data_only=True)
    for s in verify:
        assert list(s.values) == [tuple(None if x == '' else x for x in r) for r in expected[s.title]]
    verify.close()
    result = {'status': 'passed', 'businessWrites': 0, 'sourceWorkbook': data['sourceWorkbook'],
              'liveCheckedAt': live['checkedAt'], 'fastByCurrentCarrier': dict(stats),
              'b14ExpiryStatesInReportScope': dict(retained), 'rows': details, 'excel': str(excel)}
    with (folder / 'result.json').open('x') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(json.dumps({**result, 'rows': None}, ensure_ascii=False))

if __name__ == '__main__':
    main()
