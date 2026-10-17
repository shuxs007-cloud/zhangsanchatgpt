"""Offline first-appearance search; originals, business DBs and Task F stay read-only.

Presence search intentionally examines every text cell, including hidden sheets
and sheets without an ICCID header. It does not infer phone/card relationships.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import io
import json
import os
import re
import signal
import sqlite3
import sys
import time
import zipfile
from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from lxml import etree as E
import opensms_task_i as I

TZ = ZoneInfo('Asia/Hong_Kong')
PARSER = E.XMLParser(resolve_entities=False, no_network=True)
READ_TIMEOUT_SECONDS = 30


@contextmanager
def local_file_read(path, seconds=READ_TIMEOUT_SECONDS):
    """Never hydrate macOS placeholders; bound a worker's local file read."""
    if sys.platform == 'darwin' and getattr(Path(path).stat(), 'st_flags', 0) & 0x40000000:
        # SF_DATALESS from the macOS SDK sys/stat.h; opening may hydrate online.
        raise OSError('本地云文件占位符未下载；不触发下载，列为覆盖缺口')
    previous = signal.getsignal(signal.SIGALRM)
    if signal.getitimer(signal.ITIMER_REAL)[0]:
        raise RuntimeError('文件读取已有计时器，拒绝覆盖')
    def expired(signum, frame):
        raise TimeoutError(f'本地文件读取超过{seconds}秒；列为覆盖缺口')
    signal.signal(signal.SIGALRM, expired)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        yield
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def dump(path, value):
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2))
    tmp.replace(path)


def sha256(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def previous_month(day):
    year, month = (day.year - 1, 12) if day.month == 1 else (day.year, day.month - 1)
    return day.replace(year=year, month=month,
                       day=min(day.day, calendar.monthrange(year, month)[1]))


def time_hint(path):
    """Do not use mtime, workbook template metadata or unrelated chat folders."""
    path = Path(path)
    candidates = [(path.name, '文件名')]
    candidates += [(x, '归档运行编号') for x in reversed(path.parts[:-1])
                   if re.search(r'20\d{6}T\d{6}(?:\+0800|Z)?', x)]
    for name, basis in candidates:
        m = re.search(r'(?<!\d)(20\d{6})[T_](\d{6})(Z|\+0800)?(?!\d)', name)
        if m:
            try:
                stamp = datetime.strptime(m[1] + m[2], '%Y%m%d%H%M%S').replace(tzinfo=TZ)
                if m[3] == 'Z':
                    stamp += timedelta(hours=8)
                return dict(observed=stamp.isoformat(), precision='second', basis=basis)
            except ValueError:
                pass
    epoch = re.match(r'^(?:export4day|phoneConfig|export)[_-](1[6-9]\d{11})(?!\d)', path.name)
    if epoch:
        stamp = datetime.fromtimestamp(int(epoch[1]) / 1000, TZ)
        return dict(observed=stamp.isoformat(), precision='second', basis='导出文件名Unix毫秒')
    m = re.search(r'(?<!\d)(20\d{2})[-_]?([01]\d)[-_]?([0-3]\d)(?!\d)', path.name)
    if m:
        try:
            stamp = datetime(*map(int, m.groups()), tzinfo=TZ)
            return dict(observed=stamp.isoformat(), precision='day', basis='文件名日期，仅日级')
        except ValueError:
            pass
    return dict(observed=None, precision='unknown', basis='时间未知，修改时间不用于先后排序')


def normalize_list(path):
    keys, originals, invalid = [], {}, []
    for n, raw in enumerate(path.read_text(encoding='utf-8-sig').splitlines(), 1):
        if not raw.strip() or I.header(raw) == 'iccid':
            continue
        key = I.iccid_key(raw)
        if not key:
            invalid.append(dict(line=n, value=raw, reason='不是可靠的18至22位文本ICCID'))
            continue
        if key not in originals:
            keys.append(key)
            originals[key] = raw.strip()
    if not keys:
        raise ValueError('没有有效ICCID')
    return keys, originals, invalid


def fingerprint(item):
    try:
        if item.get('error'):
            raise ValueError(item['error'])
        if item['path'].startswith(I.UNAVAILABLE_PREFIX):
            raise ValueError('旧同步目录已知读取阻塞，保留为覆盖缺口，不冒充未命中')
        with local_file_read(item['path'], READ_TIMEOUT_SECONDS):
            return dict(item, signature=I.content_key(Path(item['path'])), status='ready')
    except Exception as exc:
        return dict(item, status='unreadable', error=str(exc))


def same_stat(item):
    st = Path(item['path']).stat()
    if (st.st_size, st.st_mtime_ns) != (item['size'], item['mtime']):
        raise ValueError('扫描期间源文件变化')


def possible_text_hit(xml, wanted):
    """Conservative negative filter; rich/escaped text always gets decoded."""
    return (any(prefix in xml for prefix in {k[:8].encode() for k in wanted}) or
            bool(re.search(rb'<(?:[\w.-]+:)?r(?:\s|>)', xml)) or b'&#' in xml)


def _scan_unbounded(item, keys):
    result = dict(path=item['path'], signature=item['signature'], hits=[], status='scanned', pid=os.getpid(), priority=I.WORKER_PRIORITY)
    try:
        same_stat(item)
        path, wanted = Path(item['path']), set(keys)
        if zipfile.is_zipfile(path):
            with zipfile.ZipFile(path) as z:
                if 'xl/workbook.xml' not in z.namelist():
                    raise ValueError('不是OOXML工作簿')
                matched = {}
                if 'xl/sharedStrings.xml' in z.namelist():
                    raw_strings = z.read('xl/sharedStrings.xml')
                    if possible_text_hit(raw_strings, wanted):
                        root = E.fromstring(raw_strings, PARSER)
                        for number, si in enumerate(root):
                            value = I.rich_text(si)
                            key = I.iccid_key(value)
                            if key in wanted:
                                matched[number] = (key, value)
                wb = E.fromstring(z.read('xl/workbook.xml'), PARSER)
                rels = E.fromstring(z.read('xl/_rels/workbook.xml.rels'), PARSER)
                members = {r.get('Id'): r.get('Target').lstrip('/') if r.get('Target', '').startswith('/')
                           else I.posixpath.normpath('xl/' + r.get('Target', ''))
                           for r in rels if r.get('TargetMode') != 'External'}
                shared_pattern = re.compile(rb'<(?:[\w.-]+:)?v>\s*(?:' +
                    b'|'.join(str(n).encode() for n in matched) + rb')\s*</') if matched else None
                for sheet in wb.findall('.//{' + I.NS + '}sheet'):
                    xml = z.read(members[sheet.get(I.RID)])
                    inline = b'inlineStr' in xml or b't="str"' in xml or b"t='str'" in xml
                    # Fast negative filter for plain inline text. Rich text or
                    # numeric character entities must still take the full path,
                    # because a target's digits may be split/escaped in XML.
                    if inline:
                        inline = possible_text_hit(xml, wanted)
                    if not inline and (shared_pattern is None or not shared_pattern.search(xml)):
                        continue
                    for cell in I.CELL.finditer(xml):
                        ref, kind = I.REF.search(cell[1]), I.KIND.search(cell[1])
                        body = cell[2] or b''
                        if not ref or I.FORMULA.search(body):
                            continue
                        kind = kind[1].decode() if kind else 'n'
                        key = value = None
                        if kind == 's' and matched:
                            raw = I.VALUE.search(body)
                            if raw and raw[1].strip().isdigit():
                                key, value = matched.get(int(raw[1]), (None, None))
                        elif kind in ('inlineStr', 'str'):
                            value = I.unescape(b''.join(I.TEXT.findall(body)).decode()) if kind == 'inlineStr' else (
                                I.unescape(I.VALUE.search(body)[1].decode()) if I.VALUE.search(body) else '')
                            key = I.iccid_key(value)
                        if key in wanted:
                            result['hits'].append(dict(iccid=key, raw=value, sheet=sheet.get('name'),
                                sheetState=sheet.get('state', 'visible'), cell=ref[1].decode() + ref[2].decode(),
                                row=int(ref[2])))
        elif path.suffix.lower() in ('.csv', '.tsv'):
            raw = path.read_bytes()
            encodings = ('utf-16', 'utf-8-sig', 'gb18030') if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'gb18030')
            for encoding in encodings:
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError('无法确定文本编码')
            csv.field_size_limit(16 * 1024 * 1024)
            for row, values in enumerate(csv.reader(io.StringIO(text), delimiter='\t' if path.suffix.lower() == '.tsv' else ','), 1):
                for column, value in enumerate(values, 1):
                    key = I.iccid_key(value)
                    if key in wanted:
                        result['hits'].append(dict(iccid=key, raw=value, sheet='文本表', sheetState='visible', cell=f'第{column}列', row=row))
        else:
            raise ValueError('旧二进制XLS无可靠离线读取器，未扫描内容')
        same_stat(item)
    except Exception as exc:
        result.update(status='unreadable', error=str(exc), hits=[])
    return result


def scan(item, keys):
    try:
        with local_file_read(item['path'], 180):
            return _scan_unbounded(item, keys)
    except Exception as exc:
        return dict(path=item['path'], signature=item['signature'], hits=[],
                    status='unreadable', error=str(exc), pid=os.getpid(), priority=I.WORKER_PRIORITY)


def _scan_job(job):
    return scan(*job)


def earliest_candidates(hits):
    timed = [h for h in hits if h['observed']]
    if not timed:
        return []
    def end(h):
        return datetime.fromisoformat(h['observed']) + timedelta(days=1) if h['precision'] == 'day' else datetime.fromisoformat(h['observed']) + timedelta(microseconds=1)
    earliest_end = min(map(end, timed))
    # Retain every overlapping day-level interval, never turn midnight into an
    # invented exact timestamp or arbitrarily choose between tied documents.
    selected = [h for h in timed if datetime.fromisoformat(h['observed']) < earliest_end]
    return sorted(selected, key=lambda h: (h['observed'], 'OpenSMS近10日新鲜导出_' not in h['path'], h['path'], h['sheet'], h['row'], h['cell']))


def verify_file(job):
    with local_file_read(job[0]['path'], 180):
        return _verify_file_unbounded(job)


def _verify_file_unbounded(job):
    item, hits = job
    same_stat(item)
    path = Path(item['path'])
    if zipfile.is_zipfile(path):
        # Independent parser, not the regex/shared-string search above.
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            grouped = defaultdict(list)
            for hit in hits:
                grouped[hit['sheet']].append(hit)
            for sheet, rows in grouped.items():
                needed = {h['cell']: h for h in rows}
                ws = wb[sheet]
                for row in ws.iter_rows(min_row=min(h['row'] for h in rows), max_row=max(h['row'] for h in rows)):
                    for cell in row:
                        if getattr(cell, 'coordinate', None) in needed:
                            hit = needed.pop(cell.coordinate)
                            if cell.data_type not in ('s', 'inlineStr') or I.iccid_key(cell.value) != hit['iccid'] or cell.value != hit['raw']:
                                raise ValueError(f'最早命中独立回读不一致：{sheet}!{cell.coordinate}')
                if needed:
                    raise ValueError('最早命中独立回读缺少单元格')
        finally:
            wb.close()
    else:
        # Reopen and check the requested coordinates without reusing scan().
        raw = path.read_bytes()
        for encoding in ('utf-8-sig', 'utf-16', 'gb18030'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError('独立文本回读无法确定编码')
        rows = list(csv.reader(io.StringIO(text), delimiter='\t' if path.suffix.lower() == '.tsv' else ','))
        for h in hits:
            column = int(re.fullmatch(r'第(\d+)列', h['cell'])[1])
            value = rows[h['row'] - 1][column - 1]
            if value != h['raw'] or I.iccid_key(value) != h['iccid']:
                raise ValueError('文本表独立坐标回读不一致')
    digest = sha256(path)
    same_stat(item)
    return dict(path=str(path), sha256=digest, verifiedCells=len(hits), passed=True, pid=os.getpid())


def index_reference(db_path, keys):
    """Read only the separate Task I index; never open a business/Task F DB."""
    if not db_path.exists():
        return dict(status='absent')
    with sqlite3.connect(db_path.resolve().as_uri() + '?mode=ro', uri=True, timeout=5) as db:
        db.execute('PRAGMA query_only=ON')
        meta = dict(db.execute('SELECT key,value FROM meta'))
        values = db.execute('SELECT r.iccid,MIN(f.observed),COUNT(DISTINCT f.path) FROM records r '
            'JOIN file_sheets fs ON fs.sheet_id=r.sheet_id JOIN files f ON f.id=fs.file_id '
            'WHERE f.active=1 AND r.iccid IN (' + ','.join('?' for _ in keys) + ') GROUP BY r.iccid', keys).fetchall()
        return dict(status=meta.get('index_state', 'unknown'), db=str(db_path), indexedAt=meta.get('updated_at'),
                    matchedIdentifiers=len(values), note='旧索引仅作参考；本次结论来自磁盘重新扫描，不用索引缺失证明未命中')


def probe_older_index(result_path):
    """Verify positive older index leads; not a full scan of the older period."""
    out = result_path.parent
    output = out/'older_index_evidence.json'
    if output.exists():
        existing = json.loads(output.read_text())
        if existing['monthlyResultSha256'] != sha256(result_path):
            raise ValueError('窗口外证据与本轮扫描结果哈希不一致')
        for evidence in existing['verifiedFiles']:
            if sha256(evidence['path']) != evidence['sha256']:
                raise ValueError('窗口外已核验源文件变化')
        return existing
    data = json.loads(result_path.read_text())
    keys = [q['iccid'] for q in data['queries']]
    index_path = Path(data['indexReference'].get('db', I.DEFAULT_DB))
    with sqlite3.connect(index_path.resolve().as_uri()+'?mode=ro', uri=True) as db:
        db.row_factory = sqlite3.Row
        rows = db.execute('SELECT r.iccid,f.path,f.observed FROM records r JOIN file_sheets fs ON fs.sheet_id=r.sheet_id '
            'JOIN files f ON f.id=fs.file_id WHERE f.active=1 AND f.date_basis<>? AND f.observed<? '
            'AND r.iccid IN ('+','.join('?' for _ in keys)+') ORDER BY f.observed',
            ['修改时间，非确定采集时间', data['since'][:10], *keys]).fetchall()
    earliest = {}
    for row in rows:
        hint = time_hint(row['path'])
        if hint['observed'] and hint['observed'][:10] < data['since'][:10]:
            candidate = (hint['observed'], row['path'])
            if row['iccid'] not in earliest or candidate < earliest[row['iccid']]:
                earliest[row['iccid']] = candidate
    hits, verified, errors = [], [], []
    for path in sorted({x[1] for x in earliest.values()}):
        try:
            st = Path(path).stat()
            item = fingerprint(dict(path=path, size=st.st_size, mtime=st.st_mtime_ns))
            if item['status'] != 'ready':
                raise ValueError(item['error'])
            scan_result = scan(item, keys)
            if scan_result['status'] != 'scanned':
                raise ValueError(scan_result['error'])
            proof = verify_file((item, scan_result['hits']))
            verified.append(proof)
            hits.extend(dict(h, path=path, **time_hint(path), fileSha256=proof['sha256'],
                             independentReadbackPassed=True) for h in scan_result['hits'])
        except Exception as exc:
            errors.append(dict(path=path, error=str(exc)))
    queries = []
    for q in data['queries']:
        found = earliest_candidates([h for h in hits if h['iccid'] == q['iccid']])
        if found:
            queries.append(dict(iccid=q['iccid'], input=q['input'], earliest=found))
    result = dict(status='passed' if not errors else 'partial', scope='仅核验旧索引最早的阳性线索，不等于全量扫描窗口之前所有文件',
                  queriedBefore=data['since'], matchedIdentifiers=len(queries), queries=queries,
                  verifiedFiles=verified, errors=errors, monthlyResultSha256=sha256(result_path), networkAccess=False)
    dump(output, result)
    return result


def build_report(result_path):
    import subprocess
    from opensms_task_i_pipeline import NODE, MARKER, WORKDIR
    out = result_path.parent
    result, receipt = json.loads(result_path.read_text()), json.loads((out/'run_receipt.json').read_text())
    if receipt.get('status') not in ('passed', 'passed_with_coverage_gaps') or sha256(result_path) != receipt['resultSha256']:
        raise ValueError('扫描结果未通过或来源哈希不符')
    for evidence in result['verification']['evidence']:
        if sha256(evidence['path']) != evidence['sha256']:
            raise ValueError('生成报告前最早来源文件变化')
    output = out / f"ICCID最早出现文档_{result['inputCount']}个.xlsx"
    if output.exists():
        raise ValueError('工作簿已存在，禁止覆盖')
    probe_older_index(result_path)
    subprocess.run([NODE, MARKER, '--operation-kind', 'create', '--expected-output-count', '1', '--output-format', 'xlsx'], check=True, cwd=WORKDIR)
    with (out/'report_build.log').open('w') as log:
        subprocess.run([NODE, str(Path(__file__).with_name('opensms_task_i_first_seen_report.mjs')), str(result_path), str(output)],
                       check=True, cwd=WORKDIR, stdout=log, stderr=subprocess.STDOUT)
    return check_report(result_path, output)


def check_report(result_path, output):
    from openpyxl import load_workbook
    # XLSX floating-point serial dates can round by 1 ms on independent readback.
    # Source times here have second precision; accept only this serialization noise.
    def same_time(value, observed):
        return isinstance(value, datetime) and abs((value - datetime.fromisoformat(observed).replace(tzinfo=None)).total_seconds()) <= .0011
    out = result_path.parent
    result, receipt = json.loads(result_path.read_text()), json.loads((out/'run_receipt.json').read_text())
    if sha256(result_path) != receipt['resultSha256'] or not result['verification']['passed']:
        raise ValueError('扫描结果哈希或核验状态不符')
    boundary = json.loads((out/'older_index_evidence.json').read_text()) if (out/'older_index_evidence.json').exists() else {'queries': []}
    older = {q['iccid']: q['earliest'][0]['observed'][:10] for q in boundary['queries']}
    wb = load_workbook(output, read_only=True, data_only=False, keep_links=False)
    try:
        if wb.sheetnames != ['逐卡首次出现', '原文档证据', '最早文档分组', '扫描范围与异常'] + (['窗口外更早证据'] if older else []):
            raise ValueError('报告页签不符')
        rows = list(wb['逐卡首次出现'].iter_rows(min_row=6, values_only=False))
        if len(rows) != result['inputCount']:
            raise ValueError('报告ICCID数量不符')
        for q, row in zip(result['queries'], rows):
            e = q['earliest'][0] if q['earliest'] else None
            expected = [q['input'], Path(e['path']).name if e else '未找到可定时文档',
                        e['sheet'] if e else '', e['cell'] if e else '',
                        len({h['path'] for h in q['earliest']}), q['matchedFiles'], q['unknownTimeFiles'],
                        '已有'+older[q['iccid']]+'更早证据，月内日期不是首次' if q['iccid'] in older else q['status']]
            actual = [row[i].value if row[i].value is not None else '' for i in (0,3,4,5,6,7,8,9)]
            if actual != expected or row[0].data_type != 's':
                raise ValueError('报告逐卡内容或文本精度不符')
            day = row[1].value.strftime('%Y-%m-%d') if isinstance(row[1].value, datetime) else None
            if day != q['earliestDate']:
                raise ValueError('最早日期回读不符')
            if e and e['precision'] == 'second' and not same_time(row[2].value, e['observed']):
                raise ValueError('逐卡文档时间回读不符')
        evidence = [h for q in result['queries'] for h in q['earliest']]
        evidence_rows = list(wb['原文档证据'].iter_rows(min_row=6, values_only=True))
        if len(evidence_rows) != len(evidence):
            raise ValueError('原文档证据行数不符')
        for h, row in zip(evidence, evidence_rows):
            if not same_time(row[1], h['observed']):
                raise ValueError('证据时间回读或日期格式不符')
            expected = [next(q['input'] for q in result['queries'] if q['iccid'] == h['iccid']),
                        '日级' if h['precision'] == 'day' else '秒级', h['basis'], Path(h['path']).name,
                        h['sheet'], h['cell'], h['path'], h['fileSha256'],
                        '通过' if h['independentReadbackPassed'] else '未通过']
            if [row[i] for i in (0,2,3,4,5,6,7,8,9)] != expected:
                raise ValueError('原文档证据内容不符')
        for row in wb['最早文档分组'].iter_rows(min_row=6, values_only=True):
            if not isinstance(row[0], datetime):
                raise ValueError('文档分组时间未规范为日期')
        if older:
            rows = list(wb['窗口外更早证据'].iter_rows(min_row=6, values_only=True))
            expected = [(q, h) for q in boundary['queries'] for h in q['earliest']]
            if len(rows) != len(expected):
                raise ValueError('窗口外较早证据数量或日期格式不符')
            for row, (q, h) in zip(rows, expected):
                if not same_time(row[1], h['observed']) or [row[i] for i in (0,2,3,4,5,6)] != [q['input'], Path(h['path']).name, h['sheet'], h['cell'], h['path'], h['fileSha256']]:
                    raise ValueError('窗口外较早证据内容或时间不符')
    finally:
        wb.close()
    check = dict(passed=True, identifiers=result['inputCount'], earlierEvidenceIdentifiers=len(older),
                 timestampPrecision='second', xlsxSerializationToleranceMs=1.1,
                 independentReader='openpyxl-read-only', outputSha256=sha256(output))
    dump(out/'independent_workbook_check.json', check)
    receipt.update(workbook=str(output), workbookSha256=check['outputSha256'], workbookReadback=check)
    dump(out/'run_receipt.json', receipt)
    print(json.dumps(dict(status='passed', output=str(output), identifiers=result['inputCount'], verification=check), ensure_ascii=False), flush=True)
    return check


def run_first_seen(args):
    started = time.monotonic()
    if not 1 <= args.workers <= 8:
        raise ValueError('工作进程必须1至8')
    I.validate_storage(args.db)
    keys, originals, invalid = normalize_list(args.list)
    now = datetime.now(TZ)
    since = datetime.fromisoformat(args.since).replace(tzinfo=TZ) if args.since else datetime.combine(previous_month(now.date()), datetime.min.time(), TZ)
    out = args.output_dir or args.db.parent / 'reports' / (now.strftime('%Y%m%dT%H%M%S_%f') + '_first_seen')
    if out.exists():
        raise ValueError('输出目录必须不存在，不覆盖已有结果')
    if not str(out.resolve()).startswith('/Volumes/zs01/'):
        raise ValueError('扫描输出只存zs01，不回退本地')
    out.mkdir(parents=True)
    receipt = dict(status='running', startedAt=now.isoformat(), input=str(args.list.resolve()), inputSha256=sha256(args.list),
                   outputDir=str(out), workersRequested=args.workers, networkAccess=False, browserAccess=False,
                   businessDatabaseAccess=False, taskFDatabaseAccess=False, originalsModified=False)
    dump(out / 'run_receipt.json', receipt)
    try:
        reference = index_reference(args.db, keys)
        files, discovery = I.discover(args.root or I.DEFAULT_ROOTS, [out])
        selected, outside, issues = [], [], []
        for item in files:
            hint = time_hint(item['path'])
            item = dict(item, **hint)
            if hint['observed'] and not since <= datetime.fromisoformat(hint['observed']) <= now:
                outside.append(item)
            else:
                selected.append(item)
        print(json.dumps(dict(event='first_seen_scan_started', outputDir=str(out), identifiers=len(keys), discovered=len(files),
            inScope=len(selected), outsideDateWindow=len(outside), workers=args.workers), ensure_ascii=False), flush=True)
        groups = defaultdict(list)
        with ProcessPoolExecutor(max_workers=args.workers, initializer=I.configure_worker_priority) as pool:
            for item in pool.map(fingerprint, selected, chunksize=8):
                if item['status'] != 'ready':
                    issues.append(item)
                else:
                    groups[item['signature']].append(item)
            reps = [items[0] for items in groups.values()]
            hits, scanned, pids, priorities = [], [], set(), {}
            for n, result in enumerate(pool.map(_scan_job, ((item, keys) for item in reps), chunksize=1), 1):
                pids.add(result['pid'])
                priorities[str(result['pid'])] = result['priority']
                copies = groups[result['signature']]
                for item in copies:
                    try:
                        same_stat(item)
                        if result['status'] != 'scanned':
                            raise ValueError(result['error'])
                        scanned.append(dict(path=item['path'], signature=item['signature'], observed=item['observed'],
                                            precision=item['precision'], basis=item['basis'], hitCells=len(result['hits'])))
                        hits.extend(dict(h, path=item['path'], observed=item['observed'], precision=item['precision'], basis=item['basis'],
                                         signature=item['signature'], mtime=item['mtime']) for h in result['hits'])
                    except Exception as exc:
                        issues.append(dict(item, status='unreadable', error=str(exc)))
                if n % 20 == 0 or n == len(reps):
                    print(json.dumps(dict(event='first_seen_scan_progress', completedContents=n, totalContents=len(reps),
                        matchedIdentifiers=len({h['iccid'] for h in hits}), elapsedSeconds=round(time.monotonic()-started, 1))), flush=True)
            by_key = defaultdict(list)
            for h in hits:
                by_key[h['iccid']].append(h)
            queries, verification_jobs = [], defaultdict(dict)
            file_map = {f['path']: f for f in selected}
            for key in keys:
                matches = by_key[key]
                earliest = earliest_candidates(matches)
                unknown = [h for h in matches if not h['observed']]
                for h in earliest:
                    verification_jobs[h['path']][(h['sheet'], h['cell'])] = h
                queries.append(dict(iccid=key, input=originals[key], matchedFiles=len({h['path'] for h in matches}),
                    matchedCells=len(matches), earliest=earliest, unknownTimeFiles=len({h['path'] for h in unknown}),
                    earliestDate=earliest[0]['observed'][:10] if earliest else None,
                    status='最早已定时命中；另有时间未知文档' if earliest and unknown else
                           '最早已定时命中' if earliest else '有命中但时间未知' if matches else '扫描范围内未命中'))
            verified = list(pool.map(verify_file, ((file_map[p], list(v.values())) for p, v in verification_jobs.items()), chunksize=1))
        verification_map = {v['path']: v['sha256'] for v in verified}
        for q in queries:
            for h in q['earliest']:
                h['fileSha256'] = verification_map[h['path']]
                h['independentReadbackPassed'] = True
        with (out / 'all_matches.jsonl').open('w') as stream:
            for h in hits:
                stream.write(json.dumps(h, ensure_ascii=False) + '\n')
        dump(out / 'scan_manifest.json', dict(discovery=discovery, scanned=scanned, excludedByDate=outside, unreadable=issues,
             contentGrouping='ZIP成员CRC/大小清单分组；最早命中文件另算字节SHA256并独立回读'))
        result = dict(schemaVersion=1, generatedAt=datetime.now(TZ).isoformat(), since=since.isoformat(), until=now.isoformat(),
            inputCount=len(keys), matchedCount=sum(q['matchedFiles'] > 0 for q in queries),
            datedCount=sum(bool(q['earliest']) for q in queries), queries=queries, invalidInputs=invalid,
            firstDateCounts=dict(sorted(Counter(q['earliestDate'] or '未知/未命中' for q in queries).items())),
            coverage=dict(discovered=len(files), scannedFiles=len(scanned), uniqueContents=len(groups),
                          unreadableFiles=len(issues), outsideDateWindow=len(outside), roots=discovery['roots'],
                          fileTypes=['xlsx', 'xlsm', 'xls(OOXML)', 'csv', 'tsv'],
                          excludedTypes='PDF/图片/Word/JSON/聊天纯文本不作业务表格出现证据；不访问业务数据库',
                          limitation='结论是最近一个月可读取表格中最早留存命中，不证明后台首次激活或完整生命周期；覆盖缺口见扫描清单'),
            indexReference=reference, verification=dict(passed=True, method='independent openpyxl read-only cells and file SHA256',
                        files=len(verified), cells=sum(v['verifiedCells'] for v in verified), evidence=verified),
            execution=dict(requestedWorkers=args.workers, actualWorkerPids=sorted(pids), workerPriorities=priorities,
                           durationSeconds=round(time.monotonic()-started, 3)))
        dump(out / 'result.json', result)
        receipt.update(status='passed_with_coverage_gaps' if issues else 'passed', completedAt=datetime.now(TZ).isoformat(),
                       result=str(out/'result.json'), resultSha256=sha256(out/'result.json'),
                       matchedCount=result['matchedCount'], datedCount=result['datedCount'], coverage=result['coverage'],
                       verification=result['verification'], execution=result['execution'])
        dump(out / 'run_receipt.json', receipt)
        print(json.dumps({k: v for k, v in receipt.items() if k != 'verification'}, ensure_ascii=False), flush=True)
        build_report(out/'result.json')
        return result
    except BaseException as exc:
        receipt.update(status='failed', error=str(exc), elapsedSeconds=round(time.monotonic()-started, 3))
        dump(out / 'run_receipt.json', receipt)
        raise
