#!/usr/bin/env python3
"""Generate receipt-gated Task E expiry-marker groups from the latest workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from opensms_data_anomalies import source_anomaly, anomaly_fields, assert_anomaly_identity_separation


ELIGIBLE_EXPIRY_STATUSES = ("6天", "快过期", "已过期")
NEVER_PROCESS_SOURCE_CARRIERS = ("快过期续费C", "快过期续费T")
REQUIRED_HEADERS = {"电话", "运营商", "过期状态", "截止日"}
LEDGER_SCHEMA_VERSION = 1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从本轮最终状态工作簿生成任务E到期标记分组")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--server-date", required=True)
    parser.add_argument("--tag", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--ledger", type=Path, required=True)
    parser.add_argument("--retention-receipt", type=Path, required=True)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def phone_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, int):
        raw = str(value)
    elif isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    else:
        raw = str(value).strip().removesuffix(".0")
    digits = re.sub(r"\D", "", raw)
    return digits[-10:] if len(digits) > 10 else digits


def parse_date(value: Any, epoch: datetime) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        try:
            converted = from_excel(float(text), epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            pass
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def load_ledger(path: Path) -> tuple[dict[str, Any], str | None]:
    if not path.exists():
        return {
            "schemaVersion": LEDGER_SCHEMA_VERSION,
            "dedupeKey": "phone",
            "submissions": {},
        }, None
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("schemaVersion") != LEDGER_SCHEMA_VERSION or payload.get("dedupeKey") != "phone":
        raise ValueError("任务E去重台账版本或去重键无效")
    submissions = payload.get("submissions")
    if not isinstance(submissions, dict):
        raise ValueError("任务E去重台账submissions必须为对象")
    normalized: dict[str, Any] = {}
    for key, value in submissions.items():
        phone = phone_text(key)
        if not phone or phone != phone_text((value or {}).get("phone")):
            raise ValueError(f"任务E去重台账号码键无效: {key!r}")
        if phone in normalized:
            raise ValueError(f"任务E去重台账存在重复号码: {phone}")
        normalized[phone] = value
    payload["submissions"] = normalized
    return payload, sha256(path)


def load_retention_receipt(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("task") not in {"D", "E"}:
        raise ValueError("任务E卡保留回执任务类型无效")
    if payload.get("sourceProvenance") != "phone_card_retain_all_pages":
        raise ValueError("任务E卡保留回执来源无效")
    if payload.get("endpoint") != "/phone/cardRetain/list" or int(payload.get("pageSize") or 0) != 200:
        raise ValueError("任务E卡保留回执接口或分页大小无效")
    total = int(payload.get("totalRecords") or 0)
    fetched = int(payload.get("fetchedRecords") or 0)
    records = payload.get("records")
    if total < 0 or fetched < total or not isinstance(records, list) or len(records) != fetched:
        raise ValueError("任务E卡保留回执分页或记录数量不完整")
    if int(payload.get("invalidCardRecordCount") or 0) != 0:
        raise ValueError("任务E卡保留回执存在无效卡号")
    page_count = int(payload.get("pageCount") or 0)
    pages = payload.get("pages") or []
    if page_count < 1 or len(pages) != page_count or sum(int(item.get("rowCount") or 0) for item in pages) != fetched:
        raise ValueError("任务E卡保留回执页明细不完整")
    record_phones = {phone_text(record.get("cardNo")) for record in records}
    record_phones.discard("")
    receipt_phones = {phone_text(value) for value in payload.get("cardNos") or []}
    receipt_phones.discard("")
    if record_phones != receipt_phones or len(record_phones) != int(payload.get("uniqueCardCount") or 0):
        raise ValueError("任务E卡保留回执去重卡号不一致")
    return payload


def write_exclusive(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def main() -> None:
    args = parse_args()
    if os.environ.get("OPENSMS_ENABLE_TASK_E") != "1":
        raise RuntimeError("OPENSMS_ENABLE_TASK_E 未严格等于 1，拒绝生成任务E候选")
    if os.environ.get("OPENSMS_FORCE_TASK_E") != "1":
        raise RuntimeError("OPENSMS_FORCE_TASK_E 未严格等于 1，缺少当前真实用户授权，拒绝生成任务E候选")
    server_date = date.fromisoformat(args.server_date)
    if not args.workbook.is_file():
        raise FileNotFoundError(args.workbook)

    retention = load_retention_receipt(args.retention_receipt)
    ledger, ledger_sha = load_ledger(args.ledger)
    submitted = set(ledger["submissions"])
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        if "sheet1" not in workbook.sheetnames:
            raise ValueError(f"任务E结果主表sheet1不存在: {workbook.sheetnames}")
        sheet = workbook["sheet1"]
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        ]
        missing = sorted(REQUIRED_HEADERS - set(headers))
        if missing:
            raise ValueError(f"任务E结果主表缺少字段: {missing}")

        eligible_status_counts = {status: 0 for status in ELIGIBLE_EXPIRY_STATUSES}
        protected_rows: list[dict[str, Any]] = []
        unconfigured_carrier_rows: list[dict[str, Any]] = []
        already_submitted_rows: list[dict[str, Any]] = []
        duplicate_rows: list[dict[str, Any]] = []
        candidates_by_phone: dict[str, dict[str, Any]] = {}
        invalid_rows: list[dict[str, Any]] = []
        data_anomalies = []
        workbook_phone_rows: dict[str, list[int]] = {}
        workbook_deadlines: dict[str, set[str]] = {}

        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            values = {
                headers[index]: cell.value
                for index, cell in enumerate(cells)
                if index < len(headers) and headers[index]
            }
            phone = phone_text(values.get("电话"))
            deadline = parse_date(values.get("截止日"), workbook.epoch)
            if phone:
                workbook_phone_rows.setdefault(phone, []).append(row_number)
                if deadline:
                    workbook_deadlines.setdefault(phone, set()).add(deadline.isoformat())

            status = str(values.get("过期状态") or "").strip()
            if status not in ELIGIBLE_EXPIRY_STATUSES:
                continue
            eligible_status_counts[status] += 1
            carrier = str(values.get("运营商") or "").strip()
            base = {
                "电话": phone,
                "运营商": carrier,
                "过期状态": status,
                "截止日": deadline.isoformat() if deadline else "",
                "源行号": row_number,
            }
            if not carrier:
                unconfigured_carrier_rows.append(base)
                continue
            if not phone or deadline is None:
                invalid_rows.append(base)
                data_anomalies.append(source_anomaly(row_number, values,
                    "phone_missing" if not phone else "deadline_missing_or_invalid", phone))
                continue
            if carrier in NEVER_PROCESS_SOURCE_CARRIERS:
                protected_rows.append(base)
                continue
            if phone in submitted:
                already_submitted_rows.append(base)
                continue
            prior = candidates_by_phone.get(phone)
            if prior:
                if prior["截止日"] != base["截止日"]:
                    raise ValueError(
                        f"任务E同一号码出现不同截止日: {phone}, {prior['截止日']} / {base['截止日']}"
                    )
                duplicate_rows.append(base)
                continue
            candidates_by_phone[phone] = base
    finally:
        workbook.close()

    assert_anomaly_identity_separation(data_anomalies, list(candidates_by_phone.values()))

    retention_records = retention.get("records") or []
    retention_missing_phone: set[str] = set()
    retention_missing_deadline: set[str] = set()
    retention_invalid_notification: list[dict[str, Any]] = []
    retention_conflicts: list[dict[str, Any]] = []
    retention_matches_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    retention_known_phones: set[str] = set()
    for record in retention_records:
        phone = phone_text(record.get("cardNo"))
        notification_date_text = str(record.get("expireAfter") or "").strip()
        if not phone:
            raise ValueError("任务E卡保留回执出现空卡号")
        if phone not in workbook_phone_rows:
            retention_missing_phone.add(phone)
            continue
        deadlines = workbook_deadlines.get(phone) or set()
        if not deadlines:
            retention_missing_deadline.add(phone)
            continue
        if len(deadlines) != 1:
            retention_conflicts.append({
                "电话": phone,
                "截止日": sorted(deadlines),
                "源行号": workbook_phone_rows.get(phone) or [],
            })
            continue
        retention_known_phones.add(phone)
        try:
            notification_date = date.fromisoformat(notification_date_text)
        except ValueError:
            retention_invalid_notification.append({
                "电话": phone,
                "下卡通知时间": notification_date_text,
                "截止日": next(iter(deadlines)),
                "申请时间": str(record.get("createTime") or ""),
            })
            continue
        deadline_text = next(iter(deadlines))
        if notification_date.isoformat() != deadline_text:
            continue
        key = (phone, deadline_text)
        retention_matches_by_key.setdefault(key, {
            "电话": phone,
            "下卡通知时间": notification_date.isoformat(),
            "截止日": deadline_text,
            "原到期日": str(record.get("expireBefore") or ""),
            "新到期日": notification_date.isoformat(),
            "申请时间": str(record.get("createTime") or ""),
            "申请会话ID": str(record.get("chatId") or ""),
            "文档源行号": workbook_phone_rows.get(phone) or [],
        })
    if retention_conflicts:
        raise ValueError(f"任务E卡保留号码在文档中存在不同截止日: {retention_conflicts[:20]}")
    retention_matches = sorted(
        retention_matches_by_key.values(),
        key=lambda item: (item["截止日"], item["电话"]),
    )

    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in candidates_by_phone.values():
        grouped.setdefault(row["截止日"], []).append(row)
    groups = []
    for expire_date in sorted(grouped):
        rows = sorted(grouped[expire_date], key=lambda item: (item["电话"], item["源行号"]))
        groups.append({
            "expireDate": expire_date,
            "count": len(rows),
            "cardNos": [row["电话"] for row in rows],
            "rows": rows,
        })

    payload = {
        "createdAt": datetime.now().astimezone().isoformat(),
        "task": "E",
        **anomaly_fields(data_anomalies, args.output),
        "forceGate": True,
        "automationGate": True,
        "unconfiguredCarrierPolicy": "skip_no_business_write_v1",
        "unconfiguredCarrierSkippedCount": len(unconfigured_carrier_rows),
        "unconfiguredCarrierSkippedRows": unconfigured_carrier_rows,
        "sourcePath": str(args.workbook.resolve()),
        "sourceSha256": sha256(args.workbook),
        "sourceSheet": "sheet1",
        "serverDate": server_date.isoformat(),
        "eligibleExpiryStatuses": list(ELIGIBLE_EXPIRY_STATUSES),
        "neverProcessSourceCarrierNames": list(NEVER_PROCESS_SOURCE_CARRIERS),
        "ledgerPath": str(args.ledger.resolve()),
        "ledgerSha256Before": ledger_sha,
        "ledgerDedupeKey": "phone",
        "retentionReceiptPath": str(args.retention_receipt.resolve()),
        "retentionReceiptSha256": sha256(args.retention_receipt),
        "retentionSourceProvenance": retention.get("sourceProvenance"),
        "retentionEndpoint": retention.get("endpoint"),
        "retentionPageSize": int(retention.get("pageSize") or 0),
        "retentionPageCount": int(retention.get("pageCount") or 0),
        "retentionTotalRecordCount": int(retention.get("totalRecords") or 0),
        "retentionFetchedRecordCount": int(retention.get("fetchedRecords") or 0),
        "retentionUniqueCardCount": int(retention.get("uniqueCardCount") or 0),
        "retentionWorkbookPhoneFoundCount": len(set(workbook_phone_rows) & set(retention.get("cardNos") or [])),
        "retentionWorkbookDeadlineKnownCount": len(retention_known_phones),
        "retentionWorkbookMissingPhoneCount": len(retention_missing_phone),
        "retentionWorkbookMissingPhones": sorted(retention_missing_phone),
        "retentionWorkbookMissingDeadlineCount": len(retention_missing_deadline),
        "retentionWorkbookMissingDeadlinePhones": sorted(retention_missing_deadline),
        "retentionInvalidNotificationDateRecordCount": len(retention_invalid_notification),
        "retentionInvalidNotificationDateRecords": retention_invalid_notification,
        "retentionDeadlineConflictCount": 0,
        "retentionDeadlineMatches": retention_matches,
        "retentionDeadlineMatchPhoneCount": len({item["电话"] for item in retention_matches}),
        "retentionDeadlineMatchRecordCount": len(retention_matches),
        "retentionComparisonDoesNotChangeCandidates": True,
        "eligibleStatusCounts": eligible_status_counts,
        "eligibleWorkbookRowCount": sum(eligible_status_counts.values()),
        "neverProcessSourceExcludedCount": len(protected_rows),
        "neverProcessSourceExcludedRows": protected_rows,
        "alreadySubmittedExcludedCount": len(already_submitted_rows),
        "alreadySubmittedExcludedRows": already_submitted_rows,
        "duplicateSameDeadlineCollapsedCount": len(duplicate_rows),
        "duplicateSameDeadlineCollapsedRows": duplicate_rows,
        "dateGroupCount": len(groups),
        "candidateCount": len(candidates_by_phone),
        "count": len(candidates_by_phone),
        "groups": groups,
    }
    write_exclusive(args.output, payload)
    print(json.dumps({**payload, "path": str(args.output.resolve())}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
