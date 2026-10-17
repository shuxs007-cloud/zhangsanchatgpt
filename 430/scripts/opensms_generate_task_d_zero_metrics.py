#!/usr/bin/env python3
"""Build Task D two-day-zero-reply metric candidates from this run's workbook.

This module is deliberately read-only.  Live carrier membership is checked by
the protocol preparation step after all swap-zone carrier work has finished.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DAILY_TOTAL_REPLY = re.compile(r"^(\d{2})-(\d{2})总回复量$")
NEVER_PROCESS_SOURCE_CARRIERS = ("快过期续费C", "快过期续费T")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="生成任务D连续两日总回复量为0的指标候选")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--server-date", required=True)
    parser.add_argument("--tag", required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def phone_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", "", str(value).strip()).removesuffix(".0")


def numeric_zero(value: Any) -> bool:
    if value in (None, "") or isinstance(value, bool):
        return False
    try:
        return float(str(value).strip().replace(",", "")) == 0
    except ValueError:
        return False


def resolve_daily_headers(headers: list[str], server_date: date) -> tuple[str, str, date, date]:
    candidates: list[tuple[date, str]] = []
    for header in headers:
        match = DAILY_TOTAL_REPLY.fullmatch(header)
        if not match:
            continue
        month, day = int(match.group(1)), int(match.group(2))
        years = (server_date.year, server_date.year - 1, server_date.year + 1)
        parsed = None
        for year in years:
            try:
                item = date(year, month, day)
            except ValueError:
                continue
            if item <= server_date and (server_date - item).days <= 370:
                parsed = item
                break
        if parsed is not None:
            candidates.append((parsed, header))
    candidates.sort(reverse=True)
    if len(candidates) < 2:
        raise ValueError("任务D找不到至少两个有效的MM-DD总回复量字段")
    latest_date, latest_header = candidates[0]
    previous_date, previous_header = candidates[1]
    if (latest_date - previous_date).days != 1:
        raise ValueError(
            f"任务D最新两个总回复量日期不连续: {latest_date} / {previous_date}"
        )
    if latest_date != server_date:
        raise ValueError(
            f"任务D最新总回复量日期不是后台服务器日期: latest={latest_date}, server={server_date}"
        )
    return latest_header, previous_header, latest_date, previous_date


def write_exclusive(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def main() -> None:
    args = parse_args()
    if os.environ.get("OPENSMS_ENABLE_TASK_D") != "1":
        raise RuntimeError("OPENSMS_ENABLE_TASK_D 未严格等于 1，拒绝生成任务D候选")
    if os.environ.get("OPENSMS_FORCE_TASK_D") != "1":
        raise RuntimeError("OPENSMS_FORCE_TASK_D 未严格等于 1，缺少当前真实用户授权，拒绝生成任务D候选")
    if not args.workbook.is_file():
        raise FileNotFoundError(args.workbook)
    server_date = date.fromisoformat(args.server_date)

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        if "sheet1" not in workbook.sheetnames:
            raise ValueError(f"任务D结果主表sheet1不存在: {workbook.sheetnames}")
        sheet = workbook["sheet1"]
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        ]
        for required in ("电话", "运营商", "opr", "卡槽"):
            if required not in headers:
                raise ValueError(f"任务D结果主表缺少字段: {required}")
        latest_header, previous_header, latest_date, previous_date = resolve_daily_headers(
            headers, server_date
        )
        rows: list[dict[str, Any]] = []
        never_process_rows: dict[str, list[int]] = {
            carrier: [] for carrier in NEVER_PROCESS_SOURCE_CARRIERS
        }
        seen: set[str] = set()
        unconfigured_carrier_rows = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            values = {
                headers[index]: cell.value
                for index, cell in enumerate(cells)
                if index < len(headers) and headers[index]
            }
            current_carrier = str(values.get("运营商") or "").strip()
            if not current_carrier:
                unconfigured_carrier_rows.append(row_number)
                continue
            if current_carrier in NEVER_PROCESS_SOURCE_CARRIERS:
                never_process_rows[current_carrier].append(row_number)
                continue
            if not numeric_zero(values.get(latest_header)) or not numeric_zero(
                values.get(previous_header)
            ):
                continue
            phone = phone_text(values.get("电话"))
            if not phone:
                raise ValueError(f"任务D连续两日0回复命中行缺少电话: row={row_number}")
            if phone in seen:
                raise ValueError(f"任务D指标候选存在重复号码: {phone}")
            seen.add(phone)
            rows.append(
                {
                    "电话": phone,
                    "运营商": str(values.get("运营商") or "").strip(),
                    "opr": str(values.get("opr") or "").strip(),
                    "卡槽": str(values.get("卡槽") or "").strip(),
                    latest_header: 0,
                    previous_header: 0,
                    "源行号": row_number,
                }
            )
    finally:
        workbook.close()

    payload = {
        "createdAt": datetime.now().astimezone().isoformat(),
        "task": "D",
        "forceGate": True,
        "automationGate": True,
        "phase": "swap_zero_metric",
        "unconfiguredCarrierPolicy": "skip_no_business_write_v1",
        "unconfiguredCarrierSkippedCount": len(unconfigured_carrier_rows),
        "unconfiguredCarrierSkippedRows": unconfigured_carrier_rows,
        "sourcePath": str(args.workbook.resolve()),
        "sourceSha256": sha256(args.workbook),
        "sourceSheet": "sheet1",
        "serverDate": server_date.isoformat(),
        "latestReplyHeader": latest_header,
        "previousReplyHeader": previous_header,
        "latestReplyDate": latest_date.isoformat(),
        "previousReplyDate": previous_date.isoformat(),
        "comparison": "both_numeric_equal_zero",
        "neverProcessSourceCarrierNames": list(NEVER_PROCESS_SOURCE_CARRIERS),
        "neverProcessSourceExcludedCount": sum(len(rows) for rows in never_process_rows.values()),
        "neverProcessSourceExcludedCounts": {
            carrier: len(never_process_rows[carrier])
            for carrier in NEVER_PROCESS_SOURCE_CARRIERS
        },
        "neverProcessSourceExcludedRows": never_process_rows,
        "metricCandidateCount": len(rows),
        "count": len(rows),
        "rows": rows,
        "candidates": rows,
    }
    write_exclusive(args.output, payload)
    print(json.dumps({**payload, "path": str(args.output.resolve())}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
