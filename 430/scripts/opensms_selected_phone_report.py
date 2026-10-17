#!/usr/bin/env python3
"""Fresh, read-only Task A report for an explicit phone list.

Reuse the existing acquisition/date/document verification chain. Never call the
patrol execute/candidate/write/F paths. Only selected main-sheet rows are delivered.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

from opensms_patrol_runner import (
    PatrolRunner, PatrolError, ROOT, SCRIPTS, NODE, NODE_MODULES, PYTHON,
    build_parser, sha256, atomic_json, now_hkt,
)


def phone_key(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    text = str(value).strip()
    if re.fullmatch(r"[0-9]{10}", text):
        return text
    if re.fullmatch(r"\+?1[0-9]{10}", text):
        return text[-10:]
    return None


def read_phones(path):
    lines = [line.strip() for line in Path(path).read_text().splitlines() if line.strip()]
    if not lines or any(not re.fullmatch(r"[0-9]{10}", line) for line in lines):
        raise ValueError("名单必须为每行一个10位卡号；空名单或非法条目拒绝导出")
    return list(dict.fromkeys(lines)), len(lines)


def rgb(color):
    if color is None:
        return None
    if color.type == "rgb":
        return "#" + str(color.rgb)[-6:]
    if color.type == "indexed" and color.indexed < len(COLOR_INDEX):
        return "#" + COLOR_INDEX[color.indexed][-6:]
    return None


def style_config(cell):
    font = cell.font
    result = {"numberFormat": cell.number_format,
              "font": {"name": font.name or "Arial", "size": font.sz or 11,
                       "bold": bool(font.b), "italic": bool(font.i)}}
    if rgb(font.color):
        result["font"]["color"] = rgb(font.color)
    if cell.fill.patternType == "solid" and rgb(cell.fill.fgColor):
        result["fill"] = rgb(cell.fill.fgColor)
    align = cell.alignment
    if align.horizontal in ("left", "center", "right", "general", "justify"):
        result["horizontalAlignment"] = align.horizontal
    if align.vertical in ("top", "center", "bottom"):
        result["verticalAlignment"] = align.vertical
    result["wrapText"] = bool(align.wrap_text)
    borders = {}
    for side in ("top", "bottom", "left", "right"):
        border = getattr(cell.border, side)
        if border and border.style:
            borders[side] = {"style": border.style, "color": rgb(border.color) or "#000000"}
    if borders:
        result["borders"] = borders
    return result


def pack_value(cell):
    if cell.data_type in ("f", "e"):
        raise ValueError("指定号码行含公式或错误值，不能生成可能失效的分表")
    value = cell.value
    if isinstance(value, (date, datetime)):
        return {"date": value.isoformat()}
    return value


def extract_main(path, phones):
    """Read-only extraction, exact NANP key; preserve all matching source rows."""
    book = load_workbook(path, data_only=False)
    try:
        sheet = book["sheet1"] if "sheet1" in book.sheetnames else book.worksheets[0]
        headers = [cell.value for cell in sheet[1]]
        indexes = [i for i, v in enumerate(headers) if v in ("电话", "卡号", "手机号")]
        if len(indexes) != 1:
            raise ValueError("来源缺少唯一电话列")
        index = indexes[0]
        grouped = {p: [] for p in phones}
        for row in sheet.iter_rows(min_row=2):
            key = phone_key(row[index].value)
            if key in grouped:
                grouped[key].append(row)
        rows = [tuple(sheet[1])] + [row for p in phones for row in grouped[p]]
        style_ids = {cell.style_id for row in rows for cell in row}
        styles = {}
        for row in rows:
            for cell in row:
                if cell.style_id not in styles:
                    styles[cell.style_id] = style_config(cell)
        return {
            "source": str(path), "sourceSha256": sha256(Path(path)), "sheet": sheet.title,
            "values": [[pack_value(cell) for cell in row] for row in rows],
            "styleIds": [[cell.style_id for cell in row] for row in rows], "styles": styles,
            "widths": [sheet.column_dimensions[get_column_letter(i)].width for i in range(1, len(headers)+1)],
            "counts": {p: len(grouped[p]) for p in phones}, "phoneColumn": index,
            "sourceRows": [row[0].row for row in rows], "columnCount": len(headers),
        }
    finally:
        book.close()


def value_signature(cell):
    value = None if cell.value == "" else cell.value
    if value is not None and cell.data_type == "n" and not isinstance(value, bool):
        # Excel numeric value semantics, not Python int/float binary coercion.
        # Exact decimal comparison: no tolerance and no long-number truncation.
        value = Decimal(str(value))
    return value, ("blank" if value is None else cell.data_type)


def verify_selected(source, output, phones, expected_sha, counts, with_match_sheet):
    if sha256(Path(source)) != expected_sha:
        raise ValueError("来源SHA256发生变化")
    with ZipFile(output) as zipped:
        if zipped.testzip() is not None:
            raise ValueError("交付XLSX压缩结构错误")
    original = load_workbook(source, data_only=False)
    final = load_workbook(output, data_only=False)
    try:
        expected_sheets = ["sheet1", "号码匹配核对"] if with_match_sheet else ["sheet1"]
        if final.sheetnames != expected_sheets:
            raise ValueError("输出存在范围外工作表")
        inp = original["sheet1"] if "sheet1" in original.sheetnames else original.worksheets[0]
        headers = [c.value for c in inp[1]]
        index = next(i for i, h in enumerate(headers) if h in ("电话", "卡号", "手机号"))
        grouped = {p: [] for p in phones}
        for row in inp.iter_rows(min_row=2):
            key = phone_key(row[index].value)
            if key in grouped:
                grouped[key].append(row)
        expected = [tuple(inp[1])] + [row for p in phones for row in grouped[p]]
        out = final["sheet1"]
        if out.max_row != len(expected) or out.max_column != len(headers):
            raise ValueError("交付行列数不一致")
        for row_no, (old, new) in enumerate(zip(expected, out.iter_rows()), 1):
            for a, b in zip(old, new):
                if value_signature(a) != value_signature(b) or a.number_format != b.number_format:
                    raise ValueError(f"逐单元格值、类型或数字格式不一致：行{row_no}，列{a.column}；"
                                     f"原类型={a.data_type}，新类型={b.data_type}，"
                                     f"原格式={a.number_format!r}，新格式={b.number_format!r}；"
                                     f"原值={a.value!r}，新值={b.value!r}")
                if a.fill.patternType == "solid" and rgb(a.fill.fgColor) != rgb(b.fill.fgColor):
                    raise ValueError(f"日期标色不一致：行{row_no}，列{a.column}")
        actual_counts = Counter(phone_key(out.cell(r, index+1).value) for r in range(2, out.max_row+1))
        if dict(actual_counts) != {p: n for p, n in counts.items() if n}:
            raise ValueError("存在范围外号码或遗漏匹配行")
        if with_match_sheet:
            rows = list(final["号码匹配核对"].values)
            expected_rows = [("卡号", "匹配行数", "匹配结果")] + [
                (p, counts[p], "已匹配" if counts[p] else "本轮源表未找到") for p in phones]
            if rows != expected_rows:
                raise ValueError("号码匹配核对清单不一致")
        if sha256(Path(source)) != expected_sha:
            raise ValueError("独立复核期间来源变化")
        return {"independentReadbackPassed": True, "rowCount": len(expected)-1,
                "matchedPhones": len(actual_counts), "outOfScopeRows": 0,
                "sourceUnchanged": True, "sha256": sha256(Path(output))}
    finally:
        original.close()
        final.close()


class SelectedPhoneRunner(PatrolRunner):
    def execute(self):
        raise PatrolError("指定号码只读入口禁止执行全量业务流程")

    def run_write(self, *args, **kwargs):
        raise PatrolError("指定号码只读入口禁止运营商写入")

    def run_standard_writes(self, *args, **kwargs):
        raise PatrolError("指定号码只读入口禁止普通业务写入")

    def start_background_task_f(self, *args, **kwargs):
        raise PatrolError("指定号码任务A不启动F")


def run_report(phone_file):
    phones, input_count = read_phones(phone_file)
    # Same storage guard as the standard launcher, before CDP construction.
    download_root = subprocess.check_output([str(PYTHON), str(SCRIPTS / "opensms_storage_policy.py"), "download-root"], text=True).strip()
    os.environ["OPENSMS_EXPORT_DIR"] = download_root
    args = build_parser().parse_args(["run", "--task-a-only", "--json"])
    runner = SelectedPhoneRunner(args)
    print(json.dumps({"event": "patrol_started", "runId": runner.run_id,
                      "state": str(runner.state_path), "runnerPid": os.getpid(),
                      "taskSelection": "A-selected-read-only", "requestedPhones": len(phones)}, ensure_ascii=False), file=sys.stderr, flush=True)
    runner.state.update({"selectedPhoneReport": True, "businessWritesEnabled": False,
                         "requestedPhoneCount": len(phones), "inputPhoneCount": input_count})
    runner._save_state()
    try:
        frozen = runner.receipts / "requested_phones.txt"
        with frozen.open("x") as stream:
            stream.write("\n".join(phones) + "\n")
        runner.verify_policy()
        acquire = runner.acquire()
        runner.start_document(acquire)
        document = runner.finish_document()
        runner.close_api_target()
        source = extract_main(acquire["processingSource"], phones)
        processed = extract_main(document["cardResult"], phones)
        if source["counts"] != processed["counts"]:
            raise ValueError("新鲜源与任务A处理表的指定号码集合不一致")
        paths = {
            "sourceWorkbook": str(runner.delivery / f"OpenSMS指定号码_新鲜源数据_{runner.run_id}.xlsx"),
            "finalWorkbook": str(runner.delivery / f"OpenSMS指定号码_任务A报表_{runner.run_id}.xlsx"),
        }
        spec = runner.receipts / "selected_phone_spec.json"
        atomic_json(spec, {"phones": phones, "source": source, "processed": processed, "outputs": paths})
        work = ROOT / "outputs" / f"opensms_patrol_{runner.run_id}" / "selected"
        work.mkdir(parents=True, exist_ok=True)
        module_link = work / "node_modules"
        module_link.symlink_to(NODE_MODULES, target_is_directory=True)
        builder = work / "build.mjs"
        shutil.copy2(SCRIPTS / "opensms_selected_phone_report.mjs", builder)
        runner.command("selected_phone_workbooks", [NODE, builder, spec])
        proofs = {
            "source": verify_selected(source["source"], paths["sourceWorkbook"], phones, source["sourceSha256"], source["counts"], False),
            "processed": verify_selected(processed["source"], paths["finalWorkbook"], phones, processed["sourceSha256"], processed["counts"], True),
        }
        missing = [p for p in phones if not source["counts"][p]]
        summary = {"status": "success", "runId": runner.run_id, "serverDate": acquire["serverDate"],
                   "taskSelection": "A-selected-read-only", "requestedCount": len(phones),
                   "inputCount": input_count, "matchedCount": len(phones)-len(missing),
                   "missingCount": len(missing), "matchedRows": proofs["processed"]["rowCount"],
                   "actualWriteCount": 0, "tasksBCDEFExecuted": False,
                   "independentReadbackPassed": True, "outOfScopeRows": 0,
                   **paths, "state": str(runner.state_path)}
        receipt = runner.receipts / "selected_phone_verification.json"
        atomic_json(receipt, {**summary, "proofs": proofs, "phoneListSha256": sha256(frozen), "missingPhones": missing})
        report = runner.delivery / f"OpenSMS指定号码_任务A报告_{runner.run_id}.md"
        summary["report"] = str(report)
        report.write_text("\n".join([
            "## 本轮概况", "", f"运行编号 `{runner.run_id}`；服务器日期 {acquire['serverDate']}。",
            f"指定 {len(phones)} 个号码，匹配 {summary['matchedCount']} 个／{summary['matchedRows']} 行，未找到 {len(missing)} 个。",
            "本轮重新下载源数据，复用任务A原日期补全、主表生成及独立核验。仅交付指定号码，不复用旧导出。",
            "后台接口生成全量源作为本轮只读证据；交付源表与处理表已按名单提取。未匹配不等于后台确定不存在。",
            "", "## 0回复与回复率", "", "指定号码的近10日逐日发送、回复及回复率保留在任务A主表；本轮不执行回复率路由。",
            "", "## 运营商调整", "", "只读导出，运营商写入0；任务B/C/D/E/F全部未执行。",
            "", "## 下架恢复", "", "永久关闭，未执行。",
            "", "## 余量与提醒", "", "本次为指定号码报表，不将运营商全组余量当作这些号码的余量。未匹配号码见号码匹配核对sheet。",
            "", "## 文件与回执", "", f"[指定号码源数据](<{paths['sourceWorkbook']}>)", f"[指定号码任务A报表](<{paths['finalWorkbook']}>)",
            f"[独立核验回执](<{receipt}>)", "", "## 异常", "", f"范围外行数0；未匹配号码{len(missing)}个。逐单元格值、类型、数字格式及日期标色独立回读通过。",
        ]) + "\n", encoding="utf-8")
        runner.state.update({"status": "success", "summary": summary})
        runner._save_state()
        atomic_json(runner.delivery / "selected_phone_summary.json", summary)
        return summary
    except Exception as exc:
        try:
            runner.close_api_target()
        except Exception as cleanup:
            runner.state["cleanupError"] = str(cleanup)
        runner.state.update({"status": "failed", "error": str(exc)})
        runner._save_state()
        return {"status": "failed", "runId": runner.run_id, "state": str(runner.state_path),
                "error": str(exc), "actualWriteCount": 0, "tasksBCDEFExecuted": False}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="指定号码任务A只读新鲜报表，不执行业务调整")
    parser.add_argument("--phones-file", type=Path, required=True)
    parser.add_argument("--json", action="store_true")
    opts = parser.parse_args()
    result = run_report(opts.phones_file)
    print(json.dumps(result, ensure_ascii=False))
    raise SystemExit(0 if result["status"] == "success" else 2)
