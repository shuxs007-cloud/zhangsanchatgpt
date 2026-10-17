#!/usr/bin/env python3
"""Offline device totals from a verified final workbook; no network or mutations."""
import argparse
from collections import defaultdict
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from services.opensms_document_api.pipeline.device_daily_statistics import (
    count_value, device_key, device_sort_key,
)

FIELDS = ('10日发送窗口数', '10日回复窗口数', '10日总发送量', '10日总回复量')
HEADERS = ('设备号', '机房', '号码记录数', '10日发送窗口数', '10日回复窗口数',
           '整体回复率', '10日总发送量', '10日总回复量', '数据状态')


def aggregate(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = iter(book['sheet1'].values)
        headers = list(next(rows))
        for field in (*FIELDS, '卡槽', '机房'):
            if headers.count(field) != 1:
                raise ValueError(f'必须存在唯一列：{field}')
        indexes = [headers.index(field) for field in FIELDS]
        slots, rooms = headers.index('卡槽'), headers.index('机房')
        groups = defaultdict(lambda: {'count': 0, 'rooms': set(), 'sums': [0]*4, 'bad': [0]*4})
        for row in rows:
            if not any(value is not None for value in row):
                continue
            group = groups[device_key(row[slots])]
            group['count'] += 1
            group['rooms'].add(str(row[rooms] or '未标明机房'))
            for i, index in enumerate(indexes):
                number = count_value(row[index])
                if number is None:
                    group['bad'][i] += 1
                else:
                    group['sums'][i] += number
        output = []
        for key in sorted(groups, key=device_sort_key):
            g = groups[key]
            sends, replies, sent, received = [None if g['bad'][i] else v for i, v in enumerate(g['sums'])]
            rate = replies/sends if sends and replies is not None else None
            status = '源计数缺失，相关合计留空' if any(g['bad']) else '无发送，回复率不计算' if sends == 0 else '完整'
            output.append([key, '、'.join(sorted(g['rooms'])), g['count'], sends, replies, rate, sent, received, status])
        return output
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    source, target = args.input.resolve(), args.output.resolve()
    if target == source or target.exists() or target.with_suffix('.md').exists():
        raise ValueError('禁止覆盖源文件或已有交付')
    result = aggregate(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = '设备整体回复率'
    sheet.append(HEADERS)
    for row in result:
        sheet.append(row)
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='216B78')
    for letter, width in zip('ABCDEFGHI', (14, 24, 16, 22, 22, 20, 22, 22, 34)):
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = '@'
        row[5].number_format = '0.000%'
        for cell in row:
            cell.alignment = Alignment(vertical='center')
    notes = book.create_sheet('统计口径')
    for line in (
        '来源：' + str(source),
        '直接汇总主表10日字段，不能用可见的9个日明细列替代10日总计。',
        '整体回复率=设备10日回复窗口数合计/设备10日发送窗口数合计；不是每日百分比的平均。',
        '总发送量和总回复量是另外两个指标，不作为上述窗口回复率分母、分子。',
        '包含主表全部记录与全部状态，不按手机号去重；设备号保留字母前缀。',
        '计数缺失时相关合计留空；零分母的回复率留空。',
        '这是已核验报表的离线统计，不是新鲜导出；不修改任何线上号码。',
    ):
        notes.append([line])
    notes.column_dimensions['A'].width = 120
    book.save(target)
    book.close()
    # Separate disk reads of both source and result; verify every delivered cell.
    expected = aggregate(source)
    check = load_workbook(target, read_only=True, data_only=True)
    actual = list(check['设备整体回复率'].values)
    check.close()
    assert tuple(actual[0]) == HEADERS and len(actual)-1 == len(expected)
    for saved, wanted in zip(actual[1:], expected):
        for a, b in zip(saved, wanted):
            assert abs(a-b) < 1e-12 if isinstance(b, float) else a == b
    lines = ['# 每设备整体回复率', '', f'来源：[{source.name}]({source})', '',
             '口径：主表10日回复窗口数合计 ÷ 10日发送窗口数合计。含全部状态；不是每日百分比平均。', '',
             '| ' + ' | '.join(HEADERS) + ' |', '| ' + ' | '.join(['---']*len(HEADERS)) + ' |']
    for row in result:
        cells = ['未知' if x is None else f'{x:.3%}' if i == 5 else str(x) for i, x in enumerate(row)]
        lines.append('| ' + ' | '.join(cells) + ' |')
    target.with_suffix('.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')
    totals = {name: sum(row[col] for row in result) if all(row[col] is not None for row in result) else None
              for name, col in [('发送窗口数',3),('回复窗口数',4)]}
    print(json.dumps({'devices':len(result), 'records':sum(row[2] for row in result),
                      'totals':totals, 'verification':'passed', 'xlsx':str(target),
                      'md':str(target.with_suffix('.md'))}, ensure_ascii=False))


if __name__ == '__main__':
    main()
