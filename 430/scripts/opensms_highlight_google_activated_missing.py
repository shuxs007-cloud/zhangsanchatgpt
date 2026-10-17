#!/usr/bin/env python3
"""Highlight activated Google-only records in the all-Google-missing sheet."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


ALL_SHEET = "谷歌有-系统无"
LEGACY_ALL_SHEET = "所有谷歌有-系统没有"
ACTIVATED_SHEET = "谷歌文档中有激活日期-系统无法使用"
KEY_HEADERS = ("Google页签", "Google行号", "归一化卡槽")
YELLOW = "FFFF00"
FORMULA_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_slot(value) -> str:
    raw = canonical_text(value).replace(" ", "").replace("-", "_")
    parts = []
    for part in raw.split("_"):
        match = re.fullmatch(r"(\d+)(?:\.\d+)?", part)
        parts.append(str(int(match.group(1))) if match else part.lower())
    return "_".join(parts)


def header_map(sheet) -> dict[str, int]:
    return {
        canonical_text(cell.value): cell.column
        for cell in sheet[1]
        if canonical_text(cell.value)
    }


def row_key(sheet, row_number: int, columns: dict[str, int]) -> tuple[str, str, str]:
    return (
        canonical_text(sheet.cell(row_number, columns["Google页签"]).value),
        canonical_text(sheet.cell(row_number, columns["Google行号"]).value),
        normalize_slot(sheet.cell(row_number, columns["归一化卡槽"]).value),
    )


def nonempty_row(sheet, row_number: int) -> bool:
    return any(sheet.cell(row_number, column).value not in (None, "") for column in range(1, sheet.max_column + 1))


def workbook_snapshot(workbook) -> dict:
    formulas = {}
    formula_errors = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas[f"{sheet.title}!{cell.coordinate}"] = cell.value
                if cell.value in FORMULA_ERRORS:
                    formula_errors.append(f"{sheet.title}!{cell.coordinate}")
    return {
        "sheetNames": list(workbook.sheetnames),
        "dimensions": {sheet.title: sheet.dimensions for sheet in workbook.worksheets},
        "formulas": formulas,
        "formulaErrors": formula_errors,
    }


def is_yellow(cell) -> bool:
    color = str(cell.fill.fgColor.rgb or "").upper()
    return cell.fill.fill_type == "solid" and color.endswith(YELLOW)


def run(input_path: Path, output_path: Path, receipt_path: Path) -> dict:
    input_path = input_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()
    receipt_path = receipt_path.expanduser().resolve()
    if not input_path.is_file():
        raise RuntimeError(f"输入工作簿不存在: {input_path}")
    if input_path == output_path:
        raise RuntimeError("为保留原始交付文件，输出路径不得与输入路径相同")

    input_hash = sha256(input_path)
    workbook = load_workbook(input_path, data_only=False, keep_links=True)
    before = workbook_snapshot(workbook)
    all_sheet_name = ALL_SHEET if ALL_SHEET in workbook.sheetnames else LEGACY_ALL_SHEET
    missing_sheets = [name for name in (all_sheet_name, ACTIVATED_SHEET) if name not in workbook.sheetnames]
    if missing_sheets:
        raise RuntimeError(f"缺少工作表: {missing_sheets}")

    all_sheet = workbook[all_sheet_name]
    activated_sheet = workbook[ACTIVATED_SHEET]
    all_headers = header_map(all_sheet)
    activated_headers = header_map(activated_sheet)
    missing_headers = {
        all_sheet_name: [name for name in KEY_HEADERS if name not in all_headers],
        ACTIVATED_SHEET: [name for name in KEY_HEADERS if name not in activated_headers],
    }
    if any(missing_headers.values()):
        raise RuntimeError(f"缺少匹配字段: {missing_headers}")

    all_rows: dict[tuple[str, str, str], list[int]] = {}
    for row_number in range(2, all_sheet.max_row + 1):
        if nonempty_row(all_sheet, row_number):
            all_rows.setdefault(row_key(all_sheet, row_number, all_headers), []).append(row_number)
    activated_keys = [
        row_key(activated_sheet, row_number, activated_headers)
        for row_number in range(2, activated_sheet.max_row + 1)
        if nonempty_row(activated_sheet, row_number)
    ]
    duplicate_all = {"|".join(key): rows for key, rows in all_rows.items() if len(rows) != 1}
    duplicate_activated = {
        "|".join(key): count for key, count in Counter(activated_keys).items() if count != 1
    }
    missing_matches = ["|".join(key) for key in activated_keys if len(all_rows.get(key, [])) != 1]
    if duplicate_all or duplicate_activated or missing_matches:
        raise RuntimeError(json.dumps({
            "duplicateAllKeys": duplicate_all,
            "duplicateActivatedKeys": duplicate_activated,
            "missingOrAmbiguousMatches": missing_matches,
        }, ensure_ascii=False))

    yellow_fill = PatternFill(fill_type="solid", fgColor=YELLOW)
    target_rows = sorted({all_rows[key][0] for key in activated_keys})
    for row_number in target_rows:
        for column in range(1, all_sheet.max_column + 1):
            all_sheet.cell(row_number, column).fill = yellow_fill

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{output_path.stem}.", suffix=output_path.suffix, dir=output_path.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()

    with zipfile.ZipFile(output_path) as archive:
        bad_member = archive.testzip()
    verified = load_workbook(output_path, data_only=False, keep_links=True)
    after = workbook_snapshot(verified)
    verified_sheet = verified[all_sheet_name]
    incomplete_rows = [
        row_number
        for row_number in target_rows
        if not all(is_yellow(verified_sheet.cell(row_number, column)) for column in range(1, verified_sheet.max_column + 1))
    ]
    verified.close()
    structure_unchanged = before["sheetNames"] == after["sheetNames"] and before["dimensions"] == after["dimensions"]
    formulas_unchanged = before["formulas"] == after["formulas"]
    formula_errors = after["formulaErrors"]
    passed = (
        bad_member is None
        and structure_unchanged
        and formulas_unchanged
        and not formula_errors
        and not incomplete_rows
        and len(target_rows) == len(activated_keys)
    )
    result = {
        "status": "passed" if passed else "failed",
        "input": str(input_path),
        "output": str(output_path),
        "receipt": str(receipt_path),
        "inputSha256": input_hash,
        "outputSha256": sha256(output_path),
        "activatedSubsetRows": len(activated_keys),
        "highlightedFullRows": len(target_rows),
        "yellow": YELLOW,
        "incompleteHighlightRows": incomplete_rows,
        "sheetStructureUnchanged": structure_unchanged,
        "formulasUnchanged": formulas_unchanged,
        "formulaErrorCount": len(formula_errors),
        "formulaErrors": formula_errors,
        "xlsxZipIntegrity": bad_member is None,
        "sourceUnchanged": sha256(input_path) == input_hash,
        "renderingSkipped": True,
        "renderingSkipReason": "OpenSMS巡检策略禁止打开WPS或生成预览",
    }
    if not result["sourceUnchanged"]:
        result["status"] = "failed"
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    receipt_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if result["status"] != "passed":
        raise RuntimeError(json.dumps(result, ensure_ascii=False))
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="标黄Google已激活但系统缺失记录")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("receipt", type=Path)
    args = parser.parse_args()
    print(json.dumps(run(args.input, args.output, args.receipt), ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
