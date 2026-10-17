#!/usr/bin/env python3
"""Offline, source-bound F renewal import. Never contacts or submits to a backend."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timedelta
import fcntl
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from zoneinfo import ZoneInfo

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'services/opensms_document_api/pipeline'))
from task_f_renewal_dates import VERIFIED_RULE, read_renewals, validate_renewal
from opensms_f_renewal import balance, digest, phone, save, source_date
from opensms_export_task_f_deadlines import (
    DEFAULT_DB, DEFAULT_TEMPLATE, NODE, NODE_MODULES, normalize_iccid, verify_output,
)


def read_source(source):
    """Read only A:E; passwords/emails elsewhere in the workbook never enter outputs."""
    source = Path(source).resolve()
    source_hash = digest(source)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
    records, excluded, seen = [], [], set()
    try:
        sheet = wb['Sheet']
        if [c.value for c in next(sheet.iter_rows(max_row=1, max_col=5))] != [
                '电话', '注册状态', '余额', '截止日期', '套餐类型']:
            raise ValueError('source headers changed')
        for cells in sheet.iter_rows(min_row=2, max_col=5):
            values = [c.value for c in cells]
            if all(v is None for v in values):
                continue
            item = {'sourceRow': cells[0].row}
            try:
                item['phone'] = phone(values[0])
            except ValueError as exc:
                excluded.append(item | {'reason': str(exc)})
                continue
            if item['phone'] in seen:
                raise ValueError('duplicate source phone; identity batch stopped')
            seen.add(item['phone'])
            try:
                if any(c.data_type in {'f', 'e'} for c in cells):
                    raise ValueError('formula/error source cell')
                amount = balance(values[2])
                if amount is None:
                    raise ValueError('missing balance')
                raw_plan = values[4]
                if raw_plan is None or str(raw_plan).strip() == '':
                    plan_type, defaulted = 30, True
                elif str(raw_plan).strip() in {'29', '29.0', '30', '30.0'}:
                    plan_type, defaulted = int(float(raw_plan)), False
                else:
                    raise ValueError('unsupported plan type')
                cutoff = source_date(values[3], wb.epoch)
                item.update(status=str(values[1] or ''), balanceCad=str(amount),
                            planType=plan_type, planDefaulted=defaulted,
                            sourceCutoff=cutoff.isoformat() if cutoff else None)
                records.append(item)
            except ValueError as exc:
                excluded.append(item | {'reason': str(exc)})
    finally:
        wb.close()
    if digest(source) != source_hash:
        raise ValueError('source changed during extraction')
    return {'source': str(source), 'sourceSha256': source_hash, 'sourceSheet': 'Sheet',
            'sourceRows': len(records) + len(excluded), 'rows': records, 'excluded': excluded}


def identity_source(state_file, requested):
    state_file = Path(state_file).resolve()
    state = json.loads(state_file.read_text())
    acquire = state.get('stages', {}).get('acquire', {})
    source = Path(acquire.get('sourceArchive') or '').resolve()
    source_hash = acquire.get('sourceSha256')
    if (state.get('status') != 'success' or acquire.get('sourceVerified') is not True
            or not source.is_file() or digest(source) != source_hash
            or not state.get('summary', {}).get('sourceWorkbook')
            or Path(state['summary']['sourceWorkbook']).resolve() != source):
        raise ValueError('exact successful system export identity evidence required')
    result = {}
    wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = wb['sheet1']
        headers = [c.value for c in next(sheet.iter_rows(max_row=1, max_col=8))]
        if headers != ['电话', 'ICCID', '运营商', 'opr', '卡槽', '状态', '激活日期', '截止日']:
            raise ValueError('system identity headers changed')
        for cells in sheet.iter_rows(min_row=2, max_col=8):
            values = [c.value for c in cells]
            try:
                key = phone(values[0])
            except ValueError:
                continue
            if key not in requested:
                continue
            if key in result:
                raise ValueError('system phone has multiple identities; batch stopped')
            if any(c.data_type in {'f', 'e'} for c in cells[:5]):
                raise ValueError('invalid system identity source cell')
            if not any(label.lower() in str(values[3]).lower() for label in ('kooda', 'koodo', 'telus1')):
                raise ValueError('renewal phone is not a Kooda identity')
            activation = source_date(values[6], wb.epoch)
            result[key] = {'iccid': normalize_iccid(values[1]), 'slot': str(values[4] or ''),
                           'identitySourceRow': cells[0].row,
                           'systemActivation': activation.isoformat() if activation else None}
            if not result[key]['slot']:
                raise ValueError('missing system identity slot')
    finally:
        wb.close()
    if digest(source) != source_hash:
        raise ValueError('system export changed during mapping')
    if len({r['iccid'] for r in result.values()}) != len(result):
        raise ValueError('multiple source phones share an ICCID; batch stopped')
    return result, {'identitySource': str(source), 'identitySourceSha256': source_hash,
                    'identityState': str(state_file), 'identityRunId': state['runId'],
                    'identitySourceKind': 'verified_system_export'}


def make_plan(source, state_file, database=DEFAULT_DB):
    original = read_source(source)
    mapping, identity = identity_source(state_file, {r['phone'] for r in original['rows']})
    rows, excluded = [], list(original['excluded'])
    with sqlite3.connect(Path(database).resolve().as_uri() + '?mode=ro', uri=True) as conn:
        conn.execute('BEGIN')
        base = {r[0]: r[1:] for r in conn.execute(
            'SELECT query_iccid,activation_date,cutoff_date,lookup_status FROM kooda_activation')}
        renewals = read_renewals(conn)
        events = {}
        if conn.execute("SELECT 1 FROM sqlite_master WHERE name='kooda_renewal_events'").fetchone():
            events = {r[0]: json.loads(r[1]) for r in conn.execute(
                'SELECT event_id,proof_json FROM kooda_renewal_events WHERE source_sha256=?',
                (original['sourceSha256'],))}
        for item in original['rows']:
            ident = mapping.get(item['phone'])
            if not ident:
                excluded.append(item | {'reason': 'phone missing from exact system export'})
                continue
            iccid = ident['iccid']
            raw = base.get(iccid)
            if not raw or raw[2] != 'found':
                excluded.append(item | ident | {'reason': 'no confirmed original F activation'})
                continue
            activation, base_cutoff = source_date(raw[0]), source_date(raw[1])
            if (activation is None or base_cutoff != activation + timedelta(days=29)
                    or ident['systemActivation'] != activation.isoformat()):
                raise ValueError('original activation identity/date conflict')
            prior = renewals.get(iccid, {}).get('cutoffDate', raw[1])
            event_id = hashlib.sha256(
                f"{VERIFIED_RULE}:{original['sourceSha256']}:{item['sourceRow']}:{iccid}".encode()).hexdigest()
            existing = events.get(event_id)
            if existing:
                validate_renewal(existing, iccid, raw[0], existing['cutoffDate'])
                # Re-import of the identical source never adds another 29 days.
                calculation_prior = existing['priorEffectiveCutoff']
            else:
                calculation_prior = prior
            cutoff = item['sourceCutoff']
            mode, provisional = 'source_explicit', False
            if cutoff is None:
                if balance(item['balanceCad']) > 0 and source_date(calculation_prior):
                    cutoff = (source_date(calculation_prior) + timedelta(days=29)).isoformat()
                    mode, provisional = 'positive_balance_prior_plus_29', True
                else:
                    mode = 'excluded_missing_date_without_positive_balance'
            rows.append(item | ident | {'activation': raw[0], 'baseCutoffDate': raw[1],
                'priorEffectiveCutoff': calculation_prior, 'beforeCutoff': prior,
                'cutoff': cutoff, 'dateMode': mode, 'provisional': provisional,
                'eventId': event_id, 'existingProof': existing})
    return {k: v for k, v in original.items() if k not in ('rows', 'excluded')} | identity | {
        'rule': VERIFIED_RULE, 'rows': rows, 'excluded': excluded,
        'matchedRows': len(rows), 'dateRows': sum(r['cutoff'] is not None for r in rows),
        'directDateRows': sum(r['dateMode'] == 'source_explicit' for r in rows),
        'provisionalRows': sum(r['provisional'] for r in rows),
        'noDateRows': sum(r['cutoff'] is None for r in rows),
        'planCounts': dict(Counter(str(r['planType']) for r in rows)),
        'defaultPlanRows': sum(r['planDefaulted'] for r in rows)}


def projection_hash(conn, table, columns):
    digestor = hashlib.sha256()
    for row in conn.execute(f'SELECT {",".join(columns)} FROM {table} ORDER BY query_iccid'):
        digestor.update(json.dumps(list(row), ensure_ascii=False, separators=(',', ':')).encode())
        digestor.update(b'\n')
    return digestor.hexdigest()


def commit_update(plan, database, directory, run_id):
    database, directory = Path(database).resolve(), Path(directory).resolve()
    if (directory / 'database_commit.json').exists() or (directory / 'f_database_before.sqlite3').exists():
        raise FileExistsError('commit evidence already exists; do not replay this directory')
    with database.with_name(database.name + '.task_f_writer.lock').open('a+b') as lock:
        try:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise RuntimeError('F writer busy; no database changes made') from None
        if make_plan(plan['source'], plan['identityState'], database) != plan:
            raise ValueError('source, identity or F baseline changed before commit')
        backup = directory / 'f_database_before.sqlite3'
        with sqlite3.connect(database) as source_conn, sqlite3.connect(backup) as copy_conn:
            source_conn.backup(copy_conn)
        with backup.open('rb') as handle:
            os.fsync(handle.fileno())
        old = sqlite3.connect(backup.as_uri() + '?mode=ro', uri=True)
        old_columns = {table: [r[1] for r in old.execute(f'PRAGMA table_info({table})')
                               if r[1] not in {'balance_cad', 'plan_type'}]
                       for table in ('kooda_activation', 'kooda_inventory')}
        before_hashes = {table: projection_hash(old, table, cols) for table, cols in old_columns.items()}
        old.close()
        now = datetime.now(ZoneInfo('Asia/Shanghai')).isoformat()
        committed = []
        for row in plan['rows']:
            proof = row['existingProof']
            if row['cutoff'] and not proof:
                proof = {k: plan[k] for k in ('rule', 'source', 'sourceSha256', 'sourceSheet',
                    'identitySource', 'identitySourceSha256', 'identityRunId', 'identitySourceKind')}
                proof.update(version=2, authorization='current-user-verified-renewal-source',
                    runId=run_id, queryIccid=row['iccid'], phone=row['phone'], slot=row['slot'],
                    sourceRow=row['sourceRow'], identitySourceRow=row['identitySourceRow'],
                    activationDate=row['activation'], baseCutoffDate=row['baseCutoffDate'],
                    sourceCutoff=row['sourceCutoff'], cutoffDate=row['cutoff'],
                    priorEffectiveCutoff=row['priorEffectiveCutoff'],
                    balanceCad=row['balanceCad'], planType=row['planType'], status=row['status'],
                    dateMode=row['dateMode'], provisional=row['provisional'],
                    extensionDays=29 if row['provisional'] else 0,
                    eventId=row['eventId'], recordedAt=now)
            if proof:
                validate_renewal(proof, row['iccid'], row['activation'], row['cutoff'])
            committed.append({k: v for k, v in row.items() if k != 'existingProof'} | {'renewal': proof})
        with sqlite3.connect(database, timeout=30) as conn:
            conn.execute('PRAGMA synchronous=FULL')
            conn.execute('BEGIN IMMEDIATE')
            columns = {r[1] for r in conn.execute('PRAGMA table_info(kooda_activation)')}
            if 'balance_cad' not in columns:
                conn.execute('ALTER TABLE kooda_activation ADD COLUMN balance_cad TEXT')
            if 'plan_type' not in columns:
                conn.execute('ALTER TABLE kooda_activation ADD COLUMN plan_type INTEGER NOT NULL DEFAULT 30 CHECK (plan_type IN (29,30))')
            conn.execute('CREATE TABLE IF NOT EXISTS kooda_renewal_events (event_id TEXT PRIMARY KEY,query_iccid TEXT NOT NULL,source_sha256 TEXT NOT NULL,source_row INTEGER NOT NULL,proof_json TEXT NOT NULL,created_at TEXT NOT NULL,UNIQUE(source_sha256,source_row,query_iccid))')
            conn.execute('CREATE TABLE IF NOT EXISTS kooda_renewal_current (query_iccid TEXT PRIMARY KEY,event_id TEXT NOT NULL REFERENCES kooda_renewal_events(event_id),proof_json TEXT NOT NULL,updated_at TEXT NOT NULL)')
            for row in committed:
                changed = conn.execute('UPDATE kooda_activation SET balance_cad=?,plan_type=? WHERE query_iccid=?',
                    (row['balanceCad'], row['planType'], row['iccid'])).rowcount
                if changed != 1:
                    raise ValueError('F metadata update coverage changed')
                if row['renewal']:
                    proof = row['renewal']
                    encoded = json.dumps(proof, ensure_ascii=False, sort_keys=True)
                    conn.execute('INSERT INTO kooda_renewal_events VALUES (?,?,?,?,?,?) ON CONFLICT(event_id) DO NOTHING',
                        (proof['eventId'], row['iccid'], plan['sourceSha256'], row['sourceRow'], encoded, now))
                    conn.execute('INSERT INTO kooda_renewal_current VALUES (?,?,?,?) ON CONFLICT(query_iccid) DO UPDATE SET event_id=excluded.event_id,proof_json=excluded.proof_json,updated_at=excluded.updated_at',
                        (row['iccid'], proof['eventId'], encoded, now))
            if digest(plan['source']) != plan['sourceSha256'] or digest(plan['identitySource']) != plan['identitySourceSha256']:
                raise ValueError('source changed before transaction commit')
        with sqlite3.connect(database.as_uri() + '?mode=ro', uri=True) as check:
            check.execute('BEGIN')
            if check.execute('PRAGMA quick_check').fetchone()[0] != 'ok':
                raise ValueError('database integrity failed after commit')
            renewals = read_renewals(check)
            for row in committed:
                actual = check.execute('SELECT balance_cad,plan_type FROM kooda_activation WHERE query_iccid=?', (row['iccid'],)).fetchone()
                if actual != (row['balanceCad'], row['planType']) or (row['renewal'] and renewals.get(row['iccid']) != row['renewal']):
                    raise ValueError('independent F database readback failed')
            for table, cols in old_columns.items():
                if projection_hash(check, table, cols) != before_hashes[table]:
                    raise ValueError('original F activation/inventory fields changed')
        result = {'status': 'passed', 'runId': run_id, 'database': str(database),
                  'sourceSha256': plan['sourceSha256'], 'metadataUpdatedCount': len(committed),
                  'deadlineUpdatedCount': sum(r['cutoff'] is not None for r in committed),
                  'independentReadbackPassed': True, 'originalActivationAndInventoryUnchanged': True,
                  'backup': str(backup), 'backupSha256': digest(backup), 'records': committed}
        save(directory / 'database_commit.json', result)
        return result


def export_update(plan, committed, directory):
    from task_f_renewal_dates import read_renewals
    rows = committed['records']
    if committed.get('independentReadbackPassed') is not True:
        raise ValueError('unverified commit')
    with sqlite3.connect(Path(committed['database']).as_uri() + '?mode=ro', uri=True) as check:
        saved = read_renewals(check)
        if any(r['renewal'] and saved.get(r['iccid']) != r['renewal'] for r in rows):
            raise ValueError('F renewal changed before export')
    import_rows = sorted([[r['iccid'], r['cutoff'], r['activation']] for r in rows if r['cutoff']])
    run_id = committed['runId']
    outputs = {'importWorkbook': str(directory / f'任务F_续费截止日导入_{run_id}.xlsx'),
               'reviewWorkbook': str(directory / f'任务F_余额套餐续费核对_{run_id}.xlsx')}
    payload = {'rows': rows, 'excluded': plan['excluded'], 'importRows': import_rows,
               'sourceName': Path(plan['source']).name, 'outputs': outputs}
    save(directory / 'workbook_input.json', payload)
    with tempfile.TemporaryDirectory(prefix='.f_verified_writer_', dir=directory) as temporary:
        work = Path(temporary)
        (work / 'node_modules').symlink_to(NODE_MODULES, target_is_directory=True)
        shutil.copyfile(ROOT / 'scripts/opensms_f_verified_renewal_writer.mjs', work / 'writer.mjs')
        process = subprocess.run([str(NODE), str(work / 'writer.mjs'), str(DEFAULT_TEMPLATE),
            str(directory / 'workbook_input.json'), str(directory)], capture_output=True, text=True, timeout=180)
        if process.returncode:
            raise RuntimeError('artifact writer failed: ' + process.stderr[-1500:])
    verify_output(Path(outputs['importWorkbook']), import_rows)
    # Reopen the independently saved review workbook; no reliance on writer counters.
    with zipfile.ZipFile(outputs['reviewWorkbook']) as archive:
        if archive.testzip() is not None:
            raise ValueError('review XLSX ZIP integrity failed')
    wb = openpyxl.load_workbook(outputs['reviewWorkbook'], read_only=False, data_only=False)
    try:
        sheet = wb['续费核对']
        if sheet.max_row != len(rows) + 5:
            raise ValueError('review row count mismatch')
        for index, row in enumerate(rows, 6):
            cells = list(sheet.iter_rows(min_row=index, max_row=index))[0]
            expected = [row['phone'], row['iccid'], float(balance(row['balanceCad'])), row['planType']]
            if ([c.value for c in cells[:4]] != expected or cells[1].data_type != 's'
                or cells[1].number_format != '@'):
                raise ValueError('review identity/balance/plan mismatch')
            for column, key in ((4, 'sourceCutoff'), (5, 'beforeCutoff'), (6, 'cutoff'), (8, 'activation')):
                value = cells[column].value
                actual = value.strftime('%Y-%m-%d') if value else None
                if actual != row[key] or cells[column].number_format != 'yyyy-mm-dd':
                    raise ValueError('review dates mismatch')
            expected_label = '临时推算：库内有效截止日+29天' if row['provisional'] else ('表内日期' if row['cutoff'] else '缺日期：保持原值')
            if cells[7].value != expected_label:
                raise ValueError('review date provenance mismatch')
    finally:
        wb.close()
    return outputs | {'importCount': len(import_rows), 'reviewCount': len(rows),
        'importSha256': digest(outputs['importWorkbook']), 'reviewSha256': digest(outputs['reviewWorkbook']),
        'independentReadbackPassed': True}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--identity-state', type=Path, required=True)
    parser.add_argument('--db', type=Path, default=DEFAULT_DB)
    parser.add_argument('--directory', type=Path)
    parser.add_argument('--authorization-source', choices=['current-user'])
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()
    plan = make_plan(args.source, args.identity_state, args.db)
    counts = {k: plan[k] for k in ('sourceRows', 'matchedRows', 'dateRows', 'directDateRows',
                                   'provisionalRows', 'noDateRows', 'planCounts', 'defaultPlanRows')}
    if args.dry_run:
        print(json.dumps({'status': 'planned', 'networkAccess': False, **counts,
                          'excludedCount': len(plan['excluded'])}, ensure_ascii=False))
        return
    if args.authorization_source != 'current-user':
        raise ValueError('current human authorization required')
    from opensms_storage_policy import config, require_external
    settings = config()
    require_external(settings)
    run_id = datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%dT%H%M%S+0800')
    directory = (args.directory or Path(settings['bulkRoot']) / 'f_verified_renewals' / run_id).resolve()
    if not directory.is_relative_to(Path(settings['bulkRoot']).resolve()):
        raise ValueError('backup and artifacts must remain on verified zs01')
    directory.mkdir(parents=True, exist_ok=False)
    save(directory / 'source_plan.json', plan)
    print(json.dumps({'event': 'f_verified_renewal_started', 'runId': run_id,
                      'directory': str(directory), 'networkAccess': False}, ensure_ascii=False),
          file=sys.stderr, flush=True)
    committed = commit_update(plan, args.db, directory, run_id)
    outputs = export_update(plan, committed, directory)
    from opensms_task_f_db import export_catalog
    library = directory / '任务F_ICCID库.json'
    export_catalog(args.db, library, ROOT / 'opensms_artifacts/state/task_f_kooda_iccid_library_latest.json')
    result = {'status': 'passed', 'runId': run_id, 'database': str(args.db.resolve()),
        'networkAccess': False, 'backendSubmitted': False, 'dateFormat': 'yyyy-mm-dd',
        **counts, 'excludedCount': len(plan['excluded']), 'excluded': plan['excluded'], **outputs,
        'databaseReceipt': str(directory / 'database_commit.json'), 'databaseBackup': committed['backup'],
        'library': str(library), 'librarySha256': digest(library),
        'originalActivationDatesPreserved': True}
    save(directory / 'completion.json', result)
    print(json.dumps({k: v for k, v in result.items() if k != 'excluded'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
