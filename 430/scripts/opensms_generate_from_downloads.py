#!/usr/bin/env python3
"""Generate a local system-source bundle; Google is disabled by default."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import subprocess
import sys
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


def load_date_freezer():
    """Load sibling scripts even when this CLI is imported by absolute file path.

    Do not modify sys.path: a temporary dependency registration exists only while
    executing the cache module, then the original module registry is restored.
    """
    sibling = Path(__file__).resolve().parent

    def load(name, filename):
        spec = importlib.util.spec_from_file_location(name, sibling / filename)
        if spec is None or spec.loader is None:
            raise ImportError(f"无法载入本地日期模块: {filename}")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    dependency_name = "opensms_export_task_f_deadlines"
    dependency = load("_opensms_local_deadline_dependency", "opensms_export_task_f_deadlines.py")
    previous = sys.modules.get(dependency_name)
    try:
        sys.modules[dependency_name] = dependency
        return load("_opensms_local_frozen_dates", "opensms_report_date_cache.py").freeze
    finally:
        if previous is None:
            sys.modules.pop(dependency_name, None)
        else:
            sys.modules[dependency_name] = previous


freeze = load_date_freezer()


TZ = ZoneInfo("Asia/Hong_Kong")
ROOT = Path("/Users/zhangsan/Downloads/430")
DEFAULT_ARTIFACT_ROOT = ROOT / "opensms_artifacts"
DEFAULT_OUTPUT_DIR = ROOT / "outputs"
BUILDER = ROOT / "services/opensms_document_api/pipeline/build_card_slot_report.sh"


def now_local() -> datetime:
    return datetime.now(TZ)


def stamp(value: datetime | None = None) -> str:
    return (value or now_local()).strftime("%Y%m%dT%H%M%S%z")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_xlsx(path: Path) -> dict:
    if not path.is_file():
        raise FileNotFoundError(f"输入文件不存在: {path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"输入文件不是 xlsx: {path}")
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise RuntimeError(f"XLSX 压缩结构损坏: {path} -> {bad_member}")
    stat_result = path.stat()
    return {
        "path": str(path),
        "size": stat_result.st_size,
        "mtime": datetime.fromtimestamp(stat_result.st_mtime, TZ).isoformat(),
        "sha256": sha256(path),
        "zipIntegrity": "passed",
    }


def validate_system_kooda(path_value: str | None, expected_hash: str | None, run_id: str | None) -> dict:
    """Fail closed on missing provenance; this entrypoint never fetches inputs."""
    if not path_value or not expected_hash or not run_id or not str(run_id).strip():
        raise ValueError("本地生成必须明确提供本轮系统Kooda快照路径、SHA256和runId；不会自动查找历史快照")
    if not re.fullmatch(r"[a-fA-F0-9]{64}", expected_hash):
        raise ValueError("系统Kooda快照SHA256格式无效")
    path = Path(path_value).expanduser().resolve()
    raw = path.read_bytes()
    actual_hash = hashlib.sha256(raw).hexdigest()
    if actual_hash != expected_hash.lower():
        raise ValueError("本轮系统Kooda快照SHA256不一致")
    payload = json.loads(raw)
    if not isinstance(payload, dict) or payload.get("status") != "passed" or payload.get("runId") != run_id:
        raise ValueError("本轮系统Kooda快照status或runId不一致")
    rows, total = payload.get("systemRows"), payload.get("totalRows")
    if not isinstance(rows, list) or type(total) is not int or total < 0 or len(rows) != total:
        raise ValueError("本轮系统Kooda快照行数不完整")
    for key in ("paginationTotalStable", "paginationBackendIdsUnique", "silentBrowserMode"):
        if payload.get(key) is not True:
            raise ValueError(f"本轮系统Kooda快照未通过{key}")
    if payload.get("browserForegroundActivationAttempted") is not False:
        raise ValueError("本轮系统Kooda快照缺少后台静默证明")
    return {
        "path": str(path), "sha256": actual_hash, "runId": run_id,
        "totalRows": total, "status": "passed", "source": "system",
    }


def number(value) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("%", "").strip()
    try:
        return float(text) if text else 0.0
    except ValueError:
        return 0.0


def analyze_export(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = list(next(rows))
    indexes = {
        str(value).strip(): index
        for index, value in enumerate(headers)
        if value not in (None, "")
    }
    dates = []
    for header in headers:
        match = re.fullmatch(r"(\d{2}-\d{2})发送窗口数", str(header or "").strip())
        if match:
            dates.append(match.group(1))
    if not dates:
        workbook.close()
        raise RuntimeError("export4day 没有按日发送窗口列")

    latest = max(dates, key=lambda value: tuple(map(int, value.split("-"))))
    send_index = indexes[f"{latest}发送窗口数"]
    reply_index = indexes[f"{latest}回复窗口数"]
    total_rows = 0
    total_send = 0.0
    total_reply = 0.0
    zero_reply = 0
    for values in rows:
        total_rows += 1
        send = number(values[send_index] if send_index < len(values) else None)
        reply = number(values[reply_index] if reply_index < len(values) else None)
        total_send += send
        total_reply += reply
        if send > 70 and reply == 0:
            zero_reply += 1
    workbook.close()
    return {
        "latestDate": latest,
        "rows": total_rows,
        "send": int(total_send),
        "reply": int(total_reply),
        "replyRatePct": None if total_send <= 0 else round(total_reply / total_send * 100, 3),
        "zeroReplySendGt70": zero_reply,
    }


def parse_json_stdout(stdout: str) -> dict:
    start = stdout.find("{")
    if start < 0:
        raise RuntimeError("本地工作簿生成器没有返回 JSON")
    return json.loads(stdout[start:])


def validate_card_receipt(payload: dict, no_google: bool = False) -> None:
    """Only publish the direct local builder's complete, successful receipt."""
    if not isinstance(payload, dict) or payload.get("result") != "成功":
        raise RuntimeError("本地工作簿生成器未返回成功结果")
    checks = payload.get("checks")
    if not isinstance(checks, dict) or not checks or any(value is not True for value in checks.values()):
        raise RuntimeError("本地工作簿缺少完整通过的独立校验checks")
    if payload.get("failed_checks") != []:
        raise RuntimeError("本地工作簿复核存在失败项或缺少failed_checks")
    if not payload.get("output_path"):
        raise RuntimeError("本地工作簿生成器缺少输出路径")
    if no_google:
        expected = {"enabled": False, "status": "skipped", "reason": "user_disabled",
                    "unused_port_device_scope": "system_source_only"}
        actual = payload.get("process", {}).get("google_dependency")
        review = payload.get("verification", {}).get("google_dependency_review", {})
        if actual != expected or any(review.get(key) != value for key, value in expected.items()) or review.get("matches") is not True:
            raise RuntimeError("无Google文档缺少生成及独立回读的禁用依赖证明")


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def write_report(path: Path, payload: dict) -> None:
    source = payload["inputs"]["export"]
    association = payload["inputs"]["association"]
    no_google = payload["inputs"].get("google_dependency", {}).get("enabled") is False
    system_kooda = payload["inputs"]["systemKooda"]
    analysis = payload["analysis"]
    card = payload["cardSlotReport"]
    process = card.get("process") or {}
    enrichment = card.get("enrichment") or {}
    verification = card.get("verification") or {}
    trend = enrichment.get("decreasing_rate_counts") or {}
    checks = card.get("checks") or {}
    failed_checks = [name for name, passed in checks.items() if not passed]
    result_path = card["output_path"]
    local_stage = "0（本地下载文件模式未执行）"
    lines = [
        "## 本轮概况",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| 模式 | {payload['mode']} |",
        f"| 最新统计日 | {analysis['latestDate']} |",
        f"| 主表行数 | {analysis['rows']:,} |",
        f"| 总耗时 | {payload['durationSeconds']:.2f}秒 |",
        f"| 结果 | {card.get('result', '未知')} |",
        f"| Google匹配行数 | {process.get('google_matched_rows', '未知')} |",
        f"| Kooda系统库存行数 | {system_kooda['totalRows']} |",
        f"| 未匹配卡槽 | {process.get('unmatched_slot_count', '未知')} |",
        f"| 谷歌有-系统无 | {process.get('google_all_missing_from_main_rows', '未知')} |",
        f"| 谷歌文档中有激活日期-系统无法使用 | {process.get('google_activated_missing_from_main_rows', '未知')} |",
        f"| 已激活未上架 | {process.get('activated_unlisted_rows', '未知')} |",
        "",
        "## 0回复与回复率",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| 发送窗口数 | {analysis['send']:,} |",
        f"| 回复窗口数 | {analysis['reply']:,} |",
        f"| 回复率 | {analysis['replyRatePct']}% |",
        f"| 发送大于70且0回复 | {analysis['zeroReplySendGt70']} |",
        f"| 持续2天降低 | {trend.get('持续2天降低', 0)} |",
        f"| 持续3天降低 | {trend.get('持续3天降低', 0)} |",
        f"| 持续4天降低 | {trend.get('持续4天降低', 0)} |",
        f"| 3日/4日发送窗口数与回复率重算差异 | {verification.get('window_metric_mismatch_count', '未知')} |",
        "",
        "## 运营商调整",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| 候选数 | {local_stage} |",
        f"| 后台有效数 | {local_stage} |",
        f"| 实际写入 | {local_stage} |",
        f"| 成功数 | {local_stage} |",
        f"| 失败数 | {local_stage} |",
        f"| 独立复核结果 | {local_stage} |",
        "",
        "## 下架恢复",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| 候选数 | {local_stage} |",
        f"| 实际写入 | {local_stage} |",
        f"| 成功数 | {local_stage} |",
        f"| 失败数 | {local_stage} |",
        f"| 独立复核结果 | {local_stage} |",
        "",
        "## 余量与提醒",
        "",
        "| 分组 | 当日剩余量 | 当日总量 | 剩余比例 | 回复率 | 回复率统计日期 | 状态 |",
        "|---|---:|---:|---:|---:|---|---|",
    ]
    for group in ("A1", "A12", "A11", "B1", "B11", "B12", "C2", "A3", "B3", "B1V", "B11V", "B12V"):
        lines.append(f"| {group} | 未知 | 未知 | 未知 | 未知 | 未知 | 本地下载文件模式未查询 |")
    lines.extend([
        "",
        "## 文件与回执",
        "",
        f"[打开本轮源文件]({source['path']})",
        "",
        source["path"],
        "",
        f"[打开本轮关联表]({association['path']})" if association else "Google关联表：已禁用（用户要求），不下载、不读取。",
        "",
        association["path"] if association else "",
        "",
        f"[打开卡槽关联结果表]({result_path})",
        "",
        result_path,
        "",
        f"[打开完整报告]({path})",
        "",
        str(path),
        "",
        f"- Kooda系统来源：{system_kooda['path']}；运行编号：{system_kooda['runId']}；SHA-256：`{system_kooda['sha256']}`。",
        "- Kooda主表ICCID保持系统cardStore原始文本；激活日期、截止日逐字段优先本轮系统导出，缺失才使用文档开始前冻结的本地F日期缓存。无日期列的旧格式允许本轮系统库存。Chatr缺失日期才回退Google；开卡日期规则不变。",
        f"- 本轮只读冻结日期缓存：{payload['inputs']['reportDateCache']['path']}；SHA-256：`{payload['inputs']['reportDateCache']['sha256']}`。本入口不查询任务F、不写数据库或云端。",
        f"- 逐字段日期来源计数：{json.dumps(process.get('date_resolution', {}).get('field_counts', {}), ensure_ascii=False)}。",
        "",
        f"完整回执路径：{payload['receiptPath']}",
        "",
        "## 异常",
        "",
        f"- 工作簿复核失败项：{', '.join(failed_checks) if failed_checks else '无'}",
        f"- 公式错误：{verification.get('formula_error_matches', '未知')}",
        "- 运营商写入、下架恢复、余量与后台独立复核未执行；本入口只处理本轮下载文件。",
    ])
    if no_google:
        google_rows = ("| Google匹配行数 |", "| 未匹配卡槽 |", "| 谷歌有-系统无 |", "| 谷歌文档中有激活日期-系统无法使用 |")
        lines = [line for line in lines if not line.startswith(google_rows)]
        lines = [line.replace("Chatr缺失日期才回退Google；开卡日期规则不变。", "Chatr激活日期允许空白；截止日仅使用系统导出。开卡日期留空，不倒推日期。") for line in lines]
        lines.insert(lines.index("## 0回复与回复率"), "系统无-谷歌表有仅覆盖本轮系统源主表有效数字设备的1至64端口；不包含只存在于Google的设备。\n")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        handle.write("\n".join(lines))


def main() -> int:
    parser = argparse.ArgumentParser(
        description="根据本轮系统 export4day.xlsx、系统Kooda JSON和冻结日期缓存生成并复核文档；默认禁用Google，不联网、不执行业务。"
    )
    parser.add_argument("export", help="本轮下载的 export4day.xlsx")
    parser.add_argument("association", nargs="?", help="仅显式旧版测试模式使用的关联表；正常模式忽略且不读取")
    parser.add_argument("--no-google", action="store_true", help="显式禁用Google（默认行为）")
    parser.add_argument("--legacy-google-association", action="store_true", help="仅离线旧版测试可开启Google关联读取")
    parser.add_argument("--system-kooda", default=os.environ.get("OPENSMS_SYSTEM_KOODA_PATH"),
                        help="本轮系统Kooda库存JSON（或OPENSMS_SYSTEM_KOODA_PATH）")
    parser.add_argument("--system-kooda-sha256", default=os.environ.get("OPENSMS_SYSTEM_KOODA_SHA256"),
                        help="该JSON的SHA256（或OPENSMS_SYSTEM_KOODA_SHA256）")
    parser.add_argument("--run-id", default=os.environ.get("OPENSMS_RUN_ID"),
                        help="与系统库存一致的本轮runId（或OPENSMS_RUN_ID）")
    parser.add_argument("--report-date-cache", default=os.environ.get("OPENSMS_REPORT_DATE_CACHE_PATH"),
                        help="已有本轮冻结日期缓存JSON；不填则在本地只读冻结任务F库")
    parser.add_argument("--report-date-cache-sha256", default=os.environ.get("OPENSMS_REPORT_DATE_CACHE_SHA256"))
    parser.add_argument("--task-f-db", default=os.environ.get("OPENSMS_TASK_F_DB_PATH", str(ROOT / "opensms_artifacts/state/task_f_kooda_activation.sqlite3")),
                        help="只读日期库，缺文件生成合法空快照，不创建/迁移数据库")
    parser.add_argument("--backend-date", required=True, help="必填：本轮确认的后台服务器日期YYYY-MM-DD；不猜测本机日期")
    parser.add_argument("--mode", choices=("fast", "full"), default="fast")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--artifact-root", default=str(DEFAULT_ARTIFACT_ROOT))
    args = parser.parse_args()

    started = time.monotonic()
    export_path = Path(args.export).expanduser().resolve()
    no_google = not args.legacy_google_association or args.no_google or os.environ.get("OPENSMS_GOOGLE_ASSOCIATION_ENABLED") == "0"
    if not no_google and not args.association:
        raise ValueError("旧版关联测试模式必须明确指定association")
    association_path = None if no_google else Path(args.association).expanduser().resolve()
    if export_path == association_path:
        raise ValueError("export 与关联表不能是同一个文件")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", args.backend_date):
        raise ValueError("--backend-date 必须是 YYYY-MM-DD")
    try:
        datetime.strptime(args.backend_date, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError("--backend-date 不是有效的真实日历日期") from exc

    inputs_started = time.monotonic()
    system_kooda_receipt = validate_system_kooda(args.system_kooda, args.system_kooda_sha256, args.run_id)
    with ThreadPoolExecutor(max_workers=3, thread_name_prefix="opensms-local-input") as executor:
        export_future = executor.submit(validate_xlsx, export_path)
        association_future = executor.submit(validate_xlsx, association_path) if not no_google else None
        analysis_future = executor.submit(analyze_export, export_path)
        export_receipt, association_receipt, analysis = export_future.result(), association_future.result() if association_future else None, analysis_future.result()
    inputs_duration = time.monotonic() - inputs_started

    artifact_root = Path(args.artifact_root).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    run_stamp = stamp()
    receipt_path = artifact_root / "snapshots" / f"opensms_local_bundle_{run_stamp}_receipt.json"
    card_receipt_path = artifact_root / "snapshots" / f"opensms_card_slot_report_{run_stamp}_receipt.json"
    report_path = artifact_root / "reports" / f"opensms_local_bundle_{run_stamp}.md"
    if bool(args.report_date_cache) != bool(args.report_date_cache_sha256):
        raise ValueError("已有冻结日期缓存必须同时提供路径和SHA256")
    if args.report_date_cache:
        cache_path = Path(args.report_date_cache).expanduser().resolve()
        cache_hash = sha256(cache_path)
        if cache_hash != str(args.report_date_cache_sha256).lower():
            raise ValueError("已有冻结日期缓存SHA256不一致")
        date_cache = {"path": str(cache_path), "sha256": cache_hash, "runId": system_kooda_receipt["runId"], "provided": True}
    else:
        frozen = freeze(Path(args.task_f_db).expanduser().resolve(),
                        artifact_root / "snapshots" / f"opensms_report_date_cache_{run_stamp}_{time.time_ns()}.json",
                        system_kooda_receipt["runId"])
        date_cache = {**frozen, "path": frozen["output"], "provided": False}

    env = os.environ.copy()
    env["OPENSMS_GOOGLE_ASSOCIATION_ENABLED"] = "0" if no_google else "1"
    env.update({
        "OPENSMS_SYSTEM_KOODA_PATH": system_kooda_receipt["path"],
        "OPENSMS_SYSTEM_KOODA_SHA256": system_kooda_receipt["sha256"],
        "OPENSMS_RUN_ID": system_kooda_receipt["runId"],
        "OPENSMS_REPORT_DATE_CACHE_PATH": date_cache["path"],
        "OPENSMS_REPORT_DATE_CACHE_SHA256": date_cache["sha256"],
    })
    # This is document-only; inherited business gates are never forwarded.
    for task in "BCDEF":
        for prefix in ("OPENSMS_FORCE_TASK_", "OPENSMS_ENABLE_TASK_", "OPENSMS_EXECUTE_TASK_"):
            env.pop(f"{prefix}{task}", None)
    env.pop("OPENSMS_TASK_F_POLICY_CATALOG_SYNC", None)
    if args.mode == "fast":
        env["OPENSMS_SKIP_PREVIEWS"] = "1"
        env["OPENSMS_FAST_VERIFY"] = "1"
    else:
        env.pop("OPENSMS_SKIP_PREVIEWS", None)
        env.pop("OPENSMS_FAST_VERIFY", None)
    command = [
        str(BUILDER),
        str(export_path),
        str(association_path) if association_path else "-",
        str(output_dir),
        args.backend_date,
    ]

    build_started = time.monotonic()
    completed = subprocess.run(
        command,
        cwd=ROOT,
        env=env,
        text=True,
        capture_output=True,
        check=False,
    )
    build_duration = time.monotonic() - build_started
    if completed.returncode != 0:
        raise RuntimeError(
            "本地工作簿生成失败\n"
            f"stdout:\n{completed.stdout[-4000:]}\n"
            f"stderr:\n{completed.stderr[-4000:]}"
        )
    card_receipt = parse_json_stdout(completed.stdout)
    validate_card_receipt(card_receipt, no_google=no_google)

    result_path = Path(card_receipt["output_path"]).resolve()
    result_receipt = validate_xlsx(result_path)
    duration = time.monotonic() - started
    payload = {
        "createdAt": now_local().isoformat(),
        "runId": system_kooda_receipt["runId"],
        "serverDate": args.backend_date,
        "result": card_receipt.get("result"),
        "mode": args.mode,
        "durationSeconds": round(duration, 3),
        "durations": {
            "inputValidationAndAnalysisSeconds": round(inputs_duration, 3),
            "workbookBuildAndVerificationSeconds": round(build_duration, 3),
        },
        "inputs": {
            "export": export_receipt,
            "association": association_receipt,
            "google_dependency": card_receipt["process"]["google_dependency"],
            "systemKooda": system_kooda_receipt,
            "reportDateCache": date_cache,
        },
        "analysis": analysis,
        "cardSlotReport": card_receipt,
        "cardSlotReceiptPath": str(card_receipt_path),
        "systemKoodaSource": system_kooda_receipt,
        "output": result_receipt,
        "reportPath": str(report_path),
        "receiptPath": str(receipt_path),
        "liveBackendActionsExecuted": False,
    }
    write_json(card_receipt_path, card_receipt)
    write_report(report_path, payload)
    write_json(receipt_path, payload)
    print(json.dumps({
        "result": card_receipt.get("result"),
        "mode": args.mode,
        "durationSeconds": round(duration, 3),
        "outputPath": str(result_path),
        "reportPath": str(report_path),
        "receiptPath": str(receipt_path),
        "cardSlotReceiptPath": str(card_receipt_path),
        "systemKoodaSource": system_kooda_receipt,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"错误: {error}", file=sys.stderr)
        raise SystemExit(1)
