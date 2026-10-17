#!/usr/bin/env python3
"""Offline add-on for a supplied workbook; never exports or calls business APIs.

Normal Task A calls the same aggregation inside its existing single-save builder.
This explicit CLI creates a NEW example/copy without overwriting its input.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from services.opensms_document_api.pipeline.device_daily_statistics import (
    SHEET_NAME, aggregate_device_daily, write_device_daily_sheet,
)
from services.opensms_document_api.pipeline.verify_device_daily_statistics import review_device_daily_statistics


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonical_row(row):
    # Excel numeric 10.0 and 10 are equivalent; only None/empty-string are blank.
    return tuple(("blank", None) if value is None or value == "" else
                 ("bool", value) if isinstance(value, bool) else
                 ("number", value) if isinstance(value, (int, float)) else
                 (type(value).__name__, value) for value in row)


def run(source, output, workers=4):
    started = time.perf_counter()
    source, output = Path(source).resolve(), Path(output).resolve()
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise ValueError("输入必须为现有xlsx文件")
    if source == output or output.exists() or output.suffix.lower() != ".xlsx":
        raise ValueError("输出必须是不存在的新xlsx路径；禁止覆盖源文件或历史结果")
    before_hash = digest(source)
    book = load_workbook(source)
    main = book["sheet1"] if "sheet1" in book.sheetnames else book.active
    main_title = main.title
    original_sheets = [name for name in book.sheetnames if name != SHEET_NAME]
    source_values = list(main.iter_rows(values_only=True))
    result = aggregate_device_daily(source_values[0], source_values[1:], workers)
    receipt = write_device_daily_sheet(book, main, result)
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=".device-daily-", suffix=".xlsx", dir=output.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        book.save(temporary)
        readback = load_workbook(temporary, read_only=True)
        try:
            review = review_device_daily_statistics(readback, source_values[0], source_values[1:])
            from itertools import zip_longest
            sentinel = object()
            unchanged = {}
            for title in original_sheets:
                unchanged[title] = title in readback.sheetnames and all(
                    left is not sentinel and right is not sentinel and canonical_row(left) == canonical_row(right)
                    for left, right in zip_longest(book[title].iter_rows(values_only=True),
                                                  readback[title].iter_rows(values_only=True), fillvalue=sentinel)
                )
            if not review["matches"] or not all(unchanged.values()):
                raise RuntimeError(f"设备统计或原工作表独立回读失败: {review}; {unchanged}")
        finally:
            readback.close()
        if digest(source) != before_hash:
            raise RuntimeError("输入文件在处理期间发生变化，拒绝交付")
        os.link(temporary, output)  # exclusive publication; never replace a file
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return {"status": "passed", "mode": "offline_existing_workbook_copy_not_fresh_export",
            "source": str(source), "source_sha256": before_hash, "output": str(output),
            "output_sha256": digest(output), "main_sheet": main_title,
            "device_daily_statistics": receipt, "verification": review,
            "original_sheet_values_preserved": unchanged,
            "elapsed_seconds": round(time.perf_counter() - started, 3)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    try:
        result = run(args.input, args.output, args.workers)
    except Exception as error:
        print(json.dumps({"status": "failed", "error": str(error)}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
