#!/usr/bin/env python3
"""Persistent Task F Kooda ICCID activation-date database.

The database is deliberately independent from the per-run delivery folder. A
legacy inventory-only mode keeps terminal records once they satisfy the
current schema. Report-backfill mode instead queries only current report
Kooda rows with blank activation dates; existing valid dates are ready for
manual upload, and prior blank/not-found records may be queried again.
Transient errors never replace confirmed valid dates.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import re
import sqlite3
import sys
import tempfile
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'services/opensms_document_api/pipeline'))
from archived_system_identity import index_phones, match_archived
from report_kooda_source import is_report_source, validate_projection, missing_activation_rows, iccid_key, ELIGIBILITY
from task_f_renewal_dates import read_renewals


TERMINAL_STATUSES = ("found", "not_found", "activation_missing")
ICCID_PATTERN = re.compile(r"^\d{18,22}$")
RECORD_SCHEMA_VERSION = 2
RECORDED_WEB_FIELDS = (
    "Purchase Date",
    "Act. Request Date",
    "Activation Date",
    "SIM Number",
    "SKU",
)


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def normalize_iccid(value: Any) -> str:
    raw = re.sub(r"\s+", "", str(value or "").strip())
    raw = re.sub(r"[Ff]$", "", raw)
    return raw if ICCID_PATTERN.fullmatch(raw) else ""


def parse_activation_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip().replace(".", "/")
    if not raw:
        return None
    raw = re.sub(r"\s+", " ", raw)
    for fmt in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%b %d, %Y",
        "%B %d, %Y",
        "%d %b %Y",
        "%d %B %Y",
    ):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    match = re.search(r"((?:19|20)\d{2})[-/](\d{1,2})[-/](\d{1,2})", raw)
    if match:
        try:
            return date(*map(int, match.groups()))
        except ValueError:
            return None
    return None


def cutoff_date_for_activation(activation: date) -> date:
    # Activation day counts as day 1, so a 30-day inclusive period ends +29.
    return activation + timedelta(days=29)


def atomic_bytes(path: Path, encoded: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(handle, "wb") as stream:
            stream.write(encoded)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary_name, path)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)


def encode_json(payload: Any) -> bytes:
    return (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def atomic_json(path: Path, payload: Any) -> None:
    atomic_bytes(path, encode_json(payload))


def open_database(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path, timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA synchronous=FULL")
    connection.execute("PRAGMA busy_timeout=30000")
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS kooda_inventory (
            query_iccid TEXT PRIMARY KEY,
            source_iccid TEXT NOT NULL,
            phone TEXT,
            slot TEXT,
            opr TEXT,
            carrier_name TEXT,
            first_seen_at TEXT NOT NULL,
            last_seen_at TEXT NOT NULL,
            source_occurrences_json TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS kooda_activation (
            query_iccid TEXT PRIMARY KEY,
            source_iccid TEXT NOT NULL,
            activation_date TEXT,
            cutoff_date TEXT,
            lookup_status TEXT NOT NULL,
            queried_at TEXT NOT NULL,
            detail_sim_iccid TEXT,
            worker_id INTEGER,
            source_phone TEXT,
            source_slot TEXT,
            source_opr TEXT,
            source_carrier_name TEXT,
            purchase_date TEXT,
            purchase_date_raw TEXT,
            act_request_date TEXT,
            activation_date_raw TEXT,
            sim_number TEXT,
            sim_number_raw TEXT,
            sku TEXT,
            source_data_json TEXT NOT NULL DEFAULT '[]',
            record_schema_version INTEGER NOT NULL DEFAULT 1,
            receipt_json TEXT NOT NULL,
            CHECK (lookup_status IN ('found', 'not_found', 'activation_missing')),
            CHECK (
                (lookup_status = 'found' AND activation_date IS NOT NULL AND cutoff_date IS NOT NULL)
                OR lookup_status <> 'found'
            )
        );

        CREATE INDEX IF NOT EXISTS idx_kooda_activation_status
            ON kooda_activation(lookup_status);
        CREATE INDEX IF NOT EXISTS idx_kooda_activation_cutoff
            ON kooda_activation(cutoff_date);

        CREATE TABLE IF NOT EXISTS grabandgo_catalog_import (
            source_sha256 TEXT PRIMARY KEY,
            source_path TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            source_finished_at TEXT,
            source_row_count INTEGER NOT NULL,
            found_count INTEGER NOT NULL,
            activation_missing_count INTEGER NOT NULL,
            receipt_json TEXT NOT NULL
        );
        """
    )
    existing_columns = {
        str(row[1]) for row in connection.execute("PRAGMA table_info(kooda_activation)")
    }
    migrations = {
        "purchase_date": "TEXT",
        "purchase_date_raw": "TEXT",
        "act_request_date": "TEXT",
        "activation_date_raw": "TEXT",
        "sim_number": "TEXT",
        "sim_number_raw": "TEXT",
        "sku": "TEXT",
        "source_data_json": "TEXT NOT NULL DEFAULT '[]'",
        "record_schema_version": "INTEGER NOT NULL DEFAULT 1",
        "data_source": "TEXT",
        "source_file": "TEXT",
        "source_sha256": "TEXT",
        "vendor_ordinal": "INTEGER",
        "cost": "TEXT",
        "balance_cad": "TEXT",
        "plan_type": "INTEGER NOT NULL DEFAULT 30 CHECK (plan_type IN (29, 30))",
    }
    for column, declaration in migrations.items():
        if column not in existing_columns:
            connection.execute(f"ALTER TABLE kooda_activation ADD COLUMN {column} {declaration}")
    connection.commit()
    return connection


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def card_rows(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        rows = payload
    elif isinstance(payload, dict):
        rows = payload.get("rows") or payload.get("cards") or []
    else:
        rows = []
    return [row for row in rows if isinstance(row, dict)]


def activation_records_for_iccids(
    connection: sqlite3.Connection, query_iccids: list[str],
) -> dict[str, sqlite3.Row]:
    """Read only current inventory keys and fields used by pending calculation."""
    records: dict[str, sqlite3.Row] = {}
    # Remain below SQLite builds with the older 999-parameter limit.
    for offset in range(0, len(query_iccids), 900):
        batch = query_iccids[offset:offset + 900]
        placeholders = ",".join("?" for _ in batch)
        for record in connection.execute(
            f"""SELECT query_iccid, lookup_status, record_schema_version,
                       purchase_date, purchase_date_raw, act_request_date,
                       activation_date, activation_date_raw, sim_number,
                       sim_number_raw, sku, cutoff_date, queried_at
                FROM kooda_activation WHERE query_iccid IN ({placeholders})""",
            batch,
        ):
            records[record["query_iccid"]] = record
    return records


def valid_found_record(record: Any) -> bool:
    if not record or record["lookup_status"] != "found":
        return False
    activation = parse_activation_date(record["activation_date"])
    cutoff = parse_activation_date(record["cutoff_date"])
    return bool(activation and cutoff and cutoff == cutoff_date_for_activation(activation))


def report_slot_key(value: Any) -> str:
    raw = unicodedata.normalize("NFKC", str(value or "").strip()).replace(" ", "").replace("-", "_")
    parts = []
    for part in raw.split("_"):
        match = re.fullmatch(r"(\d+)(?:\.\d+)?", part)
        parts.append(str(int(match.group(1))) if match else part.lower())
    return "_".join(parts)


def report_phone_key(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value or ""))[-10:]


def report_iccid_key(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    normalized = normalize_iccid(value)
    return normalized if re.fullmatch(r"[0-9]{18,22}", normalized) else ""


def verified_creation_fallbacks(proof: dict, run_id: str, inventory_sha256: str) -> dict[int, dict]:
    """Only an independently verified provenance receipt may widen F's scope.

    A date merely equal to creation time is not evidence of a fallback. Keep
    real system/local dates out of the query queue even when they look alike.
    """
    review = (proof.get("verification") or {}).get("date_resolution_review") or {}
    rows = review.get("creation_date_fallback_rows", [])
    if not isinstance(rows, list):
        raise RuntimeError("任务F创建日期兜底来源清单无效")
    if not rows:
        return {}
    system = (proof.get("verification") or {}).get("system_kooda_source_review") or {}
    if (proof.get("result") != "成功" or not proof.get("checks")
            or any(value is not True for value in proof["checks"].values())
            or proof.get("failed_checks") != [] or review.get("matches") is not True
            or system.get("matches") is not True or system.get("run_id") != run_id
            or system.get("snapshot_sha256") != inventory_sha256):
        raise RuntimeError("任务F创建日期兜底来源未通过本轮独立核验")
    result = {}
    for row in rows:
        if (not isinstance(row, dict) or type(row.get("row")) is not int or row["row"] < 2
                or row["row"] in result or len(report_phone_key(row.get("phone"))) != 10
                or not parse_activation_date(row.get("activation"))
                or parse_activation_date(row.get("cutoff")) != parse_activation_date(row["activation"]) + timedelta(days=29)):
            raise RuntimeError("任务F创建日期兜底来源行缺失、重复或无效")
        result[row["row"]] = row
    return result


def export_document_backfill_rows(inventory, document, document_sha256, run_id):
    """New A rule: raw export activation missing, not final/fallback/cutoff gaps."""
    from openpyxl import load_workbook
    validate_projection(inventory, run_id)
    content = document.read_bytes()
    if hashlib.sha256(content).hexdigest() != document_sha256:
        raise RuntimeError('任务F原始10日表范围的文档SHA256不一致')
    eligible_sources = {row['sourceRow']: row for row in missing_activation_rows(inventory)}
    sources = {row['sourceRow']: row for row in inventory['systemRows']}
    selected, excluded = [], []
    counts = dict(documentKoodaRows=0, documentActivationPresentRows=0,
        documentMissingActivationRows=0, documentMissingCutoffRows=0,
        documentMissingDateRows=0, documentCreationFallbackRows=0, documentLookupNeededRows=0)
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        if 'sheet1' not in book.sheetnames:
            raise RuntimeError('任务F补查文档缺少sheet1')
        iterator = book['sheet1'].iter_rows(values_only=True)
        headers = list(next(iterator, ()))
        for name in ('电话', 'opr', '卡槽', 'ICCID', '激活日期', '截止日'):
            if headers.count(name) != 1:
                raise RuntimeError(f'任务F补查文档缺少唯一{name}列')
        indexes = {name: headers.index(name) for name in headers if name}
        seen = set()
        for number, values in enumerate(iterator, 2):
            row = {name: values[pos] for name, pos in indexes.items()}
            source = sources.get(number)
            kooda = bool(re.search(r'kooda|koodo|telus1', str(row['opr'] or ''), re.I))
            if bool(source) != kooda:
                raise RuntimeError('任务F文档与原始10日表行范围不一致')
            if not kooda:
                continue
            seen.add(number)
            if (not (row['ICCID'] == source['cardStore']
                     or (row['ICCID'] in (None, '') and source['cardStore'] in (None, '')))
                    or report_phone_key(row['电话']) != report_phone_key(source['cardNo'])
                    or report_slot_key(row['卡槽']) != report_slot_key(source['slotIdentif'])):
                raise RuntimeError('任务F文档与原始10日表ICCID/源行身份不一致')
            activation_blank = row['激活日期'] is None or (isinstance(row['激活日期'], str) and not row['激活日期'].strip())
            cutoff_blank = row['截止日'] is None or (isinstance(row['截止日'], str) and not row['截止日'].strip())
            for field, blank in (('激活日期', activation_blank), ('截止日', cutoff_blank)):
                if not blank and parse_activation_date(row[field]) is None:
                    raise RuntimeError('任务F文档非空日期无效')
            counts['documentKoodaRows'] += 1
            counts['documentActivationPresentRows'] += int(not activation_blank)
            counts['documentMissingActivationRows'] += int(activation_blank)
            counts['documentMissingCutoffRows'] += int(cutoff_blank)
            counts['documentMissingDateRows'] += int(activation_blank or cutoff_blank)
            if number not in eligible_sources:
                continue
            counts['documentLookupNeededRows'] += 1
            query = iccid_key(source['cardStore'])
            evidence = {'documentRow': number, 'sourceRowIdentity': source['sourceRowIdentity'],
                'phone': str(row['电话'] or ''), 'slot': str(row['卡槽'] or ''),
                'sourceIccid': source['cardStore'], 'queryIccid': query,
                'missingFields': ['激活日期'], 'lookupReason': ELIGIBILITY,
                'dateSource': 'raw_10day_export_activation_missing'}
            if not query:
                excluded.append({**evidence, 'reason': 'document_missing_or_invalid_iccid'})
                continue
            selected.append({**evidence, 'cardStore': source['cardStore'], 'opr': source['opr'],
                'carrierName': str(source.get('carrierName') or ''), 'sourceData': source})
        if seen != set(sources):
            raise RuntimeError('任务F文档未完整覆盖本轮10日表源行')
    finally:
        book.close()
    queries = {row['queryIccid'] for row in selected}
    metadata = {'reportBackfillMode': True, 'documentPath': str(document.resolve()),
        'documentSha256': document_sha256, 'documentRunId': run_id, **counts,
        'dateEligibilityMode': ELIGIBILITY, 'sourceScope': inventory['sourceScope'],
        'sourceWorkbookSha256': inventory['sourceWorkbookSha256'],
        'documentMissingActivationIccidCount': len(queries), 'documentMissingDateIccidCount': len(queries),
        'documentDuplicateIccidRows': len(selected) - len(queries),
        'documentExcludedCount': len(excluded), 'documentExcludedRows': excluded,
        'reportEligibleRecords': [{k: v for k, v in row.items()
            if k not in ('cardStore', 'opr', 'carrierName', 'sourceData')} for row in selected],
        'inventoryValidUniqueIccidCount': len({iccid_key(row['cardStore']) for row in sources.values()} - {''}),
        'documentKoRefreshRows': 0, 'totalLookupNeededRows': counts['documentLookupNeededRows'],
        'documentMutationAllowed': False, 'cloudUploadMode': 'manual-user-upload-only'}
    return selected, metadata


def document_backfill_rows(inventory: Any, document: Path, document_sha256: str,
                           run_id: str, *, creation_fallbacks: dict[int, dict] | None = None,
                           ko_context: dict | None = None
                           ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Query blank dates and verified creation-date estimates, never real dates."""
    from openpyxl import load_workbook

    if isinstance(inventory, dict) and is_report_source(inventory):
        # The former ko-sheet/full-system extension cannot widen this scope.
        return export_document_backfill_rows(inventory, document, document_sha256, run_id)

    if not run_id or not isinstance(inventory, dict) or inventory.get("runId") != run_id:
        raise RuntimeError("任务F补查文档与库存runId不一致")
    if inventory.get("status") != "passed" or not isinstance(inventory.get("systemRows"), list):
        raise RuntimeError("任务F补查需要本轮已通过的完整系统库存")
    for field in ("paginationTotalStable", "paginationBackendIdsUnique", "silentBrowserMode"):
        if inventory.get(field) is not True:
            raise RuntimeError(f"任务F补查库存未通过{field}")
    if inventory.get("browserForegroundActivationAttempted") is not False:
        raise RuntimeError("任务F补查库存缺少后台静默证据")
    raw_rows = inventory["systemRows"]
    if type(inventory.get("totalRows")) is not int or inventory["totalRows"] != len(raw_rows):
        raise RuntimeError("任务F补查库存systemRows数量不一致")
    if not re.fullmatch(r"[a-fA-F0-9]{64}", document_sha256 or ""):
        raise RuntimeError("任务F补查缺少有效文档SHA256")
    document_bytes = document.read_bytes()
    actual_hash = hashlib.sha256(document_bytes).hexdigest()
    if actual_hash != document_sha256.lower():
        raise RuntimeError("任务F补查文档SHA256不一致")
    indexes: dict[tuple[str, str], dict[str, Any]] = {}
    inventory_iccids = set()
    backend_ids = set()
    for row in raw_rows:
        if not isinstance(row, dict) or not re.search(r"kooda|koodo|telus1", str(row.get("opr") or ""), re.I):
            raise RuntimeError("任务F补查库存存在非Kooda记录")
        backend_id = str(row.get("id") or "").strip()
        if not backend_id or backend_id in backend_ids:
            raise RuntimeError("任务F补查库存后台ID缺失或重复")
        backend_ids.add(backend_id)
        key = (report_slot_key(row.get("slotIdentif")), report_phone_key(row.get("cardNo")))
        iccid = report_iccid_key(row.get("cardStore"))
        if iccid:
            inventory_iccids.add(iccid)
        if not key[0] or len(key[1]) != 10:
            continue
        if key in indexes:
            raise RuntimeError("任务F补查库存卡槽与号码复合键重复")
        indexes[key] = row
    archived_index = index_phones(raw_rows)
    eligible, excluded, selected = [], [], []
    kooda_count = present_count = missing_count = missing_cutoff_count = missing_date_count = 0
    creation_fallback_count = lookup_needed_count = 0
    creation_fallbacks = creation_fallbacks or {}
    missing_activation_iccids = set()
    seen_document_keys = set()
    workbook = load_workbook(io.BytesIO(document_bytes), read_only=True, data_only=False)
    try:
        if "sheet1" not in workbook.sheetnames:
            raise RuntimeError("任务F补查文档缺少指定主表sheet1")
        sheet = workbook["sheet1"]
        row_iter = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(row_iter, [])]
        for name in ("opr", "卡槽", "ICCID", "激活日期"):
            if headers.count(name) != 1:
                raise RuntimeError(f"任务F补查主表列缺失或重复:{name}")
        phone_headers = [name for name in ("电话", "号码", "手机号") if name in headers]
        if len(phone_headers) != 1 or headers.count(phone_headers[0]) != 1:
            raise RuntimeError("任务F补查主表电话列不唯一")
        index = {name: headers.index(name) for name in ("opr", "卡槽", "ICCID", "激活日期", phone_headers[0])}
        if headers.count("截止日") > 1:
            raise RuntimeError("任务F补查主表截止日列重复")
        if "截止日" in headers:
            index["截止日"] = headers.index("截止日")
        for row_number, row in enumerate(row_iter, 2):
            if not re.search(r"kooda|koodo|telus1", str(row[index["opr"]] or ""), re.I):
                continue
            kooda_count += 1
            activation = row[index["激活日期"]]
            activation_missing = activation is None or (isinstance(activation, str) and not activation.strip())
            if not activation_missing:
                if parse_activation_date(activation) is None:
                    raise RuntimeError("任务F补查主表激活日期非空但无效")
                present_count += 1
            else:
                missing_count += 1
            cutoff_missing = False
            if "截止日" in index:
                cutoff_value = row[index["截止日"]]
                cutoff_missing = cutoff_value is None or (isinstance(cutoff_value, str) and not cutoff_value.strip())
                if not cutoff_missing and parse_activation_date(cutoff_value) is None:
                    raise RuntimeError("任务F补查主表截止日非空但无效")
                missing_cutoff_count += int(cutoff_missing)
            missing_date_count += int(activation_missing or cutoff_missing)
            fallback = creation_fallbacks.get(row_number)
            if fallback:
                if (report_phone_key(row[index[phone_headers[0]]]) != report_phone_key(fallback["phone"])
                        or parse_activation_date(activation) != parse_activation_date(fallback["activation"])
                        or "截止日" not in index
                        or parse_activation_date(row[index["截止日"]]) != parse_activation_date(fallback["cutoff"])):
                    raise RuntimeError("任务F创建日期兜底来源与当前文档行不一致")
                creation_fallback_count += 1
            if not activation_missing and not cutoff_missing and not fallback:
                continue
            lookup_needed_count += 1
            raw_iccid = row[index["ICCID"]]
            iccid = report_iccid_key(raw_iccid)
            phone, slot = row[index[phone_headers[0]]], row[index["卡槽"]]
            evidence = {"documentRow": row_number, "phone": str(phone or ""), "slot": str(slot or ""),
                        "sourceIccid": raw_iccid, "queryIccid": iccid,
                        "missingFields": (["激活日期"] if activation_missing or fallback else [])
                                         + (["截止日"] if cutoff_missing or fallback else [])}
            if fallback:
                evidence["dateSource"] = "creation_date_fallback"
                evidence["lookupReason"] = "creation_estimate_is_not_authoritative_activation"
            if not iccid:
                excluded.append({**evidence, "reason": "document_missing_or_invalid_iccid"})
                continue
            key = (report_slot_key(slot), report_phone_key(phone))
            if key in seen_document_keys:
                raise RuntimeError("任务F补查文档卡槽与号码复合键重复")
            seen_document_keys.add(key)
            source = indexes.get(key)
            if source is None:
                archived = match_archived(archived_index, phone=phone, slot=slot,
                    status=row[headers.index('状态')] if '状态' in headers else None,
                    created=row[headers.index('创建时间')] if '创建时间' in headers else None,
                    iccid=raw_iccid)
                if archived is not None:
                    excluded.append({**evidence, 'backendId': archived['id'],
                                     'reason': 'archived_off_shelf_no_current_slot'})
                    continue
            if not source or report_iccid_key(source.get("cardStore")) != iccid:
                raise RuntimeError("任务F补查文档ICCID、卡槽、号码与本轮系统库存不一致")
            evidence["backendId"] = source.get("id")
            if activation_missing:
                missing_activation_iccids.add(iccid)
            eligible.append(evidence)
            selected.append({
                **evidence,
                "cardStore": source["cardStore"], "phone": str(phone or ""), "slot": str(slot or ""),
                "opr": source["opr"], "carrierName": str(source.get("carrierName") or source.get("carrireName") or ""),
                "documentRow": row_number, "sourceData": source,
            })
        ko_rows = []
        if ko_context is not None:
            from opensms_task_f_ko_refresh import selected_rows
            ko_rows = selected_rows(workbook, ko_context['proof'], inventory,
                ko_context['inventoryHash'], run_id, ko_context['cachePath'], ko_context['cacheSha256'])
            selected.extend(ko_rows)
            eligible.extend({key: value for key, value in row.items()
                             if key not in ('cardStore', 'opr', 'carrierName', 'sourceData')}
                            for row in ko_rows)
    finally:
        workbook.close()
    unique_count = len({row["queryIccid"] for row in eligible})
    return selected, {
        "reportBackfillMode": True, "documentPath": str(document.resolve()),
        "documentSha256": actual_hash, "documentRunId": run_id,
        "documentKoodaRows": kooda_count, "documentActivationPresentRows": present_count,
        "documentMissingActivationRows": missing_count,
        "documentMissingActivationIccidCount": len(missing_activation_iccids),
        "documentMissingCutoffRows": missing_cutoff_count,
        "documentMissingDateRows": missing_date_count,
        "documentCreationFallbackRows": creation_fallback_count,
        "documentLookupNeededRows": lookup_needed_count,
        **({"documentKoRefreshRows": len(ko_rows), "koRefresh": {
            key: value for key, value in ko_context.items() if key not in ('proof', 'inventoryHash')},
            "totalLookupNeededRows": lookup_needed_count + len(ko_rows)} if ko_context is not None else {}),
        "documentMissingDateIccidCount": unique_count,
        "documentDuplicateIccidRows": len(eligible) - unique_count,
        "documentExcludedCount": len(excluded), "documentExcludedRows": excluded,
        "reportEligibleRecords": eligible, "inventoryValidUniqueIccidCount": len(inventory_iccids),
        "documentMutationAllowed": False, "cloudUploadMode": "manual-user-upload-only",
    }


def prepare_pending(database: Path, cards_path: Path, output: Path, *, document: Path | None = None,
                    document_sha256: str | None = None, run_id: str | None = None,
                    document_proof: Path | None = None, document_proof_sha256: str | None = None,
                    ko_cache: Path | None = None, ko_cache_sha256: str | None = None) -> dict[str, Any]:
    inventory_bytes = cards_path.read_bytes()
    inventory = json.loads(inventory_bytes)
    inventory_hash = hashlib.sha256(inventory_bytes).hexdigest()
    report_mode = any(value is not None for value in (document, document_sha256, run_id))
    report_metadata = {}
    if report_mode:
        if document is None or not document_sha256 or not run_id:
            raise RuntimeError("任务F报表补查必须同时提供document/document-sha256/run-id")
        fallback_rows = {}
        if document_proof is not None or document_proof_sha256 is not None:
            if (document_proof is None or not document_proof_sha256
                    or file_sha256(document_proof) != document_proof_sha256):
                raise RuntimeError("任务F文档来源回执SHA256不一致")
            fallback_rows = verified_creation_fallbacks(load_json(document_proof), run_id, inventory_hash)
        ko_context = None
        if ko_cache is not None or ko_cache_sha256 is not None:
            if ko_cache is None or not ko_cache_sha256 or document_proof is None:
                raise RuntimeError('ko补查必须同时绑定冻结库与文档独立回执')
            from opensms_task_f_ko_refresh import pending_context
            ko_context = pending_context(load_json(document_proof), inventory, inventory_hash,
                                         run_id, ko_cache, ko_cache_sha256)
            ko_context.update(proof=load_json(document_proof), inventoryHash=inventory_hash)
        source_rows, report_metadata = document_backfill_rows(inventory, document, document_sha256, run_id,
                                                              creation_fallbacks=fallback_rows, ko_context=ko_context)
        if document_proof is not None:
            report_metadata.update(documentProofPath=str(document_proof.resolve()),
                documentProofSha256=document_proof_sha256,
                dateEligibilityMode=ELIGIBILITY if is_report_source(inventory) else "missing_or_verified_creation_fallback")
    else:
        if ko_cache is not None or ko_cache_sha256 is not None:
            raise RuntimeError('ko补查不得脱离本轮报表')
        if document_proof is not None or document_proof_sha256 is not None:
            raise RuntimeError("任务F来源回执必须绑定报表模式")
        source_rows = card_rows(inventory)
    grouped: dict[str, list[dict[str, Any]]] = {}
    invalid_rows: list[dict[str, Any]] = []
    for row in source_rows:
        source_iccid = str(
            row.get("cardStore")
            or row.get("cardStorage")
            or row.get("iccid")
            or row.get("ICCID")
            or ""
        ).strip()
        query_iccid = normalize_iccid(source_iccid)
        if not query_iccid:
            invalid_rows.append({"sourceIccid": source_iccid, "row": row})
            continue
        grouped.setdefault(query_iccid, []).append(row)

    created_at = now_iso()
    pending: list[dict[str, Any]] = []
    known: list[dict[str, Any]] = []
    schema_upgrade_pending = 0
    refresh_existing_missing = 0
    with open_database(database) as connection:
        existing_records = activation_records_for_iccids(connection, list(grouped))
        inventory_values = []
        for query_iccid, occurrences in grouped.items():
            first = occurrences[0]
            source_iccid = str(first.get("cardStore") or first.get("cardStorage") or first.get("iccid") or "").strip()
            inventory_values.append(
                (
                    query_iccid,
                    source_iccid,
                    str(first.get("phone") or first.get("cardNo") or ""),
                    str(first.get("slot") or first.get("cardSlot") or ""),
                    str(first.get("opr") or ""),
                    str(first.get("carrierName") or first.get("sourceCarrierName") or ""),
                    created_at,
                    created_at,
                    json.dumps(occurrences, ensure_ascii=False, separators=(",", ":")),
                ),
            )
            existing = existing_records.get(query_iccid)
            item = {
                "queryIccid": query_iccid,
                "sourceIccid": source_iccid,
                "phone": str(first.get("phone") or first.get("cardNo") or ""),
                "slot": str(first.get("slot") or first.get("cardSlot") or ""),
                "opr": str(first.get("opr") or ""),
                "carrierName": str(first.get("carrierName") or first.get("sourceCarrierName") or ""),
                "sourceOccurrenceCount": len(occurrences),
                "sourceData": first.get("sourceData") if isinstance(first.get("sourceData"), dict) else first,
                "sourceOccurrences": occurrences,
            }
            existing_schema_version = int(existing["record_schema_version"] or 1) if existing else 0
            keep_existing = (bool(existing and parse_activation_date(existing['activation_date']))
                if is_report_source(inventory) else (valid_found_record(existing) if report_mode else bool(
                existing and existing["lookup_status"] in TERMINAL_STATUSES
                and existing_schema_version >= RECORD_SCHEMA_VERSION)))
            if keep_existing:
                known.append({
                    **item,
                    "lookupStatus": existing["lookup_status"],
                    "purchaseDate": existing["purchase_date"],
                    "purchaseDateRaw": existing["purchase_date_raw"],
                    "actRequestDate": existing["act_request_date"],
                    "activationDate": existing["activation_date"],
                    "activationDateRaw": existing["activation_date_raw"],
                    "simNumber": existing["sim_number"],
                    "simNumberRaw": existing["sim_number_raw"],
                    "sku": existing["sku"],
                    "cutoffDate": existing["cutoff_date"],
                    "queriedAt": existing["queried_at"],
                    "recordSchemaVersion": existing_schema_version,
                })
            else:
                if report_mode and existing:
                    item["refreshReason"] = "report_missing_date_and_database_without_valid_found_date"
                    item["previousLookupStatus"] = existing["lookup_status"]
                    refresh_existing_missing += 1
                elif report_mode:
                    item["refreshReason"] = "report_missing_date_and_no_local_record"
                elif existing and existing["lookup_status"] in TERMINAL_STATUSES:
                    item["refreshReason"] = "record_schema_upgrade_required"
                    schema_upgrade_pending += 1
                pending.append(item)
        # The conflict clause intentionally preserves first_seen_at, avoiding
        # one SELECT per ICCID while retaining its original inventory history.
        connection.executemany(
            """
            INSERT INTO kooda_inventory(
                query_iccid, source_iccid, phone, slot, opr, carrier_name,
                first_seen_at, last_seen_at, source_occurrences_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(query_iccid) DO UPDATE SET
                source_iccid=excluded.source_iccid,
                phone=excluded.phone,
                slot=excluded.slot,
                opr=excluded.opr,
                carrier_name=excluded.carrier_name,
                last_seen_at=excluded.last_seen_at,
                source_occurrences_json=excluded.source_occurrences_json
            """,
            inventory_values,
        )
        connection.commit()

    payload = {
        "status": "passed",
        "createdAt": created_at,
        "database": str(database.resolve()),
        "sourceCards": str(cards_path.resolve()),
        "sourceCardsSha256": inventory_hash,
        "sourceRowCount": len(source_rows),
        "validUniqueIccidCount": len(grouped),
        "invalidIccidCount": len(invalid_rows),
        "alreadyQueriedCount": len(known),
        "pendingCount": len(pending),
        "schemaUpgradePendingCount": schema_upgrade_pending,
        "recordSchemaVersion": RECORD_SCHEMA_VERSION,
        "recordedWebFields": list(RECORDED_WEB_FIELDS),
        "fullSourceDataIncluded": True,
        "terminalStatuses": list(TERMINAL_STATUSES),
        "trailingFNormalization": "strip_exactly_one_final_F_case_insensitive",
        "databaseLookupMode": "current-iccid-batched-indexed-select",
        "databaseLookupBatchCount": (len(grouped) + 899) // 900,
        "inventoryWriteMode": "single-transaction-executemany",
        "rows": pending,
        "alreadyQueried": known,
        "invalidRows": invalid_rows,
    }
    if report_mode:
        payload.update({**report_metadata, "readyForUploadCount": len(known),
                        "refreshExistingMissingCount": refresh_existing_missing,
                        "readyForUploadLegacySchemaCount": sum(
                            row["recordSchemaVersion"] < RECORD_SCHEMA_VERSION for row in known),
                        "reuseRecords": known})
    if file_sha256(cards_path) != inventory_hash:
        raise RuntimeError("任务F补查准备期间库存文件SHA256发生变化，禁止继续查询")
    atomic_json(output, payload)
    return payload


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def import_catalog(database: Path, catalog_path: Path, receipt_path: Path) -> dict[str, Any]:
    """Import a complete Grab&Go SIM-card snapshot as the authoritative local catalog."""
    # Hash and parse the same bytes to avoid a second disk read and ensure the
    # imported rows are exactly the content identified by the receipt SHA.
    catalog_bytes = catalog_path.read_bytes()
    source_hash = hashlib.sha256(catalog_bytes).hexdigest()
    if database.exists():
        cached_connection = sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True, timeout=30)
        try:
            try:
                cached = cached_connection.execute(
                    "SELECT receipt_json FROM grabandgo_catalog_import WHERE source_sha256 = ?",
                    (source_hash,),
                ).fetchone()
            except sqlite3.OperationalError as error:
                if "no such table" not in str(error):
                    raise
                cached = None
            if cached:
                prior_receipt = json.loads(cached[0])
                if (
                    prior_receipt.get("status") == "passed"
                    and prior_receipt.get("sourceSha256") == source_hash
                    and prior_receipt.get("invalidCount") == 0
                    and prior_receipt.get("sourceRowCount") == prior_receipt.get("normalizedUniqueIccidCount")
                ):
                    result = {
                        **prior_receipt,
                        "alreadyImported": True,
                        "database": str(database.resolve()),
                        "catalog": str(catalog_path.resolve()),
                        "receipt": str(receipt_path.resolve()),
                        "sourceHashVerifiedAt": now_iso(),
                        "importValidationMode": "sha256-matched-committed-import",
                        "normalizationSkippedForUnchangedCatalog": True,
                    }
                    atomic_json(receipt_path, result)
                    return result
        finally:
            cached_connection.close()
    payload = json.loads(catalog_bytes)
    if not isinstance(payload, dict) or payload.get("status") != "passed":
        raise RuntimeError("任务F附件不是通过状态的Grab&Go完整清单")
    rows = payload.get("rows")
    if not isinstance(rows, list):
        raise RuntimeError("任务F附件缺少rows数组")
    source_total = int(payload.get("total") or 0)
    extracted = int(payload.get("extractedRowCount") or 0)
    if payload.get("rowCountMatchesTotal") is not True or source_total != len(rows) or extracted != len(rows):
        raise RuntimeError(
            f"任务F附件完整性失败: total={source_total}, extracted={extracted}, rows={len(rows)}"
        )
    if payload.get("missingRequiredFields") not in ([], None):
        raise RuntimeError("任务F附件声明缺少必要字段")

    required_fields = (
        "lblPurchaseDate", "lblActivationRequestDate", "lblActivateDate",
        "lblSimCardNumber", "lblSKU", "lblCost",
    )
    normalized: list[dict[str, Any]] = []
    seen: set[str] = set()
    invalid_rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows, start=1):
        fields = row.get("fields") if isinstance(row, dict) else None
        if not isinstance(fields, dict) or any(name not in fields for name in required_fields):
            invalid_rows.append({"row": row_index, "reason": "missing_field_nodes"})
            continue
        raw_sim = str(fields.get("lblSimCardNumber") or "").strip()
        query_iccid = normalize_iccid(raw_sim)
        if not query_iccid:
            invalid_rows.append({"row": row_index, "reason": "invalid_sim_number"})
            continue
        if query_iccid in seen:
            invalid_rows.append({"row": row_index, "reason": "duplicate_normalized_iccid", "queryIccid": query_iccid})
            continue
        seen.add(query_iccid)
        activation_raw = str(fields.get("lblActivateDate") or "").strip()
        activation = parse_activation_date(activation_raw)
        if activation_raw and activation is None:
            invalid_rows.append({"row": row_index, "reason": "invalid_activation_date", "queryIccid": query_iccid})
            continue
        purchase_raw = str(fields.get("lblPurchaseDate") or "").strip()
        purchase = parse_activation_date(purchase_raw)
        if purchase_raw and purchase is None:
            invalid_rows.append({"row": row_index, "reason": "invalid_purchase_date", "queryIccid": query_iccid})
            continue
        normalized.append({
            "queryIccid": query_iccid,
            "sourceIccid": raw_sim,
            "status": "found" if activation else "activation_missing",
            "purchaseDate": purchase.isoformat() if purchase else "",
            "purchaseDateRaw": purchase_raw,
            "actRequestDate": str(fields.get("lblActivationRequestDate") or "").strip(),
            "activationDate": activation.isoformat() if activation else "",
            "activationDateRaw": activation_raw,
            "cutoffDate": cutoff_date_for_activation(activation).isoformat() if activation else "",
            "simNumber": query_iccid,
            "simNumberRaw": raw_sim,
            "sku": str(fields.get("lblSKU") or "").strip(),
            "cost": str(fields.get("lblCost") or "").strip(),
            "ordinal": int(row.get("ordinal") or row_index),
            "sourceRow": row,
        })
    if invalid_rows or len(normalized) != len(rows):
        raise RuntimeError(f"任务F附件存在{len(invalid_rows)}条无效或重复记录，已拒绝整批导入")

    imported_at = now_iso()
    inserted = 0
    updated = 0
    unchanged = 0
    preserved_valid = 0
    with open_database(database) as connection:
        # Prevent two catalog importers from both passing the ledger check.
        connection.execute("BEGIN IMMEDIATE")
        previous_import = connection.execute(
            "SELECT receipt_json FROM grabandgo_catalog_import WHERE source_sha256 = ?",
            (source_hash,),
        ).fetchone()
        if previous_import:
            prior_receipt = json.loads(previous_import["receipt_json"])
            result = {
                **prior_receipt,
                "status": "passed",
                "alreadyImported": True,
                "database": str(database.resolve()),
                "catalog": str(catalog_path.resolve()),
                "receipt": str(receipt_path.resolve()),
            }
            atomic_json(receipt_path, result)
            return result

        for item in normalized:
            existing = connection.execute(
                "SELECT * FROM kooda_activation WHERE query_iccid = ?",
                (item["queryIccid"],),
            ).fetchone()
            # An old attachment may predate a subsequent successful lookup.
            # Never downgrade valid dates to blank or overwrite a confirmed
            # incremental result merely because a catalog file was re-issued.
            if valid_found_record(existing) and (
                item["status"] != "found" or existing["data_source"] == "grabandgo_incremental_lookup"
            ):
                unchanged += 1
                preserved_valid += 1
                continue
            inventory = connection.execute(
                "SELECT * FROM kooda_inventory WHERE query_iccid = ?",
                (item["queryIccid"],),
            ).fetchone()
            source_data = json.dumps(item["sourceRow"], ensure_ascii=False, separators=(",", ":"))
            receipt_json = json.dumps({
                "source": "grabandgo_complete_catalog_attachment",
                "sourceSha256": source_hash,
                **item,
            }, ensure_ascii=False, separators=(",", ":"))
            comparable = (
                item["status"], item["activationDate"] or None, item["cutoffDate"] or None,
                item["purchaseDate"], item["purchaseDateRaw"], item["actRequestDate"],
                item["activationDateRaw"], item["simNumber"], item["simNumberRaw"],
                item["sku"], item["cost"], "grabandgo_complete_catalog_attachment", source_hash,
            )
            if existing:
                previous = (
                    existing["lookup_status"], existing["activation_date"], existing["cutoff_date"],
                    existing["purchase_date"] or "", existing["purchase_date_raw"] or "",
                    existing["act_request_date"] or "", existing["activation_date_raw"] or "",
                    existing["sim_number"] or "", existing["sim_number_raw"] or "",
                    existing["sku"] or "", existing["cost"] or "",
                    existing["data_source"] or "", existing["source_sha256"] or "",
                )
                if previous == comparable:
                    unchanged += 1
                else:
                    updated += 1
            else:
                inserted += 1
            connection.execute(
                """
                INSERT INTO kooda_activation(
                    query_iccid, source_iccid, activation_date, cutoff_date, lookup_status,
                    queried_at, detail_sim_iccid, worker_id, source_phone, source_slot,
                    source_opr, source_carrier_name, purchase_date, purchase_date_raw,
                    act_request_date, activation_date_raw, sim_number, sim_number_raw, sku,
                    source_data_json, record_schema_version, receipt_json, data_source,
                    source_file, source_sha256, vendor_ordinal, cost
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(query_iccid) DO UPDATE SET
                    source_iccid=excluded.source_iccid,
                    activation_date=excluded.activation_date,
                    cutoff_date=excluded.cutoff_date,
                    lookup_status=excluded.lookup_status,
                    queried_at=excluded.queried_at,
                    detail_sim_iccid=excluded.detail_sim_iccid,
                    source_phone=excluded.source_phone,
                    source_slot=excluded.source_slot,
                    source_opr=excluded.source_opr,
                    source_carrier_name=excluded.source_carrier_name,
                    purchase_date=excluded.purchase_date,
                    purchase_date_raw=excluded.purchase_date_raw,
                    act_request_date=excluded.act_request_date,
                    activation_date_raw=excluded.activation_date_raw,
                    sim_number=excluded.sim_number,
                    sim_number_raw=excluded.sim_number_raw,
                    sku=excluded.sku,
                    source_data_json=excluded.source_data_json,
                    record_schema_version=excluded.record_schema_version,
                    receipt_json=excluded.receipt_json,
                    data_source=excluded.data_source,
                    source_file=excluded.source_file,
                    source_sha256=excluded.source_sha256,
                    vendor_ordinal=excluded.vendor_ordinal,
                    cost=excluded.cost
                """,
                (
                    item["queryIccid"], item["sourceIccid"], item["activationDate"] or None,
                    item["cutoffDate"] or None, item["status"],
                    str(payload.get("finishedAt") or imported_at), item["queryIccid"], None,
                    str(inventory["phone"] if inventory else existing["source_phone"] if existing else ""),
                    str(inventory["slot"] if inventory else existing["source_slot"] if existing else ""),
                    str(inventory["opr"] if inventory else existing["source_opr"] if existing else ""),
                    str(inventory["carrier_name"] if inventory else existing["source_carrier_name"] if existing else ""),
                    item["purchaseDate"], item["purchaseDateRaw"], item["actRequestDate"],
                    item["activationDateRaw"], item["simNumber"], item["simNumberRaw"], item["sku"],
                    source_data, RECORD_SCHEMA_VERSION, receipt_json,
                    "grabandgo_complete_catalog_attachment", str(catalog_path.resolve()), source_hash,
                    item["ordinal"], item["cost"],
                ),
            )
        found_count = sum(item["status"] == "found" for item in normalized)
        activation_missing_count = len(normalized) - found_count
        result = {
            "status": "passed",
            "alreadyImported": False,
            "createdAt": imported_at,
            "database": str(database.resolve()),
            "catalog": str(catalog_path.resolve()),
            "sourceSha256": source_hash,
            "sourceFinishedAt": payload.get("finishedAt"),
            "sourceRowCount": len(rows),
            "normalizedUniqueIccidCount": len(normalized),
            "trailingFNormalization": "strip_exactly_one_final_F_case_insensitive",
            "insertedCount": inserted,
            "updatedCount": updated,
            "unchangedCount": unchanged,
            "preservedValidLookupCount": preserved_valid,
            "foundCount": found_count,
            "activationMissingCount": activation_missing_count,
            "activationField": "lblActivateDate",
            "cutoffRule": "activation_date_plus_29_days_inclusive_30_day_period",
            "invalidCount": 0,
            "importValidationMode": "complete-row-validation-and-transaction",
            "normalizationSkippedForUnchangedCatalog": False,
            "receipt": str(receipt_path.resolve()),
        }
        connection.execute(
            """
            INSERT INTO grabandgo_catalog_import(
                source_sha256, source_path, imported_at, source_finished_at,
                source_row_count, found_count, activation_missing_count, receipt_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                source_hash, str(catalog_path.resolve()), imported_at,
                str(payload.get("finishedAt") or ""), len(rows), found_count,
                activation_missing_count,
                json.dumps(result, ensure_ascii=False, separators=(",", ":")),
            ),
        )
        connection.commit()
    atomic_json(receipt_path, result)
    return result


def import_lookup(database: Path, lookup_path: Path, receipt_path: Path) -> dict[str, Any]:
    lookup_bytes = lookup_path.read_bytes()
    payload = json.loads(lookup_bytes)
    rows = payload.get("results") or payload.get("rows") or []
    if not isinstance(rows, list):
        raise RuntimeError("任务F查询结果缺少 results 数组")
    terminal_rows: list[dict[str, Any]] = []
    transient_rows: list[dict[str, Any]] = []
    imported = 0
    found = 0
    unchanged = 0
    upgraded = 0
    refreshed_missing = 0
    refreshed_blank = 0
    with open_database(database) as connection:
        for row in rows:
            if not isinstance(row, dict):
                continue
            query_iccid = normalize_iccid(row.get("queryIccid") or row.get("iccid"))
            status = str(row.get("status") or "").strip()
            if not query_iccid or status not in TERMINAL_STATUSES:
                transient_rows.append({"queryIccid": query_iccid, "status": status, "error": row.get("error")})
                continue
            activation = parse_activation_date(row.get("activationDate"))
            cutoff = cutoff_date_for_activation(activation) if activation else None
            if status == "found" and activation is None:
                transient_rows.append({"queryIccid": query_iccid, "status": status, "error": "found_without_valid_activation_date"})
                continue
            if status != "found" and activation is not None:
                transient_rows.append({"queryIccid": query_iccid, "status": status,
                                       "error": "non_found_status_with_activation_date"})
                continue
            source_occurrences = row.get("sourceOccurrences")
            inventory = connection.execute(
                "SELECT * FROM kooda_inventory WHERE query_iccid = ?",
                (query_iccid,),
            ).fetchone()
            if not isinstance(source_occurrences, list) or not source_occurrences:
                try:
                    source_occurrences = json.loads(inventory["source_occurrences_json"]) if inventory else []
                except (TypeError, json.JSONDecodeError):
                    source_occurrences = []
            if not isinstance(source_occurrences, list) or not source_occurrences:
                transient_rows.append({"queryIccid": query_iccid, "status": status, "error": "missing_full_source_data"})
                continue
            presence = row.get("webFieldPresence") if isinstance(row.get("webFieldPresence"), dict) else {}
            if status in {"found", "activation_missing"}:
                missing_fields = [
                    name for name in ("purchaseDate", "actRequestDate", "activationDate", "simNumber", "sku")
                    if presence.get(name) is not True
                ]
                if missing_fields:
                    transient_rows.append({
                        "queryIccid": query_iccid,
                        "status": status,
                        "error": f"missing_web_field_nodes:{','.join(missing_fields)}",
                    })
                    continue
            sim_number = normalize_iccid(row.get("simNumber"))
            if status in {"found", "activation_missing"} and sim_number != query_iccid:
                transient_rows.append({
                    "queryIccid": query_iccid,
                    "status": status,
                    "error": "sim_number_does_not_match_query_iccid",
                })
                continue
            source_iccid = str(row.get("sourceIccid") or (inventory["source_iccid"] if inventory else query_iccid))
            detail_iccid = normalize_iccid(row.get("detailSimIccid"))
            if normalize_iccid(source_iccid) != query_iccid:
                transient_rows.append({"queryIccid": query_iccid, "status": status,
                                       "error": "source_iccid_does_not_match_query_iccid"})
                continue
            if status in {"found", "activation_missing"} and detail_iccid != query_iccid:
                transient_rows.append({"queryIccid": query_iccid, "status": status,
                                       "error": "detail_sim_iccid_does_not_match_query_iccid"})
                continue
            existing = connection.execute(
                "SELECT * FROM kooda_activation WHERE query_iccid = ?",
                (query_iccid,),
            ).fetchone()
            existing_schema_version = int(existing["record_schema_version"] or 1) if existing else 0
            existing_valid = valid_found_record(existing)
            incoming_found_upgrade = status == "found" and activation is not None and not existing_valid
            refresh_blank_attempt = bool(payload.get("reportBackfillMode") is True
                                         and existing and not existing_valid
                                         and status in {"activation_missing", "not_found"})
            keep_existing = existing and (
                (existing_valid and status != "found")
                or (existing_schema_version >= RECORD_SCHEMA_VERSION and not incoming_found_upgrade
                    and not refresh_blank_attempt)
            )
            if keep_existing:
                unchanged += 1
                terminal_rows.append({
                    "queryIccid": query_iccid,
                    "status": existing["lookup_status"],
                    "purchaseDate": existing["purchase_date"],
                    "actRequestDate": existing["act_request_date"],
                    "activationDate": existing["activation_date"],
                    "simNumber": existing["sim_number"],
                    "sku": existing["sku"],
                    "cutoffDate": existing["cutoff_date"],
                    "recordSchemaVersion": existing_schema_version,
                    "result": "already_terminal_kept",
                    "activationDateUpdated": False,
                })
                continue
            record_json = json.dumps(row, ensure_ascii=False, separators=(",", ":"))
            values = (
                query_iccid,
                source_iccid,
                activation.isoformat() if activation else None,
                cutoff.isoformat() if cutoff else None,
                status,
                str(row.get("queriedAt") or now_iso()),
                normalize_iccid(row.get("detailSimIccid")) or None,
                row.get("workerId"),
                str(row.get("phone") or (inventory["phone"] if inventory else "")),
                str(row.get("slot") or (inventory["slot"] if inventory else "")),
                str(row.get("opr") or (inventory["opr"] if inventory else "")),
                str(row.get("carrierName") or (inventory["carrier_name"] if inventory else "")),
                str(row.get("purchaseDate") or ""),
                str(row.get("purchaseDateRaw") or row.get("purchaseDate") or ""),
                str(row.get("actRequestDate") or ""),
                str(row.get("activationDateRaw") or row.get("activationDate") or ""),
                sim_number or None,
                str(row.get("simNumberRaw") or row.get("simNumber") or ""),
                str(row.get("sku") or ""),
                json.dumps(source_occurrences, ensure_ascii=False, separators=(",", ":")),
                RECORD_SCHEMA_VERSION,
                record_json,
                "grabandgo_incremental_lookup",
                None,
                None,
                None,
                str(row.get("cost") or ""),
            )
            if existing:
                connection.execute(
                    """
                    UPDATE kooda_activation SET
                        source_iccid=?, activation_date=?, cutoff_date=?, lookup_status=?, queried_at=?,
                        detail_sim_iccid=?, worker_id=?, source_phone=?, source_slot=?, source_opr=?,
                        source_carrier_name=?, purchase_date=?, purchase_date_raw=?, act_request_date=?,
                        activation_date_raw=?, sim_number=?, sim_number_raw=?, sku=?, source_data_json=?,
                        record_schema_version=?, receipt_json=?, data_source=?, source_file=?,
                        source_sha256=?, vendor_ordinal=?, cost=?
                    WHERE query_iccid=?
                    """,
                    (*values[1:], query_iccid),
                )
                if existing_schema_version < RECORD_SCHEMA_VERSION:
                    upgraded += 1
                if incoming_found_upgrade:
                    refreshed_missing += 1
                    result = "refreshed_missing_activation_to_found"
                elif refresh_blank_attempt:
                    refreshed_blank += 1
                    result = "refreshed_blank_lookup_record"
                else:
                    result = "upgraded_to_current_record_schema"
            else:
                connection.execute(
                    """
                INSERT INTO kooda_activation(
                    query_iccid, source_iccid, activation_date, cutoff_date,
                    lookup_status, queried_at, detail_sim_iccid, worker_id,
                    source_phone, source_slot, source_opr, source_carrier_name,
                    purchase_date, purchase_date_raw, act_request_date,
                    activation_date_raw, sim_number, sim_number_raw, sku,
                    source_data_json, record_schema_version, receipt_json, data_source,
                    source_file, source_sha256, vendor_ordinal, cost
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    values,
                )
            imported += 1
            found += int(status == "found")
            terminal_rows.append({
                "queryIccid": query_iccid,
                "status": status,
                "purchaseDate": str(row.get("purchaseDateRaw") or row.get("purchaseDate") or ""),
                "actRequestDate": str(row.get("actRequestDate") or ""),
                "activationDate": activation.isoformat() if activation else None,
                "simNumber": sim_number or None,
                "sku": str(row.get("sku") or ""),
                "cutoffDate": cutoff.isoformat() if cutoff else None,
                "recordSchemaVersion": RECORD_SCHEMA_VERSION,
                "sourceOccurrenceCount": len(source_occurrences),
                "result": result if existing else "inserted",
                "activationDateUpdated": bool(status == "found" and activation is not None
                    and (not existing or existing["activation_date"] != activation.isoformat()
                         or existing["cutoff_date"] != cutoff.isoformat() or not existing_valid)),
            })
        connection.commit()
        counts = {
            row["lookup_status"]: row["count"]
            for row in connection.execute(
                "SELECT lookup_status, COUNT(*) AS count FROM kooda_activation GROUP BY lookup_status"
            )
        }
        total = int(connection.execute("SELECT COUNT(*) FROM kooda_activation").fetchone()[0])
        current_schema_count = int(connection.execute(
            "SELECT COUNT(*) FROM kooda_activation WHERE record_schema_version >= ?",
            (RECORD_SCHEMA_VERSION,),
        ).fetchone()[0])

    receipt = {
        "status": "passed" if not transient_rows else "passed_with_transient_errors",
        "createdAt": now_iso(),
        "database": str(database.resolve()),
        "lookupReceipt": str(lookup_path.resolve()),
        "lookupSha256": hashlib.sha256(lookup_bytes).hexdigest(),
        "runId": payload.get("runId") or payload.get("documentRunId"),
        "documentRunId": payload.get("documentRunId"),
        "resultCount": len(rows),
        "terminalImportedCount": imported,
        "foundImportedCount": found,
        "alreadyTerminalKeptCount": unchanged,
        "schemaUpgradedCount": upgraded,
        "refreshedMissingActivationCount": refreshed_missing,
        "refreshedBlankRecordCount": refreshed_blank,
        "activationDateUpdatedCount": sum(row.get("activationDateUpdated") is True for row in terminal_rows),
        "reportBackfillMode": payload.get("reportBackfillMode") is True,
        "transientErrorCount": len(transient_rows),
        "databaseTerminalCount": total,
        "databaseCurrentSchemaCount": current_schema_count,
        "databaseStatusCounts": counts,
        "recordSchemaVersion": RECORD_SCHEMA_VERSION,
        "recordedWebFields": list(RECORDED_WEB_FIELDS),
        "fullSourceDataIncluded": True,
        "cutoffRule": "activation_date_plus_29_days_inclusive_30_day_period",
        "rows": terminal_rows,
        "transientErrors": transient_rows,
    }
    atomic_json(receipt_path, receipt)
    return receipt


def retract_lookup(database: Path, lookup_path: Path, receipt_path: Path) -> dict[str, Any]:
    """Remove only terminal rows whose stored receipt exactly matches one bad lookup receipt."""
    payload = load_json(lookup_path)
    rows = payload.get("results") or payload.get("rows") or []
    if not isinstance(rows, list):
        raise RuntimeError("任务F撤销回执缺少 results 数组")
    candidates: list[tuple[str, str, str]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        query_iccid = normalize_iccid(row.get("queryIccid") or row.get("iccid"))
        status = str(row.get("status") or "").strip()
        if query_iccid and status in TERMINAL_STATUSES:
            candidates.append((query_iccid, status, json.dumps(row, ensure_ascii=False, separators=(",", ":"))))

    deleted: list[str] = []
    absent: list[str] = []
    mismatches: list[dict[str, Any]] = []
    with open_database(database) as connection:
        for query_iccid, status, expected_receipt in candidates:
            existing = connection.execute(
                "SELECT lookup_status, receipt_json FROM kooda_activation WHERE query_iccid = ?",
                (query_iccid,),
            ).fetchone()
            if not existing:
                absent.append(query_iccid)
                continue
            if existing["lookup_status"] != status or existing["receipt_json"] != expected_receipt:
                mismatches.append({
                    "queryIccid": query_iccid,
                    "lookupStatus": status,
                    "databaseStatus": existing["lookup_status"],
                    "reason": "stored_terminal_does_not_exactly_match_lookup_receipt",
                })
                continue
            connection.execute("DELETE FROM kooda_activation WHERE query_iccid = ?", (query_iccid,))
            deleted.append(query_iccid)
        if mismatches:
            connection.rollback()
            raise RuntimeError(f"任务F撤销发现{len(mismatches)}条数据库回执不一致，已回滚且未删除任何记录")
        connection.commit()
        remaining = int(connection.execute("SELECT COUNT(*) FROM kooda_activation").fetchone()[0])

    receipt = {
        "status": "passed",
        "createdAt": now_iso(),
        "database": str(database.resolve()),
        "lookupReceipt": str(lookup_path.resolve()),
        "candidateCount": len(candidates),
        "deletedCount": len(deleted),
        "absentCount": len(absent),
        "remainingTerminalCount": remaining,
        "exactStoredReceiptMatchRequired": True,
        "deletedQueryIccids": deleted,
        "alreadyAbsentQueryIccids": absent,
    }
    atomic_json(receipt_path, receipt)
    return receipt


def database_stats(database: Path) -> dict[str, Any]:
    with open_database(database) as connection:
        inventory = int(connection.execute("SELECT COUNT(*) FROM kooda_inventory").fetchone()[0])
        terminal = int(connection.execute("SELECT COUNT(*) FROM kooda_activation").fetchone()[0])
        found = int(connection.execute("SELECT COUNT(*) FROM kooda_activation WHERE lookup_status='found'").fetchone()[0])
        current_schema = int(connection.execute(
            "SELECT COUNT(*) FROM kooda_activation WHERE record_schema_version >= ?",
            (RECORD_SCHEMA_VERSION,),
        ).fetchone()[0])
        latest = connection.execute("SELECT MAX(queried_at) FROM kooda_activation").fetchone()[0]
        catalog = connection.execute(
            """
            SELECT source_sha256, source_path, imported_at, source_finished_at,
                   source_row_count, found_count, activation_missing_count
            FROM grabandgo_catalog_import
            ORDER BY imported_at DESC LIMIT 1
            """
        ).fetchone()
    return {
        "status": "passed",
        "database": str(database.resolve()),
        "inventoryCount": inventory,
        "terminalLookupCount": terminal,
        "foundActivationCount": found,
        "currentRecordSchemaCount": current_schema,
        "recordSchemaVersion": RECORD_SCHEMA_VERSION,
        "recordedWebFields": list(RECORDED_WEB_FIELDS),
        "latestQueriedAt": latest,
        "catalogImported": catalog is not None,
        "catalogSourceSha256": catalog["source_sha256"] if catalog else None,
        "catalogSourcePath": catalog["source_path"] if catalog else None,
        "catalogImportedAt": catalog["imported_at"] if catalog else None,
        "catalogSourceFinishedAt": catalog["source_finished_at"] if catalog else None,
        "catalogSourceRowCount": int(catalog["source_row_count"] or 0) if catalog else 0,
        "catalogFoundCount": int(catalog["found_count"] or 0) if catalog else 0,
        "catalogActivationMissingCount": int(catalog["activation_missing_count"] or 0) if catalog else 0,
    }


def export_catalog(database: Path, output: Path, latest: Path | None = None) -> dict[str, Any]:
    """Export the current durable Task F ICCID library as auditable JSON."""
    with open_database(database) as connection:
        connection.execute('BEGIN')
        database_rows = connection.execute(
            """
            SELECT query_iccid, source_iccid, activation_date, cutoff_date,
                   lookup_status, queried_at, detail_sim_iccid, worker_id,
                   source_phone, source_slot, source_opr, source_carrier_name,
                   purchase_date, purchase_date_raw, act_request_date,
                   activation_date_raw, sim_number, sim_number_raw, sku, cost,
                   source_data_json, record_schema_version, data_source,
                   source_file, source_sha256, vendor_ordinal, balance_cad, plan_type
            FROM kooda_activation
            ORDER BY
                CASE WHEN vendor_ordinal IS NULL THEN 1 ELSE 0 END,
                vendor_ordinal,
                query_iccid
            """
        ).fetchall()
        renewals = read_renewals(connection)

    rows: list[dict[str, Any]] = []
    status_counts: dict[str, int] = {}
    for record in database_rows:
        try:
            source_data = json.loads(record["source_data_json"] or "[]")
        except json.JSONDecodeError:
            source_data = []
        status = str(record["lookup_status"] or "")
        status_counts[status] = status_counts.get(status, 0) + 1
        rows.append({
            "queryIccid": record["query_iccid"],
            "sourceIccid": record["source_iccid"],
            "status": status,
            "purchaseDate": record["purchase_date"],
            "purchaseDateRaw": record["purchase_date_raw"],
            "actRequestDate": record["act_request_date"],
            "activationDate": record["activation_date"],
            "activationDateRaw": record["activation_date_raw"],
            "cutoffDate": record["cutoff_date"],
            "effectiveCutoffDate": renewals.get(record["query_iccid"], {}).get("cutoffDate", record["cutoff_date"]),
            "renewal": renewals.get(record["query_iccid"]),
            "simNumber": record["sim_number"],
            "simNumberRaw": record["sim_number_raw"],
            "sku": record["sku"],
            "cost": record["cost"],
            "balanceCad": record["balance_cad"],
            "planType": record["plan_type"],
            "queriedAt": record["queried_at"],
            "detailSimIccid": record["detail_sim_iccid"],
            "workerId": record["worker_id"],
            "sourcePhone": record["source_phone"],
            "sourceSlot": record["source_slot"],
            "sourceOpr": record["source_opr"],
            "sourceCarrierName": record["source_carrier_name"],
            "dataSource": record["data_source"],
            "sourceFile": record["source_file"],
            "sourceSha256": record["source_sha256"],
            "vendorOrdinal": record["vendor_ordinal"],
            "recordSchemaVersion": int(record["record_schema_version"] or 0),
            "sourceData": source_data,
        })

    query_iccids = [str(row["queryIccid"] or "") for row in rows]
    payload = {
        "status": "passed",
        "createdAt": now_iso(),
        "database": str(database.resolve()),
        "library": str(output.resolve()),
        "latestLibrary": str(latest.resolve()) if latest else str(output.resolve()),
        "recordCount": len(rows),
        "uniqueNormalizedIccidCount": len(set(query_iccids)),
        "statusCounts": status_counts,
        "foundCount": int(status_counts.get("found", 0)),
        "activationMissingCount": int(status_counts.get("activation_missing", 0)),
        "notFoundCount": int(status_counts.get("not_found", 0)),
        "recordSchemaVersion": RECORD_SCHEMA_VERSION,
        "recordedWebFields": list(RECORDED_WEB_FIELDS),
        "includesCompleteSourceData": True,
        "trailingFNormalization": "strip_exactly_one_final_F_case_insensitive",
        "activationField": "lblActivateDate",
        "cutoffRule": "activation_date_plus_29_days_inclusive_30_day_period",
        "effectiveCutoffRule": "authorized_renewal_overlay_else_original_activation_plus_29",
        "renewalCount": len(renewals),
        "rows": rows,
    }
    if len(query_iccids) != len(set(query_iccids)) or any(not ICCID_PATTERN.fullmatch(value) for value in query_iccids):
        raise RuntimeError("任务F最新ICCID库存在无效或重复归一化ICCID，已拒绝导出")
    encoded = encode_json(payload)
    atomic_bytes(output, encoded)
    if latest and latest.resolve() != output.resolve():
        # Both destinations must contain the very same snapshot. Serialize it
        # once; keep atomic publication and fsync for each separate file.
        atomic_bytes(latest, encoded)
    return payload


def snapshot_database(database: Path, output: Path) -> dict[str, Any]:
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise RuntimeError(f"拒绝覆盖任务F数据库快照: {output}")
    source = sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True, timeout=30)
    source.execute("PRAGMA query_only=ON")
    try:
        target = sqlite3.connect(output)
        try:
            source.backup(target)
            target.commit()
        finally:
            target.close()
    finally:
        source.close()
    return {"status": "passed", "database": str(database.resolve()), "snapshot": str(output.resolve())}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="OpenSMS任务F Kooda激活日期数据库")
    subparsers = parser.add_subparsers(dest="command", required=True)
    for name in ("init", "stats"):
        command = subparsers.add_parser(name)
        command.add_argument("--db", type=Path, required=True)
    pending = subparsers.add_parser("pending")
    pending.add_argument("--db", type=Path, required=True)
    pending.add_argument("--cards", type=Path, required=True)
    pending.add_argument("--output", type=Path, required=True)
    pending.add_argument("--document", type=Path)
    pending.add_argument("--document-sha256")
    pending.add_argument("--document-proof", type=Path)
    pending.add_argument("--document-proof-sha256")
    pending.add_argument("--ko-cache", type=Path)
    pending.add_argument("--ko-cache-sha256")
    pending.add_argument("--run-id")
    importer = subparsers.add_parser("import")
    importer.add_argument("--db", type=Path, required=True)
    importer.add_argument("--lookup", type=Path, required=True)
    importer.add_argument("--receipt", type=Path, required=True)
    catalog = subparsers.add_parser("import-catalog")
    catalog.add_argument("--db", type=Path, required=True)
    catalog.add_argument("--catalog", type=Path, required=True)
    catalog.add_argument("--receipt", type=Path, required=True)
    catalog_export = subparsers.add_parser("export-catalog")
    catalog_export.add_argument("--db", type=Path, required=True)
    catalog_export.add_argument("--output", type=Path, required=True)
    catalog_export.add_argument("--latest", type=Path)
    retract = subparsers.add_parser("retract")
    retract.add_argument("--db", type=Path, required=True)
    retract.add_argument("--lookup", type=Path, required=True)
    retract.add_argument("--receipt", type=Path, required=True)
    snapshot = subparsers.add_parser("snapshot")
    snapshot.add_argument("--db", type=Path, required=True)
    snapshot.add_argument("--output", type=Path, required=True)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if args.command == "init":
        with open_database(args.db):
            pass
        result = database_stats(args.db)
    elif args.command == "stats":
        result = database_stats(args.db)
    elif args.command == "pending":
        result = prepare_pending(args.db, args.cards, args.output, document=args.document,
                                 document_sha256=args.document_sha256, run_id=args.run_id,
                                 document_proof=args.document_proof,
                                 document_proof_sha256=args.document_proof_sha256,
                                 ko_cache=args.ko_cache, ko_cache_sha256=args.ko_cache_sha256)
    elif args.command == "import":
        result = import_lookup(args.db, args.lookup, args.receipt)
    elif args.command == "import-catalog":
        result = import_catalog(args.db, args.catalog, args.receipt)
    elif args.command == "export-catalog":
        result = export_catalog(args.db, args.output, args.latest)
    elif args.command == "retract":
        result = retract_lookup(args.db, args.lookup, args.receipt)
    elif args.command == "snapshot":
        result = snapshot_database(args.db, args.output)
    else:
        raise AssertionError(args.command)
    stdout_result = result
    if args.command in {"pending", "import", "export-catalog"}:
        # Durable receipts retain full row/source evidence. The runner consumes
        # those files; avoid printing thousands of already-known ICCIDs again.
        row_fields = {"rows", "alreadyQueried", "invalidRows", "transientErrors",
                      "reportEligibleRecords", "documentExcludedRows", "reuseRecords"}
        stdout_result = {key: value for key, value in result.items() if key not in row_fields}
        stdout_result["rowsOmittedFromStdout"] = True
        stdout_result["detailPath"] = str((args.receipt if args.command == "import" else args.output).resolve())
    print(json.dumps(stdout_result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
