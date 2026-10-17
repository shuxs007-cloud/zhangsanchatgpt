#!/usr/bin/env python3
"""Offline per-device ten-day summary and daily metrics in Task A column order."""
import argparse
from collections import defaultdict
from datetime import date, timedelta
import json
from pathlib import Path
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.opensms_device_overall_report import aggregate
from services.opensms_document_api.pipeline.device_daily_statistics import aggregate_device_daily, FIELDS, device_key


def carrier_counts(headers, rows):
    if headers.count('opr') != 1 or headers.count('卡槽') != 1:
        raise ValueError('卡数统计需要唯一opr与卡槽列')
    opr_index, slot_index = headers.index('opr'), headers.index('卡槽')
    counts = defaultdict(lambda: [0, 0])
    for row in rows:
        opr = str(row[opr_index] or '').casefold()
        kooda = any(alias in opr for alias in ('kooda', 'koodo', 'telus1'))
        chatr = 'chatr' in opr
        if kooda and chatr:
            raise ValueError('OPR同时命中两种卡类型，不能重复计数')
        counts[device_key(row[slot_index])][0] += int(kooda)
        counts[device_key(row[slot_index])][1] += int(chatr)
    return counts


def prepare(source, server_date):
    overall = aggregate(source)
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        iterator = iter(book['sheet1'].values)
        headers = next(iterator)
        source_rows = list(iterator)
        counts = carrier_counts(headers, source_rows)
        daily = aggregate_device_daily(headers, source_rows, workers=4)
    finally:
        book.close()
    dates = [(server_date-timedelta(days=i)).strftime('%m-%d') for i in range(10)]
    if dates[0] not in daily['receipt']['dates']:
        raise ValueError('服务器日期没有对应明细，拒绝将旧日期称为今天')
    index = {header: i for i, header in enumerate(daily['headers'])}
    by_device = {row[0]: row for row in daily['rows']}
    out_headers = ['设备号','机房','号码记录数','Kooda卡数','Chatr卡数','今天回复率','昨天回复率',
                   '10日发送窗口数','10日回复窗口数','10日回复率','10日总发送量','10日总回复量']
    out_headers += [day+field for day in dates for field in FIELDS]
    result = []
    for row in overall:
        values = by_device[row[0]]
        def metric(day, field):
            pos = index.get(day+field)
            return values[pos] if pos is not None else None
        result.append([*row[:3], *counts[row[0]], metric(dates[0],'回复率'),metric(dates[1],'回复率'),
                       row[3],row[4],row[5],row[6],row[7],
                       *(metric(day,field) for day in dates for field in FIELDS)])
    return out_headers, result, [d for d in dates if d not in daily['receipt']['dates']]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--server-date',type=date.fromisoformat,required=True)
    parser.add_argument('--verify-only',action='store_true')
    args = parser.parse_args()
    source,target = args.input.resolve(),args.output.resolve()
    headers,rows,missing = prepare(source,args.server_date)
    if args.verify_only:
        wb = load_workbook(target,read_only=True,data_only=True)
        try:
            actual = list(wb['设备回复率统计'].values)
            assert list(actual[0]) == headers and len(actual)-1 == len(rows)
            for saved,wanted in zip(actual[1:],rows):
                for a,b in zip(saved,wanted):
                    if b is None:
                        assert a in (None,''), (a,b)
                    elif isinstance(b,float):
                        assert abs(a-b)<1e-12,(a,b)
                    else:
                        assert a == b,(a,b)
        finally:
            wb.close()
        print(json.dumps({'verification':'passed','devices':len(rows),'missing_daily_dates':missing},ensure_ascii=False))
        return
    if target.exists() or target == source:
        raise ValueError('禁止覆盖源或已有文件')
    wb = Workbook()
    ws = wb.active
    ws.title = '设备回复率统计'
    ws.append(headers)
    for row in rows:
        ws.append(row)
    last_row = len(rows)+1
    for r in range(2,last_row+1):
        for c in [10,*range(15,63,5)]:
            send,reply = get_column_letter(c-2),get_column_letter(c-1)
            ws.cell(r,c,f'=IF(OR({send}{r}="",{reply}{r}="",{send}{r}=0),"",{reply}{r}/{send}{r})')
        ws.cell(r,6,f'=IF(O{r}="","",O{r})')
        ws.cell(r,7,f'=IF(T{r}="","",T{r})')
    ws.freeze_panes = 'H2'
    ws.auto_filter.ref = ws.dimensions
    for c,header in enumerate(headers,1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = 20 if c>3 else (15 if c!=2 else 22)
        cell=ws.cell(1,c)
        cell.font=Font(name='Arial',bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='216B78' if c<=12 else ('336699' if (c-13)//5%2==0 else '487B60'))
        for r in range(2,last_row+1):
            cell=ws.cell(r,c)
            cell.font=Font(name='Arial',size=11)
            if '回复率' in header:
                cell.number_format='0.000%'
            elif c>2:
                cell.number_format='#,##0'
            if any(header.startswith(d) for d in missing):
                cell.fill=PatternFill('solid',fgColor='FFF2CC')
    notes=wb.create_sheet('统计口径')
    for text in [f'源文件：{source}',f'报表服务器日期（今天）：{args.server_date}',
                 f'昨天：{args.server_date-timedelta(days=1)}',
                 '每设备一行，包含主表全部号码记录及状态；字母设备前缀不合并到数字设备。',
                 'Kooda卡数按OPR中的Kooda/Koodo/Telus1归类；Chatr卡数按OPR中的Chatr归类，大小写不敏感。',
                 '卡数口径与号码记录数一致：按源记录统计、不去重，不代表实时物理插卡数；其他或空白OPR不计入这两列。',
                 '近10日总计直接汇总主表10日字段；逐日字段按源日期分别汇总。',
                 '回复率=同设备同期间回复窗口数合计/发送窗口数合计，不平均百分比。',
                 '总发送量、总回复量独立求和，不混用作窗口回复率分子分母。',
                 '零分母或缺失计数的回复率留空；源缺失日期整组浅黄留空，不能当0。',
                 '源表缺失的每日明细日期：'+('、'.join(missing) or '无'),
                 '本文件是已有报表离线统计，不是当前新鲜导出；今天是源报表服务器日期。']:
        notes.append([text])
    notes.column_dimensions['A'].width=130
    target.parent.mkdir(parents=True,exist_ok=True)
    # Persist formulas with independently computed cached results so non-Excel
    # readers display values immediately, without requiring a desktop app.
    import xlsxwriter
    out = xlsxwriter.Workbook(target)
    for original in wb:
        dest = out.add_worksheet(original.title)
        if original.title == ws.title:
            dest.freeze_panes(1,7)
            dest.autofilter(0,0,last_row-1,len(headers)-1)
        for letter,dimension in original.column_dimensions.items():
            dest.set_column(f'{letter}:{letter}',dimension.width)
        formats = {}
        for row in original:
            for cell in row:
                if cell.value is None:
                    continue
                props = {'font_name':'Arial','font_size':11}
                if cell.row == 1 and original.title == ws.title:
                    props.update(bold=True,font_color='#FFFFFF',bg_color='#'+cell.fill.fgColor.rgb[-6:])
                elif cell.fill.patternType == 'solid':
                    props['bg_color']='#'+cell.fill.fgColor.rgb[-6:]
                if cell.number_format != 'General':
                    props['num_format']=cell.number_format
                key = tuple(sorted(props.items()))
                fmt = formats.setdefault(key,out.add_format(props)) if key not in formats else formats[key]
                if cell.data_type == 'f':
                    cached=rows[cell.row-2][cell.column-1]
                    dest.write_formula(cell.row-1,cell.column-1,cell.value,fmt,'' if cached is None else cached)
                else:
                    dest.write(cell.row-1,cell.column-1,cell.value,fmt)
    out.close()
    wb.close()
    print(json.dumps({'output':str(target),'devices':len(rows),'missing_daily_dates':missing,'formula_caches':'computed; verify-only required'},ensure_ascii=False))


if __name__ == '__main__':
    main()
