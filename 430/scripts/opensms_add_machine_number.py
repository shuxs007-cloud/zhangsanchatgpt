#!/usr/bin/env python3
"""Offline addition; preserves historical input and verifies all original cells."""
import argparse
import json
from pathlib import Path
import sys
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from services.opensms_document_api.pipeline.machine_number import add_machine_number, verify_machine_number

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    a=p.parse_args()
    if a.output.exists() or a.input.resolve()==a.output.resolve():
        raise ValueError('禁止覆盖历史文件')
    w=load_workbook(a.input)
    receipt=add_machine_number(w['sheet1'])
    a.output.parent.mkdir(parents=True,exist_ok=True)
    w.save(a.output);w.close()
    before=load_workbook(a.input,read_only=True,data_only=False)
    after=load_workbook(a.output,read_only=True,data_only=False)
    assert before.sheetnames==after.sheetnames
    for name in before.sheetnames:
        old=list(before[name].values); new=list(after[name].values)
        if name=='sheet1':
            headers=list(new[0]);verify_machine_number(headers,new[1:])
            pos=headers.index('机器号码')
            new=[r[:pos]+r[pos+1:] for r in new]
        assert old==new, name+'原有内容发生变化'
    before.close();after.close()
    print(json.dumps({'output':str(a.output.resolve()),'verification':'passed',**receipt},ensure_ascii=False))

if __name__=='__main__':
    main()
