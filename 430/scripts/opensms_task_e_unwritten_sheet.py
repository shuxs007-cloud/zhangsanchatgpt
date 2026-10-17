"""Local-only Task E exception tab, bound to this run's confirmed transactions.

The small new tab is authored with Artifact Tool. Existing worksheet, formula,
style, drawing and relationship parts are copied byte-for-byte, except the
three package metadata members needed to register the new tab.
"""
from __future__ import annotations

import argparse
from datetime import date
import json
from pathlib import Path
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile

from opensms_refresh_workbook_carriers_fast import (
    MAIN_NS, add_sheet_metadata, atomic_json, cell_value, copy_with_replacements,
    parse_sheet_rows, sha256_file, shared_strings, workbook_sheets,
)

SHEET = '任务E未写入'
HEADERS = ('手机号', '截止日', '未写入原因')
REASON = '后台未命中（未写入）'
RUNTIME = Path('/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies')


def require(condition, message):
    if not condition:
        raise ValueError(f'任务E未写入表：{message}')


def confirmed_rows(source: Path, manifest_path: Path, receipt_path: Path, run_id: str):
    require(manifest_path.name == f'opensms_task_e_{run_id}_manifest.json'
            and receipt_path.name == f'opensms_task_e_{run_id}_receipt.json', '回执不是本轮')
    manifest = json.loads(manifest_path.read_text())
    receipt = json.loads(receipt_path.read_text())
    digest = sha256_file(source)
    for item in (manifest, receipt):
        require(item.get('task') == 'E' and Path(item.get('sourcePath', '')).resolve() == source.resolve()
                and item.get('sourceSha256') == digest, '源文件或SHA256不匹配')
    require(Path(receipt.get('manifestPath', '')).resolve() == manifest_path.resolve(), '候选清单绑定不符')
    require(manifest.get('forceGate') is True and manifest.get('automationGate') is True, '缺少本轮E授权')
    require(receipt.get('status') in {'passed', 'passed_with_not_found', 'no_candidate'}
            and all(receipt.get(key) is True for key in (
                'businessReceiptPassed', 'transactionalVerificationPassed', 'ledgerVerificationPassed')),
            '事务或台账尚未确认，不能将不确定记录归为未写入')
    owners = {}
    for group in manifest.get('groups') or []:
        deadline = group.get('expireDate')
        require(isinstance(deadline, str) and date.fromisoformat(deadline).isoformat() == deadline, '截止日无效')
        for phone in group.get('cardNos') or []:
            require(isinstance(phone, str) and re.fullmatch(r'\d{10}', phone) and phone not in owners, '候选号码无效或重复')
            owners[phone] = deadline
    seen, rows = set(), []
    for batch in receipt.get('batches') or []:
        phones, missing = batch.get('cardNos') or [], batch.get('notFound')
        require(isinstance(missing, list), '未命中列表缺失')
        require(len(set(phones)) == len(phones) and not (seen & set(phones)), '批次号码重复')
        seen.update(phones)
        require(len(set(missing)) == len(missing) and set(missing) <= set(phones), '未命中号码不属于提交批次')
        require(all(owners.get(phone) == batch.get('expireDate') for phone in phones), '批次截止日与候选不符')
        response = batch.get('serverResponse') or {}
        raw = response.get('data') or {}
        require(response.get('code') == 200 and isinstance(raw, dict), '缺少成功同步事务响应')
        for key, expected in (('inputCount', len(phones)), ('matchedCardCount', len(phones) - len(missing))):
            require(type(raw.get(key)) is int and raw[key] == expected and batch.get(key) == expected, '输入/命中计数不符')
        require(raw.get('notFound') == missing and type(raw.get('updatedChatCount')) is int
                and raw['updatedChatCount'] >= 0, '原始响应与未命中明细不符')
        rows.extend(dict(phone=phone, expireDate=batch['expireDate'], reason='backend_not_found') for phone in missing)
    key = lambda row: (row['expireDate'], row['phone'], row['reason'])
    require(seen == set(owners) and len(owners) == manifest.get('candidateCount'), '批次未完整覆盖候选')
    require(receipt.get('inputCount') == len(owners) and receipt.get('notWrittenCount') == len(rows)
            and receipt.get('submittedCount') == len(owners) - len(rows), '汇总计数不符')
    require(sorted(rows, key=key) == sorted(receipt.get('notWrittenCards') or [], key=key), '汇总未写入明细不符')
    return sorted(rows, key=key)


def verify_sheet(workbook: Path, expected_rows):
    # Independent reader; never authors or rewrites an existing workbook.
    import openpyxl
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    try:
        require(SHEET in book.sheetnames, '最终合并表缺少独立sheet')
        cells = list(book[SHEET].iter_rows())
        actual = [tuple(cell.value for cell in row) for row in cells]
        expected = [HEADERS] + [(r['phone'], r['expireDate'], REASON) for r in expected_rows]
        require(actual == expected, '独立回读手机号、截止日、原因或数量不一致')
        require(all(row[0].data_type == 's' for row in cells[1:]), '手机号不是文本')
        require(not any(cell.data_type in {'e', 'f'} for row in cells for cell in row), '异常sheet出现公式或错误单元格')
    finally:
        book.close()


def graft_sheet(source: Path, thin: Path, output: Path, rows):
    require(source.resolve() != output.resolve(), '不得覆盖源工作簿')
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(thin) as authored:
        sheets = workbook_sheets(original)
        require(all(name != SHEET for name, _ in sheets), '拒绝混入旧未写入sheet')
        # Reuse the existing report's header/body styles, retaining ISO text dates.
        source_rows, _ = parse_sheet_rows(original.read(sheets[0][1]), shared_strings(original))
        header = source_rows.get(1, {})
        phone_col = next((col for col, (_, value, _) in header.items() if value in {'电话', '手机号'}), None)
        header_style = header.get(phone_col, ('', '', '0'))[2]
        body = next((source_rows[n] for n in sorted(source_rows) if n > 1), {})
        body_style = body.get(phone_col, ('', '', '0'))[2]
        strings = shared_strings(authored)
        xml = ET.fromstring(authored.read(workbook_sheets(authored)[0][1]))
        for cell in xml.iter(f'{{{MAIN_NS}}}c'):
            value = cell_value(cell, strings)
            reference = cell.get('r')
            cell.clear()
            cell.attrib.update(r=reference, t='inlineStr', s=header_style if re.fullmatch(r'[A-Z]+1', reference) else body_style)
            ET.SubElement(ET.SubElement(cell, f'{{{MAIN_NS}}}is'), f'{{{MAIN_NS}}}t').text = value
        for row in xml.iter(f'{{{MAIN_NS}}}row'):
            row.attrib.pop('s', None)
        for col in xml.iter(f'{{{MAIN_NS}}}col'):
            col.attrib.pop('style', None)
        replacements, member = add_sheet_metadata(original, original.namelist(), SHEET)
        sheet_xml = ET.tostring(xml, encoding='utf-8', xml_declaration=True)
        copy_with_replacements(source, output, replacements, {member: sheet_xml})
        with zipfile.ZipFile(output) as result:
            require(set(result.namelist()) == set(original.namelist()) | {member}, '工作簿成员变化越界')
            require(all(original.read(name) == result.read(name) for name in original.namelist()
                        if name not in replacements), '无关工作簿内容发生变化')
    verify_sheet(output, rows)


def run(args):
    rows = confirmed_rows(args.input, args.manifest, args.receipt, args.run_id)
    if not rows:
        return {'passed': True, 'rowCount': 0, 'sheet': None, 'output': str(args.input)}
    args.runtime.mkdir(parents=True, exist_ok=True)
    builder = args.runtime / 'opensms_task_e_unwritten_sheet.mjs'
    shutil.copy2(Path(__file__).with_suffix('.mjs'), builder)
    modules = args.runtime / 'node_modules'
    if not modules.exists():
        modules.symlink_to(RUNTIME / 'node/node_modules', target_is_directory=True)
    payload, thin = args.runtime / 'rows.json', args.runtime / 'sheet.xlsx'
    atomic_json(payload, rows)
    subprocess.run([str(RUNTIME / 'node/bin/node'), str(builder), str(payload), str(thin)], check=True, capture_output=True, text=True)
    graft_sheet(args.input, thin, args.output, rows)
    result = {'passed': True, 'runId': args.run_id, 'sheet': SHEET, 'rowCount': len(rows),
              'output': str(args.output), 'outputSha256': sha256_file(args.output),
              'sourceSha256': sha256_file(args.input), 'manifestSha256': sha256_file(args.manifest),
              'receiptSha256': sha256_file(args.receipt), 'independentReadbackPassed': True,
              'unrelatedPartsUnchanged': True}
    atomic_json(args.verify_receipt, result)
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ('input', 'manifest', 'receipt', 'output', 'runtime', 'verify-receipt'):
        parser.add_argument('--' + name, type=Path, required=True)
    parser.add_argument('--run-id', required=True)
    print(json.dumps(run(parser.parse_args()), ensure_ascii=False))
