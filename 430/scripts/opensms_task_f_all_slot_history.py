#!/usr/bin/env python3
"""Offline full Task F slot history: read-only sources, bounded parallel scanning.

Full row evidence is stored in per-content gzip JSONL files; Excel is a compact
card/slot index. No network, no production writes, no inferred current location.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict, OrderedDict
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from datetime import date, datetime
import gzip
import hashlib
import io
import json
import os
from pathlib import Path
import posixpath
import re
import sqlite3
import subprocess
import time
import zipfile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import opensms_task_f_slot_history as base

TARGETS = set()
CACHE = OrderedDict()
OUTPUT = None
WINDOW = None


def frozen_targets(path):
    result = {}
    with sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True) as conn:
        conn.execute('PRAGMA query_only=ON')
        conn.execute('BEGIN')
        for key, slot, status, timestamp in conn.execute(
                'SELECT query_iccid,source_slot,lookup_status,queried_at FROM kooda_activation'):
            if base.normalize_iccid(key) != key:
                raise ValueError('Invalid production ICCID')
            result[key] = dict(activationSlot=slot or '', lookupStatus=status or '',
                               queriedAt=timestamp or '', inventorySlot='', inventorySeenAt='', origin='当前任务F库')
        for key, slot, timestamp in conn.execute('SELECT query_iccid,slot,last_seen_at FROM kooda_inventory'):
            if base.normalize_iccid(key) != key:
                raise ValueError('Invalid inventory ICCID')
            result.setdefault(key, dict(activationSlot='', lookupStatus='仅库存', queriedAt='', origin='当前任务F库'))
            result[key].update(inventorySlot=slot or '', inventorySeenAt=timestamp or '')
    return result


def discover_json(roots, start, end):
    existing = [str(Path(p).resolve()) for p in roots if Path(p).is_dir()]
    proc = subprocess.run(['rg', '--files', '--hidden', '--no-ignore', *existing,
                           '-g', '*inventory*.json', '-g', '*lookup*.json', '-g', '!._*'],
                          capture_output=True, text=True)
    if proc.returncode not in (0, 1):
        raise RuntimeError('JSON discovery failed: ' + proc.stderr[:1000])
    selected = []
    excluded = Counter()
    for name in sorted(set(proc.stdout.splitlines())):
        p = Path(name)
        if set(p.parts) & base.EXCLUDED_PARTS or p.name.endswith('.pagination.json'):
            excluded['excluded'] += 1
            continue
        if not ('task_f' in p.name or 'system_kooda_inventory' in p.name):
            excluded['notTaskF'] += 1
            continue
        st = p.stat()
        day, basis = base.file_date(p, st)
        if not start <= day <= end:
            excluded['outsideWindow'] += 1
            continue
        selected.append(dict(path=str(p), date=day, dateBasis=basis, size=st.st_size,
                             mtimeNs=st.st_mtime_ns, format='json'))
    return selected, dict(discovered=len(set(proc.stdout.splitlines())), excluded=dict(excluded))


def init_worker(targets, output, start, end):
    global TARGETS, OUTPUT, WINDOW, CACHE
    TARGETS, OUTPUT, WINDOW = set(targets), Path(output), (start, end)
    CACHE = OrderedDict()


def fast_sheet(xml, strings, shared_hash):
    """Decode only the explicitly named identity/slot columns, never numeric ICCIDs."""
    key = (shared_hash, hashlib.sha256(xml).hexdigest())
    if key in CACHE:
        CACHE.move_to_end(key)
        return CACHE[key]
    namespaces = next(etree.iterparse(io.BytesIO(xml), events=('start',),
                                     resolve_entities=False, no_network=True))[1].nsmap
    ids, slots, header_row = [], [], 0
    for fragment in base.list_first_rows(xml, 50):
        rownum, values = base.cell_values(fragment, strings, namespaces)
        a = [col for col, v in values.items() if base.header(v[0]) in base.ICCID_HEADERS]
        b = [col for col, v in values.items() if base.header(v[0]) in base.SLOT_HEADERS]
        if a and b:
            ids, slots, header_row = a, b, rownum
            break
    if not header_row:
        return [], [], '无同表明确ICCID及卡槽列（不猜测关联）'
    # A cell must carry a selected exact column reference. Never scan across >.
    cols = b'|'.join(re.escape(c.encode()) for c in ids + slots)
    cell_pattern = re.compile(rb'<(?:[\w.-]+:)?c\b(?=[^>]*\br=["\'](?:' + cols
                              + rb')[0-9]+["\'])[^>]*(?:/>|>.*?</(?:[\w.-]+:)?c\s*>)', re.S)
    evidence, unresolved = [], []
    for match in base.ROW.finditer(xml):
        selected = b''.join(m.group() for m in cell_pattern.finditer(match.group()))
        if not selected:
            continue
        rownum_match = re.search(rb'\br=["\']([0-9]+)["\']', match.group().split(b'>', 1)[0])
        rownum = int(rownum_match.group(1)) if rownum_match else 0
        if rownum <= header_row:
            continue
        _, values = base.cell_values(b'<row r="' + str(rownum).encode() + b'">' + selected + b'</row>',
                                     strings, namespaces)
        for col in ids:
            value, kind, formula, ref = values.get(col, ('', 'n', False, ''))
            iccid = base.normalize_iccid(value)
            if iccid not in TARGETS:
                continue
            row = dict(iccid=iccid, row=rownum, iccidCell=ref)
            if kind not in ('s', 'str', 'inlineStr') or formula:
                unresolved.append(dict(row, reason='ICCID非可靠文本或公式'))
                continue
            valid_slots = []
            for slot_col in slots:
                sv, sk, sf, sr = values.get(slot_col, ('', 'n', False, ''))
                if not sf and sk in ('s', 'str', 'inlineStr') and base.SLOT.fullmatch(sv.strip()):
                    valid_slots.append((sv.strip(), sr))
            if not valid_slots:
                unresolved.append(dict(row, reason='卡槽空白或不是完整卡槽'))
            for slot, ref in valid_slots:
                evidence.append(dict(row, slot=slot, slotCell=ref,
                                     ambiguousColumns=len(ids) != 1 or len(slots) != 1))
    result = (evidence, unresolved, '')
    CACHE[key] = result
    # Bounded cache: two large sheets per worker, not the entire history in RAM.
    while len(CACHE) > 2:
        CACHE.popitem(last=False)
    return result


def json_records(document):
    """Known Task F schemas only; no cookie/login/credential or arbitrary object walk."""
    if not isinstance(document, dict):
        raise ValueError('回执不是对象')
    arrays = [key for key in ('rows', 'results') if isinstance(document.get(key), list)]
    if not arrays:
        raise ValueError('未识别任务F rows/results结构')
    for name in arrays:
        for index, obj in enumerate(document[name]):
            if not isinstance(obj, dict):
                continue
            yield f'{name}[{index}]', obj
            if isinstance(obj.get('sourceData'), dict):
                yield f'{name}[{index}].sourceData', obj['sourceData']
            if isinstance(obj.get('sourceOccurrences'), list):
                for sub, occ in enumerate(obj['sourceOccurrences']):
                    if isinstance(occ, dict):
                        yield f'{name}[{index}].sourceOccurrences[{sub}]', occ


def json_identity(obj):
    values = [obj.get(k) for k in ('queryIccid', 'sourceIccid', 'cardStore') if obj.get(k)]
    keys = [base.normalize_iccid(v) for v in values]
    if not keys or any(k is None for k in keys) or len(set(keys)) != 1:
        return None
    return keys[0]


def scan_file(item):
    digest = item['sha256']
    detail_path = OUTPUT / '逐条证据' / (digest + '.jsonl.gz')
    aggregate = {}
    unresolved_counts = Counter()
    discovered = set()
    sheets = []
    evidence_count, unresolved_count = 0, 0
    result = dict(sha256=digest, format=item['format'])
    def emit(stream, row, sheet='', locator='', kind=''):
        nonlocal evidence_count, unresolved_count
        row = dict(row, sheet=sheet, locator=locator)
        stream.write(json.dumps(row, ensure_ascii=False, separators=(',', ':')) + '\n')
        if 'slot' not in row:
            unresolved_count += 1
            unresolved_counts[row['reason']] += 1
            return
        evidence_count += 1
        k = (row['iccid'], row['slot'], sheet, bool(row.get('ambiguousColumns')), kind)
        if k not in aggregate:
            aggregate[k] = dict(count=0, example=locator)
        aggregate[k]['count'] += 1
    try:
        with gzip.open(detail_path, 'xt', encoding='utf-8', compresslevel=1) as stream:
            stream.write(json.dumps(dict(recordType='source', sha256=digest,
                                         sourceFile=item['path'], format=item['format']), ensure_ascii=False) + '\n')
            if item['format'] == 'json':
                doc = json.loads(Path(item['path']).read_text(encoding='utf-8'))
                status = doc.get('status', '')
                if status not in ('passed', 'success', 'ok', 'completed', 'passed_with_transient_errors'):
                    raise ValueError('未通过的文件级状态: ' + str(status))
                kind = '任务F查询回执原始来源' if 'results' in doc else '系统任务F库存原始记录'
                for locator, obj in json_records(doc):
                    iccid = json_identity(obj)
                    if not iccid:
                        continue
                    discovered.add(iccid)
                    sv = obj.get('slot') or obj.get('slotIdentif') or ''
                    if not isinstance(sv, str) or not base.SLOT.fullmatch(sv.strip()):
                        emit(stream, dict(iccid=iccid, reason='原始记录卡槽为空或格式不完整'), locator=locator)
                    else:
                        emit(stream, dict(iccid=iccid, slot=sv.strip(), ambiguousColumns=False,
                                          observedAt=doc.get('createdAt', ''), lookupStatus=obj.get('status', ''),
                                          sourceKind=kind), locator=locator, kind=kind)
                result['documentCreatedAt'] = doc.get('createdAt', '')
            else:
                if not zipfile.is_zipfile(item['path']):
                    raise ValueError('不是OOXML工作簿，可能为下载响应或不支持的旧二进制格式')
                with zipfile.ZipFile(item['path']) as book:
                    parser = etree.XMLParser(resolve_entities=False, no_network=True)
                    wb = etree.fromstring(book.read('xl/workbook.xml'), parser)
                    rels = etree.fromstring(book.read('xl/_rels/workbook.xml.rels'), parser)
                    members = {r.get('Id'): (posixpath.normpath('xl/' + r.get('Target', ''))
                               if not r.get('Target', '').startswith('/') else r.get('Target').lstrip('/'))
                               for r in rels if r.get('TargetMode') != 'External'}
                    raw = book.read('xl/sharedStrings.xml') if 'xl/sharedStrings.xml' in book.namelist() else b''
                    strings = [base.rich_text(si) for si in etree.fromstring(raw, parser)] if raw else []
                    shash = hashlib.sha256(raw).hexdigest()
                    sheets_xml = wb.findall('.//{' + base.NS + '}sheet')
                    if not sheets_xml:
                        raise ValueError('没有可识别的标准工作表')
                    for sheet in sheets_xml:
                        rid = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                        ev, un, reason = fast_sheet(book.read(members[rid]), strings, shash)
                        name = sheet.get('name')
                        sheets.append(dict(name=name, evidenceRows=len(ev), unresolvedRows=len(un), skipReason=reason))
                        for row in ev:
                            emit(stream, row, name, f"{row['iccidCell']} -> {row['slotCell']}")
                        for row in un:
                            emit(stream, row, name, row['iccidCell'])
            st = Path(item['path']).stat()
            if (st.st_size, st.st_mtime_ns) != (item['size'], item['mtimeNs']):
                raise ValueError('扫描中源文件变化，不接受部分证据')
    except Exception as exc:
        result['error'] = type(exc).__name__ + ': ' + str(exc)
        aggregate = {}
        evidence_count = 0
    result.update(evidenceRows=evidence_count, unresolvedRows=unresolved_count,
                  unresolvedReasons=dict(unresolved_counts), discoveredIccids=sorted(discovered),
                  detailFile=str(detail_path), sheets=sheets,
                  aggregates=[dict(iccid=k[0], slot=k[1], sheet=k[2], ambiguous=k[3], sourceKind=k[4], **v)
                              for k, v in aggregate.items()])
    return result


def make_audit_db(path):
    conn = sqlite3.connect(path)
    conn.execute('PRAGMA journal_mode=WAL')
    conn.execute('CREATE TABLE slots (iccid TEXT, slot TEXT, kind TEXT, ambiguous INTEGER, '
                 'observations INTEGER, contents INTEGER, first_date TEXT, last_date TEXT, '
                 'example_file TEXT, digest TEXT, sheet TEXT, locator TEXT, '
                 'PRIMARY KEY(iccid,slot,kind,ambiguous))')
    conn.execute('CREATE TABLE files (path TEXT PRIMARY KEY, date TEXT, date_basis TEXT, digest TEXT, '
                 'format TEXT, status TEXT, error TEXT, details TEXT, evidence_rows INTEGER)')
    return conn


def record_result(conn, result, aliases):
    # The full gzip ledger keeps every row; only the summary is reduced here.
    conn.executemany('INSERT INTO files VALUES (?,?,?,?,?,?,?,?,?)', [
        (x['path'], x['date'], x['dateBasis'], x['sha256'], x['format'],
         'failed' if result.get('error') else 'scanned', result.get('error', ''),
         result['detailFile'], result['evidenceRows']) for x in aliases])
    if result.get('error'):
        conn.commit()
        return
    grouped = {}
    for row in result['aggregates']:
        kinds = defaultdict(list)
        for alias in aliases:
            kind = row['sourceKind'] or base.source_kind(alias['path'], row['sheet'])
            kinds[kind].append(alias)
        for kind, group in kinds.items():
            latest = max(group, key=lambda x: (x['date'], x['path']))
            k = (row['iccid'], row['slot'], kind, row['ambiguous'])
            if k not in grouped:
                grouped[k] = [0, min(x['date'] for x in group), latest['date'], latest['path'], row['sheet'], row['example']]
            grouped[k][0] += row['count']
    values = [(*key, val[0], 1, val[1], val[2], val[3], result['sha256'], val[4], val[5])
              for key, val in grouped.items()]
    conn.executemany('''INSERT INTO slots VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
      ON CONFLICT(iccid,slot,kind,ambiguous) DO UPDATE SET
        observations=slots.observations+excluded.observations, contents=slots.contents+1,
        first_date=MIN(slots.first_date,excluded.first_date), last_date=MAX(slots.last_date,excluded.last_date),
        example_file=CASE WHEN excluded.last_date>=slots.last_date THEN excluded.example_file ELSE slots.example_file END,
        digest=CASE WHEN excluded.last_date>=slots.last_date THEN excluded.digest ELSE slots.digest END,
        sheet=CASE WHEN excluded.last_date>=slots.last_date THEN excluded.sheet ELSE slots.sheet END,
        locator=CASE WHEN excluded.last_date>=slots.last_date THEN excluded.locator ELSE slots.locator END''', values)
    conn.commit()


def card_summaries(conn, targets):
    stats = {}
    for iccid, slot, kind, ambiguous in conn.execute('SELECT iccid,slot,kind,ambiguous FROM slots ORDER BY iccid,slot'):
        s = stats.setdefault(iccid, dict(slots=set(), system=set(), kinds=set(), qualified=set()))
        s['slots'].add(slot)
        s['kinds'].add(kind)
        if kind in ('系统任务F库存原始记录', '任务F查询回执原始来源') and not ambiguous:
            s['system'].add(slot)
        if not ambiguous and not kind.startswith('本地库衍生'):
            s['qualified'].add(slot)
    rows = []
    for key, target in sorted(targets.items()):
        s = stats.get(key, dict(slots=set(), system=set(), kinds=set(), qualified=set()))
        slots = sorted(s['slots'])
        status = '未找到' if not slots else '唯一历史卡槽' if len(slots) == 1 else '多卡槽历史记录（不代表同时占用）'
        if slots and not s['qualified']:
            status = '仅衍生表或模糊列，待核对'
        rows.append(dict(iccid=key, origin=target['origin'], lookupStatus=target['lookupStatus'],
                         status=status, slots=slots, systemSlots=sorted(s['system']),
                         sourceKinds=sorted(s['kinds']), activationSlot=target.get('activationSlot', ''),
                         inventorySlot=target.get('inventorySlot', '')))
    return rows


def canonical_row(row):
    return json.dumps([None if v == '' else v for v in row], ensure_ascii=False, separators=(',', ':')).encode() + b'\n'


def export_excel(path, tables):
    wb = Workbook(write_only=True)
    expected = {}
    headfont = Font(name='Arial', bold=True, color='FFFFFF')
    bodyfont = Font(name='Arial', size=11)
    fill = PatternFill('solid', fgColor='284B63')
    for name, headers, rows in tables:
        ws = wb.create_sheet(name)
        ws.freeze_panes = 'A2'
        for i, title in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 55 if '文件' in title else 28
        digest, count = hashlib.sha256(), 0
        for row in iter_rows_with_header(headers, rows):
            if count >= 1048576:
                raise ValueError(f'{name}超过Excel行上限；完整证据仍在gzip账本')
            cells = []
            for value in row:
                cell = WriteOnlyCell(ws, value=value)
                cell.font = headfont if count == 0 else bodyfont
                if count == 0:
                    cell.fill = fill
                if isinstance(value, str):
                    if len(value) > 32767:
                        raise ValueError('禁止截断Excel长字段')
                    cell.data_type, cell.number_format = 's', '@'
                cells.append(cell)
            ws.append(cells)
            digest.update(canonical_row(row))
            count += 1
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count}'
        expected[name] = (count, digest.hexdigest())
    wb.save(path)
    wb.close()
    checked = load_workbook(path, read_only=True, data_only=False)
    try:
        for name, (wanted_count, wanted_hash) in expected.items():
            digest, count = hashlib.sha256(), 0
            for row in checked[name].iter_rows():
                if any(c.data_type in ('f', 'e') for c in row):
                    raise ValueError('输出公式或Excel错误')
                if count and name not in ('扫描文件清单', '说明') and row[0].data_type != 's':
                    raise ValueError('ICCID必须为文本')
                digest.update(canonical_row([c.value for c in row]))
                count += 1
            if count != wanted_count or digest.hexdigest() != wanted_hash:
                raise ValueError('独立回读值或数量不一致: ' + name)
    finally:
        checked.close()
    with zipfile.ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError('ZIP CRC错误')
    return dict(sheets={k: v[0]-1 for k, v in expected.items()},
                sha256=hashlib.sha256(path.read_bytes()).hexdigest(), independentlyVerified=True)


def iter_rows_with_header(headers, rows):
    yield headers
    yield from rows


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--db', type=Path, default=base.DB)
    p.add_argument('--root', action='append')
    p.add_argument('--start', type=date.fromisoformat, required=True)
    p.add_argument('--end', type=date.fromisoformat, required=True)
    p.add_argument('--workers', type=int, default=8)
    p.add_argument('--output', type=Path, required=True)
    args = p.parse_args()
    if not 1 <= args.workers <= 8 or args.start > args.end:
        p.error('workers需为1至8，起止日期需有效')
    started = time.monotonic()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=False)
    (out / '逐条证据').mkdir()
    targets = frozen_targets(args.db)
    current_count = len(targets)
    base.write_json(out / 'target_snapshot.json', dict(database=str(args.db), readOnly=True, records=targets))
    start, end = str(args.start), str(args.end)
    roots = args.root or base.DEFAULT_ROOTS
    files, discovery = base.discover(roots, start, end, out)
    discovery['excluded']['指定日期范围外'] = discovery['excluded'].pop('一个月范围外', 0)
    for f in files:
        f['format'] = 'excel'
    json_files, json_discovery = discover_json(roots, start, end)
    files += json_files
    def progress(**info):
        info['updatedAt'] = datetime.now().astimezone().isoformat()
        tmp = out / 'progress.tmp'
        tmp.write_text(json.dumps(info, ensure_ascii=False), encoding='utf-8')
        tmp.replace(out / 'progress.json')
        print(json.dumps(info, ensure_ascii=False), flush=True)
    progress(event='discovered', currentTaskFCards=current_count, excelFiles=len(files)-len(json_files),
             jsonFiles=len(json_files), workers=args.workers)
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        fingerprints = list(pool.map(base.fingerprint, files))
    groups = defaultdict(list)
    errors = []
    for f in fingerprints:
        if f.get('error'):
            errors.append(dict(file=f['path'], error=f['error']))
        else:
            groups[(f['format'], f['sha256'])].append(f)
    progress(event='hash_complete', uniqueContents=len(groups), duplicateFiles=len(files)-len(groups)-len(errors))
    conn = make_audit_db(out / 'history_index.sqlite3')
    completed, total_ev, total_unresolved = 0, 0, 0
    last_progress = time.monotonic()
    # JSON first adds genuine historical Task F identities absent from today's DB.
    for format_name in ('json', 'excel'):
        phase = [aliases for (kind, _), aliases in groups.items() if kind == format_name]
        with ProcessPoolExecutor(max_workers=args.workers, initializer=init_worker,
                                 initargs=(set(targets), str(out), start, end)) as pool:
            # Bounded in-flight work avoids unbounded result/future memory.
            pending = []
            iterator = iter(phase)
            for _ in range(min(args.workers * 2, len(phase))):
                aliases = next(iterator)
                pending.append((pool.submit(scan_file, aliases[0]), aliases))
            while pending:
                future, aliases = pending.pop(0)
                result = future.result()
                record_result(conn, result, aliases)
                if result.get('error'):
                    errors.extend(dict(file=a['path'], error=result['error']) for a in aliases)
                else:
                    for key in result['discoveredIccids']:
                        targets.setdefault(key, dict(origin='仅历史任务F回执', lookupStatus='当前库无记录',
                                                     activationSlot='', inventorySlot=''))
                    total_ev += result['evidenceRows']
                    total_unresolved += result['unresolvedRows']
                # Small per-file audit metadata, never credential-bearing JSON.
                base.write_json(out / '逐条证据' / (result['sha256'] + '.meta.json'),
                                {k: v for k, v in result.items() if k not in ('aggregates', 'discoveredIccids')})
                completed += 1
                if time.monotonic() - last_progress > 15 or completed == len(groups):
                    progress(event='scan_progress', phase=format_name, completed=completed, total=len(groups),
                             targetCards=len(targets), evidenceRows=total_ev, fileErrors=len(errors))
                    last_progress = time.monotonic()
                try:
                    aliases = next(iterator)
                    pending.append((pool.submit(scan_file, aliases[0]), aliases))
                except StopIteration:
                    pass
    integrity = conn.execute('PRAGMA integrity_check').fetchone()[0]
    if integrity != 'ok':
        raise ValueError('审计索引完整性校验失败')
    rows = card_summaries(conn, targets)
    base.write_json(out / 'results.json', rows)
    base.write_json(out / 'files.json', fingerprints)
    headers = ['ICCID', '范围', '本地查询状态', '历史查询结论', '历史卡槽数', '历史卡槽集合',
               '系统库存或F回执直接证据卡槽', '来源类型', '查询前激活库卡槽', '查询前库存卡槽']
    def card_rows(selected):
        for r in selected:
            yield [r['iccid'], r['origin'], r['lookupStatus'], r['status'], len(r['slots']), '；'.join(r['slots']),
                   '；'.join(r['systemSlots']), '；'.join(r['sourceKinds']), r['activationSlot'], r['inventorySlot']]
    detail_headers = ['ICCID', '历史卡槽', '来源类型', '列关联待核对', '去重内容内观测次数', '不同内容份数',
                      '最早来源文件日期', '最晚来源文件日期', '示例来源文件', '来源SHA256', '工作表', '单元格或JSON位置']
    instructions = [
        ['范围', f'{start} 至 {end}，含首尾；本地任务F全库和历史F回执中的ICCID'],
        ['身份匹配', '18至22位文本ICCID精确匹配，允许去末尾一个F/f；拒绝数字、公式、尾号匹配'],
        ['卡槽', '保留r前缀和前导零；多卡槽为多次历史记录，不等于当前位置或同时占用'],
        ['来源优先', '原始系统任务F库存/查询回执单列；Google、处理表、库衍生表不等同系统原始证据'],
        ['日期依据', '以文件名或归档目录日期筛选；无日期时用mtime。不是位置生效时间'],
        ['完整证据', '逐条证据/<SHA256>.jsonl.gz保留每条精确行/JSON位置；扫描文件清单列出所有副本'],
        ['回填', '未更新生产数据库，未修改历史Excel，未访问后台；未找到不证明从未上架'],
        ['空或未知表头', '无同表ICCID和卡槽列不猜测匹配；逐文件meta.json记录未解析的工作表'],
        ['重复计数', '相同文件内容仅解析一次，观测次数不乘副本数；同一原始行及其来源子对象可重复出现'],
        ['错误文件数', len(errors)],
    ]
    tables = [('全部任务F卡', headers, card_rows(rows)),
              ('未激活卡子集', headers, card_rows(r for r in rows if r['lookupStatus'] == 'activation_missing')),
              ('多卡槽待核对', headers, card_rows(r for r in rows if len(r['slots']) > 1)),
              ('未找到卡槽', headers, card_rows(r for r in rows if not r['slots'])),
              ('卡槽来源汇总', detail_headers, conn.execute('SELECT * FROM slots ORDER BY iccid,slot,kind,ambiguous')),
              ('扫描文件清单', ['来源文件', '日期', '日期依据', 'SHA256', '格式', '状态', '错误', '完整逐条证据', '有效证据行数'],
               conn.execute('SELECT * FROM files ORDER BY path')),
              ('说明', ['项目', '说明'], instructions)]
    progress(event='writing_and_verifying', targetCards=len(rows), evidenceRows=total_ev)
    excel = out / '任务F全部卡_近三个月卡槽历史.xlsx'
    verified = export_excel(excel, tables)
    counts = dict(Counter(r['status'] for r in rows))
    missing_subset = [r for r in rows if r['lookupStatus'] == 'activation_missing']
    summary = dict(status='completed_with_file_errors' if errors else 'completed', start=start, end=end,
                   currentDatabaseCards=current_count, historicalOnlyCards=len(rows)-current_count,
                   targetCards=len(rows), matchedCards=sum(bool(r['slots']) for r in rows), counts=counts,
                   hasDirectSystemEvidenceCards=sum(bool(r['systemSlots']) for r in rows),
                   googleOnlyCards=sum(bool(r['sourceKinds']) and all('Google' in k for k in r['sourceKinds']) for r in rows),
                   activationMissingSubset=dict(total=len(missing_subset), matched=sum(bool(r['slots']) for r in missing_subset),
                                                directSystem=sum(bool(r['systemSlots']) for r in missing_subset)),
                   selectedFiles=len(files), excelFiles=len(files)-len(json_files), jsonFiles=len(json_files),
                   uniqueContents=len(groups), evidenceRows=total_ev, unresolvedRows=total_unresolved,
                   scanErrors=errors, discovery=discovery, jsonDiscovery=json_discovery,
                   workers=args.workers, readOnlyProduction=True, databaseUpdated=False, backendAccessed=False,
                   workbook=str(excel), outputVerification=verified, auditIndex=str(out/'history_index.sqlite3'),
                   durationSeconds=round(time.monotonic()-started, 2),
                   scriptSha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    base.write_json(out / 'summary.json', summary)
    conn.execute('PRAGMA wal_checkpoint(TRUNCATE)')
    conn.close()
    report = ['# 任务F近三个月卡槽历史查询', '', f'范围：{start} 至 {end}（含首尾）。', '',
              f'当前库 {current_count} 张，额外历史F回执 {len(rows)-current_count} 张，总计 {len(rows)} 张。', '',
              '| 历史查询结论 | 卡数 |', '|---|---:|', *[f'| {k} | {v} |' for k,v in counts.items()], '',
              f'有原始系统库存或F回执直接卡槽证据：{summary["hasDirectSystemEvidenceCards"]} 张。',
              f'其中未激活 {len(missing_subset)} 张，找到 {summary["activationMissingSubset"]["matched"]} 张。',
              f'Excel {summary["excelFiles"]} 份、JSON {len(json_files)} 份；内容去重后 {len(groups)} 份。',
              f'文件级错误 {len(errors)} 份。所有错误明细见扫描文件清单和summary.json，不能宣称错误文件已读完。', '',
              '本轮仅查找历史，不回填、不改后台。多卡槽可能是调机或来源冲突；文件日期不等于卡槽生效时间。',
              'ICCID按文本精确匹配，所有输出已逐格独立回读及ZIP校验。完整逐条来源在gzip账本。', '',
              f'[结果Excel]({excel})', f'[审计汇总]({out/"summary.json"})', f'[查询索引]({out/"history_index.sqlite3"})']
    (out / '报告.md').write_text('\n'.join(report) + '\n', encoding='utf-8')
    progress(event='completed', **{k:v for k,v in summary.items() if k not in ('scanErrors','discovery','jsonDiscovery')})
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
