#!/usr/bin/env python3
"""Reconcile a hash-bound final report with Task F SQLite; never query or upload."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timedelta
import hashlib
import io
import json
from pathlib import Path
import re
import sqlite3
import sys
import zipfile
import warnings

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


RUN_RE = re.compile(r"\d{8}T\d{6}[+-]\d{4}")
DETAIL_HEADERS = ["源行", "电话", "卡槽", "原ICCID", "去F查询ICCID", "系统激活日期", "系统截止日",
                  "本地激活日期", "本地截止日", "查询时间", "状态", "Purchase Date", "Act. Request Date",
                  "Activation Date", "SIM Number", "SKU", "本地数据来源"]
UPLOAD_HEADERS = ["ICCID", "截止日", "激活日期"]
DB_FIELDS = ["query_iccid", "source_iccid", "activation_date", "cutoff_date", "lookup_status", "queried_at",
             "purchase_date", "purchase_date_raw", "act_request_date", "activation_date_raw", "sim_number",
             "sim_number_raw", "sku", "data_source", "source_data_json", "record_schema_version"]


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def encoded(value) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n").encode("utf-8")


def normalize_iccid(value) -> str:
    if not isinstance(value, str):
        return ""
    normalized = re.sub(r"[Ff]$", "", value)
    return normalized if re.fullmatch(r"[0-9]{18,22}", normalized) else ""


def iso_date(value) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if not isinstance(value, str) or not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value):
        raise ValueError("日期不是有效 yyyy-MM-dd")
    parsed = date.fromisoformat(value)
    if parsed.isoformat() != value:
        raise ValueError("日期不是有效自然日")
    return value


def valid_found(record: dict | None) -> bool:
    if not record or record.get("lookup_status") != "found":
        return False
    try:
        activation, cutoff = iso_date(record.get("activation_date")), iso_date(record.get("cutoff_date"))
        return bool(activation and cutoff and date.fromisoformat(cutoff) == date.fromisoformat(activation) + timedelta(days=29))
    except ValueError:
        return False


def read_report(state_path: Path, *, source_kind: str = "final") -> tuple[dict, dict, list[dict]]:
    if source_kind not in {"final", "system-export"}:
        raise ValueError("不支持的文档来源")
    state_bytes = state_path.read_bytes()
    state = json.loads(state_bytes)
    run_id = state.get("runId")
    if not isinstance(run_id, str) or not RUN_RE.fullmatch(run_id) or state.get("status") != "success":
        raise ValueError("只接受 runId 有效且已成功完成的巡检状态")
    if state_path.name != f"opensms_patrol_runner_state_{run_id}.json" or state_path.parent.name != "receipts":
        raise ValueError("状态文件路径与 runId 不一致")
    stages = state.get("stages", {})
    source = stages.get("acquire" if source_kind == "system-export" else "refresh", {})
    document = Path(source.get("sourceArchive" if source_kind == "system-export" else "finalWorkbook") or "").resolve()
    expected_hash = source.get("sourceSha256" if source_kind == "system-export" else "finalWorkbookSha256")
    if source_kind == "system-export" and source.get("sourceVerified") is not True:
        raise ValueError("原始报表未通过新鲜源校验")
    if not re.fullmatch(r"[a-fA-F0-9]{64}", expected_hash or ""):
        raise ValueError("最终报表缺少独立刷新 SHA256 绑定")
    if document.parent != state_path.parent.parent or document.suffix.lower() != ".xlsx" or run_id not in document.name:
        raise ValueError("最终报表不是本轮交付目录中的 XLSX")
    data = document.read_bytes()
    if sha256_bytes(data) != expected_hash.lower():
        raise ValueError("最终报表 SHA256 与成功巡检状态不一致")
    database_value = stages.get("taskF", {}).get("database")
    if not database_value and source_kind == "system-export":
        database_value = str(Path(__file__).resolve().parents[1] / "opensms_artifacts/state/task_f_kooda_activation.sqlite3")
    if not isinstance(database_value, str) or not database_value:
        raise ValueError("成功巡检状态未记录任务 F 数据库路径")
    database = Path(database_value).resolve()
    if not database.is_file():
        raise ValueError("本地任务 F 数据库不存在，禁止创建")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    counts = Counter()
    rows = []
    try:
        if "sheet1" not in workbook.sheetnames:
            raise ValueError("最终报表缺少 sheet1")
        iterator = workbook["sheet1"].iter_rows(values_only=True)
        headers = list(next(iterator))
        required = ["电话", "卡槽", "opr", "激活日期", "截止日", "ICCID"]
        if any(headers.count(name) != 1 for name in required) or (source_kind == "final" and headers[-1] != "ICCID"):
            raise ValueError("最终主表列头缺失、重复或 ICCID 不是末列")
        index = {name: headers.index(name) for name in required}
        present_keys = set()
        for row_number, values in enumerate(iterator, 2):
            if not re.search(r"kooda|koodo|telus1", str(values[index["opr"]] or ""), re.I):
                continue
            counts["documentKoodaRows"] += 1
            activation = iso_date(values[index["激活日期"]])
            if activation:
                counts["documentActivationPresentRows"] += 1
                present_keys.add(normalize_iccid(values[index["ICCID"]]))
                continue
            counts["documentMissingActivationRows"] += 1
            source_iccid = values[index["ICCID"]]
            query_iccid = normalize_iccid(source_iccid)
            rows.append({"documentRow": row_number, "phone": str(values[index["电话"]] or ""),
                         "slot": values[index["卡槽"]], "opr": values[index["opr"]], "sourceIccid": source_iccid,
                         "queryIccid": query_iccid, "systemActivationDate": activation,
                         "systemCutoffDate": iso_date(values[index["截止日"]]),
                         "status": "awaiting_local_match" if query_iccid else "excluded_invalid_iccid"})
        if any(row["queryIccid"] and row["queryIccid"] in present_keys for row in rows):
            raise ValueError("同一ICCID同时存在空激活和已有激活行，禁止覆盖系统日期")
    finally:
        workbook.close()
    metadata = {"runId": run_id, "statePath": str(state_path), "stateSha256": sha256_bytes(state_bytes),
                "sourceDocument": str(document), "sourceDocumentSha256": expected_hash.lower(),
                "database": str(database), "sourceKind": source_kind, **{name: counts[name] for name in (
                    "documentKoodaRows", "documentActivationPresentRows", "documentMissingActivationRows")}}
    return state, metadata, rows


def read_database(database: Path, query_iccids: list[str]) -> dict[str, dict]:
    records = {}
    connection = sqlite3.connect(database.as_uri() + "?mode=ro", uri=True, timeout=30)
    try:
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA query_only=ON")
        connection.execute("BEGIN")
        for offset in range(0, len(query_iccids), 800):
            batch = query_iccids[offset:offset + 800]
            query = "SELECT " + ",".join(DB_FIELDS) + " FROM kooda_activation WHERE query_iccid IN (" + ",".join("?" for _ in batch) + ")"
            for record in connection.execute(query, batch):
                value = dict(record)
                key = value["query_iccid"]
                if key in records or key not in batch:
                    raise ValueError("数据库精确 ICCID 匹配重复或身份不一致")
                raw_source = value.pop("source_data_json")
                value["sourceData"] = json.loads(raw_source) if raw_source else []
                records[key] = value
    finally:
        connection.close()
    return records


def detail_values(row: dict) -> list:
    record = row.get("databaseRecord") or {}
    raw_iccid = None if row["sourceIccid"] is None else str(row["sourceIccid"])
    return [row["documentRow"], row["phone"], row["slot"], raw_iccid, row["queryIccid"],
            row["systemActivationDate"], row["systemCutoffDate"], record.get("activation_date"),
            record.get("cutoff_date"), record.get("queried_at"), row["status"],
            record.get("purchase_date_raw") or record.get("purchase_date"), record.get("act_request_date"),
            record.get("activation_date_raw") or record.get("activation_date"),
            record.get("sim_number_raw") or record.get("sim_number"), record.get("sku"), record.get("data_source")]


def write_workbook(path: Path, sheet_name: str, headers: list, rows: list[list], *, upload: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append([date.fromisoformat(value) if upload and i in (1, 2) else value for i, value in enumerate(row)])
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top")
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise ValueError("单元格超过 Excel 文本上限；完整来源在 JSON 中，禁止静默截断")
                cell.data_type = "s"
            if upload and cell.column in (2, 3):
                cell.number_format = "yyyy-mm-dd"
            elif upload or cell.column in (2, 4, 5, 15):
                cell.number_format = "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        letter = cells[0].column_letter
        sheet.column_dimensions[letter].width = 23 if upload else 25
    if not upload:
        sheet.column_dimensions["A"].width = 10
    with path.open("xb") as handle:
        workbook.save(handle)
    workbook.close()


def verify_workbook(path: Path, sheet_name: str, headers: list, rows: list[list], *, upload: bool = False) -> None:
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("核对工作簿压缩结构异常")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != [sheet_name]:
            raise ValueError("核对工作簿工作表结构异常")
        sheet = workbook[sheet_name]
        actual = list(sheet.iter_rows())
        if [cell.value for cell in actual[0]] != headers or len(actual) != len(rows) + 1:
            raise ValueError("核对工作簿行数或表头不一致")
        for cells, expected in zip(actual[1:], rows):
            for index, (cell, value) in enumerate(zip(cells, expected)):
                decoded = iso_date(cell.value) if upload and index in (1, 2) else cell.value
                if decoded != (None if value == "" else value) or cell.data_type in {"f", "e"}:
                    raise ValueError("核对工作簿独立回读值不一致")
                if value and ((upload and index == 0) or (not upload and index in (1, 3, 4, 14))):
                    if cell.data_type != "s" or cell.number_format != "@":
                        raise ValueError("ICCID/电话/SIM 文本精度校验失败")
    finally:
        workbook.close()


def reconcile(state_path: Path, output_dir: Path, *, source_kind: str = "final") -> dict:
    state_path, output_dir = state_path.resolve(), output_dir.resolve()
    if output_dir.exists():
        raise FileExistsError("输出目录已存在，禁止覆盖原报表或历史结果")
    if not output_dir.parent.is_dir():
        raise ValueError("输出目录的父目录必须已存在")
    _, metadata, rows = read_report(state_path, source_kind=source_kind)
    keys = sorted({row["queryIccid"] for row in rows if row["queryIccid"]})
    records = read_database(Path(metadata["database"]), keys)
    ready = set()
    pending = set()
    for row in rows:
        key = row["queryIccid"]
        record = records.get(key)
        row["databaseRecord"] = record
        if not key:
            continue
        if valid_found(record):
            if row["systemCutoffDate"] and row["systemCutoffDate"] != record["cutoff_date"]:
                raise ValueError("系统已有截止日与F库不一致，禁止用导入表覆盖系统日期")
            row["status"] = "ready_for_manual_upload"
            ready.add(key)
        else:
            row["status"] = "pending_no_local_record" if record is None else "pending_no_valid_local_date"
            pending.add(key)
    status_counts = dict(Counter(row["status"] for row in rows))
    upload_rows = [[key, records[key]["cutoff_date"], records[key]["activation_date"]] for key in sorted(ready)]
    detail_rows = [detail_values(row) for row in rows]
    if sha256_bytes(Path(metadata["sourceDocument"]).read_bytes()) != metadata["sourceDocumentSha256"]:
        raise ValueError("读取期间原报表发生变化，停止导出")
    output_dir.mkdir()
    paths = {"reconciliationWorkbook": output_dir / "任务F_报表空激活日期_本地库核对.xlsx",
             "manualUploadWorkbook": output_dir / "任务F_截止日激活日期_待人工上传.xlsx",
             "report": output_dir / "reconcile.md", "receipt": output_dir / "reconcile.json"}
    write_workbook(paths["reconciliationWorkbook"], "空激活日期核对", DETAIL_HEADERS, detail_rows)
    if source_kind == "system-export" and upload_rows:
        from opensms_export_task_f_deadlines import export_deadlines, DEFAULT_TEMPLATE
        export_deadlines(Path(metadata["database"]), DEFAULT_TEMPLATE, paths["manualUploadWorkbook"],
                         output_dir / "template_verification.json", selected_rows=upload_rows,
                         extra_receipt={"scope": "source_document_missing_activation_only",
                                        **{key: value for key, value in metadata.items() if key != "database"}})
    else:
        from opensms_export_task_f_deadlines import write_template_workbook, DEFAULT_TEMPLATE
        verified_template = write_template_workbook(DEFAULT_TEMPLATE, upload_rows, paths["manualUploadWorkbook"])
        (output_dir / "template_verification.json").write_bytes(encoded({
            **verified_template, "scope": "source_document_missing_activation_only", **metadata}))
    verify_workbook(paths["reconciliationWorkbook"], "空激活日期核对", DETAIL_HEADERS, detail_rows)
    verify_workbook(paths["manualUploadWorkbook"], "截止日", UPLOAD_HEADERS, upload_rows, upload=True)
    if sha256_bytes(Path(metadata["sourceDocument"]).read_bytes()) != metadata["sourceDocumentSha256"]:
        raise ValueError("导出期间原报表发生变化，未生成成功回执")
    result = {"status": "success", **metadata, "createdAt": datetime.now().astimezone().isoformat(),
              "queryIccidCount": len(keys), "duplicateIccidRows": sum(bool(row["queryIccid"]) for row in rows) - len(keys),
              "readyForUploadCount": len(ready), "pendingCount": len(pending),
              "excludedCount": status_counts.get("excluded_invalid_iccid", 0), "statusCounts": status_counts,
              "pendingIccids": sorted(pending), "rows": rows,
              "databaseRowsSha256": sha256_bytes(encoded(records)),
              "databaseReadOnly": True, "sourceDocumentUnchanged": True,
              "networkAccess": False, "browserAccess": False, "cloudUploadAttempted": False,
              "documentDateSource": "system", "taskFDateUsage": "local_library_for_manual_upload_only",
              "verificationPassed": True, "files": {key: str(path) for key, path in paths.items()},
              "sha256": {key: sha256_bytes(path.read_bytes()) for key, path in paths.items() if path.suffix == ".xlsx"}}
    report = (f"# 任务 F 本地核对报告\n\n运行：{metadata['runId']}\n\n"
              f"Kooda 主表 {metadata['documentKoodaRows']} 行；系统激活日期为空 {len(rows)} 行。\n\n"
              f"合法唯一 ICCID {len(keys)}；本地已有有效日期、可人工上传 {len(ready)}；本地缺口 {len(pending)}；排除 {result['excludedCount']}。\n\n"
              "## 为什么原报表仍为空\n\n主报表的 Kooda 日期来自本轮系统字段，任务 F 仅更新独立本地库，不回填原报表。"
              "本次以所选来源文档激活日期为空的ICCID为范围，不导出文档外的全库记录。"
              "原始系统报表的空白不会被已生成合并表中的本地补值或创建日期推算掩盖。"
              "任务F库里的有效日期可供人工导入；仍无有效日期的记录仅列在核对表和本报告，不放进导入表。"
              "下一次新鲜主报表可用本地库补系统空白，但不覆盖系统已有日期。\n\n"
              "本次只读核对：未联网、未重新查询、未上传云端、未修改原工作簿或数据库。"
              "本地缺口仅记录为 pending，未执行任何补查。完整五项来源字段及源数据保存在 JSON。\n\n"
              "## 仍缺日期的ICCID\n\n" + ("\n".join(f"- {key}" for key in sorted(pending)) or "无") + "\n\n"
              "## 文件\n\n" + "\n".join(f"- [{path.name}]({path})" for path in paths.values()) + "\n")
    with paths["report"].open("xb") as handle:
        handle.write(report.encode("utf-8"))
    with paths["receipt"].open("xb") as handle:
        handle.write(encoded(result))
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="只读核对最终报表缺激活日期的 Kooda 与本地库；不联网、不上传")
    parser.add_argument("--state", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--source-kind", choices=["final", "system-export"], default="final")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    try:
        output_dir = args.output_dir
        if output_dir is None:
            base = Path(__file__).resolve().parents[1] / "outputs/task_f_document_missing"
            base.mkdir(parents=True, exist_ok=True)
            output_dir = base / datetime.now().astimezone().strftime("%Y%m%dT%H%M%S_%f%z")
        result = reconcile(args.state, output_dir, source_kind=args.source_kind)
        summary = {key: value for key, value in result.items() if key not in {"rows", "pendingIccids"}}
        print(json.dumps(summary, ensure_ascii=False) if args.json else (
            f"核对完成：空激活 {result['documentMissingActivationRows']} 行，待人工上传 {result['readyForUploadCount']}，"
            f"本地缺口 {result['pendingCount']}。完整报告：{result['files']['report']}"))
        return 0
    except (OSError, ValueError, sqlite3.Error, KeyError, TypeError) as exc:
        print(json.dumps({"status": "failed", "error": str(exc), "networkAccess": False,
                          "databaseReadOnly": True, "cloudUploadAttempted": False}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    sys.exit(main())
