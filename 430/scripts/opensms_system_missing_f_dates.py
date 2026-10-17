"""Read-only, full-source-row difference export; never query or upload."""
import hashlib
import json
from pathlib import Path
import warnings
from openpyxl import load_workbook
from opensms_task_f_reconcile_document import normalize_iccid, iso_date, valid_found, read_database
from opensms_export_task_f_deadlines import DEFAULT_TEMPLATE, write_template_workbook, verify_output

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / 'opensms_artifacts/state/task_f_kooda_activation.sqlite3'

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def select_rows(headers, source_rows, records):
    for name in ('ICCID','激活日期'):
        if headers.count(name) != 1:
            raise ValueError('系统源列缺失或重复:'+name)
    iccid, activation = headers.index('ICCID'), headers.index('激活日期')
    result=[]
    for values in source_rows:
        # An invalid nonblank date is not a missing date.
        if iso_date(values[activation]) is not None:
            continue
        key=normalize_iccid(values[iccid])
        record=records.get(key)
        if not valid_found(record):
            continue
        if '截止日' in headers:
            if headers.count('截止日') != 1:
                raise ValueError('系统源列重复:截止日')
            cutoff = iso_date(values[headers.index('截止日')])
            if cutoff and cutoff != record['cutoff_date']:
                raise ValueError('系统已有截止日与F库不一致，禁止用导入表覆盖系统日期')
        result.append((list(values),key,record))
    return result

def template_rows(selected):
    """Import one row per ICCID; all source occurrences stay in the bound detail JSON."""
    unique = {}
    for _, key, record in selected:
        row = [key, record['cutoff_date'], record['activation_date']]
        if key in unique and unique[key] != row:
            raise ValueError('同ICCID的F库日期冲突，停止导出')
        unique[key] = row
    return [unique[key] for key in sorted(unique)]

def read_source(source):
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        wb = load_workbook(source, read_only=True, data_only=False)
    try:
        iterator = wb['sheet1'].iter_rows(values_only=True)
        return list(next(iterator)), list(iterator)
    finally:
        wb.close()

def read_frozen_dates(cache, run_id):
    path = Path(cache.get('path') or '')
    if (cache.get('runId') != run_id or cache.get('frozenBeforeDocument') is not True
        or not path.is_file() or digest(path) != cache.get('sha256')):
        raise ValueError('本轮冻结F日期来源缺失或哈希变化')
    payload = json.loads(path.read_text())
    if (payload.get('runId') != run_id or payload.get('status') != 'passed'
        or payload.get('readOnly') is not True or payload.get('snapshotConsistent') is not True
        or not isinstance(payload.get('records'), list)
        or payload.get('recordCount') != len(payload['records'])):
        raise ValueError('本轮冻结F日期来源结构无效')
    records = {}
    for row in payload['records']:
        # Manual renewal has its own explicit overwrite authority and receipts.
        # Never feed it to automatic original-activation/+29 gap imports.
        if row.get('renewal') is not None:
            continue
        key = normalize_iccid(row.get('queryIccid'))
        record = {'query_iccid': key, 'lookup_status': 'found',
                  'activation_date': row.get('activationDate'), 'cutoff_date': row.get('cutoffDate'),
                  'data_source': 'current_run_pre_document_frozen_cache'}
        if not key or key in records or not valid_found(record):
            raise ValueError('本轮冻结F日期来源存在重复或无效记录')
        records[key] = record
    return records

def export_difference(source, source_sha, run_id, delivery, receipts, *, preview=False,
                      template=DEFAULT_TEMPLATE, report_date_cache=None):
    source, delivery, receipts=Path(source),Path(delivery),Path(receipts)
    output=delivery/f'OpenSMS_系统缺激活日期_F库已有_{run_id}.xlsx'
    receipt=receipts/f'opensms_system_missing_f_dates_{run_id}.json'
    if output.exists() or receipt.exists():
        if receipt.is_file() and output.is_file():
            saved=json.loads(receipt.read_text())
            if (saved.get('runId')==run_id and saved.get('sourceSha256')==source_sha
                and saved.get('independentReadbackPassed') is True
                and saved.get('source')==str(source)
                and saved.get('output')==str(output)
                and digest(source)==source_sha and digest(output)==saved.get('sha256')
                and digest(saved['databaseSnapshot'])==saved.get('databaseSnapshotSha256')):
                if report_date_cache is not None:
                    read_frozen_dates(report_date_cache, run_id)
                    if saved.get('reportDateCacheSha256') != report_date_cache.get('sha256'):
                        raise ValueError('既有差异导出不是本轮文档前冻结F日期，禁止复用')
                if saved.get('formatVersion') == 'task_f_deadline_template_v2':
                    if (digest(saved['sourceDetail']) != saved.get('sourceDetailSha256')
                        or digest(saved['template']) != saved.get('templateSha256')):
                        raise ValueError('模板或完整来源回执变化，禁止复用')
                return saved
        raise FileExistsError('既有差异导出未完整核验，禁止覆盖或复用')
    if digest(source)!=source_sha:
        raise ValueError('系统源SHA256不一致')
    headers, source_rows = read_source(source)
    if headers.count('ICCID')!=1 or headers.count('激活日期')!=1:
        raise ValueError('系统源列缺失或重复')
    keys=sorted({normalize_iccid(row[headers.index('ICCID')]) for row in source_rows if iso_date(row[headers.index('激活日期')]) is None})
    allowed_keys = set(keys)
    records = (read_database(DB, [key for key in keys if key]) if report_date_cache is None
               else {key: record for key, record in read_frozen_dates(report_date_cache, run_id).items()
                     if key in allowed_keys})
    snapshot=receipts/f'opensms_system_missing_f_dates_{run_id}_database.json'
    with snapshot.open('x') as handle: json.dump(records,handle,ensure_ascii=False,indent=2)
    selected=select_rows(headers,source_rows,records)
    extra=['F库查询ICCID','F库激活日期','F库截止日','F库查询时间','F库Purchase Date','F库Act. Request Date','F库Activation Date','F库SIM Number','F库SKU','F库数据来源']
    rows=[]
    for raw,key,record in selected:
        rows.append(raw+[key,record['activation_date'],record['cutoff_date'],record.get('queried_at'),record.get('purchase_date_raw') or record.get('purchase_date'),record.get('act_request_date'),record.get('activation_date_raw'),record.get('sim_number_raw') or record.get('sim_number'),record.get('sku'),record.get('data_source')])
    def serial(value):
        return value.isoformat() if hasattr(value,'isoformat') else value
    rows=[[serial(value) for value in row] for row in rows]
    payload={'headers':headers+extra,'rows':rows,'textColumns':[i for i,name in enumerate(headers+extra) if name in ('电话','ICCID','卡槽') or name.startswith('F库')]}
    inputs=receipts/f'opensms_system_missing_f_dates_{run_id}_rows.json'
    with inputs.open('x') as handle: json.dump(payload,handle,ensure_ascii=False)
    imports = template_rows(selected)
    written = write_template_workbook(Path(template), imports, output)
    # Independently recompute from the frozen input, never from a changing live DB.
    checked_headers, checked_rows = read_source(source)
    checked_records = json.loads(snapshot.read_text())
    checked_selected = select_rows(checked_headers, checked_rows, checked_records)
    if len(checked_selected) != len(selected) or template_rows(checked_selected) != imports:
        raise ValueError('源文件与冻结F库独立重算不一致')
    verify_output(output, template_rows(checked_selected))
    if digest(source)!=source_sha: raise ValueError('导出期间系统源变化')
    result={'status':'passed','runId':run_id,'rowCount':len(rows),'uniqueIccidCount':len(imports),
        'reportDateCacheSha256': (report_date_cache or {}).get('sha256'),
        'resultsApplyNextExport': report_date_cache is not None,
        'exportedRowCount':len(imports),'formatVersion':'task_f_deadline_template_v2',
        'template':written['template'],'templateSha256':written['templateSha256'],
        'headers':written['headers'],'sheet':written['sheet'],'dateFormat':written['dateFormat'],
        'sourceDetail':str(inputs),'sourceDetailSha256':digest(inputs),
        'output':str(output),'sha256':digest(output),'source':str(source),'sourceSha256':source_sha,'databaseSnapshot':str(snapshot),'databaseSnapshotSha256':digest(snapshot),'databaseReadOnly':True,'networkAccess':False,'cloudUploadAttempted':False,'independentReadbackPassed':True,'receipt':str(receipt)}
    with receipt.open('x') as handle: json.dump(result,handle,ensure_ascii=False,indent=2)
    return result
