"""Read-only selection from the current final workbook, never a second export."""
import argparse
import hashlib
import json
import time
from copy import copy
from itertools import zip_longest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

PREFIXES = (
    '60902', '55972', '54782', '65440', '54422', '56580', '55802', '64106',
    '58828', '57156', '64416', '56924', '65370', '60710', '53564', '65476',
    '62050', '50360', '55072', '53322',
)


def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def selected_rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    if headers.count('卡槽') != 1:
        raise ValueError('指定端口导出要求唯一卡槽列')
    index = headers.index('卡槽')
    yield tuple(sheet[1])
    for row in sheet.iter_rows(min_row=2):
        value = '' if row[index].value is None else str(row[index].value).strip()
        if value.startswith(PREFIXES):
            yield row


def signature(cell, style_cache=None):
    # Explicit empty strings may round-trip as blank cells in Excel.
    value = None if cell.value == '' else cell.value
    kind = 'blank' if value is None else cell.data_type
    # Cache only immutable style descriptions, separately for each independently
    # reopened workbook. Cross-workbook style IDs are NOT comparable.
    key = cell.style_id
    if style_cache is None or key not in style_cache:
        style = (cell.number_format, str(cell.font), str(cell.fill),
                 str(cell.border), str(cell.alignment), str(cell.protection))
        if style_cache is not None:
            style_cache[key] = style
    else:
        style = style_cache[key]
    return (value, kind, *style)


def verify(source, output):
    # Materialize blank cells as Cell objects so default styles/protection are
    # compared as well. Read-only EmptyCell lacks protection entirely.
    original = load_workbook(source, data_only=False)
    result = load_workbook(output, data_only=False)
    try:
        if result.sheetnames != ['sheet1']:
            raise ValueError('指定端口文件工作表不正确')
        count = 0
        original_styles, result_styles = {}, {}
        for expected, actual in zip_longest(selected_rows(original['sheet1']), result['sheet1'].iter_rows()):
            if expected is None or actual is None or (
                tuple(signature(cell, original_styles) for cell in expected)
                != tuple(signature(cell, result_styles) for cell in actual)
            ):
                raise ValueError('指定端口文件逐行值、类型或样式核验失败')
            count += 1
        return count - 1
    finally:
        original.close()
        result.close()


def export_subset(source, output, receipt, run_id, expected_sha):
    started = time.perf_counter()
    source, output, receipt = map(lambda p: Path(p).resolve(), (source, output, receipt))
    if not run_id or digest(source) != expected_sha:
        raise ValueError('指定端口来源运行编号为空或最终表SHA256不符')
    if output.exists() or receipt.exists() or source in (output, receipt):
        raise FileExistsError('禁止覆盖已有文件或来源')
    book = load_workbook(source, data_only=False)
    target = Workbook()
    out = target.active
    out.title = 'sheet1'
    counts = dict.fromkeys(PREFIXES, 0)
    # The destination has its own style table. Register components once per
    # source style, then reuse the destination-local immutable StyleArray.
    target_styles = {}
    try:
        sheet = book['sheet1']
        slot_index = [cell.value for cell in sheet[1]].index('卡槽')
        for name, dimension in sheet.column_dimensions.items():
            out.column_dimensions[name] = copy(dimension)
        out.freeze_panes = sheet.freeze_panes
        for number, row in enumerate(selected_rows(sheet), 1):
            if number > 1:
                slot = str(row[slot_index].value).strip()
                counts[next(prefix for prefix in PREFIXES if slot.startswith(prefix))] += 1
            out.row_dimensions[number].height = sheet.row_dimensions[row[0].row].height
            for column, cell in enumerate(row, 1):
                if cell.data_type in ('f', 'e'):
                    raise ValueError('指定端口分表含公式或错误值，拒绝生成可能失效引用')
                dest = out.cell(number, column, cell.value)
                dest.data_type = cell.data_type
                style_id = cell.style_id
                if style_id not in target_styles:
                    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
                        setattr(dest, attribute, copy(getattr(cell, attribute)))
                    dest.number_format = cell.number_format
                    target_styles[style_id] = copy(dest._style)
                else:
                    dest._style = copy(target_styles[style_id])
                if cell.comment:
                    dest.comment = copy(cell.comment)
        out.auto_filter.ref = out.dimensions
        with output.open('xb') as stream:
            target.save(stream)
    finally:
        book.close()
        target.close()
    with ZipFile(output) as archive:
        if archive.testzip() is not None:
            raise ValueError('指定端口文件ZIP验证失败')
    count = verify(source, output)
    if digest(source) != expected_sha or count != sum(counts.values()):
        raise ValueError('指定端口来源变化或数量不一致')
    payload = {'status': 'passed', 'runId': run_id, 'source': str(source), 'sourceSha256': expected_sha,
               'output': str(output), 'outputSha256': digest(output), 'receipt': str(receipt),
               'field': '卡槽', 'matchMode': 'trimmed_text_startswith', 'prefixes': list(PREFIXES),
               'rowCount': count, 'counts': counts, 'independentReadbackPassed': True,
               'sourceUnchanged': True, 'networkAccess': False,
               'performance': {'durationSeconds': round(time.perf_counter() - started, 3),
                               'copiedDistinctStyles': len(target_styles),
                               'styleCacheScope': 'workbook_local_only',
                               'independentSourceReopen': True,
                               'singleWorkbookWriter': True}}
    with receipt.open('x', encoding='utf-8') as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2)
    return payload


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for key in ('source', 'output', 'receipt', 'run-id', 'source-sha256'):
        parser.add_argument('--' + key, required=True)
    args = parser.parse_args()
    print(json.dumps(export_subset(args.source, args.output, args.receipt, args.run_id, args.source_sha256), ensure_ascii=False))
