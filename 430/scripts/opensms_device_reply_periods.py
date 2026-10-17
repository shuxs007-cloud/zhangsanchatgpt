"""Display-only device reply periods over one bound Task A export.

No network, browser, business database or workbook writes. Normal generation
uses already loaded scan-cache rows; the audit independently reopens the source.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, timedelta
import json
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from opensms_export_scan_cache import load_scan_cache, sha256
from services.opensms_document_api.pipeline.device_daily_statistics import (
    UNKNOWN_DEVICE, count_value, device_key, device_sort_key,
)
from services.opensms_document_api.pipeline.room_mapping import (
    CONFIG_SHA256, LABELS, room_for_device,
)

CONTRACT = "device-window-reply-today-yesterday-previous7-v1"
TITLE = "### 每台设备：今日、昨日及前7天回复率"
FIELDS = ("发送窗口数", "回复窗口数")
PERIOD_LABELS = {"today": "今日", "yesterday": "昨日", "previous7": "前7天"}


def aggregate(headers, rows, server_date):
    today = date.fromisoformat(server_date)
    periods = {
        "today": [today.isoformat()],
        "yesterday": [(today - timedelta(days=1)).isoformat()],
        "previous7": [(today - timedelta(days=i)).isoformat() for i in range(7, 0, -1)],
    }
    headers, rows = tuple(headers), tuple(rows)
    if headers.count("卡槽") != 1:
        raise ValueError("设备期间回复率需要唯一卡槽列")
    all_days = tuple(dict.fromkeys(day for days in periods.values() for day in days))
    columns = {}
    for day in all_days:
        for field in FIELDS:
            label = day[5:] + field
            if headers.count(label) > 1:
                raise ValueError(f"设备期间回复率发现重复指标列：{label}")
            columns[day, field] = headers.index(label) if label in headers else None
    slot = headers.index("卡槽")
    keys = tuple(device_key(row[slot] if slot < len(row) else None) for row in rows)
    sizes = Counter(keys)
    buckets = {key: {day: {field: 0 for field in FIELDS} for day in all_days} for key in sizes}
    invalid = {key: Counter() for key in sizes}
    # One source-row pass. Do not average source percentages or deduplicate phones.
    for key, row in zip(keys, rows):
        for (day, field), position in columns.items():
            if position is None:
                buckets[key][day][field] = None
                continue
            value = count_value(row[position] if position < len(row) else None)
            if value is None:
                invalid[key][day, field] += 1
                buckets[key][day][field] = None
            elif buckets[key][day][field] is not None:
                buckets[key][day][field] += value
    result = []
    for key in sorted(sizes, key=device_sort_key):
        metrics = {}
        for period, days in periods.items():
            totals = {}
            for field in FIELDS:
                values = [buckets[key][day][field] for day in days]
                totals[field] = None if any(value is None for value in values) else sum(values)
            missing = [day[5:] + field for day in days for field in FIELDS if columns[day, field] is None]
            invalid_cells = sum(invalid[key][day, field] for day in days for field in FIELDS)
            send, reply = (totals[field] for field in FIELDS)
            reasons = []
            if missing:
                reasons.append("缺少日期指标列")
            if invalid_cells:
                reasons.append("源计数缺失或无效")
            if send == 0:
                reasons.append("发送窗口数为0")
            metrics[period] = {
                "send": send, "reply": reply,
                "ratePct": reply / send * 100 if send and reply is not None else None,
                "missingColumns": missing, "invalidCountCells": invalid_cells,
                "unknownReason": "、".join(reasons) or None,
            }
        result.append({
            "device": key, "room": LABELS.get(room_for_device(key), "未配置"),
            "recordCount": sizes[key], "periods": metrics,
        })
    return {
        "contract": CONTRACT, "serverDate": today.isoformat(), "periods": periods,
        "roomConfigSha256": CONFIG_SHA256, "sourceRows": len(rows),
        "deviceCount": len(sizes) - int(UNKNOWN_DEVICE in sizes),
        "unrecognizedDeviceRows": sizes.get(UNKNOWN_DEVICE, 0), "rows": result,
    }


def build(headers, rows, run_id, server_date, source_sha256, observed_at):
    return {**aggregate(headers, rows, server_date), "runId": run_id,
            "sourceSha256": source_sha256, "observedAt": observed_at}


def verify_rows(summary, headers, rows, run_id, server_date, source_sha256, observed_at):
    expected = build(headers, rows, run_id, server_date, source_sha256, observed_at)
    if summary != expected:
        raise ValueError("设备今日/昨日/前7天统计、日期或来源绑定核验失败")


def verify_source(summary, source):
    """Read the raw workbook afresh, not generated totals or cached percentages."""
    from openpyxl import load_workbook
    source = Path(source)
    before = sha256(source)
    if before != summary.get("sourceSha256"):
        raise ValueError("设备期间回复率原始源表SHA256不一致")
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        iterator = book.active.iter_rows(values_only=True)
        headers = next(iterator)
        verify_rows(summary, headers, tuple(iterator), summary["runId"], summary["serverDate"],
                    before, summary["observedAt"])
    finally:
        book.close()
    if sha256(source) != before:
        raise ValueError("设备期间回复率核验期间源表变化")


def report_lines(summary):
    if not summary:  # Historical reports do not invent absent evidence.
        return []
    if summary.get("contract") != CONTRACT or summary.get("roomConfigSha256") != CONFIG_SHA256:
        raise ValueError("设备期间回复率契约或机房配置不一致")
    periods = summary["periods"]
    lines = ["", TITLE, "",
             f"- 今日（平台业务日期）：{periods['today'][0]}；昨日：{periods['yesterday'][0]}；前7天：{periods['previous7'][0]}至{periods['previous7'][-1]}，不含今天。",
             "- 今日为本轮导出时累计；昨日及前7天取本轮源表对应自然日数据，不是昨日同期快照。",
             "- 回复率=同设备、同期间回复窗口数合计÷发送窗口数合计。前7天先分别求和再相除，不平均每日或号码百分比，不使用总短信量代替窗口数。",
             f"- 包含全部状态的{summary['sourceRows']}条源记录，不按手机号去重；{summary['deviceCount']}台可识别设备，另有{summary['unrecognizedDeviceRows']}条未识别设备记录。字母设备前缀保留，未列机房显示未配置。",
             "- 缺少日期列、任一源计数缺失/非法或发送窗口数为0时，相关回复率显示未知，不用部分日期或0补齐。", "",
             "| 设备号 | 机房 | 今日回复率 | 昨日回复率 | 前7天回复率 | 未知原因 |",
             "| --- | --- | ---: | ---: | ---: | --- |"]
    for row in summary["rows"]:
        metrics = row["periods"]
        rates = ["未知" if metrics[key]["ratePct"] is None else f"{metrics[key]['ratePct']:.3f}%"
                 for key in PERIOD_LABELS]
        reasons = "；".join(f"{label}：{metrics[key]['unknownReason']}"
                           for key, label in PERIOD_LABELS.items() if metrics[key]["unknownReason"])
        lines.append("| " + " | ".join([row["device"], row["room"], *rates, reasons or "无"]) + " |")
    return lines


def supplement_report(state_path, output):
    """Explicit offline report addition; never overwrite a delivered report/state."""
    state_path, output = Path(state_path).resolve(), Path(output).resolve()
    state_hash = sha256(state_path)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    acquire = state.get("stages", {}).get("acquire", {})
    original = Path(state.get("stages", {}).get("finalize", {}).get("report", "")).resolve()
    if (state.get("status") != "success" or acquire.get("serverDateVerified") is not True
            or not original.is_file() or state["runId"] not in original.name):
        raise ValueError("只接受精确成功轮次、服务器日期已核验的原完整MD")
    if output.suffix.lower() != ".md":
        raise ValueError("设备期间补充报告必须另存为MD")
    if output == original or output.exists() or output.with_suffix(".json").exists():
        raise ValueError("禁止覆盖历史报告或已有补充交付")
    archive, source = Path(acquire["sourceArchive"]), Path(acquire["processingSource"])
    expected = acquire["sourceSha256"]
    if sha256(archive) != expected or sha256(source) != expected:
        raise ValueError("本轮归档与处理源SHA256未绑定成功轮次")
    raw = load_scan_cache(acquire["scanCache"], expected_source=source)
    export = json.loads(Path(acquire["exportReceipt"]).read_text(encoding="utf-8"))
    if export.get("serverDate") != acquire["serverDate"] or not export.get("downloadedAt"):
        raise ValueError("本轮导出日期或采集时间未确认")
    summary = build(raw["headers"], raw["rows"], state["runId"], acquire["serverDate"],
                    expected, export["downloadedAt"])
    verify_source(summary, archive)
    original_hash = sha256(original)
    text = original.read_text(encoding="utf-8")
    anchor = "### 设备0回复"
    if TITLE in text or text.count(anchor) != 1:
        raise ValueError("原完整报告已含设备期间统计或章节位置不唯一")
    section = ("\n".join(report_lines(summary)) + "\n\n"
               "- 本文件为原成功轮次的离线统计补充，不是新的后台导出；原业务结果、Excel和历史报告不变。\n\n")
    rendered = text.replace(anchor, section + anchor, 1)
    if rendered.replace(section, "", 1) != text:
        raise ValueError("离线补充影响了原完整报告内容")
    if (sha256(state_path) != state_hash or sha256(original) != original_hash
            or sha256(archive) != expected or sha256(source) != expected):
        raise ValueError("离线补充期间来源状态或原报告变化")
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("x", encoding="utf-8") as stream:
        stream.write(rendered)
        stream.flush()
        os.fsync(stream.fileno())
    if output.read_text(encoding="utf-8") != rendered:
        raise ValueError("设备期间完整MD独立回读失败")
    proof = {"status": "passed", "mode": "offline_report_supplement", "runId": state["runId"],
             "state": str(state_path), "stateSha256": state_hash,
             "originalReport": str(original), "originalReportSha256": original_hash,
             "source": str(archive), "sourceSha256": expected, "output": str(output),
             "outputSha256": sha256(output), "summary": summary,
             "independentSourceReadbackPassed": True, "independentReportReadbackPassed": True,
             "networkAccess": False, "databaseAccess": False, "businessWrites": 0}
    with output.with_suffix(".json").open("x", encoding="utf-8") as stream:
        json.dump(proof, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.flush()
        os.fsync(stream.fileno())
    if json.loads(output.with_suffix(".json").read_text(encoding="utf-8")) != proof:
        raise ValueError("设备期间补充回执独立回读失败")
    return {key: proof[key] for key in ("status", "mode", "runId", "output", "outputSha256",
            "independentSourceReadbackPassed", "independentReportReadbackPassed", "businessWrites")} | {
                "deviceCount": summary["deviceCount"], "sourceRows": summary["sourceRows"],
                "serverDate": summary["serverDate"], "periods": summary["periods"]}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--state", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(supplement_report(args.state, args.output), ensure_ascii=False))
