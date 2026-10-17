#!/usr/bin/env python3
"""Read-only deadline/carrier reconciliation. Never generates write candidates."""
import argparse
from collections import Counter
from datetime import date
import json
from pathlib import Path

from openpyxl import load_workbook
from opensms_generate_task_b_candidates import parse_date, sha256, deadline_status, phone_text

ISOLATION = {"快过期隔离区C", "快过期隔离区T"}
EXPIRED = {"过期下架卡C", "过期下架卡T"}


def review(workbook_path, server_date):
    workbook_path = Path(workbook_path).resolve()
    reference = date.fromisoformat(server_date)
    before = sha256(workbook_path)
    book = load_workbook(workbook_path, read_only=True, data_only=False)
    rows, invalid, mismatches = [], [], []
    scanned = 0
    try:
        iterator = book["sheet1"].iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        for name in ("电话", "运营商", "截止日", "过期状态", "卡槽"):
            if headers.count(name) != 1:
                raise ValueError(f"截止日核查字段缺失或重复: {name}")
        for number, values in enumerate(iterator, 2):
            scanned += 1
            row = dict(zip(headers, values))
            raw = row.get("截止日")
            cutoff = parse_date(raw, book.epoch)
            status = str(row.get("过期状态") or "").strip()
            carrier = str(row.get("运营商") or "").strip()
            item = {"row": number, "phone": phone_text(row.get("电话")), "carrier": carrier,
                    "slot": str(row.get("卡槽") or ""), "cutoff": cutoff.isoformat() if cutoff else "",
                    "remainingDays": (cutoff - reference).days if cutoff else None,
                    "storedStatus": status, "expectedStatus": deadline_status(cutoff, reference)}
            blank = raw is None or (isinstance(raw, str) and not raw.strip())
            if cutoff is None and (not blank or status or carrier in ISOLATION | EXPIRED):
                invalid.append({**item, "reason": "截止日缺失或无效，不能确认分区"})
            if item["storedStatus"] != item["expectedStatus"]:
                mismatches.append(item)
            if cutoff and ((carrier in ISOLATION and item["remainingDays"] > 1)
                           or (carrier in EXPIRED and item["remainingDays"] > 0)):
                rows.append({**item, "reason": "截止日未达到当前分区转入条件"})
    finally:
        book.close()
    if sha256(workbook_path) != before:
        raise ValueError("截止日核查期间工作簿改变，拒绝混用")
    return {"scope": "provided_workbook_only_not_live_backend", "businessWrites": 0,
            "sourceWorkbook": str(workbook_path), "sourceSha256": before, "serverDate": server_date,
            "scannedRows": scanned, "prematureCarrierCount": len(rows), "rows": rows,
            "byCarrier": dict(Counter(item["carrier"] for item in rows)),
            "byCutoff": dict(sorted(Counter(item["cutoff"] for item in rows).items())),
            "invalidDeadlineCount": len(invalid), "invalidDeadlines": invalid,
            "statusMismatchCount": len(mismatches), "statusMismatches": mismatches}


def report_lines(result):
    def safe(value):
        return str(value if value is not None else "未知").replace("|", "\\|").replace("\n", " ").replace("\r", " ")
    lines = ["### 截止日与现有分区核对", "",
             f"- 依据：工作簿‘截止日’减网页服务器日期 {result['serverDate']}；不使用激活后天数、koo是否2天或颜色。",
             f"- 已在隔离/过期区但未达到转入日期：{result['prematureCarrierCount']}行；截止日缺失/无效待核查：{result['invalidDeadlineCount']}行；状态文字与截止日不一致：{result['statusMismatchCount']}行。",
             "- 本项只核对文档分区，不产生运营商或通知写入；不能据此认定原始转入时间。恢复目标未确认时不得自动猜区。"]
    if result['rows'] or result['invalidDeadlines'] or result['statusMismatches']:
        lines += ["", "| 电话 | 文档运营商 | 卡槽 | 截止日 | 剩余天数 | 文档状态 | 截止日计算状态 | 原因 | 主表行 |",
                  "| --- | --- | --- | --- | ---: | --- | --- | --- | ---: |"]
        listed = {}
        for item in result['statusMismatches']:
            listed[item['row']] = {**item, "reason": "过期状态文字与截止日不一致"}
        for item in result['rows'] + result['invalidDeadlines']:
            listed[item['row']] = item
        for item in sorted(listed.values(), key=lambda value: value['row']):
            lines.append("| " + " | ".join(safe(item[key]) for key in
                         ("phone", "carrier", "slot", "cutoff", "remainingDays", "storedStatus", "expectedStatus", "reason", "row")) + " |")
    return lines


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', required=True, type=Path)
    parser.add_argument('--server-date', required=True)
    parser.add_argument('--output-dir', required=True, type=Path)
    args = parser.parse_args()
    result = review(args.workbook, args.server_date)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    data = args.output_dir / '截止日分区核查.json'
    report = args.output_dir / '截止日分区核查.md'
    for file, content in ((data, json.dumps(result, ensure_ascii=False, indent=2)),
                          (report, '# 截止日分区离线核查\n\n非新鲜导出；仅核查指定历史文件，未改后台或原文档。\n\n' + '\n'.join(report_lines(result)))):
        with file.open('x', encoding='utf-8') as handle:
            handle.write(content + '\n')
    print(json.dumps({key: value for key, value in result.items() if key not in {'rows', 'invalidDeadlines', 'statusMismatches'}} | {'report': str(report), 'receipt': str(data)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
