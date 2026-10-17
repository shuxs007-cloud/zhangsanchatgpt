#!/usr/bin/env python3
"""Programmatic, concurrent audit for a completed OpenSMS patrol run.

The auditor is intentionally read-only. It consumes the files and JSON receipts
already produced by a patrol, validates them with deterministic rules, and
emits one compact summary plus one detailed JSON receipt. It never opens WPS,
logs in to OpenSMS, changes carriers, or executes Task B/Task C/Task D/Task E/Task F.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import multiprocessing
import os
import re
import sqlite3
import sys
import threading
import time
import unicodedata
import zipfile
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor, as_completed
from contextlib import ExitStack
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta, timezone
from functools import lru_cache
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Iterable
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from opensms_temporary_carrier_hold import temporary_carrier_held, isolation_carrier_held

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'services/opensms_document_api/pipeline'))
from archived_system_identity import index_phones, match_archived
from google_room_sheets import select_room_sheets
from task_f_renewal_dates import validate_renewal


PASS = "passed"
WARN = "warning"
FAIL = "failed"
SKIP = "skipped"
PENDING = "pending"

REQUIRED_BALANCE_GROUPS = ["A1", "A12", "A11", "B1", "B11", "B12", "C2", "A3", "B3", "B1V", "B11V", "B12V"]
REQUIRED_RESULT_SHEETS = {
    "sheet1",
    "处理汇总",
    "未匹配卡槽",
    "电信商_OPR预警",
    "谷歌有-系统无",
    "谷歌文档中有激活日期-系统无法使用",
    "系统无-谷歌表有",
    "已激活未上架",
    "未上架机器汇总",
    "A11连续低于3%",
}
GOOGLE_ONLY_RESULT_SHEETS = {
    "未匹配卡槽", "电信商_OPR预警", "谷歌有-系统无",
    "谷歌文档中有激活日期-系统无法使用",
}
GOOGLE_ONLY_MAIN_HEADERS = {"Google電信商/價格", "OPR匹配预警"}
NO_GOOGLE_DEPENDENCY = {
    "enabled": False, "status": "skipped", "reason": "user_disabled",
    "unused_port_device_scope": "system_source_only",
}
CHATR_GOOGLE_DATE_SCOPE = "chatr_blank_activation_and_cutoff_only"
CHATR_GOOGLE_DATE_SOURCE = "current_run_system_then_fresh_google"
CHATR_GOOGLE_RESOLUTION_SOURCE = "system_export_then_frozen_local_or_chatr_google"
CHATR_GOOGLE_COMMANDS = {"google_chatr_download", "chatr_google_date_snapshot"}
CHATR_GOOGLE_SOURCE_URL = "https://docs.google.com/spreadsheets/d/1GiA9uR_2IQ4xntP_aWvCSaWfmC6WVQ7SrnXsy8YyyvM/export?format=xlsx"
OBSOLETE_RENAMED_RESULT_SHEETS = {
    "ko未激活表",
    "未使用端口",
    "所有谷歌有-系统没有",
    "设备缺失端口",
    "Google有日期_主表缺失",
    "谷歌文档中有激活日期-系统没有",
}
REQUIRED_MAIN_HEADERS = {
    "电话",
    "运营商",
    "过期状态",
    "koo是否2天",
    "激活后天数",
    "使用天数",
    "opr",
    "卡槽",
    "连续低于x%",
    "连续0回复",
    "开卡日期",
    "激活日期",
    "截止日",
    "换号状态",
    "机房",
    "10日回复率",
    "3日发送窗口数",
    "3日回复率",
    "4日发送窗口数",
    "4日回复率",
    "回复率持续降低",
}
RATE_COLUMNS_AFTER_TEN_DAY = [
    "3日发送窗口数",
    "3日回复率",
    "4日发送窗口数",
    "4日回复率",
    "回复率持续降低",
]
FORMULA_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
NUMERIC_TEXT_PATTERN = re.compile(
    r"^[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?:[eE][+-]?\d+)?$"
)
REPLY_RATE_SEND_THRESHOLD = 50
REPLY_RATE_SEND_GATED_SOURCE_PATTERN = re.compile(
    r"^(?:A1|A11|A12|B1|B11|B12|B1V|B11V|B12V)(?:-|\s|$)",
    re.IGNORECASE,
)
REPLY_RATE_SEND_GATED_AUTO_ROUTES = {
    "a1_kini_latest_lt4_to_b12",
    "a_zone_lt4_to_a12",
    "a_zone_gte4_to_a1",
    "b1_b11_b12_today_gte50_zero_to_e",
    "chatr_to_e",
}
CHATR_ZERO_REPLY_SOURCE_PATTERN = re.compile(r"^(?:A1|A11|A12|A3)(?:-|\s|$)", re.IGNORECASE)
TASK_B_RETAINED_SOURCE_CARRIERS = ("B14-C快下保留卡", "B14-T快下保留卡")
TASK_B_RENEWAL_SOURCE_CARRIERS = ("快过期续费C", "快过期续费T")

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"m": MAIN_NS, "r": DOC_REL_NS, "p": PKG_REL_NS}


def resolve_worker_count(value: str | int | None) -> int:
    raw = str(value if value is not None else "auto").strip().lower()
    ceiling = max(1, min(16, os.cpu_count() or 1))
    if raw in {"auto", "0"}:
        # v61 same-input 1/2/4/8 tests after removing redundant XML work.
        # Four balanced these checks best; eight added contention, not speed.
        return min(4, ceiling)
    if raw == "max":
        return ceiling
    try:
        workers = int(raw)
    except ValueError as exc:
        raise ValueError("--workers必须是正整数或auto") from exc
    if workers < 1 or workers > ceiling:
        raise ValueError(f"--workers必须为1..{ceiling}（受CPU及16工作器上限约束）")
    return workers


@dataclass
class CheckResult:
    name: str
    category: str
    status: str
    message: str
    duration_ms: int = 0
    thread: str = ""
    details: dict[str, Any] = field(default_factory=dict)


@dataclass
class AuditTask:
    name: str
    category: str
    fn: Callable[[], CheckResult]
    # Only explicitly enumerated, pure file checks may cross this boundary.
    # Never serialize closures, open workbooks, DB connections or site clients.
    cpu_call: tuple[str, tuple[Any, ...]] | None = None


class ActivityMeter:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self.active = 0
        self.maximum = 0
        self.threads: set[str] = set()

    def enter(self) -> None:
        with self._lock:
            self.active += 1
            self.maximum = max(self.maximum, self.active)
            self.threads.add(threading.current_thread().name)

    def leave(self) -> None:
        with self._lock:
            self.active -= 1


class PathResolver:
    def __init__(self, config_path: Path, delivery_dir: Path | None) -> None:
        self.config_path = config_path.resolve()
        self.config_dir = self.config_path.parent
        self.delivery_dir = delivery_dir.resolve() if delivery_dir else None

    def resolve(self, value: str | Path | None, required: bool = True) -> Path | None:
        if value is None:
            if required:
                raise FileNotFoundError("路径未配置")
            return None
        raw = Path(value).expanduser()
        candidates: list[Path] = []
        if self.delivery_dir:
            # A nested same-run attempt may legitimately share a basename
            # with its archived failed predecessor. Never substitute that
            # predecessor for the exact path already inside this delivery.
            if raw.is_absolute() and raw.resolve().is_relative_to(self.delivery_dir):
                candidates.append(raw)
            basename = raw.name
            candidates.extend(
                [
                    self.delivery_dir / basename,
                    self.delivery_dir / "receipts" / basename,
                    self.delivery_dir / "previews" / basename,
                ]
            )
        candidates.append(raw)
        if not raw.is_absolute():
            candidates.append(self.config_dir / raw)
        for candidate in candidates:
            if candidate.exists():
                return candidate.resolve()
        if self.delivery_dir and raw.name:
            matches = [
                item
                for item in self.delivery_dir.rglob(raw.name)
                if item.is_file() and not item.name.startswith("._")
            ]
            if len(matches) == 1:
                return matches[0].resolve()
        if required:
            raise FileNotFoundError(f"文件不存在: {value}")
        return None

    def glob(self, pattern: str) -> list[Path]:
        roots = [self.config_dir]
        if self.delivery_dir:
            roots.extend([self.delivery_dir, self.delivery_dir / "receipts"])
        found: dict[str, Path] = {}
        for root in roots:
            if not root.exists():
                continue
            for item in root.glob(pattern):
                if item.is_file() and not item.name.startswith("._"):
                    found[str(item.resolve())] = item.resolve()
        return sorted(found.values(), key=lambda item: item.stat().st_mtime)


def now_iso() -> str:
    return datetime.now().astimezone().isoformat()


def read_json(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise ValueError(f"JSON顶层必须是对象: {path}")
    return value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_hash(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    return text if re.fullmatch(r"[0-9a-f]{64}", text) else None


def parse_iso(value: Any) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    if text.endswith("Z"):
        text = text[:-1] + "+00:00"
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


def status_result(
    name: str,
    category: str,
    errors: list[str],
    warnings: list[str],
    details: dict[str, Any],
    success_message: str,
) -> CheckResult:
    if errors:
        return CheckResult(name, category, FAIL, "；".join(errors), details=details)
    if warnings:
        return CheckResult(name, category, WARN, "；".join(warnings), details=details)
    return CheckResult(name, category, PASS, success_message, details=details)


def candidate_count(data: dict[str, Any]) -> int:
    for key in ("count", "candidateCount", "total"):
        value = data.get(key)
        if value is not None:
            return int(value)
    for key in ("candidates", "rows", "phones"):
        value = data.get(key)
        if isinstance(value, list):
            return len(value)
    return 0


def candidate_phones(data: dict[str, Any]) -> set[str]:
    if data.get("count") is not None and int(data["count"]) == 0:
        return set()
    rows = data.get("candidates") or data.get("rows") or []
    result: set[str] = set()
    for row in rows:
        if isinstance(row, str):
            phone = row
        elif isinstance(row, dict):
            phone = row.get("phone") or row.get("电话") or row.get("cardNo")
        else:
            continue
        digits = re.sub(r"\D", "", str(phone or ""))
        if digits:
            result.add(digits)
    return result


def xlsx_sheet_map(archive: zipfile.ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
    result: dict[str, str] = {}
    sheets = workbook.find("m:sheets", NS)
    if sheets is None:
        return result
    for sheet in sheets:
        rel_id = sheet.attrib.get(f"{{{DOC_REL_NS}}}id")
        target = targets.get(rel_id or "")
        if not target:
            continue
        if target.startswith("/"):
            archive_path = target.lstrip("/")
        else:
            archive_path = str(PurePosixPath("xl") / target)
            if archive_path.startswith("xl/xl/"):
                archive_path = archive_path[3:]
        result[sheet.attrib["name"]] = archive_path
    return result


def xlsx_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    return [
        "".join(node.text or "" for node in item.iter(f"{{{MAIN_NS}}}t"))
        for item in root.findall("m:si", NS)
    ]


def column_index(reference: str) -> int:
    letters = re.match(r"[A-Z]+", reference.upper())
    if not letters:
        return -1
    result = 0
    for char in letters.group(0):
        result = result * 26 + ord(char) - ord("A") + 1
    return result - 1


def short_phone(value: Any) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-10:] if len(digits) > 10 else digits


def xlsx_main_phone_carriers(path: Path) -> dict[str, list[str]]:
    with zipfile.ZipFile(path) as archive:
        sheet_map = xlsx_sheet_map(archive)
        if not sheet_map:
            raise ValueError("结果表没有工作表")
        shared = xlsx_shared_strings(archive)
        root = ET.fromstring(archive.read(next(iter(sheet_map.values()))))
        rows = root.findall(".//m:sheetData/m:row", NS)
        if not rows:
            raise ValueError("结果表主表没有数据")

        cell_tag = f"{{{MAIN_NS}}}c"

        def row_values(row: ET.Element, columns: set[int] | None = None) -> dict[int, str]:
            # The carrier readback needs two columns, not every date/metric cell.
            # Still parse the entire XML and preserve missing/duplicate semantics.
            values = {}
            for cell in row.iterfind(cell_tag):
                index = column_index(cell.attrib.get("r", ""))
                if columns is None or index in columns:
                    values[index] = cell_value(cell, shared)
            return values

        header_values = row_values(rows[0])
        header_positions = {value: index for index, value in header_values.items()}
        phone_col = header_positions.get("电话")
        carrier_col = header_positions.get("运营商")
        if phone_col is None or carrier_col is None:
            raise ValueError("结果表主表缺少电话或运营商列")
        selected_columns = {phone_col, carrier_col}
        result: dict[str, list[str]] = {}
        for row in rows[1:]:
            values = row_values(row, selected_columns)
            phone = short_phone(values.get(phone_col, ""))
            if not phone:
                continue
            result.setdefault(phone, []).append(values.get(carrier_col, ""))
        supplemental_path = sheet_map.get("后台更新未在源表")
        if supplemental_path:
            supplemental_root = ET.fromstring(archive.read(supplemental_path))
            supplemental_rows = supplemental_root.findall(".//m:sheetData/m:row", NS)
            if supplemental_rows:
                supplemental_header = row_values(supplemental_rows[0])
                supplemental_positions = {value: index for index, value in supplemental_header.items()}
                supplemental_phone_col = supplemental_positions.get("电话")
                supplemental_carrier_col = supplemental_positions.get("最新运营商")
                if supplemental_phone_col is None or supplemental_carrier_col is None:
                    raise ValueError("补充状态表缺少电话或最新运营商列")
                selected_columns = {supplemental_phone_col, supplemental_carrier_col}
                for row in supplemental_rows[1:]:
                    values = row_values(row, selected_columns)
                    phone = short_phone(values.get(supplemental_phone_col, ""))
                    if phone and phone not in result:
                        result[phone] = [values.get(supplemental_carrier_col, "")]
        return result


def cell_value(cell: ET.Element, shared: list[str]) -> str:
    cell_type = cell.attrib.get("t")
    value = cell.find(f"{{{MAIN_NS}}}v")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.iter(f"{{{MAIN_NS}}}t"))
    if value is None or value.text is None:
        return ""
    if cell_type == "s":
        try:
            return shared[int(value.text)]
        except (ValueError, IndexError):
            return ""
    return value.text


def numeric_text_status(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value).strip()
    if not normalized:
        return "not_numeric_text"
    percent = normalized.endswith("%")
    numeric_text = normalized[:-1].strip() if percent else normalized
    if not NUMERIC_TEXT_PATTERN.fullmatch(numeric_text):
        return "not_numeric_text"
    canonical = numeric_text.replace(",", "")
    unsigned = re.sub(r"^[+-]", "", canonical)
    integer_part = re.split(r"[.eE]", unsigned, maxsplit=1)[0]
    if len(integer_part) > 1 and integer_part.startswith("0"):
        return "protected_leading_zero"
    mantissa = re.split(r"[eE]", canonical, maxsplit=1)[0]
    digits = re.sub(r"^[+-]", "", mantissa).replace(".", "").lstrip("0") or "0"
    if len(digits) > 15:
        return "protected_precision_risk"
    try:
        parsed = float(canonical)
    except ValueError:
        return "not_numeric_text"
    if not math.isfinite(parsed):
        return "not_numeric_text"
    if not percent and parsed.is_integer() and abs(parsed) > 9_007_199_254_740_991:
        return "protected_precision_risk"
    return "convertible"


def inspect_result_workbook(path: Path, server_date: str | None,
                            google_enabled: bool = True, google_difference_enabled: bool = False) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"XLSX压缩成员损坏: {bad_member}")
        names = set(archive.namelist())
        required_package = {"[Content_Types].xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"}
        missing_package = sorted(required_package - names)
        if missing_package:
            raise ValueError(f"XLSX结构缺失: {', '.join(missing_package)}")
        sheet_map = xlsx_sheet_map(archive)
        shared = xlsx_shared_strings(archive)
        if not sheet_map:
            raise ValueError("结果表没有工作表")
        main_name = next(iter(sheet_map))
        main_root = ET.fromstring(archive.read(sheet_map[main_name]))
        header_row = main_root.find(".//m:sheetData/m:row", NS)
        headers = [cell_value(cell, shared) for cell in header_row.findall("m:c", NS)] if header_row is not None else []
        max_row = 0
        for row in main_root.findall(".//m:sheetData/m:row", NS):
            try:
                max_row = max(max_row, int(row.attrib.get("r", "0")))
            except ValueError:
                continue
        formula_errors: list[dict[str, str]] = []
        numeric_text_counts = {
            "convertible": 0,
            "protected_leading_zero": 0,
            "protected_precision_risk": 0,
            "protected_text_format_mismatch": 0,
        }
        numeric_text_examples: list[dict[str, str]] = []
        # Private to this single independent scan: no cross-run/file cache.
        # Repeated strings (0, group names, dates) have identical classification.
        classify_text = lru_cache(maxsize=4096)(numeric_text_status)
        cell_tag = f"{{{MAIN_NS}}}c"
        value_tag = f"{{{MAIN_NS}}}v"
        formula_tag = f"{{{MAIN_NS}}}f"
        for sheet_name, archive_path in sheet_map.items():
            # The main sheet was already fully parsed above for headers/rows.
            root = main_root if sheet_name == main_name else ET.fromstring(archive.read(archive_path))
            for cell in root.iter(cell_tag):
                cell_type = cell.attrib.get("t")
                if cell_type == "e":
                    value = cell.find(value_tag)
                    error = value.text if value is not None else "unknown"
                    formula_errors.append({"sheet": sheet_name, "cell": cell.attrib.get("r", ""), "error": error})
                if cell_type not in {"s", "str", "inlineStr"} or cell.find(formula_tag) is not None:
                    continue
                value = cell_value(cell, shared)
                status = classify_text(value)
                if status == "not_numeric_text":
                    continue
                numeric_text_counts[status] += 1
                if len(numeric_text_examples) < 30:
                    numeric_text_examples.append(
                        {
                            "sheet": sheet_name,
                            "cell": cell.attrib.get("r", ""),
                            "value": value,
                            "status": "preserved_as_text",
                        }
                    )

    header_positions = {value: index for index, value in enumerate(headers)}
    required_sheets = REQUIRED_RESULT_SHEETS if google_enabled else REQUIRED_RESULT_SHEETS - GOOGLE_ONLY_RESULT_SHEETS
    if google_difference_enabled:
        required_sheets = required_sheets | {"谷歌有-系统无"}
    missing_sheets = sorted(required_sheets - set(sheet_map))
    unexpected_renamed_sheets = sorted(OBSOLETE_RENAMED_RESULT_SHEETS & set(sheet_map))
    missing_headers = sorted(REQUIRED_MAIN_HEADERS - set(headers))
    order_errors: list[str] = []
    if "谷歌有-系统无" in sheet_map and "ko系统有-未激活" in sheet_map:
        names = list(sheet_map)
        if names.index("ko系统有-未激活") != names.index("谷歌有-系统无") + 1:
            order_errors.append("ko系统有-未激活未紧邻谷歌有-系统无右侧")
    if google_enabled or google_difference_enabled:
        names = list(sheet_map)
        if "系统无-谷歌表有" in names and "谷歌有-系统无" in names:
            if names.index("谷歌有-系统无") != names.index("系统无-谷歌表有") + 1:
                order_errors.append("谷歌有-系统无未紧邻系统无-谷歌表有右侧")
    if not google_enabled:
        forbidden_sheets = (GOOGLE_ONLY_RESULT_SHEETS | {"所有谷歌有-系统没有"}) & set(sheet_map)
        if google_difference_enabled:
            forbidden_sheets -= {"谷歌有-系统无"}
        forbidden_headers = GOOGLE_ONLY_MAIN_HEADERS & set(headers)
        if forbidden_sheets:
            order_errors.append("Google依赖关闭但仍存在关联工作表：" + "、".join(sorted(forbidden_sheets)))
        if forbidden_headers:
            order_errors.append("Google依赖关闭但仍存在关联列：" + "、".join(sorted(forbidden_headers)))
    redundant_date_headers = {"系统导出激活日期", "系统导出截止日"} & set(headers)
    if redundant_date_headers:
        order_errors.append("主表仍包含已删除冗余日期列：" + "、".join(sorted(redundant_date_headers)))
    if "换号状态" in header_positions and header_positions["换号状态"] != 3:
        order_errors.append("换号状态不是第4列")
    if "koo是否2天" in header_positions and "激活后天数" in header_positions:
        if header_positions["激活后天数"] != header_positions["koo是否2天"] + 1:
            order_errors.append("激活后天数未紧跟koo是否2天")
    if '使用天数' in header_positions:
        if header_positions.get('激活后天数', -2) + 1 != header_positions['使用天数']:
            order_errors.append('使用天数未紧跟激活后天数')
        if '需要保留时间' in header_positions and header_positions['需要保留时间'] != header_positions['使用天数'] + 1:
            order_errors.append('需要保留时间未紧跟使用天数')
    if "10日回复率" in header_positions:
        start = header_positions["10日回复率"] + 1
        if headers[start : start + len(RATE_COLUMNS_AFTER_TEN_DAY)] != RATE_COLUMNS_AFTER_TEN_DAY:
            order_errors.append("10日回复率后的五列顺序不正确")

    daily_send_headers = [
        value for value in headers if re.fullmatch(r"\d{2}-\d{2}发送窗口数", value)
    ]
    continuity_error = None
    if len(daily_send_headers) < 4:
        continuity_error = "动态日期发送列少于4天"
    elif server_date:
        reference = datetime.strptime(server_date, "%Y-%m-%d").date()
        parsed_dates = []
        for value in daily_send_headers[:4]:
            month, day = map(int, value[:5].split("-"))
            choices = []
            for year in (reference.year - 1, reference.year, reference.year + 1):
                try:
                    date_value = reference.replace(year=year, month=month, day=day)
                except ValueError:
                    continue
                choices.append(date_value)
            parsed_dates.append(min(choices, key=lambda item: abs((item - reference).days)))
        if any(parsed_dates[index] - parsed_dates[index + 1] != timedelta(days=1) for index in range(3)):
            continuity_error = "最近4个统计日期不是连续自然日"

    return {
        "sheets": list(sheet_map),
        "sheet_count": len(sheet_map),
        "main_sheet": main_name,
        "main_headers": headers,
        "main_rows": max(max_row - 1, 0),
        "missing_sheets": missing_sheets,
        "unexpected_renamed_sheets": unexpected_renamed_sheets,
        "missing_headers": missing_headers,
        "order_errors": order_errors,
        "formula_errors": formula_errors,
        "numeric_text_review": {
            "matches": True,
            "conversion_enabled": False,
            "preserved_text_numeric_cells": sum(numeric_text_counts.values()),
            "examples": numeric_text_examples,
            "policy": "preserve_source_cell_types",
        },
        "daily_send_headers": daily_send_headers,
        "date_continuity_error": continuity_error,
    }


def check_xlsx(
    name: str,
    path: Path,
    expected_hash: str | None = None,
    inspect_result: bool = False,
    server_date: str | None = None,
    google_enabled: bool = True,
    google_difference_enabled: bool = False,
) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    details: dict[str, Any] = {"path": str(path)}
    if not path.is_file():
        return CheckResult(name, "files", FAIL, f"文件不存在: {path}", details=details)
    details["size"] = path.stat().st_size
    if path.stat().st_size == 0:
        errors.append("文件大小为0")
    actual_hash = sha256_file(path)
    details["sha256"] = actual_hash
    if expected_hash and actual_hash != expected_hash:
        errors.append("SHA-256与本轮配置不一致")
    try:
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
            if bad_member:
                errors.append(f"XLSX压缩成员损坏: {bad_member}")
            if "[Content_Types].xml" not in archive.namelist():
                errors.append("不是完整的XLSX文件")
    except (zipfile.BadZipFile, OSError) as exc:
        errors.append(f"XLSX无法读取: {exc}")
    if inspect_result and not errors:
        try:
            workbook = inspect_result_workbook(path, server_date, google_enabled, google_difference_enabled)
            details["workbook"] = workbook
            if workbook["missing_sheets"]:
                errors.append("缺少工作表: " + ", ".join(workbook["missing_sheets"]))
            if workbook["unexpected_renamed_sheets"]:
                errors.append("旧工作表仍存在: " + ", ".join(workbook["unexpected_renamed_sheets"]))
            if workbook["missing_headers"]:
                errors.append("缺少主表列: " + ", ".join(workbook["missing_headers"]))
            errors.extend(workbook["order_errors"])
            if workbook["formula_errors"]:
                errors.append(f"发现公式错误单元格 {len(workbook['formula_errors'])} 个")
            if workbook["date_continuity_error"]:
                errors.append(workbook["date_continuity_error"])
        except (ET.ParseError, KeyError, ValueError, zipfile.BadZipFile) as exc:
            errors.append(f"工作簿结构检查失败: {exc}")
    return status_result(name, "files", errors, warnings, details, "文件完整性和结构通过")


def check_export_receipt(path: Path, expected_hash: str, source_path: Path) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    data = read_json(path)
    saved = data.get("saved") or data.get("afterLatest") or {}
    details = {
        "path": str(path),
        "started_at": data.get("startedAt"),
        "server_date": data.get("serverDate"),
        "saved_size": saved.get("size"),
        "task_id_present": "任务ID" in str(data.get("bodyText") or "") or bool(data.get("taskId")),
    }
    if not data.get("startedAt"):
        errors.append("导出回执缺少开始时间")
    if not data.get("serverDate"):
        errors.append("导出回执缺少服务器日期")
    if not saved.get("size"):
        errors.append("导出回执没有完成文件")
    elif int(saved["size"]) != source_path.stat().st_size:
        errors.append("导出回执文件大小与源表不一致")
    if sha256_file(source_path) != expected_hash:
        errors.append("源表哈希与配置不一致")
    if saved.get("saveAsError") or data.get("saved", {}).get("saveAsError"):
        warnings.append("浏览器saveAs失败，但已观察到下载文件")
    return status_result("后台导出回执", "freshness", errors, warnings, details, "本轮导出回执通过")


def check_google_receipt(path: Path, expected_hash: str, association_path: Path) -> CheckResult:
    errors: list[str] = []
    data = read_json(path)
    details = {
        "path": str(path),
        "downloaded_at": data.get("downloadedAt"),
        "source_url": data.get("sourceUrl"),
        "sha256": data.get("sha256"),
    }
    if not data.get("downloadedAt"):
        errors.append("Google表回执缺少下载时间")
    if "docs.google.com/spreadsheets" not in str(data.get("sourceUrl") or ""):
        errors.append("Google表来源地址不正确")
    if normalize_hash(data.get("sha256")) != expected_hash:
        errors.append("Google表回执哈希与配置不一致")
    if sha256_file(association_path) != expected_hash:
        errors.append("Google表文件哈希与配置不一致")
    return status_result("Google关联表回执", "freshness", errors, [], details, "本轮Google关联表通过")


def check_server_dates(paths: list[Path], expected_date: str | None) -> CheckResult:
    errors: list[str] = []
    records = [read_json(path) for path in paths]
    by_requested: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        by_requested.setdefault(str(record.get("requestedPath") or ""), []).append(record)
    selected: dict[str, dict[str, Any]] = {}
    for requested in ("/day", "/phoneConfig"):
        matches = by_requested.get(requested, [])
        if not matches:
            errors.append(f"缺少{requested}服务器日期回执")
            continue
        if len(matches) > 1:
            matches.sort(key=lambda item: parse_iso(item.get("checkedAt")) or datetime.min.replace(tzinfo=timezone.utc))
        selected[requested] = matches[-1]
    dates = {record.get("serverDate") for record in selected.values() if record.get("serverDate")}
    if len(dates) != 1:
        errors.append("/day与/phoneConfig服务器日期不一致")
    actual_date = next(iter(dates), None)
    if expected_date and actual_date != expected_date:
        errors.append("服务器日期与文档作业日期不一致")
    details = {
        "paths": [str(path) for path in paths],
        "day": selected.get("/day", {}).get("serverDate"),
        "phone_config": selected.get("/phoneConfig", {}).get("serverDate"),
    }
    return status_result("服务器日期", "freshness", errors, [], details, "两个后台服务器日期一致")


def unwrap_card_receipt(data: dict[str, Any]) -> dict[str, Any]:
    return data.get("cardSlotReport") or data


def check_card_receipt(path: Path, google_enabled: bool = True) -> CheckResult:
    data = unwrap_card_receipt(read_json(path))
    errors: list[str] = []
    warnings: list[str] = []
    checks = data.get("checks") or {}
    failed = [name for name, value in checks.items() if value is not True]
    failed.extend(str(value) for value in data.get("failed_checks") or [])
    if data.get("result") not in {"成功", "success"}:
        errors.append("卡槽关联处理结果不是成功")
    if not checks:
        errors.append("卡槽关联回执没有程序检查项")
    if failed:
        errors.append("卡槽关联检查未通过: " + ", ".join(sorted(set(failed))))
    process = data.get("process") or {}
    if process.get('room_configuration') is not None:
        room_review = (data.get('verification') or {}).get('room_configuration_review') or {}
        if checks.get('最新机房配置及全部机房列独立回读') is not True or room_review.get('matches') is not True:
            errors.append('最新机房配置缺少独立回读')
        try:
            from opensms_verify_card_slot_fast import review_room_configuration
            book = load_workbook(process['output_path'], read_only=True, data_only=True)
            try:
                room_summary = dict(book['处理汇总'].iter_rows(max_col=2, values_only=True))
                actual_room_review = review_room_configuration(book, process, room_summary)
                if actual_room_review != room_review:
                    errors.append('最新机房配置独立审计重算与回执不一致')
            finally:
                book.close()
        except Exception as exc:
            errors.append(f'最新机房配置独立审计失败：{exc}')
    difference = process.get("google_system_missing_sheet")
    ko_proof = process.get("ko_unactivated_sheet")
    if ko_proof:
        ko_review = (data.get("verification") or {}).get("ko_unactivated_sheet_review") or {}
        if (checks.get("ko系统有-未激活独立核验") is not True or ko_review.get("matches") is not True
                or any(ko_review.get(key) != value for key, value in ko_proof.items())):
            errors.append("ko系统有-未激活缺少独立逐行回读")
    if difference:
        review = (data.get("verification") or {}).get("google_system_missing_sheet_review") or {}
        if (checks.get("所有谷歌有系统没有独立差集核验") is not True or review.get("matches") is not True
                or any(review.get(key) != value for key, value in difference.items())):
            errors.append("Google系统差集未通过独立来源、逐行及样式核验")
    if process.get('creation_usage_days') is not None:
        usage = process['creation_usage_days']
        usage_review = ((data.get('verification') or {}).get('creation_usage_days_review') or {})
        if (checks.get('使用天数独立重算') is not True or usage_review.get('matches') is not True
                or any(usage_review.get(key) != value for key, value in usage.items())):
            errors.append('使用天数源创建时间独立重算未通过')
    device_daily = process.get("device_daily_statistics")
    retention = process.get("card_retention_report")
    if retention is not None:
        retention_review = ((data.get("verification") or {}).get("card_retention_report_review") or {})
        if (checks.get("需要保留时间及卡保留全量明细") is not True
                or retention_review.get("matches") is not True
                or any(retention_review.get(key) != value for key, value in retention.items())):
            errors.append("任务A卡保留全量明细及申请时间独立核验未通过")
        source = Path(retention.get("source_path") or "")
        if not source.is_file() or sha256_file(source) != retention.get("source_sha256"):
            errors.append("任务A卡保留来源SHA256不一致")
    if device_daily is not None:
        device_review = ((data.get("verification") or {}).get("device_daily_statistics_review") or {})
        if (checks.get("设备每日统计独立重算") is not True or device_review.get("matches") is not True
                or device_daily.get("source_rows") != process.get("main_rows")
                or any(device_daily.get(key) != device_review.get(key)
                       for key in ("source_rows", "device_count", "date_count", "unrecognized_device_rows"))):
            errors.append("设备每日统计源行数、设备数、日期数或独立重算未通过")
    if not google_enabled:
        dependency = process.get("google_dependency") or {}
        dependency_review = ((data.get("verification") or {}).get("google_dependency_review") or {})
        if (dependency.get("enabled") is not False or dependency_review.get("enabled") is not False
                or any(dependency.get(key) != value or dependency_review.get(key) != value
                for key, value in NO_GOOGLE_DEPENDENCY.items())
                or dependency_review.get("matches") is not True):
            errors.append("Google依赖关闭及系统源设备范围未通过独立复核")
    if google_enabled and process.get("main_rows") != process.get("matched_rows"):
        errors.append("主表行数与匹配行数不一致")
    if google_enabled and int(process.get("unmatched_slot_count") or 0) != 0:
        warnings.append(f"存在未匹配卡槽 {process.get('unmatched_slot_count')} 个")
    unused_review = ((data.get("verification") or {}).get("unused_port_review") or {})
    if unused_review.get("matches") is not True:
        errors.append("系统无-谷歌表有独立复核未通过")
    system_kooda_review = ((data.get("verification") or {}).get("system_kooda_source_review") or {})
    if system_kooda_review.get("matches") is not True or system_kooda_review.get("source") != "system":
        errors.append("Kooda本轮系统ICCID、激活日期、截止日及开卡留空独立复核未通过")
    resolution = process.get("date_resolution") or {}
    resolution_review = ((data.get("verification") or {}).get("date_resolution_review") or {})
    if resolution:
        scoped_chatr = (process.get("google_chatr_date_fallback") or {}).get("enabled") is True
        expected_source = ("system_export_then_frozen_local_or_google" if google_enabled else
                           CHATR_GOOGLE_RESOLUTION_SOURCE if scoped_chatr else "system_export_then_frozen_local")
        if (resolution.get("source") != expected_source
                or resolution_review.get("matches") is not True
                or resolution.get("field_counts") != resolution_review.get("field_counts")
                or resolution.get("cache") != resolution_review.get("cache")
                or resolution.get("new_query_results_apply_next_export") is not True
                or resolution.get("redundant_system_date_columns_removed") is not True):
            errors.append("日期逐字段优先级、冻结缓存或冗余列删除独立复核未通过")
        chatr_dates = resolution.get("chatr_google_dates") or {}
        if chatr_dates.get("enabled") is True:
            chatr_review = (data.get("verification") or {}).get(
                "google_chatr_date_fallback_review" if scoped_chatr else "chatr_google_dates_review") or {}
            if (resolution_review.get("chatr_google_dates") != chatr_dates
                    or chatr_review.get("matches") is not True
                    or any(chatr_review.get(key) != value for key, value in chatr_dates.items())
                    or chatr_dates.get("scope") != CHATR_GOOGLE_DATE_SCOPE
                    or (not scoped_chatr and chatr_dates.get("match_rule") != "unique_normalized_slot")
                    or (scoped_chatr and (process.get("google_chatr_date_fallback") != chatr_dates
                                          or chatr_dates.get("source") != CHATR_GOOGLE_DATE_SOURCE))):
                errors.append("Chatr Google空日期补缺范围与来源独立复核未通过")
    task_f_red_review = ((data.get("verification") or {}).get("task_f_missing_activation_red_review") or {})
    if task_f_red_review.get("matches") is not True:
        errors.append("Kooda缺失激活日期整行标红或恢复后取消红色复核未通过")
    task_f_new_card_review = ((data.get("verification") or {}).get("task_f_missing_dates_new_card_review") or {})
    if task_f_new_card_review.get("matches") is not True:
        errors.append("Kooda激活日期与截止日双空白时未换号新卡独立复核未通过")
    details = {
        "path": str(path),
        "checks_total": len(checks),
        "checks_passed": sum(value is True for value in checks.values()),
        "main_rows": process.get("main_rows"),
        "matched_rows": process.get("matched_rows"),
        "server_date": process.get("expiry_reference_date"),
        "unmatched_slots": process.get("unmatched_slot_count"),
        "unused_port_device_count": process.get("unused_port_device_count"),
        "unused_port_count": process.get("unused_port_count"),
        "system_missing_port_count": process.get("system_missing_port_count"),
        "registration_failed_port_count": process.get("registration_failed_port_count"),
        "system_kooda_source_review": system_kooda_review,
        "task_f_kooda_date_review": system_kooda_review,
        "task_f_missing_activation_red_review": task_f_red_review,
        "task_f_missing_dates_new_card_review": task_f_new_card_review,
        "date_resolution_review": resolution_review,
    }
    return status_result("卡槽关联完整复核", "document", errors, warnings, details, "卡槽关联程序复核全部通过")


def check_dispatch_receipt(
    path: Path,
    source_hash: str,
    association_hash: str,
    result_hash: str,
    carrier_refresh_configured: bool = False,
    google_enabled: bool = True,
) -> CheckResult:
    data = read_json(path)
    errors: list[str] = []
    warnings: list[str] = []
    job = data.get("job") or {}
    inputs = data.get("inputs") or {}
    outputs = job.get("output_files") or {}
    if data.get("result") not in {"成功", "success"}:
        errors.append("文档调度结果不是成功")
    if data.get("processing_mode") not in {"remote_api", "local_fallback", "local"}:
        errors.append("文档处理模式未知")
    if data.get("processing_mode") == "local_fallback":
        warnings.append("局域网接口不可用，本轮使用本地回退")
    if normalize_hash((inputs.get("export4day") or {}).get("sha256")) != source_hash:
        errors.append("调度回执源表哈希不一致")
    if google_enabled and normalize_hash((inputs.get("association") or {}).get("sha256")) != association_hash:
        errors.append("调度回执Google表哈希不一致")
    if not google_enabled and inputs.get("association"):
        errors.append("Google依赖关闭但文档调度仍带关联输入")
    remote_result_hash = normalize_hash((outputs.get("result") or {}).get("sha256"))
    output_hash = normalize_hash(data.get("output_sha256"))
    cell_type_policy = data.get("cell_type_policy") or {}
    post_refresh_transition = False
    if output_hash and output_hash != result_hash:
        if carrier_refresh_configured:
            post_refresh_transition = True
        else:
            errors.append("文档调度结果哈希与最终结果不一致")
    if remote_result_hash and remote_result_hash != result_hash:
        if carrier_refresh_configured:
            post_refresh_transition = True
        else:
            errors.append("远端结果哈希与下载结果不一致")
    if cell_type_policy.get("numeric_text_conversion_enabled") is not False:
        errors.append("文档调度未明确关闭文本型数字转换")
    if job and job.get("status") != "succeeded":
        errors.append("远端文档作业未成功")
    job_checks = job.get("checks") or {}
    if job_checks and not all(value is True for value in job_checks.values()):
        errors.append("远端文档作业检查存在失败")
    details = {
        "path": str(path),
        "processing_mode": data.get("processing_mode"),
        "api_base_url": data.get("api_base_url"),
        "job_id": job.get("job_id"),
        "job_status": job.get("status"),
        "server_date": (inputs.get("server_date") or job.get("server_date")),
        "remote_checks": len(job_checks),
        "cell_type_policy": cell_type_policy,
        "post_refresh_transition": post_refresh_transition,
    }
    success_message = (
        "文档调度原始结果与写入后运营商状态刷新链路已衔接"
        if post_refresh_transition
        else "文档接口及下载哈希通过"
    )
    return status_result("文档接口调度", "document", errors, warnings, details, success_message)


def check_carrier_refresh(
    config: dict[str, Any],
    resolver: PathResolver,
    result_path: Path,
) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    actual_writes = 0
    for group in config.get("groups") or []:
        if not group.get("write"):
            continue
        write_path = resolver.resolve(group.get("write"))
        write_data = read_json(write_path)
        before = write_data.get("before") or write_data.get("prep") or {}
        actual_writes += int(before.get("eligibleCount") or 0)

    refresh_value = config.get("carrierRefreshReceipt")
    readback_value = config.get("carrierRefreshReadbackReceipt")
    details: dict[str, Any] = {
        "actual_writes": actual_writes,
        "refresh_configured": bool(refresh_value),
        "readback_configured": bool(readback_value),
        "result_path": str(result_path),
    }
    if actual_writes > 0 and (not refresh_value or not readback_value):
        errors.append("存在运营商实际写入，但缺少写入后后台批量回读或文档刷新回执")
        return status_result("运营商最新状态回写文档", "document", errors, warnings, details, "运营商最新状态已回写文档")
    if not refresh_value and not readback_value:
        return status_result("运营商最新状态回写文档", "document", errors, warnings, details, "本轮无运营商写入，无需刷新文档")

    refresh_path = resolver.resolve(refresh_value)
    readback_path = resolver.resolve(readback_value)
    refresh = read_json(refresh_path)
    readback = read_json(readback_path)
    details.update({
        "refresh_path": str(refresh_path),
        "readback_path": str(readback_path),
        "candidate_count": refresh.get("candidateCount"),
        "backend_resolved": refresh.get("backendResolvedCount"),
        "workbook_matched": refresh.get("matchedWorkbookRowCount"),
        "updated_cells": refresh.get("updatedCellCount"),
        "already_current": refresh.get("alreadyCurrentCount"),
    })
    if readback.get("passed") is not True:
        errors.append("写入后后台批量只读回读未通过")
    if readback.get("silentBrowserMode") is not True or readback.get("browserForegroundActivationAttempted") is not False:
        errors.append("写入后后台批量回读未证明全程静默且未抢占焦点")
    for key, label in (
        ("missingCount", "后台缺失"),
        ("duplicateBackendCount", "后台重复"),
        ("mismatchCount", "后台目标不一致"),
    ):
        if int(readback.get(key) or 0) != 0:
            errors.append(f"{label}{readback.get(key)}条")
    if refresh.get("passed") is not True:
        errors.append("运营商最新状态文档刷新回执未通过")
    for key, label in (
        ("backendMissingCount", "刷新回执后台缺失"),
        ("backendDuplicateCount", "刷新回执后台重复"),
        ("missingInWorkbookCount", "表内缺失"),
        ("duplicateWorkbookCount", "表内重复"),
        ("stillStaleCount", "刷新后仍不一致"),
    ):
        if int(refresh.get(key) or 0) != 0:
            errors.append(f"{label}{refresh.get(key)}条")
    candidate_count_value = int(refresh.get("candidateCount") or 0)
    matched_count = int(refresh.get("mergedWorkbookCoveredCount") or refresh.get("matchedWorkbookRowCount") or 0)
    updated_count = int(refresh.get("updatedCellCount") or 0)
    already_count = int(refresh.get("alreadyCurrentCount") or 0)
    supplemental_count = int(refresh.get("supplementalRowCount") or 0)
    if candidate_count_value != matched_count or candidate_count_value != updated_count + already_count + supplemental_count:
        errors.append("文档刷新候选、匹配、改值和原值一致数量不守恒")
    configured_output = resolver.resolve(refresh.get("outputPath"), required=False)
    if configured_output != result_path.resolve():
        errors.append("文档刷新回执未指向最终卡槽结果表")

    workbook_rows = xlsx_main_phone_carriers(result_path)
    mismatches: list[dict[str, Any]] = []
    for row in readback.get("rows") or []:
        phone = short_phone(row.get("phone"))
        expected = str(row.get("currentCarrierName") or "")
        actual = workbook_rows.get(phone) or []
        if len(actual) != 1 or actual[0] != expected:
            mismatches.append({"phone": phone, "expected": expected, "actual": actual})
    details["independent_workbook_mismatch_count"] = len(mismatches)
    details["independent_workbook_mismatches"] = mismatches[:30]
    if mismatches:
        errors.append(f"最终结果表仍有{len(mismatches)}个运营商值与后台回读不一致")
    return status_result("运营商最新状态回写文档", "document", errors, warnings, details, "运营商最新状态已回写文档并独立核验")


def check_analysis_receipt(path: Path) -> CheckResult:
    data = read_json(path)
    errors: list[str] = []
    details = {
        "path": str(path),
        "rows": data.get("rows"),
        "latest_date": data.get("latestDate"),
        "current_zero": data.get("currentZero"),
    }
    if not data.get("latestDate"):
        errors.append("分析回执缺少最新统计日期")
    if int(data.get("rows") or 0) <= 0:
        errors.append("分析回执号码行数无效")
    overall = data.get("overall") or {}
    if overall.get("send") is None or overall.get("reply") is None:
        errors.append("分析回执缺少整体发送或回复统计")
    if data.get("scopeComparison"):
        try:
            from opensms_scope_comparison import verify
            verify(data["scopeComparison"])
            details["scopeComparisonVerified"] = True
            device_periods = data["scopeComparison"].get("deviceReplyPeriods")
            if device_periods is not None:
                from opensms_device_reply_periods import verify_source
                verify_source(device_periods, data["scopeComparison"]["inputs"]["source"])
                details["deviceReplyPeriodsVerified"] = True
                details["deviceReplyPeriodsDeviceCount"] = device_periods["deviceCount"]
        except (ValueError, KeyError, TypeError, OSError) as exc:
            errors.append(f"整体/机房同期对比核验失败：{exc}")
    return status_result("源报表分析", "analysis", errors, [], details, "源报表分析回执通过")


def check_balance_receipt(path: Path) -> CheckResult:
    data = read_json(path)
    errors: list[str] = []
    warnings: list[str] = []
    required = list(data.get("requiredGroups") or REQUIRED_BALANCE_GROUPS)
    if required != REQUIRED_BALANCE_GROUPS:
        errors.append("余量检查分组与固定范围不一致")
    records = data.get("records") or []
    by_group: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        by_group.setdefault(str(record.get("group") or ""), []).append(record)
    for group in REQUIRED_BALANCE_GROUPS:
        matches = by_group.get(group, [])
        if len(matches) != 1:
            errors.append(f"余量分组{group}记录数为{len(matches)}")
            continue
        record = matches[0]
        total = record.get("dailyTotal")
        remaining = record.get("dailyRemaining")
        percent = record.get("percentRemaining")
        if total is None or float(total) <= 0:
            warnings.append(f"{group}当日总量无效，比例未知")
        elif remaining is None or percent is None:
            errors.append(f"{group}缺少余量或比例")
        else:
            expected = float(remaining) / float(total) * 100
            if abs(expected - float(percent)) > 1e-6:
                errors.append(f"{group}剩余比例重算不一致")
    reply_rates = data.get("replyRates") or []
    rates_by_group: dict[str, list[dict[str, Any]]] = {}
    for record in reply_rates:
        rates_by_group.setdefault(str(record.get("group") or ""), []).append(record)
    stat_dates: set[str] = set()
    for group in REQUIRED_BALANCE_GROUPS:
        matches = rates_by_group.get(group, [])
        if len(matches) != 1:
            errors.append(f"回复率分组{group}记录数为{len(matches)}")
            continue
        item = matches[0]
        if item.get("status") != "ok":
            errors.append(f"{group}按日回复率查询失败")
        if item.get("metric") != "firstReplyRate":
            errors.append(f"{group}使用的不是首次回复率")
        if item.get("statDate"):
            stat_dates.add(str(item["statDate"]))
    if len(stat_dates) > 1:
        warnings.append("各组最新回复率统计日期不一致，已按每组实时查询的最新行保留")
    if data.get("missingRequiredGroups"):
        errors.append("存在缺失余量分组")
    if data.get("systemIssues"):
        warnings.append(f"余量检查报告了{len(data['systemIssues'])}个系统问题")
    if "carrierReplyRates" in data:
        try:
            from opensms_carrier_reply_comparison import verify_comparison
            comparison = verify_comparison(data)
            if comparison['currentUnknownCount']:
                warnings.append(f"{comparison['currentUnknownCount']}个运营商当前首次回复率未知")
        except (ValueError, KeyError, TypeError, OSError) as exc:
            errors.append(f"运营商昨日同期回复率独立核验失败: {exc}")
    if (data.get("autoQuota") or {}).get("enabled") is not False:
        errors.append("自动加量必须保持关闭")
    if data.get("fallback"):
        warnings.append("余量数据使用了回退来源")
    details = {
        "path": str(path),
        "groups": REQUIRED_BALANCE_GROUPS,
        "reply_rate_dates": sorted(stat_dates),
        "low20": data.get("low20") or [],
        "low10": data.get("low10") or [],
        "auto_quota_enabled": (data.get("autoQuota") or {}).get("enabled"),
    }
    return status_result("余量与按日回复率", "balance", errors, warnings, details, "十二组余量和首次回复率通过")


def check_no_write_stage_proof(group, candidate, write, verify, validated_count, resolver):
    """An empty write set is not a successful verification of any backend row."""
    errors = []
    before = write.get("before") or {}
    if not (
        write.get("noWrite") is True
        and write.get("actualWriteCount") == 0
        and int(before.get("eligibleCount") or 0) == 0
        and write.get("updateResponses") == []
        and write.get("after") is None
        and write.get("independentReadbackRequired") is False
        and write.get("independentReadbackPerformed") is False
        and validated_count == 0
        and verify.get("total") == 0
        and verify.get("passed") is None
        and verify.get("verificationStatus") == "not_required"
        and verify.get("status") == "no_candidate"
        and verify.get("independentReadbackPerformed") is False
    ):
        return ["空写入快速路径契约不完整，不能省略独立复核"]
    proof_path = resolver.resolve(group.get("stageReadbackReceipt"), required=False)
    if not proof_path or not proof_path.is_file():
        return ["空写入组缺少后续新鲜阶段全量回读"]
    if sha256_file(proof_path) != group.get("stageReadbackSha256"):
        errors.append("阶段回读SHA-256不一致")
    proof = read_json(proof_path)
    rows = proof.get("rows") or []
    if not (
        proof.get("passed") is True
        and proof.get("silentBrowserMode") is True
        and proof.get("browserForegroundActivationAttempted") is False
        and proof.get("candidateCount") == proof.get("backendValidCount") == len(rows)
        and int(proof.get("failureCount") or 0) == 0
    ):
        errors.append("阶段全量回读未通过完整性校验")
    written_at, checked_at = parse_iso(write.get("ts")), parse_iso(proof.get("checkedAt"))
    if not written_at or not checked_at or checked_at < written_at:
        errors.append("阶段回读不是写入判定后的新鲜回读")
    original = {short_phone(phone) for phone in candidate_phones(candidate)}
    already_rows = before.get("alreadyTarget") or []
    excluded_rows = before.get("nonSource") or []
    observed = already_rows + excluded_rows
    observed_phones = [short_phone(row.get("phone")) for row in observed]
    if len(original) != candidate_count(candidate) or len(observed_phones) != len(set(observed_phones)) or set(observed_phones) != original:
        errors.append("空写入来源判定未精确覆盖所有候选")
    if int(before.get("missingCount") or 0) or int(before.get("duplicateBackendCount") or 0):
        errors.append("空写入来源判定存在缺失或重复")
    by_phone = {}
    for row in rows:
        by_phone.setdefault(short_phone(row.get("phone")), []).append(row)
    target = str((write.get("target") or {}).get("name") or "")
    already = {short_phone(row.get("phone")) for row in already_rows}
    for item in observed:
        phone = short_phone(item.get("phone"))
        hits = by_phone.get(phone) or []
        if len(hits) != 1 or hits[0].get("backendMatchCount") != 1 or hits[0].get("passed") is not True:
            errors.append("阶段回读未唯一覆盖空写入候选")
            continue
        row = hits[0]
        source_id = (item.get("row") or {}).get("id")
        if not source_id or str(row.get("backendId")) != str(source_id):
            errors.append("阶段回读后台身份与提交前判定不一致")
        if phone in already and (not target or row.get("liveCarrierName") != target):
            errors.append("已在目标区的号码在阶段回读中不再匹配目标")
    return errors


def check_group(title: str, group: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    candidate_path = resolver.resolve(group.get("candidate"))
    candidate = read_json(candidate_path)
    count = candidate_count(candidate)
    details: dict[str, Any] = {
        "title": title,
        "candidate_path": str(candidate_path),
        "candidate_count": count,
        "actual_writes": 0,
        "successful_writes": 0,
        "failed_writes": 0,
    }
    validated_count = count
    if group.get("validated"):
        validated_path = resolver.resolve(group.get("validated"))
        validated = read_json(validated_path)
        validated_count = candidate_count(validated)
        details["validated_path"] = str(validated_path)
        details["validated_count"] = validated_count
        original_count = validated.get("originalCount")
        if original_count is not None and int(original_count) != count:
            errors.append("验证候选数与原候选数不一致")
    write_data: dict[str, Any] | None = None
    if group.get("write"):
        write_path = resolver.resolve(group.get("write"))
        write_data = read_json(write_path)
        details["write_path"] = str(write_path)
        before = write_data.get("before") or write_data.get("prep") or {}
        after = write_data.get("after") or {}
        before_count = before.get("candidateCount")
        if before_count is not None and int(before_count) != count:
            errors.append("写入回执候选数与候选文件不一致")
        actual = int(before.get("eligibleCount") or 0)
        already = int(before.get("alreadyTargetCount") or 0)
        target_count = int(after.get("targetCount") or 0)
        success = max(min(target_count - already, actual), 0)
        details.update(
            {
                "actual_writes": actual,
                "successful_writes": success,
                "failed_writes": max(actual - success, 0),
                "already_target": already,
            }
        )
        target_name = str((write_data.get("target") or {}).get("name") or "")
        if actual > 0 and target_name.startswith("E-"):
            backup_path = resolver.resolve(write_data.get("originBackupPath"), required=False)
            if not backup_path:
                errors.append("转入E前缺少原运营商备份")
    elif count > 0:
        errors.append("有候选但缺少写入回执")

    if group.get("verify"):
        verify_path = resolver.resolve(group.get("verify"))
        verify = read_json(verify_path)
        details["verify_path"] = str(verify_path)
        no_write_path = (write_data or {}).get("noWrite") is True
        if no_write_path:
            errors.extend(check_no_write_stage_proof(group, candidate, write_data, verify, validated_count, resolver))
            details["per_group_independent_readback_performed"] = False
            details["stage_readback_path"] = group.get("stageReadbackReceipt")
        elif verify.get("passed") is not True:
            errors.append("独立只读复核未通过")
        for key, label in (
            ("stillNotTargetCount", "仍非目标"),
            ("missingCount", "复核缺失"),
            ("duplicateBackendCount", "复核重复"),
        ):
            if int(verify.get(key) or 0) != 0:
                errors.append(f"{label}{verify.get(key)}条")
        verify_total = verify.get("total")
        # The per-group independent readback receives the validated write set
        # only.  Rows already at the target are covered by the later final
        # all-candidate readback and must not be double-counted here.
        expected_verify_total = validated_count
        if verify_total is not None and int(verify_total) != expected_verify_total:
            errors.append("独立复核总数与后台有效候选数不一致")
        details["verified_target"] = int(verify.get("targetCount") or 0)
    elif count > 0:
        errors.append("有候选但缺少独立复核回执")

    if not group.get("write") and not group.get("verify") and group.get("result"):
        warnings.append(str(group["result"]))
    unidentified = int(candidate.get("unidentifiedOriginCount") or 0)
    if unidentified:
        warnings.append(f"{unidentified}条E区记录缺少可追溯原运营商，已按规则保留")
    return status_result(title, "carrier", errors, warnings, details, "候选、写入和独立回读一致")


def check_candidate_disjoint(groups: list[dict[str, Any]], resolver: PathResolver) -> CheckResult:
    owners: dict[str, list[str]] = {}
    inspected = 0
    superseded_groups: list[str] = []
    for group in groups:
        title = str(group.get("title") or "未命名组")
        if title.startswith(("任务B", "任务C", "任务D")):
            continue
        if group.get("superseded") is True:
            superseded_groups.append(title)
            continue
        path = resolver.resolve(group.get("candidate"))
        data = read_json(path)
        for phone in candidate_phones(data):
            owners.setdefault(phone, []).append(title)
            inspected += 1
    duplicates = {phone: titles for phone, titles in owners.items() if len(titles) > 1}
    errors = [f"发现{len(duplicates)}个号码同轮进入多个运营商候选组"] if duplicates else []
    details = {
        "candidate_memberships": inspected,
        "unique_phones": len(owners),
        "duplicate_phone_count": len(duplicates),
        "duplicate_groups": sorted({title for titles in duplicates.values() for title in titles}),
        "superseded_groups_excluded": superseded_groups,
    }
    return status_result("候选互斥性", "carrier", errors, [], details, "本轮运营商候选互斥")


def audit_a1_kini_candidate(candidate: dict[str, Any]) -> list[str]:
    """Independently recompute the narrow Kini rule without network access."""
    route = str(candidate.get("route") or "")
    if re.match(r"^task_[bcde]_", route) or str(candidate.get("task") or "").upper() in {"B", "C", "D", "E"}:
        return []
    rows = candidate.get("rows") or candidate.get("candidates") or []
    is_a1 = lambda value: bool(re.match(r"^A1(?:-|\s|$)", str(value or ""), re.I))
    source_of = lambda row: str(row.get("原运营商") or row.get("运营商") or row.get("sourceCarrierName") or "")
    is_kini = lambda value: bool(re.search(r"\bkini\b", str(value or ""), re.I))
    if route != "a1_kini_latest_lt4_to_b12":
        if any(is_a1(source_of(row)) and is_kini(row.get("opr")) for row in rows):
            return ["A1＋Kini错误进入其它普通路由，未使用独立B12规则"]
        return []
    sources = candidate.get("sourceCarrierNames") or []
    latest = str(candidate.get("latestDate") or "")
    if (candidate.get("a1KiniRuleVersion") != "a1-kini-latest-send50-rate-lt4-to-b12-v1"
            or candidate.get("a1KiniReplyRateThresholdExclusive") != 0.04
            or candidate.get("replyRateLatestSendThresholdInclusive") != 50
            or candidate.get("replyRateWindow") != "latest_day"
            or not re.fullmatch(r"\d{2}-\d{2}", latest)
            or len(sources) != 1 or not is_a1(sources[0])
            or not re.match(r"^B12(?:-|\s|$)", str(candidate.get("targetCarrierName") or ""), re.I)):
        return ["A1＋Kini独立候选门禁、A1来源或标准B12目标不正确"]
    errors = []
    for row in rows:
        raw = [row.get("发送窗口数"), row.get("回复窗口数")]
        try:
            if any(isinstance(value, bool) or not isinstance(value, (str, int, float)) or str(value).strip() == "" for value in raw):
                raise ValueError("missing window count")
            send, reply = map(float, raw)
            valid = (math.isfinite(send) and math.isfinite(reply)
                     and send.is_integer() and reply.is_integer()
                     and send >= 50 and 0 <= reply <= send
                     and int(reply) * 100 < int(send) * 4)
        except (TypeError, ValueError, OverflowError):
            valid = False
        if (source_of(row) != sources[0] or not is_kini(row.get("opr"))
                or re.search(r"chatr", str(row.get("opr") or ""), re.I)
                or row.get("latestDate") != latest or row.get("统计日期") != latest or not valid):
            errors.append(f"A1＋Kini候选不符合最新日发送>=50且回复率严格<4%: {row.get('phone') or row.get('电话') or '未知'}")
    return errors


def check_reply_rate_send_gate(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    inspected_paths: set[Path] = set()
    inspected_routes: set[str] = set()
    inspected_rows = 0
    invalid_rows: list[dict[str, Any]] = []
    invalid_chatr_rows: list[dict[str, Any]] = []

    paths: list[Path] = []
    for group in config.get("groups") or []:
        if group.get("superseded") is True:
            continue
        path = resolver.resolve(group.get("candidate"), required=False)
        if path is not None:
            paths.append(path)

    configured_manifest = config.get("carrierRoutingManifest")
    manifest_path = resolver.resolve(configured_manifest, required=False) if configured_manifest else None
    if manifest_path is not None:
        manifest = read_json(manifest_path)
        if int(manifest.get("replyRateLatestSendThresholdInclusive") or 0) != REPLY_RATE_SEND_THRESHOLD:
            errors.append("运营商候选清单的回复型流转发送门槛不是>=50")
        expected_groups = {
            "A1", "A11", "A12", "B1", "B11", "B12", "B1V", "B11V", "B12V",
        }
        if set(manifest.get("replyRateSendGateGroups") or []) != expected_groups:
            errors.append("运营商候选清单的回复型发送门槛来源组缺少A区、B区或B区V组")
        if set(manifest.get("replyRateSendGateRoutes") or []) != REPLY_RATE_SEND_GATED_AUTO_ROUTES:
            errors.append("运营商候选清单的回复型发送门槛路由不正确")
        if set(manifest.get("chatrZeroReplySendGateGroups") or []) != {"A1", "A11", "A12", "A3"}:
            errors.append("运营商候选清单的Chatr零回复来源组不是A1/A11/A12/A3")
        if not str(manifest.get("aZoneChatrRestTarget") or "").upper().startswith("E-C"):
            errors.append("运营商候选清单的A区Chatr零回复目标不是E-C休息区")
        for route in (manifest.get("outputs") or {}).values():
            path = resolver.resolve(route.get("path"), required=False)
            if path is not None:
                paths.append(path)

    for path in paths:
        if path in inspected_paths:
            continue
        inspected_paths.add(path)
        candidate = read_json(path)
        errors.extend(f"{path.name}: {error}" for error in audit_a1_kini_candidate(candidate))
        route = str(candidate.get("route") or "")
        if route not in REPLY_RATE_SEND_GATED_AUTO_ROUTES:
            continue
        inspected_routes.add(route)
        if float(candidate.get("replyRateLatestSendThresholdInclusive") or 0) != REPLY_RATE_SEND_THRESHOLD:
            errors.append(f"{path.name}缺少回复判断前最新日发送窗口数>=50门禁")
        if route == "chatr_to_e" and not str(candidate.get("targetCarrierName") or "").upper().startswith("E-C"):
            errors.append(f"{path.name}的Chatr零回复目标不是E-C休息区")
        for row in candidate.get("rows") or candidate.get("candidates") or []:
            source = str(row.get("原运营商") or row.get("运营商") or row.get("sourceCarrierName") or "")
            if route == "chatr_to_e":
                reply = row.get("回复窗口数")
                opr = str(row.get("opr") or row.get("OPR") or "")
                try:
                    reply_count = float(reply)
                except (TypeError, ValueError):
                    reply_count = -1
                if (
                    not CHATR_ZERO_REPLY_SOURCE_PATTERN.match(source)
                    or "chatr" not in opr.casefold()
                    or reply_count != 0
                ):
                    invalid_chatr_rows.append({
                        "phone": str(row.get("phone") or row.get("电话") or ""),
                        "source": source,
                        "opr": opr,
                        "reply": reply,
                    })
            elif not REPLY_RATE_SEND_GATED_SOURCE_PATTERN.match(source):
                continue
            inspected_rows += 1
            try:
                send_count = float(row.get("发送窗口数"))
            except (TypeError, ValueError):
                send_count = -1
            if send_count < REPLY_RATE_SEND_THRESHOLD:
                invalid_rows.append({
                    "phone": str(row.get("phone") or row.get("电话") or ""),
                    "source": source,
                    "send": None if send_count < 0 else send_count,
                    "route": route,
                })
    if invalid_rows:
        errors.append(f"发现{len(invalid_rows)}个回复型候选未满足最新日发送>=50")
    if invalid_chatr_rows:
        errors.append(f"发现{len(invalid_chatr_rows)}个Chatr零回复候选不符合来源、OPR或零回复条件")
    details = {
        "threshold_inclusive": REPLY_RATE_SEND_THRESHOLD,
        "inspected_candidate_files": len(inspected_paths),
        "inspected_routes": sorted(inspected_routes),
        "inspected_send_gated_rows": inspected_rows,
        "invalid_candidate_count": len(invalid_rows),
        "invalid_sample": invalid_rows[:20],
        "invalid_chatr_candidate_count": len(invalid_chatr_rows),
        "invalid_chatr_sample": invalid_chatr_rows[:20],
    }
    return status_result(
        "回复型流转发送窗口数门禁",
        "carrier",
        errors,
        [],
        details,
        "仅回复型流转执行最新日发送窗口数>=50门禁，B1/B11十日发送规则未被误加门槛",
    )


check_source_transfer_send_gate = check_reply_rate_send_gate


def check_rest_zone_entry_enabled(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    manifest_path = resolver.resolve(config.get("carrierRoutingManifest"), required=False)
    if manifest_path is None:
        return status_result(
            "休息区转入关闭门禁",
            "carrier",
            ["缺少统一运营商候选清单"],
            [],
            {},
            "无法审计休息区转入关闭状态",
        )
    manifest = read_json(manifest_path)
    enabled = manifest.get("restZoneEntryEnabled")
    if not isinstance(enabled, bool):
        errors.append("统一候选清单未声明有效休息区转入开关")
    expected_routes = {
        "b1_b11_b12_today_gte50_zero_to_e",
        "chatr_to_e",
    }
    if sorted(manifest.get("restZoneEntryDisabledRoutes") or []) != sorted([] if enabled else expected_routes):
        errors.append("休息区关闭路由与开关状态不一致")
    outputs = manifest.get("outputs") or {}
    inspected: dict[str, int] = {}
    for route in sorted(expected_routes):
        output = outputs.get(route) or {}
        candidate_path = resolver.resolve(output.get("path"), required=False)
        if candidate_path is None:
            errors.append(f"休息区候选文件缺失: {route}")
            continue
        candidate = read_json(candidate_path)
        rows = candidate.get("rows") or candidate.get("candidates") or []
        inspected[route] = len(rows)
        if candidate.get("restZoneEntryEnabled") is not enabled:
            errors.append(f"{candidate_path.name}休息区转入开关不一致")
        if rows and not enabled:
            errors.append(f"{candidate_path.name}关闭状态仍存在休息区转入候选")
        if candidate.get("route") != route or output.get("count") != len(rows):
            errors.append(f"{candidate_path.name}路由或候选计数不一致")
        expected_target = "E-C休息区" if route == "chatr_to_e" else "E-T休息区"
        if candidate.get("targetCarrierName") != expected_target or output.get("target") != expected_target:
            errors.append(f"{candidate_path.name}休息区目标错误")
        pattern = r"^(A1|A11|A12|A3)(?:-|\s|$)" if route == "chatr_to_e" else r"^(B1|B11|B12)(?:-|\s|$)"
        for row in rows:
            source = str(row.get("原运营商") or row.get("运营商") or row.get("sourceCarrierName") or "")
            try:
                send = float(row.get("发送窗口数"))
                reply = float(row.get("回复窗口数"))
            except (TypeError, ValueError):
                send, reply = -1, -1
            if (not re.match(pattern, source, re.IGNORECASE)
                    or not (send >= REPLY_RATE_SEND_THRESHOLD and reply == 0)
                    or (route == "chatr_to_e" and "chatr" not in str(row.get("opr") or "").lower())):
                errors.append(f"{candidate_path.name}存在不满足来源、发送>=50且回复=0的候选")
                break
    held_b = int(manifest.get("restZoneEntryBQualifiedHeldCount") or 0)
    held_chatr = int(manifest.get("restZoneEntryChatrQualifiedHeldCount") or 0)
    if held_b < 0 or held_chatr < 0:
        errors.append("休息区关闭保持号码数不可为负数")
    if not enabled and manifest.get("restZoneEntryCandidateCount") != 0:
        errors.append("关闭状态休息区转入候选总数必须为0")
    if manifest.get("restZoneEntryCandidateCount") != sum(inspected.values()):
        errors.append("休息区转入总候选数不一致")
    details = {
        "rest_zone_entry_enabled": manifest.get("restZoneEntryEnabled"),
        "disabled_routes": manifest.get("restZoneEntryDisabledRoutes"),
        "candidate_counts": inspected,
        "b_qualified_held": held_b,
        "chatr_qualified_held": held_chatr,
        "e_reply_recovery_unchanged": True,
    }
    return status_result(
        "休息区转入开关门禁",
        "carrier",
        errors,
        [],
        details,
        "E-T/E-C转入开关、候选条件及V组保护已核验；E区有回复恢复流程保留",
    )


# Backward-compatible import names for older local tests/tools.
check_rest_zone_entry_pause = check_rest_zone_entry_enabled
check_rest_zone_entry_disabled = check_rest_zone_entry_enabled


def check_v_zone_priority_rule(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    invalid_rows: list[dict[str, Any]] = []

    def metric(value: Any) -> float:
        try:
            return float(value)
        except (TypeError, ValueError):
            return float("-inf")

    manifest_path = resolver.resolve(config.get("carrierRoutingManifest"), required=False)
    if manifest_path is None:
        return status_result(
            "B区V组固定优先级",
            "carrier",
            ["缺少统一运营商候选清单"],
            [],
            {},
            "V组规则无法审计",
        )
    manifest = read_json(manifest_path)
    expected_priority = [
        "b1_b11_b12_today_gte50_zero_to_e",
        "b1v_b11v_use2_lt240_to_b11v",
        "b1v_b11v_gte240_to_b12v",
    ]
    if manifest.get("vZoneRoutePriority") != expected_priority:
        errors.append("V组候选路由优先级不是0回复保持、koo是否2天=2天转B11V、10日>=240转B12V")
    if manifest.get("candidateWorkbookPasses") != 1:
        errors.append("V组候选未证明与主路由共用一次工作簿扫描")
    if manifest.get("vZoneCandidateGenerationMergedWithMainRouting") is not True:
        errors.append("V组候选未证明已合并到主路由生成步骤")
    if manifest.get("vZoneWriteExecutionSerial") is not True:
        errors.append("V组运营商写入未声明严格串行")
    if manifest.get("useDaysFieldEquals") != "2天":
        errors.append("V组koo是否2天字段门禁不是精确等于2天")
    rules = manifest.get("vZonePriorityRules") or []
    if [rule.get("priority") for rule in rules] != [1, 2, 3, "hold"]:
        errors.append("V组结构化优先级规则不完整")

    outputs = manifest.get("outputs") or {}
    inspected_rows = 0
    route_counts: dict[str, int] = {}
    for route in expected_priority:
        output = outputs.get(route) or {}
        candidate_path = resolver.resolve(output.get("path"), required=False)
        if candidate_path is None:
            errors.append(f"V组优先级候选文件缺失: {route}")
            continue
        candidate = read_json(candidate_path)
        rows = candidate.get("rows") or candidate.get("candidates") or []
        route_counts[route] = len(rows)
        if int(candidate.get("vZonePriority") or 0) != expected_priority.index(route) + 1:
            errors.append(f"{candidate_path.name}缺少正确的V组优先级标记")
        if route == expected_priority[1] and candidate.get("useDaysFieldEquals") != "2天":
            errors.append(f"{candidate_path.name}缺少koo是否2天字段精确等于2天门禁")
        for row in rows:
            source = str(row.get("原运营商") or row.get("运营商") or "")
            is_v_source = bool(re.match(r"^(?:B1V|B11V|B12V)(?:-|\s|$)", source, re.I))
            if route == expected_priority[0] and not is_v_source:
                continue
            inspected_rows += 1
            phone = str(row.get("phone") or row.get("电话") or "")
            send = metric(row.get("发送窗口数"))
            reply = metric(row.get("回复窗口数"))
            ten_day = metric(row.get("10日发送窗口数"))
            if route == expected_priority[0]:
                invalid_rows.append({"phone": phone, "route": route, "reason": "v_group_rest_exit_forbidden"})
            elif not re.match(r"^(?:B1V|B11V)(?:-|\s|$)", source, re.I):
                invalid_rows.append({"phone": phone, "route": route, "reason": "invalid_source", "source": source})
            elif send >= 50 and reply == 0:
                invalid_rows.append({"phone": phone, "route": route, "reason": "priority_1_should_have_won"})
            elif route == expected_priority[1] and (
                str(row.get("koo是否2天") or "").strip() != "2天"
                or ten_day == float("-inf")
                or ten_day >= 240
            ):
                invalid_rows.append({"phone": phone, "route": route, "reason": "priority_2_condition"})
            elif route == expected_priority[2] and ten_day < 240:
                invalid_rows.append({"phone": phone, "route": route, "reason": "priority_3_condition"})
    if invalid_rows:
        errors.append(f"发现{len(invalid_rows)}个V组候选违反固定优先级或边界条件")
    details = {
        "route_priority": expected_priority,
        "candidate_workbook_passes": manifest.get("candidateWorkbookPasses"),
        "merged_with_main_routing": manifest.get("vZoneCandidateGenerationMergedWithMainRouting"),
        "write_execution_serial": manifest.get("vZoneWriteExecutionSerial"),
        "use_days_field_equals": manifest.get("useDaysFieldEquals"),
        "b12v_held_count": int(manifest.get("b12vHeldWhenPriority1NotMatched") or 0),
        "inspected_v_rows": inspected_rows,
        "route_counts": route_counts,
        "invalid_count": len(invalid_rows),
        "invalid_sample": invalid_rows[:20],
    }
    return status_result(
        "B区V组固定优先级",
        "carrier",
        errors,
        [],
        details,
        "V组固定优先级、一次扫描合并候选和串行写入门禁通过",
    )


def check_task_g(config: dict[str, Any], resolver: PathResolver, result_path: Path) -> CheckResult:
    from opensms_task_g import audit
    task = config.get("taskG") or {}
    if not task or str(task.get("status", "")).startswith("skipped"):
        return status_result("任务G周末休息与次日恢复", "carrier", [], [], {"status": task.get("status", "legacy_not_configured")}, "任务G本轮未执行")
    errors, details = [], {}
    try:
        if task.get("status") != "passed" or task.get("independentReadbackPassed") is not True:
            raise ValueError("任务G未完整通过独立复核")
        proof = audit(resolver.resolve(task["input"]), resolver.resolve(task["receipt"]))
        recorded = read_json(resolver.resolve(task["independentAudit"]))
        if proof != recorded:
            raise ValueError("任务G独立审计回执与重建不一致")
        receipt = read_json(resolver.resolve(task["receipt"]))
        workbook = xlsx_main_phone_carriers(result_path)
        expected = {}
        for refresh in receipt.get("refreshGroups") or []:
            for phone in refresh["phones"]:
                expected[phone] = refresh["target"]
        for phone, target in expected.items():
            if workbook.get(phone) != [target]:
                raise ValueError("任务G已确认变更未写入最终合并表")
        details = {"entered": proof["enteredCount"], "restored": proof["restoredCount"]}
    except Exception as exc:
        errors.append(str(exc))
    return status_result("任务G周末休息与次日恢复", "carrier", errors, [], details,
                         "昨日阈值、六个来源、次日台账恢复与最终表独立一致")


def check_v_group_boundary(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    from opensms_v_swap_daily import v_group
    allowed = {
        "task_b_expired_to_offline_c": ("B", "过期下架卡C"),
        "task_b_expired_to_offline_t": ("B", "过期下架卡T"),
        "task_b_tomorrow_fast_to_isolation_c": ("B", "快过期隔离区C"),
        "task_b_tomorrow_fast_to_isolation_t": ("B", "快过期隔离区T"),
        "task_c_kooda_4day_reply_rate_lt5_to_swap": ("C", "换号区T"),
        "task_c_chatr_unchanged_4day_reply_rate_lt4_to_swap": ("C", "换号区C"),
    }
    errors, inspected = [], 0
    for group in config.get("groups") or []:
        candidate = read_json(resolver.resolve(group["candidate"]))
        source_names = [row.get(key) for row in candidate.get("rows") or candidate.get("candidates") or []
                        for key in ("原运营商", "运营商", "精确原运营商", "sourceCarrierName")]
        if group.get("write"):
            written = read_json(resolver.resolve(group["write"]))
            names = {str(row["identif"]): row["name"] for row in written.get("sourceCarriers") or []}
            source_names += [names.get(str(row.get("carrireIdentif")), row.get("carrierName"))
                             for row in (written.get("before") or {}).get("eligibleRows") or []]
        if not any(v_group(name) for name in source_names):
            continue
        inspected += 1
        target = group.get("target") or candidate.get("targetCarrierName")
        exception = allowed.get(candidate.get("route"))
        if not v_group(target) and not (exception and candidate.get("task") == exception[0]
                and candidate.get("forceGate") is True and target == exception[1]
                and config.get(f"task{exception[0]}Executed") is True):
            errors.append(f"V组禁止转出路由：{candidate.get('route')} -> {target}")
    return status_result("V组转出边界", "carrier", errors, [], {"inspectedGroups": inspected},
                         "V组仅内部流转或经授权B/C条件退出；具体阈值由对应任务审计独立核验")


def check_v_group_daily(config: dict[str, Any]) -> CheckResult:
    from opensms_v_swap_daily import verify_summary
    result = config.get("vGroupSwapDaily")
    errors = []
    if not result:
        errors.append("缺少V组每日转换号统计与独立回读")
    else:
        try:
            verify_summary(result)
        except Exception as exc:
            errors.append(str(exc))
    return status_result("V组每日转换号统计", "delivery", errors, [],
                         {"today": (result or {}).get("today")}, "V组每日去重人数与台账、Excel独立回读一致")


def check_c2_routing_rule(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    configured_path = config.get("carrierRoutingManifest")
    if not configured_path:
        configured_path = next(
            (
                item.get("path")
                for item in config.get("extraReceipts") or []
                if item.get("label") == "运营商候选清单"
            ),
            None,
        )
    if not configured_path:
        return status_result(
            "C2低回复率与发送量规则",
            "carrier",
            [],
            ["历史配置未提供运营商候选清单，无法追溯C2新规则"],
            {"manifest_path": None, "legacy_config": True},
            "历史轮次不含C2新规则回执",
        )

    manifest_path = resolver.resolve(configured_path)
    manifest = read_json(manifest_path)
    outputs = manifest.get("outputs") or {}
    route = outputs.get("c2_previous_day_reply_rate_lt5_to_a11")
    if route is None:
        return status_result(
            "C2低回复率与发送量规则",
            "carrier",
            [],
            ["历史候选清单无C2前一天新路由，本轮按旧规则生成"],
            {"manifest_path": str(manifest_path), "legacy_manifest": True},
            "历史轮次不含C2新路由",
        )

    candidate_path = resolver.resolve(route.get("path"))
    candidate = read_json(candidate_path)
    rows = candidate.get("rows") or candidate.get("candidates") or []
    target = str(candidate.get("targetCarrierName") or route.get("target") or "")
    if not re.match(r"^A11(?:-|$)", target, re.IGNORECASE):
        errors.append("C2候选目标不是A11")
    if float(manifest.get("c2ReplyRateThresholdExclusive") or 0) != 0.05:
        errors.append("C2回复率阈值不是严格低于5%")
    if int(manifest.get("c2LatestDaySendAlertThresholdExclusive") or 0) != 4000:
        errors.append("C2发送量提醒阈值不是严格低于4000")
    daily_execution = manifest.get("c2DailyExecution") or {}
    if not daily_execution:
        warnings.append("C2候选清单缺少北京时间21:00前每日一次门禁回执")
    else:
        if manifest.get("c2DailyTimezone") != "Asia/Shanghai":
            errors.append("C2每日门禁时区不是Asia/Shanghai")
        if int(manifest.get("c2DailyCutoffHourExclusiveBj") or -1) != 21:
            errors.append("C2每日门禁截止时间不是北京时间21:00（排他）")
        if int(manifest.get("c2DailyMaxExecutions") or 0) != 1:
            errors.append("C2每日执行上限不是1次")
        eligible = daily_execution.get("eligible") is True
        status = daily_execution.get("status")
        if status not in {
            "evaluation_completed", "skipped_already_executed", "skipped_after_cutoff",
        }:
            errors.append(f"C2每日门禁状态无效: {status}")
        if eligible and status != "evaluation_completed":
            errors.append("C2每日门禁已授权但状态不是evaluation_completed")
        if not eligible and rows:
            errors.append("C2每日门禁未授权但仍生成了转A11候选")
        if not eligible and manifest.get("c2DailyStatisticsSendAlertTriggered") is True:
            errors.append("C2每日门禁未授权但仍触发了发送量提醒")
        evaluated_at = daily_execution.get("evaluatedAt")
        if eligible and evaluated_at:
            evaluated = datetime.fromisoformat(str(evaluated_at).replace("Z", "+00:00"))
            evaluated_bj = evaluated.astimezone(timezone(timedelta(hours=8)))
            if evaluated_bj.hour >= 21:
                errors.append("C2每日门禁在北京时间21:00（含）以后仍授权执行")
            if evaluated_bj.date().isoformat() != daily_execution.get("beijingDate"):
                errors.append("C2每日门禁北京时间日期与执行时间不一致")
        receipt_path = resolver.resolve(manifest.get("c2DailyExecutionReceipt"), required=False)
        if receipt_path is None:
            errors.append("C2每日执行回执文件缺失")
        else:
            daily_receipt = read_json(receipt_path)
            if daily_receipt.get("beijingDate") != daily_execution.get("beijingDate"):
                errors.append("C2每日执行回执与候选清单的北京时间日期不一致")
            if int(daily_receipt.get("cutoffHourExclusive") or -1) != 21:
                errors.append("C2每日执行回执的截止时间不是北京时间21:00（排他）")
            if int(daily_receipt.get("maxExecutionsPerBeijingDate") or 0) != 1:
                errors.append("C2每日执行回执的单日上限不是1次")
            if daily_execution.get("eligible") is True and daily_receipt.get("status") not in {
                "evaluation_completed", "write_started", "write_completed",
            }:
                errors.append("C2已授权执行，但每日回执未记录有效的生成或写入状态")
    invalid_rows = []
    for row in rows:
        send_count = int(row.get("前一天发送窗口数") or 0)
        reply_count = int(row.get("前一天回复窗口数") or 0)
        rate = reply_count / send_count if send_count > 0 else None
        source = str(row.get("原运营商") or row.get("运营商") or "")
        if rate is None or rate >= 0.05 or not re.match(r"^C2(?:-|$)", source, re.IGNORECASE):
            invalid_rows.append(str(row.get("phone") or row.get("电话") or ""))
    if invalid_rows:
        errors.append(f"C2候选中有{len(invalid_rows)}条不满足来源或严格低于5%条件")

    statistics_receipt_path = resolver.resolve(manifest.get("c2DailyStatisticsReceipt"), required=False)
    send_total = manifest.get("c2DailyStatisticsSendTotal")
    statistics_errors = []
    eligible = daily_execution.get("eligible") is True
    if not eligible:
        if statistics_receipt_path is not None or send_total is not None:
            statistics_errors.append("C2每日门禁未授权但仍保留本轮按日统计结果")
    elif statistics_receipt_path is None:
        statistics_errors.append("C2按日统计实时查询回执缺失")
    else:
        statistics_receipt = read_json(statistics_receipt_path)
        c2_rows = [item for item in statistics_receipt.get("replyRates") or [] if item.get("group") == "C2"]
        if len(c2_rows) != 1:
            statistics_errors.append(f"C2按日统计结果数量不是1: {len(c2_rows)}")
        else:
            c2_stat = c2_rows[0]
            daily_row = c2_stat.get("dailyRow") or {}
            if c2_stat.get("endpoint") != "/api/admin/statisticsList" or c2_stat.get("status") != "ok":
                statistics_errors.append("C2按日统计接口来源或状态无效")
            if int(daily_row.get("sendSize") or 0) != int(send_total or 0):
                statistics_errors.append("C2发送量与按日统计最新行sendSize不一致")
            if str(daily_row.get("days") or "") != str(manifest.get("c2DailyStatisticsDate") or ""):
                statistics_errors.append("C2发送量统计日期与按日统计最新行不一致")
    errors.extend(statistics_errors)
    expected_alert = (
        eligible
        and not statistics_errors
        and send_total is not None
        and int(send_total) < 4000
    )
    actual_alert = manifest.get("c2DailyStatisticsSendAlertTriggered")
    nested_alert = (manifest.get("c2DailyStatisticsSendAlert") or {}).get("triggered")
    if actual_alert is not expected_alert or nested_alert is not expected_alert:
        errors.append("C2发送量提醒结果与严格低于4000的计算不一致")
    details = {
        "manifest_path": str(manifest_path),
        "candidate_path": str(candidate_path),
        "candidate_count": len(rows),
        "target": target,
        "invalid_candidate_count": len(invalid_rows),
        "previous_day": manifest.get("c2PreviousDay"),
        "daily_statistics_date": manifest.get("c2DailyStatisticsDate"),
        "daily_statistics_send_total": send_total,
        "daily_statistics_receipt": str(statistics_receipt_path) if statistics_receipt_path else None,
        "send_alert_threshold_exclusive": 4000,
        "send_alert_triggered": actual_alert,
        "daily_execution": daily_execution,
    }
    return status_result(
        "C2低回复率与发送量规则",
        "carrier",
        errors,
        warnings,
        details,
        "C2前一天候选仅转A11，按日统计发送量、5%与4000严格边界及北京时间每日一次门禁通过",
    )


def check_forbidden_b13_targets(groups: list[dict[str, Any]], resolver: PathResolver) -> CheckResult:
    target_keys = {"target", "targetcarriername", "目标运营商"}
    hits: list[dict[str, str]] = []
    inspected_documents = 0

    def inspect(node: Any, document: str, location: str = "$") -> None:
        if isinstance(node, dict):
            for key, value in node.items():
                next_location = f"{location}.{key}"
                if str(key).replace("_", "").lower() in target_keys:
                    target = value.get("name") if isinstance(value, dict) else value
                    if isinstance(target, str) and re.match(r"^B13(?:-|$)", target.strip(), re.IGNORECASE):
                        hits.append({"document": document, "location": next_location, "target": target})
                inspect(value, document, next_location)
        elif isinstance(node, list):
            for index, value in enumerate(node):
                inspect(value, document, f"{location}[{index}]")

    for group in groups:
        for field in ("candidate", "validated", "write", "verify"):
            configured_path = group.get(field)
            if not configured_path:
                continue
            path = resolver.resolve(configured_path)
            if not path.is_file():
                continue
            inspected_documents += 1
            inspect(read_json(path), str(path))

    errors = [f"发现{len(hits)}处B13目标候选或写入"] if hits else []
    details = {
        "forbidden_group": "B13",
        "inspected_documents": inspected_documents,
        "target_hit_count": len(hits),
        "hits": hits,
    }
    return status_result(
        "B13目标写入禁令",
        "carrier",
        errors,
        [],
        details,
        "B13目标候选0、实际写入0",
    )


def check_protected_carrier_sources(groups: list[dict[str, Any]], resolver: PathResolver) -> CheckResult:
    hits: list[dict[str, str]] = []
    never_process_hits: list[dict[str, str]] = []
    task_b_exception_hits: list[dict[str, str]] = []
    task_d_exception_hits: list[dict[str, str]] = []
    task_b_retained_hits: list[dict[str, str]] = []
    inspected_documents = 0
    protected_groups = ["A21", "A2"]

    def exception_kind(document: dict[str, Any]) -> str:
        task = str(document.get("task") or "").strip().upper()
        if task == "B" and document.get("forceGate") is True:
            if (document.get("sourceCarrierScope") == "all_carriers_v1"
                    and document.get("sourceCarrierExclusions") == []
                    and re.fullmatch(r"task_b_(?:expired_to_offline|tomorrow_fast_to_isolation)_[ct]", str(document.get("route") or ""))):
                return "B_ALL"
            return "B"
        if (
            task == "D"
            and document.get("automationGate") is True
            and document.get("sourceProvenance") == "phone_card_retain_all_pages"
        ):
            return "D"
        candidate_value = document.get("candidatePath")
        if not candidate_value:
            return ""
        candidate_path = resolver.resolve(candidate_value, required=False)
        if not candidate_path or not candidate_path.is_file():
            return ""
        candidate = read_json(candidate_path)
        return exception_kind(candidate)

    def task_b_document(document: dict[str, Any], seen: set[Path] | None = None) -> bool:
        if str(document.get("task") or "").strip().upper() == "B" or re.match(r"^task_b(?:_|$)", str(document.get("route") or ""), re.IGNORECASE):
            return True
        path = resolver.resolve(document.get("candidatePath"), required=False)
        visited = set() if seen is None else seen
        if not path or path in visited:
            return False
        visited.add(path)
        return task_b_document(read_json(path), visited)

    def record(document: str, location: str, value: Any, kind: str, is_task_b: bool) -> None:
        if kind != "B_ALL" and isolation_carrier_held(value):
            hits.append({"document": document, "location": location, "source": value})
        if kind != "B_ALL" and isinstance(value, str) and value.strip() in {"快过期续费C", "快过期续费T"}:
            never_process_hits.append({"document": document, "location": location, "source": value})
        if kind != "B_ALL" and is_task_b and isinstance(value, str) and value.strip() in TASK_B_RETAINED_SOURCE_CARRIERS:
            task_b_retained_hits.append({"document": document, "location": location, "source": value})
        if isinstance(value, str) and re.match(r"^(?:A21|A2)(?:-|\s|$)", value.strip(), re.IGNORECASE):
            target = task_b_exception_hits if kind in {"B", "B_ALL"} else task_d_exception_hits if kind == "D" else hits
            target.append({"document": document, "location": location, "source": value})

    for group in groups:
        for field in ("candidate", "validated", "write"):
            configured_path = group.get(field)
            if not configured_path:
                continue
            path = resolver.resolve(configured_path)
            if not path.is_file():
                continue
            inspected_documents += 1
            document = read_json(path)
            kind = exception_kind(document)
            is_task_b = str(group.get("title") or "").startswith("任务B") or task_b_document(document)
            for index, value in enumerate(document.get("sourceCarrierNames") or []):
                record(str(path), f"$.sourceCarrierNames[{index}]", value, kind, is_task_b)
            for index, carrier in enumerate(document.get("sourceCarriers") or []):
                if isinstance(carrier, dict):
                    record(str(path), f"$.sourceCarriers[{index}].name", carrier.get("name"), kind, is_task_b)
            for collection in ("rows", "candidates"):
                for index, row in enumerate(document.get(collection) or []):
                    if not isinstance(row, dict):
                        continue
                    for key in ("运营商", "原运营商", "精确原运营商", "A区精确来源", "sourceCarrierName", "carrierName"):
                        record(str(path), f"$.{collection}[{index}].{key}", row.get(key), kind, is_task_b)
            for index, row in enumerate(((document.get("before") or {}).get("eligibleRows") or [])):
                if isinstance(row, dict):
                    record(str(path), f"$.before.eligibleRows[{index}].carrierName", row.get("carrierName"), kind, is_task_b)

    errors = [f"发现{len(hits)}处A21/A2来源候选或写入"] if hits else []
    if never_process_hits:
        errors.append(f"发现{len(never_process_hits)}处快过期续费C/T来源候选或写入")
    if task_b_retained_hits:
        errors.append(f"发现{len(task_b_retained_hits)}处任务B保留卡B14-C/T来源候选或写入")
    details = {
        "protected_groups": protected_groups,
        "inspected_documents": inspected_documents,
        "source_hit_count": len(hits),
        "hits": hits,
        "task_b_exception_source_hit_count": len(task_b_exception_hits),
        "task_b_exception_hits": task_b_exception_hits,
        "task_d_exception_source_hit_count": len(task_d_exception_hits),
        "task_d_exception_hits": task_d_exception_hits,
        "never_process_source_carriers": ["快过期续费C", "快过期续费T"],
        "never_process_source_hit_count": len(never_process_hits),
        "never_process_hits": never_process_hits,
        "task_b_retained_source_carriers": list(TASK_B_RETAINED_SOURCE_CARRIERS),
        "task_b_retained_source_hit_count": len(task_b_retained_hits),
        "task_b_retained_hits": task_b_retained_hits,
    }
    return status_result(
        "A21/A2来源运营商保护",
        "carrier",
        errors,
        [],
        details,
        "普通及其它任务保护保持；任务B全运营商到期分流单独核验，历史任务B按原保护规则审计",
    )


def check_protected_a21_sources(groups: list[dict[str, Any]], resolver: PathResolver) -> CheckResult:
    return check_protected_carrier_sources(groups, resolver)


def check_offline(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    probe = config.get('expiredCardProbe')
    if probe and probe.get('status') != 'skipped_manual_restore_scope':
        if probe.get('background') is True:
            from opensms_expired_card_probe_background import verify_dispatch
            errors, warnings = [], []
            try:
                if probe.get('status') == 'failed_dispatch_no_business':
                    path = Path(probe['dispatchFailureReceipt'])
                    if hashlib.sha256(path.read_bytes()).hexdigest() != probe['dispatchFailureSha256']:
                        raise ValueError('background dispatch failure hash changed')
                    warnings.append('后台存活测试未派发；主报表不受影响，未执行测试')
                else:
                    verify_dispatch(probe)
            except Exception as exc:
                errors.append('后台存活测试派发核验失败：' + str(exc))
            return status_result('下架恢复', 'offline', errors, warnings, probe,
                                 '仅核验后台派发；测试与主文档独立，不将派发当作恢复成功')
        from opensms_expired_card_probe_audit import verify_receipt
        errors, warnings = [], []
        try:
            receipt = verify_receipt(probe)
            if receipt.get('pendingCount') or receipt.get('blockedCount'):
                warnings.append('过期区异步恢复尚未确认，已持久记录且禁止重放')
            if receipt.get('unknownDateCount') or receipt.get('ineligibleStatusCount') or receipt.get('invalidIdentityCount'):
                warnings.append('缺有效截止日或异常/其他状态记录未操作，见独立范围统计')
        except Exception as exc:
            errors.append('过期区存活测试独立核验失败：' + str(exc))
        return status_result('下架恢复', 'offline', errors, warnings, probe,
                             '窄范围存活测试身份、一次性意图、接受与独立回读核验通过')
    offline = config.get("offline") or {}
    if config.get("offlineRecoveryEnabled") is False or offline.get("status") == "disabled":
        details = {
            "enabled": False,
            "candidate_count": 0,
            "actual_writes": 0,
            "retry_count": 0,
            "readback_count": 0,
        }
        return status_result(
            "下架恢复",
            "offline",
            [],
            [],
            details,
            "下架恢复功能已取消；候选、写入、重试和回读均为0",
        )
    errors: list[str] = []
    candidate_path = resolver.resolve(offline.get("candidate"))
    candidate = read_json(candidate_path)
    count = candidate_count(candidate)
    final_path = resolver.resolve(offline.get("finalVerify"))
    final = read_json(final_path)
    final_count = int(final.get("candidateCount") or 0)
    if final_count not in {0, count} and count != 0:
        errors.append("下架恢复复核数量与候选数量不一致")
    if final.get("status") not in {None, "success", "成功"}:
        errors.append("下架恢复最终复核状态失败")
    if int(final.get("inactiveCount") or 0) != 0:
        errors.append(f"仍有{final.get('inactiveCount')}条未恢复")
    details = {
        "candidate_path": str(candidate_path),
        "final_verify_path": str(final_path),
        "candidate_count": count,
        "final_active": final.get("alreadyActiveCount"),
        "final_inactive": final.get("inactiveCount"),
    }
    return status_result("下架恢复", "offline", errors, [], details, "下架恢复最终状态通过")


def task_b_source_date(value: Any, epoch: datetime) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return None
    text = str(value).strip()
    if isinstance(value, (int, float)) or re.fullmatch(r"\d+(?:\.0+)?", text):
        try:
            converted = from_excel(float(value), epoch)
            return converted.date() if isinstance(converted, datetime) else None
        except (TypeError, ValueError, OverflowError):
            return None
    match = re.fullmatch(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?", text)
    if match:
        try:
            return date(*(int(part) for part in match.groups()))
        except ValueError:
            pass
    return None


def task_b_source_phone(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", "", str(value).strip()).removesuffix(".0")


def audit_source_anomaly(row_number, row, reason, phone):
    return {"sourceRow": row_number, "phone": phone, "slot": str(row.get("卡槽") or "").strip(),
            "sourceCarrierName": str(row.get("运营商") or "").strip(), "reason": reason}


def review_data_anomaly_declaration(manifest, expected_rows, candidate_phones=()):
    policy = manifest.get("dataAnomalyPolicy")
    if policy is None:
        return [] if not expected_rows and not manifest.get("dataAnomalySkippedCount") and not manifest.get("dataAnomalySkippedRows") else ["旧契约不能追认异常数据跳过"]
    counts = {}
    for row in expected_rows:
        counts[row["reason"]] = counts.get(row["reason"], 0) + 1
    def phone_key(value):
        text = str(value or '').strip()
        return text[-10:] if re.fullmatch(r'1?\d{10}', text) else text
    excluded = [phone_key(row['phone']) for row in expected_rows if phone_key(row['phone'])]
    if len(excluded) != len(set(excluded)) or set(excluded).intersection(phone_key(value) for value in candidate_phones):
        return ["异常排除不得掩盖重复或冲突号码身份"]
    return [] if (policy == "skip_row_data_anomalies_v1"
        and type(manifest.get("dataAnomalySkippedCount")) is int
        and manifest.get("dataAnomalySkippedCount") == len(expected_rows)
        and manifest.get("dataAnomalySkippedRows") == expected_rows
        and manifest.get("dataAnomalyCountsByReason") == counts) else ["异常数据排除身份、原因、集合或计数与独立源表不一致"]


def task_ce_source_date(value, epoch):
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) or re.fullmatch(r"\d+(?:\.0+)?", str(value).strip()):
        try:
            result = from_excel(float(value), epoch)
            return result.date() if isinstance(result, datetime) else result
        except (ValueError, TypeError, OverflowError):
            return None
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", str(value))
    try:
        return date(*(int(part) for part in match.groups())) if match else None
    except ValueError:
        return None


def task_c_source_anomalies(manifest, source_rows, epoch, daily_reason):
    """Independently rebuild just the new row-error partition, not write outcomes."""
    def number(value):
        try:
            return None if isinstance(value, bool) or value in (None, "") else float(str(value).strip().replace(',', ''))
        except (TypeError, ValueError):
            return None
    expected = []
    for index, row in source_rows.items():
        carrier = str(row.get('运营商') or '').strip()
        status = str(row.get('过期状态') or '').strip()
        deadline = task_ce_source_date(row.get('截止日'), epoch)
        if (not carrier or isolation_carrier_held(carrier) or temporary_carrier_held(carrier)
                or carrier in {'快过期续费C', '快过期续费T', '过期下架卡C', '过期下架卡T', '快过期隔离区C', '快过期隔离区T'}
                or status == '已过期' or (status == '快过期' and deadline and deadline.isoformat() == manifest['tomorrowDate'])
                or re.match(r'^(A21|A2)(?:-|\s|$)', carrier, re.I)):
            continue
        opr = str(row.get('opr') or '')
        is_t = bool(re.search(r'kooda|koodo|telus1|kini', opr, re.I))
        is_c = bool(re.search(r'chatr', opr, re.I))
        send = number(row.get('4日发送窗口数'))
        raw_rate = str(row.get('4日回复率')).strip() if row.get('4日回复率') is not None else ''
        rate = number(raw_rate.removesuffix('%'))
        if rate is not None and (raw_rate.endswith('%') or rate > 1):
            rate /= 100
        ten_send = number(row.get('10日发送窗口数'))
        if (not (is_t or is_c) or send is None or send <= 0 or rate is None or rate < 0
                or (re.match(r'^(A1|A11|A12|B1|B11|B12|B1V|B11V|B12V)(?:-|\s|$)', carrier, re.I) and send < 50)
                or ten_send is None or not math.isfinite(ten_send) or ten_send <= 360):
            continue
        is_t_match, is_c_match = is_t and rate < .05, is_c and rate < .04
        if not (is_t_match or is_c_match):
            continue
        change = str(row.get('换号状态') or '').strip()
        phone = task_b_source_phone(row.get('电话'))
        if is_c_match and change not in ('', '未换号'):
            if change != '已换号':
                expected.append(audit_source_anomaly(index, row, 'change_status_unknown', phone))
            continue
        if daily_reason(row, .05 if is_t_match else .04):
            continue
        if not phone:
            expected.append(audit_source_anomaly(index, row, 'phone_missing', phone))
    return expected


def review_task_b_source_coverage(
    config: dict[str, Any], manifest: dict[str, Any], task_groups: list[dict[str, Any]], resolver: PathResolver,
) -> tuple[list[str], dict[str, Any]]:
    """Rebuild B's partitions from its hash-bound input, independently of its generator."""
    errors: list[str] = []
    details: dict[str, Any] = {"source_coverage_independently_verified": False}

    def require(condition: bool, message: str) -> None:
        if not condition:
            errors.append(message)

    source = resolver.resolve(manifest.get("sourcePath"))
    source_sha = sha256_file(source)
    all_sources = manifest.get("sourceCarrierScope") == "all_carriers_v1"
    require(manifest.get("sourceCarrierScope") in (None, "legacy_protected_sources_v1", "all_carriers_v1"), "任务B来源范围版本无效")
    if all_sources:
        require(manifest.get("sourceCarrierExclusions") == [], "任务B全运营商范围不得排除来源")
    source_coverage = {}
    skip_unconfigured = manifest.get("unconfiguredCarrierPolicy") == "skip_no_business_write_v1"
    unconfigured_rows = []
    skip_anomalies = manifest.get("dataAnomalyPolicy") == "skip_row_data_anomalies_v1"
    expected_anomalies = []
    due_anomaly_counts = {"expired": 0, "tomorrow": 0}
    normalized_due_phones: set[str] = set()
    require(manifest.get("sourceSha256") == source_sha, "任务B输入源表SHA-256不一致")
    require(manifest.get("retainedSourceProtectionEnforced") is (not all_sources), "任务B清单来源保护声明不符")
    if config.get("taskBInputWorkbook"):
        configured_source = resolver.resolve(config["taskBInputWorkbook"])
        require(sha256_file(configured_source) == source_sha, "任务B清单源表与本轮任务B输入不一致")
    server_date = date.fromisoformat(manifest["serverDate"])
    tomorrow = server_date + timedelta(days=1)
    require(manifest.get("tomorrowDate") == tomorrow.isoformat(), "任务B次日日期不一致")
    if config.get("serverDate"):
        require(config["serverDate"] == server_date.isoformat(), "任务B服务器日期与本轮配置不一致")
    retained = {name: [] for name in TASK_B_RETAINED_SOURCE_CARRIERS}
    renewal = {name: [] for name in TASK_B_RENEWAL_SOURCE_CARRIERS}
    retained_branch = {"expired": 0, "tomorrow": 0}
    non_protected = {"expired": [], "tomorrow": []}
    temporary_held = []
    isolation_held = []
    hold_time = datetime.fromisoformat(manifest["temporaryCarrierHoldEvaluatedAt"]) if manifest.get("temporaryCarrierHoldEvaluatedAt") else None
    expected = {branch: {family: {} for family in ("C", "T")} for branch in non_protected}
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook["sheet1"]
        headers = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(values_only=True))]
        if not {"电话", "运营商", "卡槽", "过期状态", "截止日"}.issubset(headers):
            raise ValueError("任务B输入主表缺少覆盖审计必需字段")
        epoch = workbook.epoch
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            row = dict(zip(headers, values))
            carrier = str(row.get("运营商") or "").strip()
            coverage = source_coverage.setdefault(carrier, {
                "sourceCarrierName": carrier, "sourceRowCount": 0,
                "expiredSourceCount": 0, "expiredCandidateCount": 0,
                "tomorrowSourceCount": 0, "tomorrowCandidateCount": 0,
                **({"expiredDataAnomalyCount": 0, "tomorrowDataAnomalyCount": 0} if skip_anomalies else {}),
            })
            coverage["sourceRowCount"] += 1
            if skip_unconfigured and not carrier:
                unconfigured_rows.append(row_number)
                continue
            status = str(row.get("过期状态") or "").strip()
            deadline = task_b_source_date(row.get("截止日"), epoch)
            raw_deadline = row.get("截止日")
            blank_deadline = raw_deadline is None or (isinstance(raw_deadline, str) and not raw_deadline.strip())
            if deadline is None and (not blank_deadline or status):
                if skip_anomalies:
                    expected_anomalies.append(audit_source_anomaly(row_number, row, "deadline_missing_or_invalid", task_b_source_phone(row.get("电话"))))
                    continue
                errors.append(f"任务B源表第{row_number}行截止日缺失或无效，禁止凭状态改区")
                continue
            remaining = (deadline - server_date).days if deadline else None
            status = ("已过期" if remaining <= 0 else "快过期" if remaining <= 2 else "6天" if remaining <= 6 else "") if remaining is not None else ""
            branch = "expired" if deadline and deadline <= server_date else "tomorrow" if deadline == tomorrow else None
            if not branch:
                continue
            if not all_sources and isolation_carrier_held(carrier):
                isolation_held.append(row_number)
                continue
            if not all_sources and hold_time is not None and temporary_carrier_held(carrier, hold_time):
                temporary_held.append(row_number)
                continue
            if not all_sources and carrier in renewal:
                renewal[carrier].append(row_number)
                continue
            if not all_sources and carrier in retained:
                retained[carrier].append(row_number)
                retained_branch[branch] += 1
                continue
            non_protected[branch].append(row_number)
            coverage[f"{branch}SourceCount"] += 1
            phone = task_b_source_phone(row.get("电话"))
            opr = str(row.get("opr") or "").strip()
            is_c = bool(re.search("chatr", opr, re.IGNORECASE))
            # Task B only: user-confirmed Lucky expiry routing, not an F alias.
            is_t = bool(re.search(r"kooda|koodo|kini|telus1|\blucky\b", opr, re.IGNORECASE))
            if not phone or is_c == is_t or (all_sources and not carrier):
                if skip_anomalies and not (is_c and is_t):
                    expected_anomalies.append(audit_source_anomaly(row_number, row,
                        "phone_missing" if not phone else "opr_unclassified", phone))
                    coverage[f"{branch}DataAnomalyCount"] += 1
                    due_anomaly_counts[branch] += 1
                    continue
                errors.append(f"任务B源表第{row_number}行非保护到期号码缺号或OPR无法唯一分区")
                continue
            family = "C" if is_c else "T"
            if all_sources:
                normalized = re.sub(r"\D", "", phone)[-10:]
                require(bool(normalized) and normalized not in normalized_due_phones,
                        "任务B全运营商源表到期号码标准化后无效或跨分支重复")
                normalized_due_phones.add(normalized)
            expected[branch][family][row_number] = {
                "电话": phone, "运营商": carrier, "opr": opr, "过期状态": status,
                "截止日": deadline.isoformat() if deadline else "", "卡槽": str(row.get("卡槽") or "").strip(),
            }
            coverage[f"{branch}CandidateCount"] += 1
    finally:
        workbook.close()

    errors.extend(review_data_anomaly_declaration(manifest, expected_anomalies,
        [row['电话'] for branch in expected.values() for family in branch.values() for row in family.values()]))
    if skip_anomalies:
        for branch in due_anomaly_counts:
            require(manifest.get(f"{branch}DataAnomalyCount") == due_anomaly_counts[branch], "任务B分支异常排除数与源表不一致")
    for prefix, names, excluded in (
        ("retained", TASK_B_RETAINED_SOURCE_CARRIERS, retained),
        ("neverProcess", TASK_B_RENEWAL_SOURCE_CARRIERS, renewal),
    ):
        require(manifest.get(f"{prefix}SourceCarrierNames") == ([] if all_sources else list(names)), f"任务B {prefix} 精确保留来源名单不一致")
        require(manifest.get(f"{prefix}SourceExcludedRows") == excluded, f"任务B {prefix} 保护排除源行与源表不一致")
        require(manifest.get(f"{prefix}SourceExcludedCounts") == {name: len(rows) for name, rows in excluded.items()},
                f"任务B {prefix} 分组保护排除数与源表不一致")
        require(manifest.get(f"{prefix}SourceExcludedCount") == sum(map(len, excluded.values())),
                f"任务B {prefix} 保护排除总数与源表不一致")

    candidate_paths: list[Path] = []
    actual_branch_counts = {"expired": 0, "tomorrow": 0}
    require(manifest.get("temporaryB11HeldRows", []) == temporary_held, "任务B临时B11保护行不符")
    require(manifest.get("isolationCarrierHeldRows", []) == isolation_held, "任务B三方隔离保护行不符")
    require(manifest.get("isolationCarrierHeldCount", 0) == len(isolation_held), "任务B三方隔离保护计数不符")
    require(manifest.get("temporaryB11HeldCount", 0) == len(temporary_held), "任务B临时B11保护计数不符")
    for branch in non_protected:
        require(manifest.get(f"{branch}RetainedSourceExcludedCount") == retained_branch[branch],
                f"任务B {branch} 保留卡排除数与源表不一致")
        require(manifest.get(f"{branch}NonProtectedSourceCount") == len(non_protected[branch]),
                f"任务B {branch} 非保护来源总数与源表不一致")
        seen_phones: set[str] = set()
        for family in ("C", "T"):
            candidate_path = resolver.resolve((manifest.get(f"{branch}CandidatePaths") or {}).get(family))
            candidate_paths.append(candidate_path)
            candidate = read_json(candidate_path)
            label = f"任务B {branch}/{family}"
            require(candidate.get("task") == "B" and candidate.get("forceGate") is True, f"{label}候选缺少强制门禁")
            require(candidate.get("sourceSha256") == source_sha and resolver.resolve(candidate.get("sourcePath")) == source,
                    f"{label}候选未绑定同一源表")
            require(candidate.get("sourceSheet") == "sheet1" and candidate.get("serverDate") == server_date.isoformat()
                    and candidate.get("tomorrowDate") == tomorrow.isoformat(), f"{label}候选来源页或日期不一致")
            require(candidate.get("retainedSourceCarrierNames") == ([] if all_sources else list(TASK_B_RETAINED_SOURCE_CARRIERS))
                    and candidate.get("retainedSourceProtectionEnforced") is (not all_sources), f"{label}候选来源保护声明不符")
            if all_sources:
                require(candidate.get("sourceCarrierScope") == "all_carriers_v1" and candidate.get("sourceCarrierExclusions") == [],
                        f"{label}候选未使用全运营商范围")
            route = f"task_b_expired_to_offline_{family.lower()}" if branch == "expired" else f"task_b_tomorrow_fast_to_isolation_{family.lower()}"
            target = f"过期下架卡{family}" if branch == "expired" else f"快过期隔离区{family}"
            require(candidate.get("route") == route and candidate.get("targetCarrierName") == target, f"{label}候选路由或目标不一致")
            rows = candidate.get("rows")
            if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
                errors.append(f"{label}候选行无效")
                continue
            require(candidate.get("candidates") == rows, f"{label}两份候选集合不一致")
            require(candidate.get("candidateCount") == len(rows) and candidate.get("count") == len(rows), f"{label}候选计数与明细不一致")
            require((manifest.get(f"{branch}CandidateCounts") or {}).get(family) == len(rows), f"{label}清单候选数与文件不一致")
            actual_branch_counts[branch] += len(rows)
            actual: dict[int, dict[str, str]] = {}
            for row in rows:
                row_number = row.get("源行号")
                if type(row_number) is not int or row_number in actual:
                    errors.append(f"{label}候选源行缺失、无效或重复")
                    continue
                phone = task_b_source_phone(row.get("电话"))
                require(bool(phone) and phone not in seen_phones, f"{label}候选号码为空或重复")
                seen_phones.add(phone)
                carrier = str(row.get("运营商") or "").strip()
                if not all_sources:
                    require(carrier not in retained and carrier not in renewal, f"{label}候选含保护来源")
                require(row.get("targetCarrierName") == target, f"{label}候选行目标不一致")
                actual[row_number] = {
                    "电话": phone, "运营商": carrier, "opr": str(row.get("opr") or "").strip(),
                    "过期状态": str(row.get("过期状态") or "").strip(),
                    "截止日": str(row.get("截止日") or ""), "卡槽": str(row.get("卡槽") or "").strip(),
                }
            require(actual == expected[branch][family], f"{label}候选未精确覆盖源表非保护号码")
            require(candidate.get("sourceCarrierNames") == sorted({row["运营商"] for row in actual.values() if row["运营商"]}),
                    f"{label}候选来源名单与明细不一致")
        require(manifest.get(f"{branch}ClassifiedCandidateCount") == actual_branch_counts[branch]
                and manifest.get(f"{branch}CandidateCount") == actual_branch_counts[branch], f"任务B {branch} 候选合计不一致")
        require(actual_branch_counts[branch] + due_anomaly_counts[branch] == len(non_protected[branch]), f"任务B {branch} 非保护来源未完整覆盖")

    manifest_paths = [resolver.resolve(value) for value in manifest.get("candidatePaths") or []]
    group_paths = [resolver.resolve(group.get("candidate")) for group in task_groups]
    require(len(set(candidate_paths)) == 4 and sorted(manifest_paths) == sorted(candidate_paths), "任务B候选总清单与四分支不一致")
    require(sorted(group_paths) == sorted(candidate_paths), "任务B实际审计组未精确覆盖清单四组候选")
    if all_sources:
        if skip_unconfigured:
            require(manifest.get("unconfiguredCarrierSkippedRows") == unconfigured_rows
                    and manifest.get("unconfiguredCarrierSkippedCount") == len(unconfigured_rows),
                    "任务B未配置运营商跳过集合/计数与独立源表不一致")
            for candidate_path in candidate_paths:
                require(read_json(candidate_path).get("unconfiguredCarrierPolicy") == "skip_no_business_write_v1",
                        "任务B候选未声明未配置运营商不写入规则")
        require(manifest.get("sourceCarrierCoverage") == [source_coverage[name] for name in sorted(source_coverage)],
                "任务B逐运营商覆盖清单与独立源表重算不一致")
        require(manifest.get("sourceCarrierCoveragePassed") is True
                and all(row["expiredSourceCount"] == row["expiredCandidateCount"] + row.get("expiredDataAnomalyCount", 0)
                        and row["tomorrowSourceCount"] == row["tomorrowCandidateCount"] + row.get("tomorrowDataAnomalyCount", 0) for row in source_coverage.values()),
                "任务B逐运营商到期号码未完整覆盖")
    details.update({
        "source_carrier_scope": manifest.get("sourceCarrierScope"),
        "unconfigured_carrier_skipped_count": len(unconfigured_rows),
        "data_anomaly_skipped_count": len(expected_anomalies),
        "source_carrier_coverage": [source_coverage[name] for name in sorted(source_coverage)] if all_sources else None,
        "source_path": str(source), "source_sha256": source_sha,
        "source_coverage_independently_verified": not errors,
        "source_retained_excluded_counts": {name: len(rows) for name, rows in retained.items()},
        "source_retained_excluded_count": sum(map(len, retained.values())),
        "source_renewal_excluded_count": sum(map(len, renewal.values())),
        "source_non_protected_counts": {branch: len(rows) for branch, rows in non_protected.items()},
        "source_retained_excluded_branch_counts": retained_branch,
        "actual_classified_candidate_counts": actual_branch_counts,
    })
    return errors, details


def check_task_b(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    executed = config.get("taskBExecuted") is True
    task_groups = [group for group in config.get("groups") or [] if str(group.get("title") or "").startswith("任务B")]
    errors: list[str] = []
    if executed and len(task_groups) != 4:
        errors.append("任务B已授权执行，但没有四组完整结果")
    if not executed and task_groups:
        errors.append("任务B未授权，但配置中存在任务B执行组")
    details = {
        "executed": executed,
        "group_count": len(task_groups),
        "status": config.get("taskBStatus"),
        "expired_non_protected_source_count": 0,
        "expired_classified_candidate_count": 0,
        "expired_coverage_passed": None,
        "tomorrow_non_protected_source_count": 0,
        "tomorrow_classified_candidate_count": 0,
        "tomorrow_coverage_passed": None,
    }
    manifest_value = config.get("taskBManifest")
    if executed and not manifest_value:
        errors.append("任务B已授权执行但缺少候选清单")
    elif executed:
        manifest_path = resolver.resolve(manifest_value)
        manifest = read_json(manifest_path)
        source_count = int(manifest.get("expiredNonProtectedSourceCount") or 0)
        candidate_count_value = int(manifest.get("expiredClassifiedCandidateCount") or 0)
        coverage_passed = manifest.get("expiredCoveragePassed") is True
        tomorrow_source_count = int(manifest.get("tomorrowNonProtectedSourceCount") or 0)
        tomorrow_candidate_count = int(manifest.get("tomorrowClassifiedCandidateCount") or 0)
        tomorrow_coverage_passed = manifest.get("tomorrowCoveragePassed") is True
        details.update({
            "manifest_path": str(manifest_path),
            "expired_non_protected_source_count": source_count,
            "expired_classified_candidate_count": candidate_count_value,
            "expired_coverage_passed": coverage_passed,
            "tomorrow_non_protected_source_count": tomorrow_source_count,
            "tomorrow_classified_candidate_count": tomorrow_candidate_count,
            "tomorrow_coverage_passed": tomorrow_coverage_passed,
            "never_process_source_carriers": manifest.get("neverProcessSourceCarrierNames") or [],
            "never_process_source_excluded_count": int(manifest.get("neverProcessSourceExcludedCount") or 0),
            "retained_source_carriers": manifest.get("retainedSourceCarrierNames") or [],
            "retained_source_excluded_count": int(manifest.get("retainedSourceExcludedCount") or 0),
        })
        if manifest.get("task") != "B" or manifest.get("forceGate") is not True:
            errors.append("任务B候选清单缺少强制授权门禁")
        all_sources = manifest.get("sourceCarrierScope") == "all_carriers_v1"
        if set(manifest.get("neverProcessSourceCarrierNames") or []) != (set() if all_sources else {"快过期续费C", "快过期续费T"}):
            errors.append("任务B没有同时排除快过期续费C/T")
        if not coverage_passed or source_count != candidate_count_value + manifest.get("expiredDataAnomalyCount", 0):
            errors.append("任务B未完整覆盖来源范围内的全部已过期号码")
        if not tomorrow_coverage_passed or tomorrow_source_count != tomorrow_candidate_count + manifest.get("tomorrowDataAnomalyCount", 0):
            errors.append("任务B未完整覆盖来源范围内的全部明日到期号码")
        if manifest.get("allDueCoveragePassed") is not True:
            errors.append("任务B已过期与明日到期综合完整覆盖门禁未通过")
        try:
            source_errors, source_details = review_task_b_source_coverage(config, manifest, task_groups, resolver)
            errors.extend(source_errors)
            details.update(source_details)
        except (OSError, ValueError, TypeError, KeyError, StopIteration, zipfile.BadZipFile) as exc:
            errors.append(f"任务B源表独立覆盖核验失败: {exc}")
            details["source_coverage_independently_verified"] = False
    success = "任务B授权及源表已过期/明日到期完整覆盖通过" if executed else "任务B保持关闭"
    return status_result("任务B授权", "task_b", errors, [], details, success)


def task_c_rate(value: Any) -> float | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    text = str(value).strip()
    percent = text.endswith("%")
    if percent:
        text = text[:-1]
    try:
        number = float(text)
    except (TypeError, ValueError):
        return None
    if number < 0:
        return None
    return number / 100 if percent or number > 1 else number


def check_task_c(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    executed = config.get("taskCExecuted") is True
    task_groups = [
        group for group in config.get("groups") or []
        if str(group.get("title") or "").startswith("任务C")
    ]
    errors: list[str] = []
    warnings: list[str] = []
    manifest_value = config.get("taskCManifest")
    if executed and len(task_groups) != 2:
        errors.append("任务C已授权执行，但没有两组完整结果")
    if not executed and task_groups:
        errors.append("任务C未授权，但配置中存在任务C执行组")
    if not executed and manifest_value:
        errors.append("任务C未授权，但配置中存在任务C候选清单")

    details: dict[str, Any] = {
        "executed": executed,
        "group_count": len(task_groups),
        "status": config.get("taskCStatus"),
        "manifest_path": None,
        "candidate_count": 0,
        "targets": None,
        "invalid_candidate_count": 0,
        "overlap_count": 0,
        "task_b_priority_invalid_count": 0,
        "ten_day_send_invalid_count": 0,
    }
    if not executed:
        success = "任务C保持关闭"
        return status_result("任务C授权与规则", "task_c", errors, warnings, details, success)
    if not manifest_value:
        errors.append("任务C已授权执行，但缺少候选清单")
        return status_result("任务C授权与规则", "task_c", errors, warnings, details, "任务C授权与规则通过")

    manifest_path = resolver.resolve(manifest_value)
    manifest = read_json(manifest_path)
    details["manifest_path"] = str(manifest_path)
    if manifest.get("task") != "C" or manifest.get("forceGate") is not True:
        errors.append("任务C候选清单缺少强制授权门禁")
    if float(manifest.get("koodaReplyRateThresholdExclusive") or 0) != 0.05:
        errors.append("任务C Kooda/Koodo阈值不是严格低于5%")
    if float(manifest.get("chatrReplyRateThresholdExclusive") or 0) != 0.04:
        errors.append("任务C Chatr阈值不是严格低于4%")
    if float(manifest.get("replyRateSendThresholdInclusive") or 0) != 50:
        errors.append("任务C A区、B区及B区V组来源的4日发送门槛不是>=50")
    if manifest.get("tenDaySendThresholdExclusive") != 360 or manifest.get("tenDaySendMetric") != "10日发送窗口数":
        errors.append("任务C候选清单缺少10日发送窗口数严格>360门槛")
    if set(manifest.get("replyRateSendGateGroups") or []) != {
        "A1", "A11", "A12", "B1", "B11", "B12", "B1V", "B11V", "B12V",
    }:
        errors.append("任务C回复率发送门槛来源组不正确")
    if manifest.get("mutuallyExclusive") is not True:
        errors.append("任务C候选清单未声明两组互斥")
    if manifest.get("taskBPriorityLockEnforced") is not True:
        errors.append("任务C候选清单未声明任务B过期优先级锁定")
    server_date = str(manifest.get("serverDate") or "")
    tomorrow_date = str(manifest.get("tomorrowDate") or "")
    try:
        expected_tomorrow = (
            datetime.strptime(server_date, "%Y-%m-%d") + timedelta(days=1)
        ).strftime("%Y-%m-%d")
    except ValueError:
        expected_tomorrow = ""
    if not server_date or tomorrow_date != expected_tomorrow:
        errors.append("任务C候选清单的服务器日期或明日日期无效")

    targets = manifest.get("targetCarrierNames") or {}
    expected_targets = {"C": "换号区C", "T": "换号区T"}
    details["targets"] = targets
    if targets != expected_targets:
        errors.append(f"任务C目标分区不正确: {targets or '空'}")

    expected_source_path = resolver.resolve(
        config.get("taskCInputWorkbook") or config.get("cardResult"),
        required=False,
    )
    manifest_source_path = resolver.resolve(manifest.get("sourcePath"), required=False)
    if manifest_source_path is None:
        errors.append("任务C候选源表不存在")
    if expected_source_path and manifest_source_path and expected_source_path != manifest_source_path:
        errors.append("任务C候选不是从本轮卡槽关联结果表生成")
    if manifest_source_path and normalize_hash(manifest.get("sourceSha256")) != sha256_file(manifest_source_path):
        errors.append("任务C候选源表SHA-256不一致")

    # Old immutable receipts retain their original contract. Current runner
    # always requires v1, so stripping new candidate/manifest flags cannot pass.
    daily_required = config.get("taskCRecentDailyReplyGuardRequired") is True or "recentDailyReplyRateGuard" in manifest
    daily_source = {}
    daily_base = {}
    daily_headers = []
    if daily_required:
        try:
            day = date.fromisoformat(server_date)
            previous = day - timedelta(days=1)
            daily_headers = [day.strftime("%m-%d回复率"), previous.strftime("%m-%d回复率")]
            daily_base = {"version": 1, "serverDate": server_date, "previousDate": previous.isoformat(),
                          "headers": daily_headers, "comparison": "any_strictly_greater_than", "unknownPolicy": "exclude"}
            if manifest.get("recentDailyReplyRateGuard") != {**daily_base, "thresholds": {"T": 0.05, "C": 0.04}}:
                errors.append("任务C近两日回复率保护清单门禁不一致")
            if config.get("serverDate") and config["serverDate"] != server_date:
                errors.append("任务C近两日回复率保护服务器日期与本轮不一致")
            wb = load_workbook(manifest_source_path, read_only=True, data_only=False)
            try:
                ws = wb["sheet1"]
                headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(ws.iter_rows())]
                if any(headers.count(header) != 1 for header in daily_headers):
                    raise ValueError("服务器当天/前一天回复率列缺失或重复")
                daily_source = {index: dict(zip(headers, values)) for index, values in
                                enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)}
                daily_epoch = wb.epoch
            finally:
                wb.close()
        except Exception as exc:
            errors.append(f"任务C近两日回复率保护源表独立回读失败: {exc}")

    def independent_daily_reason(row, threshold):
        rates = []
        for header in daily_headers:
            raw = row.get(header)
            text = str(raw).strip() if isinstance(raw, (int, float, str)) and not isinstance(raw, bool) else ""
            numeric = text[:-1].strip() if text.endswith('%') else text
            valid = re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:e[+-]?\d+)?", numeric, re.IGNORECASE)
            number = float(numeric) if valid else float('nan')
            rate = number / 100 if text.endswith('%') or number > 1 else number
            rates.append(rate if math.isfinite(rate) and 0 <= rate <= 1 else None)
        if any(rate is not None and rate > threshold for rate in rates):
            return 'above_threshold'
        return 'missing_or_invalid' if not rates or any(rate is None for rate in rates) else None

    if manifest.get("unconfiguredCarrierPolicy") == "skip_no_business_write_v1":
        expected_unconfigured = [index for index, row in daily_source.items()
                                 if not str(row.get("运营商") or "").strip()]
        if (not daily_required or manifest.get("unconfiguredCarrierSkippedRows") != expected_unconfigured
                or manifest.get("unconfiguredCarrierSkippedCount") != len(expected_unconfigured)):
            errors.append("任务C未配置运营商跳过集合/计数与独立源表不一致")
        details["unconfigured_carrier_skipped_count"] = len(expected_unconfigured)

    daily_invalid = []
    daily_excluded_indices = set()
    if daily_required:
        excluded = manifest.get('recentDailyReplyRateExcludedRows')
        if not isinstance(excluded, dict) or set(excluded) != {'above_threshold', 'missing_or_invalid'}:
            errors.append('任务C近两日回复率保护缺少排除明细')
            excluded = {}
        counts = {}
        for reason, indices in excluded.items():
            if not isinstance(indices, list):
                errors.append('任务C近两日回复率保护排除明细无效')
                continue
            counts[reason] = len(indices)
            for index in indices:
                row = daily_source.get(index, {}) if type(index) is int else {}
                opr = str(row.get('opr') or '')
                threshold = 0.05 if re.search(r'kooda|koodo|telus1|kini', opr, re.I) else 0.04
                if not row or index in daily_excluded_indices or independent_daily_reason(row, threshold) != reason:
                    errors.append('任务C近两日回复率保护排除明细与源表不一致')
                if type(index) is int:
                    daily_excluded_indices.add(index)
        if manifest.get('recentDailyReplyRateExcludedCounts') != counts or manifest.get('recentDailyReplyRateExcludedCount') != sum(counts.values()):
            errors.append('任务C近两日回复率保护排除计数不一致')

    expected_routes = {
        "task_c_kooda_4day_reply_rate_lt5_to_swap": (0.05, re.compile(r"kooda|koodo|telus1|kini", re.IGNORECASE), "换号区T"),
        "task_c_chatr_unchanged_4day_reply_rate_lt4_to_swap": (0.04, re.compile(r"chatr", re.IGNORECASE), "换号区C"),
    }
    if manifest.get('dataAnomalyPolicy') is not None or manifest.get('dataAnomalySkippedRows'):
        expected_anomalies = task_c_source_anomalies(manifest, daily_source, daily_epoch, independent_daily_reason) if daily_source else []
        errors.extend(review_data_anomaly_declaration(manifest, expected_anomalies))
        details['data_anomaly_skipped_count'] = len(expected_anomalies)
    owners: dict[str, list[str]] = {}
    invalid_rows: list[str] = []
    task_b_priority_invalid_rows: list[str] = []
    ten_day_send_invalid_rows: list[str] = []
    seen_routes: set[str] = set()
    for group in task_groups:
        candidate_path = resolver.resolve(group.get("candidate"))
        candidate = read_json(candidate_path)
        route = str(candidate.get("route") or "")
        seen_routes.add(route)
        if candidate.get("task") != "C" or candidate.get("forceGate") is not True:
            errors.append(f"任务C候选缺少授权标记: {candidate_path.name}")
        if candidate.get("taskBPriorityLockEnforced") is not True:
            errors.append(f"任务C候选缺少任务B过期优先级锁定: {candidate_path.name}")
        if candidate.get("serverDate") != server_date or candidate.get("tomorrowDate") != tomorrow_date:
            errors.append(f"任务C候选日期与manifest不一致: {candidate_path.name}")
        if normalize_hash(candidate.get("sourceSha256")) != normalize_hash(manifest.get("sourceSha256")):
            errors.append(f"任务C候选源表SHA-256与manifest不一致: {candidate_path.name}")
        candidate_source_path = resolver.resolve(candidate.get("sourcePath"), required=False)
        if manifest_source_path and candidate_source_path != manifest_source_path:
            errors.append(f"任务C候选源表路径与manifest不一致: {candidate_path.name}")
        if route not in expected_routes:
            errors.append(f"任务C候选路由无效: {route or '空'}")
            continue
        threshold, opr_pattern, expected_target = expected_routes[route]
        if daily_required and candidate.get('recentDailyReplyRateGuard') != {**daily_base, 'thresholdExclusive': threshold}:
            errors.append(f'任务C近两日回复率保护候选门禁不一致: {candidate_path.name}')
        if str(candidate.get("targetCarrierName") or "").strip() != expected_target:
            errors.append(f"任务C候选目标与OPR分区不一致: {candidate_path.name}")
        if float(candidate.get("replyRateThresholdExclusive") or 0) != threshold:
            errors.append(f"任务C候选阈值不一致: {candidate_path.name}")
        if float(candidate.get("replyRateSendThresholdInclusive") or 0) != 50:
            errors.append(f"任务C候选缺少4日发送>=50门槛: {candidate_path.name}")
        if candidate.get("tenDaySendThresholdExclusive") != 360 or candidate.get("tenDaySendMetric") != "10日发送窗口数":
            errors.append(f"任务C候选缺少10日发送窗口数严格>360门槛: {candidate_path.name}")
        rows = candidate.get("rows") or candidate.get("candidates") or []
        details["candidate_count"] += len(rows)
        for row in rows:
            phone = str(row.get("电话") or row.get("phone") or "").strip()
            if daily_required:
                index = row.get('源行号')
                source_row = daily_source.get(index, {}) if type(index) is int else {}
                # Compare the candidate with independently reopened source data;
                # a forged low rate in JSON cannot evade this check.
                source_phone = str(source_row.get('电话') or '').strip().removesuffix('.0')
                if (not source_row or source_phone != phone or index in daily_excluded_indices
                        or str(source_row.get('运营商') or '').strip() != str(row.get('运营商') or '').strip()
                        or str(source_row.get('opr') or '').strip() != str(row.get('opr') or '').strip()
                        or any(source_row.get(header) != row.get(header) or type(source_row.get(header)) is bool
                               for header in daily_headers)
                        or independent_daily_reason(source_row, threshold) is not None
                        or independent_daily_reason(row, threshold) is not None):
                    daily_invalid.append(phone or f'row:{index}')
            owners.setdefault(phone, []).append(route)
            send_count = row.get("4日发送窗口数")
            source = str(row.get("运营商") or row.get("原运营商") or "")
            expiry_status = str(row.get("过期状态") or "").strip()
            deadline = str(row.get("截止日") or "").strip()[:10]
            if (
                source in {"过期下架卡C", "过期下架卡T", "快过期隔离区C", "快过期隔离区T"}
                or expiry_status == "已过期"
                or (expiry_status == "快过期" and deadline == tomorrow_date)
            ):
                task_b_priority_invalid_rows.append(phone or f"row:{row.get('源行号') or 'unknown'}")
            try:
                parsed_send = float(send_count)
                send_valid = (
                    parsed_send >= 50
                    if REPLY_RATE_SEND_GATED_SOURCE_PATTERN.match(source)
                    else parsed_send > 0
                )
            except (TypeError, ValueError):
                send_valid = False
            raw_ten_day_send = row.get("10日发送窗口数")
            try:
                ten_day_send = float(str(raw_ten_day_send).strip().replace(",", ""))
                ten_day_send_valid = math.isfinite(ten_day_send) and ten_day_send > 360
            except (TypeError, ValueError):
                ten_day_send_valid = False
            if not ten_day_send_valid:
                ten_day_send_invalid_rows.append(phone or f"row:{row.get('源行号') or 'unknown'}")
            rate = task_c_rate(row.get("4日回复率"))
            opr = str(row.get("opr") or "")
            valid = send_valid and ten_day_send_valid and rate is not None and rate < threshold and bool(opr_pattern.search(opr))
            if route.endswith("chatr_unchanged_4day_reply_rate_lt4_to_swap"):
                valid = valid and str(row.get("换号状态") or "").strip() in {"", "未换号"}
            if not phone or not source.strip() or not valid:
                invalid_rows.append(phone or f"row:{row.get('源行号') or 'unknown'}")

    if seen_routes != set(expected_routes):
        errors.append(f"任务C候选路由不完整: {sorted(seen_routes)}")
    overlaps = {phone: routes for phone, routes in owners.items() if phone and len(routes) > 1}
    if {short_phone(phone) for phone in owners}.intersection(short_phone(row.get('phone')) for row in manifest.get('dataAnomalySkippedRows') or [] if isinstance(row, dict)):
        errors.append('任务C异常排除与候选出现同一号码身份，不能继续')
    if overlaps:
        errors.append(f"任务C两组候选存在{len(overlaps)}个重叠号码")
    if invalid_rows:
        errors.append(f"任务C候选中有{len(invalid_rows)}条不满足严格边界或未换号条件")
    if ten_day_send_invalid_rows:
        errors.append(f"任务C候选中有{len(ten_day_send_invalid_rows)}条未满足10日发送窗口数严格>360")
    if task_b_priority_invalid_rows:
        errors.append(
            f"任务C候选中有{len(task_b_priority_invalid_rows)}条覆盖任务B已过期或明日到期结果"
        )
    details["invalid_candidate_count"] = len(invalid_rows)
    details["overlap_count"] = len(overlaps)
    details["task_b_priority_invalid_count"] = len(task_b_priority_invalid_rows)
    details["ten_day_send_invalid_count"] = len(ten_day_send_invalid_rows)
    details['recent_daily_reply_guard_required'] = daily_required
    details['recent_daily_reply_invalid_count'] = len(daily_invalid)
    if daily_invalid:
        errors.append(f'任务C近两日回复率保护独立核验失败: {len(daily_invalid)}条')
    success = "任务C当前消息授权、任务B过期优先级锁定、10日发送>360、A/B来源4日发送>=50、4日严格边界、未换号条件、目标一致性及候选互斥通过"
    return status_result("任务C授权与规则", "task_c", errors, warnings, details, success)


def check_task_d_retention_counts(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    """Recompute the new D contract from local immutable evidence, never query."""
    errors: list[str] = []
    details: dict[str, Any] = {"contractVersion": config.get("taskDRetentionSummaryContractVersion")}
    fields = ("taskDRetentionSummary", "taskDRetentionPopulationReceipt", "taskDRetentionSummaryReceipt", "taskDRetentionGroupsReceipt")
    version = config.get("taskDRetentionSummaryContractVersion")
    if type(version) is int and version == 2:
        # v1 below is retained strictly for immutable historical receipts.
        # Current D still receives full independent notification auditing.
        if (config.get("taskDRetentionPopulationEnabled") is not False
                or config.get("taskDRetentionPopulationStatus") != "removed_by_user"
                or config.get("taskDRetentionPopulationReadExecuted") is not False):
            errors.append("已删除保留区统计必须明确关闭且未读取")
        if any(config.get(key) is not None for key in fields):
            errors.append("已删除保留区统计不得带入新数量或历史查询回执")
        if config.get("taskDRetentionCarrierTransferEnabled") is not False:
            errors.append("旧转入保留区功能不得恢复")
        for group in config.get("groups") or []:
            if str(group.get("route") or "").startswith("task_d_retain_") or str(group.get("title") or "").startswith("任务D"):
                errors.append("已删除保留区统计不得包含任务D运营商写入组")
        details.update(population_enabled=False, population_read_executed=False)
        return status_result("已删除保留区统计关闭核验", "task_d_retention", errors, [], details,
                             "保留区查询、统计及数量栏目已删除；两项下卡通知另行独立核验")
    if version is None:
        if any(config.get(key) is not None for key in fields):
            errors.append("新保留区统计缺少版本契约，不能按历史报告跳过")
        return status_result("任务D新增入区与结束存量", "task_d_retention", errors, [], details, "历史报告未声明保留区统计新契约")
    if version != 1:
        errors.append("任务D保留区统计契约版本无效")
    if config.get("taskDExecuted") is not True:
        if any(config.get(key) is not None for key in fields):
            errors.append("任务D未执行却出现保留区查询或统计，必须保持未知/null")
        return status_result("任务D新增入区与结束存量", "task_d_retention", errors, [], details, "任务D未执行，保留区统计未知且没有额外读取")
    if config.get("taskDCurrentUserAuthorized") is not True:
        errors.append("任务D保留区统计没有当前用户授权")
    try:
        paths = {key: resolver.resolve(config.get(key)) for key in fields if key != "taskDRetentionSummary"}
        stored = read_json(paths["taskDRetentionSummaryReceipt"])
        groups = read_json(paths["taskDRetentionGroupsReceipt"])
        run_id = str(config.get("runId") or "")
        if not run_id or stored.get("runId") != run_id or groups.get("runId") != run_id:
            errors.append("保留区统计/写入分组不是本轮runId")
        if sha256_file(paths["taskDRetentionGroupsReceipt"]) != stored.get("writeGroupsSha256"):
            errors.append("保留区写入分组SHA不一致")
        if sha256_file(paths["taskDRetentionPopulationReceipt"]) != stored.get("populationSha256"):
            errors.append("保留区全量回执SHA不一致")
        module_path = Path(__file__).with_name("opensms_task_d_retention_summary.py")
        spec = importlib.util.spec_from_file_location("opensms_task_d_retention_summary_audit", module_path)
        module = importlib.util.module_from_spec(spec)
        assert spec and spec.loader
        spec.loader.exec_module(module)
        bundle_groups = groups.get("groups") or []
        declared_groups = [group for group in config.get("groups") or [] if str(group.get("title") or "").startswith("任务D")]
        transfer_removed = config.get("taskDRetentionCarrierTransferEnabled") is False
        if transfer_removed:
            if (groups.get("retentionCarrierTransferEnabled") is not False
                    or groups.get("retentionCarrierTransferStatus") != "removed_by_user"
                    or bundle_groups or declared_groups):
                errors.append("任务D转入保留区工作已删除，统计不得包含运营商写入组")
        elif len(bundle_groups) != 2 or len(declared_groups) != 2:
            errors.append("保留区统计分组未绑定最终配置中的两个任务D实际写入组")
        else:
            def bound_group_facts(group):
                candidate_path = resolver.resolve(group.get("candidate"))
                candidate = read_json(candidate_path)
                facts = {"route": candidate.get("route"), "target": candidate.get("targetCarrierName")}
                for key in ("candidate", "validated", "write", "verify"):
                    facts[key] = sha256_file(resolver.resolve(group[key])) if group.get(key) else None
                return facts
            bundle_facts = sorted((bound_group_facts(group) for group in bundle_groups), key=lambda item: str(item["route"]))
            declared_facts = sorted((bound_group_facts(group) for group in declared_groups), key=lambda item: str(item["route"]))
            if bundle_facts != declared_facts:
                errors.append("保留区统计分组与任务D实际候选/写入/验证SHA不一致，疑似调包")
        completed = module.instant(groups.get("businessCompletedAt"), "任务D业务完成时间")
        for key, phase in (("taskDSwapExpiryReceipt", "swap_zero"), ("taskDRetentionExpiryReceipt", "retention_verified")):
            expiry = read_json(resolver.resolve(config.get(key)))
            if expiry.get("task") != "D" or expiry.get("phase") != phase or expiry.get("status") not in {"passed", "no_candidate", "recovered_transactional_receipt"} or expiry.get("businessReceiptPassed") is not True:
                errors.append("保留区统计早于任务D两组下卡通知业务完成")
            for time_key in ("createdAt", "completedAt", "ts"):
                if expiry.get(time_key) and module.instant(expiry[time_key], "任务D到期回执时间") > completed:
                    errors.append("保留区统计业务完成时间早于到期通知回执")
        rebuilt = module.build_summary(
            paths["taskDRetentionPopulationReceipt"], paths["taskDRetentionGroupsReceipt"], run_id,
            path_resolver=lambda value: resolver.resolve(value),
        )
        compare_keys = ("schemaVersion", "task", "status", "runId", "newTransferredCount", "totalInRetentionCount", "byCarrier", "asOf", "businessCompletedAt", "countingUnit", "populationSha256", "writeGroupsSha256", "newTransfers", "withinRetentionTransfers")
        for key in compare_keys:
            if stored.get(key) != rebuilt.get(key):
                errors.append(f"保留区统计独立重算不一致: {key}")
        # Paths may resolve to an archived byte-identical copy. The source
        # hashes, route identities, targets and derived counts must not drift.
        stable_proofs = lambda value: [{key: item.get(key) for key in ("route", "target", "candidateSha256", "validatedSha256", "writeSha256", "verifySha256", "newTransferredCount")} for item in value.get("proofGroups") or []]
        if stable_proofs(stored) != stable_proofs(rebuilt):
            errors.append("保留区逐路由SHA证据被修改或重算不一致")
        compact = config.get("taskDRetentionSummary")
        if not isinstance(compact, dict):
            errors.append("最终配置缺少保留区统计摘要")
        else:
            for key in ("status", "newTransferredCount", "totalInRetentionCount", "byCarrier", "asOf", "countingUnit", "populationSha256"):
                if compact.get(key) != rebuilt.get(key):
                    errors.append(f"最终配置保留区摘要不一致: {key}")
        details.update({"new_transferred_count": rebuilt["newTransferredCount"], "total_in_retention_count": rebuilt["totalInRetentionCount"], "by_carrier": rebuilt["byCarrier"], "as_of": rebuilt["asOf"], "counting_unit": "unique_phone10"})
    except (OSError, ValueError, TypeError, KeyError, AssertionError) as exc:
        errors.append(f"保留区统计证据不完整，禁止猜测为0: {exc}")
    return status_result("任务D新增入区与结束存量", "task_d_retention", errors, [], details, "逐号写前来源、真实写入、独立回读及结束全页存量重算一致")


def validate_task_d_filtered_partition(retention: dict[str, Any], run_id: str) -> dict[str, Any]:
    """Rebuild the partition from raw paginated query evidence, never skip claims."""
    retired_carriers = {
        "2093378431675928577": ("B14-C快下保留卡", "chatr"),
        "2088852494502465536": ("B14-T快下保留卡", "kooda"),
    }
    def require(condition: bool, message: str) -> None:
        if not condition:
            raise ValueError("任务D筛选分区证明: " + message)

    def number(value: Any, label: str) -> int:
        require(type(value) is int and value >= 0, label + "必须为非负整数")
        return value

    def cards(values: Any, label: str) -> set[str]:
        require(isinstance(values, list), label + "缺少号码列表")
        require(all(isinstance(value, str) and re.fullmatch(r"\d{10}", value) for value in values), label + "号码格式无效")
        require(len(values) == len(set(values)), label + "号码重复")
        return set(values)

    def identif(row: dict[str, Any]) -> str:
        return str(row.get("carrireIdentif") or row.get("carrierIdentif") or row.get("identif") or row.get("id") or "").strip()

    def scalar_id(value: Any, label: str) -> str:
        require((isinstance(value, str) and bool(value.strip()))
                or (type(value) is int and abs(value) <= 9007199254740991), label + "无效")
        return str(value).strip()

    def absent(value: Any) -> bool:
        return value is None or (isinstance(value, str) and not value.strip())

    def unconfigured(row: dict[str, Any]) -> bool:
        return (all(absent(row.get(key)) for key in (
                    "carrireIdentif", "carrierIdentif", "identif", "name", "carrierName", "carrireName", "carrier"))
                and row.get("displayStatus") == "未配置运营商"
                and row.get("showAble") is True and row.get("showInTable") is True
                and row.get("userAble") is False)

    require(bool(run_id) and retention.get("runId") == run_id, "回执运行编号不一致")
    require(retention.get("task") == "D" and retention.get("sourceProvenance") == "phone_card_retain_all_pages"
            and retention.get("endpoint") == "/phone/cardRetain/list" and retention.get("pageSize") == 200,
            "卡保留来源或每页200条门禁无效")
    source = cards(retention.get("cardNos"), "卡保留来源")
    total = number(retention.get("totalRecords"), "卡保留total")
    require(number(retention.get("fetchedRecords"), "卡保留实读数") == total, "卡保留记录未完整读取")
    records = retention.get("records")
    require(isinstance(records, list) and len(records) == total, "卡保留记录明细缺失或数量不符")
    require(all(isinstance(row, dict) and re.fullmatch(r"\d{10}", str(row.get("cardNo") or "")) for row in records), "卡保留记录含非法号码")
    require({row["cardNo"] for row in records} == source and retention.get("uniqueCardCount") == len(source), "卡保留去重来源与记录不符")
    page_count = max(1, (total + 199) // 200)
    source_pages = retention.get("pages")
    require(retention.get("pageCount") == page_count and isinstance(source_pages, list) and len(source_pages) == page_count,
            "卡保留分页数量不符")
    require(all(isinstance(page, dict) and page.get("pageNum") == index
                and page.get("rowCount") == min(200, max(0, total - (index - 1) * 200))
                for index, page in enumerate(source_pages, 1)), "卡保留分页序号或行数不完整")
    proof = retention.get("filteredQueryProof") or {}
    require(proof.get("version") == 1 and isinstance(proof.get("queries"), list), "缺少版本1原始查询分页")
    anomaly_policy = proof.get("dataAnomalyPolicy")
    require(anomaly_policy in (None, "skip_row_data_anomalies_v1")
            and retention.get("dataAnomalyPolicy") == anomaly_policy, "异常跳过契约未绑定原始查询")
    skip_anomalies = anomaly_policy == "skip_row_data_anomalies_v1"
    required_filters = {"koo": "kooda", "telus1": "kooda", "kini": "kooda", "chatr": "chatr"}
    covered = {key: set() for key in required_filters}
    matched: dict[str, dict[str, Any]] = {}
    backend_owners: dict[str, str] = {}
    for query in proof["queries"]:
        require(isinstance(query, dict), "查询结构无效")
        filter_value = query.get("filter")
        require(isinstance(filter_value, str) and filter_value in required_filters, "筛选条件无效")
        opr_class = required_filters[filter_value]
        require(query.get("method") == "POST" and query.get("endpoint") == "/phone/config/list"
                and query.get("oprClass") == opr_class, "号码查询端点或OPR分组无效")
        chunk = cards(query.get("cardNos"), "查询批次")
        require(0 < len(chunk) <= 500 and chunk.issubset(source), "查询批次超限或不是来源子集")
        require(not (covered[filter_value] & chunk), "同一筛选条件重复覆盖来源号码")
        covered[filter_value].update(chunk)
        pages = query.get("pages")
        require(isinstance(pages, list) and bool(pages) and isinstance(pages[0], dict), "号码查询分页缺失")
        query_total = number(pages[0].get("total"), "号码查询total")
        require(len(pages) == max(1, (query_total + 499) // 500), "号码查询遗漏分页")
        query_seen: set[str] = set()
        for index, page in enumerate(pages, 1):
            require(isinstance(page, dict) and page.get("pageNum") == index and page.get("pageSize") == 500
                    and page.get("code") == 200 and page.get("total") == query_total, "号码查询分页顺序、状态或total不一致")
            rows = page.get("rows")
            require(isinstance(rows, list) and len(rows) == min(500, max(0, query_total - (index - 1) * 500)), "号码查询分页行数不完整")
            for row in rows:
                require(isinstance(row, dict), "后台号码行无效")
                phone = short_phone(row.get("cardNo") or row.get("phone") or row.get("cardPhone") or row.get("tel"))
                require(phone in chunk and phone not in query_seen, "后台号码不属于查询批次或重复")
                query_seen.add(phone)
                backend_id = scalar_id(row.get("id"), "后台ID")
                require(backend_id not in backend_owners or backend_owners[backend_id] == phone, "后台ID对应多个号码")
                backend_owners[backend_id] = phone
                opr = str(row.get("opr") or "")
                is_c = "chatr" in opr.lower()
                is_t = re.search(r"kooda|koodo|kini|telus1", opr.lower()) is not None
                require(is_c != is_t and ("chatr" if is_c else "kooda") == opr_class
                        and filter_value in opr.lower(), "后台OPR不符合精确分组或筛选条件")
                skip = unconfigured(row)
                raw_ids = [row.get(key) for key in ("carrireIdentif", "carrierIdentif", "identif") if not absent(row.get(key))]
                raw_carrier_id = row.get("carrireIdentif") or row.get("carrierIdentif") or row.get("identif")
                valid_carrier_id = ((isinstance(raw_carrier_id, str) and bool(raw_carrier_id.strip()))
                                    or (type(raw_carrier_id) is int and abs(raw_carrier_id) <= 9007199254740991))
                if skip_anomalies:
                    canonical_ids = [str(value).strip() if ((isinstance(value, str) and value.strip())
                        or (type(value) is int and abs(value) <= 9007199254740991)) else json.dumps(value, sort_keys=True) for value in raw_ids]
                    require(len(set(canonical_ids)) <= 1, "后台运营商身份字段冲突")
                    require(row.get("displayStatus") != "未配置运营商" or (not raw_ids
                        and all(absent(row.get(key)) for key in ("name", "carrierName", "carrireName", "carrier"))), "未配置运营商状态与身份字段冲突")
                else:
                    require(row.get("displayStatus") != "未配置运营商" or skip, "未配置运营商状态与身份字段冲突")
                identity = {"backendId": backend_id, "opr": opr, "oprClass": opr_class,
                            "sourceCarrierIdentif": "" if skip or (skip_anomalies and not valid_carrier_id) else scalar_id(raw_carrier_id, "后台运营商标识")}
                if skip_anomalies:
                    identity.update(slot=str(row.get("slotIdentif") or "").strip(),
                                    rawCarrierIds=[row.get(key) for key in ("carrireIdentif", "carrierIdentif", "identif")])
                if skip:
                    identity.update(unconfigured=True, slot=str(row.get("slotIdentif") or "").strip())
                elif skip_anomalies and not valid_carrier_id:
                    identity.update(dataAnomaly="carrier_identity_missing", slot=str(row.get("slotIdentif") or "").strip(),
                                    rawCarrierIds=[row.get(key) for key in ("carrireIdentif", "carrierIdentif", "identif")])
                else:
                    require(bool(identity["sourceCarrierIdentif"]), "后台运营商标识缺失")
                    retired = retired_carriers.get(identity["sourceCarrierIdentif"])
                    if retired:
                        require(opr_class == retired[1]
                                and all(absent(row.get(key)) or scalar_id(row[key], "已删除运营商标识") == identity["sourceCarrierIdentif"]
                                        for key in ("carrireIdentif", "carrierIdentif", "identif"))
                                and all(absent(row.get(key)) or str(row[key]).strip() == retired[0]
                                        for key in ("name", "carrierName", "carrireName", "carrier")),
                                "已删除B14排除的身份、名称或OPR冲突")
                        identity.update(retired=True, retiredCarrierName=retired[0],
                                        slot=str(row.get("slotIdentif") or "").strip())
                require(phone not in matched or matched[phone] == identity, "别名查询发生跨组或后台身份冲突")
                matched[phone] = identity
    require(all(values == source for values in covered.values()), "并非每个必需筛选条件都完整覆盖来源")
    carrier_rows = retention.get("carrierConfigs")
    require(isinstance(carrier_rows, list), "缺少权威运营商配置")
    carrier_names: dict[str, list[str]] = {}
    for row in carrier_rows:
        require(isinstance(row, dict), "运营商配置行无效")
        name = str(row.get("name") or row.get("carrierName") or row.get("carrireName") or row.get("carrier") or "").strip()
        carrier_names.setdefault(identif(row), []).append("" if row.get("delData") is True else name)
    for item in matched.values():
        if item.get("unconfigured") is True or item.get("dataAnomaly"):
            item["sourceCarrierName"] = ""
            continue
        names = carrier_names.get(item["sourceCarrierIdentif"], [])
        if item.get("retired") is True:
            require(not names, "已删除B14标识仍出现在当前配置，不能安全排除")
            item["sourceCarrierName"] = ""
            continue
        require(len(names) <= 1, "运营商配置存在重复身份，不能作为单条异常跳过")
        if skip_anomalies and (not names or not names[0]):
            item.update(dataAnomaly="source_carrier_unresolved", sourceCarrierName="")
            continue
        require(len(names) == 1 and bool(names[0]), "命中号码的源运营商无法唯一精确解析")
        item["sourceCarrierName"] = names[0]
    protected = {phone for phone, item in matched.items() if item["sourceCarrierName"] in {"快过期续费C", "快过期续费T"}}
    excluded_rows = retention.get("neverProcessSourceExcludedRows")
    require(isinstance(excluded_rows, list) and len(excluded_rows) == len(protected), "续费保护证据数量不符")
    excluded_cards: set[str] = set()
    for row in excluded_rows:
        require(isinstance(row, dict), "续费保护行无效")
        phone = str(row.get("cardNo") or "")
        require(phone in protected and phone not in excluded_cards, "续费保护清单重复或包含非保护号码")
        excluded_cards.add(phone)
        require(all(str(row.get(key) or "") == matched[phone][key]
                    for key in ("backendId", "sourceCarrierIdentif", "sourceCarrierName")), "续费保护证据与原始查询身份不一致")
    require(retention.get("neverProcessSourceExcludedCount") == len(protected), "续费保护计数不符")
    unmatched = source - set(matched)
    require(cards(retention.get("filteredUnmatchedCards"), "筛选未命中") == unmatched
            and retention.get("filteredUnmatchedCount") == len(unmatched), "筛选未命中清单不是来源减实际命中")
    unconfigured_cards = {phone for phone, item in matched.items() if item.get("unconfigured") is True}
    skipped_rows = retention.get("unconfiguredCarrierSkippedRows", [])
    require(isinstance(skipped_rows, list) and all(isinstance(row, dict) for row in skipped_rows), "未配置运营商排除明细无效")
    require(cards([row.get("cardNo") for row in skipped_rows], "未配置运营商排除") == unconfigured_cards
            and cards(retention.get("unconfiguredCarrierSkippedCards", []), "未配置运营商排除清单") == unconfigured_cards
            and number(retention.get("unconfiguredCarrierSkippedCount", 0), "未配置运营商排除数") == len(unconfigured_cards),
            "未配置运营商排除集合不是原始完整查询的精确子集")
    for row in skipped_rows:
        expected = matched[row["cardNo"]]
        require(row.get("backendId") == expected["backendId"] and row.get("slot") == expected["slot"]
                and row.get("opr") == expected["opr"].strip() and row.get("reason") == "carrier_unconfigured",
                "未配置运营商排除明细身份不一致")
    retired_cards = {phone for phone, item in matched.items() if item.get("retired") is True}
    retired_rows = retention.get("retiredCarrierSkippedRows", [])
    require(isinstance(retired_rows, list) and all(isinstance(row, dict) for row in retired_rows), "已删除运营商排除明细无效")
    require(cards([row.get("cardNo") for row in retired_rows], "已删除运营商排除") == retired_cards
            and cards(retention.get("retiredCarrierSkippedCards", []), "已删除运营商排除清单") == retired_cards
            and number(retention.get("retiredCarrierSkippedCount", 0), "已删除运营商排除数") == len(retired_cards),
            "已删除运营商排除集合不是原始完整查询的精确子集")
    for row in retired_rows:
        expected = matched[row["cardNo"]]
        require(all(row.get(key) == expected[key] for key in (
                    "backendId", "slot", "sourceCarrierIdentif", "retiredCarrierName"))
                and row.get("opr") == expected["opr"].strip() and row.get("reason") == "carrier_retired_b14",
                "已删除运营商排除明细身份不一致")
    anomaly_cards = {phone for phone, item in matched.items() if item.get("dataAnomaly")}
    anomaly_rows = retention.get("dataAnomalySkippedRows", [])
    require(isinstance(anomaly_rows, list) and all(isinstance(row, dict) for row in anomaly_rows), "异常数据排除明细无效")
    require(cards([row.get("cardNo") for row in anomaly_rows], "异常数据排除") == anomaly_cards
            and cards(retention.get("dataAnomalySkippedCards", []), "异常数据排除清单") == anomaly_cards
            and number(retention.get("dataAnomalySkippedCount", 0), "异常数据排除数") == len(anomaly_cards), "异常数据排除集合不符")
    for row in anomaly_rows:
        expected = matched[row["cardNo"]]
        require(all(row.get(key) == expected[key] for key in ("backendId", "sourceCarrierIdentif"))
                and row.get("slot") == expected.get("slot", "") and row.get("opr") == expected["opr"].strip()
                and row.get("reason") == expected["dataAnomaly"], "异常数据排除身份或原因不一致")
    actionable = set(matched) - protected - unconfigured_cards - retired_cards - anomaly_cards
    return {"source": source, "matched": matched, "protected": protected, "unmatched": unmatched,
            "unconfigured": unconfigured_cards, "retired": retired_cards, "dataAnomalies": anomaly_cards,
            "actionable": actionable, "byOpr": {key: {phone for phone in actionable if matched[phone]["oprClass"] == key}
                                                for key in ("chatr", "kooda")}}


def validate_task_d_notification_rollup(config: dict[str, Any], resolver: PathResolver) -> dict[str, Any]:
    """Independently recalculate the new rollup, without its producer module."""
    def require(ok, message):
        if not ok:
            raise ValueError("任务D通知去重统计：" + message)

    summary = config.get("taskDNotificationSummary")
    require(isinstance(summary, dict) and summary.get("status") == "passed"
            and summary.get("version") == 1 and summary.get("runId") == config.get("runId"), "缺失或运行身份错误")
    path = resolver.resolve(summary.get("receipt"), required=False)
    require(path is not None and sha256_file(path) == summary.get("receiptSha256"), "统计回执哈希不符")
    require(read_json(path) == {k: v for k, v in summary.items() if k not in {"receipt", "receiptSha256"}}, "报告与统计回执内容不符")
    manifest_path = resolver.resolve(config.get("taskDManifest"))
    manifest = read_json(manifest_path)
    run_id = config.get("runId")
    require(manifest.get("runId") == run_id and isinstance(run_id, str) and bool(run_id)
            and run_id in manifest_path.name, "准备清单不是本轮")
    preparation_dir = manifest_path.resolve().parent
    receipt_dir = preparation_dir
    if preparation_dir.parent.name == "task_d_preparation_attempts":
        receipt_dir = preparation_dir.parent.parent
        require(re.fullmatch(r"[0-9]+", preparation_dir.name) is not None
                and receipt_dir.name == "receipts" and receipt_dir.parent.name == run_id,
                "准备尝试目录不属于本轮")
    populations, phase_rows, evidence = [], [], []
    for phase, notification_key, receipt_key in (
        ("swap_zero", "swapNotificationPath", "taskDSwapExpiryReceipt"),
        ("retention_verified", "retentionNotificationPath", "taskDRetentionExpiryReceipt"),
    ):
        notification_path = resolver.resolve(manifest.get(notification_key))
        receipt_path = resolver.resolve(config.get(receipt_key))
        notification, receipt = read_json(notification_path), read_json(receipt_path)
        resolved_receipt = receipt_path.resolve()
        continuation_receipt = (phase == "retention_verified"
            and resolved_receipt.parent.parent.name == "task_d_continuation_attempts"
            and re.fullmatch(r"[0-9]+", resolved_receipt.parent.name) is not None
            and resolved_receipt.parent.parent.parent == receipt_dir
            and receipt_dir.name == "receipts" and receipt_dir.parent.name == run_id)
        require(notification_path.resolve().parent == preparation_dir
                and (resolved_receipt.parent == receipt_dir or continuation_receipt)
                and all(run_id in p.name for p in (notification_path, receipt_path)),
                "源回执不是本轮目录/运行编号")
        if phase == "retention_verified":
            require(notification.get("runId") == run_id, "卡保留通知运行编号不符")
        else:
            source_workbook = resolver.resolve(manifest.get("sourcePath"))
            require(run_id in source_workbook.name and resolver.resolve(notification.get("sourcePath")) == source_workbook
                    and sha256_file(source_workbook) == manifest.get("sourceSha256") == notification.get("sourceSha256"),
                    "通知源工作簿绑定不符")
        for source in (notification_path, receipt_path):
            evidence.append({"path": str(source.resolve()), "sha256": sha256_file(source)})
        cards = notification.get("cardNos")
        require(isinstance(cards, list) and all(isinstance(x, str) and re.fullmatch(r"[0-9]{10}", x) for x in cards)
                and len(set(cards)) == len(cards), "通知号码缺失或重复")
        batches = receipt.get("batches")
        require(isinstance(batches, list) and receipt.get("validateOnly") is not True
                and receipt.get("phase") == phase, "分支回执不完整")
        actual, updated = [], 0
        for index, batch in enumerate(batches, 1):
            values = batch.get("cardNos")
            require(isinstance(values, list) and 0 < len(values) <= 20
                    and batch.get("batchNumber") == index, "批大小或顺序错误")
            actual.extend(values)
            raw = batch.get("serverResponse") or {}
            data = raw.get("data") or {}
            require(raw.get("code") == 200 and data.get("notFound") == []
                    and all(type(data.get(k)) is int and data[k] == len(values) for k in ("inputCount", "matchedCardCount"))
                    and type(data.get("updatedChatCount")) is int and data["updatedChatCount"] >= 0, "原始事务回执未确认")
            require(batch.get("expireDate") == notification.get("expireDate"), "批次到期日不符")
            updated += data["updatedChatCount"]
        require(len(actual) == len(cards) and set(actual) == set(cards), "逐批提交身份不是通知清单")
        populations.append(set(cards))
        phase_rows.append({"phase": phase, "expireDate": notification["expireDate"], "candidateCount": len(cards),
                           "confirmedPhoneCount": len(cards), "batchCount": len(batches), "updatedChatCount": updated})
    evidence.extend([{"path": str(manifest_path.resolve()), "sha256": sha256_file(manifest_path)},
                     {"path": str(source_workbook.resolve()), "sha256": sha256_file(source_workbook)}])
    expected = {"confirmedSubmissionCount": sum(map(len, populations)),
                "uniqueNotifiedPhoneCount": len(populations[0] | populations[1]),
                "overlapPhoneCount": len(populations[0] & populations[1]),
                "byPhase": phase_rows, "evidence": evidence,
                "countingUnit": "unique_phone10", "scope": "task_d_only_before_later_tasks",
                "verificationMode": "synchronous_transactional_business_receipt", "independentExpiryReadbackAvailable": False,
                "lastTaskDNotificationByDate": [
                    {"expireDate": phase_rows[0]["expireDate"], "phoneCount": len(populations[0] - populations[1])},
                    {"expireDate": phase_rows[1]["expireDate"], "phoneCount": len(populations[1])}]}
    for key, value in expected.items():
        require(summary.get(key) == value, f"{key}与独立重算不一致")
    return {key: summary[key] for key in ("confirmedSubmissionCount", "uniqueNotifiedPhoneCount", "overlapPhoneCount")}


def check_task_d(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    executed = config.get("taskDExecuted") is True
    authorized = config.get("taskDCurrentUserAuthorized") is True
    groups = [
        group for group in config.get("groups") or []
        if str(group.get("title") or "").startswith("任务D")
    ]
    details: dict[str, Any] = {
        "executed": executed,
        "current_user_authorized": authorized,
        "status": config.get("taskDStatus"),
        "group_count": len(groups),
        "retention_unique_cards": 0,
        "carrier_candidate_count": 0,
        "swap_notification_count": 0,
    }
    if not executed:
        if authorized:
            errors.append("任务D已获当前用户授权但未执行")
        if groups:
            errors.append("任务D未执行却出现任务D运营商候选或写入分组")
        if any(config.get(key) for key in (
            "taskDManifest",
            "taskDRetentionReceipt",
            "taskDSwapExpiryReceipt",
            "taskDRetentionExpiryReceipt",
            "taskDRetentionPopulationReceipt",
            "taskDRetentionSummaryReceipt",
            "taskDRetentionGroupsReceipt",
            "taskDRetentionSummary",
            "taskDNotificationSummary",
        )):
            errors.append("任务D未执行却配置了任务D候选、读取或写入回执")
        return status_result(
            "任务D下发下卡通知",
            "task_d",
            errors,
            warnings,
            details,
            "任务D无当前真实用户授权，已完整跳过且未触碰任务D业务",
        )
    if not authorized:
        errors.append("任务D已执行但缺少当前真实用户授权状态")
    manifest_path = resolver.resolve(config.get("taskDManifest"), required=False)
    if not manifest_path:
        errors.append("任务D缺少完整准备清单")
        return status_result("任务D下发下卡通知", "task_d", errors, warnings, details, "任务D规则与回执通过")
    manifest = read_json(manifest_path)
    transfer_removed = config.get("taskDRetentionCarrierTransferEnabled") is False
    if manifest.get("unconfiguredCarrierPolicy") == "skip_no_business_write_v1":
        try:
            metric = read_json(resolver.resolve(manifest.get("metricAnalysisPath")))
            metric_source = resolver.resolve(metric.get("sourcePath"))
            if (metric.get("unconfiguredCarrierPolicy") != "skip_no_business_write_v1"
                    or sha256_file(metric_source) != normalize_hash(metric.get("sourceSha256"))
                    or normalize_hash(metric.get("sourceSha256")) != normalize_hash(manifest.get("sourceSha256"))):
                raise ValueError("任务D指标来源或空运营商跳过门禁不一致")
            workbook = load_workbook(metric_source, read_only=True, data_only=True)
            try:
                sheet = workbook["sheet1"]
                headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows())]
                if headers.count("运营商") != 1:
                    raise ValueError("任务D指标源表运营商列缺失或重复")
                carrier_index = headers.index("运营商")
                skipped_indices = [index for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)
                                   if not str(row[carrier_index] or "").strip()]
            finally:
                workbook.close()
            if (metric.get("unconfiguredCarrierSkippedRows") != skipped_indices
                    or metric.get("unconfiguredCarrierSkippedCount") != len(skipped_indices)
                    or manifest.get("swapMetricUnconfiguredCarrierSkippedCount") != len(skipped_indices)
                    or any(not str(row.get("运营商") or "").strip() for row in metric.get("rows") or [])):
                raise ValueError("任务D指标空运营商跳过集合/计数与独立源表不一致")
            details["swap_metric_unconfigured_carrier_skipped_count"] = len(skipped_indices)
        except Exception as exc:
            errors.append(f"任务D空运营商独立核验失败: {exc}")
    if transfer_removed:
        if (manifest.get("retentionCarrierTransferEnabled") is not False
                or manifest.get("retentionCarrierTransferStatus") != "removed_by_user"
                or groups or any(manifest.get(key) for key in ("chatrCandidatePath", "koodaCandidatePath", "carrierCandidateCount"))
                or config.get("taskDCarrierCandidateCount") != 0 or config.get("taskDCarrierActualWrite") != 0):
            errors.append("任务D转入保留区工作已删除，候选和运营商写入必须为0")
    elif len(groups) != 2:
        errors.append("任务D缺少Chatr与Kooda/Koodo/Kini两组运营商结果")
    details["manifest_path"] = str(manifest_path)
    if (
        manifest.get("task") != "D"
        or manifest.get("forceGate") is not True
        or manifest.get("automationGate") is not True
    ):
        errors.append("任务D准备清单缺少当前用户授权门禁或程序门禁")
    if manifest.get("readOnlyPreparationConcurrent") is not True:
        errors.append("任务D未声明只读准备阶段并发")
    server_date = str(config.get("serverDate") or "")
    if manifest.get("serverDate") != server_date or manifest.get("swapExpireDate") != server_date:
        errors.append("任务D换号区到期日不是本轮服务器日期")
    try:
        plus_seven = (datetime.strptime(server_date, "%Y-%m-%d") + timedelta(days=7)).strftime("%Y-%m-%d")
    except ValueError:
        plus_seven = ""
        errors.append("任务D服务器日期格式无效")
    if manifest.get("retentionExpireDate") != plus_seven:
        errors.append("任务D卡保留到期日不是服务器日期+7天")

    retention_path = resolver.resolve(config.get("taskDRetentionReceipt"), required=False)
    if not retention_path:
        errors.append("任务D缺少卡保留全部分页回执")
        retention = {}
    else:
        retention = read_json(retention_path)
        if retention.get("sourceProvenance") != "phone_card_retain_all_pages":
            errors.append("任务D卡保留分页来源证明无效")
        if int(retention.get("pageSize") or 0) != 200:
            errors.append("任务D卡保留分页大小不是200")
        if int(retention.get("fetchedRecords") or 0) < int(retention.get("totalRecords") or 0):
            errors.append("任务D卡保留记录未读取全部页")
    try:
        partition = validate_task_d_filtered_partition(retention, str(config.get("runId") or ""))
    except (ValueError, TypeError, KeyError, AttributeError) as exc:
        errors.append(str(exc))
        partition = {"source": set(), "matched": {}, "protected": set(), "unmatched": set(), "unconfigured": set(), "retired": set(),
                     "actionable": set(), "byOpr": {"chatr": set(), "kooda": set()}}
    actionable_retention_cards = partition["actionable"]
    count_fields = {
        "retentionUniqueCardCount": len(partition["source"]),
        "retentionMatchedCardCount": len(partition["matched"]),
        "retentionActionableCardCount": len(actionable_retention_cards),
        "filteredUnmatchedCount": len(partition["unmatched"]),
        "neverProcessSourceExcludedCount": len(partition["protected"]),
        "chatrCandidateCount": 0 if transfer_removed else len(partition["byOpr"]["chatr"]),
        "koodaCandidateCount": 0 if transfer_removed else len(partition["byOpr"]["kooda"]),
    }
    for key, expected in count_fields.items():
        if type(manifest.get(key)) is not int or manifest.get(key) != expected:
            errors.append(f"任务D准备清单{key}与原始筛选查询重算不一致")
        report_key = "taskD" + key[0].upper() + key[1:]
        if type(config.get(report_key)) is not int or config.get(report_key) != expected:
            errors.append(f"任务D报告{report_key}与原始筛选查询重算不一致")
    for owner, prefix in ((manifest, ""), (config, "taskD")):
        anomaly_count_key = prefix + ("DataAnomalySkippedCount" if prefix else "dataAnomalySkippedCount")
        anomaly_cards_key = prefix + ("DataAnomalySkippedCards" if prefix else "dataAnomalySkippedCards")
        anomaly_policy_key = prefix + ("DataAnomalyPolicy" if prefix else "dataAnomalyPolicy")
        anomaly_values = owner.get(anomaly_cards_key, [])
        if (type(owner.get(anomaly_count_key, 0)) is not int
                or owner.get(anomaly_count_key, 0) != len(partition["dataAnomalies"])
                or not isinstance(anomaly_values, list) or any(not isinstance(value, str) for value in anomaly_values)
                or len(anomaly_values) != len(partition["dataAnomalies"]) or set(anomaly_values) != partition["dataAnomalies"]
                or owner.get(anomaly_policy_key) != retention.get("dataAnomalyPolicy")):
            errors.append("任务D异常数据跳过集合/契约与原始查询不一致")
        count_key = prefix + ("UnconfiguredCarrierSkippedCount" if prefix else "unconfiguredCarrierSkippedCount")
        cards_key = prefix + ("UnconfiguredCarrierSkippedCards" if prefix else "unconfiguredCarrierSkippedCards")
        count, values = owner.get(count_key, 0), owner.get(cards_key, [])
        if type(count) is not int or count != len(partition["unconfigured"]):
            errors.append(f"任务D{count_key}与原始筛选重算不一致")
        if (not isinstance(values, list) or any(not isinstance(value, str) for value in values)
                or len(values) != len(partition["unconfigured"]) or set(values) != partition["unconfigured"]):
            errors.append(f"任务D{cards_key}与原始筛选重算不一致")
        retired_count_key = prefix + ("RetiredCarrierSkippedCount" if prefix else "retiredCarrierSkippedCount")
        retired_cards_key = prefix + ("RetiredCarrierSkippedCards" if prefix else "retiredCarrierSkippedCards")
        retired_count, retired_values = owner.get(retired_count_key, 0), owner.get(retired_cards_key, [])
        if type(retired_count) is not int or retired_count != len(partition["retired"]):
            errors.append(f"任务D{retired_count_key}与原始筛选重算不一致")
        if (not isinstance(retired_values, list) or any(not isinstance(value, str) for value in retired_values)
                or len(retired_values) != len(partition["retired"]) or set(retired_values) != partition["retired"]):
            errors.append(f"任务D{retired_cards_key}与原始筛选重算不一致")
    for key, expected in (("filteredUnmatchedCards", partition["unmatched"]),
                          ("neverProcessSourceExcludedCards", partition["protected"])):
        values = manifest.get(key)
        if not isinstance(values, list) or len(values) != len(expected) or set(values) != expected:
            errors.append(f"任务D准备清单{key}与原始筛选查询重算不一致")
        report_key = "taskD" + key[0].upper() + key[1:]
        report_values = config.get(report_key)
        if not isinstance(report_values, list) or len(report_values) != len(expected) or set(report_values) != expected:
            errors.append(f"任务D报告{report_key}与原始筛选查询重算不一致")
    details.update({"retention_unique_cards": len(partition["source"]),
                    "retention_matched_cards": len(partition["matched"]),
                    "retention_actionable_cards": len(actionable_retention_cards),
                    "filtered_unmatched_count": len(partition["unmatched"]),
                    "filtered_unmatched_cards": sorted(partition["unmatched"]),
                    "unconfigured_carrier_skipped_count": len(partition["unconfigured"]),
                    "unconfigured_carrier_skipped_cards": sorted(partition["unconfigured"]),
                    "retired_carrier_skipped_count": len(partition["retired"]),
                    "retired_carrier_skipped_cards": sorted(partition["retired"]),
                    "data_anomaly_skipped_count": len(partition["dataAnomalies"]),
                    "protected_excluded_count": len(partition["protected"])})

    expected_routes = {} if transfer_removed else {
        "task_d_retain_chatr_to_b14_c": "B14-C快下保留卡",
        "task_d_retain_kooda_to_b14_t": "B14-T快下保留卡",
    }
    owners: dict[str, list[str]] = {}
    seen_routes: set[str] = set()
    for group in groups:
        candidate_path = resolver.resolve(group.get("candidate"))
        candidate = read_json(candidate_path)
        route = str(candidate.get("route") or "")
        seen_routes.add(route)
        if (
            candidate.get("task") != "D"
            or candidate.get("forceGate") is not True
            or candidate.get("automationGate") is not True
        ):
            errors.append(f"任务D运营商候选缺少当前用户授权门禁或程序门禁: {candidate_path.name}")
        if candidate.get("sourceProvenance") != "phone_card_retain_all_pages":
            errors.append(f"任务D运营商候选缺少全部分页来源证明: {candidate_path.name}")
        if str(candidate.get("targetCarrierName") or "") != expected_routes.get(route):
            errors.append(f"任务D运营商路由或目标不正确: {candidate_path.name}")
        opr_class = "chatr" if route == "task_d_retain_chatr_to_b14_c" else "kooda"
        rows = candidate.get("rows")
        if not isinstance(rows, list):
            errors.append(f"任务D候选缺少逐号证据: {candidate_path.name}")
            rows = []
        if candidate.get("candidates") is not None and candidate["candidates"] != rows:
            errors.append(f"任务D候选两份逐号列表不一致: {candidate_path.name}")
        row_cards = [str(row.get("cardNo") or "") for row in rows]
        if (len(row_cards) != len(set(row_cards)) or set(row_cards) != partition["byOpr"][opr_class]
                or any(candidate.get(key) != len(rows) for key in ("count", "candidateCount"))):
            errors.append(f"任务D {opr_class}候选未精确等于筛选命中减保护集合")
        for row, card in zip(rows, row_cards):
            owners.setdefault(card, []).append(route)
            expected = partition["matched"].get(card) or {}
            if any(str(row.get(key) or "") != str(expected.get(key) or "") for key in (
                "backendId", "opr", "oprClass", "sourceCarrierIdentif", "sourceCarrierName",
            )) or row.get("expectedTargetCarrierName") != expected_routes.get(route):
                errors.append(f"任务D运营商候选身份与筛选原始回执不符: {card}")
            if any(str(row[key]) != card for key in ("电话", "phone") if key in row):
                errors.append(f"任务D运营商候选手机号字段冲突: {card}")
    if seen_routes != set(expected_routes):
        errors.append(f"任务D运营商候选路由不完整: {sorted(seen_routes)}")
    if not transfer_removed and set(owners) != actionable_retention_cards:
        errors.append("任务D两类运营商候选未完整覆盖筛选实际命中且非续费保护的卡保留卡号")
    overlaps = {card: routes for card, routes in owners.items() if len(routes) != 1}
    if overlaps:
        errors.append(f"任务D运营商候选存在{len(overlaps)}个重叠卡号")
    details["carrier_candidate_count"] = len(owners)

    notification_path = resolver.resolve(manifest.get("retentionNotificationPath"), required=False)
    if not notification_path:
        errors.append("任务D缺少卡保留通知清单")
    else:
        notification = read_json(notification_path)
        if transfer_removed and (notification.get("retentionCarrierTransferEnabled") is not False
                or notification.get("retentionCarrierTransferStatus") != "removed_by_user"
                or notification.get("requiredCarrierTargets")):
            errors.append("任务D通知清单仍依赖已删除的B14运营商修改")
        notification_cards = notification.get("cardNos") or []
        anomaly_skip_cards = notification.get("dataAnomalySkippedCards", [])
        if (type(notification.get("dataAnomalySkippedCount", 0)) is not int
                or notification.get("dataAnomalySkippedCount", 0) != len(partition["dataAnomalies"])
                or not isinstance(anomaly_skip_cards, list) or any(not isinstance(value, str) for value in anomaly_skip_cards)
                or len(anomaly_skip_cards) != len(partition["dataAnomalies"]) or set(anomaly_skip_cards) != partition["dataAnomalies"]
                or notification.get("dataAnomalyPolicy") != retention.get("dataAnomalyPolicy")):
            errors.append("任务D通知清单异常数据排除与原始查询不一致")
        skip_cards = notification.get("unconfiguredCarrierSkippedCards", [])
        if (notification.get("unconfiguredCarrierSkippedCount", 0) != len(partition["unconfigured"])
                or not isinstance(skip_cards, list) or any(not isinstance(value, str) for value in skip_cards)
                or len(skip_cards) != len(partition["unconfigured"]) or set(skip_cards) != partition["unconfigured"]):
            errors.append("任务D通知清单的未配置运营商排除与原始查询不一致")
        retired_skip_cards = notification.get("retiredCarrierSkippedCards", [])
        if (type(notification.get("retiredCarrierSkippedCount", 0)) is not int
                or notification.get("retiredCarrierSkippedCount", 0) != len(partition["retired"])
                or not isinstance(retired_skip_cards, list) or any(not isinstance(value, str) for value in retired_skip_cards)
                or len(retired_skip_cards) != len(partition["retired"]) or set(retired_skip_cards) != partition["retired"]):
            errors.append("任务D通知清单的已删除运营商排除与原始查询不一致")
        if (notification.get("phase") != "retention_verified" or notification.get("expireDate") != plus_seven
                or notification.get("task") != "D" or notification.get("forceGate") is not True
                or notification.get("automationGate") is not True):
            errors.append("任务D卡保留通知阶段、日期或门禁错误")
        if (len(notification_cards) != len(actionable_retention_cards)
                or set(notification_cards) != actionable_retention_cards
                or candidate_count(notification) != len(actionable_retention_cards)):
            errors.append("任务D卡保留通知未精确覆盖筛选命中且非保护的号码")
        if (not retention_path or notification.get("retentionReceiptSha256") != sha256_file(retention_path)
                or resolver.resolve(notification.get("retentionReceiptPath"), required=False) != retention_path):
            errors.append("任务D卡保留通知未绑定原始完整筛选回执")

    swap_notification_path = resolver.resolve(manifest.get("swapNotificationPath"), required=False)
    if not swap_notification_path:
        errors.append("任务D缺少换号区通知清单")
        swap_notification = {}
    else:
        swap_notification = read_json(swap_notification_path)
        if swap_notification.get("phase") != "swap_zero" or swap_notification.get("expireDate") != server_date:
            errors.append("任务D换号区通知阶段或到期日错误")
        latest = str(swap_notification.get("latestReplyHeader") or "")
        previous = str(swap_notification.get("previousReplyHeader") or "")
        for row in swap_notification.get("rows") or []:
            if float(row.get(latest, -1)) != 0 or float(row.get(previous, -1)) != 0:
                errors.append("任务D换号区通知包含非连续两日0回复号码")
                break
            if str(row.get("liveCarrierName") or "") not in {"换号区C", "换号区T"}:
                errors.append("任务D换号区通知包含后台实时非换号区号码")
                break
    swap_count = candidate_count(swap_notification)
    details["swap_notification_count"] = swap_count

    for phase, configured_path, expected_count in (
        ("swap_zero", config.get("taskDSwapExpiryReceipt"), swap_count),
        ("retention_verified", config.get("taskDRetentionExpiryReceipt"), len(actionable_retention_cards)),
    ):
        receipt_path = resolver.resolve(configured_path, required=False)
        if not receipt_path:
            errors.append(f"任务D {phase}缺少到期标记回执")
            continue
        receipt = read_json(receipt_path)
        if receipt.get("status") not in {"passed", "no_candidate", "recovered_transactional_receipt"}:
            errors.append(f"任务D {phase}到期标记状态未通过")
        if receipt.get("businessReceiptPassed") is not True:
            errors.append(f"任务D {phase}业务回执未通过")
        expiry_verified = receipt.get("independentReadbackPassed") is True or (
            receipt.get("transactionalVerificationPassed") is True
            and receipt.get("verificationMode") == "synchronous_transactional_business_receipt"
            and receipt.get("independentReadbackAvailable") is False
        )
        if not expiry_verified:
            errors.append(f"任务D {phase}到期标记事务校验未通过")
        actual_input = sum(int(batch.get("inputCount") or 0) for batch in receipt.get("batches") or [])
        if actual_input != expected_count:
            errors.append(f"任务D {phase}到期标记输入数与候选数不一致")
        for batch in receipt.get("batches") or []:
            if batch.get("notFound"):
                errors.append(f"任务D {phase}存在未命中卡号")
            if int(batch.get("matchedCardCount") or 0) != int(batch.get("inputCount") or 0):
                errors.append(f"任务D {phase}业务命中数与输入数不一致")

    if config.get("taskDNotificationSummaryContractVersion") == 1:
        try:
            details["notification_rollup"] = validate_task_d_notification_rollup(config, resolver)
        except (ValueError, TypeError, KeyError, AttributeError, OSError) as exc:
            errors.append(str(exc))

    return status_result(
        "任务D下发下卡通知",
        "task_d",
        errors,
        warnings,
        details,
        "任务D两日0回复、卡保留全部分页和到期标记事务回执校验通过；转入保留区工作已删除"
        if transfer_removed else "历史任务D两日0回复、卡保留全部分页、B14-C/T写入和到期标记事务回执校验通过",
    )


def check_task_e(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    executed = config.get("taskEExecuted") is True
    authorized = config.get("taskECurrentUserAuthorized") is True
    details: dict[str, Any] = {
        "executed": executed,
        "current_user_authorized": authorized,
        "status": config.get("taskEStatus"),
        "eligible_workbook_rows": int(config.get("taskEEligibleWorkbookRowCount") or 0),
        "candidate_count": int(config.get("taskECandidateCount") or 0),
        "already_submitted_excluded": int(config.get("taskEAlreadySubmittedExcludedCount") or 0),
        "protected_excluded": int(config.get("taskENeverProcessSourceExcludedCount") or 0),
        "processed_count": int(config.get("taskEProcessedCount") or 0),
        "submitted_count": int(config.get("taskESubmittedCount") or 0),
        "not_written_count": int(config.get("taskENotWrittenCount") or 0),
        "matched_no_change_count": int(config.get("taskEMatchedNoChangeCount") or 0),
        "date_group_count": int(config.get("taskEDateGroupCount") or 0),
        "retention_records": int(config.get("taskERetentionTotalRecordCount") or 0),
        "retention_match_phones": int(config.get("taskERetentionDeadlineMatchPhoneCount") or 0),
    }
    configured_artifacts = any(config.get(key) for key in (
        "taskEManifest", "taskEReceipt", "taskELedger", "taskELedgerSnapshot", "taskERetentionReceipt",
    ))
    if not executed:
        if authorized:
            errors.append("任务E已获当前用户授权但未执行")
        if configured_artifacts:
            errors.append("任务E未执行却配置了候选、回执或去重台账")
        if config.get("taskERetentionReadExecuted") is True:
            errors.append("任务E未执行却读取了卡保留记录")
        if any(details[key] for key in (
            "eligible_workbook_rows",
            "candidate_count",
            "already_submitted_excluded",
            "protected_excluded",
            "processed_count",
            "submitted_count",
            "not_written_count",
            "matched_no_change_count",
            "date_group_count",
            "retention_records",
            "retention_match_phones",
        )):
            errors.append("任务E未执行却出现业务数量")
        return status_result(
            "任务E按截止日下卡通知",
            "task_e",
            errors,
            warnings,
            details,
            "任务E无当前真实用户授权，已完整跳过且未触碰候选、台账、后台查询或到期标记",
        )
    if not authorized:
        errors.append("任务E已执行但缺少当前真实用户授权状态")
    if config.get("taskERetentionReadExecuted") is not True:
        errors.append("任务E已执行但未记录卡保留全部分页读取")

    manifest_path = resolver.resolve(config.get("taskEManifest"), required=False)
    receipt_path = resolver.resolve(config.get("taskEReceipt"), required=False)
    if not manifest_path:
        errors.append("任务E缺少按截止日分组候选清单")
        return status_result("任务E按截止日下卡通知", "task_e", errors, warnings, details, "任务E规则与回执通过")
    if not receipt_path:
        errors.append("任务E缺少到期标记事务回执")
        return status_result("任务E按截止日下卡通知", "task_e", errors, warnings, details, "任务E规则与回执通过")

    manifest = read_json(manifest_path)
    receipt = read_json(receipt_path)
    details["manifest_path"] = str(manifest_path)
    details["receipt_path"] = str(receipt_path)
    if (
        manifest.get("task") != "E"
        or manifest.get("forceGate") is not True
        or manifest.get("automationGate") is not True
    ):
        errors.append("任务E候选清单缺少当前用户授权门禁或程序门禁")
    if manifest.get("ledgerDedupeKey") != "phone" or config.get("taskELedgerDedupeKey") != "phone":
        errors.append("任务E永久去重键不是phone")
    if set(manifest.get("eligibleExpiryStatuses") or []) != {"6天", "快过期", "已过期"}:
        errors.append("任务E候选状态范围不完整")

    retention_path = resolver.resolve(config.get("taskERetentionReceipt"), required=False)
    if not retention_path:
        errors.append("任务E缺少卡保留全部分页只读回执")
        retention = {}
    else:
        retention = read_json(retention_path)
        if retention.get("task") not in {"D", "E"}:
            errors.append("任务E卡保留回执任务类型无效")
        if retention.get("sourceProvenance") != "phone_card_retain_all_pages":
            errors.append("任务E卡保留分页来源证明无效")
        if retention.get("endpoint") != "/phone/cardRetain/list" or int(retention.get("pageSize") or 0) != 200:
            errors.append("任务E卡保留接口或分页大小无效")
        records = retention.get("records") or []
        fetched_records = int(retention.get("fetchedRecords") or 0)
        total_records = int(retention.get("totalRecords") or 0)
        if fetched_records < total_records or len(records) != fetched_records:
            errors.append("任务E卡保留记录未完整读取全部页")
        if int(retention.get("invalidCardRecordCount") or 0) != 0:
            errors.append("任务E卡保留回执含无效卡号")
        if normalize_hash(manifest.get("retentionReceiptSha256")) != sha256_file(retention_path):
            errors.append("任务E卡保留回执SHA-256与候选清单不一致")
        if manifest.get("retentionReceiptPath") != str(retention_path.resolve()):
            errors.append("任务E候选清单未指向当前卡保留回执")
        for manifest_key, receipt_key in (
            ("retentionPageCount", "pageCount"),
            ("retentionTotalRecordCount", "totalRecords"),
            ("retentionFetchedRecordCount", "fetchedRecords"),
            ("retentionUniqueCardCount", "uniqueCardCount"),
        ):
            if int(manifest.get(manifest_key) or 0) != int(retention.get(receipt_key) or 0):
                errors.append(f"任务E卡保留数量不一致: {manifest_key}")
        for config_key, manifest_key in (
            ("taskERetentionPageCount", "retentionPageCount"),
            ("taskERetentionTotalRecordCount", "retentionTotalRecordCount"),
            ("taskERetentionFetchedRecordCount", "retentionFetchedRecordCount"),
            ("taskERetentionUniqueCardCount", "retentionUniqueCardCount"),
        ):
            if int(config.get(config_key) or 0) != int(manifest.get(manifest_key) or 0):
                errors.append(f"任务E配置卡保留数量不一致: {config_key}")

    input_path = resolver.resolve(config.get("taskEInputWorkbook"), required=False)
    if not input_path:
        errors.append("任务E缺少输入工作簿")
    elif sha256_file(input_path) != normalize_hash(manifest.get("sourceSha256")):
        errors.append("任务E输入工作簿SHA-256与候选清单不一致")
    if input_path and retention:
        expected_unconfigured_rows = []
        expected_anomalies = []
        workbook = load_workbook(input_path, read_only=True, data_only=True)
        try:
            if "sheet1" not in workbook.sheetnames:
                errors.append("任务E输入工作簿缺少sheet1")
                workbook_deadlines = {}
            else:
                sheet = workbook["sheet1"]
                headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
                positions = {name: index for index, name in enumerate(headers) if name}
                if "电话" not in positions or "截止日" not in positions:
                    errors.append("任务E输入工作簿缺少电话或截止日")
                    workbook_deadlines = {}
                else:
                    workbook_deadlines: dict[str, set[str]] = {}
                    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                        source_values = dict(zip(headers, row))
                        source_carrier = str(source_values.get('运营商') or '').strip()
                        source_phone = short_phone(task_b_source_phone(source_values.get('电话')))
                        if (source_carrier and str(source_values.get('过期状态') or '').strip() in {'6天', '快过期', '已过期'}
                                and (not source_phone or task_ce_source_date(source_values.get('截止日'), workbook.epoch) is None)):
                            expected_anomalies.append(audit_source_anomaly(row_number, source_values,
                                'phone_missing' if not source_phone else 'deadline_missing_or_invalid', source_phone))
                        if (not str(source_values.get("运营商") or "").strip()
                                and str(source_values.get("过期状态") or "").strip() in {"6天", "快过期", "已过期"}):
                            expected_unconfigured_rows.append(row_number)
                        phone = short_phone(row[positions["电话"]] if positions["电话"] < len(row) else None)
                        raw_deadline = row[positions["截止日"]] if positions["截止日"] < len(row) else None
                        if not phone or raw_deadline in (None, ""):
                            continue
                        if isinstance(raw_deadline, datetime):
                            deadline = raw_deadline.date().isoformat()
                        elif hasattr(raw_deadline, "isoformat"):
                            deadline = raw_deadline.isoformat()
                        else:
                            match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", str(raw_deadline))
                            deadline = (
                                f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
                                if match else ""
                            )
                        if deadline:
                            workbook_deadlines.setdefault(phone, set()).add(deadline)
        finally:
            workbook.close()
        if manifest.get('dataAnomalyPolicy') is not None or manifest.get('dataAnomalySkippedRows'):
            errors.extend(review_data_anomaly_declaration(manifest, expected_anomalies))
            details['data_anomaly_skipped_count'] = len(expected_anomalies)
        if manifest.get("unconfiguredCarrierPolicy") == "skip_no_business_write_v1":
            skipped = manifest.get("unconfiguredCarrierSkippedRows") or []
            if (not isinstance(skipped, list) or any(not isinstance(row, dict) for row in skipped)
                    or [row.get("源行号") for row in skipped] != expected_unconfigured_rows
                    or manifest.get("unconfiguredCarrierSkippedCount") != len(expected_unconfigured_rows)
                    or any(str(row.get("运营商") or "").strip() for row in skipped)):
                errors.append("任务E未配置运营商跳过集合/计数与独立源表不一致")
            details["unconfigured_carrier_skipped_count"] = len(expected_unconfigured_rows)
        expected_matches = set()
        for record in retention.get("records") or []:
            phone = short_phone(record.get("cardNo"))
            expire_after = str(record.get("expireAfter") or "")
            deadlines = workbook_deadlines.get(phone) or set()
            if len(deadlines) > 1:
                errors.append(f"任务E卡保留号码在文档中存在不同截止日: {phone}")
            if len(deadlines) == 1 and expire_after in deadlines:
                expected_matches.add((phone, expire_after))
        actual_matches = {
            (short_phone(row.get("电话")), str(row.get("截止日") or ""))
            for row in manifest.get("retentionDeadlineMatches") or []
        }
        actual_matches.discard(("", ""))
        if actual_matches != expected_matches:
            errors.append("任务E卡保留新到期日与文档截止日一致提醒清单不正确")
        if int(manifest.get("retentionDeadlineMatchPhoneCount") or 0) != len({phone for phone, _ in expected_matches}):
            errors.append("任务E卡保留日期一致提醒号码数不正确")
        if int(manifest.get("retentionDeadlineMatchRecordCount") or 0) != len(expected_matches):
            errors.append("任务E卡保留日期一致提醒记录数不正确")
        if manifest.get("retentionComparisonDoesNotChangeCandidates") is not True:
            errors.append("任务E卡保留日期比较未声明不改变写入候选")
        details["retention_deadline_match_phones"] = len({phone for phone, _ in expected_matches})
        details["retention_deadline_match_records"] = len(expected_matches)
        if int(config.get("taskERetentionDeadlineMatchPhoneCount") or 0) != len({phone for phone, _ in expected_matches}):
            errors.append("任务E配置日期一致提醒号码数不正确")

    owners: dict[str, str] = {}
    candidate_rows = 0
    for group in manifest.get("groups") or []:
        expire_date = str(group.get("expireDate") or "")
        try:
            datetime.strptime(expire_date, "%Y-%m-%d")
        except ValueError:
            errors.append(f"任务E存在无效截止日组: {expire_date or '空'}")
        cards = [short_phone(value) for value in group.get("cardNos") or []]
        rows = group.get("rows") or []
        candidate_rows += len(cards)
        if int(group.get("count") or 0) != len(cards) or len(rows) != len(cards):
            errors.append(f"任务E截止日组计数不一致: {expire_date}")
        row_map = {short_phone(row.get("电话")): row for row in rows}
        for card in cards:
            if not card:
                errors.append(f"任务E截止日组包含空号码: {expire_date}")
                continue
            if card in owners:
                errors.append(f"任务E号码跨截止日组重复: {card}")
            owners[card] = expire_date
            row = row_map.get(card) or {}
            if row.get("截止日") != expire_date or row.get("过期状态") not in {"6天", "快过期", "已过期"}:
                errors.append(f"任务E号码状态或截止日证明无效: {card}")
            if row.get("运营商") in {"快过期续费C", "快过期续费T"}:
                errors.append(f"任务E候选包含快过期续费C/T保护号码: {card}")
            if manifest.get("unconfiguredCarrierPolicy") == "skip_no_business_write_v1" and not str(row.get("运营商") or "").strip():
                errors.append("任务E候选包含未配置运营商号码")
    if candidate_rows != int(manifest.get("candidateCount") or 0):
        errors.append("任务E候选总数与日期组不一致")
    if {short_phone(phone) for phone in owners}.intersection(short_phone(row.get('phone')) for row in manifest.get('dataAnomalySkippedRows') or [] if isinstance(row, dict)):
        errors.append('任务E异常排除与候选出现同一号码身份，不能继续')
    if len(manifest.get("groups") or []) != int(manifest.get("dateGroupCount") or 0):
        errors.append("任务E截止日组数不一致")
    if int(config.get("taskECandidateCount") or 0) != len(owners):
        errors.append("任务E配置候选数与清单不一致")
    if int(config.get("taskEDateGroupCount") or 0) != len(manifest.get("groups") or []):
        errors.append("任务E配置截止日组数与清单不一致")
    if int(manifest.get("neverProcessSourceExcludedCount") or 0) != len(manifest.get("neverProcessSourceExcludedRows") or []):
        errors.append("任务E续费保护排除数量不一致")
    if int(manifest.get("alreadySubmittedExcludedCount") or 0) != len(manifest.get("alreadySubmittedExcludedRows") or []):
        errors.append("任务E已提交排除数量不一致")
    raw_reconciled = (
        len(owners)
        + int(manifest.get("neverProcessSourceExcludedCount") or 0)
        + int(manifest.get("alreadySubmittedExcludedCount") or 0)
        + int(manifest.get("duplicateSameDeadlineCollapsedCount") or 0)
        + int(manifest.get("unconfiguredCarrierSkippedCount") or 0)
        + int(manifest.get("dataAnomalySkippedCount") or 0)
    )
    if raw_reconciled != int(manifest.get("eligibleWorkbookRowCount") or 0):
        errors.append("任务E文档命中原始行未与候选、保护、已提交和重复折叠完整核对")

    if receipt.get("task") != "E" or receipt.get("status") not in {"passed", "passed_with_not_found", "no_candidate"}:
        errors.append("任务E事务回执状态未通过")
    if receipt.get("businessReceiptPassed") is not True:
        errors.append("任务E业务回执未通过")
    if receipt.get("transactionalVerificationPassed") is not True:
        errors.append("任务E事务计数校验未通过")
    if receipt.get("ledgerVerificationPassed") is not True:
        errors.append("任务E永久去重台账校验未通过")
    if receipt.get("independentReadbackAvailable") is not False:
        errors.append("任务E未明确后台无到期标记独立查询接口")
    recovery_cards: dict[str, str] = {}
    recovery_path = resolver.resolve(receipt.get("recoveryProgressPath"), required=False)
    configured_recovery_path = resolver.resolve(config.get("taskERecoveryProof"), required=False)
    if configured_recovery_path != recovery_path:
        errors.append("任务E恢复证明配置与事务回执不一致")
    if recovery_path:
        recovery_proof = read_json(recovery_path)
        pending = recovery_proof.get("pendingBatch") or {}
        error_text = str(recovery_proof.get("error") or "")
        response_data = ((pending.get("response") or {}).get("data") or {})
        uncertain = recovery_proof.get("uncertainSubmission") is True or (
            recovery_proof.get("status") == "failed"
            and re.search(r"HTTP\s+5\d\d|gateway|time[- ]?out|timeout|ECONNRESET|socket hang up", error_text, re.I)
        )
        pending_date = str(pending.get("expireDate") or "")
        pending_cards = [short_phone(value) for value in pending.get("cardNos") or []]
        response_not_found = [short_phone(value) for value in response_data.get("notFound") or []]
        reconciled_partial = (
            recovery_proof.get("status") == "failed"
            and "业务回执不完整" in error_text
            and int(response_data.get("inputCount") or -1) == len(pending_cards)
            and int(response_data.get("matchedCardCount") or 0) + len(response_not_found) == len(pending_cards)
            and all(card in pending_cards for card in response_not_found)
            and int(response_data.get("updatedChatCount") if response_data.get("updatedChatCount") is not None else -1) >= 0
        )
        if (not uncertain and not reconciled_partial) or not pending_date or not pending_cards or any(not card for card in pending_cards):
            errors.append("任务E恢复证明不是有效的不确定提交或已核对部分业务回执")
        elif len(set(pending_cards)) != len(pending_cards):
            errors.append("任务E恢复证明包含重复号码")
        else:
            if any(owners.get(card) != pending_date for card in pending_cards):
                errors.append("任务E恢复证明的全部批次号码与本轮新鲜候选截止日不一致")
            recoverable_cards = [card for card in pending_cards if not reconciled_partial or card not in response_not_found]
            recovery_cards = {card: pending_date for card in recoverable_cards}
            if any(owners.get(card) != pending_date for card in recovery_cards):
                errors.append("任务E恢复证明与本轮新鲜候选的号码或截止日不一致")
        recovery_hash = normalize_hash(receipt.get("recoveryProgressSha256"))
        if recovery_hash and sha256_file(recovery_path) != recovery_hash:
            errors.append("任务E恢复证明SHA-256与事务回执不一致")
    elif int(receipt.get("recoveryCandidateCount") or 0):
        errors.append("任务E记录恢复候选但缺少恢复证明")
    if int(receipt.get("recoveryCandidateCount") or 0) != len(recovery_cards):
        errors.append("任务E恢复候选数与恢复证明不一致")
    if int(config.get("taskERecoveryCandidateCount") or 0) != len(recovery_cards):
        errors.append("任务E配置恢复候选数与恢复证明不一致")
    processed_cards: set[str] = set()
    submitted_cards: set[str] = set()
    not_written_cards: set[str] = set()
    not_written_rows: list[tuple[str, str, str]] = []
    matched_no_change_rows: list[tuple[str, str, str]] = []
    recovery_submitted_cards: set[str] = set()
    idempotent_recovery_no_change_count = 0
    input_count = 0
    for batch in receipt.get("batches") or []:
        card_values = [short_phone(value) for value in batch.get("cardNos") or []]
        cards = set(card_values)
        cards.discard("")
        if len(cards) != len(card_values):
            errors.append("任务E事务回执批次包含空号码或重复号码")
        if len(cards) > 20:
            errors.append("任务E单批超过20个号码，存在网关超时风险")
        if processed_cards & cards:
            errors.append("任务E事务回执批次存在重复号码")
        processed_cards.update(cards)
        batch_input = int(batch.get("inputCount") or 0)
        input_count += batch_input
        raw_not_found = [short_phone(value) for value in batch.get("notFound") or []]
        not_found = {card for card in raw_not_found if card}
        if len(not_found) != len(raw_not_found) or not_found - cards:
            errors.append("任务E批次未写入号码列表无效")
        matched_cards = cards - not_found
        matched_count = int(batch.get("matchedCardCount") or 0)
        if batch_input != len(cards) or matched_count != len(matched_cards) or matched_count + len(not_found) != batch_input:
            errors.append("任务E批次输入、命中和未写入数量无法核对")
        if submitted_cards & matched_cards or not_written_cards & not_found:
            errors.append("任务E成功或未写入号码跨批次重复")
        submitted_cards.update(matched_cards)
        not_written_cards.update(not_found)
        expected_batch_not_written = sorted(
            (card, str(batch.get("expireDate") or ""), "backend_not_found") for card in not_found
        )
        actual_batch_not_written = sorted(
            (
                short_phone(item.get("phone")),
                str(item.get("expireDate") or ""),
                str(item.get("reason") or ""),
            )
            for item in batch.get("notWrittenCards") or []
        )
        if int(batch.get("notWrittenCount") or 0) != len(not_found) or actual_batch_not_written != expected_batch_not_written:
            errors.append("任务E批次未写入明细与后台未命中列表不一致")
        not_written_rows.extend(expected_batch_not_written)
        recovery_retry = batch.get("recoveryRetry") is True
        updated_count = int(batch.get("updatedChatCount") if batch.get("updatedChatCount") is not None else -1)
        expected_no_change = sorted(
            (card, str(batch.get("expireDate") or ""), "backend_matched_no_change") for card in matched_cards
        ) if updated_count == 0 else []
        actual_no_change = sorted(
            (
                short_phone(item.get("phone")),
                str(item.get("expireDate") or ""),
                str(item.get("reason") or ""),
            )
            for item in batch.get("matchedNoChangeCards") or []
        )
        if int(batch.get("matchedNoChangeCount") or 0) != len(expected_no_change) or actual_no_change != expected_no_change:
            errors.append("任务E批次已匹配但无需变化明细不一致")
        matched_no_change_rows.extend(expected_no_change)
        if recovery_retry:
            recovery_submitted_cards.update(cards)
        if batch.get("idempotentRecoveryNoChange") is True:
            if not recovery_retry or updated_count != 0:
                errors.append("任务E幂等无变化标记与恢复回执不一致")
            else:
                idempotent_recovery_no_change_count += len(matched_cards)
        elif recovery_retry and matched_cards and updated_count == 0:
            errors.append("任务E恢复批次更新为0但未标记幂等无变化")
        if updated_count < 0:
            errors.append("任务E批次更新会话数无效")
        if int(batch.get("ledgerRecordedCount") or 0) != len(matched_cards):
            errors.append("任务E批次成功台账数量与实际命中号码数不一致")
        if any(batch.get("expireDate") != owners.get(card) for card in cards):
            errors.append("任务E批次到期日与候选日期组不一致")
    if processed_cards != set(owners):
        errors.append("任务E事务回执未完整覆盖去重后候选")
    if input_count != int(config.get("taskEProcessedCount") or 0):
        errors.append("任务E事务输入数与配置处理数不一致")
    if len(submitted_cards) != int(config.get("taskESubmittedCount") or 0):
        errors.append("任务E实际命中数与配置成功写入数不一致")
    if len(not_written_cards) != int(config.get("taskENotWrittenCount") or 0):
        errors.append("任务E未写入数与配置不一致")
    top_not_written_rows = sorted(
        (
            short_phone(item.get("phone")),
            str(item.get("expireDate") or ""),
            str(item.get("reason") or ""),
        )
        for item in receipt.get("notWrittenCards") or []
    )
    config_not_written_rows = sorted(
        (
            short_phone(item.get("phone")),
            str(item.get("expireDate") or ""),
            str(item.get("reason") or ""),
        )
        for item in config.get("taskENotWrittenCards") or []
    )
    if sorted(not_written_rows) != top_not_written_rows or top_not_written_rows != config_not_written_rows:
        errors.append("任务E未写入逐号明细在批次、回执和配置间不一致")
    if config.get("taskENotWrittenSheetRequired") is True:
        try:
            from opensms_task_e_unwritten_sheet import verify_sheet
            final_workbook = resolver.resolve(config.get("cardResult"))
            expected = [dict(phone=phone, expireDate=deadline, reason=reason)
                        for phone, deadline, reason in sorted(not_written_rows, key=lambda row: (row[1], row[0], row[2]))]
            verify_sheet(final_workbook, expected)
            details["not_written_sheet_independent_readback_passed"] = True
        except Exception as exc:
            errors.append(f"任务E未写入sheet独立核验失败: {exc}")
    if int(receipt.get("submittedCount") or 0) != len(submitted_cards):
        errors.append("任务E事务回执成功写入数不一致")
    if int(receipt.get("notWrittenCount") or 0) != len(not_written_cards):
        errors.append("任务E事务回执未写入数不一致")
    top_no_change_rows = sorted(
        (
            short_phone(item.get("phone")),
            str(item.get("expireDate") or ""),
            str(item.get("reason") or ""),
        )
        for item in receipt.get("matchedNoChangeCards") or []
    )
    config_no_change_rows = sorted(
        (
            short_phone(item.get("phone")),
            str(item.get("expireDate") or ""),
            str(item.get("reason") or ""),
        )
        for item in config.get("taskEMatchedNoChangeCards") or []
    )
    if sorted(matched_no_change_rows) != top_no_change_rows or top_no_change_rows != config_no_change_rows:
        errors.append("任务E已匹配但无需变化明细在批次、回执和配置间不一致")
    if int(receipt.get("matchedNoChangeCount") or 0) != len(top_no_change_rows):
        errors.append("任务E事务回执已匹配但无需变化数量不一致")
    if int(config.get("taskEMatchedNoChangeCount") or 0) != len(top_no_change_rows):
        errors.append("任务E配置已匹配但无需变化数量不一致")
    if receipt.get("status") == "passed_with_not_found" and not not_written_cards:
        errors.append("任务E无未写入号码却标记为部分通过")
    if receipt.get("status") == "passed" and not_written_cards:
        errors.append("任务E存在未写入号码却标记为完全通过")
    if int(receipt.get("markerBatchSize") or 0) != 20:
        errors.append("任务E回执的防504单批上限不是20")
    if not set(recovery_cards).issubset(recovery_submitted_cards):
        errors.append("任务E当前恢复证明号码未被恢复批次完整覆盖")
    if int(receipt.get("recoveryRetriedCount") or 0) != len(recovery_submitted_cards):
        errors.append("任务E恢复重试数与批次不一致")
    if int(config.get("taskERecoveryRetriedCount") or 0) != len(recovery_submitted_cards):
        errors.append("任务E配置恢复重试数与批次不一致")
    if int(receipt.get("idempotentRecoveryNoChangeCount") or 0) != idempotent_recovery_no_change_count:
        errors.append("任务E幂等无变化数与批次不一致")
    if int(config.get("taskEIdempotentRecoveryNoChangeCount") or 0) != idempotent_recovery_no_change_count:
        errors.append("任务E配置幂等无变化数与批次不一致")

    ledger_path = resolver.resolve(config.get("taskELedgerSnapshot"), required=False)
    if not ledger_path:
        ledger_path = resolver.resolve(config.get("taskELedger"), required=False)
    if submitted_cards:
        if not ledger_path:
            errors.append("任务E有实际提交但永久去重台账不存在")
        else:
            ledger = read_json(ledger_path)
            if ledger.get("schemaVersion") != 1 or ledger.get("dedupeKey") != "phone":
                errors.append("任务E永久去重台账结构无效")
            submissions = ledger.get("submissions") or {}
            missing = sorted(submitted_cards - set(submissions))
            if missing:
                errors.append(f"任务E永久去重台账缺少{len(missing)}个已提交号码")
            ledger_hash = normalize_hash(receipt.get("ledgerSha256After"))
            if ledger_hash and sha256_file(ledger_path) != ledger_hash:
                errors.append("任务E永久去重台账快照SHA-256与事务回执不一致")
    details["ledger_path"] = str(ledger_path) if ledger_path else None
    details["receipt_input_count"] = input_count
    return status_result(
        "任务E按截止日下卡通知",
        "task_e",
        errors,
        warnings,
        details,
        "任务E状态范围、截止日分组、永久手机号去重、事务回执与保护区门禁通过",
    )


def check_requested_downlist(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    downlist = config.get("downlist") or {}
    errors: list[str] = []
    candidate_path = resolver.resolve(downlist.get("candidate"))
    write_path = resolver.resolve(downlist.get("write"))
    verify_path = resolver.resolve(downlist.get("verify"))
    candidate = read_json(candidate_path)
    write = read_json(write_path)
    verify = read_json(verify_path)
    count = candidate_count(candidate)
    before = write.get("before") or {}
    after = write.get("after") or {}
    actual = int(before.get("eligibleCount") or 0)
    if int(before.get("candidateCount") or 0) != count:
        errors.append("二号机房下架写入回执候选数不一致")
    if int(before.get("backendValidCount") or 0) != count:
        errors.append("二号机房下架存在后台缺失或重复记录")
    if int(after.get("downCount") or 0) != count:
        errors.append("二号机房下架写入后仍有号码未下架")
    if int(after.get("stillActiveCount") or 0) != 0:
        errors.append("二号机房写入回执仍有上架号码")
    if verify.get("passed") is not True:
        errors.append("二号机房下架独立只读回读失败")
    if int(verify.get("candidateCount") or 0) != count:
        errors.append("二号机房下架独立回读数量不一致")
    if int(verify.get("downCount") or 0) != count:
        errors.append("二号机房下架独立回读未确认全部下架")
    details = {
        "candidate_path": str(candidate_path),
        "write_path": str(write_path),
        "verify_path": str(verify_path),
        "candidate_count": count,
        "backend_valid_count": int(before.get("backendValidCount") or 0),
        "already_down_count": int(before.get("alreadyDownCount") or 0),
        "actual_writes": actual,
        "successful_writes": actual if not errors else max(actual - int(after.get("stillActiveCount") or 0), 0),
        "failed_writes": int(after.get("stillActiveCount") or 0),
        "verified_down_count": int(verify.get("downCount") or 0),
    }
    return status_result("二号机房全量下架", "offline", errors, [], details, "二号机房候选、写入及独立回读通过")


def parse_sha256s(path: Path) -> list[tuple[str, str]]:
    entries: list[tuple[str, str]] = []
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        match = re.match(r"^([0-9a-fA-F]{64})\s+[* ]?(.+?)\s*$", line)
        if match:
            entries.append((match.group(1).lower(), match.group(2)))
    return entries


def task_f_artifact_reference(value: Any) -> bool:
    """Inspect an artifact reference as text only; never probe the F resource."""
    text = str(value or "")
    return bool(re.search(r"task[_-]?f|任务F|grabandgo|grab.?and.?go|ICCID库|截止日模板", text, re.IGNORECASE)
                or Path(text).suffix.lower() in {".db", ".sqlite", ".sqlite3"})


def google_artifact_reference(value: Any) -> bool:
    """Classify a reference without opening any Google file or receipt."""
    return bool(re.search(r"google|谷歌|association", str(value or ""), re.IGNORECASE))


def google_disabled_scope_requested(config: dict[str, Any]) -> bool:
    # Missing fields belong to historical runs; explicit but inconsistent
    # declarations enter this audit and fail instead of receiving an exemption.
    return config.get("googleAssociationEnabled") is False or bool(config.get("googleDisabledReason"))


def chatr_google_scope_requested(config: dict[str, Any]) -> bool:
    return config.get("googleChatrDateFallbackEnabled") is True


def audit_chatr_source_date(value: Any, epoch: datetime) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    parsed = value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value == int(value) and re.fullmatch(r"(?:19|20)\d{6}", str(int(value))):
            parsed = datetime.strptime(str(int(value)), "%Y%m%d")
        elif math.isfinite(value) and 20000 <= value <= 60000:
            parsed = from_excel(value, epoch)
    elif isinstance(value, str):
        raw = unicodedata.normalize("NFKC", value.strip()).strip("'\"")
        raw = raw.replace("年", "-").replace("月", "-").replace("日", "")
        if re.fullmatch(r"\d{8}", raw):
            parsed = datetime.strptime(raw, "%Y%m%d")
        else:
            match = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(.*)$", raw)
            if not match:
                raise ValueError("Chatr Google日期必须是无歧义完整年份日期")
            year, month, day, tail = match.groups()
            parsed = datetime.fromisoformat(f"{int(year):04d}-{int(month):02d}-{int(day):02d}{tail}")
    if isinstance(parsed, datetime):
        parsed = parsed.replace(tzinfo=None)
        return parsed.date().isoformat() if parsed.time() == datetime.min.time() else parsed.isoformat()
    if isinstance(parsed, date):
        return parsed.isoformat()
    raise ValueError("Chatr Google源日期无效")


def audit_chatr_snapshot_rows(source: Path, snapshot: dict[str, Any]) -> list[str]:
    """Re-read original cells to prove carrier scope and complete slot coverage."""
    errors: list[str] = []
    workbook = load_workbook(source, read_only=True, data_only=True)
    expected: dict[tuple[str, int], dict[str, Any]] = {}
    totals = {"totalSourceRows": 0, "chatrSourceRows": 0, "ignoredNonChatrRows": 0,
              "excludedInvalidSlotRows": 0}
    try:
        try:
            sheets = list(select_room_sheets(workbook))
        except ValueError as exc:
            return [f"Chatr日期Google来源范围错误: {exc}"]
        if snapshot.get("sourceSheets") != sheets:
            return ["Chatr日期快照与Google必需机房页签不一致"]
        for name in sheets:
            rows = workbook[name].iter_rows(values_only=True)
            header = [str(value).strip() if value is not None else "" for value in next(rows)]
            columns = (snapshot.get("sourceColumns") or {}).get(name) or {}
            required = ("卡槽", "電信商/價格", "激活日期", "預計過期日")
            if any(header.count(field) != 1 or columns.get(field) != header.index(field) + 1 for field in required):
                errors.append("Chatr日期快照的Google源字段位置不一致")
                continue
            for number, values in enumerate(rows, start=2):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                totals["totalSourceRows"] += 1
                carrier = values[header.index("電信商/價格")]
                normalized_carrier = unicodedata.normalize("NFKC", str(carrier or "")).casefold()
                if "chatr" not in normalized_carrier or re.search(r"kood[ao]|telus1", normalized_carrier):
                    totals["ignoredNonChatrRows"] += 1
                    continue
                totals["chatrSourceRows"] += 1
                raw_slot = values[header.index("卡槽")]
                slot = task_f_report_slot(raw_slot)
                if not re.fullmatch(r"\d+_\d+", slot):
                    totals["excludedInvalidSlotRows"] += 1
                    continue
                item = {"sourceSheet": name, "sourceRow": number, "rawSlot": raw_slot,
                        "normalizedSlot": slot, "carrier": carrier}
                for source_field, target in (("激活日期", "activationDate"), ("預計過期日", "cutoffDate")):
                    value = values[header.index(source_field)]
                    item[target] = audit_chatr_source_date(value, workbook.epoch)
                    item["raw" + target[0].upper() + target[1:]] = value.isoformat() if isinstance(value, (date, datetime)) else value
                expected[name, number] = item
    finally:
        workbook.close()
    seen: set[tuple[str, int]] = set()
    for record in snapshot.get("records") or []:
        provenance = record.get("provenance") or []
        for item in provenance:
            identity = (item.get("sourceSheet"), item.get("sourceRow"))
            raw = expected.get(identity)
            if identity in seen or raw is None or any(item.get(key) != value for key, value in (raw or {}).items()):
                errors.append("Chatr日期快照来源行与原始Google单元格不一致或重复")
            seen.add(identity)
        for field in ("activationDate", "cutoffDate"):
            dates = {item.get(field) for item in provenance if item.get(field) not in (None, "")}
            if len(dates) > 1 or record.get(field) != next(iter(dates), None):
                errors.append("Chatr日期快照归并字段与Google来源行不一致或冲突")
    if seen != set(expected):
        errors.append("Chatr日期快照未完整覆盖Google中的有效Chatr卡槽来源行")
    if any(snapshot.get(key) != value for key, value in totals.items()):
        errors.append("Chatr日期快照来源总行数或运营商排除数量不一致")
    return errors


def verified_chatr_continuation_logs(config: dict[str, Any], resolver: PathResolver,
                                    *, log_root: Path | None = None) -> list[str]:
    """Only immutable copies of the two already authorized Chatr command logs."""
    run_id = config.get("runId", "")
    if (not isinstance(run_id, str) or not re.fullmatch(r"\d{8}T\d{6}\+0800", run_id)
            or config.get("googleChatrDateFallbackEnabled") is not True
            or config.get("googleAssociationEnabled") is not False
            or config.get("googleChatrDateFallbackScope") != CHATR_GOOGLE_DATE_SCOPE):
        return []
    if not resolver.delivery_dir:
        return []
    delivery = Path(resolver.delivery_dir).resolve()
    root = delivery / "receipts"
    state_path = root / f"opensms_patrol_runner_state_{run_id}.json"
    if delivery.name != run_id or not state_path.is_file():
        return []
    state = read_json(state_path)
    if state.get("runId") != run_id or Path(state.get("delivery") or "").resolve() != delivery:
        return []
    logs = (log_root if log_root is not None else
            Path(__file__).resolve().parents[1] / "opensms_artifacts/logs/patrol_runner") / run_id
    names = {"google_chatr_download.log", "chatr_google_date_snapshot.log"}
    verified = []
    for attempt in state.get("taskDContinuationAttempts") or []:
        directory = Path(attempt.get("attemptDirectory") or "").resolve()
        if (attempt.get("authorizationSource") != "current-user"
                or directory.parent != root / "task_d_continuation_attempts"
                or re.fullmatch(r"[0-9]+", directory.name) is None):
            continue
        for evidence in attempt.get("archive") or []:
            source = Path(evidence.get("source") or "").resolve()
            archive = Path(evidence.get("archive") or "").resolve()
            digest = normalize_hash(evidence.get("sha256"))
            if (source.name not in names or source.parent != logs.resolve()
                    or archive.parent != directory
                    or re.fullmatch(r"[0-9]+_previous_" + re.escape(source.name), archive.name) is None
                    or not digest or not source.is_file() or not archive.is_file()):
                continue
            if sha256_file(source) == sha256_file(archive) == digest:
                verified.append(str(archive))
    return verified


def check_chatr_google_provenance(config: dict[str, Any], resolver: PathResolver,
                                 inputs: dict[str, Any]) -> CheckResult:
    """Admit only this run's narrowly scoped, SHA-bound Chatr date inputs."""
    errors: list[str] = []
    details: dict[str, Any] = {"allowed_artifact_paths": [], "google_artifacts_read": False}
    if (config.get("googleChatrDateFallbackEnabled") is not True
            or config.get("googleAssociationEnabled") is not False
            or config.get("googleChatrDateFallbackScope") != CHATR_GOOGLE_DATE_SCOPE
            or config.get("chatrDocumentDateSource") != CHATR_GOOGLE_DATE_SOURCE):
        errors.append("Chatr Google日期补缺未获得明确且限定范围的配置")
    paths: dict[str, Path] = {}
    for key, hash_key in (("chatrGoogleSource", "chatrGoogleSourceSha256"),
                          ("chatrGoogleDownloadReceipt", "chatrGoogleDownloadReceiptSha256"),
                          ("chatrGoogleDates", "chatrGoogleDatesSha256")):
        digest = normalize_hash(config.get(hash_key))
        if not config.get(key) or not digest:
            errors.append(f"Chatr日期来源缺少{key}路径或有效SHA256")
            continue
        if not errors:
            path = resolver.resolve(config[key], required=False)
            details["google_artifacts_read"] = details["google_artifacts_read"] or bool(path)
            if not path or sha256_file(path) != digest:
                errors.append(f"Chatr日期来源{key}文件或SHA256与本轮配置不一致")
            else:
                paths[key] = path
    if errors:
        return status_result("Chatr Google日期补缺来源", "freshness", errors, [], details, "")

    source = paths["chatrGoogleSource"]
    download_path = paths["chatrGoogleDownloadReceipt"]
    snapshot_path = paths["chatrGoogleDates"]
    download = read_json(download_path)
    snapshot = read_json(snapshot_path)
    run_id = config.get("runId")
    source_hash = normalize_hash(config.get("chatrGoogleSourceSha256"))
    snapshot_hash = normalize_hash(config.get("chatrGoogleDatesSha256"))

    def verify_copy_path(value: Any, digest: str, label: str, *, optional_original: bool = False) -> bool:
        if not isinstance(value, str) or not value or not Path(value).is_absolute():
            errors.append(f"{label}缺少明确绝对路径")
            return False
        raw = Path(value).expanduser()
        # Probe the real declared working/download location before the archive
        # resolver: a corrupt local copy must not be hidden by a good archive.
        actual = raw.resolve() if raw.is_file() else resolver.resolve(value, required=False)
        if actual is not None and sha256_file(actual) != digest:
            errors.append(f"{label}副本SHA256与本轮归档不一致")
            return False
        if actual is None and not optional_original:
            errors.append(f"{label}副本或等价归档缺失")
            return False
        return True

    source_working = config.get("processingChatrGoogleSource") or str(source)
    snapshot_working = config.get("processingChatrGoogleDates") or str(snapshot_path)
    verify_copy_path(source_working, source_hash, "Chatr Google处理源")
    verify_copy_path(snapshot_working, snapshot_hash, "Chatr日期处理快照")
    verify_copy_path(download.get("outputPath"), source_hash, "Chatr Google原始下载", optional_original=True)
    same_path = lambda left, right: bool(left and right and Path(left).expanduser().resolve() == Path(right).expanduser().resolve())
    downloaded = parse_iso(download.get("downloadedAt"))
    started = parse_iso(download.get("startedAt"))
    created = parse_iso(snapshot.get("createdAt"))
    run_started = parse_iso(config.get("runStartedAt") or config.get("startedAt"))
    if not run_started and re.fullmatch(r"\d{8}T\d{6}[+-]\d{4}", str(run_id or "")):
        run_started = datetime.strptime(run_id, "%Y%m%dT%H%M%S%z")
    if (not run_id or download.get("runId") != run_id or snapshot.get("runId") != run_id
            or not started or not downloaded or not created or not started <= downloaded <= created
            or (run_started and started < run_started)):
        errors.append("Chatr Google下载与快照未绑定本轮运行编号或新鲜时间顺序")
    if (download.get("sourceUrl") != CHATR_GOOGLE_SOURCE_URL
            or normalize_hash(download.get("sha256")) != source_hash
            or download.get("size") != source.stat().st_size):
        errors.append("Chatr Google下载回执来源、文件路径、大小或SHA256不一致")
    if (snapshot.get("schemaVersion") != 1 or snapshot.get("status") != PASS
            or snapshot.get("source") != "google_chatr_dates"
            or normalize_hash(snapshot.get("sourceSha256")) != source_hash
            or not same_path(snapshot.get("sourcePath"), source_working)
            or snapshot.get("readOnly") is not True or snapshot.get("snapshotConsistent") is not True
            or snapshot.get("sourceHashVerified") is not True):
        errors.append("Chatr日期快照未证明本轮Google源文件及只读一致性")
    records = snapshot.get("records")
    if (not isinstance(records, list) or type(snapshot.get("recordCount")) is not int
            or snapshot.get("recordCount") != len(records or [])):
        errors.append("Chatr日期快照记录结构或数量不一致")
        records = []
    seen: set[str] = set()
    for record in records:
        if not isinstance(record, dict):
            errors.append("Chatr日期快照包含非对象记录")
            break
        slot = record.get("normalizedSlot")
        provenance = record.get("provenance")
        if (not isinstance(slot, str) or not re.fullmatch(r"\d+_\d+", slot)
                or slot != task_f_report_slot(slot) or slot in seen
                or not isinstance(provenance, list) or not provenance):
            errors.append("Chatr日期快照包含重复/非规范卡槽或缺少来源行")
            break
        seen.add(slot)
        for item in [record, *provenance]:
            carrier = unicodedata.normalize("NFKC", str(item.get("carrier") or "")) if isinstance(item, dict) else ""
            if (not isinstance(item, dict) or not re.search(r"chatr", carrier, re.IGNORECASE)
                    or re.search(r"kood[ao]|telus1", carrier, re.IGNORECASE)
                    or item.get("normalizedSlot") != slot
                    or item.get("sourceSheet") not in (snapshot.get("sourceSheets") or [])
                    or type(item.get("sourceRow")) is not int or item.get("sourceRow", 0) < 2):
                errors.append("Chatr日期快照含Kooda/非Chatr来源或无效原始行身份")
                break
            for field in ("activationDate", "cutoffDate"):
                value = item.get(field)
                if value not in (None, "") and (not isinstance(value, str) or not parse_iso(value)):
                    errors.append("Chatr日期快照含无效激活日期或截止日")
    if not errors:
        errors.extend(audit_chatr_snapshot_rows(source, snapshot))
        if sha256_file(source) != source_hash:
            errors.append("Chatr Google源在独立快照复核期间发生变化")
    card = unwrap_card_receipt(read_json(inputs["card_receipt"]))
    process = card.get("process") or {}
    verification = card.get("verification") or {}
    proof = process.get("google_chatr_date_fallback") or {}
    review = verification.get("google_chatr_date_fallback_review") or {}
    resolution = process.get("date_resolution") or {}
    expected_proof = {"enabled": True, "status": PASS, "run_id": run_id,
                      "scope": CHATR_GOOGLE_DATE_SCOPE, "source": CHATR_GOOGLE_DATE_SOURCE,
                      "snapshot_sha256": snapshot_hash, "source_sha256": source_hash,
                      "record_count": len(records)}
    if (proof.get("enabled") is not True or review.get("enabled") is not True
            or any(proof.get(key) != value or review.get(key) != value for key, value in expected_proof.items())
            or review.get("matches") is not True
            or any(review.get(key) != value for key, value in proof.items())
            or resolution.get("source") != CHATR_GOOGLE_RESOLUTION_SOURCE
            or resolution.get("chatr_google_dates") != proof
            or (verification.get("date_resolution_review") or {}).get("chatr_google_dates") != proof
            or not same_path(proof.get("snapshot_path"), snapshot_working)
            or not same_path(proof.get("source_path"), source_working)):
        errors.append("Chatr日期生成与独立回读未一致绑定限定范围、来源及快照")
    field_counts = proof.get("field_counts") or {}
    all_counts = resolution.get("field_counts") or {}
    if (set(field_counts) != {"激活日期", "截止日"}
            or any(type(value) is not int or value < 0 for value in field_counts.values())
            or any((all_counts.get(field) or {}).get("google_chatr", 0) != value
                   for field, value in field_counts.items())):
        errors.append("Chatr日期补值数量与独立逐字段来源不一致")
    population_keys = ("chatr_rows", "matched_rows", "unmatched_rows", "filled_rows",
                       "missing_activation_rows", "missing_cutoff_rows")
    if (any(type(proof.get(key)) is not int or proof[key] < 0 or type(review.get(key)) is not int
            for key in population_keys)
            or proof.get("chatr_rows", -1) != proof.get("matched_rows", 0) + proof.get("unmatched_rows", 0)
            or proof.get("filled_rows", -1) > proof.get("matched_rows", -1)
            or any(proof.get(key, -1) > proof.get("chatr_rows", -1)
                   for key in ("missing_activation_rows", "missing_cutoff_rows"))):
        errors.append("Chatr主表总量、匹配、未匹配或补值后缺日期数量不一致")
    elif (any(value > proof["filled_rows"] for value in field_counts.values() if type(value) is int)
          or sum(value for value in field_counts.values() if type(value) is int) < proof["filled_rows"]
          or any(proof[missing] + (field_counts.get(field) or 0) > proof["chatr_rows"]
                 for field, missing in (("激活日期", "missing_activation_rows"), ("截止日", "missing_cutoff_rows")))):
        errors.append("Chatr实际补值字段数与补值行数或剩余缺字段数不一致")
    dispatch_inputs = (read_json(inputs["dispatch_receipt"]).get("inputs") or {})
    dispatched = dispatch_inputs.get("chatr_google_dates") or {}
    if (normalize_hash(dispatched.get("sha256")) != snapshot_hash
            or not same_path(dispatched.get("path"), snapshot_working)):
        errors.append("文档调度未绑定本轮Chatr日期快照路径和SHA256")
    commands = config.get("commandTimings") or []
    completed = {row.get("command") for row in commands
                 if isinstance(row, dict) and row.get("exitCode") == 0}
    if not CHATR_GOOGLE_COMMANDS <= completed:
        errors.append("缺少本轮Google下载及Chatr快照成功命令证据")
    details.update({"run_id": run_id, "scope": CHATR_GOOGLE_DATE_SCOPE,
                    "source": CHATR_GOOGLE_DATE_SOURCE, "record_count": len(records),
                    "source_sheets": snapshot.get("sourceSheets"), "field_counts": field_counts,
                    "snapshot_sha256": snapshot_hash, "source_sha256": source_hash,
                    "processing_source_path": source_working, "processing_snapshot_path": snapshot_working,
                    **{key: proof.get(key) for key in population_keys},
                    "downloaded_at": download.get("downloadedAt"), "created_at": snapshot.get("createdAt")})
    if not errors:
        archived_logs = verified_chatr_continuation_logs(config, resolver)
        details["archived_chatr_log_paths"] = archived_logs
        details["allowed_artifact_paths"] = [str(path) for path in paths.values()] + archived_logs
    return status_result("Chatr Google日期补缺来源", "freshness", errors, [], details,
                         "Chatr空白激活日期和截止日的本轮Google下载、快照及独立回读通过")


def check_ko_unactivated(config: dict[str, Any], inputs: dict[str, Any]) -> CheckResult:
    from ko_unactivated import review, SCOPE, SHEET, SLOT_SOURCE, REPORT_SCOPE, REPORT_SLOT_SOURCE
    from opensms_verify_card_slot_fast import load_system_kooda_snapshot
    errors, details = [], {}
    try:
        card = unwrap_card_receipt(read_json(inputs['card_receipt']))
        proof = (card.get('process') or {}).get('ko_unactivated_sheet') or {}
        verified = (card.get('verification') or {}).get('ko_unactivated_sheet_review') or {}
        if (proof != config.get('koUnactivatedSheet') or proof.get('scope') not in (SCOPE, REPORT_SCOPE)
                or proof.get('sheet') != SHEET or proof.get('slot_source') != (
                    REPORT_SLOT_SOURCE if proof.get('scope') == REPORT_SCOPE else SLOT_SOURCE)
                or verified.get('matches') is not True
                or any(verified.get(key) != value for key, value in proof.items())):
            raise ValueError('ko系统有-未激活生成及初始核验回执不一致')
        system = load_system_kooda_snapshot(Path(config['koUnactivatedSystemPath']),
            config['koUnactivatedSystemSha256'], config['runId'])
        book = load_workbook(inputs['result'], read_only=True, data_only=False)
        try:
            details = review(book, proof, config['reportDateCachePath'], config['reportDateCacheSha256'], config['runId'], True, system=system)
        finally:
            book.close()
        if details.get('matches') is not True:
            errors.append('最终刷新表ko未激活记录与本轮F冻结库及系统交集不一致')
    except (ValueError, KeyError, OSError, RuntimeError, TypeError) as exc:
        errors.append(str(exc))
    return status_result('ko系统有-未激活', 'document', errors, [], details, 'F未激活与本轮系统ICCID交集及最终表独立核验通过')


def check_google_difference(config: dict[str, Any], inputs: dict[str, Any]) -> CheckResult:
    from google_system_missing import review, SCOPE
    from opensms_verify_card_slot_fast import timestamp_value
    errors, details = [], {}
    try:
        card = unwrap_card_receipt(read_json(inputs["card_receipt"]))
        proof = (card.get("process") or {}).get("google_system_missing_sheet") or {}
        verified = (card.get("verification") or {}).get("google_system_missing_sheet_review") or {}
        if (config.get("googleSystemMissingSheetEnabled") is not True
                or config.get("googleAssociationEnabled") is not False
                or config.get("googleChatrDateFallbackEnabled") is not True
                or config.get("googleSystemMissingSheet") != proof
                or proof.get("scope") != SCOPE or proof.get("run_id") != config.get("runId")
                or proof.get("source_sha256") != config.get("chatrGoogleSourceSha256")
                or proof.get("snapshot_sha256") != config.get("chatrGoogleDatesSha256")
                or verified.get("matches") is not True):
            raise ValueError("Google差集范围、来源或本轮独立回读不一致")
        with inputs["source"].open("rb") as stream:
            source = load_workbook(stream, read_only=True, data_only=True)
            try:
                iterator = source.active.iter_rows(values_only=True)
                headers, rows = list(next(iterator)), list(iterator)
            finally:
                source.close()
        book = load_workbook(inputs["result"], read_only=True, data_only=False)
        try:
            details = review(book, headers, rows, proof, proof["snapshot_path"],
                             config["chatrGoogleDatesSha256"], config["runId"], timestamp_value)
        finally:
            book.close()
        if details.get("matches") is not True:
            errors.append("最终写后刷新表Google差集与本轮原始两份来源不一致")
    except (ValueError, KeyError, OSError) as exc:
        errors.append(str(exc))
    return status_result("所有谷歌有系统没有", "document", errors, [], details,
                         "本轮Google与系统卡槽差集及最终工作表独立重算通过")


def check_google_dependency(config: dict[str, Any], resolver: PathResolver,
                            inputs: dict[str, Any], chatr_provenance: CheckResult | None = None) -> CheckResult:
    """Independently prove a skipped dependency from ordinary run evidence."""
    errors: list[str] = []
    chatr_enabled = chatr_google_scope_requested(config)
    chatr_check = (chatr_provenance or check_chatr_google_provenance(config, resolver, inputs)) if chatr_enabled else None
    allowed_paths = set((chatr_check.details if chatr_check else {}).get("allowed_artifact_paths") or [])
    if chatr_check and chatr_check.status == FAIL:
        errors.append(chatr_check.message)
    if not chatr_enabled and (config.get("googleChatrDateFallbackEnabled") not in (None, False)
                             or any(config.get(key) for key in (
                                 "chatrGoogleSource", "chatrGoogleDownloadReceipt", "chatrGoogleDates"))):
        errors.append("未明确启用Chatr Google日期补缺却存在日期来源输入；未访问这些文件")
    def allowed_reference(value: Any) -> bool:
        if not allowed_paths or not value:
            return False
        return str(resolver.resolve(value, required=False)) in allowed_paths
    if config.get("googleAssociationEnabled") is not False or config.get("googleDisabledReason") != "user_disabled":
        errors.append("Google关闭模式或原因不一致")
    for key in ("association", "googleSha256", "googleReceipt", "googleWorkbook", "googleArchive"):
        if config.get(key):
            errors.append(f"Google依赖关闭但配置仍引用{key}；未访问该文件")
    for item in config.get("extraReceipts") or []:
        if ((google_artifact_reference(item.get("label")) or google_artifact_reference(item.get("path")))
                and not allowed_reference(item.get("path"))):
            errors.append("Google依赖关闭但额外回执仍引用Google产物；未访问该文件")
    commands = config.get("commandTimings")
    if not isinstance(commands, list) or not commands:
        errors.append("缺少本轮命令执行证据，不能确认Google未运行")
    else:
        for item in commands:
            command = str(item.get("command") or "")
            if google_artifact_reference(command) and not (allowed_paths and command in CHATR_GOOGLE_COMMANDS):
                errors.append("本轮命令存在Google采集或关联文件访问：" + command)
    for key in ("association", "google_receipt"):
        if inputs.get(key):
            errors.append(f"Google依赖关闭但审计输入仍引用{key}；未访问该文件")
    card = unwrap_card_receipt(read_json(inputs["card_receipt"]))
    process = card.get("process") or {}
    verification = card.get("verification") or {}
    proof = verification.get("google_dependency_review") or {}
    for label, dependency in (("生成回执", process.get("google_dependency") or {}),
                              ("独立回读", proof)):
        if dependency.get("enabled") is not False or any(dependency.get(key) != value for key, value in NO_GOOGLE_DEPENDENCY.items()):
            errors.append(f"{label}未证明Google跳过及系统主表设备范围")
    if proof.get("matches") is not True:
        errors.append("Google关闭独立回读未通过")
    system_review = verification.get("system_kooda_source_review") or {}
    if not config.get("runId") or system_review.get("run_id") != config.get("runId"):
        errors.append("Google关闭文档的系统来源独立回读未绑定本轮运行编号")
    resolution = process.get("date_resolution") or {}
    if not chatr_enabled and ((process.get("google_chatr_date_fallback") or {}).get("enabled") is True
                              or (resolution.get("chatr_google_dates") or {}).get("enabled") is True):
        errors.append("未明确启用Chatr Google日期补缺却存在补值回执")
    if resolution.get("association"):
        errors.append("日期来源仍引用Google关联文件；未访问该文件")
    field_counts = resolution.get("field_counts") or {}
    def google_count_present(value: Any) -> bool:
        if isinstance(value, dict):
            return any((str(key).lower() in ({"google"} if chatr_enabled else {"google", "google_chatr"})
                        and child not in (None, 0, False, ""))
                       or google_count_present(child) for key, child in value.items())
        return isinstance(value, list) and any(google_count_present(child) for child in value)
    if google_count_present(field_counts):
        errors.append("关闭Google后日期来源仍包含Google补值")
    if chatr_enabled and any((values or {}).get("google_chatr", 0)
                             for field, values in field_counts.items() if field not in {"激活日期", "截止日"}):
        errors.append("Chatr Google日期补缺越界到开卡日期或其他字段")
    dispatch = read_json(inputs["dispatch_receipt"])
    dispatch_inputs = dispatch.get("inputs") or {}
    dependency = dispatch_inputs.get("google_dependency") or dispatch.get("google_dependency") or {}
    if dependency.get("enabled") is not False or any(dependency.get(key) != value for key, value in NO_GOOGLE_DEPENDENCY.items()):
        errors.append("文档调度未证明Google关闭")
    if dispatch_inputs.get("association"):
        errors.append("文档调度仍引用Google关联文件；未访问该文件")
    if not chatr_enabled and dispatch_inputs.get("chatr_google_dates"):
        errors.append("未明确启用Chatr Google日期补缺却存在调度快照输入")
    processing_mode = dispatch.get("processing_mode")
    if processing_mode not in {"local", "local_fallback"}:
        errors.append("Google关闭流程未使用本机文档处理")
    if dispatch.get("api_attempts") or dispatch.get("job") or dispatch.get("api_base_url"):
        errors.append("Google关闭流程存在远端文档执行证据")
    if processing_mode == "local_fallback":
        # The dispatcher also uses this mode for --no-google/--force-local.
        # Accept that existing producer contract only with successful local
        # execution evidence and an explicit record of zero remote attempts.
        local_result = dispatch.get("local_result")
        local_result = local_result if isinstance(local_result, dict) else {}
        local_checks = local_result.get("checks")
        builder = dispatch.get("local_builder")
        if (dispatch.get("result") not in {"成功", "success"}
                or local_result.get("result") not in {"成功", "success"}
                or not isinstance(builder, str) or not Path(builder).is_absolute()
                or dispatch.get("api_attempts") != []
                or not isinstance(local_checks, dict) or not local_checks
                or not all(value is True for value in local_checks.values())
                or local_result.get("failed_checks") != []
                or any(not local_result.get(key) or local_result.get(key) != dispatch.get(key)
                       for key in ("output_path", "receipt_path"))):
            errors.append("本机回退缺少成功执行、检查或输出一致性证据")
    # Only inspect this run's directory entries, never open matching artifacts.
    delivery_dir = inputs.get("delivery_dir")
    unexpected: list[str] = []
    if delivery_dir:
        root = Path(delivery_dir)
        unexpected = [str(path.relative_to(root)) for path in root.rglob("*")
                      if not path.name.startswith("._") and google_artifact_reference(path.name)
                      and str(path.resolve()) not in allowed_paths]
        if unexpected:
            errors.append("本轮交付目录混入Google文件或回执；未访问这些文件")
    details = {**NO_GOOGLE_DEPENDENCY, "skip_contract_verified": not errors,
               "google_artifacts_read": bool(chatr_check and chatr_check.details.get("google_artifacts_read")),
               "unexpected_artifacts": unexpected,
               "chatr_date_fallback_enabled": chatr_enabled,
               "chatr_date_fallback_review": chatr_check.details if chatr_check else None,
               "google_dependency_review": proof, "processing_mode": processing_mode}
    if errors:
        return status_result("Google依赖关闭", "freshness", errors, [], details, "")
    return CheckResult("Google依赖关闭", "freshness", SKIP,
                       ("Google全面关联按用户要求跳过；仅Chatr空日期使用本轮Google来源，命令、独立回读及交付范围已核验"
                        if chatr_enabled else "Google依赖按用户要求跳过；命令、输入、独立回读及交付范围已核验"), details=details)


def check_delivery(delivery_dir: Path, hash_workers: int = 2, task_a_only: bool = False,
                   google_disabled: bool = False,
                   allowed_google_paths: Iterable[str] = ()) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    sums_path = delivery_dir / "SHA256SUMS.txt"
    files_path = delivery_dir / "FILES.tsv"
    if not sums_path.is_file():
        errors.append("缺少SHA256SUMS.txt")
        entries: list[tuple[str, str]] = []
    else:
        entries = parse_sha256s(sums_path)
        if not entries:
            errors.append("SHA256SUMS.txt没有有效记录")
    if not files_path.is_file():
        warnings.append("缺少FILES.tsv")

    def verify_entry(entry: tuple[str, str]) -> tuple[str, str | None]:
        expected, relative = entry
        target = delivery_dir / relative
        if (google_disabled and google_artifact_reference(relative)
                and str(target.resolve()) not in set(allowed_google_paths)):
            return relative, "Google依赖关闭但交付清单包含Google产物；未访问该文件"
        if task_a_only and task_f_artifact_reference(relative):
            return relative, "仅任务A交付清单包含F或数据库产物；未访问该文件"
        if not target.is_file():
            return relative, "文件不存在"
        actual = sha256_file(target)
        return relative, None if actual == expected else "SHA-256不一致"

    failures: dict[str, str] = {}
    if entries:
        with ThreadPoolExecutor(max_workers=max(1, min(hash_workers, len(entries))), thread_name_prefix="delivery-hash") as pool:
            for relative, error in pool.map(verify_entry, entries):
                if error:
                    failures[relative] = error
    if failures:
        errors.append(f"交付目录有{len(failures)}个文件校验失败")
    details = {
        "delivery_dir": str(delivery_dir),
        "checksum_entries": len(entries),
        "checksum_failures": failures,
        "hash_workers": max(1, min(hash_workers, len(entries))) if entries else 0,
    }
    return status_result("最终交付校验", "delivery", errors, warnings, details, "最终交付文件哈希全部通过")


def extra_receipt(config: dict[str, Any], keywords: Iterable[str]) -> str | None:
    for item in config.get("extraReceipts") or []:
        label = str(item.get("label") or "")
        if all(keyword in label for keyword in keywords):
            return item.get("path")
    return None


def newest(paths: list[Path]) -> Path | None:
    return max(paths, key=lambda item: item.stat().st_mtime) if paths else None


def infer_delivery_dir(config: dict[str, Any], resolver_seed: PathResolver) -> Path | None:
    source = resolver_seed.resolve(config.get("sourceArchive"), required=False)
    result = resolver_seed.resolve(config.get("cardResult"), required=False)
    if source and result and source.parent == result.parent:
        return source.parent
    return source.parent if source else None


def discover_inputs(
    config_path: Path,
    config: dict[str, Any],
    delivery_dir: Path | None,
    server_date_args: list[Path],
) -> tuple[PathResolver, dict[str, Any]]:
    manual_export = config.get("runType") in {"manual_export", "manual_export_task_b"}
    seed = PathResolver(config_path, delivery_dir)
    if delivery_dir is None:
        delivery_dir = infer_delivery_dir(config, seed)
    resolver = PathResolver(config_path, delivery_dir)
    source = resolver.resolve(config.get("sourceArchive"))
    result = resolver.resolve(config.get("cardResult"))
    card_receipt = resolver.resolve(config.get("cardReceipt"))
    balance = resolver.resolve(config.get("balance"), required=not manual_export)
    analysis = resolver.resolve(config.get("analysis"), required=not manual_export)

    export_value = extra_receipt(config, ["导出", "回执"])
    google_disabled = google_disabled_scope_requested(config)
    dispatch_value = extra_receipt(config, ["文档", "回执"])
    export_receipt = resolver.resolve(export_value, required=False) or newest(resolver.glob("opensms_export_once_*.json"))
    google_receipt = None
    if not google_disabled:
        google_value = extra_receipt(config, ["Google", "回执"])
        google_receipt = resolver.resolve(google_value, required=False) or newest(resolver.glob("opensms_google_card_slots_*_receipt.json"))
    dispatch_receipt = resolver.resolve(dispatch_value, required=False) or newest(resolver.glob("opensms_document_dispatch_*_receipt.json"))
    if not export_receipt or (not google_disabled and not google_receipt) or not dispatch_receipt:
        missing = [
            label
            for label, path in (
                ("导出回执", export_receipt),
                ("Google回执", google_receipt),
                ("文档调度回执", dispatch_receipt),
            )
            if not path and not (google_disabled and label == "Google回执")
        ]
        raise FileNotFoundError("、".join(missing) + "未找到")
    association = None
    if not google_disabled:
        google_data = read_json(google_receipt)
        association = resolver.resolve(config.get("association"), required=False)
        if not association:
            association = resolver.resolve(google_data.get("outputPath"), required=False)
        if not association:
            association = newest([path for path in resolver.glob("opensms_google_card_slots_*.xlsx") if "receipt" not in path.name])
        if not association:
            raise FileNotFoundError("本轮Google关联表未找到")

    if server_date_args:
        server_receipts = [Path(value).expanduser().resolve() for value in server_date_args]
    else:
        server_receipts = resolver.glob("opensms_server_date_*.json")
        selected: dict[str, Path] = {}
        for path in server_receipts:
            try:
                requested = str(read_json(path).get("requestedPath") or "")
            except (OSError, ValueError, json.JSONDecodeError):
                continue
            if requested in {"/day", "/phoneConfig"}:
                selected[requested] = path
        server_receipts = list(selected.values())
    dispatch_data = read_json(dispatch_receipt)
    server_date = str(((dispatch_data.get("inputs") or {}).get("server_date") or "")) or None
    return resolver, {
        "source": source,
        "result": result,
        "association": association,
        "card_receipt": card_receipt,
        "balance": balance,
        "analysis": analysis,
        "export_receipt": export_receipt,
        "google_receipt": google_receipt,
        "dispatch_receipt": dispatch_receipt,
        "server_receipts": server_receipts,
        "server_date": server_date,
        "delivery_dir": delivery_dir,
    }


CPU_CHECKS = frozenset({
    "check_xlsx", "check_carrier_refresh", "check_task_g",
    "check_google_difference", "check_ko_unactivated",
    "check_protected_carrier_sources", "check_forbidden_b13_targets",
})


def run_cpu_check(name: str, arguments: tuple[Any, ...]) -> tuple[CheckResult, int]:
    """Spawned workers independently reopen this audit's immutable evidence."""
    if name not in CPU_CHECKS:
        raise ValueError("CPU审计仅允许明确列出的本地只读核验")
    return globals()[name](*arguments), os.getpid()


def run_task(task: AuditTask, meter: ActivityMeter,
             cpu_pool: ProcessPoolExecutor | None = None) -> CheckResult:
    meter.enter()
    started = time.perf_counter()
    cpu_pid = None
    try:
        if cpu_pool is not None and task.cpu_call is not None:
            result, cpu_pid = cpu_pool.submit(run_cpu_check, *task.cpu_call).result()
        else:
            result = task.fn()
    except Exception as exc:  # One broken receipt must not hide the other checks.
        result = CheckResult(task.name, task.category, FAIL, f"检查程序异常: {type(exc).__name__}: {exc}")
    finally:
        duration_ms = round((time.perf_counter() - started) * 1000)
        thread_name = threading.current_thread().name
        meter.leave()
    result.duration_ms = duration_ms
    result.thread = thread_name + (f"/cpu:{cpu_pid}" if cpu_pid is not None else "")
    return result


def execute_audit_tasks(tasks: list[AuditTask], workers: int,
                        executor: str = "hybrid") -> tuple[list[CheckResult], ActivityMeter, dict[str, Any]]:
    """One bounded dispatch window; processes do not add a second job queue.

    Coordinator threads wait on CPU jobs, so at most `workers` checks run at
    once in total, not workers threads PLUS workers independent CPU checks.
    All jobs finish before cross-checks, either audit receipt or the next pass.
    Task F, DB checks and shared Chatr provenance retain their original code.
    """
    if executor not in {"threads", "hybrid"}:
        raise ValueError("未知本地审计执行器")
    if type(workers) is not int or not 1 <= workers <= min(16, os.cpu_count() or 1):
        raise ValueError("本地审计工作器数超出CPU或固定上限")
    meter = ActivityMeter()
    results: list[CheckResult] = []
    cpu_workers = min(workers, 4, sum(task.cpu_call is not None for task in tasks)) if executor == "hybrid" and workers > 1 else 0
    with ExitStack() as stack:
        # macOS spawn avoids forking a thread holding an input/cache lock.
        cpu_pool = stack.enter_context(ProcessPoolExecutor(
            max_workers=cpu_workers, mp_context=multiprocessing.get_context("spawn"))) if cpu_workers else None
        pool = stack.enter_context(ThreadPoolExecutor(max_workers=workers, thread_name_prefix="opensms-audit"))
        futures = [pool.submit(run_task, task, meter, cpu_pool) for task in tasks]
        for future in as_completed(futures):
            results.append(future.result())
    pids = sorted({int(item.thread.rsplit("/cpu:", 1)[1]) for item in results if "/cpu:" in item.thread})
    return results, meter, {
        "executor": "hybrid" if cpu_workers else "threads",
        "configured_cpu_workers": cpu_workers,
        "observed_cpu_processes": len(pids),
        "cpu_process_ids": pids,
        "cpu_checks": sum("/cpu:" in item.thread for item in results),
        "all_joined_before_cross_check": True,
        "total_inflight_limit": workers,
        "workbook_writers": 0,
        "business_writes": 0,
    }


def cross_check(
    config: dict[str, Any],
    inputs: dict[str, Any],
    results: list[CheckResult],
) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    by_name = {item.name: item for item in results}
    workbook = (by_name.get("卡槽关联结果表") or CheckResult("", "", FAIL, "")).details.get("workbook") or {}
    card = (by_name.get("卡槽关联完整复核") or CheckResult("", "", FAIL, "")).details
    analysis = (by_name.get("源报表分析") or CheckResult("", "", FAIL, "")).details
    row_values = [value for value in (workbook.get("main_rows"), card.get("main_rows"), analysis.get("rows")) if value is not None]
    if row_values and len({int(value) for value in row_values}) != 1:
        errors.append("源分析、卡槽回执和结果主表行数不一致")
    server = (by_name.get("服务器日期") or CheckResult("", "", FAIL, "")).details
    card_date = card.get("server_date")
    if card_date and server.get("day") and card_date != server.get("day"):
        errors.append("卡槽过期基准日期与后台/day日期不一致")
    if int(config.get("taskBExecuted") is True) == 0 and any(
        str(group.get("title") or "").startswith("任务B") for group in config.get("groups") or []
    ):
        errors.append("任务B授权状态不一致")
    if int(config.get("taskCExecuted") is True) == 0 and any(
        str(group.get("title") or "").startswith("任务C") for group in config.get("groups") or []
    ):
        errors.append("任务C授权状态不一致")
    details = {"row_counts": row_values, "server_date": server.get("day"), "card_server_date": card_date}
    return status_result("跨文件一致性", "cross_check", errors, warnings, details, "核心行数、日期和授权状态一致")


def _check_task_f_legacy(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    warnings: list[str] = []
    executed = config.get("taskFExecuted") is True
    authorized = config.get("taskFCurrentUserAuthorized") is True
    details: dict[str, Any] = {
        "executed": executed,
        "authorized": authorized,
        "default": "off",
        "login_required": config.get("taskFLoginRequired") is True,
    }
    if not executed:
        if authorized:
            errors.append("任务F未执行但授权状态为真")
        for key in (
            "taskFInventoryCount", "taskFAlreadyQueriedCount", "taskFPendingCount",
            "taskFFoundImportedCount", "taskFTransientErrorCount",
        ):
            if int(config.get(key) or 0) != 0:
                errors.append(f"任务F未授权但{key}不是0")
        for key in (
            "taskFLoginReceipt", "taskFInventoryReceipt", "taskFPendingReceipt",
            "taskFLookupReceipt", "taskFImportReceipt", "taskFDatabaseSnapshot",
        ):
            if config.get(key):
                errors.append(f"任务F未授权但生成了{key}")
        return status_result("任务F Kooda激活日期数据库", "task_f", errors, warnings, details, "任务F默认关闭且未触碰")

    if not authorized:
        errors.append("任务F执行但缺少当前真实用户授权")
    if config.get("taskFLoginRequired") is True:
        errors.append("任务F完成配置仍标记需要登录")

    paths: dict[str, Path | None] = {}
    for key in (
        "taskFLoginReceipt", "taskFInventoryReceipt", "taskFPendingReceipt",
        "taskFLookupReceipt", "taskFImportReceipt", "taskFDatabaseSnapshot",
    ):
        paths[key] = resolver.resolve(config.get(key), required=False)
        if paths[key] is None:
            errors.append(f"任务F缺少{key}")
    if errors:
        return status_result("任务F Kooda激活日期数据库", "task_f", errors, warnings, details, "任务F回执不完整")

    login = read_json(paths["taskFLoginReceipt"])
    inventory = read_json(paths["taskFInventoryReceipt"])
    pending = read_json(paths["taskFPendingReceipt"])
    lookup = read_json(paths["taskFLookupReceipt"])
    imported = read_json(paths["taskFImportReceipt"])
    if login.get("status") not in {"authenticated", "passed"} or login.get("loginStatus") not in {None, "authenticated"}:
        errors.append("任务F Grab&Go登录回执未通过")
    if not str(login.get("browser") or "").startswith("BitBrowser 430"):
        errors.append("任务F登录回执未证明OpenSMS库存仍使用BitBrowser 430")
    if login.get("silentBrowserMode") is not True or login.get("browserForegroundActivationAttempted") is not False:
        errors.append("任务F登录检查未证明后台静默且未抢占焦点")
    if inventory.get("status") != "passed" or int(inventory.get("pageSize") or 0) != 500 or inventory.get("oprSearch") != "koo":
        errors.append("任务F OpenSMS Kooda全量清单回执无效")
    if inventory.get("fullSourceDataIncluded") is not True:
        errors.append("任务F库存回执未声明保留完整OpenSMS源数据")
    inventory_rows = [row for row in inventory.get("rows") or [] if isinstance(row, dict)]
    if any(not isinstance(row.get("sourceData"), dict) for row in inventory_rows):
        errors.append("任务F库存回执存在缺失完整OpenSMS源对象的ICCID")
    if int(inventory.get("missingCardStoreCount") or 0) or int(inventory.get("invalidCardStoreCount") or 0):
        errors.append("任务F Kooda清单存在缺失或无效卡存储")
    if int(inventory.get("uniqueIccidCount") or 0) != int(config.get("taskFInventoryCount") or 0):
        errors.append("任务F Kooda唯一ICCID计数不一致")
    if pending.get("status") != "passed" or pending.get("terminalStatuses") != ["found", "not_found", "activation_missing"]:
        errors.append("任务F数据库待查去重回执无效")
    expected_web_fields = ["Purchase Date", "Act. Request Date", "Activation Date", "SIM Number", "SKU"]
    if int(pending.get("recordSchemaVersion") or 0) != 2 or pending.get("recordedWebFields") != expected_web_fields:
        errors.append("任务F待查回执不是完整字段结构v2")
    if int(pending.get("alreadyQueriedCount") or 0) != int(config.get("taskFAlreadyQueriedCount") or 0):
        errors.append("任务F已查跳过计数不一致")
    if int(pending.get("pendingCount") or 0) != int(config.get("taskFPendingCount") or 0):
        errors.append("任务F待查计数不一致")
    if lookup.get("status") != "passed" or lookup.get("loginStatus") != "authenticated":
        errors.append("任务F Grab&Go查询回执未完整通过")
    if not str(lookup.get("browser") or "").startswith("BitBrowser 430") or int(lookup.get("perWorkerDelayMs") or 0) < 1000:
        errors.append("任务F浏览器或每线程等待时间不符合规则")
    if lookup.get("silentBrowserMode") is not True or lookup.get("browserForegroundActivationAttempted") is not False:
        errors.append("任务F查询未证明后台静默且未抢占焦点")
    if lookup.get("transportPriority") != ["same-origin-http", "background-page-javascript"]:
        errors.append("任务F未按同源HTTP优先、后台页面JavaScript回退执行")
    accepted_lookup_modes = {
        "task-f-grabandgo-read-only-webforms-bitbrowser-430",
        "task-f-grabandgo-read-only-webforms-independent-http-sessions",
    }
    if lookup.get("mode") not in accepted_lookup_modes:
        errors.append("任务F查询模式不是受控只读模式")
    if (
        int(lookup.get("recordSchemaVersion") or 0) != 2
        or lookup.get("recordedWebFields") != expected_web_fields
        or lookup.get("fullSourceDataIncluded") is not True
    ):
        errors.append("任务F查询回执未声明网页五字段与完整源数据结构")
    for row in [item for item in lookup.get("results") or [] if isinstance(item, dict)]:
        if not isinstance(row.get("sourceOccurrences"), list) or not row.get("sourceOccurrences"):
            errors.append(f"任务F查询结果缺少完整源记录: {row.get('queryIccid') or 'unknown'}")
            break
        if row.get("status") in {"found", "activation_missing"}:
            presence = row.get("webFieldPresence") if isinstance(row.get("webFieldPresence"), dict) else {}
            if any(presence.get(key) is not True for key in ("purchaseDate", "actRequestDate", "activationDate", "simNumber", "sku")):
                errors.append(f"任务F查询结果缺少网页字段节点: {row.get('queryIccid') or 'unknown'}")
                break
    if lookup.get("mode") == "task-f-grabandgo-read-only-webforms-bitbrowser-430" and (
        not login.get("cdpEndpoint") or login.get("cdpEndpoint") != lookup.get("cdpEndpoint")
    ):
        errors.append("任务F浏览器回退模式的登录与查询没有复用同一个BitBrowser 430 CDP会话")
    requested_workers = int(lookup.get("requestedWorkers") or 0)
    effective_workers = int(lookup.get("effectiveWorkers") or 0)
    input_count = int(lookup.get("inputCount") or 0)
    if not 1 <= effective_workers <= 8:
        errors.append("任务F实际并发工作线程数无效")
    if input_count >= 8 and requested_workers != 8:
        errors.append("任务F待查数不少于8但未请求8个工作器")
    if lookup.get("adaptiveConcurrency") is not True or lookup.get("sharedSessionForcesSerial") is not True:
        errors.append("任务F查询回执未声明自适应并发与共享会话强制串行")
    if effective_workers > 1 and (
        lookup.get("independentAspNetSessions") is not True
        or int(lookup.get("uniqueAspNetSessionCount") or 0) != effective_workers
    ):
        errors.append("任务F多工作器未证明每个工作器拥有独立ASP.NET会话")
    if lookup.get("workerSessionMode") == "independent-form-login-unique-aspnet-session-per-worker" and (
        lookup.get("independentFormLoginAttempted") is not True
        or lookup.get("independentFormLoginSucceeded") is not True
    ):
        errors.append("任务F独立表单登录会话缺少成功证明")
    if effective_workers == 1 and requested_workers > 1 and (
        lookup.get("workerSessionMode") != "shared-aspnet-session-serial-worker"
        or lookup.get("concurrencyLimitedReason") != "shared_server_session_forces_serial_execution"
    ):
        errors.append("任务F共享会话降为单工作器时缺少准确原因")
    if lookup.get("sharedAspNetSessionConcurrentWorkers") is True:
        errors.append("任务F禁止在共享ASP.NET会话上并发查询")
    if lookup.get("strictIdentityGuard") != "queryIccid_equals_listSimNumber_equals_detailSimIccid":
        errors.append("任务F未证明查询、列表和详情ICCID三重一致性保护")
    if lookup.get("identityMismatchStoredAsTerminal") is not False:
        errors.append("任务F身份不一致可能被错误写为数据库终态")
    if lookup.get("sharedSessionNotFoundRequiresSerialConfirmation") is not True:
        errors.append("任务F未声明共享会话not_found必须串行确认")
    if lookup.get("sharedAspNetSessionConcurrentWorkers") is True:
        for row in [item for item in lookup.get("results") or [] if isinstance(item, dict)]:
            if row.get("status") == "not_found" and row.get("sharedSessionSerialConfirmation") is not True:
                errors.append(f"任务F共享会话not_found缺少串行确认: {row.get('queryIccid') or 'unknown'}")
                break
    if int(lookup.get("inputCount") or 0) != int(pending.get("pendingCount") or 0):
        errors.append("任务F Grab&Go输入数与待查数不一致")
    if lookup.get("batchMode") is not True or imported.get("batchMode") is not True:
        errors.append("任务F全量查询未启用分批数据库断点")
    if lookup.get("databaseCheckpointAfterEveryBatch") is not True or imported.get("databaseCheckpointAfterEveryBatch") is not True:
        errors.append("任务F未证明每批完成后立即写入数据库")
    if int(lookup.get("batchCount") or 0) != int(lookup.get("completedBatchCount") or 0):
        errors.append("任务F查询批次未全部完成")
    if imported.get("status") != "passed" or int(imported.get("transientErrorCount") or 0) != 0:
        errors.append("任务F数据库导入存在瞬时失败")
    if imported.get("cutoffRule") != "activation_date_plus_29_days_inclusive_30_day_period":
        errors.append("任务F截止日计算口径无效")
    if (
        int(imported.get("recordSchemaVersion") or 0) != 2
        or imported.get("recordedWebFields") != expected_web_fields
        or imported.get("fullSourceDataIncluded") is not True
    ):
        errors.append("任务F数据库导入回执未证明网页五字段与完整源数据已保存")
    if int(imported.get("foundImportedCount") or 0) != int(config.get("taskFFoundImportedCount") or 0):
        errors.append("任务F激活日期导入计数不一致")

    snapshot = paths["taskFDatabaseSnapshot"]
    try:
        connection = sqlite3.connect(f"file:{snapshot.resolve()}?mode=ro", uri=True)
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        terminal_count = int(connection.execute("SELECT COUNT(*) FROM kooda_activation").fetchone()[0])
        invalid_cutoffs = int(connection.execute(
            "SELECT COUNT(*) FROM kooda_activation WHERE lookup_status='found' AND (activation_date IS NULL OR cutoff_date IS NULL)"
        ).fetchone()[0])
        activation_columns = {
            str(row[1]) for row in connection.execute("PRAGMA table_info(kooda_activation)")
        }
        required_columns = {
            "purchase_date", "purchase_date_raw", "act_request_date", "activation_date_raw",
            "sim_number", "sim_number_raw", "sku", "source_data_json", "record_schema_version",
        }
        current_schema_count = int(connection.execute(
            "SELECT COUNT(*) FROM kooda_activation WHERE record_schema_version >= 2"
        ).fetchone()[0]) if required_columns.issubset(activation_columns) else 0
        connection.close()
        if integrity != "ok" or invalid_cutoffs or not required_columns.issubset(activation_columns):
            errors.append("任务F数据库快照完整性或found日期字段无效")
        if current_schema_count != int(config.get("taskFDatabaseCurrentSchemaCount") or 0):
            errors.append("任务F数据库完整记录结构计数不一致")
    except sqlite3.Error as exc:
        terminal_count = 0
        errors.append(f"任务F数据库快照无法读取: {exc}")
    details.update({
        "inventory_count": config.get("taskFInventoryCount"),
        "already_queried_count": config.get("taskFAlreadyQueriedCount"),
        "pending_count": config.get("taskFPendingCount"),
        "found_imported_count": config.get("taskFFoundImportedCount"),
        "effective_workers": lookup.get("effectiveWorkers"),
        "per_worker_delay_ms": lookup.get("perWorkerDelayMs"),
        "database_terminal_count": terminal_count,
        "record_schema_version": config.get("taskFRecordSchemaVersion"),
        "recorded_web_fields": config.get("taskFRecordedWebFields"),
        "database_current_schema_count": config.get("taskFDatabaseCurrentSchemaCount"),
    })
    return status_result("任务F Kooda激活日期数据库", "task_f", errors, warnings, details, "任务F登录、去重、查询和数据库回执通过")


def task_f_document_boundary_matches(config: dict[str, Any]) -> bool:
    """Audit F timing without treating its separate library as report dates."""
    if config.get("taskFBackground"):
        result_boundary = (config.get("taskFResultsApplyNextExport") is True
                           and config.get("taskFResultsWriteBackToDocument") is False)
    elif config.get("taskFDocumentDateSource") == "system":
        result_boundary = (
            config.get("taskFResultsApplyNextExport") is False
            and config.get("taskFResultsWriteBackToDocument") is False
        )
    else:
        # Historical receipts retain their original post-document contract.
        result_boundary = config.get("taskFResultsApplyNextExport") is True
    return (
        config.get("taskFPostDocumentSupplemental") is True
        and config.get("taskFDocumentGeneratedBeforeLookup") is True
        and result_boundary
    )


def check_date_resolution(config: dict[str, Any], resolver: PathResolver,
                          inputs: dict[str, Any]) -> CheckResult:
    """Bind independent per-field readback to this run's immutable inputs."""
    data = unwrap_card_receipt(read_json(inputs["card_receipt"]))
    process = data.get("process") or {}
    resolution = process.get("date_resolution") or {}
    if not resolution:
        if config.get("reportDateCachePath") or config.get("taskFBackground"):
            return CheckResult("日期来源与冻结缓存", "document", FAIL, "本轮缺少逐字段日期来源独立复核")
        return CheckResult("日期来源与冻结缓存", "document", SKIP, "历史回执未启用逐字段冻结缓存")
    errors: list[str] = []
    verification = (data.get("verification") or {}).get("date_resolution_review") or {}
    if verification.get("matches") is not True:
        errors.append("日期来源独立回读未通过")
    for key in ("field_counts", "cache"):
        if resolution.get(key) != verification.get(key):
            errors.append(f"日期来源{key}与独立回读不一致")
    google_disabled = google_disabled_scope_requested(config)
    expected_source = (CHATR_GOOGLE_RESOLUTION_SOURCE if chatr_google_scope_requested(config) else
                       "system_export_then_frozen_local" if google_disabled else "system_export_then_frozen_local_or_google")
    if resolution.get("source") != expected_source:
        errors.append("日期来源不是系统优先、Kooda冻结本地及限定Chatr Google逐字段补缺"
                      if chatr_google_scope_requested(config) else
                      "日期来源不是系统优先、冻结本地逐字段补缺" if google_disabled else
                      "日期来源不是系统优先、冻结本地/Google逐字段补缺")
    if chatr_google_scope_requested(config):
        proof = process.get("google_chatr_date_fallback") or {}
        independent = (data.get("verification") or {}).get("google_chatr_date_fallback_review") or {}
        if (proof.get("enabled") is not True or proof.get("run_id") != config.get("runId")
                or proof.get("scope") != CHATR_GOOGLE_DATE_SCOPE
                or proof.get("source") != CHATR_GOOGLE_DATE_SOURCE
                or resolution.get("chatr_google_dates") != proof
                or verification.get("chatr_google_dates") != proof
                or independent.get("matches") is not True
                or any(independent.get(key) != value for key, value in proof.items())):
            errors.append("Chatr日期补缺未独立核验本轮运行、限定范围与来源")
    if (resolution.get("new_query_results_apply_next_export") is not True
            or resolution.get("redundant_system_date_columns_removed") is not True):
        errors.append("日期来源未证明本轮冻结与冗余列删除")
    source_bindings = [("source_export", "source")]
    if not google_disabled:
        source_bindings.append(("association", "association"))
    elif resolution.get("association"):
        errors.append("Google关闭时日期来源仍引用关联文件；未访问该文件")
    for metadata_key, input_key in source_bindings:
        source = resolution.get(metadata_key) or {}
        expected = inputs.get(input_key)
        actual = resolver.resolve(source.get("path"), required=False)
        # Generator inputs may be local working copies of the archived fresh
        # inputs; bind byte identity, not a basename or path spelling.
        expected_digest = sha256_file(Path(expected)) if expected and Path(expected).is_file() else None
        if (not expected_digest or normalize_hash(source.get("sha256")) != expected_digest
                or (actual is not None and sha256_file(actual) != expected_digest)):
            errors.append(f"日期来源{metadata_key}未绑定本轮新鲜输入")
    cache = resolution.get("cache") or {}
    cache_path = resolver.resolve(cache.get("path"), required=False)
    configured = resolver.resolve(config.get("reportDateCachePath"), required=False)
    cache_hash = normalize_hash(cache.get("sha256"))
    snapshot: dict[str, Any] = {}
    renewal_count = 0
    if (not cache_path or not configured or cache_path != configured or not cache_hash
            or cache_hash != normalize_hash(config.get("reportDateCacheSha256"))
            or sha256_file(cache_path) != cache_hash):
        errors.append("日期冻结缓存路径/SHA256与本轮配置不一致")
    else:
        snapshot = read_json(cache_path)
        if (cache.get("run_id") != config.get("runId")
                or (snapshot.get("runId") or snapshot.get("run_id")) != config.get("runId")):
            errors.append("日期冻结缓存未绑定本轮运行编号")
        if not parse_iso(cache.get("created_at")):
            errors.append("日期冻结缓存缺少有效冻结时间")
        records = snapshot.get("records")
        if (snapshot.get("schemaVersion") != 1 or snapshot.get("status") != PASS
                or not isinstance(records, list) or snapshot.get("recordCount") != len(records)
                or snapshot.get("createdAt") != cache.get("created_at")):
            errors.append("日期冻结缓存结构、数量或冻结时间不一致")
        else:
            seen: set[str] = set()
            for record in records:
                query = task_f_report_iccid(record.get("queryIccid")) if isinstance(record, dict) else None
                try:
                    activation = task_f_report_date(record.get("activationDate")) if isinstance(record, dict) else None
                    cutoff = task_f_report_date(record.get("cutoffDate")) if isinstance(record, dict) else None
                except ValueError:
                    activation = cutoff = None
                pair_valid = bool(activation and cutoff)
                if pair_valid and "renewal" in record:
                    try:
                        validate_renewal(record["renewal"], query, activation.isoformat(), cutoff.isoformat())
                    except (ValueError, TypeError):
                        pair_valid = False
                    else:
                        renewal_count += 1
                elif pair_valid:
                    pair_valid = cutoff == activation + timedelta(days=29)
                if (not query or query != record.get("queryIccid") or query in seen or not pair_valid):
                    errors.append("日期冻结缓存含重复/非规范ICCID或无效日期配对")
                    break
                seen.add(query)
    details = {"source": resolution.get("source"), "field_counts": resolution.get("field_counts"),
               "cache": cache, "frozen_cache_verified": not errors,
               "validated_renewal_record_count": renewal_count,
               "background_results_apply_next_export": True}
    return status_result("日期来源与冻结缓存", "document", errors, [], details,
                         "服务器优先、逐字段补缺与本轮冻结缓存绑定通过")


def task_f_report_iccid(value: Any) -> str | None:
    """Independent text-only ICCID normalization; never round numeric cells."""
    if not isinstance(value, str):
        return None
    raw = value.strip()
    if raw[-1:].lower() == "f":
        raw = raw[:-1]
    return raw if re.fullmatch(r"[0-9]{18,22}", raw) else None


def task_f_report_slot(value: Any) -> str:
    raw = unicodedata.normalize("NFKC", str(value or "").strip()).replace(" ", "").replace("-", "_")
    parts = []
    for part in raw.split("_"):
        matched = re.fullmatch(r"(\d+)(?:\.\d+)?", part)
        parts.append(str(int(matched.group(1))) if matched else part.lower())
    return "_".join(parts)


def task_f_report_date(value: Any) -> date | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            return date.fromisoformat(value)
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:\d{2})?", value):
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
    raise ValueError("非空日期不是有效ISO日期")


def review_task_f_report_backfill(
    config: dict[str, Any], pending: dict[str, Any], inventory: dict[str, Any],
    lookup: dict[str, Any], resolver: PathResolver,
) -> tuple[list[str], dict[str, Any]]:
    """Reread the hash-bound report, independently proving every F identity."""
    errors: list[str] = []
    details: dict[str, Any] = {"report_backfill_mode": True}
    if config.get("taskFReportBackfillMode") is not True or pending.get("reportBackfillMode") is not True:
        return ["任务F报表补查模式与配置不一致"], details
    frozen_fallback = bool(config.get("reportDateCachePath"))
    if (config.get("taskFDocumentDateSource") not in {"system", "system_export_then_frozen_local_or_google", "system_export_then_frozen_local", "system_first_frozen_cache", "system_then_frozen_local_cache"}
        or config.get("taskFResultsApplyNextExport") is not frozen_fallback
        or config.get("taskFResultsWriteBackToDocument") is not False
        or config.get("taskFCloudUploadAttempted") is not False):
        errors.append("任务F报表补查未证明本轮无主表回填及无云端上传，或下轮冻结缓存边界不一致")
    if pending.get("documentMutationAllowed") is not False or pending.get("cloudUploadMode") != "manual-user-upload-only":
        errors.append("任务F补查清单未证明文档只读和仅供手工上传")
    inventory_path = resolver.resolve(config.get("taskFInventoryReceipt"), required=False)
    pending_inventory = resolver.resolve(pending.get("sourceCards"), required=False)
    if (not inventory_path or inventory_path != pending_inventory
        or normalize_hash(pending.get("sourceCardsSha256")) != sha256_file(inventory_path)):
        errors.append("任务F补查清单未绑定完整系统库存文件SHA256")
    expected_run = str(config.get("runId") or "")
    if (not expected_run or pending.get("documentRunId") != expected_run
        or config.get("taskFDocumentRunId") != expected_run):
        errors.append("任务F补查文档未绑定本轮运行编号")
    document = resolver.resolve(pending.get("documentPath"), required=False)
    configured_document = resolver.resolve(config.get("taskFDocumentPath"), required=False)
    digest = normalize_hash(pending.get("documentSha256"))
    if (not document or not digest or document != configured_document
        or digest != normalize_hash(config.get("taskFDocumentSha256"))
        or sha256_file(document) != digest):
        return errors + ["任务F补查文档路径或SHA256绑定无效"], details
    receipt = resolver.resolve(config.get("taskFDocumentReceipt"), required=False)
    if not receipt or check_card_receipt(receipt, not google_disabled_scope_requested(config)).status == FAIL:
        return errors + ["任务F补查文档缺少已通过的独立回读回执"], details
    proof = unwrap_card_receipt(read_json(receipt))
    system_review = (proof.get("verification") or {}).get("system_kooda_source_review") or {}
    proof_document = resolver.resolve(proof.get("output_path"), required=False)
    if system_review.get("run_id") != expected_run or not proof_document or sha256_file(proof_document) != digest:
        return errors + ["任务F文档独立回读回执与本轮文档字节或运行编号不一致"], details
    from report_kooda_source import is_report_source, validate_projection, missing_activation_rows, ELIGIBILITY
    report_only = is_report_source(inventory)
    provenance_mode = pending.get("dateEligibilityMode") in ("missing_or_verified_creation_fallback", ELIGIBILITY)
    fallback_rows = {}
    if provenance_mode:
        if (resolver.resolve(pending.get("documentProofPath"), required=False) != receipt
                or normalize_hash(pending.get("documentProofSha256")) != sha256_file(receipt)):
            return errors + ["任务F兜底来源回执路径或SHA256不一致"], details
        fallback_errors, fallback_rows = ([], {}) if report_only else task_f_creation_fallback_rows(proof, expected_run,
                                                                       sha256_file(inventory_path))
        if fallback_errors:
            return errors + fallback_errors, details

    def identity_key(slot: Any, phone: Any) -> tuple[str, str] | None:
        normalized_slot = task_f_report_slot(slot)
        digits = re.sub(r"\D", "", str(phone or ""))
        return (normalized_slot, digits[-10:]) if normalized_slot and len(digits) >= 10 else None

    raw_index: dict[tuple[str, str], dict[str, Any]] = {}
    backend_ids: set[str] = set()
    system_rows = inventory.get("systemRows")
    if report_only:
        try:
            validate_projection(inventory, expected_run)
        except RuntimeError as exc:
            return errors + [str(exc)], details
        if pending.get('dateEligibilityMode') != ELIGIBILITY or pending.get('sourceScope') != inventory['sourceScope']:
            return errors + ['任务F补查未绑定10日表且双端缺激活日期规则'], details
    report_sources = {row['sourceRow']: row for row in missing_activation_rows(inventory)} if report_only else {}
    if (not isinstance(system_rows, list) or inventory.get("runId") != expected_run
        or type(inventory.get("totalRows")) is not int or len(system_rows) != inventory.get("totalRows")
        or (not report_only and (inventory.get("paginationTotalStable") is not True
        or inventory.get("paginationBackendIdsUnique") is not True))):
        return errors + ["任务F报表补查库存缺少本轮完整原始系统快照"], details
    for raw in system_rows:
        if not isinstance(raw, dict):
            errors.append("任务F原始系统库存包含非对象行")
            continue
        backend_id = str(raw.get("sourceRowIdentity") if report_only else raw.get("id") or "")
        if (not backend_id or backend_id in backend_ids
            or not re.search(r"kooda|koodo|telus1", str(raw.get("opr") or ""), re.I)):
            errors.append("任务F原始系统库存后台ID缺失/重复或包含非Kooda行")
        backend_ids.add(backend_id)
        key = identity_key(raw.get("slotIdentif"), raw.get("cardNo"))
        if key:
            if key in raw_index and not report_only:
                errors.append("任务F原始系统库存复合身份重复")
            raw_index[key] = raw

    eligible: dict[int, dict[str, Any]] = {}
    excluded: dict[int, dict[str, Any]] = {}
    archived_index = index_phones([row for row in system_rows if isinstance(row, dict)])
    document_keys: set[tuple[str, str] | None] = set()
    counts = {"documentKoodaRows": 0, "documentActivationPresentRows": 0, "documentMissingActivationRows": 0}
    if provenance_mode:
        counts.update(documentCreationFallbackRows=0, documentLookupNeededRows=0)
    missing_activation_queries: set[str] = set()
    has_cutoff = False
    workbook = load_workbook(document, read_only=True, data_only=False)
    try:
        if "sheet1" not in workbook.sheetnames:
            return errors + ["任务F补查文档缺少sheet1"], details
        rows = workbook["sheet1"].iter_rows(values_only=True)
        headers = list(next(rows, ()))
        required = ("电话", "卡槽", "opr", "激活日期", "ICCID")
        if any(headers.count(name) != 1 for name in required):
            return errors + ["任务F补查文档必要列缺失或重复"], details
        positions = {name: headers.index(name) for name in required}
        for name in ('状态', '创建时间'):
            if name in headers:
                positions[name] = headers.index(name)
        has_cutoff = headers.count("截止日") == 1
        if headers.count("截止日") > 1:
            return errors + ["任务F补查文档截止日列重复"], details
        if has_cutoff:
            positions["截止日"] = headers.index("截止日")
            counts.update({"documentMissingDateRows": 0, "documentMissingCutoffRows": 0})
        for row_number, values in enumerate(rows, 2):
            row = {name: values[index] if index < len(values) else None for name, index in positions.items()}
            if not re.search(r"kooda|koodo|telus1", str(row["opr"] or ""), re.I):
                continue
            counts["documentKoodaRows"] += 1
            try:
                activation = task_f_report_date(row["激活日期"])
                cutoff = task_f_report_date(row["截止日"]) if has_cutoff else None
            except ValueError:
                errors.append(f"任务F文档第{row_number}行非空无效激活日期或截止日不得当作空白")
                continue
            if activation:
                counts["documentActivationPresentRows"] += 1
            else:
                counts["documentMissingActivationRows"] += 1
            missing_fields = ([] if activation else ["激活日期"]) + (["截止日"] if has_cutoff and not cutoff else [])
            if has_cutoff and not cutoff:
                counts["documentMissingCutoffRows"] += 1
            if has_cutoff and missing_fields:
                counts["documentMissingDateRows"] += 1
            fallback = fallback_rows.get(row_number)
            if fallback:
                if (re.sub(r"\D", "", str(row["电话"] or ""))[-10:] != re.sub(r"\D", "", str(fallback["phone"]))[-10:]
                        or activation != task_f_report_date(fallback["activation"])
                        or not has_cutoff or cutoff != task_f_report_date(fallback["cutoff"])):
                    errors.append(f"任务F第{row_number}行兜底来源与文档不一致")
                    continue
                counts["documentCreationFallbackRows"] += 1
                missing_fields = ["激活日期", "截止日"]
            if report_only:
                if row_number not in report_sources:
                    continue
                missing_fields = ['激活日期']
            if not missing_fields:
                continue
            if provenance_mode:
                counts["documentLookupNeededRows"] += 1
            query = task_f_report_iccid(row["ICCID"])
            expected = {"documentRow": row_number, "phone": str(row["电话"] or ""),
                        "slot": str(row["卡槽"] or ""), "sourceIccid": row["ICCID"], "queryIccid": query}
            if report_only:
                expected.update(sourceRowIdentity=report_sources[row_number]['sourceRowIdentity'],
                    lookupReason=ELIGIBILITY, dateSource='raw_10day_export_activation_missing')
            if has_cutoff:
                expected["missingFields"] = missing_fields
            if fallback:
                expected.update(dateSource="creation_date_fallback",
                                lookupReason="creation_estimate_is_not_authoritative_activation")
            if not query:
                excluded[row_number] = expected
                continue
            key = identity_key(row["卡槽"], row["电话"])
            if key in document_keys and not report_only:
                errors.append(f"任务F文档第{row_number}行复合身份重复")
            document_keys.add(key)
            raw = report_sources[row_number] if report_only else raw_index.get(key)
            if raw is None:
                archived = match_archived(archived_index, phone=row['电话'], slot=row['卡槽'],
                    status=row.get('状态'), created=row.get('创建时间'), iccid=row['ICCID'])
                if archived is not None:
                    excluded[row_number] = {**expected, 'backendId': archived['id'],
                                           'reason': 'archived_off_shelf_no_current_slot'}
                    continue
            if (not raw or not re.search(r"kooda|koodo|telus1", str(raw.get("opr") or ""), re.I)
                or task_f_report_iccid(raw.get("cardStore")) != query
                or (not report_only and raw.get("id") in (None, ""))):
                errors.append(f"任务F文档第{row_number}行与系统库存精确身份不一致")
                continue
            if report_only:
                expected.update(sourceRowIdentity=raw['sourceRowIdentity'], lookupReason=ELIGIBILITY,
                                dateSource='raw_10day_export_activation_missing')
            else:
                expected["backendId"] = raw["id"]
            expected["sourceData"] = raw
            eligible[row_number] = expected
            if report_only or not activation:
                missing_activation_queries.add(query)
        ko_context = pending.get('koRefresh')
        if ko_context is not None:
            from opensms_task_f_ko_refresh import CONTRACT, selected_rows
            if not provenance_mode or ko_context.get('contract') != CONTRACT:
                errors.append('ko补查缺少独立回执来源契约')
            else:
                try:
                    extra = selected_rows(workbook, proof, inventory, sha256_file(inventory_path),
                        expected_run, Path(ko_context['cachePath']), ko_context['cacheSha256'])
                    for row in extra:
                        eligible[row['documentRow']] = row
                    counts['documentKoRefreshRows'] = len(extra)
                    counts['totalLookupNeededRows'] = counts['documentLookupNeededRows'] + len(extra)
                except (ValueError, RuntimeError, KeyError, OSError) as exc:
                    errors.append(f'ko补查独立来源核验失败:{exc}')
    finally:
        workbook.close()
    if sha256_file(document) != digest:
        errors.append("任务F补查文档在独立扫描期间变化")

    def same_row(actual: Any, expected: dict[str, Any], with_source: bool = False) -> bool:
        if not isinstance(actual, dict):
            return False
        matches = (actual.get("documentRow") == expected["documentRow"]
                   and identity_key(actual.get("slot"), actual.get("phone")) == identity_key(expected["slot"], expected["phone"])
                   and actual.get("sourceIccid") == expected["sourceIccid"]
                   and (actual.get("queryIccid") or None) == expected["queryIccid"])
        if expected.get("backendId") is not None and str(actual.get("backendId")) != str(expected["backendId"]):
            matches = False
        if report_only and actual.get('sourceRowIdentity') != expected.get('sourceRowIdentity'):
            matches = False
        if with_source and actual.get("sourceData") != expected.get("sourceData"):
            matches = False
        if has_cutoff and actual.get("missingFields") != expected.get("missingFields"):
            matches = False
        if provenance_mode and any(actual.get(key) != expected.get(key) for key in ("dateSource", "lookupReason")):
            matches = False
        if expected.get('documentSheet') and any(actual.get(key) != expected.get(key)
                for key in ('documentSheet', 'sheetRow', 'systemMatches')):
            matches = False
        return matches

    declared_eligible = pending.get("reportEligibleRecords")
    declared_excluded = pending.get("documentExcludedRows")
    for name, actual_rows, expected_rows in (("合格", declared_eligible, eligible), ("排除", declared_excluded, excluded)):
        if (not isinstance(actual_rows, list) or len(actual_rows) != len(expected_rows)
            or len({row.get("documentRow") for row in actual_rows if isinstance(row, dict)}) != len(expected_rows)
            or any(not same_row(row, expected_rows.get(row.get("documentRow"), {}))
                   for row in actual_rows if isinstance(row, dict) and row.get("documentRow") in expected_rows)
            or any(not isinstance(row, dict) or row.get("documentRow") not in expected_rows for row in actual_rows)):
            errors.append(f"任务F报表{name}行与独立扫描身份不一致")
    if isinstance(declared_excluded, list) and any(not isinstance(row, dict) or not str(row.get("reason") or "") for row in declared_excluded):
        errors.append("任务F缺失/非法ICCID排除行未记录原因")
    groups: dict[str, set[int]] = {}
    for row_number, row in eligible.items():
        groups.setdefault(row["queryIccid"], set()).add(row_number)
    query_sets: dict[str, set[str]] = {}
    for label, field in (("待查", "rows"), ("复用", "alreadyQueried")):
        records = pending.get(field)
        keys: set[str] = set()
        if not isinstance(records, list):
            errors.append(f"任务F{label}清单不是数组")
            records = []
        for record in records:
            if not isinstance(record, dict):
                errors.append(f"任务F{label}清单包含非对象")
                continue
            query = record.get("queryIccid")
            if query in keys or query not in groups:
                errors.append(f"任务F{label}ICCID重复或不属于主表空激活日期/截止日行")
            keys.add(query)
            occurrences = record.get("sourceOccurrences") or []
            expected_rows = groups.get(query, set())
            if (not isinstance(occurrences, list) or len(occurrences) != len(expected_rows)
                or {row.get("documentRow") for row in occurrences if isinstance(row, dict)} != expected_rows
                or any(not isinstance(row, dict) or row.get("documentRow") not in eligible
                       or not same_row(row, eligible.get(row.get("documentRow"), {}), with_source=True) for row in occurrences)):
                errors.append(f"任务F{label}ICCID源出现位置或完整系统对象不一致")
            if label == "复用":
                try:
                    activation, cutoff = task_f_report_date(record.get("activationDate")), task_f_report_date(record.get("cutoffDate"))
                    valid = bool(activation) if report_only else record.get("lookupStatus") == "found" and activation and cutoff == activation + timedelta(days=29)
                except ValueError:
                    valid = False
                if not valid:
                    errors.append("任务F复用清单包含无有效激活日期/截止日的记录")
        query_sets[label] = keys
    if query_sets["待查"] & query_sets["复用"] or query_sets["待查"] | query_sets["复用"] != set(groups):
        errors.append("任务F有效日期复用与待查不是主表合格ICCID的互斥完整分区")
    if pending.get("reuseRecords") != pending.get("alreadyQueried"):
        errors.append("任务F待人工上传复用别名与已有日期清单不一致")
    counts.update({"documentExcludedCount": len(excluded), "documentMissingActivationIccidCount": len(missing_activation_queries),
                   "validUniqueIccidCount": len(groups), "documentDuplicateIccidRows": len(eligible) - len(groups),
                   "alreadyQueriedCount": len(query_sets["复用"]), "readyForUploadCount": len(query_sets["复用"]),
                   "pendingCount": len(query_sets["待查"])})
    if has_cutoff:
        counts["documentMissingDateIccidCount"] = len(groups)
    refresh_existing = sum(isinstance(row, dict) and row.get("previousLookupStatus") is not None for row in pending.get("rows") or [])
    if pending.get("refreshExistingMissingCount") != refresh_existing:
        errors.append("任务F历史空日期/未找到重查计数不一致")
    for name, expected in counts.items():
        if type(pending.get(name)) is not int or pending[name] != expected:
            errors.append(f"任务F报表补查{name}计数不一致")
    for pending_name, config_name in (("readyForUploadCount", "taskFReadyForUploadCount"),
                                       ("documentExcludedCount", "taskFDocumentExcludedCount"),
                                       ("documentMissingActivationIccidCount", "taskFDocumentMissingActivationIccidCount")):
        if counts[pending_name] != config.get(config_name):
            errors.append(f"任务F报表补查{config_name}配置计数不一致")
    results = lookup.get("results") or []
    result_keys = [row.get("queryIccid") for row in results if isinstance(row, dict)]
    if len(result_keys) != len(results) or len(result_keys) != len(set(result_keys)) or set(result_keys) != query_sets["待查"]:
        errors.append("任务F查询结果ICCID与本轮报表待查清单不完全一致")
    details.update({"document_path": str(document), "document_sha256": digest, "document_run_id": expected_run,
                    "document_counts": counts, "cloud_upload_attempted": False, "matches": not errors})
    return errors, details


def task_f_creation_fallback_rows(proof: dict, run_id: str, inventory_hash: str) -> tuple[list[str], dict]:
    """Independent provenance check; date equality alone never authorizes lookup."""
    review = (proof.get("verification") or {}).get("date_resolution_review") or {}
    rows = review.get("creation_date_fallback_rows", [])
    if not isinstance(rows, list):
        return ["任务F兜底来源清单无效"], {}
    if not rows:
        return [], {}
    system = (proof.get("verification") or {}).get("system_kooda_source_review") or {}
    if (proof.get("result") != "成功" or not proof.get("checks")
            or any(value is not True for value in proof["checks"].values())
            or proof.get("failed_checks") != [] or review.get("matches") is not True
            or system.get("matches") is not True or system.get("run_id") != run_id
            or system.get("snapshot_sha256") != inventory_hash):
        return ["任务F兜底来源未独立核验或不属于本轮库存"], {}
    result = {}
    try:
        for row in rows:
            if (not isinstance(row, dict) or type(row.get("row")) is not int or row["row"] < 2
                    or row["row"] in result
                    or len(re.sub(r"\D", "", str(row.get("phone") or ""))[-10:]) != 10
                    or not task_f_report_date(row.get("activation"))
                    or task_f_report_date(row.get("cutoff")) != task_f_report_date(row["activation"]) + timedelta(days=29)):
                return ["任务F兜底来源行缺失、重复或日期无效"], {}
            result[row["row"]] = row
    except (ValueError, TypeError):
        return ["任务F兜底来源日期无效"], {}
    return [], result


def task_a_only_scope_requested(config: dict[str, Any]) -> bool:
    return (config.get("taskAOnly") is True or config.get("taskSelection") == "A"
            or config.get("taskFSkippedReason") == "task_a_only")


def check_task_a_only_scope(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    """Verify the skip contract without resolving or reading any Task F artifact."""
    errors: list[str] = []
    if not (config.get("taskAOnly") is True and config.get("taskSelection") == "A"
            and config.get("taskFSkippedReason") == "task_a_only"
            and config.get("taskFStageStatus") == SKIP):
        errors.append("仅任务A模式、任务选择与任务F跳过原因/状态不一致")
    for key in ("taskBExecuted", "taskCExecuted", "taskDExecuted", "taskEExecuted",
                "taskFExecuted", "taskDCurrentUserAuthorized", "taskECurrentUserAuthorized",
                "taskFCurrentUserAuthorized", "taskFAutomaticCatalogSync"):
        if config.get(key) is not False:
            errors.append(f"仅任务A未证明{key}=false")
    for key in ("taskBCurrentUserAuthorized", "taskCCurrentUserAuthorized"):
        if config.get(key) not in (None, False):
            errors.append(f"仅任务A仍存在{key}授权")
    # These legacy taskF-prefixed fields describe the ordinary A document, not
    # F inventory, local reuse, lookup, or database execution.
    allowed_document_fields = {
        "taskFDocumentReceipt", "taskFDocumentDateSource", "taskFDocumentMissingActivationRedRows",
        "taskFDocumentFullRedRows", "taskFDocumentRedReviewPassed", "taskFMissingActivationFillRgb",
    }
    allowed_skip_fields = {"taskFStatus", "taskFStageStatus", "taskFSkippedReason"}
    populated_f_fields = [key for key, value in config.items()
                          if key.startswith("taskF") and key not in allowed_document_fields | allowed_skip_fields
                          and value not in (None, False, 0, "", [], {})]
    if populated_f_fields:
        errors.append("仅任务A存在F业务路径、数量或执行证据：" + "、".join(sorted(populated_f_fields)))
    if any(str(group.get("task") or "").upper() in set("BCDEF")
           or re.match(r"任务[BCDEF]", str(group.get("title") or ""))
           for group in config.get("groups") or []):
        errors.append("仅任务A仍包含任务B/C/D/E/F候选或执行分组")
    if any(config.get(key) for key in ("taskBManifest", "taskCManifest", "taskDManifest", "taskEManifest",
                                        "taskDRetentionReceipt", "taskDSwapExpiryReceipt", "taskDRetentionExpiryReceipt",
                                        "taskEReceipt", "taskELedger", "taskELedgerSnapshot", "taskERetentionReceipt")):
        errors.append("仅任务A仍配置任务B/C/D/E候选、执行回执或台账")
    forbidden_receipts = [str(item.get("label") or item.get("path"))
                          for item in config.get("extraReceipts") or []
                          if task_f_artifact_reference(item.get("label"))
                          or task_f_artifact_reference(item.get("path"))]
    if forbidden_receipts:
        errors.append("仅任务A声明了F业务产物或执行回执；未读取：" + "、".join(forbidden_receipts))
    timings = config.get("commandTimings")
    forbidden_commands: list[str] = []
    if not isinstance(timings, list) or not timings:
        errors.append("仅任务A缺少本轮实际子程序命令记录，无法证明F未运行")
    else:
        for item in timings:
            if not isinstance(item, dict) or not item.get("command"):
                errors.append("仅任务A子程序命令记录格式无效")
                continue
            command = str(item["command"])
            argv = item.get("argv") or []
            if (re.search(r"task[_-]?[bcdef]|任务[BCDEF]|grabandgo", command, re.IGNORECASE)
                    or any(re.search(r"(?:force|enable|execute)[_-]task[_-][bcdef]", str(arg), re.IGNORECASE)
                           for arg in argv)):
                forbidden_commands.append(command)
    if forbidden_commands:
        errors.append("仅任务A存在已执行的任务B/C/D/E/F命令：" + "、".join(forbidden_commands))
    safety_value = config.get("safetyReceipt") or extra_receipt(config, ["本轮安全门禁回执"])
    if not safety_value or task_f_artifact_reference(safety_value):
        errors.append("仅任务A缺少可读取的本轮安全门禁回执")
    else:
        # Resolve only the known run-local receipt paths; never recursively
        # search delivery directories that might contain historical F data.
        raw_safety_path = Path(safety_value).expanduser()
        safety_candidates = []
        if resolver.delivery_dir:
            safety_candidates.extend([resolver.delivery_dir / "receipts" / raw_safety_path.name,
                                      resolver.delivery_dir / raw_safety_path.name])
        safety_candidates.append(raw_safety_path if raw_safety_path.is_absolute()
                                 else resolver.config_dir / raw_safety_path)
        safety_path = next((path.resolve() for path in safety_candidates if path.is_file()), None)
        if not safety_path:
            errors.append("仅任务A安全门禁回执不存在")
        else:
            safety = read_json(safety_path)
            authorization = safety.get("taskAuthorization") or {}
            if not config.get("runId") or safety.get("runId") != config.get("runId"):
                errors.append("仅任务A安全门禁回执未绑定本轮运行编号")
            expected = {"taskAOnly": True, "taskSelection": "A", "taskFSkippedReason": "task_a_only", "taskFStageStatus": SKIP,
                        "taskBExecuted": False, "taskCExecuted": False, "taskDExecuted": False,
                        "taskEExecuted": False, "taskFExecuted": False, "taskFCurrentUserAuthorized": False,
                        "taskDCurrentUserAuthorized": False, "taskECurrentUserAuthorized": False,
                        "taskFAutomaticCatalogSync": False, "passed": True}
            for key, value in expected.items():
                if authorization.get(key) != value:
                    errors.append(f"仅任务A安全门禁回执{key}不一致")
    details = {"task_a_only": True, "task_selection": "A", "task_status": SKIP,
               "reason": "task_a_only", "executed": False, "authorized": False,
               "automatic_catalog_sync": False, "f_artifacts_read": False,
               "f_inventory_and_gap_checked": False, "forbidden_commands": forbidden_commands,
               "populated_f_fields": populated_f_fields, "skip_contract_verified": not errors}
    if errors:
        return status_result("仅任务A范围与任务F跳过门禁", "task_f", errors, [], details, "")
    return CheckResult("任务F完整跳过", "task_f", SKIP,
                       "仅任务A模式（task_a_only）；已核验无F授权、命令及实执行回执，F库与缺口未检查", details=details)


def check_task_f_no_missing_dates(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    errors: list[str] = []
    receipt = resolver.resolve(config.get("taskFDocumentReceipt"), required=False)
    if not receipt or check_card_receipt(receipt, not google_disabled_scope_requested(config)).status == FAIL:
        errors.append("任务F无缺口跳过缺少已核验文档")
        review = {}
    else:
        data = unwrap_card_receipt(read_json(receipt))
        review = ((data.get("verification") or {}).get("system_kooda_source_review") or {})
        if review.get("run_id") != config.get("runId"):
            errors.append("任务F无缺口独立回读不属于本轮")
    report_only = review.get('source_kind') == 'verified_10day_export'
    if report_only:
        try:
            from report_kooda_source import missing_activation_rows, verify_projection_workbook, iccid_key
            source = resolver.resolve(review['snapshot_path'])
            if sha256_file(source) != review['snapshot_sha256']:
                raise ValueError('10日表Kooda来源SHA256不一致')
            inventory = read_json(source)
            raw = resolver.resolve(inventory['sourceWorkbookPath'])
            verify_projection_workbook(inventory, raw, config['runId'])
            cache = resolver.resolve(config['reportDateCachePath'])
            if sha256_file(cache) != config['reportDateCacheSha256']:
                raise ValueError('F冻结库SHA256不一致')
            known = {row['queryIccid'] for row in read_json(cache)['records'] if row.get('activationDate')}
            if any(iccid_key(row['cardStore']) and iccid_key(row['cardStore']) not in known
                   for row in missing_activation_rows(inventory)):
                raise ValueError('10日表和F库仍均缺激活日期，不得跳过F')
        except (ValueError, RuntimeError, KeyError, OSError) as exc:
            errors.append(str(exc))
    if not report_only and (type(review.get("missing_either_date_rows")) is not int or review["missing_either_date_rows"] != 0):
        errors.append("任务F无缺口跳过与独立回读日期缺口计数不一致")
    if not report_only and config.get("taskFSkippedReason") == "no_missing_authoritative_dates" and receipt:
        proof = unwrap_card_receipt(read_json(receipt))
        ko_review = (proof.get('verification') or {}).get('ko_unactivated_sheet_review') or {}
        if ko_review.get('matches') is True and ko_review.get('row_count', 0) > 0:
            errors.append('ko系统有-未激活仍有记录，不得以主表无缺口跳过F')
        fallback_errors, fallback_rows = task_f_creation_fallback_rows(proof, str(config.get("runId") or ""),
                                                                      review.get("snapshot_sha256"))
        errors.extend(fallback_errors)
        if fallback_rows:
            document = resolver.resolve(proof.get("output_path"), required=False)
            if not document:
                errors.append("任务F无权威日期缺口跳过缺少文档")
            else:
                book = load_workbook(document, read_only=True, data_only=False)
                try:
                    rows = book["sheet1"].iter_rows(values_only=True)
                    headers = list(next(rows, ()))
                    pos = headers.index("opr")
                    if any(number in fallback_rows and re.search(r"kooda|koodo|telus1", str(values[pos] or ""), re.I)
                           for number, values in enumerate(rows, 2)):
                        errors.append("任务F仍有Kooda创建日期兜底行却跳过补查")
                finally:
                    book.close()
    for key in ("taskFExecuted", "taskFOnlineLookupTriggered", "taskFCloudUploadAttempted", "taskFBackground"):
        if config.get(key) is not False:
            errors.append(f"任务F无缺口仍有{key}执行声明")
    if (config.get("taskFStageStatus") != SKIP
            or config.get("taskFDocumentMissingDateRows") != (review.get('missing_either_date_rows') if report_only else 0)
            or config.get("taskFIncrementalWorkbook")
            or config.get("taskFNewActivationImportedCount") not in (None, 0)):
        errors.append("任务F无缺口跳过仍声明后台入库或新增文件")
    details = {"reason": config.get("taskFSkippedReason"), "online_executed": False,
               "document_missing_date_rows": review.get("missing_either_date_rows"),
               "incremental_file_generated": False}
    if errors:
        return status_result("任务F无日期缺口", "task_f", errors, [], details, "")
    return CheckResult("任务F无日期缺口", "task_f", SKIP,
                       ("已独立核验本轮10日表与F库均缺激活日期的有效ICCID为0，无需F联网"
                        if report_only else "已独立核验补缺后日期完整，无需F联网，不生成或提示新增文件"), details=details)


def check_task_f_background(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    """Validate a detached job without treating dispatch as F completion.

    Deliberately do not load the changing child state or library. The main
    document and manifest must remain auditable after the worker advances.
    """
    errors: list[str] = []
    if config.get("taskFBackgroundLaunchReceipt"):
        return check_task_f_background_launch_failure(config, resolver)
    job_path = resolver.resolve(config.get("taskFBackgroundJob"), required=False)
    digest = normalize_hash(config.get("taskFBackgroundJobSha256"))
    if not job_path or not digest or sha256_file(job_path) != digest:
        return CheckResult("任务F独立后台作业", "task_f", FAIL, "任务F不可变启动回执缺失或SHA256不符")
    job = read_json(job_path)
    run_id = config.get("runId")
    if not run_id or (job.get("parentRunId") or job.get("runId")) != run_id:
        errors.append("任务F后台作业未绑定本轮父运行编号")
    for name, expected in (("taskFMainReportIndependent", True),
                           ("taskFResultsApplyNextExport", True),
                           ("taskFResultsWriteBackToDocument", False),
                           ("taskFCloudUploadAttempted", False)):
        if config.get(name) is not expected:
            errors.append(f"任务F后台边界{name}不一致")
    for name, expected in (("mainReportIndependent", True), ("resultsApplyNextExport", True),
                           ("resultsWriteBackToDocument", False), ("cloudUploadAttempted", False)):
        if job.get(name) is not expected:
            errors.append(f"任务F启动回执{name}不一致")
    if not (config.get("taskFAutomaticCatalogSync") or config.get("taskFCurrentUserAuthorized")):
        errors.append("任务F后台作业没有当前执行依据")
    if job.get("requestedWorkers") != 8 or config.get("taskFRequestedWorkers") != 8:
        errors.append("任务F后台请求并发不是8个隔离工作器")
    document = resolver.resolve(config.get("taskFDocumentPath"), required=False)
    document_digest = normalize_hash(config.get("taskFDocumentSha256"))
    if (not document or not document_digest or sha256_file(document) != document_digest
            or resolver.resolve(job.get("documentPath"), required=False) != document
            or normalize_hash(job.get("documentSha256")) != document_digest
            or job.get("documentRunId") != run_id):
        errors.append("任务F后台作业未绑定本轮已生成文档的确切字节")
    receipt = resolver.resolve(config.get("taskFDocumentReceipt"), required=False)
    if (not receipt or check_card_receipt(receipt, not google_disabled_scope_requested(config)).status == FAIL
            or config.get("taskFDocumentGeneratedBeforeLookup") is not True):
        errors.append("任务F后台启动前文档独立回读未证明通过")
    if job.get('koRefreshEnabled') is True and receipt and document:
        from opensms_task_f_ko_refresh import selected_rows
        try:
            bound = job['inputs']
            for key in ('reportDateCache', 'inventory'):
                if sha256_file(Path(bound[key]['path'])) != bound[key]['sha256']:
                    raise ValueError('ko后台输入SHA256不一致')
            book = load_workbook(document, read_only=True, data_only=False)
            try:
                ko_rows = selected_rows(book, unwrap_card_receipt(read_json(receipt)),
                    read_json(Path(bound['inventory']['path'])), bound['inventory']['sha256'], run_id,
                    Path(bound['reportDateCache']['path']), bound['reportDateCache']['sha256'])
            finally:
                book.close()
            if job.get('documentKoRefreshRows') != len(ko_rows):
                errors.append('ko后台候选行数与独立回读不一致')
        except (KeyError, ValueError, RuntimeError, OSError) as exc:
            errors.append(f'ko后台派发独立核验失败:{exc}')
    for field in ("reportDateCachePath", "reportDateCacheSha256"):
        if not config.get(field) or job.get(field) != config[field]:
            errors.append(f"任务F后台作业未绑定{field}")
    delivery = resolver.delivery_dir or resolver.resolve(config.get("deliveryDir"), required=False)
    mutable_paths = [config.get("taskFBackgroundState"), config.get("taskFBackgroundReport")]
    if any(not value or not Path(str(value)).is_absolute() for value in mutable_paths):
        errors.append("任务F后台状态/报告未配置独立绝对路径")
    if delivery and any(Path(str(value)).resolve().is_relative_to(delivery.resolve())
                        for value in mutable_paths if value):
        errors.append("任务F后台可变产物不得写入主交付目录及审计清单")
    status = str(config.get("taskFStageStatus") or "pending")
    if status not in {"queued", "running", "pending"}:
        errors.append("主报告仅能审计任务F不可变待完成启动状态，不能推断后台终态")
    if config.get("taskFIncrementalWorkbook") or config.get("taskFNewActivationImportedCount") not in (None, 0):
        errors.append("后台待完成时主报告不得声明已新入库或增量文件已生成")
    details = {"background": True, "dispatch_contract_verified": not errors,
               "completion_verified": False, "task_status": status,
               "mutable_child_state_read": False, "main_report_independent": True,
               "job_path": str(job_path), "job_sha256": digest,
               "background_state": config.get("taskFBackgroundState"),
               "background_report": config.get("taskFBackgroundReport")}
    if errors:
        return status_result("任务F独立后台作业", "task_f", errors, [], details, "")
    return CheckResult("任务F独立后台作业", "task_f", PENDING,
                       "已核验独立启动边界；后台查询/入库仍待完成，不计入通过或失败", details=details)


def check_task_f_background_launch_failure(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    """A proved detached-launch failure is an F warning, not a failed report."""
    errors: list[str] = []
    receipt = resolver.resolve(config.get("taskFBackgroundLaunchReceipt"), required=False)
    expected_hash = normalize_hash(config.get("taskFBackgroundLaunchReceiptSha256"))
    if not receipt or not expected_hash or sha256_file(receipt) != expected_hash:
        return CheckResult("任务F后台启动异常", "task_f", FAIL, "后台启动失败回执缺失或SHA256不符")
    failure = read_json(receipt)
    for key, expected in (("schemaVersion", 1), ("runId", config.get("runId")),
                           ("status", FAIL), ("reason", "background_launch_failed"),
                           ("executed", False), ("workerSpawnConfirmed", False),
                           ("mainReportIndependent", True), ("cloudUploadAttempted", False),
                           ("resultsWriteBackToDocument", False)):
        if failure.get(key) != expected:
            errors.append(f"后台启动异常{key}回执不一致")
    if not failure.get("error") or failure.get("error") != config.get("taskFBackgroundLaunchError"):
        errors.append("后台启动失败未记录确切错误")
    if (config.get("taskFStageStatus") != FAIL or config.get("taskFExecuted") is not False
            or config.get("taskFMainReportIndependent") is not True):
        errors.append("后台启动失败却声明已完成执行或主表不独立")
    document = resolver.resolve(failure.get("documentPath"), required=False)
    if not document or sha256_file(document) != normalize_hash(failure.get("documentSha256")):
        errors.append("后台启动失败未绑定本轮已生成文档")
    proof_path = resolver.resolve(config.get("taskFDocumentReceipt"), required=False)
    if not proof_path or check_card_receipt(proof_path, not google_disabled_scope_requested(config)).status == FAIL:
        errors.append("后台启动失败时主表独立回读未通过")
    else:
        proof = unwrap_card_receipt(read_json(proof_path))
        reviewed = (proof.get("verification") or {}).get("system_kooda_source_review") or {}
        if reviewed.get("run_id") != config.get("runId"):
            errors.append("后台启动失败主表回读不属于本轮")
    if (not config.get("reportDateCacheSha256")
            or failure.get("reportDateCacheSha256") != config["reportDateCacheSha256"]):
        errors.append("后台启动失败未绑定本轮冻结日期快照")
    details = {"background": True, "task_status": FAIL, "completion_verified": False,
               "main_report_independent": True, "launch_failure_receipt": str(receipt),
               "error": failure.get("error")}
    return status_result("任务F后台启动异常", "task_f", errors,
                         ["任务F后台启动失败；主表已独立完成，后台错误见报告，不计为查询成功"], details, "")


def check_task_f(config: dict[str, Any], resolver: PathResolver) -> CheckResult:
    """Audit the default-on, post-document Task F supplemental workflow."""
    if task_a_only_scope_requested(config):
        return check_task_a_only_scope(config, resolver)
    if config.get("taskFSkippedReason") in {"no_missing_activation_after_fallback", "no_missing_authoritative_dates"}:
        return check_task_f_no_missing_dates(config, resolver)
    if config.get("taskFBackground"):
        return check_task_f_background(config, resolver)
    errors: list[str] = []
    warnings: list[str] = []
    executed = config.get("taskFExecuted") is True
    authorized = config.get("taskFCurrentUserAuthorized") is True
    automatic = config.get("taskFAutomaticCatalogSync") is True
    details: dict[str, Any] = {
        "executed": executed,
        "authorized": authorized,
        "automatic_catalog_sync": automatic,
        "default": "on_post_document_catalog_incremental",
        "login_required": config.get("taskFLoginRequired") is True,
    }
    if not executed:
        errors.append("任务F默认目录同步未执行")
        return status_result(
            "任务F Kooda激活日期数据库", "task_f", errors, warnings, details,
            "任务F默认目录同步缺失",
        )
    if not (automatic or authorized):
        errors.append("任务F既非默认目录同步，也没有当前真实用户诊断授权")
    if automatic and not task_f_document_boundary_matches(config):
        errors.append("任务F默认流程未证明先生成文档、后后台补查及对应文档日期来源边界")
    if config.get("taskFLoginRequired") is True:
        errors.append("任务F完成配置仍标记需要登录")

    required_keys = (
        "taskFCatalogReceipt", "taskFLoginReceipt", "taskFInventoryReceipt",
        "taskFPendingReceipt", "taskFLookupReceipt", "taskFImportReceipt",
        "taskFDatabaseSnapshot", "taskFIccidLibrary", "taskFLatestIccidLibrary",
    )
    paths: dict[str, Path | None] = {}
    for key in required_keys:
        paths[key] = resolver.resolve(config.get(key), required=False)
        if paths[key] is None:
            errors.append(f"任务F缺少{key}")
    if errors:
        return status_result(
            "任务F Kooda激活日期数据库", "task_f", errors, warnings, details,
            "任务F回执或最新库不完整",
        )

    catalog = read_json(paths["taskFCatalogReceipt"])
    login = read_json(paths["taskFLoginReceipt"])
    inventory = read_json(paths["taskFInventoryReceipt"])
    pending = read_json(paths["taskFPendingReceipt"])
    lookup = read_json(paths["taskFLookupReceipt"])
    imported = read_json(paths["taskFImportReceipt"])
    library = read_json(paths["taskFIccidLibrary"])
    latest_library = read_json(paths["taskFLatestIccidLibrary"])

    catalog_total = int(catalog.get("sourceRowCount") or 0)
    catalog_found = int(catalog.get("foundCount") or 0)
    catalog_missing = int(catalog.get("activationMissingCount") or 0)
    if catalog.get("status") != "passed":
        errors.append("任务F权威附件事务导入回执未通过")
    if (
        catalog_total <= 0
        or int(catalog.get("normalizedUniqueIccidCount") or 0) != catalog_total
        or catalog_found + catalog_missing != catalog_total
        or int(catalog.get("invalidCount") or 0) != 0
        or catalog.get("trailingFNormalization") != "strip_exactly_one_final_F_case_insensitive"
        or catalog.get("activationField") != "lblActivateDate"
        or catalog.get("cutoffRule") != "activation_date_plus_29_days_inclusive_30_day_period"
    ):
        errors.append("任务F权威附件数量、唯一ICCID、空激活日期或日期口径不一致")
    if catalog_total != int(config.get("taskFCatalogSourceRowCount") or 0):
        errors.append("任务F权威附件总数与配置不一致")
    if catalog_found != int(config.get("taskFCatalogFoundCount") or 0):
        errors.append("任务F权威附件有激活日期数量与配置不一致")
    if catalog_missing != int(config.get("taskFCatalogActivationMissingCount") or 0):
        errors.append("任务F权威附件未激活数量与配置不一致")

    if (
        inventory.get("status") != "passed"
        or int(inventory.get("pageSize") or 0) != 500
        or inventory.get("oprSearch") != "koo"
    ):
        errors.append("任务F OpenSMS Kooda全量清单回执无效")
    if inventory.get("silentBrowserMode") is not True or inventory.get("browserForegroundActivationAttempted") is not False:
        errors.append("任务F OpenSMS Kooda库存读取未证明后台静默")
    if inventory.get("fullSourceDataIncluded") is not True:
        errors.append("任务F库存回执未声明保留完整OpenSMS源数据")
    inventory_rows = [row for row in inventory.get("rows") or [] if isinstance(row, dict)]
    if any(not isinstance(row.get("sourceData"), dict) for row in inventory_rows):
        errors.append("任务F库存回执存在缺失完整OpenSMS源对象的ICCID")
    if int(inventory.get("missingCardStoreCount") or 0):
        warnings.append("任务F Kooda清单存在无卡存储行，已单列且无法形成ICCID联网候选")
    if int(inventory.get("invalidCardStoreCount") or 0):
        errors.append("任务F Kooda清单存在无效卡存储ICCID")
    if int(inventory.get("uniqueIccidCount") or 0) != int(config.get("taskFInventoryCount") or 0):
        errors.append("任务F Kooda唯一ICCID计数不一致")

    expected_fields = ["Purchase Date", "Act. Request Date", "Activation Date", "SIM Number", "SKU"]
    if (
        pending.get("status") != "passed"
        or pending.get("terminalStatuses") != ["found", "not_found", "activation_missing"]
        or pending.get("trailingFNormalization") != "strip_exactly_one_final_F_case_insensitive"
        or int(pending.get("recordSchemaVersion") or 0) != 2
        or pending.get("recordedWebFields") != expected_fields
    ):
        errors.append("任务F数据库差集回执无效")
    if int(pending.get("alreadyQueriedCount") or 0) != int(config.get("taskFAlreadyQueriedCount") or 0):
        errors.append("任务F数据库已有ICCID计数不一致")
    if int(pending.get("pendingCount") or 0) != int(config.get("taskFPendingCount") or 0):
        errors.append("任务F联网增量ICCID计数不一致")
    report_backfill = config.get("taskFReportBackfillMode") is True or pending.get("reportBackfillMode") is True
    if report_backfill:
        report_errors, report_details = review_task_f_report_backfill(config, pending, inventory, lookup, resolver)
        errors.extend(report_errors)
        details["report_backfill_review"] = report_details

    input_count = int(lookup.get("inputCount") or 0)
    no_network_gap = input_count == 0
    if no_network_gap:
        if login.get("status") != "not_required_no_missing_iccid" or login.get("loginStatus") != "not_required":
            errors.append("任务F无ICCID缺口时登录回执不应访问Grab&Go")
        if (
            lookup.get("status") != "passed"
            or lookup.get("loginStatus") != "not_required"
            or lookup.get("mode") != "task-f-local-catalog-no-network-gap"
        ):
            errors.append("任务F无ICCID缺口查询回执无效")
        if any(int(lookup.get(key) or 0) != 0 for key in (
            "effectiveWorkers", "uniqueAspNetSessionCount", "sameOriginHttpRequests",
            "backgroundJavascriptFallbackRequests",
        )):
            errors.append("任务F无ICCID缺口时产生了线程、会话或网络请求")
        if lookup.get("results") not in ([], None):
            errors.append("任务F无ICCID缺口时查询结果不为空")
    else:
        if login.get("status") not in {"authenticated", "passed"} or login.get("loginStatus") not in {None, "authenticated"}:
            errors.append("任务F Grab&Go登录回执未通过")
        if lookup.get("status") != "passed" or lookup.get("loginStatus") != "authenticated":
            errors.append("任务F Grab&Go查询回执未完整通过")
        if lookup.get("mode") not in {
            "task-f-grabandgo-read-only-webforms-bitbrowser-430",
            "task-f-grabandgo-read-only-webforms-independent-http-sessions",
        }:
            errors.append("任务F查询模式不是受控只读模式")
        if int(lookup.get("perWorkerDelayMs") or 0) < 1000:
            errors.append("任务F每线程查询等待时间不足1秒")
    if lookup.get("silentBrowserMode") is not True or lookup.get("browserForegroundActivationAttempted") is not False:
        errors.append("任务F查询未证明后台静默且未抢占焦点")
    if lookup.get("transportPriority") != ["same-origin-http", "background-page-javascript"]:
        errors.append("任务F未按同源HTTP优先、后台页面JavaScript回退执行")
    if (
        int(lookup.get("recordSchemaVersion") or 0) != 2
        or lookup.get("recordedWebFields") != expected_fields
        or lookup.get("fullSourceDataIncluded") is not True
    ):
        errors.append("任务F查询回执未声明网页五字段与完整源数据结构")
    for row in [item for item in lookup.get("results") or [] if isinstance(item, dict)]:
        if not isinstance(row.get("sourceOccurrences"), list) or not row.get("sourceOccurrences"):
            errors.append(f"任务F查询结果缺少完整源记录: {row.get('queryIccid') or 'unknown'}")
            break
        if row.get("status") in {"found", "activation_missing"}:
            presence = row.get("webFieldPresence") if isinstance(row.get("webFieldPresence"), dict) else {}
            if any(presence.get(key) is not True for key in ("purchaseDate", "actRequestDate", "activationDate", "simNumber", "sku")):
                errors.append(f"任务F查询结果缺少网页字段节点: {row.get('queryIccid') or 'unknown'}")
                break
    requested_workers = int(lookup.get("requestedWorkers") or 0)
    effective_workers = int(lookup.get("effectiveWorkers") or 0)
    if requested_workers != 8:
        errors.append("任务F程序未固定请求8个工作器")
    if no_network_gap:
        if effective_workers != 0:
            errors.append("任务F无差集时实际工作器必须为0")
    elif not 1 <= effective_workers <= min(8, input_count):
        errors.append("任务F实际并发工作线程数无效")
    if lookup.get("adaptiveConcurrency") is not True or lookup.get("sharedSessionForcesSerial") is not True:
        errors.append("任务F查询回执未声明自适应并发与共享会话强制串行")
    if effective_workers > 1 and (
        lookup.get("independentAspNetSessions") is not True
        or int(lookup.get("uniqueAspNetSessionCount") or 0) != effective_workers
    ):
        errors.append("任务F多工作器未证明每个工作器拥有独立ASP.NET会话")
    if effective_workers == 1 and input_count > 1 and (
        lookup.get("workerSessionMode") != "shared-aspnet-session-serial-worker"
        or lookup.get("concurrencyLimitedReason") != "shared_server_session_forces_serial_execution"
    ):
        errors.append("任务F共享会话降为单工作器时缺少准确原因")
    if lookup.get("sharedAspNetSessionConcurrentWorkers") is True:
        errors.append("任务F禁止在共享ASP.NET会话上并发查询")
    if lookup.get("strictIdentityGuard") != "queryIccid_equals_listSimNumber_equals_detailSimIccid":
        errors.append("任务F未证明查询、列表和详情ICCID三重一致性保护")
    if lookup.get("identityMismatchStoredAsTerminal") is not False:
        errors.append("任务F身份不一致可能被错误写为数据库终态")
    if int(lookup.get("inputCount") or 0) != int(pending.get("pendingCount") or 0):
        errors.append("任务F Grab&Go输入数与文档ICCID差集不一致")
    if lookup.get("batchMode") is not True or imported.get("batchMode") is not True:
        errors.append("任务F增量查询未启用分批数据库断点")
    if lookup.get("databaseCheckpointAfterEveryBatch") is not True or imported.get("databaseCheckpointAfterEveryBatch") is not True:
        errors.append("任务F未证明每批完成后立即写入数据库")
    if int(lookup.get("batchCount") or 0) != int(lookup.get("completedBatchCount") or 0):
        errors.append("任务F查询批次未全部完成")
    if imported.get("status") != "passed" or int(imported.get("transientErrorCount") or 0) != 0:
        errors.append("任务F数据库导入存在瞬时失败")
    if (
        imported.get("cutoffRule") != "activation_date_plus_29_days_inclusive_30_day_period"
        or int(imported.get("recordSchemaVersion") or 0) != 2
        or imported.get("recordedWebFields") != expected_fields
        or imported.get("fullSourceDataIncluded") is not True
    ):
        errors.append("任务F数据库导入字段或截止日口径无效")
    if int(imported.get("foundImportedCount") or 0) != int(config.get("taskFFoundImportedCount") or 0):
        errors.append("任务F激活日期导入计数不一致")

    for name, payload in (("本轮ICCID库", library), ("固定最新ICCID库", latest_library)):
        row_count = int(payload.get("recordCount") or 0)
        status_counts = payload.get("statusCounts") if isinstance(payload.get("statusCounts"), dict) else {}
        if (
            payload.get("status") != "passed"
            or row_count <= 0
            or int(payload.get("uniqueNormalizedIccidCount") or 0) != row_count
            or sum(int(value or 0) for value in status_counts.values()) != row_count
            or payload.get("trailingFNormalization") != "strip_exactly_one_final_F_case_insensitive"
            or payload.get("activationField") != "lblActivateDate"
            or payload.get("cutoffRule") != "activation_date_plus_29_days_inclusive_30_day_period"
            or payload.get("includesCompleteSourceData") is not True
        ):
            errors.append(f"任务F{name}结构或计数无效")
    if sha256_file(paths["taskFIccidLibrary"]) != sha256_file(paths["taskFLatestIccidLibrary"]):
        errors.append("任务F本轮ICCID库与固定最新库内容不一致")
    if int(library.get("recordCount") or 0) != int(config.get("taskFIccidLibraryRecordCount") or 0):
        errors.append("任务F最新ICCID库总数与配置不一致")
    if int(library.get("foundCount") or 0) != int(config.get("taskFIccidLibraryFoundCount") or 0):
        errors.append("任务F最新ICCID库有激活日期数量与配置不一致")
    if int(library.get("activationMissingCount") or 0) != int(config.get("taskFIccidLibraryActivationMissingCount") or 0):
        errors.append("任务F最新ICCID库未激活数量与配置不一致")

    snapshot = paths["taskFDatabaseSnapshot"]
    terminal_count = 0
    try:
        connection = sqlite3.connect(f"file:{snapshot.resolve()}?mode=ro", uri=True)
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        terminal_count = int(connection.execute("SELECT COUNT(*) FROM kooda_activation").fetchone()[0])
        invalid_cutoffs = int(connection.execute(
            "SELECT COUNT(*) FROM kooda_activation WHERE lookup_status='found' AND (activation_date IS NULL OR cutoff_date IS NULL)"
        ).fetchone()[0])
        columns = {str(row[1]) for row in connection.execute("PRAGMA table_info(kooda_activation)")}
        required_columns = {
            "purchase_date", "purchase_date_raw", "act_request_date", "activation_date_raw",
            "sim_number", "sim_number_raw", "sku", "cost", "source_data_json",
            "record_schema_version", "data_source", "source_file", "source_sha256", "vendor_ordinal",
        }
        current_schema_count = int(connection.execute(
            "SELECT COUNT(*) FROM kooda_activation WHERE record_schema_version >= 2"
        ).fetchone()[0]) if required_columns.issubset(columns) else 0
        if report_backfill:
            reused = {row["queryIccid"]: row for row in pending.get("alreadyQueried") or [] if isinstance(row, dict) and row.get("queryIccid")}
            reused_keys = list(reused)
            found_reused: set[str] = set()
            for offset in range(0, len(reused_keys), 900):
                batch = reused_keys[offset:offset + 900]
                placeholders = ",".join("?" for _ in batch)
                for query, status, activation, cutoff in connection.execute(
                    f"SELECT query_iccid, lookup_status, activation_date, cutoff_date FROM kooda_activation WHERE query_iccid IN ({placeholders})", batch,
                ):
                    found_reused.add(query)
                    expected = reused[query]
                    if (status != "found" or activation != expected.get("activationDate") or cutoff != expected.get("cutoffDate")):
                        errors.append("任务F复用待人工上传日期与本轮数据库快照不一致")
            if found_reused != set(reused_keys):
                errors.append("任务F复用待人工上传ICCID在本轮数据库快照缺失")
        connection.close()
        if integrity != "ok" or invalid_cutoffs or not required_columns.issubset(columns):
            errors.append("任务F数据库快照完整性、字段或found日期无效")
        if current_schema_count != int(config.get("taskFDatabaseCurrentSchemaCount") or 0):
            errors.append("任务F数据库完整记录结构计数不一致")
        if terminal_count != int(library.get("recordCount") or 0):
            errors.append("任务F数据库快照与最新ICCID库记录数不一致")
    except sqlite3.Error as exc:
        errors.append(f"任务F数据库快照无法读取: {exc}")

    details.update({
        "catalog_source_row_count": catalog_total,
        "catalog_found_count": catalog_found,
        "catalog_activation_missing_count": catalog_missing,
        "inventory_count": config.get("taskFInventoryCount"),
        "database_present_count": config.get("taskFAlreadyQueriedCount"),
        "pending_online_count": config.get("taskFPendingCount"),
        "online_lookup_triggered": config.get("taskFOnlineLookupTriggered") is True,
        "post_document_supplemental": config.get("taskFPostDocumentSupplemental") is True,
        "document_generated_before_lookup": config.get("taskFDocumentGeneratedBeforeLookup") is True,
        "results_apply_next_export": config.get("taskFResultsApplyNextExport") is True,
        "document_date_source": config.get("taskFDocumentDateSource"),
        "results_write_back_to_document": config.get("taskFResultsWriteBackToDocument"),
        "document_missing_activation_red_rows": config.get("taskFDocumentMissingActivationRedRows"),
        "document_red_review_passed": config.get("taskFDocumentRedReviewPassed") is True,
        "requested_workers": requested_workers,
        "effective_workers": effective_workers,
        "database_terminal_count": terminal_count,
        "iccid_library_record_count": library.get("recordCount"),
        "iccid_library": str(paths["taskFIccidLibrary"]),
        "latest_iccid_library": str(paths["taskFLatestIccidLibrary"]),
    })
    return status_result(
        "任务F Kooda激活日期数据库", "task_f", errors, warnings, details,
        "任务F权威附件、文档ICCID差集、增量查询和最新库通过",
    )


def build_tasks(
    config: dict[str, Any],
    resolver: PathResolver,
    inputs: dict[str, Any],
    workers: int,
) -> list[AuditTask]:
    manual_export = config.get("runType") in {"manual_export", "manual_export_task_b"}
    source_hash = normalize_hash(config.get("sourceSha256"))
    google_hash = normalize_hash(config.get("googleSha256"))
    google_enabled = not google_disabled_scope_requested(config)
    chatr_enabled = chatr_google_scope_requested(config)
    chatr_provenance_cache: list[CheckResult] = []
    chatr_provenance_lock = threading.Lock()

    def current_chatr_provenance() -> CheckResult | None:
        if not chatr_enabled:
            return None
        # Three parallel audit consumers share one raw-source scan, scoped to
        # this invocation; no historical receipt or cross-run cache is reused.
        with chatr_provenance_lock:
            if not chatr_provenance_cache:
                chatr_provenance_cache.append(check_chatr_google_provenance(config, resolver, inputs))
            return chatr_provenance_cache[0]
    if not source_hash:
        raise ValueError("配置缺少有效的本轮源表SHA-256")
    if google_enabled and not google_hash:
        raise ValueError("配置缺少有效的Google表SHA-256")
    result_hash = sha256_file(inputs["result"])
    tasks = [
        AuditTask("本轮源文件", "files", lambda: check_xlsx("本轮源文件", inputs["source"], source_hash)),
        AuditTask(
            "卡槽关联结果表",
            "files",
            lambda: check_xlsx("卡槽关联结果表", inputs["result"], result_hash, True, inputs["server_date"], google_enabled,
                               config.get("googleSystemMissingSheetEnabled") is True),
        ),
        AuditTask(
            "后台导出回执",
            "freshness",
            lambda: check_export_receipt(inputs["export_receipt"], source_hash, inputs["source"]),
        ),
        AuditTask(
            "服务器日期",
            "freshness",
            lambda: check_server_dates(inputs["server_receipts"], inputs["server_date"]),
        ),
        AuditTask("卡槽关联完整复核", "document", lambda: check_card_receipt(inputs["card_receipt"], google_enabled)),
        AuditTask("日期来源与冻结缓存", "document", lambda: check_date_resolution(config, resolver, inputs)),
        AuditTask(
            "文档接口调度",
            "document",
            lambda: check_dispatch_receipt(
                inputs["dispatch_receipt"],
                source_hash,
                google_hash,
                result_hash,
                bool(config.get("carrierRefreshReceipt")),
                google_enabled,
            ),
        ),
        AuditTask(
            "运营商最新状态回写文档",
            "document",
            lambda: check_carrier_refresh(config, resolver, inputs["result"]),
        ),
        AuditTask("任务B授权", "task_b", lambda: check_task_b(config, resolver)),
        AuditTask("任务C授权与规则", "task_c", lambda: check_task_c(config, resolver)),
        AuditTask("任务D下发下卡通知", "task_d", lambda: check_task_d(config, resolver)),
        AuditTask("已删除保留区统计关闭核验" if config.get("taskDRetentionSummaryContractVersion") == 2 else
                  "任务D新增入区与结束存量", "task_d_retention", lambda: check_task_d_retention_counts(config, resolver)),
        AuditTask("任务E按截止日下卡通知", "task_e", lambda: check_task_e(config, resolver)),
        AuditTask("任务F Kooda激活日期数据库", "task_f", lambda: check_task_f(config, resolver)),
        AuditTask("A21/A2来源运营商保护", "carrier", lambda: check_protected_carrier_sources(config.get("groups") or [], resolver)),
        AuditTask("B13目标写入禁令", "carrier", lambda: check_forbidden_b13_targets(config.get("groups") or [], resolver)),
    ]
    if config.get("taskG"):
        tasks.append(AuditTask("任务G周末休息与次日恢复", "carrier", lambda: check_task_g(config, resolver, inputs["result"])))
    if config.get("googleSystemMissingSheetEnabled") is True:
        tasks.append(AuditTask("所有谷歌有系统没有", "document", lambda: check_google_difference(config, inputs)))
    if config.get('koUnactivatedSheetEnabled') is True or config.get('koUnactivatedSheet'):
        tasks.append(AuditTask('ko系统有-未激活', 'document', lambda: check_ko_unactivated(config, inputs)))
    if google_enabled:
        tasks.extend([
            AuditTask("本轮Google关联表", "files", lambda: check_xlsx("本轮Google关联表", inputs["association"], google_hash)),
            AuditTask("Google关联表回执", "freshness", lambda: check_google_receipt(inputs["google_receipt"], google_hash, inputs["association"])),
        ])
    else:
        tasks.append(AuditTask("Google依赖关闭", "freshness", lambda: check_google_dependency(
            config, resolver, inputs, current_chatr_provenance())))
        if chatr_enabled:
            tasks.append(AuditTask("Chatr Google日期补缺来源", "freshness",
                                   current_chatr_provenance))
    if manual_export:
        tasks.extend(
            [
                AuditTask(
                    "手动导出写入范围",
                    "carrier",
                    lambda: status_result(
                        "手动导出写入范围",
                        "carrier",
                        [],
                        [],
                        {
                            "scope": config.get("runType"),
                            "task_b_executed": config.get("taskBExecuted") is True,
                        },
                        "手动导出仅执行当前授权的专项写入" if config.get("taskBExecuted") is True else "手动导出未执行运营商或上下架写入",
                    ),
                ),
                AuditTask(
                    "下架恢复",
                    "offline",
                    lambda: status_result(
                        "下架恢复",
                        "offline",
                        [],
                        [],
                        {"scope": "manual_export"},
                        "手动导出不包含下架恢复",
                    ),
                ),
            ]
        )
    else:
        tasks.extend(
            [
                AuditTask("源报表分析", "analysis", lambda: check_analysis_receipt(inputs["analysis"])),
                AuditTask("余量与按日回复率", "balance", lambda: check_balance_receipt(inputs["balance"])),
                AuditTask("候选互斥性", "carrier", lambda: check_candidate_disjoint(config.get("groups") or [], resolver)),
                AuditTask("回复型流转发送窗口数门禁", "carrier", lambda: check_reply_rate_send_gate(config, resolver)),
                AuditTask("休息区转入关闭门禁", "carrier", lambda: check_rest_zone_entry_disabled(config, resolver)),
                AuditTask("B区V组固定优先级", "carrier", lambda: check_v_zone_priority_rule(config, resolver)),
                AuditTask("C2低回复率与发送量规则", "carrier", lambda: check_c2_routing_rule(config, resolver)),
            ]
        )
    if config.get("runType") in {"task_b_only", "task_c_only"}:
        scope = config.get("runType")
        scope_label = "任务B" if scope == "task_b_only" else "任务C"
        tasks.append(
            AuditTask(
                "下架恢复",
                "offline",
                lambda: status_result(
                    "下架恢复",
                    "offline",
                    [],
                    [],
                    {"scope": scope},
                    f"{scope_label}专项执行不包含下架恢复",
                ),
            )
        )
    elif not manual_export:
        tasks.append(AuditTask("下架恢复", "offline", lambda: check_offline(config, resolver)))
    if config.get("downlist"):
        tasks.append(AuditTask("二号机房全量下架", "offline", lambda: check_requested_downlist(config, resolver)))
    if config.get("reportDetailMode") == "summary":
        tasks.extend([
            AuditTask("V组转出边界", "carrier", lambda: check_v_group_boundary(config, resolver)),
            AuditTask("V组每日转换号统计", "delivery", lambda: check_v_group_daily(config)),
        ])
    for index, group in enumerate(config.get("groups") or [], start=1):
        title = str(group.get("title") or f"运营商组{index}")
        tasks.append(AuditTask(title, "carrier", lambda group=group, title=title: check_group(title, group, resolver)))
    if inputs.get("delivery_dir"):
        tasks.append(
            AuditTask(
                "最终交付校验",
                "delivery",
                lambda: check_delivery(inputs["delivery_dir"], hash_workers=workers,
                                       task_a_only=task_a_only_scope_requested(config),
                                       google_disabled=not google_enabled,
                                       allowed_google_paths=(current_chatr_provenance().details.get(
                                           "allowed_artifact_paths", []) if chatr_enabled else [])),
            )
        )
    cpu_calls = {
        "本轮源文件": ("check_xlsx", ("本轮源文件", inputs["source"], source_hash)),
        "卡槽关联结果表": ("check_xlsx", ("卡槽关联结果表", inputs["result"], result_hash, True,
            inputs["server_date"], google_enabled, config.get("googleSystemMissingSheetEnabled") is True)),
        "运营商最新状态回写文档": ("check_carrier_refresh", (config, resolver, inputs["result"])),
        "任务G周末休息与次日恢复": ("check_task_g", (config, resolver, inputs["result"])),
        "所有谷歌有系统没有": ("check_google_difference", (config, inputs)),
        "ko系统有-未激活": ("check_ko_unactivated", (config, inputs)),
        "A21/A2来源运营商保护": ("check_protected_carrier_sources", (config.get("groups") or [], resolver)),
        "B13目标写入禁令": ("check_forbidden_b13_targets", (config.get("groups") or [], resolver)),
        "本轮Google关联表": ("check_xlsx", ("本轮Google关联表", inputs["association"], google_hash)),
    }
    for task in tasks:
        task.cpu_call = cpu_calls.get(task.name)
    return tasks


def concurrency_plan(google_enabled: bool = False, chatr_google_enabled: bool = False) -> dict[str, Any]:
    return {
        "parallel_after_run": [
            "源表、Google表、结果表的哈希与XLSX完整性" if google_enabled else
            "源表、结果表及Chatr新鲜Google来源/日期快照的哈希；全面关联跳过证明" if chatr_google_enabled else
            "源表、结果表的哈希与XLSX完整性；Google依赖跳过证明",
            "结果表工作表、表头、公式错误和日期连续性",
            "导出、Google下载、服务器日期、文档接口、分析、余量回执" if google_enabled else
            "导出、Chatr限定Google日期来源、服务器日期、本机文档、分析、余量回执；无旧Google文件混入" if chatr_google_enabled else
            "导出、服务器日期、本机文档、分析、余量回执；无Google输入或旧文件混入",
            "各运营商组候选/写入/只读复核回执的一致性",
            "候选互斥、V组固定优先级、下架恢复关闭状态、任务B/任务C授权、最终交付清单",
        ],
        "parallel_during_collection": [
            "新鲜Google表下载与/day、/phoneConfig日期读取" if google_enabled else
            "新鲜系统导出、系统库存、Chatr日期所需Google下载及/day、/phoneConfig日期读取" if chatr_google_enabled else
            "新鲜系统导出、系统库存及/day、/phoneConfig日期读取；Google下载关闭",
            "互不依赖的只读余量与按日回复率查询（建议并发2到4，避免压垮同一后台）",
            "文档作业等待期间可执行只读分析和候选预计算",
            "V组候选与标准组、A区、C2候选合并为同一次新鲜源表扫描，不另开重复文档任务",
        ],
        "must_remain_serial": [
            "后台生成export4day完成后才能分析",
            "运营商实际写入必须遵守B区 -> Chatr -> E恢复 -> A区的流转优先级",
            "V组同一号码按优先级1、2、3首个命中即停止；B12V未命中优先级1时保持原组",
            "每个写入组必须先验证，再写入，再独立回读",
            "任务B必须等卡槽结果表和服务器日期确认完成，且需要当前授权",
            "任务C必须等任务B后工作簿刷新完成，取得当前授权并实时分别唯一解析换号区T和换号区C目标",
            "最终SHA-256清单必须在所有交付文件写完之后生成",
        ],
        "recommended_workers": {
            "receipt_and_xlsx_checks": "auto（实测默认最多4）；max受逻辑CPU及16工作器上限约束",
            "local_cpu_checks": "hybrid：最多4个spawn进程；与协调线程共用总在途上限，全部收尾后交叉核验",
            "same_backend_read_queries": 4,
            "large_file_hashes": "auto（不超过文件数）",
            "live_write_pipeline": 1,
        },
    }


def find_config(delivery_dir: Path) -> Path:
    candidates = list((delivery_dir / "receipts").glob("opensms_patrol_finalize_*_config.json"))
    if not candidates:
        candidates = list(delivery_dir.glob("opensms_patrol_finalize_*_config.json"))
    if len(candidates) != 1:
        raise SystemExit(f"交付目录应有且仅有一个巡检配置，实际找到{len(candidates)}个")
    return candidates[0].resolve()


def unique_output_path(config_path: Path) -> Path:
    stamp = datetime.now().astimezone().strftime("%Y%m%dT%H%M%S%z")
    configured = os.environ.get("OPENSMS_AUDIT_OUTPUT_DIR")
    if configured:
        output_dir = Path(configured).expanduser().resolve()
    else:
        workspace_snapshots = Path.cwd() / "opensms_artifacts" / "snapshots"
        output_dir = workspace_snapshots if workspace_snapshots.is_dir() else config_path.parent
    return output_dir / f"opensms_parallel_audit_{stamp}_receipt.json"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="OpenSMS全流程程序化并行检查（只读）")
    parser.add_argument("--config", type=Path, help="巡检finalize配置；可只传--delivery-dir自动发现")
    parser.add_argument("--delivery-dir", type=Path, help="本轮最终交付目录")
    parser.add_argument("--server-date-receipt", type=Path, action="append", default=[])
    parser.add_argument("--output", type=Path, help="详细JSON回执输出路径")
    parser.add_argument(
        "--workers",
        default="auto",
        help="并发检查数；auto默认最多4；max不超过逻辑CPU数且最多16",
    )
    parser.add_argument("--warnings-as-errors", action="store_true")
    parser.add_argument("--executor", choices=("threads", "hybrid"), default="hybrid",
                        help="本地只读检查：线程协调；hybrid将明确的CPU检查交给最多4个进程")
    parser.add_argument("--show-plan", action="store_true", help="只显示并发方案")
    parser.add_argument("--json", action="store_true", help="标准输出使用精简JSON")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if args.show_plan:
        print(json.dumps(concurrency_plan(), ensure_ascii=False, indent=2))
        return 0
    try:
        args.workers = resolve_worker_count(args.workers)
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc
    delivery_dir = args.delivery_dir.expanduser().resolve() if args.delivery_dir else None
    if args.config:
        config_path = args.config.expanduser().resolve()
    elif delivery_dir:
        config_path = find_config(delivery_dir)
    else:
        raise SystemExit("必须提供--config或--delivery-dir")
    config = read_json(config_path)
    resolver, inputs = discover_inputs(config_path, config, delivery_dir, args.server_date_receipt)
    output_path = args.output.expanduser().resolve() if args.output else unique_output_path(config_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    tasks = build_tasks(config, resolver, inputs, args.workers)
    started_at = now_iso()
    wall_started = time.perf_counter()
    results, meter, scheduling = execute_audit_tasks(tasks, args.workers, args.executor)
    results.sort(key=lambda item: (item.category, item.name))
    results.append(cross_check(config, inputs, results))
    wall_ms = round((time.perf_counter() - wall_started) * 1000)

    failed = [item for item in results if item.status == FAIL]
    warnings = [item for item in results if item.status == WARN]
    passed = [item for item in results if item.status == PASS]
    skipped = [item for item in results if item.status == SKIP]
    pending = [item for item in results if item.status == PENDING]
    group_results = [item for item in results if item.category == "carrier" and item.name != "候选互斥性"]
    actual = sum(int(item.details.get("actual_writes") or 0) for item in group_results)
    success = sum(int(item.details.get("successful_writes") or 0) for item in group_results)
    write_failures = sum(int(item.details.get("failed_writes") or 0) for item in group_results)
    overall = FAIL if failed or (args.warnings_as_errors and warnings) else WARN if warnings else PASS
    payload = {
        "schema_version": 1,
        "started_at": started_at,
        "finished_at": now_iso(),
        "result": overall,
        "summary": {
            "checks_total": len(results),
            "passed": len(passed),
            "warnings": len(warnings),
            "failed": len(failed),
            "skipped": len(skipped),
            "pending": len(pending),
            "actual_writes": actual,
            "successful_writes": success,
            "failed_writes": write_failures,
            "source_file": str(inputs["source"]),
            "result_file": str(inputs["result"]),
            "delivery_dir": str(inputs["delivery_dir"]) if inputs.get("delivery_dir") else None,
        },
        "timing": {
            "wall_ms": wall_ms,
            "serial_sum_ms": sum(item.duration_ms for item in results),
            "estimated_parallel_saving_ms": max(sum(item.duration_ms for item in results) - wall_ms, 0),
            "configured_workers": args.workers,
            "observed_worker_threads": len(meter.threads),
            "max_observed_concurrency": meter.maximum,
            "local_scheduling": scheduling,
        },
        "checks": [asdict(item) for item in results],
        "concurrency_plan": concurrency_plan(not google_disabled_scope_requested(config), chatr_google_scope_requested(config)),
        "config_path": str(config_path),
        "audit_receipt_path": str(output_path),
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    concise = {
        "状态": {PASS: "成功", WARN: "成功（有提示）", FAIL: "失败"}[overall],
        "检查": f"{len(passed)}/{len(results)}通过，{len(warnings)}提示，{len(failed)}失败，{len(skipped)}跳过，{len(pending)}后台待完成",
        "实际修改": actual,
        "修改成功": success,
        "修改失败": write_failures,
        "耗时毫秒": wall_ms,
        "最大并发": meter.maximum,
        "结果文件": str(inputs["result"]),
        "详细回执": str(output_path),
        "异常": [item.message for item in failed] or [item.message for item in warnings] or [],
    }
    if args.json:
        print(json.dumps(concise, ensure_ascii=False, indent=2))
    else:
        for label, value in concise.items():
            if label == "异常":
                print(f"{label}：{'；'.join(value) if value else '无'}")
            else:
                print(f"{label}：{value}")
    return 2 if overall == FAIL else 1 if overall == WARN and args.warnings_as_errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
