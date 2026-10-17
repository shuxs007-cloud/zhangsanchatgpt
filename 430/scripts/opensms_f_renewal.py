#!/usr/bin/env python3
"""Explicit, source-bound renewal dates. Never query Grab or alter raw activation records."""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import sys
import os
import fcntl
import tempfile
import shutil
import subprocess
import openpyxl
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / 'opensms_artifacts/state/task_f_kooda_activation.sqlite3'
ELIGIBLE = {'充值成功', '不需要充值'}
RULE = 'current_source_cutoff_plus_30_if_balance_gte_30_else_source_cutoff'
sys.path.insert(0, str(ROOT / 'services/opensms_document_api/pipeline'))
from task_f_renewal_dates import read_renewals, validate_renewal

def digest(path):
    with Path(path).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()

def save(path, value):
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, 'w', encoding='utf-8') as f:
        json.dump(value, f, ensure_ascii=False, indent=2)
        f.flush(); os.fsync(f.fileno())

def phone(value):
    if isinstance(value, bool):
        raise ValueError('invalid phone')
    if isinstance(value, (int, float)):
        if value != int(value) or abs(value) > 2**53:
            raise ValueError('phone precision')
        value = str(int(value))
    if not isinstance(value, str):
        raise ValueError('phone type')
    digits = re.sub(r'[+ ()-]', '', value.strip())
    if re.fullmatch(r'1[0-9]{10}', digits):
        digits = digits[1:]
    if not re.fullmatch(r'[0-9]{10}', digits):
        raise ValueError('phone format')
    return digits

def balance(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        raise ValueError('invalid balance')
    cleaned = re.sub(r'^(?:CAD\s*|C\$\s*|\$\s*)', '', str(value).strip(), flags=re.I)
    try:
        number = Decimal(cleaned)
    except InvalidOperation:
        raise ValueError('invalid balance') from None
    if not number.is_finite() or number < 0:
        raise ValueError('invalid balance')
    return number

def source_date(value, epoch=openpyxl.utils.datetime.CALENDAR_WINDOWS_1900):
    """Calendar dates, not UTC-converted timestamps. Reject ambiguous day/month text."""
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        raise ValueError('invalid date')
    if isinstance(value, (int, float)):
        if not 35000 <= value <= 90000:
            raise ValueError('invalid Excel date')
        return from_excel(value, epoch=epoch).date()
    if not isinstance(value, str):
        raise ValueError('invalid date type')
    text = value.strip()
    if not text:
        return None
    # ISO date/datetime: retain the source's written calendar day.
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:?\d{2})?)?', text):
        if len(text) > 10:
            return datetime.fromisoformat(text.replace('Z', '+00:00')).date()
        return date.fromisoformat(text)
    for fmt in ('%Y/%m/%d', '%Y.%m.%d', '%Y年%m月%d日', '%b %d, %Y',
                '%B %d, %Y', '%b %d %Y', '%B %d %Y', '%d %b %Y', '%d %B %Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    match = re.fullmatch(r'(\d{1,2})[/-](\d{1,2})[/-](\d{4})', text)
    if match:
        first, second, year = map(int, match.groups())
        if first <= 12 and second <= 12 and first != second:
            raise ValueError('ambiguous day/month date')
        return date(year, second, first) if first > 12 else date(year, first, second)
    raise ValueError('unsupported or invalid date format')

def plan(source):
    source = Path(source).resolve(); source_hash = digest(source)
    wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = wb['Sheet']
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        if headers[:4] != ['电话', '注册状态', '余额', '截止日期']:
            raise ValueError('source headers changed')
        rows, excluded, seen = [], [], set()
        status_counts = Counter()
        for cells in sheet.iter_rows(min_row=2, max_col=4):
            raw_phone, status, raw_balance, raw_date = [c.value for c in cells]
            status_counts[str(status or '空白')] += 1
            if status not in ELIGIBLE:
                continue
            item = {'sourceRow': cells[0].row, 'status': status}
            try:
                if any(c.data_type in {'f', 'e'} for c in cells):
                    raise ValueError('formula/error source cell')
                item['phone'] = phone(raw_phone)
                if item['phone'] in seen:
                    raise ValueError('duplicate eligible phone')
                seen.add(item['phone'])
                amount, cutoff = balance(raw_balance), source_date(raw_date, wb.epoch)
                item['balanceCad'] = str(amount) if amount is not None else None
                if cutoff is None:
                    raise ValueError('missing source cutoff')
                if amount is None:
                    raise ValueError('missing balance')
                extend = amount >= Decimal(30)
                item.update(sourceCutoff=cutoff.isoformat(), extensionDays=30 if extend else 0,
                            cutoff=(cutoff + timedelta(days=30 if extend else 0)).isoformat())
                rows.append(item)
            except ValueError as exc:
                item['reason'] = str(exc); excluded.append(item)
        count = sheet.max_row - 1
    finally:
        wb.close()
    if digest(source) != source_hash:
        raise ValueError('source changed during extraction')
    return {'version': 1, 'scope': 'current_user_authorized_renewal', 'source': str(source),
            'sourceSha256': source_hash, 'sourceSheet': 'Sheet', 'sourceRange': f'A1:D{count+1}',
            'rule': RULE, 'sourceRows': count, 'eligibleRows': len(rows)+len(excluded),
            'statusCounts': dict(status_counts), 'rows': rows, 'excluded': excluded,
            'computableRows': len(rows), 'extensionRows': sum(r['extensionDays']==30 for r in rows),
            'directDateRows': sum(r['extensionDays']==0 for r in rows),
            'dateCounts': dict(Counter(r['cutoff'] for r in rows))}

def build_records(source_plan, mapping, database=DB):
    if source_plan.get('rule') != RULE or source_plan.get('scope') != 'current_user_authorized_renewal':
        raise ValueError('renewal scope not authorized')
    live = {r['phone']: r for r in mapping['rows']}
    if len(live) != len(mapping['rows']):
        raise ValueError('duplicate backend phone')
    result, excluded = [], []
    with sqlite3.connect(Path(database).resolve().as_uri()+'?mode=ro', uri=True) as conn:
        base = {r[0]: r[1:] for r in conn.execute('SELECT query_iccid,activation_date,cutoff_date,lookup_status FROM kooda_activation')}
        for r in source_plan['rows']:
            row = live.get(r['phone'])
            if not row:
                excluded.append(r | {'reason': 'phone not found in backend'}); continue
            old = base.get(row['iccid'])
            if not old or old[2] != 'found':
                excluded.append(r | {'reason': 'no original F activation record'}); continue
            if row['activation'] != old[0] or date.fromisoformat(old[1]) != date.fromisoformat(old[0])+timedelta(days=29):
                raise ValueError('original activation identity/date conflict')
            result.append(r | {'iccid': row['iccid'], 'activation': old[0], 'baseCutoffDate': old[1],
                               'backendId': row['id'], 'slot': row['slot'], 'backendCutoffBefore': row['cutoff']})
    if len({r['iccid'] for r in result}) != len(result):
        raise ValueError('duplicate renewal ICCID')
    return result, excluded

def commit_renewals(source_plan, mapping, directory, run_id, database=DB):
    if not re.fullmatch(r'\d{8}T\d{6}\+0800', run_id):
        raise ValueError('invalid run ID')
    if digest(source_plan['source']) != source_plan['sourceSha256'] or plan(source_plan['source']) != source_plan:
        raise ValueError('source independent reconstruction changed')
    directory, database = Path(directory), Path(database)
    backup = directory / 'f_database_before.sqlite3'
    receipt_file = directory / 'database_commit.json'
    if backup.exists() or receipt_file.exists():
        raise FileExistsError('renewal commit cannot replay; inspect exact receipt')
    lock_path = database.with_name(database.name+'.task_f_writer.lock')
    with lock_path.open('a+b') as lock:
        # Never stop a running F worker to acquire its write lock.
        try:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise RuntimeError('F database writer busy; no changes made') from None
        rows, excluded = build_records(source_plan, mapping, database)
        if not rows:
            raise ValueError('no verifiable renewal rows')
        fd=os.open(backup, os.O_CREAT|os.O_EXCL|os.O_WRONLY, 0o600); os.close(fd)
        with sqlite3.connect(database) as conn, sqlite3.connect(backup) as snapshot:
            conn.backup(snapshot)
        with backup.open('rb') as handle:
            os.fsync(handle.fileno())
        created_at = datetime.now().astimezone().isoformat()
        records=[]
        for r in rows:
            event_id=hashlib.sha256(f"{source_plan['sourceSha256']}:{r['sourceRow']}:{r['iccid']}".encode()).hexdigest()
            proof={'version':1, 'authorization':'current-user-renewal-source', 'rule':RULE,
                   'runId':run_id, 'source':source_plan['source'], 'sourceSha256':source_plan['sourceSha256'],
                   'sourceSheet':'Sheet', 'sourceRow':r['sourceRow'], 'status':r['status'], 'balanceCad':r['balanceCad'],
                   'sourceCutoff':r['sourceCutoff'], 'extensionDays':r['extensionDays'], 'cutoffDate':r['cutoff'],
                   'activationDate':r['activation'], 'baseCutoffDate':r['baseCutoffDate'],
                   'queryIccid':r['iccid'], 'phone':r['phone'], 'slot':r['slot'], 'backendId':r['backendId'],
                   'eventId':event_id, 'recordedAt':created_at}
            validate_renewal(proof,r['iccid'],r['activation'],r['cutoff'])
            records.append(r | {'renewal':proof})
        with sqlite3.connect(database, timeout=30) as conn:
            conn.execute('PRAGMA synchronous=FULL'); conn.execute('BEGIN IMMEDIATE')
            conn.execute('CREATE TABLE IF NOT EXISTS kooda_renewal_events (event_id TEXT PRIMARY KEY,query_iccid TEXT NOT NULL,source_sha256 TEXT NOT NULL,source_row INTEGER NOT NULL,proof_json TEXT NOT NULL,created_at TEXT NOT NULL,UNIQUE(source_sha256,source_row,query_iccid))')
            conn.execute('CREATE TABLE IF NOT EXISTS kooda_renewal_current (query_iccid TEXT PRIMARY KEY,event_id TEXT NOT NULL REFERENCES kooda_renewal_events(event_id),proof_json TEXT NOT NULL,updated_at TEXT NOT NULL)')
            for r in records:
                proof = r['renewal']; encoded=json.dumps(proof,ensure_ascii=False,sort_keys=True)
                exists=conn.execute('SELECT proof_json FROM kooda_renewal_events WHERE event_id=?',(proof['eventId'],)).fetchone()
                if exists:
                    previous=json.loads(exists[0])
                    if any(previous[k]!=proof[k] for k in proof if k not in {'recordedAt','runId'}):
                        raise ValueError('existing renewal source identity changed')
                    proof=previous; encoded=exists[0]; r['renewal']=proof
                else:
                    conn.execute('INSERT INTO kooda_renewal_events VALUES (?,?,?,?,?,?)',
                        (proof['eventId'],r['iccid'],source_plan['sourceSha256'],r['sourceRow'],encoded,created_at))
                conn.execute('INSERT INTO kooda_renewal_current VALUES (?,?,?,?) ON CONFLICT(query_iccid) DO UPDATE SET event_id=excluded.event_id,proof_json=excluded.proof_json,updated_at=excluded.updated_at',
                    (r['iccid'],proof['eventId'],encoded,created_at))
            if digest(source_plan['source'])!=source_plan['sourceSha256']:
                raise ValueError('source changed before commit')
        # New read-only connection and raw-table comparison: not a commit return value.
        with sqlite3.connect(database.resolve().as_uri()+'?mode=ro',uri=True) as check, sqlite3.connect(backup.resolve().as_uri()+'?mode=ro',uri=True) as old:
            if check.execute('PRAGMA quick_check').fetchone()[0]!='ok':
                raise ValueError('renewal database integrity failed')
            for table in ('kooda_activation','kooda_inventory'):
                if check.execute(f'SELECT * FROM {table} ORDER BY query_iccid').fetchall()!=old.execute(f'SELECT * FROM {table} ORDER BY query_iccid').fetchall():
                    raise ValueError('original F table changed')
            renewals=read_renewals(check)
            if any(renewals.get(r['iccid'])!=r['renewal'] for r in records):
                raise ValueError('renewal independent database readback failed')
        receipt={'status':'passed','runId':run_id,'sourceSha256':source_plan['sourceSha256'],
                 'updatedCount':len(records),'excluded':excluded,'database':str(database.resolve()),
                 'backup':str(backup.resolve()),'backupSha256':digest(backup),'originalTablesUnchanged':True,
                 'independentReadbackPassed':True,'records':records}
        save(receipt_file,receipt)
        fcntl.flock(lock.fileno(),fcntl.LOCK_UN)
    return receipt

def export_renewals(directory):
    from opensms_export_task_f_deadlines import DEFAULT_TEMPLATE, check_template, verify_output, NODE, NODE_MODULES
    directory=Path(directory)
    committed=json.loads((directory/'database_commit.json').read_text())
    if committed.get('status')!='passed' or committed.get('independentReadbackPassed') is not True:
        raise ValueError('unverified database commit')
    rows=sorted([[r['iccid'],r['cutoff'],r['activation']] for r in committed['records']])
    with sqlite3.connect(DB.resolve().as_uri()+'?mode=ro',uri=True) as conn:
        saved=read_renewals(conn)
        if any(saved.get(r['iccid'])!=r['renewal'] for r in committed['records']):
            raise ValueError('database renewal changed before export')
    output=directory/f"任务F_续费截止日_{committed['runId']}.xlsx"
    if output.exists():
        raise FileExistsError('no overwrite of renewal workbook')
    check_template(DEFAULT_TEMPLATE)
    with tempfile.TemporaryDirectory(prefix='.renewal_writer_',dir=directory) as temporary:
        work=Path(temporary); (work/'node_modules').symlink_to(NODE_MODULES,target_is_directory=True)
        writer=work/'writer.mjs'; shutil.copyfile(ROOT/'scripts/opensms_task_f_deadline_writer.mjs',writer)
        source=work/'rows.json'; save(source,rows)
        env=dict(os.environ,OPENSMS_RENEWAL_RENDER_PATH=str(directory/'preview.png'))
        completed=subprocess.run([str(NODE),str(writer),str(DEFAULT_TEMPLATE),str(source),str(output)],env=env,capture_output=True,text=True,timeout=120)
        if completed.returncode:
            raise RuntimeError('artifact-tool renewal writer failed: '+completed.stderr[-800:])
        verified=json.loads(completed.stdout.strip().splitlines()[-1])
        if verified.get('status')!='success' or verified.get('recordCount')!=len(rows):
            raise ValueError('renewal writer count mismatch')
        verify_output(output,rows)
    os.chmod(output,0o600)
    result={'version':1,'scope':'current_user_authorized_renewal','runId':committed['runId'],
            'file':str(output.resolve()),'fileSha256':digest(output),'fileRows':len(rows),
            'sourceSha256':committed['sourceSha256'],'independentReadbackPassed':True,
            'dateFormat':'yyyy-mm-dd','records':committed['records']}
    save(directory/'renewal_input.json',result)
    return result

def finalize_renewals(directory):
    """Independent saved evidence reconstruction; never sends a business request."""
    from opensms_export_task_f_deadlines import verify_output
    from opensms_verify_card_slot_fast import load_frozen_date_snapshot
    directory=Path(directory)
    input=json.loads((directory/'renewal_input.json').read_text())
    receipt=json.loads((directory/'backend_import_receipt.json').read_text())
    committed=json.loads((directory/'database_commit.json').read_text())
    source=json.loads((directory/'source_plan.json').read_text())
    if (receipt.get('status')!='passed' or receipt.get('independentReadbackPassed') is not True
        or receipt.get('verifiedDatePairs')!=input['fileRows'] or receipt.get('ownedHiddenPageClosed') is not True
        or receipt.get('fileSha256')!=input['fileSha256'] or digest(input['file'])!=input['fileSha256']
        or source['sourceSha256']!=input['sourceSha256'] or plan(source['source'])!=source):
        raise ValueError('final source/import verification failed')
    rows=sorted([[r['iccid'],r['cutoff'],r['activation']] for r in input['records']])
    verify_output(Path(input['file']),rows)
    expected={r['phone']:r for r in input['records']}
    files=sorted(directory.glob('renewal_readback_*.json'),key=lambda p:int(p.stem.rsplit('_',1)[1]))
    if len(files)<2:
        raise ValueError('two distinct readbacks required')
    for file in files[-2:]:
        actual=json.loads(file.read_text())
        if len(actual)!=len(expected) or len({r['phone'] for r in actual})!=len(expected):
            raise ValueError('readback coverage mismatch')
        for row in actual:
            r=expected.get(row['phone'])
            if not r or any(row[k]!=r[v] for k,v in (('id','backendId'),('iccid','iccid'),('slot','slot'),('activation','activation'),('cutoff','cutoff'))):
                raise ValueError('independent saved backend readback mismatch')
    cache=directory/'f_effective_dates_readback.json'
    reviewed=load_frozen_date_snapshot(cache,digest(cache),input['runId'])
    with sqlite3.connect(DB.resolve().as_uri()+'?mode=ro',uri=True) as conn:
        renewals=read_renewals(conn)
        for r in input['records']:
            if renewals.get(r['iccid'])!=r['renewal']:
                raise ValueError('current renewal library readback mismatch')
            effective=reviewed['index'].get(r['iccid'])
            if not effective or effective['截止日'].strftime('%Y-%m-%d')!=r['cutoff'] or effective['激活日期'].strftime('%Y-%m-%d')!=r['activation']:
                raise ValueError('effective report dates mismatch')
    result={'status':'passed','completed':True,'runId':input['runId'],'sourceSha256':input['sourceSha256'],
            'eligibleRows':source['eligibleRows'],'updatedLocalLibrary':committed['updatedCount'],
            'importFileRows':input['fileRows'],'importRequestsSent':receipt['importRequestsSent'],
            'acceptedRows':receipt.get('acceptedRows'),'verifiedDatePairs':input['fileRows'],
            'extensionRows':source['extensionRows'],'directDateRows':source['directDateRows'],
            'excludedMissingSourceCutoff':sum(r['reason']=='missing source cutoff' for r in source['excluded']),
            'originalActivationDatesPreserved':True,'originalFQueryTablesUnchanged':committed['originalTablesUnchanged'],
            'dateFormat':'yyyy-mm-dd','independentReadbackPassed':True,
            'priorImportResolution':receipt.get('priorResolution'),'output':input['file'],
            'sha256':input['fileSha256'],'localDatabaseBackup':committed['backup'],
            'sourceVerification':'Sheet!A1:D774','backendReceipt':str(directory/'backend_import_receipt.json'),
            'effectiveFrozenDatesSha256':digest(cache),'readbacks':[str(p) for p in files[-2:]],
            'finishedAt':datetime.now().astimezone().isoformat()}
    save(directory/'completion.json',result)
    return result

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['plan','commit','export','finalize'])
    parser.add_argument('--source', type=Path)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--directory',type=Path)
    parser.add_argument('--mapping',type=Path)
    parser.add_argument('--run-id')
    parser.add_argument('--authorization-source')
    args = parser.parse_args()
    if args.command=='plan':
        result=plan(args.source); save(args.output,result)
    else:
        if args.authorization_source!='current-user':
            raise ValueError('current human authorization required')
        if args.command=='commit':
            result=commit_renewals(json.loads((args.directory/'source_plan.json').read_text()),json.loads(args.mapping.read_text()),args.directory,args.run_id)
        elif args.command=='export':
            result=export_renewals(args.directory)
        else:
            result=finalize_renewals(args.directory)
    print(json.dumps({k:v for k,v in result.items() if k not in {'rows','records','excluded','source'}}, ensure_ascii=False))

if __name__ == '__main__':
    main()
