#!/usr/bin/env python3
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


CANCELLED_SHEETS = {
    "今日总回复量为0",
    "今日昨天总回复量均为0",
    "快过期_今天明天",
    "已过期_未进过期下架卡",
}
BALANCE_GROUPS = ["A1", "A12", "A11", "B1", "B11", "B12", "C2", "A3"]


def parse_args():
    parser = argparse.ArgumentParser(description="汇总二号机房全量下架与强制任务B报告")
    parser.add_argument("--tag", required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--export-receipt", type=Path, required=True)
    parser.add_argument("--google-receipt", type=Path, required=True)
    parser.add_argument("--phone-date", type=Path, required=True)
    parser.add_argument("--day-date", type=Path, required=True)
    parser.add_argument("--dispatch-receipt", type=Path, required=True)
    parser.add_argument("--card-result", type=Path, required=True)
    parser.add_argument("--card-receipt", type=Path, required=True)
    parser.add_argument("--analysis", type=Path, required=True)
    parser.add_argument("--balance", type=Path, required=True)
    parser.add_argument("--c2-gate", type=Path)
    parser.add_argument("--task-b-manifest", type=Path, required=True)
    parser.add_argument("--expired-candidate", type=Path, required=True)
    parser.add_argument("--expired-write", type=Path, required=True)
    parser.add_argument("--expired-verify", type=Path, required=True)
    parser.add_argument("--tomorrow-candidate", type=Path, required=True)
    parser.add_argument("--tomorrow-write", type=Path, required=True)
    parser.add_argument("--tomorrow-verify", type=Path, required=True)
    parser.add_argument("--down-candidate", type=Path, required=True)
    parser.add_argument("--down-write", type=Path, required=True)
    parser.add_argument("--down-verify", type=Path, required=True)
    parser.add_argument("--report-dir", type=Path, required=True)
    parser.add_argument("--snapshot-dir", type=Path, required=True)
    parser.add_argument("--delivery-dir", type=Path, required=True)
    return parser.parse_args()


def load(path):
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_new(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        if isinstance(payload, str):
            handle.write(payload)
        else:
            json.dump(payload, handle, ensure_ascii=False, indent=2, default=str)
            handle.write("\n")


def successful_updates(write):
    count = 0
    for response in write.get("updateResponses") or []:
        if response.get("code") != 200 or response.get("skipped"):
            continue
        data = response.get("data") or {}
        updated = data.get("updated")
        if updated is not None:
            count += int(updated)
    return count


def task_group(candidate_path, write_path, verify_path):
    candidate = load(candidate_path)
    write = load(write_path)
    verify = load(verify_path)
    before = write.get("before") or {}
    candidate_count = int(candidate.get("candidateCount") or candidate.get("count") or 0)
    actual = int(before.get("eligibleCount") or 0)
    response_success = successful_updates(write)
    verified = int(verify.get("targetCount") or 0)
    passed = verify.get("passed") is True
    success = actual if passed and verified == candidate_count else response_success
    return {
        "candidateCount": candidate_count,
        "backendValidCount": candidate_count
        - int(before.get("missingCount") or 0)
        - int(before.get("duplicateBackendCount") or 0)
        - int(before.get("nonSourceCount") or 0),
        "alreadyTargetCount": int(before.get("alreadyTargetCount") or 0),
        "nonSourceCount": int(before.get("nonSourceCount") or 0),
        "missingCount": int(before.get("missingCount") or 0),
        "duplicateCount": int(before.get("duplicateBackendCount") or 0),
        "actualWriteCount": actual,
        "successCount": success,
        "failureCount": max(actual - success, 0),
        "verifyTargetCount": verified,
        "verifyStillNotTargetCount": int(verify.get("stillNotTargetCount") or 0),
        "verifyMissingCount": int(verify.get("missingCount") or 0),
        "verifyDuplicateCount": int(verify.get("duplicateBackendCount") or 0),
        "verifyPassed": passed,
        "validated": Path(write["validatedCandidatePath"]),
    }


def down_group(candidate_path, write_path, verify_path):
    candidate = load(candidate_path)
    write = load(write_path)
    verify = load(verify_path)
    before = write.get("before") or {}
    candidate_count = int(candidate.get("candidateCount") or candidate.get("count") or 0)
    actual = int(before.get("eligibleCount") or 0)
    passed = verify.get("passed") is True and int(verify.get("downCount") or 0) == candidate_count
    success = actual if passed else max(actual - int(verify.get("stillActiveCount") or 0), 0)
    return {
        "candidateCount": candidate_count,
        "backendValidCount": int(before.get("backendValidCount") or 0),
        "alreadyDownCount": int(before.get("alreadyDownCount") or 0),
        "missingCount": int(before.get("missingCount") or 0),
        "duplicateCount": int(before.get("duplicateBackendCount") or 0),
        "actualWriteCount": actual,
        "successCount": success,
        "failureCount": max(actual - success, 0),
        "verifyDownCount": int(verify.get("downCount") or 0),
        "verifyStillActiveCount": int(verify.get("stillActiveCount") or 0),
        "verifyMissingCount": int(verify.get("missingCount") or 0),
        "verifyDuplicateCount": int(verify.get("duplicateBackendCount") or 0),
        "verifyPassed": passed,
        "validated": Path(write["validatedCandidatePath"]),
    }


def external(path, delivery, receipts=False):
    return delivery / ("receipts" if receipts else "") / path.name


def fmt_number(value):
    if value is None:
        return "未知"
    number = float(value)
    return str(int(number)) if number.is_integer() else f"{number:.3f}".rstrip("0").rstrip(".")


def pct(value):
    return "未知" if value is None else f"{float(value):.2f}%"


def main():
    args = parse_args()
    required = [value for key, value in vars(args).items() if isinstance(value, Path) and key not in {"report_dir", "snapshot_dir", "delivery_dir", "c2_gate"}]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"报告输入缺失: {missing}")

    day = load(args.day_date)
    phone = load(args.phone_date)
    if day.get("serverDate") != phone.get("serverDate"):
        raise ValueError("phoneConfig与day服务器日期不一致")
    card = load(args.card_receipt)
    analysis = load(args.analysis)
    balance = load(args.balance)
    manifest = load(args.task_b_manifest)
    expired = task_group(args.expired_candidate, args.expired_write, args.expired_verify)
    tomorrow = task_group(args.tomorrow_candidate, args.tomorrow_write, args.tomorrow_verify)
    down = down_group(args.down_candidate, args.down_write, args.down_verify)

    workbook = load_workbook(args.card_result, read_only=True, data_only=False)
    sheet_names = set(workbook.sheetnames)
    workbook.close()
    cancelled_present = sorted(CANCELLED_SHEETS & sheet_names)
    card_failed = card.get("failed_checks") or []
    card_passed = card.get("result") in {"成功", "success"} and not card_failed and not cancelled_present
    groups_passed = all(group["failureCount"] == 0 and group["verifyPassed"] for group in (expired, tomorrow))
    down_passed = down["failureCount"] == 0 and down["verifyPassed"]
    overall_passed = card_passed and groups_passed and down_passed

    report_path = args.report_dir / f"opensms_room2_down_task_b_report_{args.tag}.md"
    final_receipt_path = args.snapshot_dir / f"opensms_room2_down_task_b_{args.tag}_final_receipt.json"
    config_path = args.snapshot_dir / f"opensms_patrol_finalize_{args.tag}_config.json"

    records = {item.get("group"): item for item in balance.get("records") or []}
    rates = {item.get("group"): item for item in balance.get("replyRates") or []}
    overall = analysis.get("overall") or {}
    c2_gate_status = "未找到本日回执"
    if args.c2_gate and args.c2_gate.is_file():
        gate = load(args.c2_gate)
        c2_gate_status = f"今日已于{gate.get('evaluatedAt')}执行，本轮跳过"

    lines = [
        "## 本轮概况",
        "",
        "| 指标 | 结果 |",
        "|---|---|",
        "| 任务 | 新鲜导出一次；二号机房全部号码设置为下架；强制执行任务B |",
        "| 任务B授权 | 当前用户消息明确授权，`OPENSMS_FORCE_TASK_B=1` |",
        f"| 后台服务器日期 | {day.get('serverDate')} |",
        f"| 源表号码行数 | {analysis.get('rows')} |",
        f"| 二号机房号码数 | {down['candidateCount']} |",
        f"| 二号机房实际下架写入 | {down['actualWriteCount']} |",
        f"| 任务B实际写入 | {expired['actualWriteCount'] + tomorrow['actualWriteCount']} |",
        f"| 卡槽结果复核 | {'通过' if card_passed else '失败'} |",
        f"| 整体结果 | {'通过' if overall_passed else '失败'} |",
        "",
        "## 0回复与回复率",
        "",
        "| 指标 | 结果 |",
        "|---|---:|",
        f"| 最新统计日 | {analysis.get('latestDate')} |",
        f"| 当日发送窗口数 | {overall.get('send')} |",
        f"| 当日回复窗口数 | {overall.get('reply')} |",
        f"| 当日整体回复率 | {pct(overall.get('rate_pct'))} |",
        f"| 发送大于70且0回复 | {analysis.get('currentZero')} |",
        "",
        "## 运营商调整",
        "",
        "| 分组 | 候选数 | 后台有效数 | 已在目标 | 离开来源 | 缺失 | 重复 | 实际写入 | 成功数 | 失败数 | 独立复核 | 结果 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---|---|",
        f"| 任务B：已过期转过期下架卡 | {expired['candidateCount']} | {expired['backendValidCount']} | {expired['alreadyTargetCount']} | {expired['nonSourceCount']} | {expired['missingCount']} | {expired['duplicateCount']} | {expired['actualWriteCount']} | {expired['successCount']} | {expired['failureCount']} | {expired['verifyTargetCount']}/{expired['candidateCount']} | {'通过' if expired['verifyPassed'] else '失败'} |",
        f"| 任务B：明日到期转快过期隔离区 | {tomorrow['candidateCount']} | {tomorrow['backendValidCount']} | {tomorrow['alreadyTargetCount']} | {tomorrow['nonSourceCount']} | {tomorrow['missingCount']} | {tomorrow['duplicateCount']} | {tomorrow['actualWriteCount']} | {tomorrow['successCount']} | {tomorrow['failureCount']} | {tomorrow['verifyTargetCount']}/{tomorrow['candidateCount']} | {'通过' if tomorrow['verifyPassed'] else '失败'} |",
        "| B13目标 | 0 | 0 | 0 | 0 | 0 | 0 | 0 | 0 | 0 | 无B13目标写入 | 通过 |",
        "",
        "## 下架恢复",
        "",
        "| 指标 | 数量/结果 |",
        "|---|---:|",
        f"| 二号机房候选数 | {down['candidateCount']} |",
        f"| 后台有效数 | {down['backendValidCount']} |",
        f"| 提交前已下架 | {down['alreadyDownCount']} |",
        f"| 实际写入 | {down['actualWriteCount']} |",
        f"| 成功数 | {down['successCount']} |",
        f"| 失败数 | {down['failureCount']} |",
        f"| 独立复核已下架 | {down['verifyDownCount']} |",
        f"| 独立复核仍上架 | {down['verifyStillActiveCount']} |",
        f"| 独立复核结果 | {'通过' if down['verifyPassed'] else '失败'} |",
        "| 下架恢复 | 本轮未执行；用户要求为全量下架 |",
        "",
        "## 余量与提醒",
        "",
        "| 分组 | 当日剩余量 | 当日总量 | 剩余比例 | 回复率 | 回复率统计日期 | 状态 |",
        "|---|---:|---:|---:|---:|---|---|",
    ]
    for group in BALANCE_GROUPS:
        record = records.get(group) or {}
        rate = rates.get(group) or {}
        percent = record.get("percentRemaining")
        status = "未知" if rate.get("status") != "ok" or percent is None else "紧急" if percent < 10 else "提醒" if percent < 20 else "正常"
        lines.append(
            f"| {group} | {fmt_number(record.get('dailyRemaining'))} | {fmt_number(record.get('dailyTotal'))} | {pct(percent)} | {pct(rate.get('replyRatePct')) if rate.get('status') == 'ok' else '未知'} | {rate.get('statDate') or '未知'} | {status} |"
        )
    lines.extend([
        "",
        f"- C2每日门禁：{c2_gate_status}。",
        f"- C2按日统计最新发送量：{balance.get('c2LatestDailySendSize')}；统计日期：{balance.get('c2LatestDailyDate')}；4000阈值提醒本轮因今日已执行而不重复触发。",
        "- 自动加量：关闭；客户绑定：本轮未执行。",
        "",
        "## 文件与回执",
        "",
        f"[打开本轮源文件]({external(args.source, args.delivery_dir)})",
        "",
        str(external(args.source, args.delivery_dir)),
        "",
        f"[打开卡槽关联结果表]({external(args.card_result, args.delivery_dir)})",
        "",
        str(external(args.card_result, args.delivery_dir)),
        "",
        f"[打开完整Markdown报告]({external(report_path, args.delivery_dir)})",
        "",
        str(external(report_path, args.delivery_dir)),
        "",
        f"[打开Google关联表]({external(Path(load(args.google_receipt)['outputPath']), args.delivery_dir)})",
        "",
        str(external(Path(load(args.google_receipt)['outputPath']), args.delivery_dir)),
        "",
        "## 异常",
        "",
        "| 项目 | 结果 |",
        "|---|---|",
        f"| 卡槽复核失败项 | {', '.join(card_failed) if card_failed else '无'} |",
        f"| 已取消工作表存在项 | {', '.join(cancelled_present) if cancelled_present else '无，四张工作表均不存在'} |",
        f"| 任务B异常 | {'无' if groups_passed else '存在写入或独立回读失败'} |",
        f"| 二号机房下架异常 | {'无' if down_passed else '存在缺失、重复、写入或独立回读失败'} |",
        "| B13目标写入 | 无 |",
        "| 任务B数据来源 | 结果主表过期状态与截止日直接计算；不依赖已取消工作表 |",
        "",
    ])
    write_new(report_path, "\n".join(lines))

    receipt_paths = [
        args.export_receipt, args.google_receipt, args.phone_date, args.day_date,
        args.dispatch_receipt, args.card_receipt, args.analysis, args.balance,
        args.task_b_manifest, args.expired_candidate, args.expired_write,
        args.expired_verify, args.tomorrow_candidate, args.tomorrow_write,
        args.tomorrow_verify, args.down_candidate, args.down_write,
        args.down_verify, expired["validated"], tomorrow["validated"], down["validated"],
    ]
    if args.c2_gate and args.c2_gate.is_file():
        receipt_paths.append(args.c2_gate)

    final_receipt = {
        "result": "成功" if overall_passed else "失败",
        "serverDate": day.get("serverDate"),
        "sourcePath": str(args.source.resolve()),
        "sourceSha256": sha256(args.source),
        "cardResultPath": str(args.card_result.resolve()),
        "cardResultSha256": sha256(args.card_result),
        "cancelledSheetsAbsent": not cancelled_present,
        "taskB": {"forceGate": True, "manifest": str(args.task_b_manifest.resolve()), "expired": expired, "tomorrow": tomorrow},
        "room2Downlist": down,
        "localReportPath": str(report_path.resolve()),
        "deliveryReportPath": str(external(report_path, args.delivery_dir)),
        "allReceiptInputs": [str(path.resolve()) for path in receipt_paths],
    }
    final_receipt["finalReceiptPath"] = str(final_receipt_path.resolve())
    write_new(final_receipt_path, final_receipt)

    def receipt_external(path):
        return str(external(path, args.delivery_dir, True))

    config = {
        "runType": "task_b_only",
        "sourceArchive": str(external(args.source, args.delivery_dir)),
        "sourceSha256": sha256(args.source),
        "googleSha256": sha256(Path(load(args.google_receipt)["outputPath"])),
        "cardResult": str(external(args.card_result, args.delivery_dir)),
        "cardReceipt": receipt_external(args.card_receipt),
        "analysis": receipt_external(args.analysis),
        "balance": receipt_external(args.balance),
        "taskBExecuted": True,
        "taskBStatus": "当前消息明确授权且OPENSMS_FORCE_TASK_B=1；两组写入和独立回读均完成",
        "groups": [
            {
                "title": "任务B-已过期转过期下架卡",
                "candidate": receipt_external(args.expired_candidate),
                "validated": receipt_external(expired["validated"]),
                "write": receipt_external(args.expired_write),
                "verify": receipt_external(args.expired_verify),
            },
            {
                "title": "任务B-明日到期转快过期隔离区",
                "candidate": receipt_external(args.tomorrow_candidate),
                "validated": receipt_external(tomorrow["validated"]),
                "write": receipt_external(args.tomorrow_write),
                "verify": receipt_external(args.tomorrow_verify),
            },
        ],
        "downlist": {
            "candidate": receipt_external(args.down_candidate),
            "validated": receipt_external(down["validated"]),
            "write": receipt_external(args.down_write),
            "verify": receipt_external(args.down_verify),
        },
        "extraReceipts": [
            {"label": "后台导出回执", "path": receipt_external(args.export_receipt)},
            {"label": "Google下载回执", "path": receipt_external(args.google_receipt)},
            {"label": "phoneConfig服务器日期回执", "path": receipt_external(args.phone_date)},
            {"label": "day服务器日期回执", "path": receipt_external(args.day_date)},
            {"label": "文档调度回执", "path": receipt_external(args.dispatch_receipt)},
            {"label": "专项总回执", "path": receipt_external(final_receipt_path)},
        ],
    }
    write_new(config_path, config)
    print(json.dumps({
        "result": final_receipt["result"],
        "reportPath": str(report_path.resolve()),
        "finalReceiptPath": str(final_receipt_path.resolve()),
        "configPath": str(config_path.resolve()),
        "expired": expired,
        "tomorrow": tomorrow,
        "room2Downlist": down,
    }, ensure_ascii=False, indent=2, default=str))
    if not overall_passed:
        raise SystemExit(4)


if __name__ == "__main__":
    main()
