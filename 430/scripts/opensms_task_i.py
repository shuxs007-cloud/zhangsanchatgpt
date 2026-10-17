#!/usr/bin/env python3
"""Task I: offline, incremental document identity index and daily SMS statistics.

No network, no business DB, no edits to originals or Task F. SQLite belongs only
to Task I. Counts use exact text ICCIDs / normalized North-American phone keys.
"""
from __future__ import annotations
import argparse, csv, fcntl, hashlib, io, json, os, posixpath, re, sqlite3, sys
from html import unescape
import subprocess, time, zipfile
from collections import Counter, defaultdict, deque
from concurrent.futures import ProcessPoolExecutor
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from lxml import etree as E

DEFAULT_DB=Path('/Volumes/zs01/chatgpt/OpenSMS/task_i/document_index.sqlite3')
CPU_CORES=os.cpu_count() or 1
DEFAULT_WORKERS=7
DEFAULT_QUERY_WORKERS=7
WORKER_PRIORITY=None
DEFAULT_ROOTS=['/Users/zhangsan/Downloads','/Users/zhangsan/Documents','/Users/zhangsan/Desktop','/Volumes/zs01']
# This migrated sync directory has repeatedly stalled even ZIP-directory reads.
# It is explicitly accounted for as unindexed, never reported as no matches.
UNAVAILABLE_PREFIX='/Users/zhangsan/Documents/文稿 - 张三的MacBook Air - 1/'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
ROW=re.compile(rb'<(?:[\w.-]+:)?row\b[^>]*(?:/>|>.*?</(?:[\w.-]+:)?row\s*>)',re.S)
CELL=re.compile(rb'<(?:[\w.-]+:)?c\b([^>]*)(?:/>|>(.*?)</(?:[\w.-]+:)?c\s*>)',re.S)
REF=re.compile(rb'\br=["\']([A-Z]+)([0-9]+)["\']')
KIND=re.compile(rb'\bt=["\']([^"\']+)["\']')
VALUE=re.compile(rb'<(?:[\w.-]+:)?v(?:\s[^>]*)?>(.*?)</(?:[\w.-]+:)?v>',re.S)
TEXT=re.compile(rb'<(?:[\w.-]+:)?t(?:\s[^>]*)?>(.*?)</(?:[\w.-]+:)?t>',re.S)
FORMULA=re.compile(rb'<(?:[\w.-]+:)?f(?:\s|>)')
EXCLUDED={'node_modules','.git','.venv','__pycache__','tests','fixtures','.Trash','.Trashes','task_i'}
IC_HEADERS={'iccid','系统iccid','系統iccid','卡存储','卡存儲','simnumber','sim卡号','sim卡號','queryiccid','sourceiccid'}
PH_HEADERS={'电话','電話','手机号','手機號','电话号码','電話號碼','手机号码','電話號碼','phone','phoneno','phonenumber','phonechat','recipient','from','to','destination'}
METRICS=('发送窗口数','回复窗口数','总发送量','总回复量')

def configure_worker_priority():
    """Raise only this Task I worker, with no sudo or system-wide changes.

    On macOS, USER_INTERACTIVE is the highest normal application thread QoS.
    Negative process nice may require privileges; record denial, never claim it
    succeeded. This is not real-time scheduling and not a CPU-usage guarantee.
    """
    global WORKER_PRIORITY
    result=dict(requested='highest-normal-application-priority',threadQos='unavailable')
    if sys.platform=='darwin':
        try:
            import ctypes
            system=ctypes.CDLL('/usr/lib/libSystem.B.dylib',use_errno=True)
            system.pthread_set_qos_class_self_np.argtypes=[ctypes.c_uint,ctypes.c_int]
            system.pthread_set_qos_class_self_np.restype=ctypes.c_int
            code=system.pthread_set_qos_class_self_np(0x21,0)
            result['threadQos']='USER_INTERACTIVE' if code==0 else 'not-applied'
            result['threadQosReturnCode']=code
        except (OSError,AttributeError) as e:result['threadQosError']=str(e)
    try:
        os.setpriority(os.PRIO_PROCESS,0,-20)
        result['processNiceRequested']=-20
    except (PermissionError,OSError) as e:
        result['processNiceRequested']=-20;result['processNiceError']=str(e)
    try:result['processNiceActual']=os.getpriority(os.PRIO_PROCESS,0)
    except (AttributeError,OSError):result['processNiceActual']=None
    WORKER_PRIORITY=result
    return result

def iccid_key(value):
    if not isinstance(value,str):return None
    v=value.strip()
    if v[-1:] in ('f','F'):v=v[:-1]
    return v if re.fullmatch(r'[0-9]{18,22}',v) else None

def phone_key(value):
    if value is None:return None
    s=str(value).strip()
    if not re.fullmatch(r'[+0-9 ()-]+',s):return None
    s=re.sub(r'[^0-9]','',s)
    if len(s)==11 and s[0]=='1':s=s[1:]
    return s if len(s)==10 else None

def header(value):return re.sub(r'\s+','',str(value or '')).lower()

def observed_hint(path):
    for part in reversed(path.parts):
        m=re.search(r'(20\d{6})T(\d{6})(\+0800|Z)?',part)
        if m:
            dt=datetime.strptime(m[1]+m[2],'%Y%m%d%H%M%S')
            if m[3]=='Z':dt+=timedelta(hours=8)
            return dt.isoformat(),'文件名或归档运行编号'
        m=re.search(r'(20\d{2})[-_]?([01]\d)[-_]?([0-3]\d)(?!\d)',part)
        if m:
            try:return datetime(*map(int,m.groups())).isoformat(),'文件名或归档日期'
            except ValueError:pass
    return datetime.fromtimestamp(path.stat().st_mtime).isoformat(),'修改时间，非确定采集时间'

def discover(roots,omit=()):
    proc=subprocess.run(['rg','--files','--hidden','--no-ignore',*roots,
        '-g','*.xlsx','-g','*.xlsm','-g','*.xls','-g','*.csv','-g','*.tsv',
        '-g','!._*','-g','!~$*','-g','!**/node_modules/**','-g','!**/.git/**','-g','!**/.venv/**'],
        capture_output=True,text=True)
    if proc.returncode not in (0,1):raise RuntimeError(proc.stderr)
    files=[];seen=set();excluded=Counter()
    for s in proc.stdout.splitlines():
        p=Path(s)
        if EXCLUDED & set(p.parts) or any(x in p.parents for x in omit):
            excluded['开发、测试、任务I输出']+=1;continue
        try:
            resolved=str(p.resolve());st=p.stat()
            if resolved in seen:excluded['符号链接或同一真实路径']+=1;continue
            seen.add(resolved);observed,basis=observed_hint(p)
            files.append(dict(path=s,size=st.st_size,mtime=st.st_mtime_ns,observed=observed,dateBasis=basis))
        except OSError as e:files.append(dict(path=s,error=str(e),size=0,mtime=0,observed='',dateBasis=''))
    return files,dict(roots=roots,discovered=len(proc.stdout.splitlines()),excluded=dict(excluded))

def rich_text(e):return ''.join(t.text or '' for t in e.iter() if E.QName(t).localname=='t')

def read_cells(fragment,strings,namespaces):
    decl=' '.join('xmlns'+(':'+k if k else '')+'="'+v+'"' for k,v in namespaces.items())
    row=E.fromstring(('<wrapper '+decl+'>').encode()+fragment+b'</wrapper>',E.XMLParser(resolve_entities=False,no_network=True))[0]
    values={}
    for c in row:
        if E.QName(c).localname!='c':continue
        ref=c.get('r','');col=re.sub(r'[0-9]','',ref);kind=c.get('t','n')
        f=c.find('{'+NS+'}f');v=c.find('{'+NS+'}v');raw=v.text or '' if v is not None else ''
        value=strings[int(raw)] if kind=='s' and raw else rich_text(c) if kind=='inlineStr' else raw
        values[col]=(value,kind,f is not None,ref)
    return int(row.get('r','0')),values

def read_selected_cells(fragment,strings,selected):
    """Decode only needed columns, without building a DOM for every data row."""
    row_id=re.search(rb'\br=["\']([0-9]+)["\']',fragment.split(b'>',1)[0])
    values={}
    for c in CELL.finditer(fragment):
        ref=REF.search(c[1])
        if not ref:continue
        col=ref[1].decode()
        if col not in selected:continue
        kind=KIND.search(c[1]);kind=kind[1].decode() if kind else 'n'
        body=c[2] or b'';raw=VALUE.search(body)
        raw=unescape(raw[1].decode()) if raw else ''
        value=strings[int(raw)] if kind=='s' and raw else unescape(b''.join(TEXT.findall(body)).decode()) if kind=='inlineStr' else raw
        values[col]=(value,kind,bool(FORMULA.search(body)),ref[1].decode()+ref[2].decode())
    return int(row_id[1]) if row_id else 0,values

def peek(z,member):
    with z.open(member) as f:
        b=b''
        while len(b)<2_000_000:
            a=f.read(65536)
            if not a:break
            b+=a
            if sum(1 for _ in ROW.finditer(b))>=15:break
    opening=re.search(rb'<((?:[\w.-]+:)?worksheet)\b[^>]*>',b)
    if not opening:raise ValueError('工作表XML根节点无法识别')
    root=E.fromstring(opening.group()+b'</'+opening[1]+b'>',E.XMLParser(resolve_entities=False,no_network=True))
    return [m.group() for m in ROW.finditer(b)][:15],root.nsmap

def metric_columns(values,observed):
    groups=defaultdict(dict);dt=datetime.fromisoformat(observed)
    for col,v in values.items():
        m=re.fullmatch(r'(?:(20\d{2})[-/])?(\d{1,2})[-/](\d{1,2})(发送窗口数|回复窗口数|总发送量|总回复量)',header(v[0]))
        if not m:continue
        y=int(m[1]) if m[1] else dt.year
        try:d=date(y,int(m[2]),int(m[3]))
        except ValueError:continue
        if not m[1] and d>dt.date()+timedelta(days=1):d=date(y-1,d.month,d.day)
        groups[d.isoformat()][m[4]]=col
    return {d:g for d,g in groups.items() if set(g)==set(METRICS)}

def count_cell(v):
    s,kind,formula,ref=v
    if formula or not s:return None
    try:
        n=float(str(s).replace(',',''))
        return int(n) if n>=0 and n.is_integer() else None
    except ValueError:return None

def identify_header(fr,strings,ns,observed):
    for fragment in fr:
        n,v=read_cells(fragment,strings,ns)
        ic=[c for c,x in v.items() if header(x[0]) in IC_HEADERS]
        ph=[c for c,x in v.items() if header(x[0]) in PH_HEADERS]
        if ic or ph:return n,v,ic,ph,metric_columns(v,observed)
    return None

def row_records(rownum,v,h,ic,ph,metrics):
    # Only an explicitly named column creates an identifier index.
    ics=[(iccid_key(v.get(c,('',))[0]),v.get(c,('','n',False,'')),h[c][0]) for c in ic]
    reliable=[(k,c,label) for k,c,label in ics if k and c[1] in ('s','str','inlineStr') and not c[2]]
    phones=[(phone_key(v.get(c,('',))[0]),v.get(c,('',))[0],v.get(c,('','','',c))[3]) for c in ph]
    phones=[x for x in phones if x[0]]
    if not reliable and not phones:return []
    if not reliable and not metrics and len({x[0] for x in phones})>1:
        return [[rownum,'',key,raw,'',ref,'','','','', '{}','',''] for key,raw,ref in phones]
    phone,phone_raw,phone_cell=phones[0] if len({x[0] for x in phones})==1 else ('','','')
    meta={h[c][0]:v.get(c,('',))[0] for c in h if h[c][0] in ('卡槽','运营商','opr','创建时间')}
    daily={}
    for d,m in metrics.items():daily[d]=[count_cell(v.get(m[x],('','n',False,''))) for x in METRICS]
    ambiguous=len({k for k,c,label in reliable})>1
    issue='同条多个ICCID不一致，统计排除' if ambiguous else '多个手机号不一致' if len({x[0] for x in phones})>1 else ''
    rows=[]
    for k,c,label in reliable or [('',('','','',''),'')]:
        rows.append([rownum,k or '',phone or '',phone_raw or '',c[3],phone_cell,label,
                     meta.get('卡槽',''),meta.get('运营商',''),meta.get('创建时间',''),
                     json.dumps(daily,ensure_ascii=False,separators=(',',':')),max(metrics) if metrics else '',issue])
    if not reliable:
        for k,c,label in ics:
            if k:rows[-1][-1]='ICCID数值型或公式型，精度不可靠'
    return rows

def scan_file(item):
    result=dict(item,status='已检查',error='',sheets=[])
    p=Path(item['path'])
    try:
        if str(p).startswith(UNAVAILABLE_PREFIX):
            result.update(sha256='unavailable:'+str(p),status='旧同步目录读取阻塞，未索引');return result
        result['sha256']=content_key(p)
        if not zipfile.is_zipfile(p):
            if p.suffix.lower() not in ('.csv','.tsv'):
                result['status']='不支持的旧二进制表格';return result
            raw=p.read_bytes()
            encodings=('utf-16','utf-8-sig','gb18030') if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','gb18030')
            for encoding in encodings:
                try:txt=raw.decode(encoding);break
                except UnicodeDecodeError:continue
            else:raise ValueError('文本编码无法识别')
            delimiter='\t' if p.suffix.lower()=='.tsv' else ','
            csv.field_size_limit(16*1024*1024)
            reader=csv.reader(io.StringIO(txt),delimiter=delimiter)
            rows=list(reader)
            selected=None
            for i,row in enumerate(rows[:15],1):
                h={str(j):(s,'str',False,str(j)) for j,s in enumerate(row)}
                ic=[c for c,x in h.items() if header(x[0]) in IC_HEADERS];ph=[c for c,x in h.items() if header(x[0]) in PH_HEADERS]
                if ic or ph:selected=(i,h,ic,ph,metric_columns(h,item['observed']));break
            if selected:
                hn,h,ic,ph,metrics=selected;records=[]
                for n,row in enumerate(rows[hn:],hn+1):
                    v={str(j):(s,'str',False,str(j)) for j,s in enumerate(row)}
                    records.extend(row_records(n,v,h,ic,ph,metrics))
                result['sheets'].append(dict(name='文本表',sha256=result['sha256'],records=records,status='已索引',rowCount=len(rows),headers={c:x[0] for c,x in h.items()}))
            return result
        with zipfile.ZipFile(p) as z:
            if 'xl/workbook.xml' not in z.namelist():
                result['status']='非OOXML工作簿，未索引';return result
            raw=z.read('xl/sharedStrings.xml') if 'xl/sharedStrings.xml' in z.namelist() else b''
            strings=[rich_text(si) for si in E.fromstring(raw)] if raw else []
            wb=E.fromstring(z.read('xl/workbook.xml'));rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'))
            members={r.get('Id'):r.get('Target').lstrip('/') if r.get('Target','').startswith('/') else posixpath.normpath('xl/'+r.get('Target','')) for r in rels if r.get('TargetMode')!='External'}
            for s in wb.findall('.//{'+NS+'}sheet'):
                name=s.get('name');member=members[s.get(RID)]
                fr,ns=peek(z,member);selected=identify_header(fr,strings,ns,item['observed'])
                if not selected:
                    result['sheets'].append(dict(name=name,status='无明确ICCID或手机号列',records=[]));continue
                hn,h,ic,ph,metrics=selected;xml=z.read(member)
                shash=hashlib.sha256(raw+xml).hexdigest()
                selected_cols=set(ic+ph)|{c for c,v in h.items() if v[0] in ('卡槽','运营商','opr','创建时间')}|{c for d in metrics.values() for c in d.values()}
                records=[];count=0
                for rm in ROW.finditer(xml):
                    count+=1
                    n,v=read_selected_cells(rm.group(),strings,selected_cols)
                    if n<=hn:continue
                    records.extend(row_records(n,v,h,ic,ph,metrics))
                result['sheets'].append(dict(name=name,sha256=shash,records=records,status='已索引',rowCount=count,headers={c:x[0] for c,x in h.items()}))
        st=p.stat()
        if (st.st_size,st.st_mtime_ns)!=(item['size'],item['mtime']):raise ValueError('扫描期间源文档发生变化')
    except Exception as e:result.update(status='读取失败',error=str(e))
    return result

def content_key(path):
    """OOXML member CRC/size index, not a claim of byte-for-byte file SHA256.

    Actual worksheet content uses SHA256 after parsing. This first-stage ZIP
    identity avoids reading large irrelevant bodies and ignores ZIP timestamps.
    """
    if str(path).startswith(UNAVAILABLE_PREFIX):return 'unavailable:'+str(path)
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            signature=sorted((x.filename,x.CRC,x.file_size) for x in z.infolist())
        return hashlib.sha256(json.dumps(signature,separators=(',',':')).encode()).hexdigest()
    if path.suffix.lower() in ('.csv','.tsv'):
        with path.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
    return 'unsupported:'+str(path)

SCHEMA='''
CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY,value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS contents(sha256 TEXT PRIMARY KEY,status TEXT,error TEXT,indexed_at TEXT);
CREATE TABLE IF NOT EXISTS files(id INTEGER PRIMARY KEY,path TEXT UNIQUE,sha256 TEXT,size INTEGER,mtime INTEGER,observed TEXT,date_basis TEXT,active INTEGER DEFAULT 1,status TEXT,error TEXT);
CREATE TABLE IF NOT EXISTS sheets(id INTEGER PRIMARY KEY,sha256 TEXT UNIQUE,headers TEXT,row_count INTEGER);
CREATE TABLE IF NOT EXISTS file_sheets(file_id INTEGER,sheet_id INTEGER,sheet_name TEXT,UNIQUE(file_id,sheet_name));
CREATE TABLE IF NOT EXISTS records(id INTEGER PRIMARY KEY,sheet_id INTEGER,row_number INTEGER,iccid TEXT,phone TEXT,phone_raw TEXT,iccid_cell TEXT,phone_cell TEXT,iccid_label TEXT,slot TEXT,carrier TEXT,created_raw TEXT,daily_json TEXT,latest_day TEXT,issue TEXT,UNIQUE(sheet_id,row_number,iccid,phone));
CREATE INDEX IF NOT EXISTS file_sheets_sheet ON file_sheets(sheet_id,file_id);
CREATE INDEX IF NOT EXISTS files_content ON files(sha256);
'''

def bounded_scan(pool,items,workers):
    """Bound pending parse results so a slow DB writer cannot exhaust RAM."""
    source=iter(items);pending=deque()
    for _ in range(workers*2):
        try:pending.append(pool.submit(scan_file,next(source)))
        except StopIteration:break
    while pending:
        yield pending.popleft().result()
        try:pending.append(pool.submit(scan_file,next(source)))
        except StopIteration:pass

def validate_storage(db_path):
    mounts=subprocess.run(['mount'],capture_output=True,text=True,check=True).stdout
    if not any(' on /Volumes/zs01 (' in line for line in mounts.splitlines()):
        raise ValueError('zs01未正确挂载，任务I停止；不在本地建立替代大库')
    if not str(db_path.resolve()).startswith('/Volumes/zs01/'):
        raise ValueError('任务I索引固定保存在zs01，禁止大库回退本地')

def index(args):
    started=time.monotonic();validate_storage(args.db)
    if not args.db.parent.exists():args.db.parent.mkdir(parents=True)
    with (args.db.parent/'index.lock').open('a') as lock:
        try:fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError:raise RuntimeError('任务I索引正在构建，不启动第二份')
        db=sqlite3.connect(args.db);db.executescript(SCHEMA)
        db.execute('PRAGMA journal_mode=DELETE');db.execute('PRAGMA synchronous=FULL')
        meta=dict(db.execute('SELECT key,value FROM meta'))
        existing={r[0]:r for r in db.execute('SELECT path,size,mtime,sha256,status FROM files')}
        if args.manifest:
            doc=json.loads(args.manifest.read_text());files=doc['files'];discovery=doc['discovery']
        else:files,discovery=discover(args.root or DEFAULT_ROOTS,[Path(args.omit)] if args.omit else [])
        changed=[f for f in files if f['path'] not in existing or (f.get('size'),f.get('mtime'))!=existing[f['path']][1:3] or existing[f['path']][4]=='读取失败']
        known={x[0] for x in db.execute("SELECT sha256 FROM contents WHERE status<>'读取失败'")}
        now=datetime.now().isoformat();before_records=db.execute('SELECT COUNT(*) FROM records').fetchone()[0]
        before_max_id=db.execute('SELECT COALESCE(MAX(id),0) FROM records').fetchone()[0]
        active_before=set(r[0] for r in db.execute('SELECT path FROM files WHERE active=1'))
        coverage_changed=active_before!={f['path'] for f in files}
        db.execute('UPDATE files SET active=0')
        for f in files:db.execute('UPDATE files SET active=1,observed=?,date_basis=? WHERE path=?',(f.get('observed',''),f.get('dateBasis',''),f['path']))
        db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)',('roots',json.dumps(discovery,ensure_ascii=False)))
        db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)',('index_state','building'))
        db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)',('content_key_definition','OOXML: SHA256 of sorted ZIP member CRC32 and uncompressed size manifest; sheets: SHA256 of XML and shared strings. CSV/TSV: file SHA256.'))
        if len(changed)>500:
            db.execute('DROP INDEX IF EXISTS records_iccid')
            db.execute('DROP INDEX IF EXISTS records_phone')
        db.commit();t=time.monotonic()
        print(json.dumps(dict(event='task_i_index_started',documents=len(files),changed=len(changed),unchanged=len(files)-len(changed),workers=args.workers),ensure_ascii=False),flush=True)
        # Reuse any immutable file with already indexed identical content. Only
        # newly discovered/modified files are opened on subsequent refreshes.
        def persist(result):
            sha=result.get('sha256','');path=result['path'];status=result['status'];error=result.get('error','')
            db.execute('INSERT INTO files(path,sha256,size,mtime,observed,date_basis,active,status,error) VALUES(?,?,?,?,?,?,1,?,?) ON CONFLICT(path) DO UPDATE SET sha256=excluded.sha256,size=excluded.size,mtime=excluded.mtime,observed=excluded.observed,date_basis=excluded.date_basis,active=1,status=excluded.status,error=excluded.error',
                (path,sha,result.get('size',0),result.get('mtime',0),result.get('observed',''),result.get('dateBasis',''),status,error))
            file_id=db.execute('SELECT id FROM files WHERE path=?',(path,)).fetchone()[0]
            db.execute('DELETE FROM file_sheets WHERE file_id=?',(file_id,))
            for s in result.get('sheets',[]):
                if not s.get('sha256'):continue
                db.execute('INSERT OR IGNORE INTO sheets(sha256,headers,row_count) VALUES(?,?,?)',(s['sha256'],json.dumps(s['headers'],ensure_ascii=False),s['rowCount']))
                sheet_id=db.execute('SELECT id FROM sheets WHERE sha256=?',(s['sha256'],)).fetchone()[0]
                db.execute('INSERT OR REPLACE INTO file_sheets VALUES(?,?,?)',(file_id,sheet_id,s['name']))
                if not db.execute('SELECT 1 FROM records WHERE sheet_id=? LIMIT 1',(sheet_id,)).fetchone():
                    db.executemany('INSERT OR IGNORE INTO records(sheet_id,row_number,iccid,phone,phone_raw,iccid_cell,phone_cell,iccid_label,slot,carrier,created_raw,daily_json,latest_day,issue) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',([sheet_id,*r] for r in s['records']))
            db.execute('INSERT OR REPLACE INTO contents VALUES(?,?,?,?)',(sha,status,error,now))
        # Hashes saved by a previous offline scan are accepted only if current
        # size and mtime still match. They avoid re-reading identical copies.
        hash_map={}
        if args.scan_evidence:
            ev=json.loads(args.scan_evidence.read_text())
            for r in ev['results']:
                for p in r.get('copies',[r['path']]):hash_map[p]=(r.get('size'),r.get('mtime'),r.get('sha256'))
        groups=defaultdict(list)
        for f in changed:
            cached=hash_map.get(f['path'])
            if cached and cached[:2]==(f.get('size'),f.get('mtime')):digest=cached[2]
            else:
                try:digest=content_key(Path(f['path']))
                except OSError:digest=f['path']
            groups[digest].append(f)
        todo=[];copies={}
        for sha,g in groups.items():
            representative=max(g,key=lambda f:(f.get('dateBasis')!='修改时间，非确定采集时间','/chatgpt/OpenSMS/' in f['path'],f.get('observed','')))
            if sha in known:
                template=db.execute('SELECT id,status,error FROM files WHERE sha256=? LIMIT 1',(sha,)).fetchone()
                links=list(db.execute('SELECT sheet_id,sheet_name FROM file_sheets WHERE file_id=?',(template[0],)))
                for f in g:
                    persist(dict(f,sha256=sha,status=template[1],error=template[2],sheets=[]))
                    fid=db.execute('SELECT id FROM files WHERE path=?',(f['path'],)).fetchone()[0]
                    db.executemany('INSERT OR IGNORE INTO file_sheets VALUES(?,?,?)',((fid,*link) for link in links))
            else:todo.append(representative);copies[representative['path']]=g
        db.commit();print(json.dumps(dict(event='task_i_unique_contents',toScan=len(todo),duplicateCopies=sum(len(g)-1 for g in copies.values()),elapsed=round(time.monotonic()-t,1))),flush=True)
        stats=Counter()
        with ProcessPoolExecutor(max_workers=args.workers,initializer=configure_worker_priority) as pool:
            for n,result in enumerate(bounded_scan(pool,todo,args.workers),1):
                persist(result);stats[result['status']]+=1
                src_id=db.execute('SELECT id FROM files WHERE path=?',(result['path'],)).fetchone()[0]
                for f in copies[result['path']]:
                    if f['path']==result['path']:continue
                    persist(dict(f,sha256=result.get('sha256',''),status=result['status'],error=result.get('error',''),sheets=[]))
                    fid=db.execute('SELECT id FROM files WHERE path=?',(f['path'],)).fetchone()[0]
                    db.execute('INSERT OR IGNORE INTO file_sheets SELECT ?,sheet_id,sheet_name FROM file_sheets WHERE file_id=?',(fid,src_id))
                if n%20==0:
                    db.commit()
                    print(json.dumps(dict(event='task_i_index_progress',completed=n,total=len(todo),elapsed=round(time.monotonic()-t,1)),ensure_ascii=False),flush=True)
        db.commit()
        print(json.dumps(dict(event='task_i_building_query_indexes',elapsed=round(time.monotonic()-t,1))),flush=True)
        temporary=args.db.parent/'index_tmp';temporary.mkdir(exist_ok=True)
        os.environ['SQLITE_TMPDIR']=str(temporary)
        db.execute('PRAGMA temp_store=FILE')
        db.execute('CREATE INDEX IF NOT EXISTS records_iccid ON records(iccid)');db.commit()
        db.execute('CREATE INDEX IF NOT EXISTS records_phone ON records(phone)');db.commit()
        # Full quick_check walks the entire 7+ GiB database. Routine refreshes
        # validate the document registry and committed new records instead;
        # explicit --full-audit retains the complete database integrity check.
        # Never label the incremental check as a newly passed full audit.
        after_records=db.execute('SELECT COUNT(*) FROM records').fetchone()[0]
        full_audit=getattr(args,'full_audit',False) or not meta.get('updated_at')
        if full_audit:
            check=db.execute('PRAGMA quick_check').fetchone()[0]
            if check=='ok':
                db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)',('last_full_audit_at',now))
        else:
            checks={table:db.execute('PRAGMA quick_check('+table+')').fetchone()[0]
                for table in ('meta','contents','files','sheets','file_sheets')}
            check='ok' if all(v=='ok' for v in checks.values()) else json.dumps(checks)
            # Restrict identity/metric validation to rows added in this refresh.
            invalid=db.execute("SELECT COUNT(*) FROM records WHERE id>? AND (row_number<1 OR (iccid='' AND phone='') OR NOT json_valid(daily_json))",(before_max_id,)).fetchone()[0]
            if invalid:check='invalid newly indexed records: '+str(invalid)
        if check!='ok':
            db.commit();db.close();raise RuntimeError('任务I索引核验失败，未发布ready版本: '+check)
        version=now if changed or coverage_changed or not meta.get('updated_at') else meta['updated_at']
        db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)',('updated_at',version))
        db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)',('index_state','ready'))
        db.commit()
        receipt=dict(status='passed' if check=='ok' else 'failed',db=str(args.db),documentCount=db.execute('SELECT COUNT(*) FROM files WHERE active=1').fetchone()[0],
            indexedDocuments=db.execute("SELECT COUNT(DISTINCT f.id) FROM files f JOIN file_sheets fs ON f.id=fs.file_id WHERE f.active=1").fetchone()[0],
            recordCount=db.execute('SELECT COUNT(*) FROM records').fetchone()[0],iccidCount=db.execute("SELECT COUNT(DISTINCT iccid) FROM records WHERE iccid<>''").fetchone()[0],phoneCount=db.execute("SELECT COUNT(DISTINCT phone) FROM records WHERE phone<>''").fetchone()[0],
            statusCounts=dict(db.execute('SELECT status,COUNT(*) FROM files WHERE active=1 GROUP BY status')),
            quickCheck=check if full_audit else 'not-run-full-database',incrementalCheck=check,
            validationMode='full-database' if full_audit else 'registry-and-new-records',
            lastFullAuditAt=now if full_audit else meta.get('last_full_audit_at'),
            changedDocuments=len(changed),parsedContents=len(todo),addedRecords=after_records-before_records,
            indexedAt=version,requestedWorkers=args.workers,durationSeconds=round(time.monotonic()-started,3),discovery=discovery)
        if args.receipt:args.receipt.write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
        print(json.dumps(receipt,ensure_ascii=False),flush=True);db.close()
        return receipt

def query(db_path,keys,include_locations=True):
    """Read one consistent, immutable index snapshot; never change the DB."""
    db=sqlite3.connect(Path(db_path).resolve().as_uri()+'?mode=ro',uri=True,timeout=30)
    try:
        db.execute('PRAGMA query_only=ON');db.execute('BEGIN')
        return _query_snapshot(db,keys,include_locations)
    finally:db.close()

def _slot_sort(value):
    return tuple((0,int(p)) if p.isdigit() else (1,p.casefold()) for p in re.split(r'(\d+)',value))

def merge_slot_summaries(items):
    """Preserve every literal historical slot; don't silently break time ties."""
    history=sorted({slot for item in items for slot in item['historicalSlots']},key=_slot_sort)
    at=lambda item:item.get('latestSlotObservedAt',item.get('historicalLastSlotObservedAt'))
    observed=max((at(item) for item in items if at(item)),default=None)
    tied=[item for item in items if observed and at(item)==observed]
    latest=sorted({slot for item in tied for slot in item.get('latestSlots',item.get('historicalLastSlots',[]))},key=_slot_sort)
    sources={json.dumps(row,ensure_ascii=False,sort_keys=True):row for item in tied for row in item.get('latestSlotSources',item.get('historicalLastSlotSources',[]))}
    return dict(historicalSlots=history,latestSlot=latest[0] if len(latest)==1 else None,latestSlots=latest,
        latestSlotObservedAt=observed,latestSlotSources=[sources[k] for k in sorted(sources)],
        latestSlotStatus='confirmed-document' if len(latest)==1 else 'conflicting-same-time' if len(latest)>1 else 'time-unknown' if history else 'no-slot-record')

def historical_slot_fields(summary):
    return dict(historicalSlots=summary['historicalSlots'],historicalLastSlot=summary['latestSlot'],historicalLastSlots=summary['latestSlots'],
        historicalLastSlotObservedAt=summary['latestSlotObservedAt'],historicalLastSlotSources=summary['latestSlotSources'],historicalLastSlotStatus=summary['latestSlotStatus'])

def attach_current_positions(result,receipt_path):
    """Current position must come from fresh number-management search XLSX.

    Independently read server-authored bytes, check scope/precision and match a
    second live search. Never substitute an A/Google/history document.
    """
    from openpyxl import load_workbook
    receipt_path=Path(receipt_path);receipt=json.loads(receipt_path.read_text())
    if receipt.get('source')!='号码管理->指定ICCID->搜索->导出搜索结果' or receipt.get('status')!='downloaded-search-readback-passed' or receipt.get('independentSearchReadback') is not True or receipt.get('businessWrites')!=0:
        raise ValueError('最新位置的号码管理导出回执未通过')
    when=datetime.fromisoformat(receipt['exportedAt'].replace('Z','+00:00'))
    age=(datetime.now(timezone.utc)-when).total_seconds()
    if not 0<=age<=1200:raise ValueError('号码管理搜索结果已超过20分钟，不能冒充本次最新位置')
    keys={iccid_key(q['query']) for q in result['queries']}
    if None in keys or not keys<=set(receipt['sourceScope']):raise ValueError('最新位置导出范围不包含本次ICCID查询')
    file=Path(receipt['path'])
    with file.open('rb') as f:digest=hashlib.file_digest(f,'sha256').hexdigest()
    if digest!=receipt['sha256']:raise ValueError('号码管理搜索导出文件SHA256不一致')
    with zipfile.ZipFile(file) as z:
        if z.testzip() is not None:raise ValueError('号码管理搜索导出ZIP校验失败')
    book=load_workbook(file,read_only=True,data_only=False);rows=[]
    try:
        for sheet in book:
            cells=next(sheet.iter_rows(),None)
            if not cells:continue
            names=[header(c.value) for c in cells]
            # Verified number-management export headers differ from A tables:
            # 卡存储 / 卡号 / 卡槽标识 map to cardStore/cardNo/slotIdentif.
            ids=[i for i,n in enumerate(names) if n in IC_HEADERS];slots=[i for i,n in enumerate(names) if n in ('卡槽','卡槽标识')];phones=[i for i,n in enumerate(names) if n in PH_HEADERS or n=='卡号']
            if not ids and not slots:continue
            if len(ids)!=1 or len(slots)!=1 or len(phones)!=1:raise ValueError('最新位置导出身份/卡槽列不唯一')
            for rn,cells in enumerate(sheet.iter_rows(min_row=2),2):
                if all(c.value is None for c in cells):continue
                cell=cells[ids[0]];key=iccid_key(cell.value)
                if cell.data_type not in ('s','str','inlineStr') or not key or key not in set(receipt['sourceScope']):raise ValueError('最新位置Excel的ICCID精度或范围无效')
                pos=cells[slots[0]]
                if pos.data_type=='f' or (pos.value is not None and not isinstance(pos.value,str)):raise ValueError('最新位置Excel的卡槽不是可靠文本')
                phone=phone_key(cells[phones[0]].value)
                if not phone and cells[phones[0]].value not in (None,''):raise ValueError('最新位置Excel的非空手机号无效')
                rows.append(dict(iccid=key,phone=phone,slot=(pos.value or '').strip(),sheet=sheet.title,row=rn,iccidCell=cell.coordinate,slotCell=pos.coordinate))
    finally:book.close()
    live=json.loads((receipt_path.parent/'search_after.json').read_text())
    actual=Counter((r['iccid'],r['phone'],r['slot']) for r in rows)
    expected=Counter((iccid_key(r['cardStore']),phone_key(r.get('cardNo')),str(r.get('slotIdentif') or '').strip()) for r in live['rows'])
    if len(rows)!=receipt['searchRows'] or actual!=expected:raise ValueError('最新位置Excel与独立搜索回读不一致')
    def update(item,matched):
        item.update(currentSlots=sorted({r['slot'] for r in matched if r['slot']},key=_slot_sort),currentSlotMatchedRows=len(matched),
            currentSlotStatus='confirmed-search-export' if any(r['slot'] for r in matched) else 'slot-missing' if matched else 'search-not-found',
            currentSlotSources=matched,currentSlotObservedAt=receipt['exportedAt'])
    for q in result['queries']:update(q,[r for r in rows if r['iccid']==iccid_key(q['query'])])
    for phone,q in result['phoneSlots'].items():update(q,[r for r in rows if r['iccid'] in keys and r['phone']==phone])
    result['currentPositionSource']=dict(status='passed',workflow=receipt['source'],path=str(file),sha256=digest,exportedAt=receipt['exportedAt'],
        searchedIccidCount=receipt['requestedIccids'],exportRows=len(rows),matchedIccidCount=len({r['iccid'] for r in rows}),independentXlsxReadback=True,businessWrites=0)
    return result

def slot_summary(evidence):
    # Position is independent of SMS-day metrics. Include inventory records,
    # but never infer a slot from another card or an ambiguous identity row.
    rows=[e for e in evidence if not e['issue'] and str(e.get('slot') or '').strip() not in ('','未知','-','—','None','null','N/A')]
    history=sorted({str(e['slot']).strip() for e in rows},key=_slot_sort)
    timed=[e for e in rows if e['observed'] and e['date_basis']!='修改时间，非确定采集时间']
    observed=max((e['observed'] for e in timed),default=None)
    last=[e for e in timed if e['observed']==observed]
    latest=sorted({str(e['slot']).strip() for e in last},key=_slot_sort)
    sources=[{k:e[k] for k in ('path','sheet_name','row_number','iccid_cell','phone_cell','iccid','phone','slot','observed','date_basis')} for e in last]
    return merge_slot_summaries([dict(historicalSlots=history,latestSlots=latest,latestSlotObservedAt=observed,latestSlotSources=sources)])

def _query_snapshot(db,keys,include_locations):
    meta=dict(db.execute('SELECT key,value FROM meta'))
    if meta.get('index_state','ready')!='ready':raise RuntimeError('任务I索引正在刷新，暂不输出不完整统计')
    if not meta.get('updated_at'):raise RuntimeError('任务I首轮索引尚未完成，暂不输出部分统计')
    output=[];observations={};sources={};phone_slots=defaultdict(list);t=time.monotonic()
    for supplied in keys:
        ic=iccid_key(supplied);ph=phone_key(supplied)
        key=ic or ph
        if not key:raise ValueError('非法查询标识: '+supplied)
        field='iccid' if ic else 'phone'
        sql=f'''SELECT r.id AS record_id,r.*,f.id,f.path,f.observed,f.date_basis,f.sha256,fs.sheet_name FROM records r
          JOIN file_sheets fs ON r.sheet_id=fs.sheet_id JOIN files f ON fs.file_id=f.id
          WHERE r.{field}=? AND f.active=1 AND f.status='已检查' ORDER BY f.observed,f.path,r.row_number'''
        columns=[d[0] for d in db.execute(sql,(key,)).description]
        evidence=[dict(zip(columns,row)) for row in db.execute(sql,(key,))]
        parsed={}
        for e in evidence:
            sources[e['path']]=dict(path=e['path'],observed=e['observed'],sha256=e['sha256'])
            if e['record_id'] not in parsed:parsed[e['record_id']]=json.loads(e['daily_json'])
            for day,values in parsed[e['record_id']].items():
                group=(supplied,day,e['phone'] or '@'+e['iccid'])
                if group not in observations:observations[group]=dict(count=0,best=None,rank=None,tied=set(),versions=set())
                state=observations[group];state['count']+=1
                if e['issue'] or None in values or e['date_basis']=='修改时间，非确定采集时间':continue
                value=tuple(values);closed=day<e['latest_day']
                direct='新鲜导出' in e['path'] or re.search(r'export4day_\d',Path(e['path']).name) is not None
                rank=(closed,e['observed'],direct)
                if closed:state['versions'].add(value)
                if state['rank'] is None or rank>state['rank']:
                    state.update(rank=rank,tied={value},best={k:e[k] for k in ('iccid','phone','path','sheet_name','row_number','iccid_cell','observed','latest_day','slot','carrier')})
                    state['best']['values']=values
                elif rank==state['rank']:state['tied'].add(value)
        item=dict(query=supplied,normalized=key,queryType=field,documentCount=len({e['path'] for e in evidence}),contentCount=len({e['sha256'] for e in evidence}),
            historicalPhones=sorted({e['phone_raw'] for e in evidence if e['phone_raw']}),locationCount=len(evidence))
        item.update(historical_slot_fields(slot_summary(evidence)))
        by_phone=defaultdict(list)
        for e in evidence:
            if e['phone']:by_phone[e['phone']].append(e)
        for phone,rows in by_phone.items():phone_slots[phone].append(slot_summary(rows))
        if include_locations:item['locations']=[{k:e[k] for k in ('path','sheet_name','row_number','iccid_cell','phone_cell','iccid','phone_raw','slot','carrier','observed','issue')} for e in evidence]
        output.append(item)
    selected=[];conflicts=[];unknown=[]
    for (key,day,phone),state in sorted(observations.items()):
        best=state['best']
        if best is None:
            unknown.append(dict(query=key,date=day,phone=phone,reason='身份歧义、指标缺失或来源时间不明'));continue
        # Closed platform days outrank same-day partial snapshots. Prefer a
        # direct export over derivative reports at the same capture timestamp.
        if len(state['tied'])>1:
            unknown.append(dict(query=key,date=day,phone=phone,reason='同最高优先级来源指标冲突'));continue
        versions=state['versions']
        if len(versions)>1:conflicts.append(dict(query=key,date=day,phone=phone,versions=sorted(versions),chosen=best['values'],path=best['path']))
        selected.append(dict(query=key,date=day,phone=phone,iccid=best['iccid'],values=best['values'],
            partial=not day<best['latest_day'],path=best['path'],sheet=best['sheet_name'],row=best['row_number'],iccidCell=best['iccid_cell'],observed=best['observed'],duplicateSnapshots=state['count'],slot=best['slot'],carrier=best['carrier']))
    for o in output:
        days=[d for d in selected if d['query']==o['query']]
        totals=[sum(d['values'][i] for d in days) for i in range(4)] if days else [None]*4
        o.update(dailyRecordCount=len(days),dates=sorted({d['date'] for d in days}),totals=totals,
                 replyRate=totals[1]/totals[0] if totals[0] else None,
                 messageReplyRate=totals[3]/totals[2] if totals[2] else None,
                 unknownDailyRecords=sum(u['query']==o['query'] for u in unknown),
                 partialDates=sorted({d['date'] for d in days if d['partial']}))
    result=dict(status='passed',inputCount=len(keys),uniqueInputCount=len(set(keys)),queries=output,daily=selected,conflicts=conflicts,unknown=unknown,sources=list(sources.values()),
        totals=[sum(d['values'][i] for d in selected) for i in range(4)],dateStart=min((d['date'] for d in selected),default=None),dateEnd=max((d['date'] for d in selected),default=None),
        querySeconds=round(time.monotonic()-t,3),indexedAt=meta['updated_at'],allLocationsStoredInTaskIIndex=True,
        phoneSlots={phone:historical_slot_fields(merge_slot_summaries(items)) for phone,items in sorted(phone_slots.items())})
    # A mixed ICCID/phone query may select the same traffic twice. Disclose this
    # and omit portfolio totals instead of silently double-counting.
    identities=[(d['date'],d['iccid'],d['phone']) for d in selected]
    if len(identities)!=len(set(identities)):result.update(totals=None,portfolioWarning='查询列表存在ICCID与手机号交叉，组合总计停用')
    return result

def _query_worker(args):
    if WORKER_PRIORITY is None:configure_worker_priority()
    result=query(*args)
    result['worker']=dict(pid=os.getpid(),priority=WORKER_PRIORITY)
    return result

def parallel_query(db_path,keys,workers=DEFAULT_QUERY_WORKERS,include_locations=False):
    """Bounded CPU process pool with separate read-only SQLite connections.

    Each card stays in one worker so daily deduplication cannot cross workers.
    Refuse to publish if workers saw different index versions. Final merging and
    file writing happen only in the coordinator, never in parallel writers.
    """
    if not 1<=workers<=8:raise ValueError('并发数必须1至8')
    t=time.monotonic();supplied=list(keys);unique=list(dict.fromkeys(supplied))
    if not unique:raise ValueError('请提供ICCID、手机号或清单')
    for key in unique:
        if not (iccid_key(key) or phone_key(key)):raise ValueError('非法查询标识: '+key)
    actual=min(workers,len(unique))
    if actual==1:parts=[_query_worker((db_path,unique,include_locations))]
    else:
        # Round-robin spreads high-match cards across workers. Only `actual`
        # jobs are submitted, and compact results avoid copying huge locations.
        jobs=[(Path(db_path),unique[n::actual],include_locations) for n in range(actual)]
        with ProcessPoolExecutor(max_workers=actual,initializer=configure_worker_priority) as pool:parts=list(pool.map(_query_worker,jobs))
    versions={part['indexedAt'] for part in parts}
    if len(versions)!=1:raise RuntimeError('查询期间索引版本变化，未交付混合版本统计；请重新查询')
    by_query={q['query']:q for part in parts for q in part['queries']}
    daily=sorted((row for part in parts for row in part['daily']),key=lambda x:(x['query'],x['date'],x['phone']))
    conflicts=sorted((row for part in parts for row in part['conflicts']),key=lambda x:(x['query'],x['date'],x['phone']))
    unknown=sorted((row for part in parts for row in part['unknown']),key=lambda x:(x['query'],x['date'],x['phone']))
    sources={row['path']:row for part in parts for row in part['sources']}
    phone_slots=defaultdict(list)
    for part in parts:
        for phone,item in part['phoneSlots'].items():phone_slots[phone].append(item)
    result=dict(status='passed',inputCount=len(supplied),uniqueInputCount=len(unique),queries=[by_query[k] for k in unique],daily=daily,conflicts=conflicts,unknown=unknown,
        sources=[sources[p] for p in sorted(sources)],totals=[sum(d['values'][i] for d in daily) for i in range(4)],
        dateStart=min((d['date'] for d in daily),default=None),dateEnd=max((d['date'] for d in daily),default=None),
        querySeconds=round(time.monotonic()-t,3),indexedAt=versions.pop(),allLocationsStoredInTaskIIndex=True,
        phoneSlots={phone:historical_slot_fields(merge_slot_summaries(items)) for phone,items in sorted(phone_slots.items())},
        execution=dict(cpuCores=CPU_CORES,requestedWorkers=workers,actualWorkers=actual,configuredCoreShare=actual/CPU_CORES,mode='cpu-process-pool-read-only' if actual>1 else 'single-read-only',workerPriorities=[part['worker'] for part in parts]))
    identities=[(d['date'],d['iccid'],d['phone']) for d in daily]
    if len(identities)!=len(set(identities)):result.update(totals=None,portfolioWarning='查询列表存在ICCID与手机号交叉，组合总计停用')
    return result

def main():
    p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='command',required=True)
    first=sub.add_parser('first-seen',help='离线扫描最近一个月所有表格文本单元格，追溯ICCID最早出现文档')
    first.add_argument('--list',type=Path,required=True);first.add_argument('--db',type=Path,default=DEFAULT_DB)
    first.add_argument('--root',action='append');first.add_argument('--workers',type=int,default=7)
    first.add_argument('--since',help='北京时间起始自然日期，例如2026-09-16；默认前一个月')
    first.add_argument('--output-dir',type=Path)
    first_report=sub.add_parser('first-seen-report',help='仅从已核验扫描结果生成首次出现工作簿，不重新扫描')
    first_report.add_argument('--result',type=Path,required=True)
    a=sub.add_parser('index');a.add_argument('--db',type=Path,default=DEFAULT_DB);a.add_argument('--root',action='append');a.add_argument('--workers',type=int,default=DEFAULT_WORKERS);a.add_argument('--omit');a.add_argument('--scan-evidence',type=Path);a.add_argument('--manifest',type=Path);a.add_argument('--receipt',type=Path);a.add_argument('--full-audit',action='store_true',help='另行执行全库完整性审计；日常使用增量检查')
    run=sub.add_parser('run',help='一键历史统计及独立核验；默认离线，明确指定时查询最新位置')
    run.add_argument('identifiers',nargs='*');run.add_argument('--list',type=Path,action='append',default=[])
    run.add_argument('--db',type=Path,default=DEFAULT_DB);run.add_argument('--root',action='append')
    run.add_argument('--workers',type=int,default=7);run.add_argument('--output-dir',type=Path)
    run.add_argument('--full-audit',action='store_true');run.add_argument('--locations',action='store_true')
    position=run.add_mutually_exclusive_group()
    position.add_argument('--current-search-receipt',type=Path,help='仅独立核验已有20分钟内的当前授权搜索导出；不会联网')
    position.add_argument('--current-positions',action='store_true',help='本次明确授权：串行查询号码管理ICCID搜索导出，并核验后生成最新位置')
    run.add_argument('--authorization-source',choices=['current-user'],help='联网最新位置查询必须注明本次用户授权，不沿用历史授权')
    q=sub.add_parser('query');q.add_argument('--db',type=Path,default=DEFAULT_DB);q.add_argument('identifiers',nargs='*');q.add_argument('--list',type=Path);q.add_argument('--output',type=Path);q.add_argument('--json',action='store_true');q.add_argument('--workers',type=int,default=DEFAULT_QUERY_WORKERS);q.add_argument('--current-search-receipt',type=Path);q.add_argument('--locations',action='store_true',help='同时导出所有文档、工作表和单元格位置；默认保留在索引中')
    current=sub.add_parser('current-export',help='当前明确授权的号码管理ICCID搜索导出，只读且串行')
    current.add_argument('--queries-json',type=Path,required=True);current.add_argument('--output-dir',type=Path,required=True);current.add_argument('--authorization-source',choices=['current-user'],required=True)
    args=p.parse_args()
    if args.command=='first-seen':
        from opensms_task_i_first_seen import run_first_seen
        run_first_seen(args)
    elif args.command=='first-seen-report':
        from opensms_task_i_first_seen import build_report
        build_report(args.result)
    elif args.command=='index':
        if not 1<=args.workers<=8:p.error('并发数必须1至8')
        index(args)
    elif args.command=='run':
        from opensms_task_i_pipeline import run_pipeline
        run_pipeline(args)
    elif args.command=='current-export':
        node='/Users/zhangsan/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
        script=Path(__file__).with_name('opensms_task_i_current_slots.mjs')
        subprocess.run([node,str(script),'export',str(args.output_dir.resolve()),str(args.queries_json.resolve()),'--current-user-authorized-task-i-search-export'],check=True,cwd='/Users/zhangsan/Downloads/430')
    else:
        keys=args.identifiers+(args.list.read_text().splitlines() if args.list else [])
        if not keys:p.error('请提供ICCID、手机号或清单')
        if not 1<=args.workers<=8:p.error('并发数必须1至8')
        result=parallel_query(args.db,keys,args.workers,args.locations)
        if args.current_search_receipt:attach_current_positions(result,args.current_search_receipt)
        if args.output:args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2))
        print(json.dumps({k:v for k,v in result.items() if k not in ('daily','queries','conflicts','unknown','sources','phoneSlots')} if args.output else result,ensure_ascii=False))

if __name__=='__main__':main()
