"""Read-only extraction of card identities from a hash-bound historical workbook."""
import hashlib
import json
import sys

from openpyxl import load_workbook


def extract(file, expected_hash, requested):
    with open(file, 'rb') as stream:
        if hashlib.file_digest(stream, 'sha256').hexdigest() != expected_hash:
            raise ValueError('SWAP_HISTORY_WORKBOOK_CHANGED')
    book = load_workbook(file, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = list(next(iterator))
        required = ['电话', '运营商', '卡槽', 'ICCID', 'opr']
        if any(headers.count(name) != 1 for name in required):
            raise ValueError('SWAP_HISTORY_WORKBOOK_HEADERS_INVALID')
        columns = {name: headers.index(name) for name in required}
        output = []
        for number, values in enumerate(iterator, 2):
            phone = str(values[columns['电话']] or '').strip()
            if phone.isascii() and phone.isdigit() and (len(phone) == 10 or (len(phone) == 11 and phone[0] == '1')):
                phone = phone[-10:]
            else:
                continue
            if phone not in requested:
                continue
            output.append({'phone': phone, 'sourceRow': number,
                           'slot': values[columns['卡槽']], 'iccid': values[columns['ICCID']],
                           'carrier': values[columns['运营商']], 'opr': values[columns['opr']]})
        return output
    finally:
        book.close()


if __name__ == '__main__':
    print(json.dumps(extract(sys.argv[1], sys.argv[2], set(json.loads(sys.argv[3]))), ensure_ascii=False))
