#!/usr/bin/env python3
"""Freeze Chatr dates from one hash-bound Google XLSX; no browser or live writes."""
from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from opensms_export_task_f_deadlines import publish_no_replace

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "services/opensms_document_api/pipeline"))
from google_room_sheets import LEGACY_ROOM_SHEETS, select_room_sheets


REQUIRED_SHEETS = LEGACY_ROOM_SHEETS  # Legacy fixture/public constant, not live selection.
REQUIRED_HEADERS = ("卡槽", "電信商/價格", "激活日期", "預計過期日")
OPTIONAL_HEADERS = {
    "iccid": ("iccid", "ICCID"),
    "phone": ("手機號", "手机号", "電話", "电话", "号码"),
}


def normalize_slot(value) -> str:
    """Keep the document pipeline's exact slot alias normalization."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u3000", "").replace(" ", "").replace("-", "_")
    normalized = []
    for part in text.split("_"):
        numeric = re.fullmatch(r"(\d+)(?:\.\d+)?", part)
        normalized.append(str(int(numeric.group(1))) if numeric else part.lower())
    return "_".join(normalized)


def normalized_iccid(value) -> str | None:
    # Numeric Excel ICCIDs may already have lost precision. Never reconstruct them.
    if not isinstance(value, str):
        return None
    text = re.sub(r"\s+", "", value)
    text = re.sub(r"[Ff]$", "", text)
    return text if re.fullmatch(r"[0-9]{18,22}", text) else None


def is_chatr(value) -> bool:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    return "chatr" in text and not re.search(r"kooda|koodo|telus1", text)


def json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def parse_date(value, epoch) -> str | None:
    """Only unambiguous calendar dates; preserve local time without inference."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    parsed = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(value):
            raise ValueError("日期不是有限数值")
        if value == int(value) and re.fullmatch(r"(?:19|20)[0-9]{6}", str(int(value))):
            parsed = datetime.strptime(str(int(value)), "%Y%m%d")
        elif 20000 <= value <= 60000:
            parsed = from_excel(value, epoch)
    elif isinstance(value, str):
        text = unicodedata.normalize("NFKC", value.strip()).strip("'\"")
        text = text.replace("年", "-").replace("月", "-").replace("日", "")
        # Year-first dates cannot silently reinterpret ambiguous 09/10/2026.
        match = re.fullmatch(
            r"([0-9]{4})[-/.]([0-9]{1,2})[-/.]([0-9]{1,2})"
            r"((?:[T ][0-9]{2}:[0-9]{2}(?::[0-9]{2}(?:\.[0-9]{1,6})?)?)?)",
            text,
        )
        if match:
            year, month, day, suffix = match.groups()
            parsed = datetime.fromisoformat(f"{int(year):04d}-{int(month):02d}-{int(day):02d}{suffix}")
        elif re.fullmatch(r"[0-9]{8}", text):
            parsed = datetime.strptime(text, "%Y%m%d")
        elif re.fullmatch(
            r"[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}"
            r"(?::[0-9]{2}(?:\.[0-9]{1,6})?)?(?:Z|[+-][0-9]{2}:[0-9]{2})", text,
        ):
            # Keep source wall-clock semantics (the report uses naive server time).
            parsed = datetime.fromisoformat(text).replace(tzinfo=None)
    if not isinstance(parsed, datetime):
        raise ValueError(f"非空日期无效或有歧义: {value!r}")
    parsed = parsed.replace(tzinfo=None)
    return parsed.date().isoformat() if parsed.time() == datetime.min.time() else parsed.isoformat()


def sheet_columns(headers: list, sheet: str) -> tuple[dict, dict]:
    names = [str(value).strip() if value is not None else "" for value in headers]
    columns = {}
    for name in REQUIRED_HEADERS:
        hits = [i + 1 for i, value in enumerate(names) if value == name]
        if len(hits) != 1:
            raise ValueError(f"Google页签 {sheet} 必要精确字段缺失或重复: {name}")
        columns[name] = hits[0]
    optional = {}
    for key, aliases in OPTIONAL_HEADERS.items():
        hits = [(name, index + 1) for index, name in enumerate(names) if name in aliases]
        if len(hits) > 1:
            raise ValueError(f"Google页签 {sheet} 可选身份字段重复: {key}")
        optional[key] = hits[0][1] if hits else None
        columns.update(hits)
    return columns, optional


def parse_source(raw: bytes) -> dict:
    """Parse the same in-memory bytes whose SHA is checked by freeze()."""
    values = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    formulas = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    try:
        source_sheets = select_room_sheets(values)
        by_slot, source_columns, excluded = {}, {}, []
        counts = {key: 0 for key in (
            "totalSourceRows", "ignoredNonChatrRows", "chatrSourceRows",
            "excludedInvalidSlotRows", "mergedDuplicateRows", "numericIccidRows",
        )}
        for sheet_name in source_sheets:
            ws, fws = values[sheet_name], formulas[sheet_name]
            columns, optional = sheet_columns([cell.value for cell in next(ws.iter_rows())], sheet_name)
            source_columns[sheet_name] = columns
            for row_number, (row, formula_row) in enumerate(
                zip(ws.iter_rows(min_row=2), fws.iter_rows(min_row=2)), start=2
            ):
                if not any(cell.value is not None and str(cell.value).strip() for cell in row):
                    continue
                counts["totalSourceRows"] += 1

                def cell_value(column):
                    return row[column - 1].value if column else None

                carrier = cell_value(columns["電信商/價格"])
                if not is_chatr(carrier):
                    counts["ignoredNonChatrRows"] += 1
                    continue
                counts["chatrSourceRows"] += 1
                raw_slot = cell_value(columns["卡槽"])
                slot = normalize_slot(raw_slot)
                provenance = {"sourceSheet": sheet_name, "sourceRow": row_number,
                              "rawSlot": json_value(raw_slot), "normalizedSlot": slot,
                              "carrier": json_value(carrier)}
                if not re.fullmatch(r"[0-9]+_[0-9]+", slot):
                    counts["excludedInvalidSlotRows"] += 1
                    excluded.append(provenance | {"reason": "invalid_or_missing_slot"})
                    continue
                for source_field, target_field in (("激活日期", "activationDate"), ("預計過期日", "cutoffDate")):
                    column = columns[source_field]
                    cell, formula_cell = row[column - 1], formula_row[column - 1]
                    if formula_cell.data_type == "f" and cell.value is None and cell.data_type not in {"str", "inlineStr"}:
                        raise ValueError(f"Google {sheet_name}!{formula_cell.coordinate} 日期公式没有缓存值")
                    try:
                        provenance[target_field] = parse_date(cell.value, values.epoch)
                    except (ValueError, OverflowError) as exc:
                        raise ValueError(f"Google {sheet_name}!{cell.coordinate} {source_field}: {exc}") from exc
                    provenance["raw" + target_field[0].upper() + target_field[1:]] = json_value(cell.value)
                iccid, phone = cell_value(optional["iccid"]), cell_value(optional["phone"])
                provenance.update({"iccid": json_value(iccid), "queryIccid": normalized_iccid(iccid),
                                   "phone": json_value(phone)})
                if isinstance(iccid, (int, float)) and not isinstance(iccid, bool):
                    counts["numericIccidRows"] += 1
                existing = by_slot.get(slot)
                if existing is None:
                    by_slot[slot] = provenance | {"provenance": [provenance]}
                    continue
                for field in ("activationDate", "cutoffDate", "queryIccid"):
                    old, new = existing[field], provenance[field]
                    if old is not None and new is not None and old != new:
                        raise ValueError(f"Google归一化卡槽 {slot} 的 {field} 冲突")
                    if old is None and new is not None:
                        existing[field] = new
                existing["provenance"].append(provenance)
                counts["mergedDuplicateRows"] += 1
        records = sorted(by_slot.values(), key=lambda row: tuple(map(int, row["normalizedSlot"].split("_"))))
        return {"sourceSheets": list(source_sheets), "sourceColumns": source_columns,
                **counts, "recordCount": len(records), "excludedRows": excluded,
                "missingActivationRows": sum(row["activationDate"] is None for row in records),
                "missingCutoffRows": sum(row["cutoffDate"] is None for row in records),
                "missingBothDateRows": sum(row["activationDate"] is None and row["cutoffDate"] is None for row in records),
                "records": records}
    finally:
        values.close()
        formulas.close()


def freeze(source: Path, source_sha256: str, output: Path, run_id: str) -> dict:
    if not run_id or not run_id.strip():
        raise ValueError("Chatr日期快照必须绑定本轮runId")
    expected = source_sha256.lower()
    if not re.fullmatch(r"[a-f0-9]{64}", expected):
        raise ValueError("Chatr日期快照缺少有效源SHA256")
    if os.path.lexists(output):
        raise FileExistsError(f"Chatr日期快照不可覆盖: {output}")
    source = source.resolve()
    raw = source.read_bytes()
    if hashlib.sha256(raw).hexdigest() != expected:
        raise ValueError("Chatr Google源SHA256不一致")
    payload = {"schemaVersion": 1, "status": "passed", "runId": run_id,
               "createdAt": datetime.now().astimezone().isoformat(),
               "source": "google_chatr_dates", "sourcePath": str(source), "sourceSha256": expected,
               "readOnly": True, "snapshotConsistent": True, "sourceHashVerified": True,
               **parse_source(raw)}
    # Detect source replacement during extraction; downstream also verifies its hash.
    if hashlib.sha256(source.read_bytes()).hexdigest() != expected:
        raise ValueError("Chatr Google源在快照提取期间发生变化")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=output.parent,
                                     prefix=".chatr_google_dates_", delete=False) as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
        handle.flush()
        os.fsync(handle.fileno())
        temporary = Path(handle.name)
    try:
        digest = hashlib.sha256(temporary.read_bytes()).hexdigest()
        publish_no_replace(temporary, output)
        if hashlib.sha256(output.read_bytes()).hexdigest() != digest or json.loads(output.read_text()) != payload:
            raise ValueError("Chatr日期快照发布后独立回读不一致")
    finally:
        temporary.unlink(missing_ok=True)
    return {key: value for key, value in payload.items() if key not in {"records", "excludedRows"}} | {
        "output": str(output.resolve()), "sha256": digest, "verificationPassed": True}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--source-sha256", required=True)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = freeze(args.source, args.source_sha256, args.output, args.run_id)
    except Exception as exc:
        print(json.dumps({"status": "failed", "runId": args.run_id, "error": str(exc)}, ensure_ascii=False))
        return 2
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
