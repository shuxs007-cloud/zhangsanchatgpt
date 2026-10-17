"""Read only: extract four-day metrics from one exact verified final workbook."""
import hashlib
import json
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from opensms_generate_task_c_candidates import parse_number, parse_rate, phone_text


def proof(file):
    p = Path(file).resolve()
    with p.open('rb') as stream:
        return {'path': str(p), 'sha256': hashlib.file_digest(stream, 'sha256').hexdigest()}


def extract(state_path):
    state_proof = proof(state_path)
    state = json.loads(Path(state_path).read_text())
    refresh = state.get('stages', {}).get('refresh', {})
    acquire = state.get('stages', {}).get('acquire', {})
    if state.get('status') != 'success' or not acquire.get('serverDateVerified'):
        raise ValueError('SWAP_METRICS_SOURCE_NOT_VERIFIED')
    workbook = Path(refresh['finalWorkbook'])
    workbook_proof = proof(workbook)
    verification_path = refresh['verifyReceipt']
    verification = json.loads(Path(verification_path).read_text())
    if (workbook_proof['sha256'] != refresh.get('finalWorkbookSha256')
            or verification.get('passed') is not True
            or verification.get('runId') != state['runId']
            or str(workbook.resolve()) != str(Path(verification['finalWorkbook']).resolve())):
        raise ValueError('SWAP_METRICS_WORKBOOK_VERIFICATION_MISMATCH')
    as_of = refresh['completedAt']
    age = (datetime.now(timezone.utc) - datetime.fromisoformat(as_of)).total_seconds()
    if not 0 <= age <= 90 * 60:
        raise ValueError('SWAP_METRICS_SOURCE_TOO_OLD')
    book = load_workbook(workbook, read_only=True, data_only=True)
    rows = []
    try:
        sheet = book.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = list(next(iterator))
        required = ['电话', 'ICCID', '卡槽', '运营商', '4日回复率', '4日发送窗口数']
        if any(headers.count(name) != 1 for name in required):
            raise ValueError('SWAP_METRICS_HEADERS_INVALID')
        cols = {name: headers.index(name) for name in required}
        for source_row, values in enumerate(iterator, 2):
            p = phone_text(values[cols['电话']])
            if not p.isascii() or not p.isdigit() or not (len(p) == 10 or (len(p) == 11 and p[0] == '1')):
                continue
            raw_rate, raw_send = values[cols['4日回复率']], values[cols['4日发送窗口数']]
            rate = None if isinstance(raw_rate, bool) else parse_rate(raw_rate)
            send = None if isinstance(raw_send, bool) else parse_number(raw_send)
            if rate is None or not math.isfinite(rate) or not 0 <= rate <= 1:
                rate = None
            if send is None or not math.isfinite(send) or send < 0 or not send.is_integer():
                send = None
            rows.append({'phone': p[-10:], 'iccid': values[cols['ICCID']],
                         'slot': values[cols['卡槽']], 'carrier': values[cols['运营商']],
                         'replyRate': rate, 'sendWindows': int(send) if send is not None else None,
                         'sourceRow': source_row})
    finally:
        book.close()
    if proof(workbook) != workbook_proof or proof(state_path) != state_proof:
        raise ValueError('SWAP_METRICS_SOURCE_CHANGED')
    return {'schemaVersion': 1, 'runId': state['runId'], 'serverDate': acquire['serverDate'],
            'asOf': as_of, 'sheet': sheet.title, 'metric': '4日回复率',
            'tieBreak': '4日发送窗口数降序，然后手机号升序',
            'sourceProofs': [state_proof, workbook_proof, proof(verification_path)],
            'workbook': workbook_proof, 'rows': rows}


if __name__ == '__main__':
    print(json.dumps(extract(sys.argv[1]), ensure_ascii=False, allow_nan=False))
