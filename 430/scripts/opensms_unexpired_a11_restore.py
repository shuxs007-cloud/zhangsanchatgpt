"""Explicit one-run correction; never part of default carrier routing."""
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
import json
import re

from openpyxl import load_workbook
from opensms_generate_task_b_candidates import parse_date, phone_text, sha256

ROUTE = "restore_unexpired_to_a11"
SOURCES = ("快过期隔离区C", "快过期隔离区T", "过期下架卡C", "过期下架卡T")


def select_rows(workbook, server_date):
    reference = date.fromisoformat(server_date)
    book = load_workbook(workbook, read_only=True, data_only=False)
    selected, excluded, seen = [], [], set()
    try:
        iterator = book["sheet1"].iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(iterator)]
        for name in ("电话", "运营商", "截止日", "卡槽", "ICCID"):
            if headers.count(name) != 1:
                raise ValueError(f"恢复字段缺失或重复: {name}")
        for index, values in enumerate(iterator, 2):
            raw = dict(zip(headers, values))
            carrier = str(raw["运营商"] or "").strip()
            if carrier not in SOURCES:
                continue
            phone = phone_text(raw["电话"])
            cutoff = parse_date(raw["截止日"], book.epoch)
            identity = phone[-10:]
            if not re.fullmatch(r"1?\d{10}", phone) or identity in seen:
                raise ValueError(f"恢复来源号码无效或重复: 主表第{index}行")
            seen.add(identity)
            row = {"电话": phone, "原运营商": carrier, "源行号": index,
                   "卡槽": str(raw["卡槽"] or "").strip(),
                   "ICCID": phone_text(raw["ICCID"]),
                   "截止日": cutoff.isoformat() if cutoff else "",
                   "距截止日天数": (cutoff - reference).days if cutoff else None}
            if cutoff is None or cutoff <= reference:
                excluded.append({**row, "原因": "截止日无效或缺失" if cutoff is None else "已到截止日，不恢复"})
            elif not row["卡槽"] or not re.fullmatch(r"\d{18,22}[Ff]?", row["ICCID"]):
                excluded.append({**row, "原因": "缺少可验证的卡槽或ICCID，不写入"})
            else:
                selected.append(row)
    finally:
        book.close()
    return selected, excluded


def build_candidate(document, acquire, run_id, target, date_receipts, output):
    if not re.match(r"^A11(?:-|\s|$)", target, re.I):
        raise ValueError("恢复目标必须是运行时唯一启用的A11")
    workbook = Path(document["cardResult"]).resolve()
    source = Path(acquire["sourceArchive"]).resolve()
    receipt = Path(document["cardReceipt"]).resolve()
    proof = json.loads(receipt.read_text())
    review = (proof.get("verification") or {}).get("date_resolution_review") or {}
    if (proof.get("result") != "成功" or proof.get("failed_checks")
            or review.get("matches") is not True
            or (review.get("source_export") or {}).get("sha256") != acquire["sourceSha256"]
            or acquire.get("serverDateVerified") is not True
            or run_id not in workbook.name or run_id not in str(source)
            or sha256(source) != acquire["sourceSha256"]):
        raise ValueError("本轮新鲜源表或日期独立验证证明不完整")
    before = sha256(workbook)
    rows, excluded = select_rows(workbook, acquire["serverDate"])
    if sha256(workbook) != before:
        raise ValueError("读取期间工作簿发生变化，停止恢复")
    result = {"task": "RESTORE_UNEXPIRED_A11", "route": ROUTE,
              "forceGate": True, "authorizationSource": "current-user", "runId": run_id,
              "createdAt": datetime.now(timezone.utc).isoformat(), "serverDate": acquire["serverDate"],
              "targetCarrierName": target, "sourceCarrierNames": list(SOURCES),
              "sourceWorkbook": str(workbook), "sourceWorkbookSha256": before,
              "sourceArchive": str(source), "sourceSha256": acquire["sourceSha256"],
              "documentReceipt": str(receipt), "documentReceiptSha256": sha256(receipt),
              "serverDateReceipts": [{"path": str(p), "sha256": sha256(Path(p))} for p in date_receipts],
              "count": len(rows), "rows": rows, "candidates": rows,
              "bySource": dict(Counter(r["原运营商"] for r in rows)),
              "excludedCount": len(excluded), "excluded": excluded}
    with Path(output).open("x", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
    return result


def report_lines(config):
    summary = config.get("unexpiredRestore")
    if not summary:
        return []
    candidate = json.loads(Path(summary["candidate"]).read_text())
    verified = {}
    if summary.get("validated"):
        proof = json.loads(Path(summary["validated"]).read_text())
        verified = {r["电话"]: r for r in proof["rows"]}
    lines = ["### 本轮明确授权：未到截止日恢复A11", "",
             "- 仅四个隔离/过期区；截止日晚于服务器当天。其它区域不动，普通流转及B/C/D/E未执行；不产生下卡通知。",
             f"- 目标：{summary['target']}；候选 {summary['candidateCount']}；实际写入 {summary['actualWrite']}；独立回读 {summary['result']}。",
             "", "| 电话 | 原运营商 | 截止日 | 剩余天数 | 恢复结果 |", "| --- | --- | --- | ---: | --- |"]
    for row in candidate["rows"]:
        status = "已写入并独立回读" if row["电话"] in verified else "未写入，详见回执"
        lines.append(f"| {row['电话']} | {row['原运营商']} | {row['截止日']} | {row['距截止日天数']} | {status} |")
    for row in candidate["excluded"]:
        lines.append(f"| {row['电话']} | {row['原运营商']} | {row['截止日']} | {row['距截止日天数']} | {row['原因']} |")
    return lines
